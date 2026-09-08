import csv
import hashlib
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_006_r4"
QA_RUN_ID = "20260907_161639"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_006_r3" / "SEO_Product_Optimization_qa_batch_006_r3.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
CUSTOMIZER_FIELDS_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_fields_audit.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_006_r4.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_006_r4.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    51: {
        "title": "Personalized Christian Scripture Floral Blanket",
        "meta_title": "Personalized Christian Scripture Floral Blanket",
        "meta_description": "Customize a Christian scripture floral blanket with required name, pink EMILY artwork, blanket sizes and verified text fields.",
        "primary": "personalized Christian scripture floral blanket",
        "secondary": "custom Bible verse floral blanket, Christian blanket with name, personalized scripture blanket",
        "cluster": "personalized Christian scripture floral blanket",
        "intent_role": "Design D10 with verified required Custom Name field up to 1000 characters.",
        "detail": "pink EMILY Christian scripture floral blanket artwork",
    },
    52: {
        "title": "Personalized Emily God Says I Am Blanket",
        "meta_title": "Personalized Emily God Says I Am Blanket",
        "meta_description": "Customize an Emily God Says I Am blanket with required name, required color choice, blanket sizes and verified text fields.",
        "primary": "personalized Emily God Says I Am blanket",
        "secondary": "God Says I Am blanket with name, custom Christian affirmation blanket, Emily scripture blanket",
        "cluster": "personalized Emily God Says I Am blanket",
        "intent_role": "Design D9 with verified required Custom Name up to 1000 characters and required Choose Color group.",
        "detail": "blue EMILY God Says I Am scripture blanket artwork with selectable color choices",
    },
    53: {
        "title": "Custom Exploding Soccer Ball Comforter Set",
        "meta_title": "Custom Exploding Soccer Ball Comforter Set",
        "meta_description": "Customize an exploding soccer ball comforter set with optional custom text, product type, size, pillowcase and sheet-cover choices.",
        "primary": "custom exploding soccer ball comforter set",
        "secondary": "exploding soccer comforter, custom soccer bedding, soccer ball comforter set",
        "cluster": "custom exploding soccer ball comforter set",
        "intent_role": "A10 exploding soccer page with verified optional Customize Your Item field up to 1000 characters.",
        "detail": "JACKSON exploding soccer ball comforter artwork with sample text only as an example",
    },
    54: {
        "title": "Custom Fiery Soccer Ball Comforter Set",
        "meta_title": "Custom Fiery Soccer Ball Comforter Set",
        "meta_description": "Customize a fiery soccer ball comforter set with optional custom text, product type, size, pillowcase and sheet-cover choices.",
        "primary": "custom fiery soccer ball comforter set",
        "secondary": "fiery soccer comforter, custom soccer ball bedding, soccer comforter with text",
        "cluster": "custom fiery soccer ball comforter set",
        "intent_role": "Design 5 fiery soccer page with verified optional Customize Your Item field up to 1000 characters; source title typo avoided.",
        "detail": "KENZO fiery soccer ball comforter artwork with sample text only as an example",
    },
    55: {
        "title": "Custom Flaming Soccer Ball Comforter Set",
        "meta_title": "Custom Flaming Soccer Ball Comforter Set",
        "meta_description": "Customize a red flaming soccer ball comforter with Michael sample script text, product type, size, pillowcase and sheet-cover choices.",
        "primary": "custom red flaming soccer ball comforter",
        "secondary": "Michael soccer comforter sample, flaming soccer ball bedding, custom soccer comforter set",
        "cluster": "custom red flaming soccer ball comforter",
        "intent_role": "A12 red flaming soccer ball page separated from product 56 by script Michael sample text and ball-trail artwork.",
        "detail": "red flaming soccer ball comforter artwork with 07 and Michael script sample text",
    },
    56: {
        "title": "Custom Flaming Soccer Bedding Set",
        "meta_title": "Custom Flaming Soccer Bedding Set",
        "meta_description": "Customize a flaming soccer bedding set with MICHAEL 07 block-style sample text, product type, size and bedding choices.",
        "primary": "custom MICHAEL 07 soccer bedding set",
        "secondary": "flaming soccer bedding with number, custom soccer comforter with name, soccer ball bedding set",
        "cluster": "custom MICHAEL 07 soccer bedding set",
        "intent_role": "Flaming soccer bedding page separated from product 55 by MICHAEL 07 block-style sample text and broader bedding-set wording.",
        "detail": "flaming soccer bedding artwork with MICHAEL 07 block-style sample text",
    },
    57: {
        "title": "Personalized Proverbs 31 Floral Butterfly Blanket",
        "meta_title": "Personalized Proverbs 31 Floral Butterfly Blanket",
        "meta_description": "Customize a Proverbs 31 floral butterfly blanket with required name, inspirational artwork and blanket size choices.",
        "primary": "personalized Proverbs 31 floral butterfly blanket",
        "secondary": "Proverbs 31 blanket with name, custom floral butterfly blanket, personalized inspirational blanket",
        "cluster": "personalized Proverbs 31 floral butterfly blanket",
        "intent_role": "Design 8 with verified required Custom Name field up to 1000 characters.",
        "detail": "Proverbs 31 floral butterfly inspirational blanket artwork with sample Sophia name",
    },
    58: {
        "title": "Personalized Floral Bible Verse Blanket",
        "meta_title": "Personalized Floral Bible Verse Blanket",
        "meta_description": "Customize a floral Bible verse blanket with required name, SARA artwork, blanket sizes and verified text fields.",
        "primary": "personalized floral Bible verse blanket",
        "secondary": "custom Bible verse blanket, floral Christian blanket with name, personalized inspirational blanket",
        "cluster": "personalized floral Bible verse blanket",
        "intent_role": "Design D3 with verified required Custom Name field up to 1000 characters.",
        "detail": "SARA floral butterfly Bible verse blanket artwork",
    },
    59: {
        "title": "Custom Purple Floral Cross Bible Verse Blanket",
        "meta_title": "Custom Purple Floral Cross Bible Verse Blanket",
        "meta_description": "Customize a purple floral cross Bible verse blanket with optional name, Christian artwork and blanket size choices.",
        "primary": "custom purple floral cross Bible verse blanket",
        "secondary": "purple floral cross blanket, custom Christian cross blanket, Bible verse blanket with name",
        "cluster": "custom purple floral cross Bible verse blanket",
        "intent_role": "Design D13 with verified optional Custom Name field up to 1000 characters.",
        "detail": "purple floral cross Bible verse blanket artwork",
    },
    60: {
        "title": "Custom Floral Cross Butterfly Bible Verse Blanket",
        "meta_title": "Custom Floral Cross Butterfly Bible Verse Blanket",
        "meta_description": "Customize a floral cross butterfly Bible verse blanket with optional name, Christian artwork and blanket size choices.",
        "primary": "custom floral cross butterfly Bible verse blanket",
        "secondary": "floral cross butterfly blanket, custom Bible verse blanket, Christian butterfly blanket",
        "cluster": "custom floral cross butterfly Bible verse blanket",
        "intent_role": "Design D6 with verified optional Custom Name field up to 1000 characters.",
        "detail": "floral cross with butterflies Bible verse blanket artwork",
    },
}


IMAGE_DETAILS = {
    51: [
        ("Person holding pink God Says I Am Christian scripture blanket with large EMILY sample name, floral border and butterflies.", "Pink God Says I Am blanket with EMILY name"),
        ("Flat view of pink scripture blanket with EMILY sample name, floral border, butterflies and Philippians 4:13 text.", "Pink scripture floral blanket flat view"),
        ("Pink God Says I Am blanket draped on couch with EMILY sample name and floral border.", "Pink EMILY scripture blanket couch mockup"),
        ("Pink scripture blanket spread on bed or sofa with God Says I Am text and EMILY sample name.", "Pink EMILY Christian blanket bed mockup"),
        ("Lifestyle reading scene with adult and child under pink EMILY scripture blanket.", "Pink EMILY scripture blanket reading scene"),
        ("Blanket size chart showing PET, BOY/GIRL, ADULTS and BED dimensions.", "Pink Christian scripture blanket size chart"),
        ("Blanket material panel showing fluffy soft texture, high quality, no pilling, no shedding and smooth feel.", "Pink scripture blanket soft material panel"),
    ],
    52: [
        ("Person holding blue God Says I Am blanket with EMILY sample name, floral line art and butterfly border.", "Blue God Says I Am blanket with EMILY name"),
        ("Flat view of blue EMILY God Says I Am blanket with iris flower line art and scripture labels.", "Blue EMILY God Says I Am blanket flat view"),
        ("Blue God Says I Am blanket spread on couch with iris flower artwork and vertical EMILY letters.", "Blue EMILY scripture blanket couch mockup"),
        ("Blue scripture blanket draped on sofa with floral border and EMILY sample name.", "Blue EMILY God Says I Am sofa mockup"),
        ("Lifestyle reading scene with adult and child under blue EMILY God Says I Am blanket.", "Blue EMILY scripture blanket reading scene"),
        ("Color choice panel showing blue, gray, red, green, purple and yellow blanket variants.", "Emily God Says I Am blanket color options"),
        ("Blanket size chart showing PET, BOY/GIRL, ADULTS and BED dimensions.", "Emily God Says I Am blanket size chart"),
        ("Blanket material panel showing fluffy soft texture, high quality, no pilling, no shedding and smooth feel.", "Emily God Says I Am blanket material panel"),
    ],
    53: [
        ("Exploding soccer ball comforter bed mockup with shattered black-gray background, JACKSON sample text and number 15.", "Exploding soccer ball comforter with JACKSON 15"),
        ("Microfiber feature panel showing exploding soccer bedding, 3D printed pattern, soft filling and breathable comfort.", "Exploding soccer comforter microfiber panel"),
        ("Bedroom mockup of exploding soccer ball comforter with JACKSON sample text and number 15 pillows.", "Exploding soccer ball bedding room mockup"),
        ("Bedding type comparison panel showing duvet cover set and comforter set options.", "Exploding soccer comforter bedding type panel"),
        ("Close-up fabric panel showing white underside, soccer ball print and fluffy microfiber filling callout.", "Exploding soccer comforter filling close-up"),
        ("Feature icon panel on exploding soccer comforter showing soft, lightweight, durable and breathable notes.", "Exploding soccer comforter feature icons"),
        ("Size dimension panel showing Twin, Full, Queen and King bedding dimensions.", "Exploding soccer comforter size chart"),
        ("Stress-free easy care panel with wrinkle-free, stain-proof, anti-pilling and color-care notes.", "Exploding soccer comforter easy care panel"),
    ],
    54: [
        ("Fiery soccer ball comforter bed mockup with orange flames, KENZO sample text and matching pillows.", "Fiery soccer ball comforter with KENZO text"),
        ("Microfiber feature panel showing fiery soccer bedding, 3D printed pattern, soft filling and breathable comfort.", "Fiery soccer comforter microfiber panel"),
        ("Pillow sham mockup showing fiery soccer ball artwork and KENZO sample text on red pillows.", "Fiery soccer KENZO pillow sham mockup"),
        ("Bedding type comparison panel showing duvet cover set and comforter set options.", "Fiery soccer comforter bedding type panel"),
        ("Feature icon panel on fiery soccer comforter showing soft, lightweight, durable and breathable notes.", "Fiery soccer comforter feature icons"),
        ("Bedroom mockup showing fiery soccer ball comforter with KENZO sample text.", "Fiery soccer ball comforter room mockup"),
        ("Stress-free easy care panel with wrinkle-free, stain-proof, anti-pilling and color-care notes.", "Fiery soccer comforter easy care panel"),
    ],
    55: [
        ("Red flaming soccer ball comforter bed mockup with 07 sample number, Michael script name and matching pillows.", "Red flaming soccer comforter with Michael 07"),
        ("Microfiber feature panel showing red flaming soccer bedding, 3D printed pattern, soft filling and breathable comfort.", "Red flaming soccer comforter microfiber panel"),
        ("Feature icon panel on red flaming soccer comforter showing soft, lightweight, durable and breathable notes.", "Red flaming soccer comforter feature icons"),
        ("Pillow sham mockup showing flaming soccer ball artwork with 07 sample number.", "Red flaming soccer pillow sham mockup"),
        ("Bedding type comparison panel showing duvet cover set and comforter set options.", "Red flaming soccer comforter bedding type panel"),
        ("Close-up fabric panel showing white underside over red soccer artwork and Michael sample text.", "Red flaming soccer comforter fabric close-up"),
        ("Bedroom mockup of red flaming soccer ball comforter with Michael script sample name.", "Red flaming soccer comforter room mockup"),
        ("Size dimension panel showing Twin, Full, Queen and King bedding dimensions.", "Red flaming soccer comforter size chart"),
    ],
    56: [
        ("Flaming soccer bedding bed mockup with MICHAEL block-style sample name, 07 number and matching pillows.", "MICHAEL 07 flaming soccer bedding set"),
        ("Microfiber feature panel showing flaming soccer bedding, 3D printed pattern, soft filling and breathable comfort.", "MICHAEL 07 soccer bedding microfiber panel"),
        ("Close-up fabric panel showing white underside over flaming soccer print and fluffy microfiber filling callout.", "MICHAEL 07 soccer comforter filling close-up"),
        ("Pillow sham mockup showing flaming soccer ball artwork with 07 sample number.", "MICHAEL 07 soccer pillow sham mockup"),
        ("Bedding type comparison panel showing duvet cover set and comforter set options.", "MICHAEL 07 soccer bedding type panel"),
        ("Bedroom mockup of flaming soccer bedding with MICHAEL block-style sample name and 07 pillows.", "MICHAEL 07 soccer bedding room mockup"),
        ("Feature icon panel on flaming soccer bedding showing soft, lightweight, durable and breathable notes.", "MICHAEL 07 soccer bedding feature icons"),
        ("Size dimension panel showing Twin, Full, Queen and King bedding dimensions.", "MICHAEL 07 soccer bedding size chart"),
    ],
    57: [
        ("Person holding Proverbs 31 floral butterfly blanket with SOPHIA sample name, peach flowers and gold corner border.", "Proverbs 31 floral butterfly blanket SOPHIA"),
        ("Flat view of SOPHIA Proverbs 31 blanket with peach flowers, monarch butterfly and gold border.", "SOPHIA Proverbs 31 blanket flat view"),
        ("Lifestyle reading scene with floral butterfly blanket, SOPHIA sample name and Proverbs 31:25 text.", "SOPHIA floral butterfly blanket reading scene"),
        ("Blanket size chart showing PET, BOY/GIRL, ADULTS and BED dimensions.", "Proverbs 31 floral butterfly blanket size chart"),
        ("Blanket material panel showing fluffy soft texture, high quality, no pilling, no shedding and smooth feel.", "Proverbs 31 blanket soft material panel"),
        ("Detail collage showing fleece folds, butterfly artwork, printed floral design and soft white backing.", "Proverbs 31 blanket texture detail collage"),
        ("Floral butterfly blanket draped on couch with SOPHIA sample name and Proverbs 31:25 text.", "SOPHIA floral butterfly blanket couch mockup"),
        ("SOPHIA floral butterfly blanket spread on sofa with peach flower artwork and gold border.", "Proverbs 31 floral butterfly blanket sofa mockup"),
    ],
    58: [
        ("Person holding SARA floral Bible verse blanket with leaf border and scripture affirmation list.", "SARA floral Bible verse blanket"),
        ("SARA floral Bible verse blanket draped on couch with scripture list and pale floral border.", "SARA Bible verse blanket couch mockup"),
        ("Blanket use panel showing 80 by 60 inch example, personalization callout and sofa, office, bed and travel icons.", "SARA Bible verse blanket size and use panel"),
        ("Detail collage showing fleece folds, scripture print, floral artwork and soft white backing.", "SARA Bible verse blanket texture detail collage"),
        ("Fleece material panel listing machine washable care, soft warmth, dense stitching and 260GSM fleece.", "SARA Bible verse blanket fleece material"),
        ("Lifestyle bed reading scene with SARA floral blanket, open book and small dog.", "SARA Bible verse blanket reading lifestyle"),
        ("SARA floral Bible verse blanket draped on armchair beside books and mug.", "SARA Bible verse blanket armchair mockup"),
        ("Close-up reading lifestyle view with SARA floral scripture blanket and mug.", "SARA Bible verse blanket close-up"),
    ],
    59: [
        ("Person holding purple floral cross blanket with Lord stood with me verse text and Isabella sample name.", "Purple floral cross blanket Isabella"),
        ("Purple floral cross blanket spread on bed with Lord stood with me verse text and Isabella sample name.", "Purple floral cross blanket bed mockup"),
        ("Purple floral cross blanket draped on couch with lilies, cross artwork and Isabella sample name.", "Purple floral cross blanket couch mockup"),
        ("Close-up reading lifestyle view with purple floral cross blanket, book and mug.", "Purple floral cross blanket reading close-up"),
        ("Lifestyle bed reading scene with purple floral cross blanket, open book and small dog.", "Purple floral cross blanket reading lifestyle"),
        ("Blanket use panel showing 80 by 60 inch example and sofa, office, bed, travel, reading and camping icons.", "Purple floral cross blanket size and use panel"),
        ("Fleece material panel listing machine washable care, soft warmth, dense stitching and 260GSM fleece.", "Purple floral cross blanket fleece material"),
        ("Detail collage showing fleece folds, purple lily artwork, cross print and soft white backing.", "Purple floral cross blanket texture detail collage"),
    ],
    60: [
        ("Person holding floral cross butterfly blanket with gold butterflies, pink floral border and Sophia sample name.", "Floral cross butterfly blanket Sophia"),
        ("Flat view of floral cross butterfly blanket with rose bouquet, gold butterflies and Sophia sample name.", "Floral cross butterfly blanket flat view"),
        ("Lifestyle reading scene with floral cross butterfly blanket and rose artwork visible.", "Floral cross butterfly blanket reading scene"),
        ("Blanket size chart showing PET, BOY/GIRL, ADULTS and BED dimensions.", "Floral cross butterfly blanket size chart"),
        ("Detail collage showing fleece folds, gold butterfly print, floral artwork and soft white backing.", "Floral cross butterfly blanket texture details"),
        ("Blanket material panel showing fluffy soft texture, high quality, no pilling, no shedding and smooth feel.", "Floral cross butterfly blanket material panel"),
        ("Floral cross butterfly blanket draped on couch with Sophia sample name and rose border.", "Floral cross butterfly blanket couch mockup"),
        ("Floral cross butterfly blanket spread on sofa with gold butterflies and pink flower border.", "Floral cross butterfly blanket sofa mockup"),
    ],
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def load_qa():
    data = json.loads(QA_DATASET.read_text(encoding="utf-8"))
    products = data["QA_Products"]
    handle_by_key = {item["product_key"]: item["handle"] for item in products}
    pos_by_key = {item["product_key"]: int(item["inventory_position"]) for item in products}
    image_by_qa_key = {img["qa_image_key"]: img for img in data["QA_Images"]}
    image_fix_by_key = {}
    observation_by_key = {}
    for image in data["QA_Images"]:
        observation_by_key[(image["product_key"], text(image["media_id"]))] = image["qa_observation"]
    for issue in data["QA_Issues"]:
        if issue["field"].startswith("image_") and issue["recommended_fix"]:
            image = image_by_qa_key.get(issue["qa_image_key"])
            if image:
                image_fix_by_key[(image["product_key"], text(image["media_id"]))] = issue["recommended_fix"]
    return {
        "data": data,
        "product_keys": [item["product_key"] for item in products],
        "handle_by_key": handle_by_key,
        "pos_by_key": pos_by_key,
        "image_fix_by_key": image_fix_by_key,
        "observation_by_key": observation_by_key,
    }


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Body (HTML)": first.get("Body (HTML)", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    return admin


def load_field_audit():
    return {int(item["inventory_position"]): item for item in json.loads(CUSTOMIZER_FIELDS_AUDIT.read_text(encoding="utf-8"))}


def field_sentence(pos, field_audit):
    item = field_audit.get(pos, {})
    fields = item.get("fields", [])
    text_inputs = [entry["fields"] for entry in fields if "textInputs[" in entry.get("path", "")]
    option_groups = [entry["fields"] for entry in fields if re.search(r"(^|\.)optionGroups\[\d+\]$", entry.get("path", ""))]
    parts = []
    for text_input in text_inputs:
        required = "required" if text_input.get("required") else "optional"
        label = text_input.get("label", "text field")
        limit = text_input.get("maxLength")
        if limit:
            parts.append(f"Personalization uses a {required} {label} field up to {limit} characters.")
        else:
            parts.append(f"Personalization uses a {required} {label} field.")
    if option_groups:
        labels = ", ".join(group.get("label", "") for group in option_groups[:4] if group.get("label"))
        if labels:
            parts.append(f"Additional option groups include: {labels}.")
    return " ".join(parts) or "No shopper text-entry field is described for this product."


def compact_observation_to_alt(observation, fallback):
    value = re.sub(r"[:;].*$", "", text(observation)).strip()
    value = re.sub(r"\s+", " ", value)
    return (value or fallback)[:125]


def component_note(pos, product_type):
    if pos in {51, 52, 57, 58, 59, 60}:
        return "Gallery panels show blanket sizes, use cases, fleece or material details and lifestyle scenes where available."
    return "Gallery panels show bedding type choices, size dimensions, feature icons, care notes and pillow/sham visuals where available."


def description(pos, admin_row, ctext):
    item = PRODUCT_UPDATES[pos]
    option_names = [v for v in [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]] if v]
    options = ", ".join(option_names) if option_names else "available product selectors"
    product_type = admin_row["Type"].lower() if admin_row["Type"] else "product"
    return (
        f"<p>{item['detail'].capitalize()} gives this Jeminise {product_type} a clear design focus for personalized gifting, faith-inspired decor, sports rooms or everyday bedding.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork focus: {item['detail']}.</li>"
        "<li>Main images and support panels show the visible motif, sample personalization, sizing, use cases, care or material details where shown.</li>"
        f"<li>{component_note(pos, product_type)}</li></ul>"
        "<h3>Options and Customization</h3>"
        f"<ul><li>Available selectors cover {options}.</li>"
        f"<li>{ctext}</li>"
        "<li>Use the product selectors to confirm size, product type and any required custom fields before checkout.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    qa = load_qa()
    field_audit = load_field_audit()
    admin = load_admin(set(qa["handle_by_key"].values()))
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ctext = field_sentence(pos, field_audit)

        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or ws.cell(row_num, idx["product_type"]).value
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, ctext)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"Buyer intent targets {item['cluster']} for US English product search; no search-volume claim is made."
        )
        ws.cell(row_num, idx["revision"]).value = "r4"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R4 deep recheck: used products_export_1.csv (sha256:{admin_hash}) as admin baseline; "
            "rewrote customer-facing descriptions, removed blanket pillowcase/sham wording, mapped verified Customizer rules, "
            "rewrote image observations/alt by visual panel, and separated overlapping keyword clusters. Still NEEDS_REVIEW; no APPROVED/import."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    image_seen = defaultdict(int)
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        media_id = text(ws.cell(row_num, idx["media_id"]).value)
        key = (pk, media_id)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        image_seen[pos] += 1
        detail_list = IMAGE_DETAILS[pos]
        detail_index = min(image_seen[pos] - 1, len(detail_list) - 1)
        observation, proposed_alt = detail_list[detail_index]
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = proposed_alt[:125]
        ws.cell(row_num, idx["revision"]).value = "r4"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R4: image observation and alt rewritten from contact-sheet review; current admin alt and image URL matched from products_export_1.csv where possible."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "Product"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = "Sang re-QA r2 SERP evidence + products_export_1.csv admin baseline + customizer_fields_audit.json"
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = "R4 separates sibling product intent by motif, sample text style and verified Customizer rules."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r4"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as a specific bedding product."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm artwork, product type, size choices "
            "and verified customizer rules."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = "Confirm size, product type, verified text fields, sample values, care panels and image accuracy."
        ws.cell(row_num, idx["emotional_social_motivation"]).value = "Choose a faith, soccer or inspirational bedding gift with clear artwork."
        ws.cell(row_num, idx["purchase_concerns"]).value = "Exact Customizer rules, included components, fabric/care claims, blanket versus comforter options and image accuracy."
        ws.cell(row_num, idx["source_refs"]).value = f"products_export_1.csv sha256:{admin_hash}; Sang QA r2 {QA_RUN_ID}; customizer_fields_audit.json"
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R4_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No search-volume source supplied; Shopify admin CSV baseline is now available for stored title/meta/body/image alt."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = qa["product_keys"][pos - 51]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = f"Storefront/product.js from Sang QA r2; products_export_1.csv sha256:{admin_hash}; customizer_fields_audit.json"
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={admin_row['Option1 Name']}, "
            f"{admin_row['Option2 Name']}, {admin_row['Option3 Name']}; status={admin_row['Status']}; "
            f"image_count={len(admin_row['images'])}; design={item['detail']}; customizer={field_sentence(pos, field_audit)}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "R4 applied deep batch recheck; needs independent QA recheck."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R4 uses contact-sheet image review, detailed Customizer field audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_006_r4_revision", "r4", "Separate 10-product deep revision after r3; source r3 workbook was not modified."),
        ("qa_batch_006_r4_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_006_r4_scope", "inventory positions 51-60", "No products outside qa_batch_006 were revised."),
        ("qa_batch_006_r4_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_006_r4_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_006_r4_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa_r2": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_006 only; inventory positions 51-60",
        "revision": "r4",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "r4_changes": [
            "Deepened product-specific descriptions for all 10 products using admin export, customizer audit and contact-sheet evidence.",
            "Rewrote 78 image observations and alt texts by visual panel instead of using repeated product-level templates.",
            "Removed blanket pillowcase/sham wording from products 51, 52, 57, 58, 59 and 60.",
            "Mapped verified customizer rules including required/optional Custom Name, optional Customize Your Item and required Choose Color on product 52.",
            "Separated soccer comforter clusters 53-56 by motif, sample text and bedding intent."
        ],
        "next_step": "Sang QA lại qa_batch_006_r4 before any approval or import.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_006_r4",
                "",
                "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r3: `{SOURCE.relative_to(ROOT)}`",
                f"- QA r2 của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r4: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Sửa trọng yếu: 78 ảnh được viết lại theo từng visual panel; blanket bỏ pillowcase/sham; product 52 ghi rõ color choice; soccer 53-56 tách intent theo motif và sample text.",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Bước tiếp theo: Sang QA lại `qa_batch_006_r4`.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
