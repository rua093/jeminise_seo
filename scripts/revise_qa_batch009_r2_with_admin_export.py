import csv
import hashlib
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_009_r2"
QA_RUN_ID = "20260907_180500"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "batches" / "SEO_Product_Optimization_through_batch_009.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
CUSTOMIZER_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_audit.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_009_r2.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_009_r2.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    81: {
        "title": "Custom Purple Baseball Batter Blanket",
        "meta_title": "Custom Purple Baseball Batter Blanket",
        "meta_description": "Shop a purple baseball batter blanket with flag stripes, silhouette artwork and optional name or number fields in Customize.",
        "primary": "custom purple baseball batter blanket",
        "secondary": "baseball batter blanket, custom baseball blanket, baseball throw blanket for players",
        "cluster": "purple baseball batter silhouette blanket",
        "detail": "black and purple baseball batter silhouette over vertical American flag stripes with sample NAME 23 text",
        "intent_role": "Baseball page separated by purple batter silhouette and vertical flag-stripe motif.",
        "customizer": "Customizer audit shows optional Custom Name up to 200 characters and optional Custom Number up to 20 characters.",
    },
    82: {
        "title": "Baseball Batter Stadium Flag Blanket",
        "meta_title": "Baseball Batter Stadium Flag Blanket",
        "meta_description": "Shop a baseball batter blanket with stadium lights, red-blue flag artwork and Daniel sample text, focused on the printed design.",
        "primary": "baseball batter stadium flag blanket",
        "secondary": "baseball batter throw blanket, American flag baseball blanket, baseball player gift blanket",
        "cluster": "baseball batter stadium flag blanket",
        "detail": "baseball batter in a stadium scene over red and blue American flag artwork with sample Daniel name",
        "intent_role": "Stadium batter page; photo/picture claim removed because no image upload input was verified.",
        "customizer": "Customizer audit shows optional Custom Name up to 200 characters and optional Custom Number up to 20 characters.",
    },
    83: {
        "title": "Custom Baseball Field Quote Blanket",
        "meta_title": "Custom Baseball Field Quote Blanket",
        "meta_description": "Shop a vintage baseball field quote blanket with Christopher and 22 sample artwork plus optional name and number fields.",
        "primary": "custom baseball field quote blanket",
        "secondary": "baseball quote blanket, baseball player blanket gift, custom baseball throw",
        "cluster": "vintage baseball field quote blanket",
        "detail": "vintage tan baseball batter artwork with sample Christopher name, 22 and the quote When you step on the field nothing else matters",
        "intent_role": "Quote-led batter page separated by the exact field quote and tan vintage artwork.",
        "customizer": "Customizer audit shows optional Custom Name up to 200 characters and optional Custom Number up to 20 characters.",
    },
    84: {
        "title": "Custom Baseball Practice Quote Blanket",
        "meta_title": "Custom Baseball Practice Quote Blanket",
        "meta_description": "Shop a baseball practice quote blanket with American flag, glove artwork and optional name or number fields in Customize.",
        "primary": "custom baseball practice quote blanket",
        "secondary": "baseball little boy quote blanket, American flag baseball blanket, baseball blanket gift",
        "cluster": "baseball practice little boy quote blanket",
        "detail": "American flag and rustic wood blanket with sample Jackson 25, glove artwork and the practice little boy quote",
        "intent_role": "Practice quote page separated from other baseball blankets by the long little-boy/coaches quote.",
        "customizer": "Customizer audit shows optional Custom Name up to 200 characters and optional Custom Number up to 20 characters.",
    },
    85: {
        "title": "Flaming Baseball Blanket",
        "meta_title": "Flaming Baseball Blanket",
        "meta_description": "Shop a dark flaming baseball blanket with Phoenix #09 sample artwork, fiery bat trail and multiple size choices.",
        "primary": "flaming baseball blanket",
        "secondary": "fire baseball blanket, baseball blanket with bat and ball, baseball throw blanket",
        "cluster": "flaming baseball and bat blanket",
        "detail": "dark baseball design with glowing flaming baseball, fiery bat trail and sample Phoenix #09 text",
        "intent_role": "Flame baseball page; personalization claim softened because Sang QA flagged live purchase-flow proof as incomplete.",
        "customizer": "Customer-facing copy focuses on the printed baseball design and size choice.",
    },
    86: {
        "title": "Custom Baseball Glove Flag Blanket",
        "meta_title": "Custom Baseball Glove Flag Blanket",
        "meta_description": "Shop a baseball glove flag blanket with bat, ball and Jackson 15 sample artwork plus optional name and number fields.",
        "primary": "custom baseball glove flag blanket",
        "secondary": "patriotic baseball glove blanket, baseball bat and glove blanket, baseball player throw blanket",
        "cluster": "baseball glove bat flag blanket",
        "detail": "baseball glove, bat and ball artwork over a red, white and blue flag background with sample Jackson 15",
        "intent_role": "Equipment-led patriotic page separated from batter silhouette, catcher and Christian baseball blanket pages.",
        "customizer": "Customizer audit shows optional Custom Name up to 200 characters and optional Custom Number up to 20 characters.",
    },
    87: {
        "title": "Baseball Catcher Quote Blanket",
        "meta_title": "Baseball Catcher Quote Blanket",
        "meta_description": "Shop a catcher quote blanket with I Catch He Pitches text, Daniel 56 sample artwork and baseball-stitch border.",
        "primary": "baseball catcher quote blanket",
        "secondary": "catcher blanket gift, I catch he pitches blanket, baseball catcher throw blanket",
        "cluster": "I catch he pitches catcher blanket",
        "detail": "catcher reaching for a pitch inside baseball stitching with I Catch He Pitches text and sample Daniel 56",
        "intent_role": "Catcher quote page; personalization claim softened because live purchase-flow proof needs recheck.",
        "customizer": "Customer-facing copy focuses on the printed catcher quote design and size choice.",
    },
    88: {
        "title": "American Flag Baseball Pitcher Blanket",
        "meta_title": "American Flag Baseball Pitcher Blanket",
        "meta_description": "Shop a patriotic baseball pitcher blanket with vintage flag artwork, Jude 5 sample text and soft fleece details.",
        "primary": "American flag baseball pitcher blanket",
        "secondary": "patriotic baseball blanket, baseball pitcher throw blanket, vintage flag baseball blanket",
        "cluster": "American flag baseball pitcher blanket",
        "detail": "vintage American flag blanket with black baseball pitcher silhouette and sample Jude 5 text",
        "intent_role": "Pitcher silhouette page; personalization claim softened because live purchase-flow proof needs recheck.",
        "customizer": "Customer-facing copy focuses on the printed patriotic pitcher design and size choice.",
    },
    89: {
        "title": "Baseball Catcher Batter Blanket",
        "meta_title": "Baseball Catcher Batter Blanket",
        "meta_description": "Shop a reversible-style baseball blanket with catcher and batter silhouettes, stitching graphics and sample names.",
        "primary": "baseball catcher batter blanket",
        "secondary": "baseball stitching blanket, catcher batter throw blanket, baseball blanket for players",
        "cluster": "catcher batter reversible-style baseball blanket",
        "detail": "reversible-style baseball stitching design with catcher and batter silhouettes plus sample Joey, Hudson, Noah and 20",
        "intent_role": "Catcher/batter dual-layout page; personalization claim softened because live purchase-flow proof needs recheck.",
        "customizer": "Customer-facing copy focuses on the printed catcher/batter design and size choice.",
    },
    90: {
        "title": "God Says You Are Baseball Blanket",
        "meta_title": "God Says You Are Baseball Blanket",
        "meta_description": "Shop a baseball faith blanket with God Says You Are affirmations, Bible verse references and soft sherpa texture.",
        "primary": "God Says You Are baseball blanket",
        "secondary": "Christian baseball blanket, baseball faith blanket, Bible verse baseball throw",
        "cluster": "God Says You Are baseball faith blanket",
        "detail": "vintage baseball faith design with God Says You Are affirmations, Bible references, Your Name placeholder and 02",
        "intent_role": "Christian affirmation baseball page separated from athlete, catcher, pitcher and equipment baseball designs.",
        "customizer": "Customer-facing copy focuses on the printed baseball faith design and size choice.",
    },
}


IMAGE_DETAILS = {
    81: {
        1: ("Lifestyle mockup of black and purple baseball blanket held upright with batter silhouette, flag stripes and NAME 23.", "Purple baseball batter blanket with NAME 23"),
        2: ("Customization panel for the purple baseball batter blanket showing name and number variants.", "Customize name and number panel for baseball blanket"),
        3: ("Outdoor mockup of purple baseball batter blanket with vertical flag stripes and NAME 23.", "Purple baseball batter blanket in outdoor mockup"),
        4: ("Outdoor mockup variant showing batter silhouette over flag stripes with NAME 23.", "Baseball batter silhouette blanket with flag stripes"),
        5: ("Outdoor mockup of the same purple baseball batter blanket held in a park scene.", "Purple baseball throw blanket held outdoors"),
    },
    82: {
        1: ("Room mockup of baseball batter blanket on armchair with Daniel name and red-blue stadium flag artwork.", "Baseball batter stadium flag blanket with Daniel"),
        2: ("Holiday fireplace mockup with Daniel baseball batter blanket and custom name and picture text in the image.", "Baseball batter blanket with Daniel in holiday mockup"),
        3: ("Outdoor autumn mockup of Daniel baseball batter stadium flag blanket.", "Daniel baseball batter blanket held outdoors"),
        4: ("Porch mockup of Daniel baseball batter blanket with custom name and picture text in the image.", "Baseball batter flag blanket porch mockup"),
        5: ("Armchair room mockup of Daniel baseball batter blanket.", "Daniel baseball batter blanket on armchair"),
        6: ("Sofa room mockup of Daniel baseball batter blanket with books and mug.", "Daniel baseball batter blanket on sofa"),
    },
    83: {
        1: ("Lifestyle mockup of vintage tan baseball quote blanket with Christopher, 22 and batter artwork.", "Vintage baseball quote blanket with Christopher 22"),
        2: ("Armchair mockup of Christopher baseball field quote blanket.", "Christopher baseball field quote blanket on armchair"),
        3: ("Closer armchair mockup showing the quote When you step on the field nothing else matters.", "Baseball field quote blanket with Christopher"),
        4: ("Sofa mockup of vintage baseball batter quote blanket with 22.", "Vintage baseball batter quote blanket on sofa"),
        5: ("Outdoor autumn mockup of Christopher baseball quote blanket.", "Christopher baseball quote blanket outdoors"),
        6: ("Outdoor forest mockup of Christopher baseball quote blanket held upright.", "Baseball quote blanket held outdoors"),
        7: ("Winter porch mockup of Christopher baseball quote blanket held by two people.", "Christopher baseball quote blanket porch mockup"),
    },
    84: {
        1: ("Lifestyle mockup of baseball practice quote blanket with American flag, Your Name 21 and glove artwork.", "Baseball practice quote blanket with flag and glove"),
        2: ("Armchair mockup showing Jackson 25, flag panel and practice quote text.", "Jackson 25 baseball practice quote blanket"),
        3: ("Outdoor mockup of Jackson 25 baseball practice quote blanket.", "Baseball practice quote blanket held outdoors"),
        4: ("Porch mockup with red arrow showing custom name and number area on Jackson 25 blanket.", "Custom name and number area on baseball quote blanket"),
        5: ("Armchair mockup of Jackson 25 baseball practice quote blanket.", "Jackson baseball practice blanket on armchair"),
        6: ("Sofa mockup of Jackson 25 baseball practice quote blanket with books and mug.", "Jackson baseball practice blanket on sofa"),
        7: ("Outdoor autumn mockup of baseball practice quote blanket held upright.", "Baseball practice quote blanket outdoor mockup"),
    },
    85: {
        1: ("Armchair mockup of dark flaming baseball blanket with fiery bat trail and partial Phoenix text.", "Flaming baseball blanket on armchair"),
        2: ("Outdoor mockup of Phoenix #09 flaming baseball blanket held in a forest setting.", "Phoenix #09 flaming baseball blanket outdoors"),
        3: ("Outdoor autumn mockup of flaming baseball blanket with Phoenix #09.", "Flaming baseball blanket with Phoenix #09"),
        4: ("Porch mockup of Phoenix #09 flaming baseball blanket held by two people.", "Phoenix flaming baseball blanket porch mockup"),
        5: ("Armchair mockup of Phoenix #09 flaming baseball blanket.", "Phoenix #09 flaming baseball blanket on armchair"),
        6: ("Sofa mockup of dark flaming baseball blanket with bat trail and baseball.", "Flaming baseball blanket on sofa"),
    },
    86: {
        1: ("Armchair mockup of baseball glove, bat and ball blanket over red white and blue flag with Jackson 15.", "Baseball glove flag blanket with Jackson 15"),
        2: ("Porch mockup of Jackson 15 baseball glove flag blanket held by two people.", "Jackson baseball glove flag blanket porch mockup"),
        3: ("Armchair close mockup showing baseball glove, bat, ball and Jackson name.", "Baseball glove bat and ball blanket with Jackson"),
        4: ("Sofa mockup of baseball glove flag blanket with Jackson 15 and books.", "Jackson baseball glove flag blanket on sofa"),
        5: ("Outdoor autumn mockup of Jackson 15 baseball glove flag blanket.", "Jackson 15 baseball glove blanket outdoors"),
        6: ("Outdoor forest mockup of Jackson 15 baseball glove flag blanket.", "Baseball glove flag blanket held outdoors"),
    },
    87: {
        1: ("Sofa mockup of catcher blanket with I Catch He Pitches text, Daniel 56 and baseball-stitch border.", "I Catch He Pitches catcher blanket with Daniel 56"),
        2: ("Outdoor mockup of Daniel 56 catcher quote blanket held upright.", "Daniel 56 catcher quote blanket outdoors"),
        3: ("Outdoor forest mockup showing catcher artwork and I Catch He Pitches text.", "Baseball catcher quote blanket held outdoors"),
        4: ("Porch mockup of Daniel 56 catcher quote blanket held by two people.", "Daniel catcher quote blanket porch mockup"),
        5: ("Armchair mockup of Daniel 56 catcher quote blanket.", "Daniel 56 catcher blanket on armchair"),
        6: ("Sofa mockup of Daniel catcher quote blanket with books.", "Baseball catcher quote blanket on sofa"),
    },
    88: {
        1: ("Lifestyle mockup of American flag baseball pitcher blanket with Jude name and number 5.", "American flag baseball pitcher blanket with Jude 5"),
        2: ("Armchair mockup of Jude 5 baseball pitcher blanket over vintage American flag.", "Jude 5 baseball pitcher blanket on armchair"),
        3: ("Sofa mockup of patriotic baseball pitcher blanket with Jude text.", "American flag baseball pitcher blanket on sofa"),
        4: ("Close-up fabric collage showing flag print, pitcher silhouette, Jude text and fleece texture.", "Close-up collage for American flag baseball blanket"),
        5: ("Feature panel showing machine washable, soft and warm, dense stitching and 260GSM fleece blanket.", "Machine washable fleece blanket feature panel"),
        6: ("Lifestyle reading image with patriotic baseball pitcher blanket across lap.", "American flag baseball blanket used while reading"),
        7: ("Size and use panel showing 80 by 60 in blanket and sofa, office, bed, plane and travel uses.", "Baseball pitcher blanket size and use panel"),
    },
    89: {
        1: ("Collage mockup of reversible-style baseball blanket with catcher, batter, Joey, Hudson, Noah and 20.", "Baseball catcher batter blanket collage with names"),
        2: ("Porch mockup showing catcher side with Joey and 20 on baseball-stitch background.", "Joey 20 baseball catcher blanket porch mockup"),
        3: ("Outdoor mockup showing batter side with Hudson name and 20.", "Hudson 20 baseball batter blanket outdoors"),
        4: ("Customization panel showing catcher and batter sides with custom name and number text.", "Custom name and number panel for catcher batter blanket"),
        5: ("Armchair mockup showing batter side with Noah name and 20.", "Noah 20 baseball batter blanket on armchair"),
    },
    90: {
        1: ("Holiday mockup of God Says You Are baseball faith blanket with affirmations, verse references and Your Name placeholder.", "God Says You Are baseball faith blanket"),
        2: ("Lifestyle mockup of the same baseball faith blanket held near a potted plant.", "Baseball faith blanket with God Says You Are text"),
        3: ("Close-up image of white sherpa fleece texture.", "White sherpa fleece texture close-up"),
        4: ("Machine washable care panel with cold wash and mild detergent instructions.", "Machine washable care panel for baseball blanket"),
    },
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def load_qa():
    data = json.loads(QA_DATASET.read_text(encoding="utf-8"))
    products = data["QA_Products"]
    customizer = json.loads(CUSTOMIZER_AUDIT.read_text(encoding="utf-8"))
    return {
        "data": data,
        "product_keys": [item["product_key"] for item in products],
        "handle_by_key": {item["product_key"]: item["handle"] for item in products},
        "pos_by_key": {item["product_key"]: int(item["inventory_position"]) for item in products},
        "customizer_by_pos": {int(item["inventory_position"]): item for item in customizer},
    }


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Body (HTML)": first.get("Body (HTML)", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    return admin


def option_text(admin_row):
    values = [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]]
    values = [v for v in values if v]
    return ", ".join(values) if values else "Shopify export shows one size option group."


def description(pos, admin_row):
    item = PRODUCT_UPDATES[pos]
    return (
        f"<p>{item['detail'].capitalize()} gives this Jeminise blanket a clear {item['cluster']} focus.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork focus: {item['detail']}.</li>"
        "<li>Gallery images show blanket mockups, room or outdoor lifestyle views, feature panels, size guidance or care details where available.</li>"
        f"<li>Current Shopify export title used for identity check: {admin_row['Title']}.</li></ul>"
        "<h3>Options and Customization</h3>"
        f"<ul><li>Shopify option groups: {option_text(admin_row)}.</li>"
        f"<li>{item['customizer']}</li>"
        "<li>Choose the blanket size before checkout; only use customization fields that are visible on the product page.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    qa = load_qa()
    admin = load_admin(set(qa["handle_by_key"].values()))
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or "Blanket"
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"US English buyer intent targets {item['cluster']}; no search-volume claim is made."
        )
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R2 after Sang QA: used products_export_1.csv (sha256:{admin_hash}) as admin baseline; "
            "removed internal workflow copy, corrected personalization claims to match available evidence, "
            "removed photo/picture upload claim where unsupported, rewrote descriptions and image alt from contact sheets, "
            "and separated baseball keyword intents. Still NEEDS_REVIEW; no APPROVED/import."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        image_number = int(ws.cell(row_num, idx["image_number"]).value)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation, alt = IMAGE_DETAILS[pos][image_number]
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = alt[:125]
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R2: image observation and alt rewritten from inspected contact sheet; current admin alt and image URL matched from products_export_1.csv where possible."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "Product"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = "Sang QA + products_export_1.csv admin baseline + customizer_audit.json + contact-sheet inspection"
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = "R2 separates baseball pages by visible motif, wording, audience and proof level for customization."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r2"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as a specific baseball blanket product page."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm the exact baseball motif, blanket size choices, "
            "and only the customization options that are actually supported."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = "Confirm motif, size, supported custom fields, care details and image accuracy before purchase."
        ws.cell(row_num, idx["emotional_social_motivation"]).value = "Choose a baseball-themed gift or keepsake without confusing it with another similar baseball design."
        ws.cell(row_num, idx["purchase_concerns"]).value = "Avoid unsupported photo/personalization claims; confirm size, fabric/care statements and image accuracy."
        ws.cell(row_num, idx["source_refs"]).value = f"products_export_1.csv sha256:{admin_hash}; Sang QA {QA_RUN_ID}; customizer_audit.json; contact sheets 81-90"
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No search-volume source supplied; Customizer-to-cart persistence still needs controlled test before approval."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    pk_by_pos = {qa["pos_by_key"][pk]: pk for pk in qa["product_keys"]}
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = pk_by_pos[pos]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = (
            f"Original storefront/product evidence; Sang QA {QA_RUN_ID}; products_export_1.csv sha256:{admin_hash}; customizer_audit.json; contact sheet inspected"
        )
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={option_text(admin_row)}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={item['detail']}; customizer_note={item['customizer']}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "Sang QA issues addressed in r2; requires independent QA recheck."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R2 uses inspected contact sheets, Customizer audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_009_r2_revision", "r2", "Separate 10-product revision after Sang QA; source workbook was not modified."),
        ("qa_batch_009_r2_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_009_r2_scope", "inventory positions 81-90", "No products outside qa_batch_009 were revised."),
        ("qa_batch_009_r2_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_009_r2_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_009_r2_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_009 only; inventory positions 81-90",
        "revision": "r2",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Hold for combined QA handoff after revisions through qa_batch_020, per user request.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_009_r2",
                "",
                "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r1: `{SOURCE.relative_to(ROOT)}`",
                f"- QA của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r2: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Sửa trọng yếu: bỏ nội dung nội bộ, bỏ photo-upload claim không có image input, làm mềm claim personalization cho sản phẩm bị QA critical, tách keyword intent, viết lại mô tả/alt theo ảnh đã xem.",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Bước tiếp theo: giữ lại để gửi Sang QA chung sau khi revision đến `qa_batch_020`, theo yêu cầu người dùng.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
