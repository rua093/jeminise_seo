import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_027"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_026.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_027.xlsx"


CHRISTIAN_REFS = "https://www.etsy.com/listing/1880095899/personalized-god-says-you-are-christian; https://macorner.co/products/christian-bible-verse-affirmations-for-girls-boys-personalized-photo-blanket-bwo241101lahn; https://www.etsy.com/market/christian_blanket"
CELTIC_REFS = "https://www.amazon.com/clp/B0DNSVY2QJ; https://www.pinterest.com/pin/562668547209120403/; https://myvikinggear.com/collections/viking-home-decoration/viking-quilt-set/; https://luvingift.com/product/viking-quilt-set-norse-celtic-yggdrasil-tree-of-life/"
RV_REFS = "https://www.amazon.com/cactus-sunset-3-piece-quilted-bedding-set/s?k=cactus+sunset+3-piece+quilted+bedding+set; https://www.etsy.com/listing/1635884318/sw-desert-sunset-duvet-cover-colorful; https://www.target.com/s/rv%2Bbedding"
PHOTO_REFS = "https://www.etsy.com/market/custom_photo_quilt; https://www.amazon.com/custom-photo-blanket/s?k=custom+photo+blanket; https://www.etsy.com/market/valentines_photo_blanket"
READING_REFS = "https://www.etsy.com/listing/4403836729/bookworm-girl-reading-blanket-custom; https://www.etsy.com/market/book_lover_blanket; https://www.amazon.com/clp/B0DJ32WQ8X; https://www.amazon.com/Just-Girl-Loves-Books-Blanket/dp/B0DMVXGZKG"


PROFILES = {
    261: ("Christian woman affirmation Bible verse blanket", "Christian affirmation Bible verse blanket", "Christian Affirmation Bible Verse Blanket", "Christian Woman Affirmation Bible Verse Blanket", "wood plank style throw blanket with seated woman artwork, custom name and colored Bible verse affirmation blocks", "personalized Christian blanket, Bible verse throw blanket, faith affirmation blanket", "EVERGREEN", CHRISTIAN_REFS),
    262: ("vintage Christian scripture photo blanket", "Christian scripture photo blanket", "Christian Scripture Photo Blanket", "Vintage Christian Woman Scripture Photo Blanket", "vintage beige throw blanket with custom photo circle, Christian scripture phrases and faith quote typography", "personalized Christian photo blanket, Bible verse blanket, faith gift blanket", "EVERGREEN", CHRISTIAN_REFS),
    263: ("colorful Celtic Tree of Life quilt", "colorful Celtic Tree of Life quilt", "Colorful Celtic Tree of Life Quilt Set", "Colorful Celtic Tree of Life Quilt Set", "teal and orange quilt with large Tree of Life medallion, Celtic knot frame and colorful leaf accents", "Yggdrasil quilt set, Celtic bedding, Tree of Life bedspread", "EVERGREEN", CELTIC_REFS),
    264: ("vintage colorful Celtic Yggdrasil tree quilt", "colorful Yggdrasil tree quilt", "Colorful Yggdrasil Tree Quilt Set", "Vintage Colorful Celtic Yggdrasil Quilt Set", "sunset-toned quilt with large leafy Tree of Life, Celtic border and stained glass style color blocks", "Tree of Life quilt, Celtic quilt set, Viking bedding", "EVERGREEN", CELTIC_REFS),
    265: ("colorful desert RV sunset quilt", "desert RV sunset quilt", "Desert RV Sunset Quilt Set", "Colorful Desert RV Sunset Quilt Set", "bright desert landscape quilt with RV camper, cactus plants, flowers, mountain layers and setting sun", "RV bedding, desert quilt set, camper bedding", "EVERGREEN", RV_REFS),
    266: ("colorful Tree of Life Yggdrasil quilt", "colorful Tree of Life quilt", "Colorful Tree of Life Quilt Set", "Colorful Yggdrasil Tree of Life Quilt Set", "black quilt with vibrant rainbow leaves, glowing sun, twisting roots and Tree of Life artwork", "Yggdrasil quilt, Celtic Tree of Life bedding, colorful quilt set", "EVERGREEN", CELTIC_REFS),
    267: ("black colorful Yggdrasil Tree of Life quilt", "Yggdrasil Tree of Life quilt", "Yggdrasil Tree of Life Quilt Set", "Black Yggdrasil Tree of Life Quilt Set", "black quilt with central Yggdrasil medallion, green Celtic border knots and orange corner knot details", "Celtic quilt set, Viking bedding, Tree of Life quilt", "EVERGREEN", CELTIC_REFS),
    268: ("custom couple portrait Valentine quilt", "custom couple photo quilt", "Custom Couple Photo Quilt", "Custom Couple Hugging Portrait Quilt", "custom photo quilt with close-up couple portrait, warm sepia tones and Happy Valentines Day script", "personalized photo quilt, Valentines Day photo blanket, couple gift quilt", "VALENTINES_DAY", PHOTO_REFS),
    269: ("Sophia bookworm cozy room blanket", "bookworm reading girl blanket", "Bookworm Reading Girl Blanket", "Sophia Bookworm Cozy Room Blanket", "bookworm throw blanket with personalized Sophia Bookworm text, girl reading in cozy chair and library room background", "personalized book lover blanket, reading blanket for girls, bookworm gift", "EVERGREEN", READING_REFS),
    270: ("dark haired girl book lover blanket", "dark haired book lover blanket", "Dark Haired Book Lover Blanket", "Dark Haired Girl Reading Book Blanket", "book lover throw blanket with dark-haired girl holding a book, stacked books, floral beige background and quote text", "reading blanket, book lover gift, personalized book blanket", "EVERGREEN", READING_REFS),
}


def qa_issues(pos, html_error):
    issues = [
        "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."
    ]
    if pos in (261, 262):
        issues.append("Christian scripture products need verse text, custom name/photo fields and faith-claim wording reviewed before import.")
        issues.append("Feature images mention fleece, washing, stitching and softness claims; confirm material, GSM and care details.")
    if pos in (263, 264, 266, 267):
        issues.append("Celtic/Yggdrasil Tree of Life quilt products are close variants; review motif separation, titles and canonical strategy.")
        issues.append("Verify quilt/bedspread construction, optional pillow shams and machine-washable claims.")
    if pos == 265:
        issues.append("RV desert quilt may overlap camping/desert bedding terms; verify product type, size chart and optional shams.")
    if pos == 268:
        issues.append("Custom photo Valentine product requires photo-upload QA, seasonal positioning review and component verification.")
    if pos in (269, 270):
        issues.append("Book lover blankets overlap batches 025-026; keep character, quote text, background and personalization fields distinct.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def product_context(pos):
    if pos in (261, 262):
        return "faith gifts, prayer spaces, bedrooms, church gifts or personalized encouragement gifts."
    if pos in (263, 264, 266, 267):
        return "Celtic bedrooms, Viking-inspired decor, spiritual home decor or Tree of Life gift bedding."
    if pos == 265:
        return "RV rooms, camper decor, southwestern bedrooms or desert travel gifts."
    if pos == 268:
        return "Valentine gifts, anniversary gifts, couple keepsakes or personalized bedroom decor."
    return "reading nooks, bedrooms, book clubs, librarian gifts or personalized reader gifts."


def alt_for_image(h1, pos, ipos):
    if pos in (263, 264, 265, 266, 267, 268):
        return {
            1: f"{h1} displayed on a bed in the main quilt mockup",
            2: f"Flat product or angled mockup of {h1}",
            3: f"Size chart or pillow sham detail for {h1}",
            4: f"Bedspread construction feature image for {h1}",
            5: f"Angled bedroom feature mockup of {h1}",
            6: f"Pillow sham and close-up artwork for {h1}",
            7: f"Fabric and product detail image for {h1}",
        }.get(ipos, f"Product detail image for {h1}")[:125]
    return {
        1: f"{h1} shown as a personalized throw blanket",
        2: f"Flat product or feature image for {h1}",
        3: f"Lifestyle or couch mockup for {h1}",
        4: f"Size guide or reading lifestyle image for {h1}",
        5: f"Fabric and fleece feature image for {h1}",
        6: f"Couch or close-up mockup for {h1}",
        7: f"Reading lifestyle mockup for {h1}",
        8: f"Additional lifestyle image for {h1}",
        9: f"Background or size option image for {h1}",
    }.get(ipos, f"Product detail image for {h1}")[:125]


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for {product_context(pos)}</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broad Christian, Celtic, RV, photo gift and book lover blanket terms for collection pages or strongest representatives.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, room decor or a personalized gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": qa_issues(pos, html_error),
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed bedding, the buyer wants a {kind} matching a specific faith, Celtic, RV, photo gift or reading motif.",
            "functional_motivation": "Find the right design, size, product type, personalization fields, care details and room fit.",
            "emotional_social_motivation": "Give a faith, couple or reader gift, decorate a themed room, or add Celtic/RV artwork to a bedroom.",
            "purchase_concerns": "Personalization spelling, photo quality, verse accuracy, product identity, material/care claims, size fit, included components and whether artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt_for_image(h1, pos, ipos), "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_027_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_027_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_027"},
        {"metric": "cumulative_products", "value": "270", "definition": "Total products included through batch_027"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[270:280]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_028",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
