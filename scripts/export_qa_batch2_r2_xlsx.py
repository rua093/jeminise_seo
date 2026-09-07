from pathlib import Path
import hashlib, json, zipfile
from openpyxl import load_workbook
import export_qa_batch1_xlsx as base

ROOT = Path(__file__).resolve().parents[1]
RUN = "20260907_123726"
base.QA_RUN_ID = RUN
base.DATA = ROOT / "seo_runs" / "jeminise.com" / "20260906_234129" / "qa" / RUN / "qa_workbook_payload.json"
base.OUTPUT = ROOT / "resutls" / "jeminise.com" / "20260906_234129" / "qa" / RUN / "SEO_QA_qa_batch_002_r2.xlsx"

if __name__ == "__main__":
    base.main()
    wb = load_workbook(base.OUTPUT, data_only=False)
    audit = {
        "path": str(base.OUTPUT), "sheet_names": wb.sheetnames,
        "expected_sheet_names": list(base.SHEETS),
        "row_counts": {ws.title: ws.max_row - 1 for ws in wb.worksheets},
        "freeze_panes": {ws.title: str(ws.freeze_panes) for ws in wb.worksheets},
        "auto_filters": {ws.title: ws.auto_filter.ref for ws in wb.worksheets},
        "formula_cells": sum(1 for ws in wb.worksheets for row in ws.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")),
        "hyperlink_cells": sum(1 for ws in wb.worksheets for row in ws.iter_rows() for cell in row if cell.hyperlink),
        "formula_error_tokens": sum(1 for ws in wb.worksheets for row in ws.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=") and any(t in cell.value for t in ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!"))),
        "xlsx_sha256": hashlib.sha256(base.OUTPUT.read_bytes()).hexdigest(),
        "source_sha256": hashlib.sha256((ROOT / "resutls" / "jeminise.com" / "20260906_234129" / "revisions" / "qa_batch_002_r2" / "SEO_Product_Optimization_qa_batch_002_r2.xlsx").read_bytes()).hexdigest(),
        "zip_bad_member": zipfile.ZipFile(base.OUTPUT).testzip(),
    }
    payload = json.loads(base.DATA.read_text(encoding="utf-8"))
    expected_counts = {"QA_Summary":13,"QA_Products":10,"QA_Criteria":110,"QA_Images":64,"QA_Issues":len(payload.get("QA_Issues", []))}
    audit["expected_row_counts"] = expected_counts
    audit["passed"] = audit["sheet_names"] == audit["expected_sheet_names"] and audit["row_counts"] == expected_counts and all(audit["freeze_panes"].values()) and all(audit["auto_filters"].values()) and audit["formula_error_tokens"] == 0 and audit["zip_bad_member"] is None
    (ROOT / "seo_runs" / "jeminise.com" / "20260906_234129" / "qa" / RUN / "spreadsheet_validation.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))
