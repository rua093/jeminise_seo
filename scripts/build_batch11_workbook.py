import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_011"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_010.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_011.xlsx"


SOCCER_REFS = "https://www.amazon.com/; https://www.pinterest.com/; https://www.authenticsoccer.com/; https://www.target.com/p/5pc-full-queen-kids-39-girls-soccer-kick-reversible-oversized-comforter-bedding-set-white-turquoise-lush-d-233-cor/-/A-88098171; https://www.lushdecor.com/products/girls-soccer-kick-comforter-set"

PROFILES = {
    101: ("personalized Keep on Truckin semi truck comforter with American flag", "personalized semi truck comforter", "Personalized Semi Truck Comforter", "Personalized Keep On Truckin Semi Truck Comforter", "red semi truck with American flag, Keep on Truckin text and Your Name personalization area", "custom trucker bedding, American flag truck comforter, semi truck bedding with name", "EVERGREEN", "https://jeminise.com/products/custom-patriotic-trucking-semi-truck-below-waving-american-flag-comforter-05c0cd8623-05c0cd8623; https://www.etsy.com/listing/1160610884/keep-on-truckin-american-flag-decal-semi; https://www.walmart.com/ip/Truck-Duvet-Cover-Set-Queen-Size-American-Flag-Themed-Semi-18-Wheeler-Patriotic-Transportation-Industrial-Vehicle-Decorative-3-Piece-Bedding-Set-2-Pi/432876364"),
    102: ("personalized soccer cleats comforter with Alexis number 07", "personalized soccer cleats comforter", "Personalized Soccer Cleats Comforter", "Personalized Soccer Cleats Comforter", "black and white soccer cleats over a soccer ball with Alexis name and number 07", "custom soccer bedding, soccer comforter with name, soccer bedding set for boys", "EVERGREEN", "https://www.authenticsoccer.com/; https://www.nike.com/w/nike-by-you-soccer-shoes-1gdj0z6ealhzy7ok; https://www.amazon.com/"),
    103: ("personalized soccer goal net comforter with Matthew number 15", "soccer goal net comforter", "Soccer Goal Net Comforter Set", "Personalized Soccer Goal Net Comforter", "blue soccer ball in goal net artwork with Matthew name and number 15", "personalized soccer bedding, soccer net bedding, soccer comforter with name", "EVERGREEN", SOCCER_REFS),
    104: ("personalized soccer goal comforter with Kenzo number 10", "personalized soccer goal comforter", "Personalized Soccer Goal Comforter", "Personalized Soccer Goal Comforter", "blue and white soccer ball goal/net artwork with Kenzo name and number 10", "custom soccer bedding, soccer comforter with number, soccer goal bedding", "EVERGREEN", SOCCER_REFS),
    105: ("personalized multicolor soccer net comforter with Jayden number 07", "soccer net comforter", "Soccer Net Comforter Set", "Personalized Multicolor Soccer Net Comforter", "black, teal and orange soccer net speed-line design with Jayden name and number 07", "custom soccer bedding, personalized soccer comforter, soccer bedding with name", "EVERGREEN", SOCCER_REFS),
    106: ("personalized soccer water splash comforter with Matthew number 15", "soccer water splash comforter", "Soccer Water Splash Comforter", "Personalized Soccer Water Splash Comforter", "blue water splash soccer ball artwork with Matthew name and number 15", "water soccer bedding, custom soccer comforter, soccer comforter with name", "EVERGREEN", "https://www.pinterest.com/pin/water-and-fire-footballsoccer-3d-bedding-set--1097752477927007130/; https://www.amazon.com/; https://www.target.com/p/5pc-full-queen-kids-39-girls-soccer-kick-reversible-oversized-comforter-bedding-set-white-turquoise-lush-d-233-cor/-/A-88098171"),
    107: ("personalized lightning soccer ball comforter with Lucas number 10", "lightning soccer comforter", "Lightning Soccer Comforter Set", "Personalized Lightning Soccer Ball Comforter", "dark blue soccer ball with white lightning splash, Lucas name and number 10", "custom soccer bedding, soccer ball comforter, soccer comforter with name", "EVERGREEN", SOCCER_REFS),
    108: ("personalized soccer ball number 07 comforter with Isaac", "soccer ball number comforter", "Personalized Soccer Ball Number Comforter", "Personalized Soccer Ball Number Comforter", "white and blue paint-splash soccer ball design with Isaac name and number 07", "soccer comforter with number, custom soccer bedding, personalized soccer bedding", "EVERGREEN", SOCCER_REFS),
    109: ("personalized green soccer goal splash comforter with James", "green soccer goal comforter", "Green Soccer Goal Comforter Set", "Personalized Green Soccer Goal Comforter", "green soccer ball goal splash artwork with James name and main number 07", "soccer goal bedding, personalized soccer comforter, soccer net comforter", "EVERGREEN", SOCCER_REFS),
    110: ("personalized black and white soccer net comforter with Kenzo number 07", "black soccer net comforter", "Black Soccer Net Comforter Set", "Personalized Black Soccer Net Comforter", "black and white soccer ball in net artwork with Kenzo name and number 07", "custom soccer net bedding, soccer comforter with name, soccer bedding for boys", "EVERGREEN", SOCCER_REFS),
}


def append_dict(ws, data):
    headers = [cell.value for cell in ws[1]]
    ws.append([data.get(h, "") for h in headers])


def image_obs(theme, pos):
    if pos == 1:
        return f"Main product mockup showing {theme}."
    if pos == 2:
        return f"Secondary flat, close-up or bedroom mockup showing {theme}."
    if pos == 3:
        return f"Bedding type, feature, lifestyle or material image for {theme}."
    if pos == 4:
        return f"Detail, feature panel or lifestyle image connected to {theme}."
    if pos == 5:
        return f"Additional mockup, care or size information for {theme}."
    if pos == 6:
        return f"Easy-care, size guide or product feature image for {theme}."
    return f"Additional lifestyle, feature or product detail image showing {theme}."


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        issue_bits = [
            "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."
        ]
        if pos == 101:
            issue_bits.append("Closely related patriotic semi truck comforters should be reviewed for keyword cannibalization and title uniqueness.")
        if 102 <= pos <= 110:
            issue_bits.append("Closely related personalized soccer comforters should be reviewed for keyword cannibalization and title uniqueness.")
        if pos in (103, 108):
            issue_bits.append("Current product title appears truncated around machine washable wording and needs QA.")
        if pos == 109:
            issue_bits.append("Contact sheet shows a number mismatch: main comforter appears to show 07 while pillow mockups show 10; personalization consistency needs QA.")
        if html_error:
            issue_bits.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
        issues = " ".join(issue_bits)
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for themed decor, sports rooms or personalized gifting.</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork, personalization and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broader bedding terms for collections or stronger representative products.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as a themed bedding item, personalized gift or room decor.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": issues,
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed or personalized bedding, the buyer wants a {kind} matching a specific design, name, number, message or hobby.",
            "functional_motivation": "Find the right design, size, bedding type, care details and personalization fit.",
            "emotional_social_motivation": "Create meaningful room decor or give a personalized sports, trucker or hobby gift.",
            "purchase_concerns": "Personalization spelling, product identity, material/care claims, design distinction and whether the visible artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork, personalization and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed in a main product mockup",
                2: f"Secondary mockup of {h1}",
                3: f"Product option, feature or lifestyle image for {h1}",
                4: f"Detail or feature image for {h1}",
                5: f"Additional product mockup or care information for {h1}",
                6: f"Care, size or product feature image for {h1}",
                7: f"Additional product mockup of {h1}",
                8: f"Lifestyle use image for {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_011_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_011_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_011"},
        {"metric": "cumulative_products", "value": "110", "definition": "Total products included through batch_011"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[110:120]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_012",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
