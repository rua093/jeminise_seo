import csv
import hashlib
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_010_r3"
QA_RUN_ID = "20260907_181000"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_010_r2" / "SEO_Product_Optimization_qa_batch_010_r2.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
CUSTOMIZER_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_audit.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_010_r3.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_010_r3.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    91: {
        "title": "Turquoise Wolf Quilt Set",
        "meta_title": "Turquoise Wolf Quilt Set",
        "meta_description": "Shop a wolf quilt set with turquoise circular artwork, matching sham mockups and printed size guidance for the quilt.",
        "primary": "turquoise wolf quilt set",
        "secondary": "boho wolf quilt set, wolf bedding set, Native inspired wolf quilt",
        "cluster": "turquoise circular wolf quilt set",
        "detail": "wolf head centered in turquoise circular boho artwork with feather-like ornaments and matching sham mockups",
        "intent_role": "Wolf quilt page separated by turquoise circular wolf artwork and sham-size panel.",
        "customizer": "An optional custom-name text field is available for shopper notes on this quilt design.",
    },
    92: {
        "title": "Neon Green Soccer Comforter Set",
        "meta_title": "Neon Green Soccer Comforter Set",
        "meta_description": "Shop a black soccer comforter set with neon green player artwork, JACKSON and 07 sample text, microfiber panels and size guide.",
        "primary": "neon green soccer comforter set",
        "secondary": "soccer player comforter set, black green soccer bedding, boys soccer comforter",
        "cluster": "black neon green soccer player comforter",
        "detail": "black comforter set with neon green soccer player artwork, glowing ball, JACKSON lettering and 07 sample number",
        "intent_role": "Soccer page separated by neon green player artwork; personalization claim removed after Sang critical finding.",
        "customizer": "JACKSON and 07 are shown as sample artwork in the gallery; enter only the information requested by the visible field.",
    },
    93: {
        "title": "Blue Semi Truck Flag Comforter",
        "meta_title": "Blue Semi Truck Flag Comforter",
        "meta_description": "Shop a blue semi truck comforter with distressed American flag artwork, YOUR NAME sample text and bedding size panels.",
        "primary": "blue semi truck flag comforter",
        "secondary": "American flag truck comforter, blue trucker bedding, semi truck comforter set",
        "cluster": "blue semi truck on distressed American flag comforter",
        "detail": "blue semi truck over a distressed American flag with YOUR NAME sample text across the bedding mockup",
        "intent_role": "Semi-truck page separated by blue cab, distressed flag background and horizontal name bar.",
        "customizer": "A required Enter Name field supports up to 35 characters, with a required product-confirmation choice.",
    },
    94: {
        "title": "Silver Semi Truck Flag Comforter",
        "meta_title": "Silver Semi Truck Flag Comforter",
        "meta_description": "Shop a silver semi truck comforter with waving American flag graphics, large YOUR NAME text and microfiber feature panels.",
        "primary": "silver semi truck flag comforter",
        "secondary": "patriotic semi truck comforter, silver trucker bedding, waving flag truck comforter",
        "cluster": "silver semi truck with waving American flag comforter",
        "detail": "silver semi truck with waving American flag graphics and large YOUR NAME sample text on a dark lower panel",
        "intent_role": "Semi-truck page separated by silver cab, waving flag graphics and large lower name placement.",
        "customizer": "A required Enter Name field supports up to 35 characters, with a required product-confirmation choice.",
    },
    95: {
        "title": "Red Semi Truck Sunset Comforter",
        "meta_title": "Red Semi Truck Sunset Comforter",
        "meta_description": "Shop a red semi truck comforter with sunset road artwork, waving American flag stripes and YOUR NAME sample text.",
        "primary": "red semi truck sunset comforter",
        "secondary": "red trucker comforter, patriotic trucking bedding, American flag semi truck bedding",
        "cluster": "red semi truck below waving flag at sunset comforter",
        "detail": "red semi truck below waving American flag stripes against a sunset road scene with YOUR NAME sample text",
        "intent_role": "Semi-truck page separated by red cab, sunset background and flag stripes above the truck.",
        "customizer": "A required Enter Name field supports up to 35 characters, with a required product-confirmation choice.",
    },
    96: {
        "title": "Custom Photo Collage Quilt Set",
        "meta_title": "Custom Photo Collage Quilt Set",
        "meta_description": "Shop a photo collage quilt set with multiple image panels, central Your Text Here area, optional shams and size chart.",
        "primary": "custom photo collage quilt set",
        "secondary": "photo collage quilt, custom text quilt set, personalized photo quilt",
        "cluster": "photo collage quilt with central text panel",
        "detail": "photo collage quilt template with multiple image panels, a central Your Text Here message area and optional sham mockups",
        "intent_role": "Photo quilt page separated by multi-photo collage layout and central custom text panel.",
        "customizer": "Custom quilt surfaces and background color choices are available for the photo collage layout.",
    },
    97: {
        "title": "God Says I Am Butterfly Blanket",
        "meta_title": "God Says I Am Butterfly Blanket",
        "meta_description": "Shop a purple God Says I Am blanket with Elizabeth sample name, butterfly artwork and Bible affirmation blocks.",
        "primary": "God Says I Am butterfly blanket",
        "secondary": "purple Christian butterfly blanket, Elizabeth affirmation blanket, Bible verse butterfly blanket",
        "cluster": "God Says I Am Elizabeth butterfly blanket",
        "detail": "purple God Says I Am design with Elizabeth sample name, butterfly artwork and Bible affirmation blocks",
        "intent_role": "Christian blanket page separated by God Says I Am wording and Elizabeth vertical name layout.",
        "customizer": "A required Custom Name field and optional Message to Seller field are available.",
    },
    98: {
        "title": "God Is Within Her Butterfly Blanket",
        "meta_title": "God Is Within Her Butterfly Blanket",
        "meta_description": "Shop a purple God Is Within Her blanket with Isabella sample name, Psalm 46:5 text and butterfly vine artwork.",
        "primary": "God Is Within Her butterfly blanket",
        "secondary": "Isabella Christian blanket, Psalm 46:5 butterfly blanket, purple faith blanket for her",
        "cluster": "God Is Within Her Isabella butterfly blanket",
        "detail": "purple God Is Within Her design with Isabella sample name, Psalm 46:5 text and butterfly vine artwork",
        "intent_role": "Christian blanket page separated by God Is Within Her heading, Psalm 46:5 and Isabella acrostic layout.",
        "customizer": "A required Custom Name field and optional Message to Seller field are available.",
    },
    99: {
        "title": "Blessed Is She Purple Cross Blanket",
        "meta_title": "Blessed Is She Purple Cross Blanket",
        "meta_description": "Shop a Blessed is she purple cross blanket with roses, butterflies, Amabel sample name and soft fleece feature panels.",
        "primary": "Blessed is she purple cross blanket",
        "secondary": "purple rose cross blanket, Christian cross blanket for her, Amabel faith blanket",
        "cluster": "Blessed is she purple rose cross blanket",
        "detail": "dark purple Blessed is she design with glowing cross, purple roses, butterflies and Amabel sample name",
        "intent_role": "Christian cross page separated by Blessed is she wording, roses and dark purple cross artwork.",
        "customizer": "Amabel appears as sample text in the artwork; follow the visible product options before checkout.",
    },
    100: {
        "title": "Estella Scripture Collage Blanket",
        "meta_title": "Estella Scripture Collage Blanket",
        "meta_description": "Shop a purple scripture collage blanket with Estella sample name, draped wooden cross, butterflies and floral accents.",
        "primary": "Estella scripture collage blanket",
        "secondary": "purple scripture blanket, Christian collage blanket, floral cross Bible verse blanket",
        "cluster": "Estella purple scripture collage blanket",
        "detail": "purple scripture collage design with Estella sample name, draped wooden cross, butterflies and floral accents",
        "intent_role": "Christian blanket page separated by scripture-collage layout, Estella name panel and draped cross motif.",
        "customizer": "A required Custom Name field supports up to 12 characters, with an optional Message to Seller field.",
    },
}


IMAGE_DETAILS = {
    91: {
        1: ("Bedroom mockup of turquoise wolf quilt set with matching wolf shams and circular boho feather artwork.", "Turquoise wolf quilt set with matching shams"),
        2: ("Close-up of wolf face artwork with turquoise eyes, line shading and ornamental forehead detail.", "Close-up of turquoise wolf face artwork"),
        3: ("Single sham mockup showing the centered wolf head and turquoise circular border.", "Wolf sham with turquoise circular artwork"),
        4: ("Bedroom lifestyle mockup of the wolf quilt on a bed with matching shams.", "Wolf quilt bedding set in bedroom mockup"),
        5: ("Size panel listing one premium quilt and optional two standard shams.", "Wolf quilt set size and optional sham panel"),
    },
    92: {
        1: ("Bed mockup of black comforter set with neon green soccer player, JACKSON text and 07 number.", "Neon green soccer comforter with JACKSON 07"),
        2: ("Feature panel showing neon soccer bedding details, microfiber label and 3D printed pattern callouts.", "Neon soccer comforter microfiber feature panel"),
        3: ("Comparison panel for duvet cover set and comforter set bedding types.", "Duvet cover and comforter set comparison panel"),
        4: ("Close-up corner showing neon soccer artwork folded back over white microfiber filling.", "Neon soccer bedding corner with microfiber filling"),
        5: ("Bedroom mockup of neon soccer comforter with JACKSON lettering and matching 07 shams.", "Neon soccer bedding set with matching shams"),
        6: ("Low-angle bed mockup with neon soccer design and icons for soft, lightweight, durable and breathable.", "Neon soccer comforter with bedding feature icons"),
        7: ("Stress-free easy-care panel listing wrinkle-free, stain-proof, anti-pilling and wash durability.", "Easy care feature panel for soccer bedding"),
        8: ("Size dimension chart for twin, full, queen and king bedding dimensions.", "Soccer comforter size dimension chart"),
    },
    93: {
        1: ("Bedroom mockup of blue semi truck comforter on distressed American flag with YOUR NAME text.", "Blue semi truck flag comforter with YOUR NAME"),
        2: ("Bedroom mockup of the blue truck flag comforter with matching shams and folded white top sheet.", "Blue truck flag bedding set with matching shams"),
        3: ("Comparison panel explaining duvet cover set versus comforter set.", "Duvet cover and comforter comparison panel"),
        4: ("Microfiber feature panel with blue semi truck pillow mockup and 3D printed pattern callouts.", "Blue semi truck comforter microfiber feature panel"),
        5: ("Low-angle blue truck bedding mockup with icons for soft, lightweight, durable and breathable.", "Blue semi truck bedding with feature icons"),
        6: ("Stress-free easy-care panel with stacked white bedding and care feature icons.", "Easy care feature panel for truck bedding"),
        7: ("Size dimension chart showing twin, full, queen and king measurements.", "Semi truck comforter size dimension chart"),
    },
    94: {
        1: ("Bedroom mockup of silver semi truck comforter with waving American flag and large YOUR NAME text.", "Silver semi truck flag comforter with YOUR NAME"),
        2: ("Bedroom mockup of silver truck flag comforter with matching pillow shams.", "Silver truck flag bedding set with shams"),
        3: ("Comparison panel explaining duvet cover set and comforter set bedding options.", "Duvet cover and comforter set comparison panel"),
        4: ("Microfiber feature panel with silver truck pillow mockup and close-up detail circles.", "Silver semi truck comforter feature panel"),
        5: ("Low-angle mockup of silver truck bedding with soft, lightweight, durable and breathable icons.", "Silver truck bedding with feature icons"),
        6: ("Stress-free easy-care panel listing wrinkle-free, stain-proof, anti-pilling and wash durability.", "Easy care feature panel for truck bedding"),
        7: ("Size dimension chart showing twin, full, queen and king bedding measurements.", "Semi truck comforter size dimension chart"),
    },
    95: {
        1: ("Bedroom mockup of red semi truck comforter below waving American flag stripes and sunset sky.", "Red semi truck sunset flag comforter"),
        2: ("Bedroom mockup of red truck sunset comforter with matching pillow shams.", "Red semi truck bedding set with shams"),
        3: ("Comparison panel explaining duvet cover set and comforter set options.", "Duvet cover and comforter set comparison panel"),
        4: ("Microfiber feature panel with red semi truck artwork and 3D printed pattern callouts.", "Red semi truck comforter feature panel"),
        5: ("Low-angle mockup of red truck bedding with soft, lightweight, durable and breathable icons.", "Red semi truck bedding with feature icons"),
        6: ("Stress-free easy-care panel with stacked white bedding and care icons.", "Easy care feature panel for truck bedding"),
        7: ("Size dimension chart showing twin, full, queen and king bedding measurements.", "Semi truck comforter size dimension chart"),
    },
    96: {
        1: ("Flat quilt mockup with custom photo collage panels and central Your Text Here message area.", "Custom photo collage quilt with Your Text Here"),
        2: ("Bedroom mockup of photo collage quilt set with matching shams and central text panel.", "Photo collage quilt bedding set with shams"),
        3: ("Bedspread features panel showing woven top, microfiber filling, back microfiber and care icons.", "Photo quilt bedspread features panel"),
        4: ("Bedroom lifestyle mockup of photo collage quilt with premium printed craft note.", "Photo collage quilt bedroom lifestyle mockup"),
        5: ("Close-up bedding mockup with lightweight, soft, anti-pill, anti-static icons and optional two pillow shams.", "Photo collage quilt with optional pillow shams"),
        6: ("High-quality fabric panel showing quilt layers and bedding use cases.", "Photo quilt fabric and use case feature panel"),
        7: ("Premium quilt set size chart with throw, twin, full, queen and king measurements.", "Photo collage quilt set size chart"),
    },
    97: {
        1: ("Flat blanket mockup with purple God Says I Am text, Elizabeth name and butterfly artwork.", "God Says I Am butterfly blanket with Elizabeth"),
        2: ("Lifestyle mockup of a person holding the purple God Says I Am Elizabeth butterfly blanket.", "Elizabeth God Says I Am blanket held up"),
        3: ("Feature panel showing fluffy soft layers, high quality stitching, no pilling, no shedding and smooth texture.", "God Says I Am blanket feature panel"),
        4: ("Collage of fabric folds, butterfly artwork, affirmation blocks and white fleece texture.", "Close-up collage for God Says I Am blanket"),
        5: ("Sofa mockup of Elizabeth God Says I Am blanket draped over a tan couch.", "Elizabeth butterfly affirmation blanket on sofa"),
        6: ("Bench mockup of the purple God Says I Am blanket with butterfly and affirmation text blocks.", "Purple God Says I Am blanket on bench"),
        7: ("Reading lifestyle image with God Says I Am blanket across a bed.", "God Says I Am blanket used while reading"),
        8: ("Blanket size chart showing 40x30, 50x40, 60x50 and 80x60 options.", "God Says I Am blanket size chart"),
    },
    98: {
        1: ("Person holding purple God Is Within Her blanket with Isabella name and butterfly vine artwork.", "God Is Within Her butterfly blanket with Isabella"),
        2: ("Bed mockup of Isabella God Is Within Her blanket with Psalm 46:5 text.", "Isabella God Is Within Her blanket on bed"),
        3: ("Reading lifestyle image with partial God Is Within Her blanket, open book and coffee.", "God Is Within Her butterfly blanket reading mockup"),
        4: ("Sofa mockup of God Is Within Her blanket with butterfly vine and Isabella acrostic blocks.", "God Is Within Her Isabella blanket on sofa"),
        5: ("Fabric collage showing purple blanket folds, affirmation text, butterfly art and white fleece.", "Close-up collage for God Is Within Her blanket"),
        6: ("Use and size panel showing 80 by 60 blanket and sofa, office, bed and travel use icons.", "God Is Within Her blanket size and use panel"),
        7: ("Feature panel showing machine washable, soft and warm, dense stitching and 260GSM fleece blanket.", "God Is Within Her fleece feature panel"),
        8: ("Close lifestyle image of God Is Within Her blanket beside an open book and mug.", "God Is Within Her blanket reading close-up"),
    },
    99: {
        1: ("Person holding dark purple Blessed is she blanket with cross, roses, butterflies and Amabel name.", "Blessed is she cross blanket with Amabel"),
        2: ("Sofa mockup of Blessed is she cross blanket draped over a beige couch.", "Blessed is she purple cross blanket on sofa"),
        3: ("Armchair mockup showing the Blessed is she text, glowing cross and purple roses.", "Purple Blessed is she cross blanket on chair"),
        4: ("Reading close-up of Blessed is she blanket with cross detail, open book and mug.", "Blessed is she cross blanket reading close-up"),
        5: ("Fabric collage showing printed verse text, cross detail, purple roses and white fleece backing.", "Close-up collage for Blessed is she blanket"),
        6: ("Feature panel showing machine washable, soft and warm, dense stitching and 260GSM fleece blanket.", "Blessed is she fleece feature panel"),
        7: ("Reading lifestyle image with Blessed is she blanket, open book, mug and small dog.", "Blessed is she blanket reading lifestyle mockup"),
        8: ("Use and size panel showing 80 by 60 blanket, text can be changed badge and personalization badge.", "Blessed is she blanket size and use panel"),
    },
    100: {
        1: ("Person holding purple scripture collage blanket with Estella name, draped cross and Bible verse blocks.", "Estella scripture collage blanket with cross"),
        2: ("Bed mockup of Estella scripture collage blanket with purple panels and draped cross artwork.", "Estella scripture collage blanket on bed"),
        3: ("Reading lifestyle image with scripture collage blanket, open book and small dog.", "Scripture collage blanket reading lifestyle mockup"),
        4: ("Sofa mockup of Estella scripture collage blanket with cross, butterflies and verse blocks.", "Estella scripture collage blanket on sofa"),
        5: ("Close reading view of purple scripture blanket with mug, open book and draped cross detail.", "Estella scripture blanket reading close-up"),
        6: ("Feature panel showing machine washable, soft and warm, dense stitching and 260GSM fleece blanket.", "Scripture collage fleece feature panel"),
        7: ("Fabric collage showing purple verse text, draped cross detail, butterfly artwork and white fleece.", "Close-up collage for scripture collage blanket"),
        8: ("Use and size panel showing 80 by 60 blanket and travel, sleeping, reading, relaxing and camping icons.", "Scripture collage blanket size and use panel"),
    },
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def load_qa():
    data = json.loads(QA_DATASET.read_text(encoding="utf-8"))
    products = data["QA_Products"]
    customizer = json.loads(CUSTOMIZER_AUDIT.read_text(encoding="utf-8"))
    return {
        "data": data,
        "product_keys": [item["product_key"] for item in products],
        "handle_by_key": {item["product_key"]: item["handle"] for item in products},
        "pos_by_key": {item["product_key"]: int(item["inventory_position"]) for item in products},
        "customizer_by_pos": {int(item["inventory_position"]): item for item in customizer},
    }


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Body (HTML)": first.get("Body (HTML)", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    return admin


def option_text(admin_row):
    values = [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]]
    values = [v for v in values if v]
    return ", ".join(values) if values else "available size options"


def customizer_sentence(audit):
    excerpt = audit.get("config_excerpt", "")
    text_matches = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"minLength":\s*(\d+),\s*"maxLength":\s*(\d+)',
        excerpt,
    )
    option_matches = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"defaultOptionId"',
        excerpt,
    )
    parts = []
    for label, required, min_length, max_length in text_matches:
        requirement = "required" if required == "true" else "optional"
        parts.append(f"Personalization uses a {requirement} {label} field, {min_length}-{max_length} characters.")
    useful_options = [
        label for label, _ in option_matches
        if label and "confirmation" not in label.lower()
    ]
    if useful_options:
        parts.append(f"Additional visible customization option: {', '.join(useful_options)}.")
    if not text_matches:
        parts.append("No shopper text-entry field is described for this product.")
    labels = " ".join(label for label, *_ in text_matches).lower()
    if "number" in labels:
        parts.append("Names and numbers visible in mockups are sample artwork unless the matching input is entered.")
    elif "name" in labels:
        parts.append("Visible names in mockups are sample artwork unless the matching input is entered.")
    else:
        parts.append("Visible sample text or photos in mockups are part of the displayed design unless a matching input is available.")
    return " ".join(parts)


def description(pos, admin_row, ctext):
    item = PRODUCT_UPDATES[pos]
    product_word = "quilt set" if "Quilt" in (admin_row["Type"] or item["title"]) else "comforter set" if "Comforter" in (admin_row["Type"] or item["title"]) else "blanket"
    return (
        f"<p>{item['title']} features {item['detail']}. "
        f"It gives shoppers a specific {product_word} design instead of a generic themed print.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork: {item['detail']}.</li>"
        "<li>Gallery images show main mockups plus close-up, feature, care, use or size panels where present.</li></ul>"
        "<h3>Options and Fit</h3>"
        f"<ul><li>Available option groups: {option_text(admin_row)}.</li>"
        f"<li>{ctext}</li>"
        "<li>Check the selected size and any visible customization field before checkout.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    qa = load_qa()
    admin = load_admin(set(qa["handle_by_key"].values()))
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ctext = customizer_sentence(qa["customizer_by_pos"].get(pos, {}))
        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or "Blanket"
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, ctext)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"US English buyer intent targets {item['cluster']}; no volume or ranking claim is made."
        )
        ws.cell(row_num, idx["keyword_evidence_level"]).value = "SERP_ONLY"
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R3 deep recheck after Sang QA batch 010: used products_export_1.csv sha256:{admin_hash}; "
            "removed internal workflow copy, corrected or removed unsafe customization claims, "
            "split similar truck and Christian blanket keyword intents, rewrote customer-facing HTML, "
            "and refreshed image observations/alts from inspected contact sheets. Still NEEDS_REVIEW."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    seen_alts = set()
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        image_number = int(ws.cell(row_num, idx["image_number"]).value)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation, alt = IMAGE_DETAILS[pos][image_number]
        if alt in seen_alts:
            alt = f"{alt} for {PRODUCT_UPDATES[pos]['cluster']}"
        seen_alts.add(alt)
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = alt[:125]
        ws.cell(row_num, idx["alt_action"]).value = "SET"
        ws.cell(row_num, idx["viewed_status"]).value = "VIEWED"
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R3: observation and alt rewritten from inspected contact sheet; admin image URL and current alt matched from Shopify CSV where available."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "PRODUCT"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = (
            "Sang QA batch 010 + products_export_1.csv admin baseline + customizer_audit.json + contact-sheet inspection"
        )
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = (
            "R3 separates nearby products by visible motif, product type, sample wording and verified customization proof level."
        )
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r3"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as a specific product page."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm the exact artwork, product type, size choices, "
            "and only the customization options that are actually supported."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = (
            "Confirm product type, visible motif, size/care panels, supported custom fields and image accuracy before purchase."
        )
        ws.cell(row_num, idx["emotional_social_motivation"]).value = (
            "Choose a gift or themed bedding item without confusing it with another similar design in the batch."
        )
        ws.cell(row_num, idx["purchase_concerns"]).value = (
            "Avoid unsupported photo, name, number or material claims; confirm exact size and product type."
        )
        ws.cell(row_num, idx["source_refs"]).value = (
            f"products_export_1.csv sha256:{admin_hash}; Sang QA {QA_RUN_ID}; customizer_audit.json; contact sheets 91-100"
        )
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R3_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = (
            "No Search Console, paid keyword volume or internal site-search export supplied; proof remains SERP-only."
        )
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    pk_by_pos = {qa["pos_by_key"][pk]: pk for pk in qa["product_keys"]}
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = pk_by_pos[pos]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = (
            f"Original storefront/product evidence; Sang QA {QA_RUN_ID}; products_export_1.csv sha256:{admin_hash}; "
            "customizer_audit.json; inspected contact sheets"
        )
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={option_text(admin_row)}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={item['detail']}; "
            f"customizer={customizer_sentence(qa['customizer_by_pos'].get(pos, {}))}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "Sang QA issues addressed in r3; independent QA recheck still required."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R3 uses inspected contact sheets, customizer audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_010_r3_revision", "r3", "Separate 10-product revision after Sang QA; source workbook was not modified."),
        ("qa_batch_010_r3_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_010_r3_scope", "inventory positions 91-100", "No products outside qa_batch_010 were revised."),
        ("qa_batch_010_r3_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_010_r3_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_010_r3_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_010 only; inventory positions 91-100",
        "revision": "r3",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Hold for combined QA handoff after revisions through qa_batch_020, per user request.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_010_r3",
                "",
                "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r2: `{SOURCE.relative_to(ROOT)}`",
                f"- QA của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r3: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Sửa trọng yếu: bỏ nội dung nội bộ, gỡ claim personalization rủi ro ở sản phẩm 92, tách keyword intent cho cụm semi-truck và Christian blanket, viết lại mô tả/alt theo contact sheets.",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Bước tiếp theo: giữ lại để gửi Sang QA chung sau khi revision đến `qa_batch_020`, theo yêu cầu người dùng.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
