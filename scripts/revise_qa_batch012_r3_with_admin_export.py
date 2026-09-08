import csv
import hashlib
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_012_r3"
QA_RUN_ID = "20260907_181200"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_012_r2" / "SEO_Product_Optimization_qa_batch_012_r2.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
CUSTOMIZER_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_audit.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_012_r3.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_012_r3.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    111: {
        "title": "Daniel Soccer Ball Comforter Set",
        "meta_title": "Daniel Soccer Ball Comforter Set",
        "meta_description": "Shop a blue soccer comforter set with Daniel sample text, number 10 ball artwork, water-splash styling and care panels.",
        "primary": "Daniel soccer ball comforter set",
        "secondary": "blue soccer comforter set, soccer number 10 bedding, water splash soccer bedding",
        "cluster": "Daniel blue water-splash soccer comforter",
        "detail": "blue water-splash soccer ball artwork with Daniel sample text and number 10",
        "intent_role": "Soccer page separated by Daniel sample text, number 10 ball artwork and blue water-splash styling.",
        "customizer": "Daniel and 10 are treated as sample artwork; the visible field is a general Customize Your Item text box.",
    },
    112: {
        "title": "Tyler Soccer Goal Comforter Set",
        "meta_title": "Tyler Soccer Goal Comforter Set",
        "meta_description": "Shop a blue soccer goal comforter set with Tyler sample text, number 10 ball, goal net artwork and size guide.",
        "primary": "Tyler soccer goal comforter set",
        "secondary": "blue soccer goal bedding, soccer net comforter set, soccer number 10 comforter",
        "cluster": "Tyler blue soccer goal net comforter",
        "detail": "blue soccer ball in goal net artwork with Tyler sample text and number 10",
        "intent_role": "Soccer page separated by Tyler sample text, number 10 ball and goal-net artwork.",
        "customizer": "Tyler and 10 are treated as sample artwork; the visible field is a general Customize Your Item text box.",
    },
    113: {
        "title": "Matthew Paint Splatter Soccer Comforter",
        "meta_title": "Matthew Paint Splatter Soccer Comforter",
        "meta_description": "Shop a white soccer comforter set with multicolor paint splatter, Matthew sample text, number 15 shams and care panels.",
        "primary": "Matthew paint splatter soccer comforter",
        "secondary": "paint splatter soccer bedding, multicolor soccer comforter, soccer ball bedding set",
        "cluster": "Matthew multicolor paint-splatter soccer comforter",
        "detail": "white soccer ball artwork with multicolor paint splatter, Matthew sample text and number 15 shams",
        "intent_role": "Soccer page separated by Matthew sample text, number 15 and multicolor paint-splatter styling.",
        "customizer": "Matthew and 15 are treated as sample artwork; the visible field is a general Customize Your Item text box.",
    },
    114: {
        "title": "Custom Softball Name Number Comforter",
        "meta_title": "Custom Softball Name Number Comforter",
        "meta_description": "Shop a yellow, black and gray softball comforter with required name and number fields, batter artwork and included pillowcases.",
        "primary": "custom softball name number comforter",
        "secondary": "softball patchwork bedding, girls softball comforter set, personalized softball comforter",
        "cluster": "custom yellow black softball comforter",
        "detail": "yellow, black and gray softball patchwork artwork with batter silhouettes and matching pillowcases",
        "intent_role": "Softball page separated by confirmed required Custom Your Name and Custom Your Number fields.",
        "customizer": "Required Custom Your Name and Custom Your Number fields are present in the customizer audit.",
    },
    115: {
        "title": "Kevin Football Flag Comforter",
        "meta_title": "Kevin Football Flag Comforter",
        "meta_description": "Shop a football comforter with distressed American flag background, Kevin sample text, number 20 shams and size guide.",
        "primary": "Kevin football flag comforter",
        "secondary": "football American flag comforter, custom football number bedding, patriotic football bedding",
        "cluster": "Kevin football American flag comforter",
        "detail": "football artwork over a distressed American flag background with Kevin sample text and number 20",
        "intent_role": "Football page separated by confirmed Enter Name and Enter Number fields plus patriotic flag artwork.",
        "customizer": "Required confirmation plus Enter Name up to 25 characters and Enter Number up to 5 characters are present.",
    },
    116: {
        "title": "Blue Geometric Semi Truck Comforter",
        "meta_title": "Blue Geometric Semi Truck Comforter",
        "meta_description": "Shop a blue semi truck comforter with dark geometric graphics, YOUR NAME sample panel, microfiber callouts and size guide.",
        "primary": "blue geometric semi truck comforter",
        "secondary": "geometric trucker bedding, blue semi truck bedding, truck comforter with name",
        "cluster": "blue geometric semi truck comforter",
        "detail": "blue semi truck artwork on a dark geometric background with YOUR NAME sample panel",
        "intent_role": "Truck page separated by blue cab artwork, geometric background and confirmed Enter Name field.",
        "customizer": "A required Enter Name field supports up to 35 characters, with a required product-confirmation choice.",
    },
    117: {
        "title": "Colorful Tree of Life Quilt Set",
        "meta_title": "Colorful Tree of Life Quilt Set",
        "meta_description": "Shop a colorful Tree of Life quilt set with rainbow leaves, flowing roots, matching shams, fabric panels and size chart.",
        "primary": "colorful tree of life quilt set",
        "secondary": "rainbow tree of life bedding, colorful Yggdrasil quilt, tree of life quilt set",
        "cluster": "colorful rainbow Tree of Life quilt set",
        "detail": "colorful Tree of Life artwork with rainbow leaves, flowing roots and matching shams",
        "intent_role": "Quilt page separated by rainbow leaf palette and flowing root Tree of Life artwork.",
        "customizer": "The visible field is a general Customize Your Quilt text box; no name or number field is claimed.",
    },
    118: {
        "title": "Celtic Yggdrasil Tree Quilt Set",
        "meta_title": "Celtic Yggdrasil Tree Quilt Set",
        "meta_description": "Shop a Celtic Yggdrasil quilt set with green tree artwork, oval knotwork frame, gold accents, matching shams and size chart.",
        "primary": "Celtic Yggdrasil tree quilt set",
        "secondary": "green tree of life quilt, Celtic tree bedding, Yggdrasil bedding set",
        "cluster": "green Celtic Yggdrasil Tree of Life quilt set",
        "detail": "green Yggdrasil Tree of Life artwork with oval knotwork frame, gold accents and matching shams",
        "intent_role": "Quilt page separated by Celtic/Yggdrasil framing, green palette and gold accent artwork.",
        "customizer": "The visible field is a general Customize Your Quilt text box; no name or number field is claimed.",
    },
    119: {
        "title": "Red Semi Truck Stone Wall Comforter",
        "meta_title": "Red Semi Truck Stone Wall Comforter",
        "meta_description": "Shop a red semi truck comforter with stone wall break-through artwork, YOUR NAME sample panel, feature callouts and size guide.",
        "primary": "red semi truck stone wall comforter",
        "secondary": "custom red trucker bedding, semi truck wall comforter, truck comforter with name",
        "cluster": "red semi truck breaking through stone wall comforter",
        "detail": "red semi truck breaking through a gray stone wall with YOUR NAME sample panel",
        "intent_role": "Truck page separated by red cab color, stone wall break-through artwork and confirmed Enter Name field.",
        "customizer": "A required Enter Name field supports up to 35 characters, with a required product-confirmation choice.",
    },
    120: {
        "title": "Black Semi Truck Stone Wall Comforter",
        "meta_title": "Black Semi Truck Stone Wall Comforter",
        "meta_description": "Shop a black semi truck comforter with stone wall break-through artwork, YOUR NAME sample panel, feature callouts and size guide.",
        "primary": "black semi truck stone wall comforter",
        "secondary": "custom black trucker bedding, black truck comforter, truck stone wall bedding",
        "cluster": "black semi truck breaking through stone wall comforter",
        "detail": "black semi truck breaking through a gray stone wall with YOUR NAME sample panel",
        "intent_role": "Truck page separated by black cab color, stone wall break-through artwork and confirmed Enter Name field.",
        "customizer": "A required Enter Name field supports up to 35 characters, with a required product-confirmation choice.",
    },
}


COMMON_DUVET_PANEL = ("Comparison panel explaining duvet cover set and comforter set bedding options.", "Duvet cover and comforter set comparison panel")
COMMON_CARE = ("Stress-free easy-care panel listing wrinkle-free, stain-proof, anti-pilling and wash durability.", "Easy care panel for comforter bedding")
COMMON_SIZE = ("Size dimension chart showing twin, full, queen and king bedding measurements.", "Comforter set size dimension chart")
COMMON_QUILT_SIZE = ("Premium quilt set size chart showing available bed sizes and measurements.", "Tree of Life quilt set size chart")
COMMON_QUILT_LAYER = ("Bedspread feature panel showing layered quilt construction and soft-touch callouts.", "Quilt set construction feature panel")


IMAGE_DETAILS = {
    111: {
        1: ("Bedroom mockup of blue water-splash soccer comforter with Daniel text, number 10 ball and matching shams.", "Daniel soccer ball comforter with number 10"),
        2: ("Feature panel showing Daniel soccer bedding with fabric and 3D printed pattern callouts.", "Daniel soccer comforter feature panel"),
        3: ("Pillow sham mockup showing Daniel text and number 10 soccer ball artwork.", "Daniel soccer pillow shams with 10"),
        4: ("Low-angle Daniel soccer bedding mockup with soft, lightweight, durable and breathable icons.", "Daniel soccer bedding with feature icons"),
        5: COMMON_DUVET_PANEL,
        6: ("Bedroom mockup of Daniel blue soccer comforter set with water-splash ball artwork.", "Daniel blue soccer comforter set on bed"),
        7: ("Close-up corner showing blue soccer artwork under white microfiber filling.", "Daniel soccer comforter microfiber close-up"),
        8: COMMON_CARE,
    },
    112: {
        1: ("Bedroom mockup of blue soccer goal comforter with Tyler text, number 10 ball and matching shams.", "Tyler soccer goal comforter with number 10"),
        2: ("Bedroom mockup of Tyler soccer goal bedding with goal-net artwork and folded top sheet.", "Tyler soccer goal bedding set on bed"),
        3: ("Feature panel showing Tyler soccer goal bedding with microfiber and 3D printed pattern callouts.", "Tyler soccer goal comforter feature panel"),
        4: ("Pillow sham mockup showing Tyler text and number 10 soccer ball artwork.", "Tyler soccer pillow shams with 10"),
        5: COMMON_DUVET_PANEL,
        6: ("Low-angle Tyler soccer goal bedding mockup with soft, lightweight, durable and breathable icons.", "Tyler soccer bedding with feature icons"),
        7: ("Close-up corner showing blue soccer goal artwork under white microfiber filling.", "Tyler soccer goal comforter close-up"),
        8: COMMON_SIZE,
    },
    113: {
        1: ("Bedroom mockup of white soccer comforter with Matthew text, number 15 and multicolor paint splatter.", "Matthew paint splatter soccer comforter"),
        2: ("Low-angle paint-splatter soccer bedding mockup with feature icons.", "Paint splatter soccer bedding with feature icons"),
        3: ("Feature panel showing Matthew soccer bedding with microfiber and 3D printed pattern callouts.", "Matthew paint splatter comforter feature panel"),
        4: ("Pillow sham mockup showing Matthew text and number 15 soccer ball artwork.", "Matthew soccer pillow shams with 15"),
        5: COMMON_DUVET_PANEL,
        6: ("Bedroom mockup of Matthew paint-splatter soccer bedding with matching shams.", "Matthew paint splatter soccer bedding set"),
        7: COMMON_SIZE,
        8: COMMON_CARE,
    },
    114: {
        1: ("Bedroom mockup of yellow, black and gray softball patchwork comforter with matching pillowcases.", "Yellow black softball comforter set"),
        2: ("Bedroom mockup of softball comforter with yellow patchwork panels and batter silhouettes.", "Softball patchwork comforter on bed"),
        3: ("Pillowcase mockup showing yellow softball batter silhouettes and patchwork background.", "Softball batter pillowcases"),
        4: ("Folded comforter mockup showing yellow softball patchwork artwork.", "Folded yellow black softball comforter"),
        5: ("Close-up fabric view of yellow and black softball patchwork design.", "Softball comforter fabric close-up"),
        6: ("Feature panel listing machine washable, soft and lightweight bedding benefits.", "Softball comforter feature panel"),
        7: ("Included items panel showing one comforter and two pillowcases.", "Softball comforter set included items"),
    },
    115: {
        1: ("Bedroom mockup of football comforter with distressed American flag background, Kevin text and 20 shams.", "Kevin football flag comforter with 20"),
        2: ("Bedroom mockup of patriotic football bedding with Kevin text and matching shams.", "Kevin football American flag bedding set"),
        3: COMMON_DUVET_PANEL,
        4: ("Low-angle football flag bedding mockup with soft, lightweight, durable and breathable icons.", "Football flag bedding with feature icons"),
        5: COMMON_CARE,
        6: COMMON_SIZE,
    },
    116: {
        1: ("Bedroom mockup of blue semi truck comforter on dark geometric background with YOUR NAME panel.", "Blue geometric semi truck comforter"),
        2: ("Bedroom mockup of blue semi truck bedding with matching shams and folded top sheet.", "Blue semi truck bedding set on bed"),
        3: COMMON_DUVET_PANEL,
        4: ("Feature panel showing blue semi truck pillow mockup, microfiber label and 3D printed pattern callouts.", "Blue semi truck comforter feature panel"),
        5: ("Low-angle blue semi truck bedding mockup with soft, lightweight, durable and breathable icons.", "Blue semi truck bedding with feature icons"),
        6: COMMON_CARE,
        7: COMMON_SIZE,
    },
    117: {
        1: ("Bedroom mockup of colorful Tree of Life quilt with rainbow leaves, flowing roots and matching shams.", "Colorful Tree of Life quilt set"),
        2: ("Room mockup showing colorful Tree of Life quilt with premium craft callout.", "Colorful Tree of Life quilt room mockup"),
        3: ("Optional pillow shams panel showing matching Tree of Life pillow designs.", "Colorful Tree of Life pillow shams"),
        4: ("Fabric and use panel showing colorful Tree of Life quilt draped on a bed.", "Colorful Tree of Life quilt fabric panel"),
        5: COMMON_QUILT_SIZE,
        6: COMMON_QUILT_LAYER,
        7: ("Bedroom mockup of rainbow Tree of Life quilt with matching shams and white furniture.", "Rainbow Tree of Life quilt bedroom mockup"),
    },
    118: {
        1: ("Bedroom mockup of green Celtic Yggdrasil quilt with oval knotwork frame and matching shams.", "Celtic Yggdrasil tree quilt set"),
        2: ("Room mockup showing green Tree of Life quilt with premium craft callout.", "Green Tree of Life quilt room mockup"),
        3: ("Optional pillow shams panel showing green Celtic Tree of Life pillow designs.", "Celtic Tree of Life pillow shams"),
        4: ("Fabric and use panel showing green Yggdrasil quilt draped on a bed.", "Green Yggdrasil quilt fabric panel"),
        5: COMMON_QUILT_SIZE,
        6: COMMON_QUILT_LAYER,
        7: ("Bedroom mockup of Celtic Yggdrasil quilt with matching shams and wood room styling.", "Green Celtic Yggdrasil quilt bedroom mockup"),
    },
    119: {
        1: ("Bedroom mockup of red semi truck breaking through gray stone wall with YOUR NAME panel.", "Red semi truck stone wall comforter"),
        2: ("Bedroom mockup of red semi truck stone wall bedding with matching shams and folded top sheet.", "Red semi truck stone wall bedding set"),
        3: COMMON_DUVET_PANEL,
        4: ("Feature panel showing red semi truck pillow mockup, microfiber label and 3D printed pattern callouts.", "Red semi truck comforter feature panel"),
        5: ("Low-angle red semi truck bedding mockup with soft, lightweight, durable and breathable icons.", "Red semi truck bedding with feature icons"),
        6: COMMON_CARE,
        7: COMMON_SIZE,
    },
    120: {
        1: ("Bedroom mockup of black semi truck breaking through gray stone wall with YOUR NAME panel.", "Black semi truck stone wall comforter"),
        2: ("Bedroom mockup of black semi truck stone wall bedding with matching shams and folded top sheet.", "Black semi truck stone wall bedding set"),
        3: COMMON_DUVET_PANEL,
        4: ("Feature panel showing black semi truck pillow mockup, microfiber label and 3D printed pattern callouts.", "Black semi truck comforter feature panel"),
        5: ("Low-angle black semi truck bedding mockup with soft, lightweight, durable and breathable icons.", "Black semi truck bedding with feature icons"),
        6: COMMON_CARE,
        7: COMMON_SIZE,
    },
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def load_qa():
    data = json.loads(QA_DATASET.read_text(encoding="utf-8"))
    products = data["QA_Products"]
    customizer = json.loads(CUSTOMIZER_AUDIT.read_text(encoding="utf-8"))
    return {
        "product_keys": [item["product_key"] for item in products],
        "handle_by_key": {item["product_key"]: item["handle"] for item in products},
        "pos_by_key": {item["product_key"]: int(item["inventory_position"]) for item in products},
        "customizer_by_pos": {int(item["inventory_position"]): item for item in customizer},
    }


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Body (HTML)": first.get("Body (HTML)", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    return admin


def option_text(admin_row):
    values = [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]]
    values = [v for v in values if v]
    return ", ".join(values) if values else "available size options"


def customizer_sentence(audit):
    excerpt = audit.get("config_excerpt", "")
    text_matches = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"minLength":\s*(\d+),\s*"maxLength":\s*(\d+)',
        excerpt,
    )
    option_matches = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"defaultOptionId"',
        excerpt,
    )
    parts = []
    for label, required, min_length, max_length in text_matches:
        requirement = "required" if required == "true" else "optional"
        parts.append(f"Personalization uses a {requirement} {label} field, {min_length}-{max_length} characters.")
    useful_options = [
        label for label, _ in option_matches
        if label and "confirmation" not in label.lower()
    ]
    if useful_options:
        parts.append(f"Additional visible customization option: {', '.join(useful_options)}.")
    if not text_matches:
        parts.append("No shopper text-entry field is described for this product.")
    labels = " ".join(label for label, *_ in text_matches).lower()
    if "number" in labels:
        parts.append("Names and numbers visible in mockups are sample artwork unless the matching input is entered.")
    elif "name" in labels:
        parts.append("Visible names in mockups are sample artwork unless the matching input is entered.")
    else:
        parts.append("Visible custom text in mockups is sample artwork unless the matching input is entered.")
    return " ".join(parts)


def description(pos, admin_row, ctext):
    item = PRODUCT_UPDATES[pos]
    return (
        f"<p>{item['title']} features {item['detail']}. "
        "This product page is written around the specific visible artwork rather than a broad template.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork: {item['detail']}.</li>"
        "<li>Gallery images show the main bed mockup plus detail, feature, care, size or bedding-type panels where present.</li></ul>"
        "<h3>Options and Fit</h3>"
        f"<ul><li>Available option groups: {option_text(admin_row)}.</li>"
        f"<li>{ctext}</li>"
        "<li>Choose the product type and size shown on the product page before checkout.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    qa = load_qa()
    admin = load_admin(set(qa["handle_by_key"].values()))
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ctext = customizer_sentence(qa["customizer_by_pos"].get(pos, {}))
        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or "Bedding"
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, ctext)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"US English buyer intent targets {item['cluster']}; no volume or ranking claim is made."
        )
        ws.cell(row_num, idx["keyword_evidence_level"]).value = "SERP_ONLY"
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R3 deep recheck after Sang QA batch 012: used products_export_1.csv sha256:{admin_hash}; "
            "removed internal workflow copy, tied personalization wording only to customizer evidence, "
            "split nearby soccer/truck/quilt intents by visible motif, and rewrote image observations/alts from inspected contact sheets. "
            "Still NEEDS_REVIEW."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    seen_alts = set()
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        image_number = int(ws.cell(row_num, idx["image_number"]).value)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation, alt = IMAGE_DETAILS[pos][image_number]
        if alt in seen_alts:
            alt = f"{alt} for {PRODUCT_UPDATES[pos]['cluster']}"
        seen_alts.add(alt)
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = alt[:125]
        ws.cell(row_num, idx["alt_action"]).value = "SET"
        ws.cell(row_num, idx["viewed_status"]).value = "VIEWED"
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R3: observation and alt rewritten from inspected contact sheet; admin image URL and current alt matched from Shopify CSV where available."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "PRODUCT"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = (
            "Sang QA batch 012 + products_export_1.csv admin baseline + customizer_audit.json + contact-sheet inspection"
        )
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = (
            "R3 separates nearby pages by visible motif, sample wording, color palette and proof level for customization."
        )
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r3"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as a specific product page."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm the exact artwork, product type, size choices, "
            "and only the customization options that are actually supported."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = (
            "Confirm product type, visible motif, size/care panels, supported custom fields and image accuracy before purchase."
        )
        ws.cell(row_num, idx["emotional_social_motivation"]).value = (
            "Choose a sports, truck or Tree of Life bedding item without confusing it with another similar design in the batch."
        )
        ws.cell(row_num, idx["purchase_concerns"]).value = (
            "Avoid unsupported name, number, photo or material claims; confirm exact size, product type and customization fields."
        )
        ws.cell(row_num, idx["source_refs"]).value = (
            f"products_export_1.csv sha256:{admin_hash}; Sang QA {QA_RUN_ID}; customizer_audit.json; contact sheets 111-120"
        )
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R3_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = (
            "No Search Console, paid keyword volume or internal site-search export supplied; proof remains SERP-only."
        )
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    pk_by_pos = {qa["pos_by_key"][pk]: pk for pk in qa["product_keys"]}
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = pk_by_pos[pos]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = (
            f"Original storefront/product evidence; Sang QA {QA_RUN_ID}; products_export_1.csv sha256:{admin_hash}; "
            "customizer_audit.json; inspected contact sheets"
        )
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={option_text(admin_row)}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={item['detail']}; "
            f"customizer={customizer_sentence(qa['customizer_by_pos'].get(pos, {}))}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "Sang QA issues addressed in r3; independent QA recheck still required."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R3 uses inspected contact sheets, customizer audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_012_r3_revision", "r3", "Separate 10-product revision after Sang QA; source workbook was not modified."),
        ("qa_batch_012_r3_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_012_r3_scope", "inventory positions 111-120", "No products outside qa_batch_012 were revised."),
        ("qa_batch_012_r3_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_012_r3_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_012_r3_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_012 only; inventory positions 111-120",
        "revision": "r3",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Hold for combined QA handoff after revisions through qa_batch_020, per user request.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_012_r3",
                "",
                "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r2: `{SOURCE.relative_to(ROOT)}`",
                f"- QA của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r3: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Sửa trọng yếu: bỏ nội dung nội bộ, chỉ dùng claim name/number khi customizer audit xác nhận, tách intent cho soccer/truck/quilt, viết lại mô tả/alt theo contact sheets.",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Bước tiếp theo: giữ lại để gửi Sang QA chung sau khi revision đến `qa_batch_020`, theo yêu cầu người dùng.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
