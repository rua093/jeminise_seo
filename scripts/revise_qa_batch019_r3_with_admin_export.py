import csv
import hashlib
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_019_r3"
QA_RUN_ID = "20260907_213200"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_019_r2" / "SEO_Product_Optimization_qa_batch_019_r2.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
CUSTOMIZER_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_audit.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_019_r3.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_019_r3.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    181: {
        "title": "Custom Baseball Glove Name Bedding",
        "meta_title": "Custom Baseball Glove Name Bedding",
        "meta_description": "Shop custom baseball glove bedding with brown glove artwork, baseballs, required name field and feature panels.",
        "primary": "custom baseball glove name bedding",
        "secondary": "baseball glove bedding, personalized baseball comforter, custom name baseball bedding",
        "cluster": "custom name brown baseball glove bedding design 06",
        "detail": "brown baseball glove bedding with close-up baseball, glove leather texture, netting and sample name and number text",
        "intent_role": "Baseball bedding page supported by a required name field and separated by brown glove close-up artwork; visible numbers are sample artwork only.",
    },
    182: {
        "title": "Custom Baseball Flag Name Bedding",
        "meta_title": "Custom Baseball Flag Name Bedding",
        "meta_description": "Shop custom baseball flag bedding with large baseball, red white blue flag artwork and required name field.",
        "primary": "custom baseball flag name bedding",
        "secondary": "American flag baseball bedding, personalized baseball bedding, baseball name comforter",
        "cluster": "custom name baseball American flag bedding design 07",
        "detail": "red, white and blue American flag bedding with a large baseball and sample cursive name near the bottom edge",
        "intent_role": "Baseball bedding page supported by a required name field and separated by oversized ball on flag layout.",
    },
    183: {
        "title": "Custom Baseball Flag Glove Bedding",
        "meta_title": "Custom Baseball Flag Glove Bedding",
        "meta_description": "Shop custom baseball flag glove bedding with American flag background, baseball artwork and required name field.",
        "primary": "custom baseball flag glove bedding",
        "secondary": "baseball flag bedding, American flag baseball bedding, personalized sports bedding",
        "cluster": "custom name baseball flag glove bedding design 08",
        "detail": "distressed American flag baseball bedding with glove corners, oversized baseball, sample name and large sample number",
        "intent_role": "Baseball bedding page supported by a required name field and separated by large sample number artwork on the baseball layout.",
    },
    184: {
        "title": "Custom Close-Up Baseball Glove Bedding",
        "meta_title": "Custom Close-Up Baseball Glove Bedding",
        "meta_description": "Shop custom baseball glove bedding with close-up glove and ball artwork, required name field and feature panels.",
        "primary": "custom close-up baseball glove bedding",
        "secondary": "baseball glove comforter, personalized baseball bedding, custom name sports bedding",
        "cluster": "custom name close-up baseball glove bedding design 09",
        "detail": "close-up brown baseball glove and ball bedding with sample name printed across the baseball",
        "intent_role": "Baseball bedding page supported by a required name field and separated by full-frame glove close-up.",
    },
    185: {
        "title": "Custom Gray Baseball Flag Bedding",
        "meta_title": "Custom Gray Baseball Flag Bedding",
        "meta_description": "Shop custom gray baseball flag bedding with glove and ball artwork, required name field and feature panels.",
        "primary": "custom gray baseball flag bedding",
        "secondary": "gray baseball bedding, American flag baseball comforter, personalized baseball bedding",
        "cluster": "custom name gray baseball American flag bedding design 10",
        "detail": "gray baseball glove and ball artwork over a distressed American flag background with a large sample name",
        "intent_role": "Baseball bedding page supported by a required name field and separated by gray glove over flag artwork.",
    },
    186: {
        "title": "Custom Baseball Pattern Name Bedding",
        "meta_title": "Custom Baseball Pattern Name Bedding",
        "meta_description": "Shop custom baseball pattern bedding with red white blue stripes, many baseballs, required name field and feature panels.",
        "primary": "custom baseball pattern name bedding",
        "secondary": "red white blue baseball bedding, baseball number bedding, personalized baseball comforter",
        "cluster": "custom name baseball pattern red blue bedding design 11",
        "detail": "red, white and blue striped bedding with repeating baseballs, star panels, sample name and sample number on shams",
        "intent_role": "Baseball bedding page supported by a required name field and separated by repeated baseball pattern; visible numbers are sample artwork only.",
    },
    187: {
        "title": "Custom Pitcher Baseball Bedding",
        "meta_title": "Custom Pitcher Baseball Bedding",
        "meta_description": "Shop custom pitcher baseball bedding with silhouette artwork, baseball stitch border and required name field.",
        "primary": "custom pitcher baseball bedding",
        "secondary": "baseball player bedding, pitcher silhouette comforter, personalized baseball bedding",
        "cluster": "custom name pitcher silhouette baseball bedding design 12",
        "detail": "cream baseball bedding with black pitcher silhouette, red baseball stitch border and sample name at the foot",
        "intent_role": "Baseball bedding page supported by a required name field and separated by pitcher silhouette artwork.",
    },
    188: {
        "title": "Custom Black Baseball Glove Bedding",
        "meta_title": "Custom Black Baseball Glove Bedding",
        "meta_description": "Shop custom black baseball glove bedding with ball artwork, required name field, sample number artwork and feature panels.",
        "primary": "custom black baseball glove bedding",
        "secondary": "black baseball bedding, baseball glove comforter, personalized sports bedding",
        "cluster": "custom name black baseball glove bedding design 13",
        "detail": "black baseball bedding with grayscale glove and ball artwork, sample cursive name and sample number",
        "intent_role": "Baseball bedding page supported by a required name field and separated by black monochrome glove design; visible numbers are sample artwork only.",
    },
    189: {
        "title": "Custom Blue Stripe Baseball Bedding",
        "meta_title": "Custom Blue Stripe Baseball Bedding",
        "meta_description": "Shop custom blue stripe baseball bedding with oversized baseball artwork, required name field and feature panels.",
        "primary": "custom blue stripe baseball bedding",
        "secondary": "blue baseball bedding, personalized baseball comforter, custom name baseball bedding",
        "cluster": "custom name blue striped baseball bedding design 14",
        "detail": "light blue striped bedding with oversized baseball graphic, baseball pillow pattern and sample cursive name",
        "intent_role": "Baseball bedding page supported by a required name field and separated by blue striped baseball layout.",
    },
    190: {
        "title": "Custom Galaxy Baseball Name Bedding",
        "meta_title": "Custom Galaxy Baseball Name Bedding",
        "meta_description": "Shop custom galaxy baseball bedding with blue space background, baseball artwork, required name field and sample number artwork.",
        "primary": "custom galaxy baseball name bedding",
        "secondary": "galaxy baseball bedding, baseball comforter, personalized baseball bedding",
        "cluster": "custom name galaxy baseball bedding design 15",
        "detail": "dark blue galaxy baseball bedding with space background, large baseball, sample cursive name and sample number",
        "intent_role": "Baseball bedding page supported by a required name field and separated by galaxy background artwork; visible numbers are sample artwork only.",
    },
}


COMMON_WEAVING = ("High-density weaving panel comparing BeddingOutlet fabric with other weave samples.", "High-density weaving feature panel")
COMMON_WASH = ("Machine washable care panel with washing machine, laundry basket and care text.", "Machine washable bedding care panel")
COMMON_ZIPPER = ("Bottom zippered closure panel showing white zipper close-up.", "Bottom zippered closure panel")


IMAGE_DETAILS = {
    181: {
        1: ("Bedroom mockup of custom brown baseball glove bedding with sample name and number.", "Custom baseball glove name bedding"),
        2: COMMON_WEAVING,
        3: COMMON_WASH,
        4: COMMON_ZIPPER,
    },
    182: {
        1: ("Bedroom mockup of custom baseball flag bedding with large baseball and sample cursive name.", "Custom baseball flag name bedding"),
        2: COMMON_WEAVING,
        3: COMMON_WASH,
        4: COMMON_ZIPPER,
    },
    183: {
        1: ("Bedroom mockup of custom baseball flag bedding with glove corners, sample name and large number.", "Custom baseball name number bedding"),
        2: COMMON_WEAVING,
        3: COMMON_WASH,
        4: COMMON_ZIPPER,
    },
    184: {
        1: ("Bedroom mockup of custom close-up baseball glove bedding with sample name on the ball.", "Custom close-up baseball glove bedding"),
        2: COMMON_WEAVING,
        3: COMMON_WASH,
        4: COMMON_ZIPPER,
    },
    185: {
        1: ("Bedroom mockup of custom gray baseball flag bedding with glove, ball and sample name.", "Custom gray baseball flag bedding"),
        2: COMMON_WEAVING,
        3: COMMON_WASH,
        4: COMMON_ZIPPER,
    },
    186: {
        1: ("Bedroom mockup of custom baseball pattern bedding with red blue stripes, many baseballs and sample name.", "Custom baseball pattern name bedding"),
        2: COMMON_WASH,
        3: COMMON_ZIPPER,
        4: COMMON_WEAVING,
    },
    187: {
        1: ("Bedroom mockup of custom pitcher baseball bedding with silhouette, stitch border and sample name.", "Custom pitcher baseball bedding"),
        2: COMMON_WEAVING,
        3: COMMON_WASH,
        4: COMMON_ZIPPER,
    },
    188: {
        1: ("Bedroom mockup of custom black baseball glove bedding with grayscale ball, sample name and number.", "Custom black baseball glove bedding"),
        2: COMMON_ZIPPER,
        3: COMMON_WEAVING,
        4: COMMON_WASH,
    },
    189: {
        1: ("Bedroom mockup of custom blue stripe baseball bedding with oversized baseball and sample name.", "Custom blue stripe baseball bedding"),
        2: COMMON_WASH,
        3: COMMON_ZIPPER,
        4: COMMON_WEAVING,
    },
    190: {
        1: ("Bedroom mockup of custom galaxy baseball bedding with large baseball, sample name and number.", "Custom galaxy baseball name bedding"),
        2: COMMON_WEAVING,
        3: COMMON_WASH,
        4: COMMON_ZIPPER,
    },
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def load_qa():
    data = json.loads(QA_DATASET.read_text(encoding="utf-8"))
    products = data["QA_Products"]
    customizer = json.loads(CUSTOMIZER_AUDIT.read_text(encoding="utf-8"))
    return {
        "product_keys": [item["product_key"] for item in products],
        "handle_by_key": {item["product_key"]: item["handle"] for item in products},
        "pos_by_key": {item["product_key"]: int(item["inventory_position"]) for item in products},
        "customizer_by_pos": {int(item["inventory_position"]): item for item in customizer},
    }


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    missing = sorted(set(handles) - set(admin))
    if missing:
        raise RuntimeError(f"Missing handles in admin export: {missing}")
    return admin


def option_text(admin_row):
    values = [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]]
    values = [v for v in values if v]
    return ", ".join(values) if values else "available size options"


def customizer_sentence(audit):
    excerpt = audit.get("config_excerpt", "")
    text_matches = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"minLength":\s*(\d+),\s*"maxLength":\s*(\d+)',
        excerpt,
    )
    option_matches = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"defaultOptionId"',
        excerpt,
    )
    parts = []
    for label, required, min_length, max_length in text_matches:
        requirement = "required" if required == "true" else "optional"
        parts.append(f"Personalization uses a {requirement} {label} field, {min_length}-{max_length} characters.")
    useful_options = [label for label, _ in option_matches if label and "confirmation" not in label.lower()]
    if useful_options:
        parts.append(f"Additional visible customization option: {', '.join(useful_options)}.")
    if not text_matches:
        parts.append("No shopper text-entry field is described for this product.")
    labels = " ".join(label for label, *_ in text_matches).lower()
    if "number" in labels:
        parts.append("Names and numbers visible in mockups are sample artwork unless the matching input is entered.")
    elif "name" in labels:
        parts.append("Visible names in mockups are sample artwork unless the matching input is entered. Any visible numbers are sample artwork only.")
    elif text_matches:
        parts.append("Visible custom text in mockups is sample artwork unless the matching input is entered.")
    else:
        parts.append("Displayed wording or artwork is part of the product design shown, not a shopper-entered text claim.")
    return " ".join(parts)


def description(pos, admin_row, ctext):
    item = PRODUCT_UPDATES[pos]
    return (
        f"<p>{item['title']} features {item['detail']}. "
        "The page copy focuses on the specific visible design, product type and verified custom fields.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork: {item['detail']}.</li>"
        "<li>Gallery images show the main baseball bedding mockup plus weave, care or zipper feature panels.</li></ul>"
        "<h3>Customization and Fit</h3>"
        f"<ul><li>Available option groups: {option_text(admin_row)}.</li>"
        f"<li>{ctext}</li>"
        "<li>Select the product type and size shown on the product page before checkout.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    qa = load_qa()
    admin = load_admin(set(qa["handle_by_key"].values()))
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ctext = customizer_sentence(qa["customizer_by_pos"].get(pos, {}))
        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or "Bedding"
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, ctext)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"US English buyer intent targets {item['cluster']}; no search-volume, trend or ranking claim is made."
        )
        ws.cell(row_num, idx["keyword_evidence_level"]).value = "SERP_ONLY"
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R3 deep recheck after Sang QA batch 019: used products_export_1.csv sha256:{admin_hash}; "
            "removed internal workflow copy, filled publishable title values, kept custom name claims only because customizer audit verifies them, "
            "split baseball intents by visible motif and rewrote image observations/alts from inspected contact sheets. Still NEEDS_REVIEW."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    seen_alts = set()
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        image_number = int(ws.cell(row_num, idx["image_number"]).value)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation, alt = IMAGE_DETAILS[pos][image_number]
        if alt in seen_alts:
            alt = f"{alt} for {PRODUCT_UPDATES[pos]['cluster']}"
        seen_alts.add(alt)
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = alt[:125]
        ws.cell(row_num, idx["alt_action"]).value = "SET"
        ws.cell(row_num, idx["viewed_status"]).value = "VIEWED"
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R3: observation and alt rechecked from inspected contact sheet; admin image URL and current alt matched from Shopify CSV where available."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "PRODUCT"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = (
            "Sang QA batch 019 + products_export_1.csv admin baseline + customizer audit + contact-sheet inspection"
        )
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = (
            "R3 separates nearby baseball pages by motif, layout, color palette and verified custom fields."
        )
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r3"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as a specific product page."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm exact baseball artwork, product type, size choices, "
            "and the required name field before purchase."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = (
            "Confirm product type, visible baseball motif, required custom name field, care panels and image accuracy before purchase."
        )
        ws.cell(row_num, idx["emotional_social_motivation"]).value = (
            "Choose a personalized baseball bedding item that matches a preferred visual style, such as flag, glove, pitcher or galaxy artwork."
        )
        ws.cell(row_num, idx["purchase_concerns"]).value = (
            "Confirm spelling before checkout; avoid unsupported team, player, league, material, custom-number or delivery claims."
        )
        ws.cell(row_num, idx["source_refs"]).value = (
            f"products_export_1.csv sha256:{admin_hash}; Sang QA {QA_RUN_ID}; customizer audit; contact sheets 181-190"
        )
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R3_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = (
            "No Search Console, paid keyword volume or internal site-search export supplied; proof remains SERP-only."
        )
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    pk_by_pos = {qa["pos_by_key"][pk]: pk for pk in qa["product_keys"]}
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = pk_by_pos[pos]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ctext = customizer_sentence(qa["customizer_by_pos"].get(pos, {}))
        ws.cell(row_num, idx["sources_accessed"]).value = (
            f"Original storefront/product evidence; Sang QA {QA_RUN_ID}; products_export_1.csv sha256:{admin_hash}; "
            "customizer audit; inspected contact sheets"
        )
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={option_text(admin_row)}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={item['detail']}; "
            f"customizer={ctext}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "Sang QA issues addressed in r3; independent QA recheck still required."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R3 uses inspected contact sheets, customizer audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_019_r3_revision", "r3", "Separate 10-product revision after Sang QA; source workbook was not modified."),
        ("qa_batch_019_r3_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_019_r3_scope", "inventory positions 181-190", "No products outside qa_batch_019 were revised."),
        ("qa_batch_019_r3_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_019_r3_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_019_r3_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_019 only; inventory positions 181-190",
        "revision": "r3",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Hold for combined QA handoff after revisions through qa_batch_020, per user request.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_019_r3",
                "",
                "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r2: `{SOURCE.relative_to(ROOT)}`",
                f"- QA của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r3: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Sửa trọng yếu: rà lại theo rút kinh nghiệm các batch trước, bỏ nội dung nội bộ, giữ title/H1 publishable, chỉ giữ custom name theo customizer audit, tách intent cho 10 thiết kế baseball theo motif nhìn thấy và viết lại mô tả/alt theo contact sheets.",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Bước tiếp theo: giữ lại để gửi Sang QA chung sau khi revision đến `qa_batch_020`, theo yêu cầu người dùng.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
