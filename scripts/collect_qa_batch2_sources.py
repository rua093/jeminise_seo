import csv
import hashlib
import json
from datetime import datetime
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
QA_RUN_ID = "20260907_091545"
QA_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID
WORKBOOK = QA_DIR / "source_snapshot" / "SEO_Product_Optimization_through_batch_034.xlsx"
INVENTORY = ROOT / "seo_runs" / SHOP / RUN_ID / "inventory.csv"
EVIDENCE_PRODUCTS = ROOT / "seo_runs" / SHOP / RUN_ID / "evidence" / "products"


def rows(workbook, name):
    sheet = workbook[name]
    headers = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    return [dict(zip(headers, (c.value for c in row))) for row in sheet.iter_rows(min_row=2)]


def fetch_live(url):
    headers = {"User-Agent": "Mozilla/5.0 independent SEO QA/1.0"}
    response = requests.get(url, headers=headers, timeout=45)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    product_response = requests.get(url + ".js", headers=headers, timeout=45)
    product_response.raise_for_status()
    product = product_response.json()
    find = lambda tag, **kwargs: soup.find(tag, **kwargs)
    canonical, description, h1, title = find("link", rel="canonical"), find("meta", attrs={"name":"description"}), find("h1"), find("title")
    return {"checked_at": datetime.now().astimezone().isoformat(), "status_code": response.status_code,
        "html_sha256": hashlib.sha256(response.content).hexdigest(),
        "title_element": title.get_text(" ", strip=True) if title else "", "h1": h1.get_text(" ", strip=True) if h1 else "",
        "meta_description": description.get("content", "").strip() if description else "",
        "canonical": canonical.get("href", "").strip() if canonical else "", "product_js": product}


def main():
    with INVENTORY.open(newline="", encoding="utf-8-sig") as handle:
        batch = list(csv.DictReader(handle))[10:20]
    keys, handles, urls = {r["product_key"] for r in batch}, {r["Handle"] for r in batch}, {r["product_url"] for r in batch}
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    submitted = {"products":[r for r in rows(wb,"SEO_Products") if r["product_key"] in keys],
        "images":[r for r in rows(wb,"Image_Audit") if r["Handle"] in handles],
        "product_evidence":[r for r in rows(wb,"Product_Evidence") if r["product_url"] in urls],
        "keyword_map":[r for r in rows(wb,"Keyword_Map") if r["product_key"] in keys],
        "buyer_search_research":[r for r in rows(wb,"Buyer_Search_Research") if r["product_key"] in keys]}
    (QA_DIR/"submitted_batch_data.json").write_text(json.dumps(submitted,ensure_ascii=False,indent=2),encoding="utf-8")
    live=[]
    for row in batch:
        record={"inventory":row}
        try: record["live"],record["fetch_error"]=fetch_live(row["product_url"]),""
        except Exception as exc: record["live"],record["fetch_error"]={},str(exc)
        matches=list(EVIDENCE_PRODUCTS.glob(f"{row['inventory_position']}_*/evidence_summary.json"))
        if matches: record["research_snapshot"]=json.loads(matches[0].read_text(encoding="utf-8"))
        live.append(record)
    (QA_DIR/"live_source_comparison.json").write_text(json.dumps(live,ensure_ascii=False,indent=2),encoding="utf-8")
    print(f"products={len(submitted['products'])} images={len(submitted['images'])} keywords={len(submitted['keyword_map'])} buyers={len(submitted['buyer_search_research'])} live={len(live)}")
    for x in live: print(x['inventory']['inventory_position'],x['live'].get('status_code'),len(x['live'].get('product_js',{}).get('media',[])),x['fetch_error'])


if __name__ == "__main__": main()
