import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "batches" / "SEO_Product_Optimization_through_batch_034.xlsx"
QA_ROOT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa"
RESULT_REVISION_ROOT = ROOT / "resutls" / SHOP / RUN_ID / "revisions"
RUN_REVISION_ROOT = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions"

QA_RUNS = {
    "qa_batch_001": "20260907_083534",
    "qa_batch_002": "20260907_091545",
    "qa_batch_003": "20260907_094127",
    "qa_batch_004": "20260907_095643",
    "qa_batch_005": "20260907_101050",
    "qa_batch_006": "20260907_102634",
    "qa_batch_007": "20260907_104027",
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def clean_phrase(value):
    value = text(value)
    replacements = [
        (r"\bDragonfly\b", ""),
        (r"\bChristian Knight Templar\b", "Christian inspirational"),
        (r"\bD[0-9]{1,2}\b", ""),
        (r"\bsnowman Christmas\b", "gingerbread Christmas"),
        (r"\bsnowman\b", "gingerbread figure"),
        (r"\balligator\b", "crocodile"),
        ("\ufffd", ""),
    ]
    for pattern, replacement in replacements:
        value = re.sub(pattern, replacement, value, flags=re.I)
    return re.sub(r"\s{2,}", " ", value).strip(" ,;.-")


def strip_custom_claims(value):
    value = text(value)
    replacements = [
        (r"\b[Pp]ersonalized\s+", ""),
        (r"\b[Pp]ersonalised\s+", ""),
        (r"\b[Cc]ustomizable\s+", ""),
        (r"\b[Cc]ustom\s+", ""),
        (r"\bwith custom name and number\b", "with sports artwork"),
        (r"\bwith name and number\b", "with sports artwork"),
        (r"\bwith custom name\b", "with themed artwork"),
        (r"\bwith name\b", "with themed artwork"),
        (r"\bcustomization\b", "product option"),
        (r"\bpersonalization\b", "product option"),
        (r"\b[Pp]ersonalize\b", "Style"),
        (r"\b[Cc]ustomize\b", "Choose"),
        (r"\bcustom-made\b", "themed"),
        (r"\bpersonalized\b", "themed"),
        (r"\bcustom\b", "themed"),
    ]
    for pattern, replacement in replacements:
        value = re.sub(pattern, replacement, value)
    return re.sub(r"\s{2,}", " ", value).strip()


def product_kind(product_type):
    raw = text(product_type).lower()
    if "comforter" in raw:
        return "comforter set"
    if "blanket" in raw:
        return "blanket"
    if "quilt" in raw:
        return "quilt set"
    return "bedding item"


def extract_detail(desc_html, title):
    desc = text(desc_html)
    match = re.search(r"Bring distinctive (.*?) artwork", desc, flags=re.I | re.S)
    if match:
        detail = re.sub(r"<.*?>", " ", match.group(1))
    else:
        match = re.search(r"<li>(.*?)</li>", desc, flags=re.I | re.S)
        detail = re.sub(r"<.*?>", " ", match.group(1)) if match else title
    return re.sub(r"\s+", " ", clean_phrase(detail))[:260] or clean_phrase(title)


def build_description(row, idx, flagged_no_custom):
    title = clean_phrase(row[idx["title_proposed"] - 1])
    detail = extract_detail(row[idx["description_proposed_html"] - 1], title)
    if flagged_no_custom:
        detail = strip_custom_claims(detail)
    kind = product_kind(row[idx["product_type"] - 1])
    return (
        f"<p>Add themed style to the bedroom with this Jeminise {kind} featuring {detail}. "
        "The artwork is shown across the main bedding mockup and coordinated pillow or sham images where included.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork focus: {detail}.</li>"
        "<li>Gallery images show the main bedding view, detail views, fabric or care graphics, and included-component visuals where available.</li></ul>"
        "<h3>Product Options</h3>"
        "<ul><li>Choose from the available size and pillowcase options shown on the product page.</li>"
        "<li>Review the live product options before checkout to confirm the selected configuration.</li></ul>"
    )


def load_qa_batch(batch_id):
    qa_run_id = QA_RUNS[batch_id]
    data = json.loads((QA_ROOT / qa_run_id / "qa_dataset.json").read_text(encoding="utf-8"))

    products = data["QA_Products"]
    product_keys = [item["product_key"] for item in products]
    product_positions = {item["product_key"]: int(item["inventory_position"]) for item in products}

    qa_image_to_media = {}
    image_observations = {}
    for image in data["QA_Images"]:
        key = (image["product_key"], text(image["media_id"]))
        qa_image_to_media[image["qa_image_key"]] = key
        image_observations[key] = image["qa_observation"]

    no_custom = set()
    product_issues = defaultdict(list)
    image_fixes = {}
    issue_counts = defaultdict(lambda: defaultdict(int))
    for issue in data["QA_Issues"]:
        pk = issue["product_key"]
        if pk:
            product_issues[pk].append(issue["issue_id"])
            issue_counts[pk][issue["severity"]] += 1
        issue_text = " ".join([issue["field"], issue["reason"], issue["submitted_value"]])
        if pk and issue["severity"] == "CRITICAL" and re.search(r"personal|custom", issue_text, re.I):
            no_custom.add(pk)
        if issue["field"].startswith("image_") and issue["qa_image_key"]:
            key = qa_image_to_media.get(issue["qa_image_key"])
            if key and issue["recommended_fix"]:
                image_fixes[key] = issue["recommended_fix"][:125]

    return {
        "qa_run_id": qa_run_id,
        "product_keys": product_keys,
        "product_positions": product_positions,
        "no_custom": no_custom,
        "product_issues": product_issues,
        "image_fixes": image_fixes,
        "image_observations": image_observations,
        "issue_counts": issue_counts,
        "raw": data,
    }


def revise_batch(batch_id):
    qa = load_qa_batch(batch_id)
    result_dir = RESULT_REVISION_ROOT / f"{batch_id}_r2"
    run_dir = RUN_REVISION_ROOT / f"{batch_id}_r2"
    result_dir.mkdir(parents=True, exist_ok=True)
    run_dir.mkdir(parents=True, exist_ok=True)

    output = result_dir / f"SEO_Product_Optimization_{batch_id}_r2.xlsx"
    shutil.copy2(SOURCE, output)
    wb = load_workbook(output)
    now = datetime.now().astimezone().isoformat()

    product_key_set = set(qa["product_keys"])
    revised_products = 0
    revised_images = 0

    ws_products = wb["SEO_Products"]
    pidx = headers(ws_products)
    handle_to_pk = {}

    for row_num in range(2, ws_products.max_row + 1):
        pk = text(ws_products.cell(row_num, pidx["product_key"]).value)
        handle_to_pk[text(ws_products.cell(row_num, pidx["Handle"]).value)] = pk
        if pk not in product_key_set:
            continue

        row_values = [ws_products.cell(row_num, col).value for col in range(1, ws_products.max_column + 1)]
        flagged = pk in qa["no_custom"]
        for col_name in [
            "title_proposed",
            "meta_title_seo",
            "meta_description_seo",
            "primary_keyword",
            "secondary_keywords",
            "meta_keyword",
            "keyword_strategy",
            "buyer_search_summary",
        ]:
            value = clean_phrase(ws_products.cell(row_num, pidx[col_name]).value)
            if flagged:
                value = strip_custom_claims(value)
            ws_products.cell(row_num, pidx[col_name]).value = value

        if qa["product_positions"].get(pk) == 24:
            ws_products.cell(row_num, pidx["primary_keyword"]).value = "gingerbread Christmas quilt"
            ws_products.cell(row_num, pidx["meta_keyword"]).value = "gingerbread Christmas quilt"
            ws_products.cell(row_num, pidx["secondary_keywords"]).value = "gingerbread bedding, Christmas village quilt, holiday quilt set"
            ws_products.cell(row_num, pidx["title_proposed"]).value = "Gingerbread Christmas Village Quilt Set"
            ws_products.cell(row_num, pidx["meta_title_seo"]).value = "Gingerbread Christmas Village Quilt Set"
            ws_products.cell(row_num, pidx["keyword_strategy"]).value = (
                "Target the visible gingerbread Christmas village motif; do not use snowman wording for this product."
            )

        ws_products.cell(row_num, pidx["description_proposed_html"]).value = build_description(row_values, pidx, flagged)
        ws_products.cell(row_num, pidx["meta_title_length"]).value = len(text(ws_products.cell(row_num, pidx["meta_title_seo"]).value))
        ws_products.cell(row_num, pidx["meta_description_length"]).value = len(text(ws_products.cell(row_num, pidx["meta_description_seo"]).value))
        ws_products.cell(row_num, pidx["revision"]).value = "r2"
        ws_products.cell(row_num, pidx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in pidx:
                ws_products.cell(row_num, pidx[col_name]).value = ""
        ws_products.cell(row_num, pidx["issues"]).value = (
            f"R2 revision after {batch_id}: removed internal SEO/QA/import language; "
            f"{'removed unverified personalization/custom claims; ' if flagged else ''}"
            "remaining admin-export, material, source-mapping and keyword evidence issues require QA recheck. "
            f"Original QA issue refs: {', '.join(qa['product_issues'][pk])}."
        )
        revised_products += 1

    ws_images = wb["Image_Audit"]
    iidx = headers(ws_images)
    for row_num in range(2, ws_images.max_row + 1):
        handle = text(ws_images.cell(row_num, iidx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        media_id = text(ws_images.cell(row_num, iidx["media_id"]).value)
        key = (pk, media_id)
        if key in qa["image_observations"]:
            ws_images.cell(row_num, iidx["observed_visual_details"]).value = qa["image_observations"][key]
        if key in qa["image_fixes"]:
            ws_images.cell(row_num, iidx["alt_proposed"]).value = qa["image_fixes"][key]
        else:
            alt = clean_phrase(ws_images.cell(row_num, iidx["alt_proposed"]).value)
            if pk in qa["no_custom"]:
                alt = strip_custom_claims(alt)
            ws_images.cell(row_num, iidx["alt_proposed"]).value = alt[:125]
        ws_images.cell(row_num, iidx["revision"]).value = "r2"
        ws_images.cell(row_num, iidx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in iidx:
                ws_images.cell(row_num, iidx[col_name]).value = ""
        ws_images.cell(row_num, iidx["issues"]).value = (
            f"R2 revision after {batch_id}: alt/observation refreshed from QA where available; "
            "requires QA recheck and admin media-alt export before approval."
        )
        revised_images += 1

    ws_keywords = wb["Keyword_Map"]
    kidx = headers(ws_keywords)
    for row_num in range(2, ws_keywords.max_row + 1):
        pk = text(ws_keywords.cell(row_num, kidx["product_key"]).value)
        if pk not in product_key_set:
            continue
        keyword = clean_phrase(ws_keywords.cell(row_num, kidx["keyword"]).value)
        if pk in qa["no_custom"]:
            keyword = strip_custom_claims(keyword)
        if qa["product_positions"].get(pk) == 24 and ws_keywords.cell(row_num, kidx["keyword_role"]).value == "PRIMARY":
            keyword = "gingerbread Christmas quilt"
        ws_keywords.cell(row_num, kidx["keyword"]).value = keyword
        ws_keywords.cell(row_num, kidx["decision_reason"]).value = (
            f"R2 after {batch_id}: revised from QA issue set; product query must match visible artwork and verified purchase features."
        )
        ws_keywords.cell(row_num, kidx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws_keywords.cell(row_num, kidx["mapping_version"]).value = "r2"

    ws_research = wb["Buyer_Search_Research"]
    ridx = headers(ws_research)
    for row_num in range(2, ws_research.max_row + 1):
        pk = text(ws_research.cell(row_num, ridx["product_key"]).value)
        if pk not in product_key_set:
            continue
        for col_name in [
            "purchase_context",
            "jtbd_statement",
            "functional_motivation",
            "emotional_social_motivation",
            "purchase_concerns",
            "seo_application",
        ]:
            value = clean_phrase(ws_research.cell(row_num, ridx[col_name]).value)
            if pk in qa["no_custom"]:
                value = strip_custom_claims(value)
            ws_research.cell(row_num, ridx[col_name]).value = value
        ws_research.cell(row_num, ridx["limitations"]).value = (
            text(ws_research.cell(row_num, ridx["limitations"]).value)
            + f" R2 note: {batch_id} QA found issues requiring recheck before approval."
        ).strip()

    ws_evidence = wb["Product_Evidence"]
    eidx = headers(ws_evidence)
    allowed_positions = set(qa["product_positions"].values())
    for row_num in range(2, ws_evidence.max_row + 1):
        evidence_id = text(ws_evidence.cell(row_num, eidx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match or int(match.group(1)) not in allowed_positions:
            continue
        ws_evidence.cell(row_num, eidx["confidence_and_reason"]).value = (
            f"R2 revised from {batch_id}; content still needs independent QA recheck and human approval before import."
        )

    ws_readme = wb["README_QA"]
    midx = headers(ws_readme)
    for metric, value, definition in [
        (f"{batch_id}_r2_revision", "r2", f"Separate 10-product revision after {batch_id}; source workbook was not modified."),
        (f"{batch_id}_r2_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        (f"{batch_id}_r2_revised_images", str(revised_images), "Image_Audit rows refreshed and kept NEEDS_REVIEW."),
        (f"{batch_id}_r2_scope", ", ".join(qa["product_keys"]), "Exact product_key scope for this revision artifact."),
    ]:
        ws_readme.append([
            metric if col == midx["metric"] else value if col == midx["value"] else definition if col == midx["definition"] else ""
            for col in range(1, ws_readme.max_column + 1)
        ])

    wb.save(output)
    load_workbook(output).save(output)

    summary = {
        "created_at": now,
        "batch_id": batch_id,
        "qa_run_id": qa["qa_run_id"],
        "source_workbook": str(SOURCE.relative_to(ROOT)),
        "revision_workbook": str(output.relative_to(ROOT)),
        "scope_rule": "official per-batch revision, exactly 10 products unless final batch is smaller",
        "inventory_positions": sorted(allowed_positions),
        "product_keys": qa["product_keys"],
        "revision": "r2",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "no_custom_claims_removed_products": len(qa["no_custom"]),
        "image_recommended_fixes_applied": len(qa["image_fixes"]),
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "requires_next_step": f"Sang QA lại {batch_id}_r2 before any human approval or import.",
        "issue_counts_by_product": {pk: dict(counts) for pk, counts in qa["issue_counts"].items()},
    }
    summary_path = run_dir / "revision_summary.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    report_path = result_dir / f"SEO_Product_Optimization_{batch_id}_r2.md"
    report_path.write_text(
        "\n".join(
            [
                f"# Revision {batch_id} r2",
                "",
                "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Workbook nguồn: `{SOURCE.relative_to(ROOT)}`",
                f"- QA nguồn: `{qa['qa_run_id']}`",
                f"- Workbook revision: `{output.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                f"- Bước tiếp theo: Sang QA lại `{batch_id}_r2` trước khi duyệt/import.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    return summary


def main():
    summaries = [revise_batch(batch_id) for batch_id in QA_RUNS]
    index_path = RUN_REVISION_ROOT / "per_batch_r2_index.json"
    index_path.parent.mkdir(parents=True, exist_ok=True)
    index_path.write_text(json.dumps(summaries, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"created_batches={len(summaries)}")
    for item in summaries:
        print(
            f"{item['batch_id']}: products={item['revised_products']} "
            f"images={item['revised_images']} workbook={item['revision_workbook']}"
        )
    print(f"index={index_path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
