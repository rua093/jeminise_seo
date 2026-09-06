import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_034"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_033.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_034.xlsx"


SOFTBALL_REFS = "https://www.amazon.com/softball-comforter/s?k=softball+comforter; https://www.etsy.com/market/softball_bedding; https://www.pinterest.com/ideas/softball-bedding/918246657688/; https://www.wayfair.com/keyword.php?keyword=softball+bedding"
PERSONALIZED_SOFTBALL_REFS = "https://www.etsy.com/market/personalized_softball_blanket; https://www.amazon.com/personalized-softball-blanket/s?k=personalized+softball+blanket; https://www.pinterest.com/search/pins/?q=personalized%20softball%20bedding; https://www.zazzle.com/softball+blankets"
CELTIC_REFS = "https://www.amazon.com/clp/B0DNSVY2QJ; https://luvingift.com/product/viking-quilt-set-norse-celtic-yggdrasil-tree-of-life/; https://myvikinggear.com/collections/viking-home-decoration/viking-quilt-set/; https://www.pinterest.com/pin/562668547209120403/"
CARDINAL_REFS = "https://www.etsy.com/market/cardinal_quilt; https://www.amazon.com/cardinal-quilt/s?k=cardinal+quilt; https://www.wayfair.com/keyword.php?keyword=cardinal+bedding; https://www.pinterest.com/ideas/cardinal-quilt/938570791549/"
CAT_CARDINAL_REFS = "https://www.etsy.com/market/cat_christmas_quilt; https://www.amazon.com/christmas-cat-quilt/s?k=christmas+cat+quilt; https://www.pinterest.com/search/pins/?q=cat%20cardinal%20christmas%20quilt; https://www.wayfair.com/keyword.php?keyword=christmas+cat+bedding"


PROFILES = {
    331: ("personalized vintage softball comforter", "personalized softball comforter", "Personalized Softball Comforter Set", "Personalized Vintage Softball Comforter Set", "dark vintage softball comforter with large yellow ball, Caitlin name, number 07 and matching personalized shams", "softball bedding with name, custom softball comforter, girls softball bedding", "EVERGREEN", PERSONALIZED_SOFTBALL_REFS),
    332: ("teal polka dot softball comforter", "teal softball comforter", "Teal Softball Comforter Set", "Teal Polka Dot Softball Comforter Set", "black polka dot and teal stripe softball comforter with Mia name, I Love Softball text and paisley border", "softball bedding for girls, softball fan bedding, teal softball bedding", "EVERGREEN", PERSONALIZED_SOFTBALL_REFS),
    333: ("fire softball comforter", "fire softball comforter", "Fire Softball Comforter Set", "Fire Softball Name and Number Comforter Set", "bright yellow softball comforter with fiery orange background, Lexie name, number 21 and matching shams", "personalized softball bedding, softball number comforter, softball bedroom decor", "EVERGREEN", PERSONALIZED_SOFTBALL_REFS),
    334: ("red glove softball comforter", "softball glove comforter", "Softball Glove Comforter Set", "Red Softball Glove Comforter Set", "red softball glove comforter with yellow ball, Reagan name, number 11 and matching personalized shams", "personalized softball comforter, softball player bedding, softball bedding with name", "EVERGREEN", PERSONALIZED_SOFTBALL_REFS),
    335: ("Celtic Tree of Life quilt", "Celtic Tree of Life quilt", "Celtic Tree of Life Quilt Set", "Celtic Tree of Life Knot Roots Quilt Set", "green and neutral Tree of Life quilt with Celtic knot roots, circular landscape medallion and knotwork border", "Tree of Life bedding, Celtic knot quilt, Yggdrasil quilt set", "EVERGREEN", CELTIC_REFS),
    336: ("winter cardinal berry quilt", "winter cardinal quilt", "Winter Cardinal Quilt Set", "Winter Cardinal Berry Branch Quilt Set", "winter cardinal quilt with red berry branches, blue gray leaves, watercolor bird artwork and matching sham", "cardinal Christmas quilt, winter bird bedding, cardinal bedding set", "CHRISTMAS", CARDINAL_REFS),
    337: ("heart cardinals Christmas quilt", "cardinal Christmas quilt", "Cardinal Christmas Quilt Set", "Heart Cardinals Christmas Quilt Set", "Christmas quilt with two red cardinals, heart-shaped berry branches, snowy village background and matching sham", "winter cardinal quilt, Christmas bird bedding, cardinal bedding set", "CHRISTMAS", CARDINAL_REFS),
    338: ("cat and cardinal Christmas quilt", "cat cardinal Christmas quilt", "Cat and Cardinal Christmas Quilt Set", "Winter Cat and Cardinal Christmas Quilt Set", "snowy Christmas quilt with gray kitten, red cardinal, winter branch artwork and It Is Well With My Soul text", "cat Christmas bedding, cardinal quilt, winter cat quilt", "CHRISTMAS", CAT_CARDINAL_REFS),
}


def qa_issues(pos, html_error):
    issues = ["Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."]
    if 331 <= pos <= 334:
        issues.append("Softball comforters are close variants and continue batch_033 cluster; review keyword cannibalization, title uniqueness and whether sample names or numbers should be removed from deploy copy.")
        issues.append("Several current storefront titles appear truncated; proposed titles should be checked against admin field limits before import.")
    if pos == 334:
        issues.append("Public JSON/contact sheet shows duplicated image sequence, 14 images instead of a unique 7-image set; media order and duplicate removal need QA.")
    if pos == 335:
        issues.append("Celtic Tree of Life overlaps previous Yggdrasil and Tree of Life products; keep knot roots and neutral green landscape motif distinct.")
    if pos in (336, 337):
        issues.append("Winter cardinal quilts are close Christmas variants; distinguish single cardinal versus heart-shaped pair and verify seasonal targeting.")
    if pos == 338:
        issues.append("Cat and cardinal Christmas quilt includes visible quote text; verify exact quote spelling and whether religious phrase should appear in SEO fields.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def product_context(pos):
    if 331 <= pos <= 334:
        return "softball bedrooms, teen sports rooms, player gifts, team-family gifts or personalized fan bedding."
    if pos == 335:
        return "Celtic bedrooms, nature-inspired rooms, spiritual decor or Tree of Life gift bedding."
    return "Christmas bedrooms, winter guest rooms, bird lovers, pet lovers or seasonal holiday bedding."


def alt_for_image(h1, pos, ipos):
    if 331 <= pos <= 334:
        return {
            1: f"{h1} displayed on a bed in the main comforter mockup",
            2: f"Angled bedroom mockup for {h1}",
            3: f"Matching pillowcase detail for {h1}",
            4: f"Folded comforter mockup for {h1}",
            5: f"Fabric close-up for {h1}",
            6: f"Soft microfiber and care feature image for {h1}",
            7: f"Included comforter and pillowcase information for {h1}",
            8: f"Duplicate main bedroom mockup for {h1}",
            9: f"Duplicate angled bedroom mockup for {h1}",
            10: f"Duplicate pillowcase detail for {h1}",
            11: f"Duplicate folded comforter mockup for {h1}",
            12: f"Duplicate fabric close-up for {h1}",
            13: f"Duplicate microfiber and care feature image for {h1}",
            14: f"Duplicate included comforter and pillowcase information for {h1}",
        }.get(ipos, f"Product detail image for {h1}")[:125]
    return {
        1: f"{h1} displayed on a bed in the main quilt mockup",
        2: f"Close-up artwork detail for {h1}",
        3: f"Matching pillow sham for {h1}",
        4: f"Bedroom lifestyle mockup for {h1}",
        5: f"Bedding set size and sham information for {h1}",
        6: f"Bedspread construction and care feature image for {h1}",
        7: f"Overhead bedroom mockup of {h1}",
    }.get(ipos, f"Product detail image for {h1}")[:125]


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for {product_context(pos)}</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broad softball, Tree of Life and Christmas quilt terms for collection pages or strongest representative products.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, room decor or a personalized gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": qa_issues(pos, html_error),
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed bedding, the buyer wants a {kind} matching a specific softball, Celtic Tree of Life or Christmas motif.",
            "functional_motivation": "Find the right design, size, product type, personalization fields, care details and room fit.",
            "emotional_social_motivation": "Give a personalized sports gift or decorate with Celtic nature, winter cardinal, pet or Christmas artwork.",
            "purchase_concerns": "Personalization spelling, sample names or numbers, product identity, material/care claims, size fit, included components and whether artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt_for_image(h1, pos, ipos), "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_034_products", "value": "8", "definition": "Products appended in this final cumulative workbook"},
        {"metric": "batch_034_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_034"},
        {"metric": "cumulative_products", "value": "338", "definition": "Total products included through batch_034"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[338:348]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "ALL_INVENTORY_DRAFTED_AWAITING_FINAL_QA",
        "awaiting_confirmation": False,
        "next_batch_id": "",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=8 appended_images={added_images}")


if __name__ == "__main__":
    main()
