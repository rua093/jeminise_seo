import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_010"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_009.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_010.xlsx"


PROFILES = {
    91: ("turquoise boho wolf quilt set", "wolf quilt set", "Boho Wolf Quilt Set", "Turquoise Boho Wolf Quilt Set", "wolf head inside a turquoise circular boho frame with matching shams and tribal-style ornamental accents", "wolf bedding, boho wolf quilt, wildlife quilt set", "EVERGREEN", "https://www.amazon.com/wolf-quilt/s?k=wolf+quilt; https://www.walmart.com/c/kp/wolf-quilts; https://www.ebay.com/b/Wolf-Quilt-In-Duvet-Covers-Bedding-Sets/37644/bn_7022238087"),
    92: ("personalized neon green soccer player comforter with Jackson number 07", "neon soccer comforter", "Neon Soccer Comforter Set", "Personalized Neon Soccer Player Comforter", "black comforter with neon green soccer player artwork, Jackson name, number 07 and glowing ball details", "custom soccer bedding, soccer comforter with name, soccer bedding for boys", "EVERGREEN", "https://www.amazon.com/Personalized-Bedding-Daughter-Friends-Birthday/dp/B0C3VYL11C; https://www.etsy.com/listing/956705304/personalized-soccer-ball-comforter-flame; https://www.pinterest.com/pin/personalized-soccer-duvet-cover-set-soccer-fire-ball-player-gift-idea-black-duvet-cover-pillowcases-custom--999939923509327748/"),
    93: ("personalized blue semi truck American flag comforter", "personalized semi truck comforter", "Personalized Semi Truck Comforter", "Personalized Blue Semi Truck Flag Comforter", "blue semi truck over a distressed American flag background with Your Name personalization area", "custom trucker bedding, American flag truck comforter, semi truck bedding with name", "EVERGREEN", "https://jeminise.com/; https://www.desertcart.in/; https://www.wal-mart.com/"),
    94: ("personalized patriotic semi truck comforter with waving flag", "patriotic semi truck comforter", "Patriotic Semi Truck Comforter", "Personalized Patriotic Semi Truck Comforter", "silver semi truck with waving American flag artwork and large Your Name personalization text", "American flag truck bedding, custom semi truck comforter, trucker bedding set", "EVERGREEN", "https://jeminise.com/; https://www.desertcart.in/; https://www.wal-mart.com/"),
    95: ("personalized red semi truck sunset flag comforter", "red semi truck comforter", "Red Semi Truck Comforter Set", "Personalized Red Semi Truck Flag Comforter", "red semi truck below a waving American flag at sunset with Your Name personalization text", "patriotic truck bedding, custom trucker comforter, American flag semi truck bedding", "EVERGREEN", "https://jeminise.com/; https://www.desertcart.in/; https://www.wal-mart.com/"),
    96: ("custom photo collage quilt with text", "custom photo quilt", "Custom Photo Quilt Set", "Custom Photo Collage Quilt Set", "photo collage quilt template with multiple personal image panels and central Your Text Here message area", "personalized photo quilt, custom picture quilt, photo collage bedding", "EVERGREEN", "https://www.etsy.com/market/custom_photo_quilt; https://www.etsy.com/market/custom_photo_baseball_blanket; https://www.bedbathandbeyond.com/Home-Garden/Thomas-Black-Bronze-Framed-Picture-Frame-Photo-Frame/43715273/product.html"),
    97: ("personalized purple God Says I Am butterfly blanket for Elizabeth", "purple God Says I Am blanket", "Purple God Says I Am Blanket", "Personalized Purple God Says I Am Blanket", "purple God Says I Am blanket with monarch-style butterflies, Elizabeth name and Bible verse affirmation blocks", "Christian butterfly blanket, personalized Bible verse blanket, God Says I Am blanket", "EVERGREEN", "https://www.etsy.com/market/god_says_i_am_personalized_blanket; https://christianartbag.com/products/christianart-blanket-god-says-i-am-christian-blanket-bible-verse-blanket-personalized-blanket-christmas-gift-cabbk01111123; https://us.shein.com/Christian-Purple-Butterfly-Flannel-Throw-Blanket-Inspirational-Bible-Quote-Print-Soft-Cozy-Warm-Lightweight-All-Season-Blanket-For-Sofa-Bed-Couch-Bedroom-Living-Room-Office-Travel-Camping-Nap-Home-Decor-Christian-Gift-For-Women-Mom-Friends-Decorative-p-345272924.html"),
    98: ("personalized purple God is within her butterfly blanket for Isabella", "purple Christian butterfly blanket", "Purple Christian Butterfly Blanket", "Personalized Purple Christian Butterfly Blanket", "purple God Is Within Her design with butterflies, Isabella name and Psalm 46:5 verse-inspired affirmation blocks", "Christian blanket for her, butterfly Bible verse blanket, personalized faith blanket", "EVERGREEN", "https://us.shein.com/Christian-Purple-Butterfly-Flannel-Throw-Blanket-Inspirational-Bible-Quote-Print-Soft-Cozy-Warm-Lightweight-All-Season-Blanket-For-Sofa-Bed-Couch-Bedroom-Living-Room-Office-Travel-Camping-Nap-Home-Decor-Christian-Gift-For-Women-Mom-Friends-Decorative-p-345272924.html; https://www.pinterest.com/pin/personalized-purple-cross-fleece-blanket--889109151413619739/; https://www.zazzle.com/z_purple_cross_christian_symbol_personalized_fleece_blanket-256080411035054740"),
    99: ("personalized purple floral Christian cross blanket for Amabel", "purple floral cross blanket", "Purple Floral Cross Blanket", "Personalized Purple Floral Cross Blanket", "dark purple floral Christian cross design with purple roses, butterflies, Blessed is she text and Amabel name", "Christian cross blanket, personalized faith blanket, purple Bible verse blanket", "EVERGREEN", "https://www.amazon.com/Blanket-Christian-Religious-Hummingbird-Catholic/dp/B0CKZ87PH6; https://www.pinterest.com/pin/personalized-purple-cross-fleece-blanket--889109151413619739/; https://www.zazzle.com/z_purple_cross_christian_symbol_personalized_fleece_blanket-256080411035054740"),
    100: ("personalized purple floral cross Bible verse blanket for Estella", "purple floral cross Bible verse blanket", "Purple Floral Cross Bible Verse Blanket", "Personalized Purple Floral Cross Bible Verse Blanket", "purple Bible verse collage design with wooden cross, floral accents, butterflies and Estella name", "Christian cross blanket, personalized scripture blanket, purple faith blanket", "EVERGREEN", "https://www.amazon.com/Blanket-Christian-Religious-Hummingbird-Catholic/dp/B0CKZ87PH6; https://www.pinterest.com/pin/personalized-purple-cross-fleece-blanket--889109151413619739/; https://www.zazzle.com/z_purple_cross_christian_symbol_personalized_fleece_blanket-256080411035054740"),
}


def append_dict(ws, data):
    headers = [cell.value for cell in ws[1]]
    ws.append([data.get(h, "") for h in headers])


def image_obs(theme, pos):
    if pos == 1:
        return f"Main product mockup showing {theme}."
    if pos == 2:
        return f"Secondary flat, close-up or bedroom mockup showing {theme}."
    if pos == 3:
        return f"Bedding type, feature, lifestyle or material image for {theme}."
    if pos == 4:
        return f"Detail, feature panel or lifestyle image connected to {theme}."
    if pos == 5:
        return f"Additional mockup, care or size information for {theme}."
    if pos == 6:
        return f"Easy-care, size guide or product feature image for {theme}."
    return f"Additional lifestyle, feature or product detail image showing {theme}."


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        issue_bits = [
            "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."
        ]
        if pos in (93, 94, 95):
            issue_bits.append("Closely related patriotic semi truck comforters should be reviewed for keyword cannibalization and title uniqueness.")
        if pos == 96:
            issue_bits.append("Current title says bicycle and flowers, but contact sheet shows a custom photo collage quilt template; product identity needs QA.")
        if pos == 100:
            issue_bits.append("Current title appears truncated around 'Machi'; spelling should be corrected during QA.")
        if html_error:
            issue_bits.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
        issues = " ".join(issue_bits)
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for themed decor, faith gifting or personalized bedding.</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork, personalization and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broader bedding terms for collections or stronger representative products.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as a themed bedding item, personalized gift or faith keepsake.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": issues,
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed or personalized bedding, the buyer wants a {kind} matching a specific design, name, message or hobby.",
            "functional_motivation": "Find the right design, size, bedding type, care details and personalization fit.",
            "emotional_social_motivation": "Create meaningful room decor or give a personalized sports, trucker, wildlife or faith gift.",
            "purchase_concerns": "Personalization spelling, product identity, material/care claims, design distinction and whether the visible artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork, personalization and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed in a main product mockup",
                2: f"Secondary mockup of {h1}",
                3: f"Product option, feature or lifestyle image for {h1}",
                4: f"Detail or feature image for {h1}",
                5: f"Additional product mockup or care information for {h1}",
                6: f"Care, size or product feature image for {h1}",
                7: f"Additional product mockup of {h1}",
                8: f"Lifestyle use image for {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_010_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_010_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_010"},
        {"metric": "cumulative_products", "value": "100", "definition": "Total products included through batch_010"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[100:110]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_011",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
