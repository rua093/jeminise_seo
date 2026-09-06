import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_030"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_029.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_030.xlsx"


READING_REFS = "https://www.etsy.com/market/book_lover_blanket; https://www.amazon.com/clp/B0DJ32WQ8X; https://linovagoods.com/products/personalized-book-blanket-custom-name-reading-throw-floral-bookshelf-gift-for-book-lovers-cozy-library-decor; https://www.etsy.com/listing/1749170765/just-a-girl-who-loves-books-blanket"
PHOTO_REFS = "https://www.etsy.com/market/mothers_day_photo_blanket; https://www.etsy.com/market/custom_photo_quilt; https://www.amazon.com/custom-photo-blanket/s?k=custom+photo+blanket; https://blog.mimeophotos.com/mothers-day-photobooks"
WOLF_REFS = "https://www.walmart.com/c/kp/southwest-quilts-patterns; https://www.etsy.com/market/queen_size_wolf_quilts; https://society6.com/collections/duvet-covers-southwestern; https://www.pinterest.com/pin/820007044635684499/"
CHRISTIAN_REFS = "https://www.etsy.com/listing/1756643894/personalized-christian-blanket-bible; https://www.amazon.com/Christian-Religious-Scripture-SunflowerPlush-Blankets/dp/B0CJR2192D; https://business.walmart.com/ip/Yellow-Sunflower-Prayer-Blanket-Bible-Verse-Healing-Inspirational-Sympathy-Fleece-Throw-Religious-Christian-Gifts-Women-Wife-Womens-Gift-Ideas-40-x50/2483775954; https://www.pinterest.com/ideas/bible-verse-blanket/910177925520/"
CACTUS_REFS = "https://www.target.com/c/bedding-home-decor/cactus/-/N-5xtv4Zg2fff; https://www.wayfair.com/keyword.php?keyword=cactus+bedding; https://www.pinterest.com/ideas/cactus-bedding/959782873627/; https://www.ebay.com/itm/257177873478"


PROFILES = {
    291: ("reading girl on book with flowers blanket", "reading girl blanket", "Reading Girl Blanket", "Reading Girl on Book with Flowers Blanket", "beige reading blanket with girl sitting on an open book, floral wreath, camera border details and custom name", "book lover blanket, personalized reading blanket, bookworm gift", "EVERGREEN", READING_REFS),
    292: ("reading girl heart bookshelf blanket", "heart bookshelf reading blanket", "Heart Bookshelf Reading Blanket", "Reading Girl Heart Bookshelf Blanket", "book lover blanket with girl inside heart bookshelf, stacked custom book titles, butterflies and custom name", "personalized book blanket, reading blanket for girls, book lover gift", "EVERGREEN", READING_REFS),
    293: ("Best Mom Ever mother daughter photo quilt", "mother daughter photo quilt", "Mother Daughter Photo Quilt", "Best Mom Ever Mother Daughter Photo Quilt", "black photo collage quilt with Best Mom Ever text, mother and daughter photos, daisy details and optional shams", "custom mom photo quilt, Mothers Day photo blanket, personalized mom quilt", "MOTHERS_DAY", PHOTO_REFS),
    294: ("seasonal bookworms Sophia reading blanket", "seasonal bookworms blanket", "Seasonal Bookworms Blanket", "Seasonal Bookworms Sophia Reading Blanket", "library reading blanket with Sophia Bookworm text, cartoon girl in cozy chair, plants and seasonal reader quotes", "bookworm blanket, personalized book lover blanket, reading nook blanket", "EVERGREEN", READING_REFS),
    295: ("southwestern wolf head quilt", "southwestern wolf quilt", "Southwestern Wolf Quilt Set", "Southwestern Wolf Head Quilt Set", "southwestern bedding with large wolf head, geometric tribal borders, neutral tan palette and matching sham artwork", "wolf bedding, southwestern quilt, geometric wolf comforter", "EVERGREEN", WOLF_REFS),
    296: ("sunflower Christian Bible verse blanket", "sunflower Christian blanket", "Sunflower Christian Blanket", "Sunflower Christian Bible Verse Blanket", "bright sunflower blanket with woman silhouette, God says you are text, Bible verse labels and custom name", "Bible verse blanket, Christian affirmation blanket, personalized faith blanket", "EVERGREEN", CHRISTIAN_REFS),
    297: ("sunflower God Says You Are blanket", "God Says You Are blanket", "God Says You Are Blanket", "Sunflower God Says You Are Blanket", "rustic sunflower blanket with woodgrain woman silhouette, God Says You Are text, Bible references and custom name", "Christian sunflower blanket, Bible verse throw blanket, personalized prayer blanket", "EVERGREEN", CHRISTIAN_REFS),
    298: ("sunflowers butterflies self affirmations blanket", "self affirmations blanket", "Self Affirmations Blanket", "Sunflowers and Butterflies Self Affirmations Blanket", "cream affirmation blanket with praying girl photo, sunflower border, butterflies, custom name and self-affirmation phrases", "personalized affirmation blanket, self love blanket, sunflower blanket", "EVERGREEN", CHRISTIAN_REFS),
    299: ("teal argyle reading girl blanket", "teal reading girl blanket", "Teal Reading Girl Blanket", "Teal Argyle Reading Girl Blanket", "teal argyle reading blanket with rear-view girl, book stacks, dragonflies, custom name and happily ever after quote", "book lover blanket, personalized reading blanket, bookish gift", "EVERGREEN", READING_REFS),
    300: ("vibrant cactus desert flowers quilt", "vibrant cactus quilt", "Vibrant Cactus Quilt Set", "Vibrant Cactus and Desert Flowers Quilt Set", "dark cactus quilt with potted succulents, bright desert flowers, deep green background and matching shams", "cactus bedding, succulent quilt, desert flower bedding", "EVERGREEN", CACTUS_REFS),
}


def qa_issues(pos, html_error):
    issues = ["Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."]
    if pos in (291, 292, 294, 299):
        issues.append("Book lover blankets overlap several previous reading-girl batches; keep open book, heart bookshelf, seasonal bookworms and teal argyle motifs distinct.")
        issues.append("Verify displayed sample names, character/background options, size variants and fleece/care claims.")
    if pos == 293:
        issues.append("Mother daughter photo quilt requires photo-upload QA, Mother's Day seasonal positioning review and optional pillow sham/component verification.")
    if pos == 295:
        issues.append("Wolf bedding should be reviewed for comforter versus quilt naming, included components and southwestern/geometric keyword overlap.")
    if pos in (296, 297, 298):
        issues.append("Christian sunflower/affirmation blankets need verse text, custom name/photo fields and faith wording reviewed before import.")
        issues.append("Feature images mention fleece, washing, stitching and softness claims; confirm exact material, GSM and care details.")
    if pos == 300:
        issues.append("Cactus quilt overlaps recent cactus/desert batches; review motif separation and canonical strategy.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def product_context(pos):
    if pos in (291, 292, 294, 299):
        return "reading nooks, bedrooms, book clubs, librarian gifts or personalized reader gifts."
    if pos == 293:
        return "Mother's Day gifts, birthday gifts for mom, mother daughter keepsakes or personalized bedroom decor."
    if pos == 295:
        return "southwestern bedrooms, rustic cabins, wolf decor or outdoors-inspired bedding."
    if pos in (296, 297, 298):
        return "faith gifts, prayer spaces, bedrooms, church gifts or personalized encouragement gifts."
    return "southwestern bedrooms, desert decor, cactus-themed rooms or garden-inspired bedding."


def alt_for_image(h1, pos, ipos):
    if pos in (293, 295, 300):
        return {
            1: f"{h1} displayed in the main bedding mockup",
            2: f"Flat product or close-up image for {h1}",
            3: f"Size chart or pillow sham detail for {h1}",
            4: f"Bedroom lifestyle mockup for {h1}",
            5: f"Fabric or component feature image for {h1}",
            6: f"Angled bedroom mockup for {h1}",
            7: f"Pillow sham or bedspread feature image for {h1}",
        }.get(ipos, f"Product detail image for {h1}")[:125]
    return {
        1: f"{h1} shown as a personalized throw blanket",
        2: f"Flat product or alternate character image for {h1}",
        3: f"Size chart, couch mockup or close-up for {h1}",
        4: f"Lifestyle reading or couch image for {h1}",
        5: f"Fabric and fleece feature image for {h1}",
        6: f"Couch or close-up mockup for {h1}",
        7: f"Reading lifestyle image for {h1}",
        8: f"Additional reading lifestyle image for {h1}",
        9: f"Size, color or family lifestyle image for {h1}",
    }.get(ipos, f"Product detail image for {h1}")[:125]


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for {product_context(pos)}</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broad reading, photo gift, wolf, Christian and cactus bedding terms for collection pages or strongest representatives.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, room decor or a personalized gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": qa_issues(pos, html_error),
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed bedding, the buyer wants a {kind} matching a specific reader, family photo, wolf, faith or cactus motif.",
            "functional_motivation": "Find the right design, size, product type, personalization fields, care details and room fit.",
            "emotional_social_motivation": "Give a personalized reader, mom, faith or nature gift, or decorate with bookish, rustic or desert artwork.",
            "purchase_concerns": "Personalization spelling, photo quality, verse text, product identity, material/care claims, size fit, included components and whether artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt_for_image(h1, pos, ipos), "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_030_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_030_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_030"},
        {"metric": "cumulative_products", "value": "300", "definition": "Total products included through batch_030"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[300:310]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_031",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
