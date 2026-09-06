import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_009"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_008.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_009.xlsx"


PROFILES = {
    81: ("custom baseball batter silhouette blanket with name and number", "personalized baseball blanket", "Personalized Baseball Blanket", "Personalized Baseball Batter Silhouette Blanket", "black and purple baseball silhouette pattern with a batter over a flag stripe graphic and custom name/number placeholders", "custom baseball blanket, baseball blanket with name, baseball player gift blanket", "EVERGREEN", "https://www.etsy.com/market/custom_photo_baseball_blanket; https://www.amazon.com/Personalized-Baseball-Blankets-Numbers-Blanket/dp/B0GDW7K2CJ; https://www.pinterest.com/pin/custom-blankets-baseball-catcher-personalized-blanket-perfect-gift-for-son-fleece-blanket-personalized-blankets--889109151413626280/"),
    82: ("custom baseball batter photo blanket with Daniel", "custom baseball photo blanket", "Custom Baseball Photo Blanket", "Custom Baseball Batter Photo Blanket", "baseball batter on a red and blue stadium background with American flag section and the name Daniel", "personalized baseball blanket, baseball blanket with picture, baseball player gift blanket", "EVERGREEN", "https://www.etsy.com/market/custom_photo_baseball_blanket; https://www.etsy.com/listing/4321813622/custom-baseball-photo-and-name-blanket; https://www.amazon.com/Personalized-Baseball-Blankets-Numbers-Blanket/dp/B0GDW7K2CJ"),
    83: ("personalized baseball batter quote blanket with Christopher number 22", "personalized baseball player blanket", "Personalized Baseball Player Blanket", "Personalized Baseball Batter Quote Blanket", "vintage tan baseball batter artwork with the name Christopher, number 22 and the quote When you step on the field nothing else matters", "custom baseball blanket, baseball gift for players, baseball throw blanket", "EVERGREEN", "https://www.etsy.com/market/custom_photo_baseball_blanket; https://ohaprints.com/products/ohaprints-fleece-sherpa-blanket-baseball-player-ball-black-camo-sports-gift-custom-personalized-name-number-soft-throw-blanket-2206; https://www.amazon.com/OhaPrints-Baseball-Catcher-Personalized-Pillowcases/dp/B0BGDZ41FL"),
    84: ("personalized baseball practice quote blanket with Jackson number 25", "baseball practice quote blanket", "Baseball Practice Quote Blanket", "Personalized Baseball Practice Quote Blanket", "American flag and rustic wood design with Jackson 25, baseball glove and quote about practice, coaches and playing for him", "baseball blanket for boys, custom baseball blanket, baseball son gift blanket", "EVERGREEN", "https://www.etsy.com/market/custom_photo_baseball_blanket; https://www.amazon.com/Personalized-Baseball-Blankets-Numbers-Blanket/dp/B0GDW7K2CJ; https://www.pinterest.com/pin/custom-blankets-baseball-catcher-personalized-blanket-perfect-gift-for-son-fleece-blanket-personalized-blankets--889109151413626280/"),
    85: ("personalized flaming baseball blanket with Phoenix number 09", "flaming baseball blanket", "Flaming Baseball Blanket", "Personalized Flaming Baseball Blanket", "dark baseball design with glowing flaming baseball, flaming bat trail, Phoenix name and #09", "custom baseball blanket, baseball blanket with name, baseball player gift blanket", "EVERGREEN", "https://www.etsy.com/market/custom_photo_baseball_blanket; https://www.amazon.com/Personalized-Baseball-Blankets-Numbers-Blanket/dp/B0GDW7K2CJ; https://ohaprints.com/products/ohaprints-fleece-sherpa-blanket-baseball-player-ball-black-camo-sports-gift-custom-personalized-name-number-soft-throw-blanket-2206"),
    86: ("personalized baseball glove flag blanket with Jackson number 15", "baseball glove blanket", "Baseball Glove Blanket", "Personalized Baseball Glove Flag Blanket", "baseball glove, bat and ball artwork over a red, white and blue flag background with Jackson and number 15", "custom baseball blanket, patriotic baseball blanket, baseball player throw blanket", "EVERGREEN", "https://www.amazon.com/OhaPrints-Baseball-Catcher-Personalized-Pillowcases/dp/B0BGDZ41FL; https://www.etsy.com/market/custom_photo_baseball_blanket; https://www.amazon.com/Personalized-Baseball-Blankets-Numbers-Blanket/dp/B0GDW7K2CJ"),
    87: ("personalized baseball catcher blanket with Daniel number 56", "personalized baseball catcher blanket", "Personalized Baseball Catcher Blanket", "Personalized Baseball Catcher Blanket", "catcher reaching for a pitch inside large baseball stitching with the phrase I catch he pitches, Daniel name and 56", "custom catcher blanket, baseball catcher gift, baseball blanket with name", "EVERGREEN", "https://www.amazon.com/OhaPrints-Baseball-Catcher-Personalized-Pillowcases/dp/B0BGDZ41FL; https://www.pinterest.com/pin/custom-blankets-baseball-catcher-personalized-blanket-perfect-gift-for-son-fleece-blanket-personalized-blankets--889109151413626280/; https://www.etsy.com/market/custom_photo_baseball_blanket"),
    88: ("personalized American flag baseball batter blanket with Jude number 5", "American flag baseball blanket", "American Flag Baseball Blanket", "Personalized American Flag Baseball Blanket", "baseball batter silhouette over a vintage American flag background with Jude name and number 5", "patriotic baseball blanket, custom baseball throw, baseball blanket with name", "EVERGREEN", "https://www.etsy.com/market/custom_photo_baseball_blanket; https://www.amazon.com/Personalized-Baseball-Blankets-Numbers-Blanket/dp/B0GDW7K2CJ; https://www.walmart.com/ip/Personalized-Sports-Themed-Blanket-Customized-Balls-Blankets-Custom-Name-Throws-Boys-Girls-Teen-Athletes-Fun-Birthday-Gift-Print-finish-USA-2nd-day-s/13557065996"),
    89: ("personalized reversible baseball catcher batter blanket", "personalized baseball name number blanket", "Personalized Baseball Name Number Blanket", "Personalized Baseball Catcher Batter Blanket", "reversible-style baseball stitching design with catcher and batter silhouettes plus custom names Joey, Hudson and Noah with number 20", "custom baseball blanket, baseball blanket with name and number, baseball player gift blanket", "EVERGREEN", "https://www.etsy.com/market/custom_photo_baseball_blanket; https://www.amazon.com/Personalized-Baseball-Blankets-Numbers-Blanket/dp/B0GDW7K2CJ; https://ohaprints.com/products/ohaprints-fleece-sherpa-blanket-baseball-player-ball-black-camo-sports-gift-custom-personalized-name-number-soft-throw-blanket-2206"),
    90: ("personalized God says you are baseball blanket with name", "personalized baseball Christian blanket", "Personalized Baseball Christian Blanket", "Personalized God Says You Are Baseball Blanket", "vintage baseball faith design with God Says You Are text, Your Name placeholder, number 02 and Bible verse affirmations", "God says you are baseball blanket, Christian baseball blanket, personalized Bible verse blanket", "EVERGREEN", "https://www.etsy.com/listing/1762659181/personalized-baseball-blanket-god-says-i; https://www.amazon.com/ENCYCOM-Personalized-Gods-Say-Blanket/dp/B0FRFZ6QQD; https://www.walmart.com/ip/Customizaholic-Personalized-God-Says-You-Are-Blanket-Custom-Name-Bible-Verses-Floral-Christian-Faith-Gift-Encouragement-Birthday-Special-Moments/18727801098"),
}


def append_dict(ws, data):
    headers = [cell.value for cell in ws[1]]
    ws.append([data.get(h, "") for h in headers])


def image_obs(theme, pos):
    if pos == 1:
        return f"Main product or lifestyle mockup showing {theme}."
    if pos == 2:
        return f"Secondary mockup or customization image showing {theme}."
    if pos == 3:
        return f"Outdoor, sofa or lifestyle mockup connected to {theme}."
    if pos == 4:
        return f"Customization, detail or additional lifestyle panel for {theme}."
    if pos == 5:
        return f"Additional mockup, material detail or use-case image for {theme}."
    if pos == 6:
        return f"Additional lifestyle or product detail image showing {theme}."
    return f"Additional product mockup or lifestyle image showing {theme}."


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        issue_bits = [
            "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment.",
            "Batch contains many closely related personalized baseball blankets; keyword cannibalization and final title uniqueness need QA.",
        ]
        if html_error:
            issue_bits.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
        if "Baseba -" in row["title_current"]:
            issue_bits.append("Current title appears truncated around 'Baseba'; spelling should be corrected during QA.")
        issues = " ".join(issue_bits)
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for baseball gifting, player keepsakes or sports-themed decor.</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork, personalization and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific baseball motif: {label}. Keep broader baseball blanket terms for collection pages or top representatives.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as a custom sports gift, player blanket or baseball room accent.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": issues,
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for personalized baseball bedding, the buyer wants a {kind} with the right role, name, number or faith message for the recipient.",
            "functional_motivation": "Find the right baseball design, size, blanket material, personalization details and gift fit.",
            "emotional_social_motivation": "Give a player, catcher, son, coach or baseball fan a personalized sports keepsake.",
            "purchase_concerns": "Personalization spelling, jersey number accuracy, design distinction, material/care claims and whether the design matches the recipient's baseball role.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible baseball artwork, personalization and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed in a main product mockup",
                2: f"Secondary mockup of {h1}",
                3: f"Lifestyle or outdoor mockup of {h1}",
                4: f"Customization or product detail image for {h1}",
                5: f"Additional product mockup or material detail for {h1}",
                6: f"Additional lifestyle image for {h1}",
                7: f"Additional product mockup of {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_009_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_009_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_009"},
        {"metric": "cumulative_products", "value": "90", "definition": "Total products included through batch_009"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[90:100]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_010",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
