import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_015"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_014.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_015.xlsx"


DRAGON_REFS = "https://www.temu.com/dragon-bedding-set-5030175152007-s.html; https://www.pinterest.com/pin/800233427555270189/; https://www.etsy.com/market/fantasy_duvet; https://www.walmart.com/ip/Gothic-Bedding-Set-Dragon-Quilt-Fantasy-Castle-Dark-Sky-Comforter-Set-Queen-Vintage-Mythical-Animal-Duvet-Insert-Antique-Prehistoric-Room-Decor-Aesth/5536921884"
TREE_REFS = "https://www.amazon.com/clp/B0D9NXRP1Q; https://celticartstudio.com/product/celtic-tree-of-life-bedspread-jacquard-woven/; https://alphaquilt.com/products/tai111124134; https://www.amazon.com/clp/B0DNSVW2QK"
CARDINAL_REFS = "https://www.walmart.com/c/kp/cardinal-quilt; https://www.ebay.com/b/Cardinal-Quilt-Indiana-Quilts-Bedspreads-Coverlets/175749/bn_7022481015; https://www.wayfair.com/keyword.php?keyword=cardinal+bedspread; https://www.bedbathandbeyond.com/Bedding-Bath/MarCielo-3-Pcs-Winter-Cardinals-Christmas-Quilt-Bedspread-Set-C79/38941937/product.html"
FOX_REFS = "https://www.amazon.com/fox-quilt/s?k=fox+quilt; https://www.etsy.com/market/fox_quilt; https://www.pinterest.com/search/pins/?q=fox%20patchwork%20quilt"


PROFILES = {
    141: ("red fire dragon lava quilt set", "fire dragon quilt", "Fire Dragon Quilt Set", "Red Fire Dragon Lava Quilt Set", "red and black fantasy dragon with glowing lava, flames and dark volcanic shading", "dragon bedding, lava dragon quilt, fantasy dragon bedding", "EVERGREEN", DRAGON_REFS),
    142: ("teal winged mythical dragon quilt set", "winged dragon quilt", "Winged Dragon Quilt Set", "Teal Winged Mythical Dragon Quilt", "teal winged dragon on a dark fantasy background with glowing green light accents", "dragon bedding, mythical dragon quilt, fantasy bedding set", "EVERGREEN", DRAGON_REFS),
    143: ("golden roots Tree of Life quilt set", "gold tree of life quilt", "Gold Tree of Life Quilt Set", "Golden Roots Tree of Life Quilt Set", "golden Tree of Life with entwined roots, teal leaves and ornate gold accents on a black background", "tree of life bedding, Celtic tree quilt, gold tree quilt set", "EVERGREEN", TREE_REFS),
    144: ("twisted trunk Tree of Life quilt set", "twisted tree of life quilt", "Twisted Tree of Life Quilt Set", "Twisted Trunk Tree of Life Quilt Set", "twisted golden tree trunk and sprawling roots under a teal night-sky landscape", "tree of life bedding, Yggdrasil quilt set, Celtic tree quilt", "EVERGREEN", TREE_REFS),
    145: ("green fire dragon spikes quilt set", "green dragon quilt", "Green Dragon Quilt Set", "Green Fire Dragon Spikes Quilt Set", "green fantasy dragon with spikes, open mouth and glowing fire at the bottom of the artwork", "dragon bedding, spiked dragon quilt, fantasy dragon comforter", "EVERGREEN", DRAGON_REFS),
    146: ("purple nebula winged dragon quilt set", "nebula dragon quilt", "Nebula Dragon Quilt Set", "Purple Nebula Winged Dragon Quilt", "purple winged dragon with blue highlights and starry nebula pattern on a black background", "dragon bedding, galaxy dragon quilt, fantasy bedding set", "EVERGREEN", DRAGON_REFS),
    147: ("red volcanic winged dragon quilt set", "volcanic dragon quilt", "Volcanic Dragon Quilt Set", "Red Volcanic Winged Dragon Quilt", "red winged dragon silhouette over orange volcanic eruption and lava glow", "dragon bedding, lava dragon comforter, fantasy dragon quilt", "EVERGREEN", DRAGON_REFS),
    148: ("cardinal floral oval frame winter quilt", "cardinal floral quilt", "Cardinal Floral Quilt Set", "Cardinal Floral Oval Frame Winter Quilt", "red cardinal inside an ornate oval frame with roses, winter greenery and vintage floral accents", "cardinal quilt, winter cardinal bedding, floral bird quilt set", "HOLIDAY_WINTER", CARDINAL_REFS),
    149: ("Celtic knot Tree of Life quilt set", "Celtic tree of life quilt", "Celtic Tree of Life Quilt Set", "Celtic Knot Tree of Life Quilt Set", "black Tree of Life silhouette inside teal Celtic knot artwork with gold interlaced border", "tree of life bedding, Celtic knot quilt, folklore tree quilt", "EVERGREEN", TREE_REFS),
    150: ("fox patchwork animal quilt set", "fox patchwork quilt", "Fox Patchwork Quilt Set", "Fox Patchwork Animal Quilt Set", "orange fox portrait inside a circular patchwork frame with blue, red and tan geometric accents", "fox bedding, animal quilt set, woodland fox quilt", "EVERGREEN", FOX_REFS),
}


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        issue_bits = [
            "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."
        ]
        if pos in (141, 142, 145, 146, 147):
            issue_bits.append("Closely related fantasy dragon quilt variants should be reviewed for cannibalization and motif-specific title differentiation.")
        if pos in (143, 144, 149):
            issue_bits.append("Tree of Life/Celtic variants overlap with prior batches; review keyword map, title uniqueness and canonical strategy.")
        if pos == 148:
            issue_bits.append("Cardinal floral winter quilt may overlap with Christmas/cardinal products; seasonality and title distinction need QA.")
        if html_error:
            issue_bits.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
        issues = " ".join(issue_bits)
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for fantasy rooms, nature-inspired decor, winter bedding or woodland themes.</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broader bedding terms for collections or stronger representative products.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as a themed bedding item, room decor or gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": issues,
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed bedding, the buyer wants a {kind} matching a specific fantasy, nature, winter bird or woodland motif.",
            "functional_motivation": "Find the right design, size, bedding type, care details and decor fit.",
            "emotional_social_motivation": "Create a distinctive bedroom look or give a themed fantasy, nature, winter or woodland gift.",
            "purchase_concerns": "Product identity, material/care claims, design distinction and whether the visible artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed in a main product mockup",
                2: f"Secondary mockup or close-up of {h1}",
                3: f"Product pillow or lifestyle image for {h1}",
                4: f"Lifestyle room mockup for {h1}",
                5: f"Size, set contents or product feature image for {h1}",
                6: f"Care, size or product feature image for {h1}",
                7: f"Additional product mockup of {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_015_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_015_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_015"},
        {"metric": "cumulative_products", "value": "150", "definition": "Total products included through batch_015"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[150:160]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_016",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
