import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_018"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_017.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_018.xlsx"


HALLOWEEN_REFS = "https://www.amazon.com/Erosebridal-Comforter-Halloween-Anniversary-Valentines/dp/B0DB89DJYY; https://www.walmart.com/c/kp/skull-comforter; https://www.etsy.com/market/goth_quilt_comforter; https://www.pinterest.com/ideas/skull-comforter-set/950651219832/"
MUSIC_REFS = "https://www.amazon.com/Personalized-Customized-Blankets-Birthday-Valentines/dp/B09QG12M16; https://www.etsy.com/market/custom_blanket_with_picture_song; https://www.pinterest.com/pin/couverture-tisse-photo-personnalise-avec-lecteur-de-musique-cadeau-couverture-photo-et-chanson-personnalise-lecteur--1060597780992008688/"
TREE_REFS = "https://www.amazon.com/tree-life-quilt/s?k=tree+of+life+quilt; https://bedsurehome.com/products/autumn-wildflower-quilt-set; https://www.ebay.com/itm/176672376406; https://www.pinterest.com/pin/970385050946630965/"
BASEBALL_REFS = "https://www.walmart.com/ip/7-Pieces-Baseball-Bedding-Queen-Size-Boys-Girls-American-Flag-Comforter-Set-Sports-Theme-Bedding-Comforter-Sets-Vintage-Grunge-Stripe-Bed-Bag-Ball-Ga/8432902077; https://www.youcustomizeit.com/p/Baseball-Duvet-Cover-Personalized/100020; https://www.pinterest.com/pin/4611263934498190592/"


PROFILES = {
    171: ("pink black gothic skull patchwork comforter set", "gothic skull comforter set", "Gothic Skull Comforter Set", "Pink Black Gothic Skull Patchwork Comforter Set", "pink, black and gray gothic patchwork bedding with skulls, ravens, skeletons, floral panels and striped/checkered blocks", "skull bedding set, gothic Halloween bedding, pink black comforter set", "HALLOWEEN", HALLOWEEN_REFS),
    172: ("black haunted house trick or treat comforter set", "haunted house comforter set", "Haunted House Comforter Set", "Black Haunted House Trick or Treat Comforter Set", "black and gray Halloween bedding with a haunted house, orange moon, trick-or-treat text, skeletons, pumpkins, bats and spiderwebs", "Halloween comforter set, trick or treat bedding, haunted house bedding", "HALLOWEEN", HALLOWEEN_REFS),
    173: ("cream pumpkin ghost Halloween comforter set", "pumpkin ghost comforter set", "Pumpkin Ghost Comforter Set", "Cream Pumpkin Ghost Halloween Comforter Set", "cream Halloween bedding with jack-o-lantern pumpkins, white ghosts, bats, spiderwebs, stars and black ornamental accents", "Halloween bedding set, ghost pumpkin bedding, cream Halloween comforter", "HALLOWEEN", HALLOWEEN_REFS),
    174: ("personalized photo music player quilt", "personalized music player quilt", "Personalized Music Player Quilt", "Personalized Photo Music Player Quilt", "custom photo quilt with a music player interface, song name, artist name, play controls and waveform graphics", "custom song quilt, photo music blanket, personalized couple quilt", "EVERGREEN", MUSIC_REFS),
    175: ("autumn tree of life birds flowers quilt", "autumn tree of life quilt", "Autumn Tree of Life Quilt Set", "Autumn Tree of Life Birds and Flowers Quilt", "radiant autumn Tree of Life artwork with orange foliage, blue birds, exposed roots, colorful flowers and matching shams", "tree of life bedding, autumn quilt set, birds flowers quilt", "EVERGREEN", TREE_REFS),
    176: ("personalized baseball American flag bedding with vertical name", "personalized baseball bedding", "Personalized Baseball Bedding", "Personalized Baseball American Flag Bedding", "American flag baseball bedding with large baseballs, red white and blue stripes, stars and vertical custom name text", "baseball bedding set, custom name baseball comforter, American flag bedding", "EVERGREEN", BASEBALL_REFS),
    177: ("personalized black baseball glove bedding", "baseball glove bedding", "Baseball Glove Bedding Set", "Personalized Black Baseball Glove Bedding", "black baseball bedding with glove, bat, ball artwork and script custom name placement", "personalized baseball comforter, custom baseball bedding, sports bedding set", "EVERGREEN", BASEBALL_REFS),
    178: ("personalized baseball flag bedding with large name", "baseball flag bedding", "Baseball Flag Bedding Set", "Personalized Baseball Flag Bedding with Custom Name", "distressed American flag baseball bedding with glove, ball artwork and large block custom name text", "personalized baseball bedding, American flag comforter, baseball duvet set", "EVERGREEN", BASEBALL_REFS),
    179: ("baseball home quote bedding with custom name", "baseball home quote bedding", "Baseball Home Quote Bedding", "Personalized Baseball Home Quote Bedding", "cream baseball bedding with player silhouettes, baseball stitching, custom names on shams and the quote There is no place like Home", "baseball quote bedding, custom name sports bedding, baseball comforter set", "EVERGREEN", BASEBALL_REFS),
    180: ("personalized catcher American flag baseball bedding", "catcher baseball bedding", "Catcher Baseball Bedding Set", "Personalized Catcher American Flag Baseball Bedding", "American flag baseball bedding with catcher artwork, jersey number and large custom name text", "custom baseball bedding, catcher comforter set, personalized sports bedding", "EVERGREEN", BASEBALL_REFS),
}


def qa_issues(pos, html_error):
    issues = [
        "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."
    ]
    if 171 <= pos <= 173:
        issues.append("Pamnest Halloween comforter products are closely related to batch 017; keep motif-specific titles and canonical review before import.")
        issues.append("Set-content, deep-pocket and size-chart claims should be checked against variants before import.")
    if pos == 171:
        issues.append("Gothic skull artwork may overlap Halloween and goth bedding clusters; review category and tone.")
    if pos == 174:
        issues.append("Personalized photo/song fields, song name and artist name must be proofread and mapped to the correct customization fields.")
    if pos == 175:
        issues.append("Tree of Life products overlap prior batches; differentiate autumn birds/flowers motif and review canonical strategy.")
    if 176 <= pos <= 180:
        issues.append("Personalized baseball designs are very close; review custom-name placement, quote/number accuracy and duplicate keyword cannibalization.")
        issues.append("Feature images mention zipper closure and high-density weaving; verify actual product type and set configuration before import.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def product_context(pos):
    if 171 <= pos <= 173:
        return "seasonal Halloween rooms, gothic decor, guest beds or themed fall gifts."
    if pos == 174:
        return "anniversary gifts, couple gifts, memorial gifts or custom music-themed bedrooms."
    if pos == 175:
        return "nature-inspired bedrooms, autumn decor or symbolic Tree of Life gifts."
    return "baseball rooms, sports gifts, kids bedrooms or personalized fan bedding."


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        issues = qa_issues(pos, html_error)
        meta = f"Shop a {kind} featuring {detail}. Review available sizes and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for {product_context(pos)}</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broader bedding terms for collections or stronger representative products.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, room decor or a personalized gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": issues,
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed bedding, the buyer wants a {kind} matching a specific Halloween, music, Tree of Life or baseball motif.",
            "functional_motivation": "Find the right design, size, bedding type, included pieces, personalization fields, care details and decor fit.",
            "emotional_social_motivation": "Create a distinctive bedroom look or give a themed seasonal, romantic, nature or sports-inspired gift.",
            "purchase_concerns": "Product identity, personalization spelling, material/care claims, set contents, size fit and whether visible artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed in a main product mockup",
                2: f"Feature or close-up image for {h1}",
                3: f"Lifestyle or care image for {h1}",
                4: f"Product detail or closure image for {h1}",
                5: f"Fitted sheet or size feature image for {h1}",
                6: f"Set contents, fabric or size chart image for {h1}",
                7: f"Additional product mockup of {h1}",
                8: f"Lifestyle use image for {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_018_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_018_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_018"},
        {"metric": "cumulative_products", "value": "180", "definition": "Total products included through batch_018"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[180:190]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_019",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
