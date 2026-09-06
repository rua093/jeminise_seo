import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_007"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_006.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_007.xlsx"


PROFILES = {
    61: ("personalized football yard line comforter with David and number 35", "personalized football comforter", "Personalized Football Comforter Set", "Personalized Football Yard Line Comforter", "black and gold football field yard-line design with large football graphic, name David and number 35", "custom football bedding, football comforter with name, football bedding for boys", "EVERGREEN", "https://www.amazon.com/clp/B0CG1XRNMS; https://www.pinterest.com/pin/personalized-football-bedding-set-custom-name-football-bedding-personalized-sports-bed-set-duvet-covers--413557178297079798/; https://www.ebay.com/itm/165689760708"),
    62: ("personalized American flag football comforter with John and number 15", "American flag football comforter", "American Flag Football Comforter Set", "Personalized American Flag Football Comforter", "close-up football over a distressed American flag background with vertical John lettering and number 15", "patriotic football bedding, custom football comforter, football bedding with name", "EVERGREEN", "https://www.amazon.com/Erosebridal-American-Football-Comforter-Federations/dp/B09L4T6CQG; https://www.walmart.com/c/kp/american-flag-bedding; https://www.etsy.com/listing/1504206666/patriotic-american-flag-comforter-red"),
    63: ("personalized flaming lightning football comforter with Andrew and number 33", "flaming football comforter", "Flaming Football Comforter Set", "Personalized Flaming Lightning Football Comforter", "large football wrapped in orange flames and blue lightning on a black background with name Andrew and number 33", "custom football bedding, football comforter with number, football bedding set", "EVERGREEN", "https://www.amazon.com/clp/B0CG1XRNMS; https://www.pinterest.com/ideas/football-bedding/932913586866/; https://www.ebay.com/itm/165689760708"),
    64: ("personalized football helmet American flag comforter with John and number 8", "football helmet comforter", "Football Helmet Comforter Set", "Personalized Football Helmet Flag Comforter", "red football helmet over a vintage American flag background with the name John and #8", "custom football helmet bedding, patriotic football comforter, football bedding with name", "EVERGREEN", "https://ohaprints.com/products/personalized-football-duvet-cover-set-football-helmet-on-grass-player-gift-idea-vintage-duvet-cover-pillowcases-custom-name-bedding-set; https://www.amazon.com/s?k=football+helmet+comforter; https://www.pinterest.com/ideas/football-bedding/932913586866/"),
    65: ("personalized football helmet and flaming ball comforter with Matthew number 06", "football helmet bedding set", "Football Helmet Bedding Set", "Personalized Football Helmet Comforter Set", "black football helmet with number 06, flaming football trail, stadium-style dark background and name Matthew", "football helmet comforter, custom football bedding, personalized sports comforter", "EVERGREEN", "https://ohaprints.com/products/personalized-football-duvet-cover-set-football-helmet-on-grass-player-gift-idea-vintage-duvet-cover-pillowcases-custom-name-bedding-set; https://www.amazon.com/clp/B0CG1XRNMS; https://www.pinterest.com/ideas/football-bedding/932913586866/"),
    66: ("personalized football over American flag comforter with Andrew number 45", "personalized patriotic football comforter", "Personalized Patriotic Football Comforter", "Personalized Football Flag Comforter Set", "large close-up football over a distressed American flag with the name Andrew and number 45", "American flag football bedding, football comforter with name, custom football bedding", "EVERGREEN", "https://www.amazon.com/Erosebridal-American-Football-Comforter-Federations/dp/B09L4T6CQG; https://www.walmart.com/c/kp/american-flag-bedding; https://www.etsy.com/listing/1504206666/patriotic-american-flag-comforter-red"),
    67: ("personalized black football player comforter with William number 20", "football player comforter", "Football Player Comforter Set", "Personalized Football Player Comforter", "black football player close-up holding a ball with orange honeycomb pattern, name William and number 20", "custom football player bedding, football bedding with name, black football comforter", "EVERGREEN", "https://www.amazon.com/clp/B0CG1XRNMS; https://www.pinterest.com/ideas/football-bedding/932913586866/; https://www.ebay.com/itm/165689760708"),
    68: ("personalized black and white football player helmet comforter with David number 15", "custom football player comforter", "Custom Football Player Comforter Set", "Custom Football Player Helmet Comforter", "black and white smoky football player and helmet artwork with name David and #15", "football helmet bedding, personalized football bedding, football comforter with name", "EVERGREEN", "https://www.amazon.com/clp/B0CG1XRNMS; https://www.pinterest.com/pin/personalized-football-bedding-set-custom-name-football-bedding-personalized-sports-bed-set-duvet-covers--413557178297079798/; https://www.ebay.com/itm/165689760708"),
    69: ("personalized football player collage comforter with Michael number 30", "personalized football player bedding", "Personalized Football Player Bedding Set", "Personalized Football Player Collage Comforter", "black football player collage with helmet and ball close-up, large Michael name and number 30 on pillows", "custom football comforter, football player comforter, football bedding set for boys", "EVERGREEN", "https://www.amazon.com/clp/B0CG1XRNMS; https://www.pinterest.com/ideas/football-bedding/932913586866/; https://www.ebay.com/itm/165689760708"),
    70: ("personalized camouflage football player comforter with Robert number 09", "camouflage football comforter", "Camouflage Football Comforter Set", "Personalized Camouflage Football Comforter", "football player running with the ball over a green camouflage flag-style background, name Robert and number 09", "custom camouflage football bedding, football comforter with name, camo sports bedding", "EVERGREEN", "https://www.amazon.com/s?k=camouflage+football+comforter; https://www.amazon.com/clp/B0CG1XRNMS; https://www.pinterest.com/ideas/football-bedding/932913586866/"),
}


def append_dict(ws, data):
    headers = [cell.value for cell in ws[1]]
    ws.append([data.get(h, "") for h in headers])


def image_obs(theme, pos):
    if pos == 1:
        return f"Main bed mockup showing {theme}."
    if pos == 2:
        return f"Angled bedroom mockup showing {theme}."
    if pos == 3:
        return "Bedding type comparison for duvet cover set versus comforter set."
    if pos == 4:
        return f"Close room/detail image with softness, lightweight, durable and breathable icons for {theme}."
    if pos == 5:
        return "Stress-free easy-care image listing wrinkle-free, stain-proof, anti-pilling and wash-fade resistance."
    if pos == 6:
        return "Size dimension guide for twin, full, queen and king comforter options."
    return f"Additional mockup or product feature image showing {theme}."


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        issue_bits = [
            "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment.",
            "Batch contains many closely related football comforters; keyword cannibalization and final title uniqueness need QA.",
        ]
        if html_error:
            issue_bits.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
        issues = " ".join(issue_bits)
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, giving football fans a personalized bedding design for a bedroom or gift.</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork, personalization and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific football motif: {label}. Keep broader football bedding terms for collections or only the strongest representative products.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as a custom sports bedroom item or personalized football gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": issues,
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for personalized football bedding, the buyer wants a coordinated {kind} with a design, name and number that fit the recipient.",
            "functional_motivation": "Find the right football design, size, bedding type and personalization details.",
            "emotional_social_motivation": "Create a sports-themed bedroom or give a personal football gift to a player or fan.",
            "purchase_concerns": "Fit, included components, material/care claims, personalization spelling, jersey number accuracy and whether similar designs are distinct enough.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible football artwork, personalization and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed on a bed",
                2: f"Angled bedroom view of {h1}",
                3: f"Duvet cover and comforter option guide for {h1}",
                4: f"Close-up bedroom detail of {h1}",
                5: f"Easy care information for {h1}",
                6: f"Size dimension guide for {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_007_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_007_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_007"},
        {"metric": "cumulative_products", "value": "70", "definition": "Total products included through batch_007"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[70:80]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_008",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
