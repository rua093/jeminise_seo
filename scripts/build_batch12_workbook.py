import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_012"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_011.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_012.xlsx"


SOCCER_REFS = "https://www.etsy.com/listing/4379556971/personalized-girl-soccer-comforter; https://www.pinterest.com/pin/personalized-soccer-quilt-set-soccer-balls-quilt-blanket-with-pillowcases-custom-name-and-number-quilt-bedding-set--999939923509360353/; https://www.amazon.com/"
SOFTBALL_REFS = "https://www.etsy.com/market/softball_comforter_sets; https://www.pinterest.com/pin/personalized-softball-quilt-set-softball-quilt-blanket-with-pillowcases-custom-name-quilt-bedding-set--999939923509361679/; https://www.youcustomizeit.com/p/Softball-Comforters-Personalized/353832"
FOOTBALL_REFS = "https://www.amazon.com/Football-Personalized-Blanket-Pillowcases-Bedspread/dp/B0BFGNX4HW; https://macorner.co/products/american-football-custom-name-number-personalized-blanket-foo201101tuht; https://luvingift.com/football/"
TRUCK_REFS = "https://www.amazon.com/clp/B0DFQ6T1C5; https://www.amazon.com/Printluxe-Trucker-Tapestry-Comforters-Bedspreads/dp/B0DFQ9TMLY; https://www.etsy.com/listing/1091390349/personalized-semi-truck-baby-boy-blanket; https://jakescabsolutions.com/"
TREE_REFS = "https://www.amazon.com/clp/B0DNSVY2QJ; https://www.etsy.com/listing/4479152685/viking-bedding-set-yggdrasil-tree-of; https://www.walmart.ca/en/ip/Viking-Quilt-Bed-Set-Tree-Of-Life-Yggdrasil-Viking-Quilt-Bedding-Set/3OFTOZZSJ3FN; https://vikingsonsofodin.com/product/viking-bedding-set-viking-tree-of-life-yggdrasil-viking-bed-set-2/"


PROFILES = {
    111: ("personalized blue soccer player comforter with Daniel number 10", "personalized soccer player comforter", "Personalized Soccer Player Comforter", "Personalized Blue Soccer Player Comforter", "blue soccer ball artwork with Daniel name, number 10 and water-splash styling", "custom soccer bedding, soccer comforter with name, soccer comforter with number", "EVERGREEN", SOCCER_REFS),
    112: ("personalized soccer goal comforter with Tyler number 10", "soccer player goal comforter", "Soccer Player Goal Comforter Set", "Personalized Soccer Goal Comforter", "blue soccer ball in goal net artwork with Tyler name and number 10", "personalized soccer bedding, soccer goal bedding, soccer comforter with name", "EVERGREEN", SOCCER_REFS),
    113: ("personalized paint splatter soccer comforter with Matthew number 15", "paint splatter soccer comforter", "Paint Splatter Soccer Comforter", "Personalized Paint Splatter Soccer Comforter", "white comforter with multicolor paint splatter soccer ball artwork, Matthew name and number 15", "custom soccer bedding, soccer ball comforter, soccer comforter with name", "EVERGREEN", SOCCER_REFS),
    114: ("yellow black softball patchwork comforter for girls", "softball comforter set for girls", "Softball Comforter Set for Girls", "Yellow Black Softball Comforter Set", "yellow, black and gray softball patchwork artwork with batter silhouettes and matching pillowcases", "softball bedding, girls softball comforter, softball bedding set", "EVERGREEN", SOFTBALL_REFS),
    115: ("personalized football American flag comforter with Kevin number 20", "football American flag comforter", "Football American Flag Comforter", "Personalized Football American Flag Comforter", "large brown football over a distressed American flag background with Kevin name and number 20", "custom football bedding, patriotic football comforter, football comforter with name", "EVERGREEN", FOOTBALL_REFS),
    116: ("personalized blue semi truck geometric comforter", "semi truck geometric comforter", "Semi Truck Geometric Comforter", "Personalized Blue Semi Truck Geometric Comforter", "blue semi truck over black and blue geometric graphics with Your Name personalization area", "custom trucker bedding, personalized semi truck comforter, semi truck bedding with name", "EVERGREEN", TRUCK_REFS),
    117: ("colorful Tree of Life quilt set with rainbow leaves", "colorful tree of life quilt", "Colorful Tree of Life Quilt Set", "Colorful Tree of Life Quilt Set", "vivid rainbow Tree of Life artwork with flowing roots, multicolor leaves and matching shams", "tree of life bedding, Yggdrasil quilt set, colorful quilt bedding", "EVERGREEN", TREE_REFS),
    118: ("Celtic Tree of Life Yggdrasil quilt set in green and gold", "Celtic tree of life quilt", "Celtic Tree of Life Quilt Set", "Celtic Tree of Life Yggdrasil Quilt Set", "green Tree of Life and roots inside a Celtic ornamental oval frame with gold accents", "Yggdrasil bedding, Viking tree of life quilt, Celtic quilt set", "EVERGREEN", TREE_REFS),
    119: ("personalized red semi truck breaking through stone wall comforter", "red semi truck wall comforter", "Red Semi Truck Wall Comforter", "Personalized Red Semi Truck Wall Comforter", "red semi truck breaking through a gray stone wall with Your Name personalization area", "custom trucker bedding, semi truck comforter with name, truck bedding set", "EVERGREEN", TRUCK_REFS),
    120: ("personalized black semi truck breaking through stone wall comforter", "black semi truck wall comforter", "Black Semi Truck Wall Comforter", "Personalized Black Semi Truck Wall Comforter", "black semi truck breaking through a gray stone wall with Your Name personalization area", "custom trucker bedding, semi truck comforter with name, truck bedding set", "EVERGREEN", TRUCK_REFS),
}


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        issue_bits = [
            "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."
        ]
        if 111 <= pos <= 113:
            issue_bits.append("Closely related personalized soccer comforters should be reviewed for keyword cannibalization and title uniqueness.")
        if pos == 114:
            issue_bits.append("Current title mentions personalized name and number, but contact sheet does not show an obvious name/number field; personalization claim needs QA.")
        if pos == 115:
            issue_bits.append("Patriotic football comforter overlaps with broader football and American flag bedding terms; review product-level keyword specificity.")
        if pos == 116:
            issue_bits.append("Closely related semi truck products should be reviewed for cannibalization against other trucker bedding pages.")
        if pos in (117, 118):
            issue_bits.append("Tree of Life/Yggdrasil variants are very close; current meta titles appear similar and should be differentiated.")
        if pos in (119, 120):
            issue_bits.append("Breaking-through-wall semi truck variants are near duplicates; color-specific titles and canonical mapping should be reviewed.")
        if html_error:
            issue_bits.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
        issues = " ".join(issue_bits)
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for themed decor, sports rooms, trucker gifts or nature-inspired bedrooms.</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork, personalization and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broader bedding terms for collections or stronger representative products.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as a themed bedding item, personalized gift or room decor.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": issues,
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed or personalized bedding, the buyer wants a {kind} matching a specific design, name, number, message or hobby.",
            "functional_motivation": "Find the right design, size, bedding type, care details and personalization fit.",
            "emotional_social_motivation": "Create meaningful room decor or give a personalized sports, trucker, Viking or nature-inspired gift.",
            "purchase_concerns": "Personalization spelling, product identity, material/care claims, design distinction and whether the visible artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork, personalization and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed in a main product mockup",
                2: f"Secondary mockup of {h1}",
                3: f"Product option, feature or lifestyle image for {h1}",
                4: f"Detail or feature image for {h1}",
                5: f"Additional product mockup or care information for {h1}",
                6: f"Care, size or product feature image for {h1}",
                7: f"Additional product mockup of {h1}",
                8: f"Lifestyle use image for {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_012_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_012_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_012"},
        {"metric": "cumulative_products", "value": "120", "definition": "Total products included through batch_012"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[120:130]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_013",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
