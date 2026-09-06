import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_021"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_020.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_021.xlsx"


BASEBALL_REFS = "https://www.amazon.com/Custom-Baseball-Bedding-Boys-Adults/dp/B0FHWPH2F3; https://www.walmart.com/c/kp/baseball-bedding-set; https://www.youcustomizeit.com/p/Baseball-Jersey-Comforters-Personalized/142431; https://ohaprints.com/products/personalized-baseball-duvet-cover-set-baseball-player-unique-gift-idea-blue-duvet-cover-pillowcases-custom-name-number-bedding-set; https://www.etsy.com/listing/4339325702/personalized-name-baseball-duvet-cover"
BASKETBALL_REFS = "https://www.amazon.com/-/he/dp/B0FXR9M3MM; https://www.etsy.com/market/personalized_basketball_comforter; https://ohaprints.com/products/ohaprints-quilt-bed-set-pillowcase-basketball-boy-player-fan-unique-gift-grey-custom-personalized-name-number-blanket-bedspread-bedding-423; https://www.walmart.com/c/kp/basketball-bedding; https://www.etsy.com/listing/1689155537/personalized-basketball-blanket-custom; https://www.zazzle.com/basketball%2Bblankets"


PROFILES = {
    201: ("American flag baseball bedding with wood plank and script name", "baseball flag bedding", "Baseball Flag Bedding Set", "American Flag Baseball Bedding with Script Name", "American flag baseball bedding with wood plank texture, large baseball, script custom name and numbered ball shams", "personalized baseball bedding, custom name baseball comforter, patriotic baseball bedding", "EVERGREEN", BASEBALL_REFS),
    202: ("neon baseball player duvet cover with custom name", "neon baseball duvet cover", "Neon Baseball Duvet Cover", "Neon Baseball Player Duvet Cover with Custom Name", "black baseball duvet cover with neon green player silhouettes, custom name and jersey number", "personalized baseball duvet, custom baseball bedding, baseball player comforter", "EVERGREEN", BASEBALL_REFS),
    203: ("basketball above hoop comforter with name and number", "basketball hoop comforter", "Basketball Hoop Comforter Set", "Basketball Above Hoop Comforter with Name and Number", "basketball comforter showing a ball above the hoop and net with custom name and jersey number", "personalized basketball bedding, basketball net comforter, custom sports bedding", "EVERGREEN", BASKETBALL_REFS),
    204: ("blue red basketball hoop close-up comforter", "basketball hoop bedding", "Basketball Hoop Bedding Set", "Blue Red Basketball Hoop Bedding with Name and Number", "blue and red basketball bedding with close-up hoop, net, ball, vertical custom name and jersey number", "personalized basketball comforter, custom basketball bedding, basketball room decor", "EVERGREEN", BASKETBALL_REFS),
    205: ("basketball paint splash comforter with name and number", "basketball paint splash comforter", "Basketball Paint Splash Comforter", "Basketball Paint Splash Comforter with Name and Number", "red and blue paint-splash basketball comforter with large ball graphic, custom name and jersey number", "personalized basketball bedding, basketball comforter set, custom name sports bedding", "EVERGREEN", BASKETBALL_REFS),
    206: ("basketball below net blanket with name and number", "basketball net blanket", "Basketball Net Blanket", "Basketball Below Net Blanket with Name and Number", "black basketball blanket with hoop, net, ball under the rim, vertical custom name and jersey number", "personalized basketball blanket, custom basketball throw, basketball gift blanket", "EVERGREEN", BASKETBALL_REFS),
    207: ("basketball court hoop comforter with custom name", "basketball court comforter", "Basketball Court Comforter Set", "Basketball Court Hoop Comforter with Custom Name", "blue-black basketball comforter with court lines, hoop, large ball and custom name", "personalized basketball bedding, basketball player comforter, custom sports bedding", "EVERGREEN", BASKETBALL_REFS),
    208: ("black basketball hoop comforter with vertical name", "black basketball comforter", "Black Basketball Comforter Set", "Black Basketball Hoop Comforter with Vertical Name", "black basketball comforter with orange ball under hoop, grid background, vertical custom name and jersey number", "personalized basketball bedding, basketball net comforter, custom name bedding", "EVERGREEN", BASKETBALL_REFS),
    209: ("burning flames basketball comforter with name and number", "flame basketball comforter", "Flame Basketball Comforter Set", "Burning Flames Basketball Comforter with Name and Number", "basketball comforter with oversized close-up ball, burning flame background, custom name and jersey number", "personalized basketball bedding, fire basketball comforter, custom sports bedding", "EVERGREEN", BASKETBALL_REFS),
    210: ("dark basketball close-up blanket with name and number", "basketball name blanket", "Basketball Name Blanket", "Dark Basketball Close-Up Blanket with Name and Number", "dark basketball blanket with close-up ball, gray player silhouettes, custom name and jersey number", "personalized basketball blanket, custom basketball throw, basketball player blanket", "EVERGREEN", BASKETBALL_REFS),
}


def qa_issues(pos, html_error):
    issues = [
        "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment.",
        "Personalized sports products are close variants; review custom-name placement, jersey number accuracy and duplicate keyword cannibalization before import.",
        "Feature images mention duvet cover/comforter/blanket options, microfiber, easy care and size dimensions; verify actual product type, variants and set configuration before deployment.",
    ]
    if pos in (201, 202):
        issues.append("Baseball designs overlap batches 018-020; keep flag/wood and neon-player motifs distinct.")
    if 203 <= pos <= 210:
        issues.append("Basketball designs share name-and-number intent; review canonical strategy and avoid repeating the same primary keyword across pages.")
    if pos in (206, 210):
        issues.append("These are blankets rather than comforter sets; verify size chart and product-type fields before import.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def product_context(pos):
    if pos in (201, 202):
        return "baseball rooms, sports gifts, kids bedrooms or personalized fan bedding."
    return "basketball rooms, player gifts, sports bedrooms or personalized fan bedding."


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for {product_context(pos)}</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broad personalized sports bedding terms for collections or the strongest representative product.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, a sports room accent or a personalized gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": qa_issues(pos, html_error),
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for personalized sports bedding, the buyer wants a {kind} matching a baseball or basketball motif, custom name and number.",
            "functional_motivation": "Find the right sport design, size, personalization fields, care details and room fit.",
            "emotional_social_motivation": "Create a sports-themed bedroom or give a custom gift to a player or fan.",
            "purchase_concerns": "Personalization spelling, jersey number, product identity, material/care claims, size fit and whether artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed in a main product mockup",
                2: f"Secondary product mockup for {h1}",
                3: f"Feature or size image for {h1}",
                4: f"Fabric or product feature image for {h1}",
                5: f"Lifestyle, care or size image for {h1}",
                6: f"Bedding type or blanket feature image for {h1}",
                7: f"Size dimension image for {h1}",
                8: f"Lifestyle use image for {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_021_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_021_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_021"},
        {"metric": "cumulative_products", "value": "210", "definition": "Total products included through batch_021"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[210:220]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_022",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
