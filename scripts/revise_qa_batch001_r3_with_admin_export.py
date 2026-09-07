import csv
import hashlib
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_001_r3"
QA_RUN_ID = "20260907_114114"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_001_r2" / "SEO_Product_Optimization_qa_batch_001_r2.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
CUSTOMIZER_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_audit.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_001_r3.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_001_r3.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    1: {
        "title": "Cardinal Sunflower Autumn Quilt Set",
        "meta_title": "Cardinal Sunflower Autumn Quilt Set",
        "meta_description": "Bring fall color to the bedroom with red cardinals, large sunflowers, autumn leaves, size choices and optional shams.",
        "primary": "cardinal sunflower autumn quilt set",
        "secondary": "autumn cardinal quilt, sunflower cardinal bedding, cardinal bird quilt set",
        "cluster": "cardinal autumn floral quilt",
        "intent_role": "Autumn cardinal and sunflower design page; internally separate from Christmas cardinal memorial quilts.",
        "detail": "red cardinals perched above large yellow sunflowers with autumn leaves",
        "options": "Choose quilt size; pillowcases are optional according to the current Shopify export.",
    },
    2: {
        "title": "Personalized Teal Softball Comforter Set",
        "meta_title": "Personalized Teal Softball Comforter Set",
        "meta_description": "Customize a teal softball comforter with name and number fields, coordinated pillowcases and sporty bedroom artwork.",
        "primary": "personalized softball comforter set",
        "secondary": "teal softball bedding set, custom softball comforter, softball bedding with name and number",
        "cluster": "personalized softball comforter",
        "intent_role": "Softball comforter page with verified required name and number text fields.",
        "detail": "teal, black and white softball patchwork with player silhouettes, number artwork and Eat Sleep Softball text",
        "options": "Choose comforter size and pillowcases. Live Customizer shows required Custom Your Name and Custom Your Number text fields, with instructions to enter NO if not wanted.",
    },
    3: {
        "title": "Christmas Cardinal Memorial Quilt Set",
        "meta_title": "Christmas Cardinal Memorial Quilt Set",
        "meta_description": "Create a comforting holiday bed with red cardinals, a decorated Christmas tree, memorial text, size choices and optional shams.",
        "primary": "Christmas cardinal memorial quilt set",
        "secondary": "cardinal Christmas quilt bedding, I am always with you quilt, cardinal remembrance bedding",
        "cluster": "Christmas cardinal memorial tree quilt",
        "intent_role": "Christmas tree cardinal memorial page; internally separate from rose cardinal remembrance design.",
        "detail": "red cardinals beside a decorated Christmas tree with I Am Always with You memorial text",
        "options": "Choose quilt size; pillowcases are optional according to the current Shopify export.",
    },
    4: {
        "title": "Cardinal Roses Memorial Quilt Set",
        "meta_title": "Cardinal Roses Memorial Quilt Set",
        "meta_description": "Choose a memorial-style cardinal quilt with a cardinal pair, red roses, remembrance text, size choices and optional shams.",
        "primary": "cardinal roses memorial quilt set",
        "secondary": "cardinal remembrance quilt, cardinal rose bedding, I am always with you cardinal quilt",
        "cluster": "cardinal roses remembrance quilt",
        "intent_role": "Rose cardinal remembrance page; internally separate from Christmas-tree cardinal memorial design.",
        "detail": "a male and female cardinal pair on a branch with red roses and I Am Always with You text",
        "options": "Choose quilt size; pillowcases are optional according to the current Shopify export.",
    },
    5: {
        "title": "Colorful Cat Patchwork Quilt Set",
        "meta_title": "Colorful Cat Patchwork Quilt Set",
        "meta_description": "Style a cat lover's room with a colorful cat face patchwork quilt, feature graphics, size choices and optional pillowcases.",
        "primary": "colorful cat patchwork quilt set",
        "secondary": "cat patchwork bedding set, colorful cat quilt, cat face quilt set",
        "cluster": "colorful cat face patchwork quilt",
        "intent_role": "Colorful cat-face design page; no chicken evidence or farmhouse-only query should be used.",
        "detail": "a large colorful cat face built from bright patchwork blocks",
        "options": "Choose quilt size; pillowcases are purchased separately according to the current Shopify export.",
    },
    6: {
        "title": "Geometric Cat Patchwork Quilt Set",
        "meta_title": "Geometric Cat Patchwork Quilt Set",
        "meta_description": "Add a calmer cat design with a geometric sitting cat quilt, warm patchwork colors, size choices and optional pillowcases.",
        "primary": "geometric cat patchwork quilt set",
        "secondary": "cat patchwork bedding set, sitting cat quilt, cat quilt bedding",
        "cluster": "geometric sitting cat quilt",
        "intent_role": "Geometric sitting-cat design page; internally separate from colorful cat-face patchwork design.",
        "detail": "a sitting cat in geometric patchwork colors on a beige leaf background",
        "options": "Choose quilt size; pillowcases are purchased separately according to the current Shopify export.",
    },
    7: {
        "title": "Celtic Fantasy Tree Quilt Set",
        "meta_title": "Celtic Fantasy Tree Quilt Set",
        "meta_description": "Add fantasy woodland style with a twisting Celtic tree quilt, exposed roots, knotwork border, size choices and pillowcase options.",
        "primary": "Celtic fantasy tree quilt set",
        "secondary": "fantasy tree quilt bedding, Celtic tree quilt, knotwork tree bedding",
        "cluster": "Celtic fantasy tree quilt",
        "intent_role": "Twisting fantasy-tree page; internally separate from green Tree of Life medallion page.",
        "detail": "a twisting tree trunk with exposed roots, teal night sky accents and a Celtic knot border",
        "options": "Choose quilt size and pillowcase quantity according to the current Shopify export.",
    },
    8: {
        "title": "Celtic Tree of Life Quilt Set",
        "meta_title": "Celtic Tree of Life Quilt Set",
        "meta_description": "Give the bedroom a Celtic focal point with a green Tree of Life medallion quilt, knotwork border, sizes and optional shams.",
        "primary": "Celtic Tree of Life quilt set",
        "secondary": "Tree of Life bedding set, Celtic knotwork quilt, green tree quilt set",
        "cluster": "green Celtic Tree of Life quilt",
        "intent_role": "Green Tree of Life medallion page; internally separate from twisting fantasy-tree page.",
        "detail": "a green Tree of Life medallion with Celtic knotwork border",
        "options": "Choose quilt size; pillowcases are optional according to the current Shopify export.",
    },
    9: {
        "title": "Farmhouse Chicken Patchwork Quilt Set",
        "meta_title": "Farmhouse Chicken Patchwork Quilt Set",
        "meta_description": "Add country charm with a farmhouse chicken quilt featuring hens, nest-and-egg artwork, patchwork accents and optional shams.",
        "primary": "farmhouse chicken patchwork quilt set",
        "secondary": "chicken bedding set, country chicken quilt, hen patchwork quilt",
        "cluster": "farmhouse chicken patchwork quilt",
        "intent_role": "Farmhouse chicken and hen design page; not used for cat-quilt evidence.",
        "detail": "farmhouse chicken and hen artwork with nest-and-egg details and black patchwork accents",
        "options": "Choose quilt size; pillowcases are purchased separately according to the current Shopify export.",
    },
    10: {
        "title": "Personalized God Says I Am Christian Bedding Set",
        "meta_title": "Personalized God Says I Am Christian Bedding Set",
        "meta_description": "Customize a God Says I Am Christian bedding set with an Enter Name field, birth-month flowers, size choices and pillowcase options.",
        "primary": "personalized God Says I Am Christian bedding set",
        "secondary": "Christian comforter set, God Says I Am bedding, personalized Bible verse bedding",
        "cluster": "personalized Christian affirmation bedding",
        "intent_role": "Personalized God Says I Am bedding page with verified Enter Name field up to 13 characters.",
        "detail": "God Says I Am Christian affirmation artwork with a custom name example, flowers and butterflies",
        "options": "Choose product type and size, pillowcases and optional flat sheet. Live Customizer shows a required Enter Name field with a 13-character limit.",
    },
}

CAT_SERP_URLS = "https://www.walmart.com/ip/20788770071; https://www.walmart.com/ip/20076702698"
CARDINAL_TREE_SERP = "https://www.target.com/s/cardinal%2Bquilt; https://www.marcielobedding.com/products/marcelo-3-pcs-winter-cardinals-christmas-quilt-bedspread-set-decor"
CARDINAL_ROSES_SERP = "https://www.reddit.com/r/quilting/comments/19fiv3f; https://www.reddit.com/r/quilting/comments/166n7be"


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def safe_len(value):
    return len(text(value))


def make_description(pos, admin):
    data = PRODUCT_UPDATES[pos]
    admin_type = admin.get("Type") or "Bedding"
    admin_title = admin.get("Title") or data["title"]
    return (
        f"<p>{data['detail'].capitalize()} gives this Jeminise {admin_type.lower()} a clear product-specific look for themed bedroom decor or gifting.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork focus: {data['detail']}.</li>"
        f"<li>Current Shopify export title: {admin_title}.</li>"
        "<li>Gallery includes the main bed mockup plus detail, feature, component or size graphics where shown.</li></ul>"
        "<h3>Options and Care</h3>"
        f"<ul><li>{data['options']}</li>"
        "<li>Use the live product selectors to confirm the final configuration before checkout.</li>"
        "<li>Care and fabric claims are kept only where shown in the current Shopify export, storefront data or gallery feature panels.</li></ul>"
    )


def load_admin_export(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)

    admin = {}
    image_by_handle_url = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            src = row.get("Image Src", "")
            if not src:
                continue
            images[normalize_url(src)] = {
                "src": src,
                "position": row.get("Image Position", ""),
                "alt": row.get("Image Alt Text", ""),
            }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Body (HTML)": first.get("Body (HTML)", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
            "rows": len(rows),
        }
        image_by_handle_url[handle] = images
    return admin, image_by_handle_url


def load_qa():
    data = json.loads(QA_DATASET.read_text(encoding="utf-8"))
    products = data["QA_Products"]
    product_keys = [item["product_key"] for item in products]
    handle_by_key = {item["product_key"]: item["handle"] for item in products}
    pos_by_key = {item["product_key"]: int(item["inventory_position"]) for item in products}

    image_fix_by_key = {}
    observation_by_key = {}
    image_number_by_media = {}
    for image in data["QA_Images"]:
        key = (image["product_key"], text(image["media_id"]))
        observation_by_key[key] = image["qa_observation"]
        image_number_by_media[key] = image.get("image_number")

    for issue in data["QA_Issues"]:
        if issue["field"].startswith("image_") and issue["recommended_fix"]:
            matching = None
            for image in data["QA_Images"]:
                if issue["qa_image_key"] == image["qa_image_key"]:
                    matching = (image["product_key"], text(image["media_id"]))
                    break
            if matching:
                image_fix_by_key[matching] = issue["recommended_fix"]

    issue_refs_by_key = defaultdict(list)
    for issue in data["QA_Issues"]:
        if issue["product_key"]:
            issue_refs_by_key[issue["product_key"]].append(issue["issue_id"])

    customizer = json.loads(CUSTOMIZER_AUDIT.read_text(encoding="utf-8"))
    customizer_by_pos = {int(item["inventory_position"]): item for item in customizer}
    return {
        "data": data,
        "product_keys": product_keys,
        "handle_by_key": handle_by_key,
        "pos_by_key": pos_by_key,
        "image_fix_by_key": image_fix_by_key,
        "observation_by_key": observation_by_key,
        "image_number_by_media": image_number_by_media,
        "issue_refs_by_key": issue_refs_by_key,
        "customizer_by_pos": customizer_by_pos,
    }


def compact_observation_to_alt(observation, fallback):
    value = re.sub(r"[:;].*$", "", text(observation)).strip()
    value = re.sub(r"\s+", " ", value)
    if not value:
        value = fallback
    return value[:125]


def update_workbook():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)

    qa = load_qa()
    handles = set(qa["handle_by_key"].values())
    admin, image_admin = load_admin_export(handles)
    now = datetime.now().astimezone().isoformat()
    admin_hash = sha256(ADMIN_EXPORT)

    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        update = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]

        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or ws.cell(row_num, idx["product_type"]).value

        ws.cell(row_num, idx["primary_keyword"]).value = update["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = update["secondary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = update["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"Buyer intent targets {update['cluster']} for US English product search; no search-volume claim is made."
        )
        ws.cell(row_num, idx["title_proposed"]).value = update["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = update["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = safe_len(update["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = update["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = safe_len(update["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = make_description(pos, admin_row)
        ws.cell(row_num, idx["meta_keyword"]).value = update["primary"]
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R3 after Sang re-QA r2: applied QA_Issues for batch 001, used products_export_1.csv "
            f"(sha256:{admin_hash}) as admin baseline, fixed image alt/observations, richer product-specific "
            "description, customizer wording where verified, and keyword/evidence mapping. Still NEEDS_REVIEW; no APPROVED/import."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matched = 0
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pk = None
        for product_key, product_handle in qa["handle_by_key"].items():
            if product_handle == handle:
                pk = product_key
                break
        if pk not in product_key_set:
            continue
        media_id = text(ws.cell(row_num, idx["media_id"]).value)
        key = (pk, media_id)
        pos = qa["pos_by_key"][pk]
        update = PRODUCT_UPDATES[pos]
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = image_admin.get(handle, {}).get(normalize_url(image_url))

        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matched += 1

        observation = qa["observation_by_key"].get(key)
        if observation:
            ws.cell(row_num, idx["observed_visual_details"]).value = observation

        proposed_alt = qa["image_fix_by_key"].get(key)
        if not proposed_alt:
            proposed_alt = compact_observation_to_alt(observation, f"{update['title']} product image")
        ws.cell(row_num, idx["alt_proposed"]).value = proposed_alt[:125]
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R3: QA r2 observation/alt applied; current admin alt and image URL matched from products_export_1.csv where possible. "
            "Requires Sang re-QA before approval."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        update = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        if role == "PRIMARY":
            keyword = update["primary"]
        else:
            secondaries = [item.strip() for item in update["secondary"].split(",")]
            keyword = secondaries[min(max(keyword_rows, 0), len(secondaries) - 1)] if secondaries else update["primary"]
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = update["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "Product"
        ws.cell(row_num, idx["decision_reason"]).value = update["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = "Sang re-QA r2 SERP evidence + products_export_1.csv admin baseline"
        ws.cell(row_num, idx["checked_at"]).value = now
        if pos == 5:
            ws.cell(row_num, idx["representative_SERP_URLs"]).value = CAT_SERP_URLS
        elif pos == 3:
            ws.cell(row_num, idx["representative_SERP_URLs"]).value = CARDINAL_TREE_SERP
        elif pos == 4:
            ws.cell(row_num, idx["representative_SERP_URLs"]).value = CARDINAL_ROSES_SERP
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = update["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = "R3 separates sibling product intent by visible motif and verified option set."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r3"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        update = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {update['cluster']} as themed bedding or a gift."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {update['cluster']}, the buyer wants a product page that clearly shows the design, "
            "available size/component choices and any verified customization fields."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = "Confirm design, size, components, fabric/care claims and customization requirements before checkout."
        ws.cell(row_num, idx["emotional_social_motivation"]).value = "Match a room theme, sports identity, remembrance message, faith gift or animal/fantasy decor preference."
        ws.cell(row_num, idx["purchase_concerns"]).value = "Fit, included or separately purchased pillowcases, verified personalization fields, care claims and image accuracy."
        ws.cell(row_num, idx["source_refs"]).value = f"products_export_1.csv sha256:{admin_hash}; Sang QA r2 {QA_RUN_ID}"
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No search-volume source supplied; Shopify admin CSV baseline is now available for stored title/meta/body/image alt."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {update['primary']} with intent role: {update['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        update = PRODUCT_UPDATES[pos]
        pk = qa["product_keys"][pos - 1]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        ws.cell(row_num, idx["sources_accessed"]).value = (
            f"Storefront/product.js from Sang QA r2; products_export_1.csv sha256:{admin_hash}"
        )
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options="
            f"{admin_row['Option1 Name']}, {admin_row['Option2 Name']}, {admin_row['Option3 Name']}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={update['detail']}."
        )
        if pos == 5:
            ws.cell(row_num, idx["SERP_evidence_references"]).value = CAT_SERP_URLS
        elif pos == 3:
            ws.cell(row_num, idx["SERP_evidence_references"]).value = CARDINAL_TREE_SERP
        elif pos == 4:
            ws.cell(row_num, idx["SERP_evidence_references"]).value = CARDINAL_ROSES_SERP
        ws.cell(row_num, idx["factual_conflicts"]).value = "R3 applied Sang QA r2 fixes; needs independent QA recheck."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R3 uses QA-observed images, live Customizer audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    rows = [
        ("qa_batch_001_r3_revision", "r3", "Separate 10-product revision after Sang re-QA r2; source r2 workbook was not modified."),
        ("qa_batch_001_r3_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_001_r3_scope", "inventory positions 1-10", "No products outside qa_batch_001 were revised."),
        ("qa_batch_001_r3_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_001_r3_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_001_r3_admin_alt_matches", str(admin_alt_matched), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]
    for metric, value, definition in rows:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa_r2": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_001 only; inventory positions 1-10",
        "revision": "r3",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matched,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Sang QA lại qa_batch_001_r3 before any approval or import.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_001_r3",
                "",
                "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r2: `{SOURCE.relative_to(ROOT)}`",
                f"- QA r2 của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r3: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matched}/{revised_images}",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Bước tiếp theo: Sang QA lại `qa_batch_001_r3`.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    update_workbook()
