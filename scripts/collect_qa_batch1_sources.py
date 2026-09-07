import csv
import hashlib
import json
from datetime import datetime
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
QA_RUN_ID = "20260907_083534"
QA_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID
WORKBOOK = QA_DIR / "source_snapshot" / "SEO_Product_Optimization_through_batch_034.xlsx"
INVENTORY = ROOT / "seo_runs" / SHOP / RUN_ID / "inventory.csv"
EVIDENCE_PRODUCTS = ROOT / "seo_runs" / SHOP / RUN_ID / "evidence" / "products"


def sheet_rows(workbook, name):
    sheet = workbook[name]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    return [dict(zip(headers, (cell.value for cell in row))) for row in sheet.iter_rows(min_row=2)]


def fetch_live(url):
    headers = {"User-Agent": "Mozilla/5.0 independent SEO QA/1.0"}
    response = requests.get(url, headers=headers, timeout=45)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    canonical = soup.find("link", rel="canonical")
    description = soup.find("meta", attrs={"name": "description"})
    h1 = soup.find("h1")
    title = soup.find("title")
    product_response = requests.get(url + ".js", headers=headers, timeout=45)
    product_response.raise_for_status()
    product = product_response.json()
    return {
        "checked_at": datetime.now().astimezone().isoformat(),
        "status_code": response.status_code,
        "html_sha256": hashlib.sha256(response.content).hexdigest(),
        "title_element": title.get_text(" ", strip=True) if title else "",
        "h1": h1.get_text(" ", strip=True) if h1 else "",
        "meta_description": description.get("content", "").strip() if description else "",
        "canonical": canonical.get("href", "").strip() if canonical else "",
        "product_js": product,
    }


def main():
    with INVENTORY.open(newline="", encoding="utf-8-sig") as handle:
        first_ten = list(csv.DictReader(handle))[:10]
    keys = {row["product_key"] for row in first_ten}
    handles = {row["Handle"] for row in first_ten}

    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    products = [row for row in sheet_rows(wb, "SEO_Products") if row["product_key"] in keys]
    images = [row for row in sheet_rows(wb, "Image_Audit") if row["Handle"] in handles]
    evidence = [row for row in sheet_rows(wb, "Product_Evidence") if row["product_url"] in {r["product_url"] for r in first_ten}]
    keywords = [row for row in sheet_rows(wb, "Keyword_Map") if row["product_key"] in keys]
    buyers = [row for row in sheet_rows(wb, "Buyer_Search_Research") if row["product_key"] in keys]

    submitted = {
        "products": products,
        "images": images,
        "product_evidence": evidence,
        "keyword_map": keywords,
        "buyer_search_research": buyers,
    }
    (QA_DIR / "submitted_batch_data.json").write_text(json.dumps(submitted, ensure_ascii=False, indent=2), encoding="utf-8")

    live = []
    for row in first_ten:
        record = {"inventory": row}
        try:
            record["live"] = fetch_live(row["product_url"])
            record["fetch_error"] = ""
        except Exception as exc:
            record["live"] = {}
            record["fetch_error"] = str(exc)
        old_matches = list(EVIDENCE_PRODUCTS.glob(f"{row['inventory_position']}_*/evidence_summary.json"))
        if old_matches:
            record["research_snapshot"] = json.loads(old_matches[0].read_text(encoding="utf-8"))
        live.append(record)
    (QA_DIR / "live_source_comparison.json").write_text(json.dumps(live, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"products={len(products)} images={len(images)} keywords={len(keywords)} buyers={len(buyers)} live={len(live)}")
    for item in live:
        source = item["inventory"]
        current = item["live"]
        print(source["inventory_position"], current.get("status_code"), current.get("h1"), len(current.get("product_js", {}).get("media", [])), item["fetch_error"])


if __name__ == "__main__":
    main()
