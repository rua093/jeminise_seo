import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import revise_qa_batch002_r3_with_admin_export as base


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_005_r3"
QA_RUN_ID = "20260907_152903"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_005_r2" / "SEO_Product_Optimization_qa_batch_005_r2.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
CUSTOMIZER_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_audit.json"
CUSTOMIZER_FIELDS_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_fields_audit.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_005_r3.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_005_r3.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    41: {
        "title": "Personalized Football Player Comforter Set",
        "meta_title": "Personalized Football Player Comforter Set",
        "meta_description": "Customize a football player comforter with required name, optional number, product type, size, pillowcase and sheet-cover choices.",
        "primary": "personalized football player comforter",
        "secondary": "custom football player bedding, football comforter with name, football bedding with number",
        "cluster": "personalized football player comforter",
        "intent_role": "Verified required Enter Name up to 25 characters and optional Enter Number up to 5 characters.",
        "detail": "American sports football player running with ball comforter artwork; visible names and numbers are samples",
    },
    42: {
        "title": "Personalized Basketball Name Number Blanket",
        "meta_title": "Personalized Basketball Name Number Blanket",
        "meta_description": "Customize a basketball blanket with optional Custom Name and Custom Number fields, size choices and verified sample artwork.",
        "primary": "personalized basketball name number blanket",
        "secondary": "custom basketball blanket, basketball blanket with name, basketball blanket with number",
        "cluster": "personalized basketball name number blanket",
        "intent_role": "Verified optional Custom Name up to 200 characters and optional Custom Number up to 20 characters; COLON 06 and RASHAD 22 are samples.",
        "detail": "basketball player blanket artwork with separate COLON 06 and RASHAD 22 sample personalization images",
    },
    43: {
        "title": "Personalized God Says I Am Christian Blanket",
        "meta_title": "Personalized God Says I Am Christian Blanket",
        "meta_description": "Customize a God Says I Am Christian blanket with required name, affirmation artwork and size choices.",
        "primary": "personalized God Says I Am Christian blanket",
        "secondary": "Christian affirmation blanket with name, God Says I Am blanket, personalized Bible verse blanket",
        "cluster": "personalized God Says I Am Christian blanket",
        "intent_role": "Corrects the r2 unsupported visual-personalization claim: only required Customize Your Name up to 1000 characters is verified.",
        "detail": "God Says I Am Christian affirmation blanket artwork with Bible references and sample portrait-style artwork",
    },
    44: {
        "title": "Custom Cardinal Flowering Branches Quilt Set",
        "meta_title": "Custom Cardinal Flowering Branches Quilt Set",
        "meta_description": "Customize a cardinal flowering branches quilt with optional quilt text, cardinal artwork, quilt sizes and pillowcase options.",
        "primary": "custom cardinal flowering branches quilt",
        "secondary": "cardinal flowering branches quilt, custom cardinal quilt set, cardinal quilt with text",
        "cluster": "custom cardinal flowering branches quilt",
        "intent_role": "Verified optional Customize Your Quilt text field up to 1000 characters from customizer_fields_audit.",
        "detail": "cardinals on flowering branches quilt artwork with matching pillow sham visuals",
    },
    45: {
        "title": "Custom Colorful Tree of Life Quilt Set",
        "meta_title": "Custom Colorful Tree of Life Quilt Set",
        "meta_description": "Customize a colorful Tree of Life quilt set with optional quilt text, mosaic sunburst artwork, sizes and pillowcase options.",
        "primary": "custom colorful Tree of Life quilt set",
        "secondary": "colorful Tree of Life quilt, custom Yggdrasil quilt set, mosaic Tree of Life bedding",
        "cluster": "custom colorful Tree of Life quilt set",
        "intent_role": "Verified optional Customize Your Quilt text field up to 1000 characters from customizer_fields_audit.",
        "detail": "colorful mosaic Tree of Life sunburst quilt artwork",
    },
    46: {
        "title": "Custom Fantasy Tree of Life Quilt Set",
        "meta_title": "Custom Fantasy Tree of Life Quilt Set",
        "meta_description": "Customize a fantasy Tree of Life quilt set with optional quilt text, central eye artwork, sizes and pillowcase options.",
        "primary": "custom fantasy Tree of Life quilt set",
        "secondary": "fantasy Tree of Life quilt, custom Celtic quilt set, central eye Yggdrasil bedding",
        "cluster": "custom fantasy Tree of Life quilt set",
        "intent_role": "Verified optional Customize Your Quilt text field up to 1000 characters from customizer_fields_audit.",
        "detail": "fantasy Tree of Life quilt artwork with central eye and Celtic-style patterns",
    },
    47: {
        "title": "Custom Celtic Yggdrasil Tree of Life Quilt Set",
        "meta_title": "Custom Celtic Yggdrasil Tree of Life Quilt Set",
        "meta_description": "Customize a Celtic Yggdrasil Tree of Life quilt set with optional quilt text, intertwined roots, sizes and pillowcase options.",
        "primary": "custom Celtic Yggdrasil Tree of Life quilt",
        "secondary": "Celtic Yggdrasil quilt set, custom Tree of Life quilt, intertwined roots bedding",
        "cluster": "custom Celtic Yggdrasil Tree of Life quilt",
        "intent_role": "Verified optional Customize Your Quilt text field up to 1000 characters from customizer_fields_audit.",
        "detail": "Celtic Yggdrasil Tree of Life quilt artwork with intertwined roots",
    },
    48: {
        "title": "Personalized Trucker Prayer Comforter Set",
        "meta_title": "Personalized Trucker Prayer Comforter Set",
        "meta_description": "Customize a Trucker's Prayer comforter with required name, product type, size, pillowcase and sheet-cover choices.",
        "primary": "personalized Trucker Prayer comforter",
        "secondary": "custom trucker prayer bedding, Christian truck comforter with name, semi truck comforter",
        "cluster": "personalized Trucker Prayer comforter",
        "intent_role": "Verified required Enter Name up to 35 characters and product confirmation option group.",
        "detail": "Christian semi truck Trucker's Prayer comforter artwork with sample name only as an example",
    },
    49: {
        "title": "Personalized Bible Emergency Numbers Blanket",
        "meta_title": "Personalized Bible Emergency Numbers Blanket",
        "meta_description": "Customize a Bible emergency numbers blanket with required name, character choices, flower options and blanket size choices.",
        "primary": "personalized Bible emergency numbers blanket",
        "secondary": "custom Christian girl blanket, Bible emergency numbers blanket with name, personalized Christian blanket for girls",
        "cluster": "personalized Bible emergency numbers blanket",
        "intent_role": "Verified required Custom Name up to 1000 characters plus skin, eye, pants, shirt, hair and flower choices.",
        "detail": "floral Bible emergency numbers blanket artwork with cartoon girl, sample Sophia name and character choices",
    },
    50: {
        "title": "Personalized Christian Affirmation Blanket for Girls",
        "meta_title": "Personalized Christian Affirmation Blanket for Girls",
        "meta_description": "Customize a Christian affirmation blanket with required name, girl character choices, flower options and blanket size choices.",
        "primary": "personalized Christian affirmation blanket for girls",
        "secondary": "Christian girl blanket with name, Dear Sophia blanket, personalized inspirational blanket",
        "cluster": "personalized Christian affirmation blanket for girls",
        "intent_role": "Distinct affirmation/girl design separated from product 49's Bible emergency numbers intent; verified required Custom Name up to 30 characters.",
        "detail": "Dear Sophia Christian affirmation blanket artwork with cartoon girl and flower graphic",
    },
}


def patch_base_globals():
    base.ROOT = ROOT
    base.SHOP = SHOP
    base.RUN_ID = RUN_ID
    base.BATCH_ID = BATCH_ID
    base.QA_RUN_ID = QA_RUN_ID
    base.SOURCE = SOURCE
    base.QA_DATASET = QA_DATASET
    base.CUSTOMIZER_AUDIT = CUSTOMIZER_AUDIT
    base.ADMIN_EXPORT = ADMIN_EXPORT
    base.RESULT_DIR = RESULT_DIR
    base.RUN_DIR = RUN_DIR
    base.OUTPUT = OUTPUT
    base.REPORT = REPORT
    base.SUMMARY = SUMMARY
    base.PRODUCT_UPDATES = PRODUCT_UPDATES


def load_field_audit():
    if not CUSTOMIZER_FIELDS_AUDIT.exists():
        return {}
    return {int(item["inventory_position"]): item for item in json.loads(CUSTOMIZER_FIELDS_AUDIT.read_text(encoding="utf-8"))}


def field_sentence(pos, field_audit):
    item = field_audit.get(pos, {})
    fields = item.get("fields", [])
    text_inputs = [entry["fields"] for entry in fields if ".textInputs[" in entry.get("path", "")]
    option_groups = [entry["fields"] for entry in fields if re.search(r"\.optionGroups\[\d+\]$", entry.get("path", ""))]
    parts = []
    for text_input in text_inputs:
        required = "required" if text_input.get("required") else "optional"
        limit = text_input.get("maxLength")
        label = text_input.get("label", "text field")
        if limit:
            parts.append(f"Live Customizer shows a {required} {label} field up to {limit} characters.")
        else:
            parts.append(f"Live Customizer shows a {required} {label} field.")
    if option_groups:
        labels = ", ".join(group.get("label", "") for group in option_groups[:6] if group.get("label"))
        if labels:
            parts.append(f"Additional verified option groups include: {labels}.")
    if pos == 43:
        parts.append("Only name text personalization is claimed for this revision.")
    if pos in {41, 42, 48}:
        parts.append("Names and numbers visible in mockups are treated as samples.")
    return " ".join(parts) or "No shopper text-entry field is used in the SEO claim for this revision."


def description(pos, admin_row, customizer):
    item = PRODUCT_UPDATES[pos]
    options = ", ".join([v for v in [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]] if v])
    product_type = admin_row["Type"].lower() if admin_row["Type"] else "product"
    current_title_line = (
        "<li>Current Shopify export title is stored in the admin baseline for recheck.</li>"
        if pos == 43
        else f"<li>Current Shopify export title: {admin_row['Title']}.</li>"
    )
    return (
        f"<p>{item['detail'].capitalize()} gives this Jeminise {product_type} a specific customer-facing design focus.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork focus: {item['detail']}.</li>"
        f"{current_title_line}"
        "<li>Gallery images include the main mockup plus detail, lifestyle, care, size, feature or material panels where shown.</li></ul>"
        "<h3>Options and Customization</h3>"
        f"<ul><li>Available option groups from Shopify export: {options}.</li>"
        f"<li>{customizer}</li>"
        "<li>Use the live selectors to confirm size, product type and any required custom fields before checkout.</li></ul>"
    )


def alt_fix(pos, original):
    value = base.text(original)
    if pos != 43:
        return value
    replacements = [
        ("Photo Blanket", "Christian Blanket"),
        ("photo blanket", "Christian blanket"),
        ("Photo blanket", "Christian blanket"),
        ("personalized God Says I Am photo blanket", "personalized God Says I Am Christian blanket"),
        ("Personalized God Says I Am photo blanket", "Personalized God Says I Am Christian blanket"),
        ("with a sample photo", "with sample portrait-style artwork"),
        ("the photo blanket", "the Christian affirmation blanket"),
        ("God Says I Am photo blanket", "God Says I Am Christian blanket"),
    ]
    for old, new in replacements:
        value = value.replace(old, new)
    value = re.sub(r"\bphoto\b", "portrait-style artwork", value, flags=re.I)
    value = re.sub(r"\bupload\b", "image submission", value, flags=re.I)
    return value


def main():
    patch_base_globals()
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)

    qa = base.load_qa()
    field_audit = load_field_audit()
    admin = base.load_admin(set(qa["handle_by_key"].values()))
    admin_hash = base.sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()
    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = base.headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = base.text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ctext = field_sentence(pos, field_audit)

        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or ws.cell(row_num, idx["product_type"]).value
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, ctext)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"Buyer intent targets {item['cluster']} for US English product search; no search-volume claim is made."
        )
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R3 after Sang re-QA r2: used products_export_1.csv (sha256:{admin_hash}) as admin baseline; "
            "removed unsupported product 43 visual-personalization claim, restored verified customizer rules, "
            "rewrote generic descriptions, and refreshed image observations/alt. Still NEEDS_REVIEW; no APPROVED/import."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = base.headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    for row_num in range(2, ws.max_row + 1):
        handle = base.text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        media_id = base.text(ws.cell(row_num, idx["media_id"]).value)
        key = (pk, media_id)
        image_url = base.text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(base.normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation = qa["observation_by_key"].get(key)
        current_observation = observation or ws.cell(row_num, idx["observed_visual_details"]).value
        if current_observation:
            ws.cell(row_num, idx["observed_visual_details"]).value = alt_fix(pos, current_observation)
        fallback = f"{PRODUCT_UPDATES[pos]['title']} product image"
        proposed_alt = qa["image_fix_by_key"].get(key) or base.compact_observation_to_alt(observation, fallback)
        ws.cell(row_num, idx["alt_proposed"]).value = alt_fix(pos, proposed_alt)[:125]
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R3: Sang QA r2 image observation/alt applied; current admin alt and image URL matched from products_export_1.csv where possible."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = base.headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = base.text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = base.text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "Product"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = "Sang re-QA r2 SERP evidence + products_export_1.csv admin baseline"
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = "R3 separates sibling product intent by visible motif and verified customizer requirements."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r3"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = base.headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = base.text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as a specific bedding product."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm the artwork, product type, size choices "
            "and verified customizer rules."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = "Confirm size, product type, verified text fields, sample values, care panels and image accuracy."
        ws.cell(row_num, idx["emotional_social_motivation"]).value = "Choose a faith, sports, wildlife or Tree of Life bedding gift with clear artwork."
        ws.cell(row_num, idx["purchase_concerns"]).value = "Unsupported photo upload, exact customizer rules, included components, fabric/care claims and image accuracy."
        ws.cell(row_num, idx["source_refs"]).value = f"products_export_1.csv sha256:{admin_hash}; Sang QA r2 {QA_RUN_ID}; customizer_fields_audit.json"
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No search-volume source supplied; Shopify admin CSV baseline is now available for stored title/meta/body/image alt."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = base.headers(ws)
    for row_num in range(2, ws.max_row + 1):
        evidence_id = base.text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = qa["product_keys"][pos - 41]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = f"Storefront/product.js from Sang QA r2; products_export_1.csv sha256:{admin_hash}; customizer_fields_audit.json"
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={admin_row['Option1 Name']}, "
            f"{admin_row['Option2 Name']}, {admin_row['Option3 Name']}; status={admin_row['Status']}; "
            f"image_count={len(admin_row['images'])}; design={item['detail']}; customizer={field_sentence(pos, field_audit)}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "R3 applied Sang QA r2 fixes; needs independent QA recheck."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R3 uses QA-observed images, detailed Customizer field audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = base.headers(ws)
    for metric, value, definition in [
        ("qa_batch_005_r3_revision", "r3", "Separate 10-product revision after Sang re-QA r2; source r2 workbook was not modified."),
        ("qa_batch_005_r3_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_005_r3_scope", "inventory positions 41-50", "No products outside qa_batch_005 were revised."),
        ("qa_batch_005_r3_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_005_r3_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_005_r3_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa_r2": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_005 only; inventory positions 41-50",
        "revision": "r3",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Sang QA lại qa_batch_005_r3 before any approval or import.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_005_r3",
                "",
                "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r2: `{SOURCE.relative_to(ROOT)}`",
                f"- QA r2 của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r3: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Sửa trọng yếu: product 43 bỏ claim cá nhân hóa bằng ảnh; products 41-50 ghi lại đúng customizer fields đã xác minh.",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Bước tiếp theo: Sang QA lại `qa_batch_005_r3`.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
