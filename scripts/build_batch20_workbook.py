import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_020"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_019.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_020.xlsx"


BASEBALL_REFS = "https://www.amazon.com/Custom-Baseball-Bedding-Boys-Adults/dp/B0FHWPH2F3; https://www.walmart.com/c/kp/baseball-bedding-set; https://www.youcustomizeit.com/p/Baseball-Jersey-Comforters-Personalized/142431; https://ohaprints.com/products/personalized-baseball-duvet-cover-set-baseball-player-unique-gift-idea-blue-duvet-cover-pillowcases-custom-name-number-bedding-set; https://www.etsy.com/listing/4339325702/personalized-name-baseball-duvet-cover; https://www.wayfair.com/keyword.php?keyword=baseball+comforter"


PROFILES = {
    191: ("red black batter baseball bedding with name", "batter baseball bedding", "Batter Baseball Bedding Set", "Red Black Batter Baseball Bedding with Custom Name", "red and black baseball bedding with a batter silhouette, large custom name and jersey number artwork", "personalized baseball bedding, custom name baseball comforter, baseball player bedding", "EVERGREEN", BASEBALL_REFS),
    192: ("brown baseball glove close-up bedding with script name", "baseball glove bedding", "Baseball Glove Bedding Set", "Brown Baseball Glove Bedding with Script Name", "brown baseball glove close-up bedding with oversized ball, leather glove details and script custom name", "personalized baseball comforter, baseball ball bedding, custom name sports bedding", "EVERGREEN", BASEBALL_REFS),
    193: ("flaming baseball batter bedding with vertical name", "flaming baseball bedding", "Flaming Baseball Bedding Set", "Flaming Baseball Batter Bedding with Custom Name", "black baseball bedding with flame effects, batter artwork, vertical custom name and jersey number", "baseball player bedding, custom baseball comforter, personalized sports bedding", "EVERGREEN", BASEBALL_REFS),
    194: ("orange fire baseball player bedding with name", "fire baseball bedding", "Fire Baseball Bedding Set", "Orange Fire Baseball Player Bedding with Name", "orange fire baseball bedding with batter stance, vertical custom name and bold jersey number", "personalized baseball bedding, baseball player comforter, custom name bedding", "EVERGREEN", BASEBALL_REFS),
    195: ("American flag baseball glove bedding with vertical name", "baseball glove flag bedding", "Baseball Glove Flag Bedding", "American Flag Baseball Glove Bedding with Name", "American flag baseball bedding with glove and ball artwork, vertical custom name and jersey number on shams", "custom baseball bedding, American flag baseball comforter, personalized sports bedding", "EVERGREEN", BASEBALL_REFS),
    196: ("black baseball glove bedding with script name and number", "black baseball glove bedding", "Black Baseball Glove Bedding", "Black Baseball Glove Bedding with Script Name", "black baseball bedding with grayscale glove and ball artwork, script custom name and jersey number", "custom baseball comforter, personalized baseball bedding, baseball glove bedding", "EVERGREEN", BASEBALL_REFS),
    197: ("lightning baseball bedding with custom name and number", "lightning baseball bedding", "Lightning Baseball Bedding Set", "Lightning Baseball Bedding with Name and Number", "stormy blue baseball bedding with lightning effects, large baseball, custom name and jersey number", "personalized baseball bedding, baseball number comforter, custom sports bedding", "EVERGREEN", BASEBALL_REFS),
    198: ("night field baseball glove bedding with custom name", "night baseball bedding", "Night Baseball Bedding Set", "Night Field Baseball Glove Bedding with Custom Name", "night-sky baseball bedding with helmet, glove, bat, ball, grass field detail and script custom name", "baseball glove comforter, personalized baseball bedding, custom sports room bedding", "EVERGREEN", BASEBALL_REFS),
    199: ("smoke baseball batter bedding with vertical name", "smoke baseball bedding", "Smoke Baseball Bedding Set", "Smoke Baseball Batter Bedding with Custom Name", "black and smoke baseball bedding with batter silhouette, vertical custom name and jersey number on pillow shams", "personalized baseball comforter, baseball player bedding, custom name bedding", "EVERGREEN", BASEBALL_REFS),
    200: ("baseball stitching bedding with script name and number", "baseball stitching bedding", "Baseball Stitching Bedding Set", "Baseball Stitching Bedding with Name and Number", "cream baseball stitching bedding with red lace detail, script custom name and jersey number", "personalized baseball bedding, baseball stitch comforter, custom name sports bedding", "EVERGREEN", BASEBALL_REFS),
}


def qa_issues(pos, html_error):
    issues = [
        "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment.",
        "Personalized baseball products from batches 018-020 are very close; review custom-name placement, jersey number accuracy and duplicate keyword cannibalization before import.",
        "Feature images mention duvet cover/comforter options, microfiber, easy care and size dimensions; verify actual product type, variants and set configuration before deployment.",
    ]
    if pos in (191, 193, 194, 199):
        issues.append("Batter/player silhouette designs overlap; keep color, fire/smoke and name orientation distinct.")
    if pos in (192, 195, 196, 198):
        issues.append("Glove-and-ball motifs overlap; review title uniqueness and canonical strategy.")
    if pos == 197:
        issues.append("Lightning/galaxy-style baseball designs may overlap with product 190; keep storm/lightning wording specific.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for baseball rooms, sports gifts, kids bedrooms or personalized fan bedding.</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broad personalized baseball bedding terms for collections or the strongest representative product.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, a sports room accent or a personalized gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": qa_issues(pos, html_error),
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for personalized sports bedding, the buyer wants a {kind} matching a baseball motif, custom name, number or player style.",
            "functional_motivation": "Find the right baseball design, size, personalization fields, care details and room fit.",
            "emotional_social_motivation": "Create a baseball-themed bedroom or give a custom sports gift to a player or fan.",
            "purchase_concerns": "Personalization spelling, jersey number, product identity, material/care claims, size fit and whether artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed in a main product mockup",
                2: f"Secondary bedroom mockup for {h1}",
                3: f"Close-up or feature image for {h1}",
                4: f"Fabric or product feature image for {h1}",
                5: f"Easy care feature image for {h1}",
                6: f"Bedding type comparison image for {h1}",
                7: f"Size dimension image for {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_020_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_020_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_020"},
        {"metric": "cumulative_products", "value": "200", "definition": "Total products included through batch_020"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[200:210]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_021",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
