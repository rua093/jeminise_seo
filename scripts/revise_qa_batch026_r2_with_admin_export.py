import csv
import hashlib
import html
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_026_r2"
SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "batches" / "SEO_Product_Optimization_through_batch_026.xlsx"
SUMMARY_PATH = ROOT / "seo_runs" / SHOP / RUN_ID / "batch_026_evidence_summary.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_026_r2.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_026_r2.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    251: ("Personalized Christian Butterfly Blanket", "personalized Christian butterfly blanket", "Christian butterfly blanket, faith affirmation blanket, custom scripture blanket", "white Christian throw blanket with butterfly wings, affirmation words, scripture references and sample name Maria", "personalized Christian butterfly affirmation blanket"),
    252: ("Personalized Reading Tree Blanket", "personalized reading tree blanket", "personalized reading blanket, book lover blanket, custom reader gift", "reading-themed throw blanket with cartoon reader, tree of books, quote text and sample name Jessica", "personalized reading tree blanket"),
    253: ("Personalized Celtic Tree of Life Quilt Set", "personalized Celtic Tree of Life quilt set", "Celtic quilt set, Tree of Life bedding, custom Yggdrasil quilt", "green and gold quilt set with Celtic panels and central Tree of Life medallion", "personalized Celtic Tree of Life quilt set"),
    254: ("Personalized Celtic Yggdrasil Quilt Set", "personalized Celtic Yggdrasil quilt set", "Yggdrasil quilt set, Celtic Tree of Life bedding, custom Celtic quilt", "dark Celtic quilt set with interlaced red and green knotwork and a glowing tree medallion", "personalized Celtic Yggdrasil quilt set"),
    255: ("Personalized Green Yggdrasil Quilt Set", "personalized green Yggdrasil quilt set", "green Yggdrasil quilt set, Celtic tree quilt, custom Tree of Life bedding", "green quilt set with circular gold Yggdrasil artwork and interlaced knot pattern", "personalized green Yggdrasil quilt set"),
    256: ("Personalized Cosmic Tree of Life Quilt Set", "personalized cosmic Tree of Life quilt set", "cosmic Tree of Life quilt, Celtic quilt set, custom bedroom quilt", "dark quilt set with cosmic Tree of Life panels, moon imagery, chain border and purple accents", "personalized cosmic Tree of Life quilt set"),
    257: ("Personalized Celtic Knot Tree Quilt Set", "personalized Celtic knot tree quilt set", "Celtic knot quilt set, Tree of Life quilt, custom Celtic bedding", "black and green quilt set with Celtic knot frame, red accents and central tree symbol", "personalized Celtic knot Tree of Life quilt set"),
    258: ("Personalized Christian Affirmations Blanket", "personalized Christian affirmations blanket", "Christian affirmation blanket, personalized scripture blanket, faith gift blanket", "white affirmation blanket with butterflies, scripture reference bubbles, photo area and sample name Laura", "personalized Christian affirmations blanket"),
    259: ("Personalized Christian Woman Affirmation Blanket", "personalized Christian woman affirmation blanket", "Christian woman blanket, affirmation throw blanket, custom scripture blanket", "white affirmation blanket with seated woman illustration, butterflies, scripture reference bubbles and sample name Jennifer", "personalized Christian woman affirmation blanket"),
    260: ("Personalized Lavender Scripture Blanket", "personalized lavender scripture blanket", "Christian lavender blanket, personalized scripture blanket, butterfly affirmation blanket", "white blanket with lavender flowers, butterflies, affirmation words, scripture references and sample name Rebecca", "personalized lavender scripture blanket"),
}


IMAGE_DETAILS = {
    251: {1: ("Person holds Christian butterfly blanket with words God says you are, colorful wings and sample name Maria.", "Personalized Christian butterfly blanket held view"), 2: ("Feature panel showing white fleece edge and machine washable, soft and warm callouts.", "Christian butterfly blanket feature panel"), 3: ("Close-up of butterfly wing design with scripture references and sample name.", "Christian butterfly blanket close-up"), 4: ("Size and use panel for Christian butterfly blanket with sofa, office, bed and travel icons.", "Christian butterfly blanket size and use panel"), 5: ("Detail collage showing butterfly artwork, printed name and white fleece texture.", "Christian butterfly blanket detail collage"), 6: ("Christian butterfly blanket draped over an armchair in a living room.", "Christian butterfly blanket chair mockup"), 7: ("Lifestyle reading scene with Christian butterfly blanket, open book and small pet.", "Christian butterfly blanket reading lifestyle")},
    252: {1: ("Person holds reading tree blanket with cartoon reader, books, quote text and sample name Jessica.", "Personalized reading tree blanket held view"), 2: ("Flat reading blanket design with tree of books and cartoon reader.", "Reading tree blanket flat design"), 3: ("Person holds reading tree blanket in a living room setting.", "Reading tree blanket room held view"), 4: ("Size chart showing four blanket sizes for the reading tree design.", "Reading tree blanket size chart"), 5: ("Reading tree blanket draped across sofa under a window.", "Reading tree blanket on sofa"), 6: ("Second sofa mockup of reading tree blanket beneath round mirror.", "Reading tree blanket sofa mockup"), 7: ("Feature panel over reading tree artwork with fluffy and no-shedding callouts.", "Reading tree blanket feature panel")},
    253: {1: ("Bed mockup of green Celtic Tree of Life quilt set with matching pillows.", "Celtic Tree of Life quilt bed view"), 2: ("Second room mockup of Celtic Tree of Life quilt with printed craft callout.", "Celtic Tree of Life quilt room mockup"), 3: ("Close bed view with optional pillow shams and Celtic panel artwork.", "Celtic Tree of Life quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and matching Tree of Life pillow.", "Celtic Tree of Life quilt fabric panel"), 5: ("Premium quilt sets sizing chart for the Celtic Tree of Life design.", "Celtic Tree of Life quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Celtic Tree of Life quilt features panel"), 7: ("Overhead bedroom mockup of Celtic Tree of Life quilt and matching pillows.", "Celtic Tree of Life quilt overhead view")},
    254: {1: ("Bed mockup of dark Celtic Yggdrasil quilt with red and green knotwork.", "Celtic Yggdrasil quilt bed view"), 2: ("Second room mockup of Celtic Yggdrasil quilt with printed craft callout.", "Celtic Yggdrasil quilt room mockup"), 3: ("Close bed view with optional pillow shams and interlaced Celtic artwork.", "Celtic Yggdrasil quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and colorful tree artwork.", "Celtic Yggdrasil quilt fabric panel"), 5: ("Premium quilt sets sizing chart for the Celtic Yggdrasil design.", "Celtic Yggdrasil quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Celtic Yggdrasil quilt features panel"), 7: ("Overhead bedroom mockup of Celtic Yggdrasil quilt and matching pillows.", "Celtic Yggdrasil quilt overhead view")},
    255: {1: ("Bed mockup of green Yggdrasil quilt with circular gold tree artwork.", "Green Yggdrasil quilt bed view"), 2: ("Second room mockup of green Yggdrasil quilt with craft callout.", "Green Yggdrasil quilt room mockup"), 3: ("Close bed view with optional pillow shams and circular tree artwork.", "Green Yggdrasil quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and green knotwork sample.", "Green Yggdrasil quilt fabric panel"), 5: ("Premium quilt sets sizing chart for the green Yggdrasil design.", "Green Yggdrasil quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Green Yggdrasil quilt features panel"), 7: ("Overhead bedroom mockup of green Yggdrasil quilt and matching pillows.", "Green Yggdrasil quilt overhead view")},
    256: {1: ("Bed mockup of cosmic Tree of Life quilt with moon panels and chain border.", "Cosmic Tree of Life quilt bed view"), 2: ("Second room mockup of cosmic Tree of Life quilt with printed craft callout.", "Cosmic Tree of Life quilt room mockup"), 3: ("Close bed view with optional pillow shams and cosmic panel artwork.", "Cosmic Tree of Life quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and chain-border artwork sample.", "Cosmic Tree of Life quilt fabric panel"), 5: ("Premium quilt sets sizing chart for the cosmic Tree of Life design.", "Cosmic Tree of Life quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Cosmic Tree of Life quilt features panel"), 7: ("Overhead bedroom mockup of cosmic Tree of Life quilt and matching pillows.", "Cosmic Tree of Life quilt overhead view")},
    257: {1: ("Bed mockup of black and green Celtic knot Tree quilt with red accents.", "Celtic knot Tree quilt bed view"), 2: ("Second room mockup of Celtic knot Tree quilt with printed craft callout.", "Celtic knot Tree quilt room mockup"), 3: ("Close bed view with optional pillow shams and Celtic knot artwork.", "Celtic knot Tree quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and green tree artwork sample.", "Celtic knot Tree quilt fabric panel"), 5: ("Premium quilt sets sizing chart for the Celtic knot Tree design.", "Celtic knot Tree quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Celtic knot Tree quilt features panel"), 7: ("Overhead bedroom mockup of Celtic knot Tree quilt and matching pillows.", "Celtic knot Tree quilt overhead view")},
    258: {1: ("Person holds Christian affirmations blanket with butterflies, scripture bubbles, photo area and sample name Laura.", "Personalized Christian affirmations blanket held view"), 2: ("Flat Christian affirmations blanket with photo area, butterflies and scripture reference bubbles.", "Christian affirmations blanket flat view"), 3: ("Christian affirmations blanket draped across sofa under a window.", "Christian affirmations blanket on sofa"), 4: ("Second sofa mockup of Christian affirmations blanket beneath round mirror.", "Christian affirmations blanket sofa mockup"), 5: ("Feature panel over affirmations artwork with fluffy and no-shedding callouts.", "Christian affirmations blanket feature panel"), 6: ("Photo guidelines panel showing acceptable and unsuitable photo examples.", "Christian affirmations blanket photo guidelines panel")},
    259: {1: ("Person holds Christian affirmation blanket with seated woman illustration, butterflies and sample name Jennifer.", "Christian woman affirmation blanket held view"), 2: ("Flat Christian woman affirmation blanket with scripture reference bubbles and butterflies.", "Christian woman affirmation blanket flat view"), 3: ("Christian woman affirmation blanket draped across sofa under a window.", "Christian woman affirmation blanket on sofa"), 4: ("Second sofa mockup of Christian woman affirmation blanket beneath round mirror.", "Christian woman affirmation blanket sofa mockup"), 5: ("Feature panel over affirmation artwork with fluffy and no-shedding callouts.", "Christian woman affirmation blanket feature panel")},
    260: {1: ("Person holds lavender scripture blanket with purple flowers, butterflies, affirmation words and sample name Rebecca.", "Personalized lavender scripture blanket held view"), 2: ("Held lavender scripture blanket with flower color examples shown on side.", "Lavender scripture blanket color option view"), 3: ("Flat lavender scripture blanket with affirmation words and scripture references.", "Lavender scripture blanket flat view"), 4: ("Lavender scripture blanket draped across sofa under a window.", "Lavender scripture blanket on sofa"), 5: ("Second sofa mockup of lavender scripture blanket beneath round mirror.", "Lavender scripture blanket sofa mockup"), 6: ("Feature panel over lavender scripture artwork with fluffy and no-shedding callouts.", "Lavender scripture blanket feature panel")},
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def parse_customizer(summary):
    raw = Path(summary["html"]["html_file"]).read_text(encoding="utf-8", errors="ignore")
    page = html.unescape(raw)
    fields = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"minLength":\s*(\d+),\s*"maxLength":\s*(\d+)',
        page,
    )
    if not fields:
        fields = re.findall(
            r'"label"\s*:\s*"([^"]+)"\s*,\s*"required"\s*:\s*(true|false)\s*,\s*"minLength"\s*:\s*(\d+)\s*,\s*"maxLength"\s*:\s*(\d+)',
            page,
        )
    options = re.findall(r'"label"\s*:\s*"([^"]+)"\s*,\s*"required"\s*:\s*(true|false)\s*,\s*"defaultOptionId"', page)
    return fields, options


def customizer_sentence(fields, options):
    parts = []
    for label, required, min_length, max_length in fields:
        requirement = "required" if required == "true" else "optional"
        parts.append(f"Personalization uses a {requirement} {label} field, {min_length}-{max_length} characters.")
    visible_options = [label for label, _ in options if label and "confirmation" not in label.lower()]
    if visible_options:
        parts.append(f"Visible required choice fields include: {', '.join(visible_options)}.")
    if fields:
        parts.append("Visible custom text or example selections in mockups are sample artwork unless the matching input is entered.")
    else:
        parts.append("No shopper text-entry field is described for this product.")
    return " ".join(parts)


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {"src": row["Image Src"], "alt": row.get("Image Alt Text", "")}
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    missing = sorted(set(handles) - set(admin))
    if missing:
        raise RuntimeError(f"Missing handles in admin export: {missing}")
    return admin


def option_text(admin_row):
    values = [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]]
    values = [v for v in values if v]
    return ", ".join(values) if values else "available size options"


def description(pos, admin_row, ctext):
    title, _, _, detail, _ = PRODUCT_UPDATES[pos]
    return (
        f"<p>{title} features {detail}. The copy stays specific to the visible artwork, product form and selectable options for this exact item.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork: {detail}.</li><li>Gallery images include the main mockup plus size, feature, care, fabric or lifestyle panels where shown.</li></ul>"
        "<h3>Options and Fit</h3>"
        f"<ul><li>Available option groups: {option_text(admin_row)}.</li><li>{ctext}</li><li>Select the product type and size shown on the product page before checkout.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    summaries = {int(item["inventory_row"]["inventory_position"]): item for item in json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))}
    positions = set(PRODUCT_UPDATES)
    handles = {summaries[pos]["inventory_row"]["Handle"] for pos in positions}
    admin = load_admin(handles)
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()
    wb = load_workbook(OUTPUT)

    pos_by_key = {summaries[pos]["inventory_row"]["product_key"]: pos for pos in positions}
    handle_by_pos = {pos: summaries[pos]["inventory_row"]["Handle"] for pos in positions}
    fields_by_pos = {}
    option_fields_by_pos = {}
    for pos in positions:
        fields_by_pos[pos], option_fields_by_pos[pos] = parse_customizer(summaries[pos])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        handle = handle_by_pos[pos]
        title, primary, secondary, detail, cluster = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        meta = f"Shop {title.lower()} with {detail}, selectable sizes and verified custom text options."[:155]
        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or summaries[pos]["inventory_row"]["product_type"]
        ws.cell(row_num, idx["title_proposed"]).value = title
        ws.cell(row_num, idx["meta_title_seo"]).value = title
        ws.cell(row_num, idx["meta_title_length"]).value = len(title)
        ws.cell(row_num, idx["meta_description_seo"]).value = meta
        ws.cell(row_num, idx["meta_description_length"]).value = len(meta)
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, customizer_sentence(fields_by_pos[pos], option_fields_by_pos[pos]))
        ws.cell(row_num, idx["primary_keyword"]).value = primary
        ws.cell(row_num, idx["secondary_keywords"]).value = secondary
        ws.cell(row_num, idx["meta_keyword"]).value = primary
        ws.cell(row_num, idx["keyword_strategy"]).value = f"Target the specific product motif: {cluster}; keep broad Christian gift, reading gift or Celtic bedding terms for collection pages."
        ws.cell(row_num, idx["buyer_search_summary"]).value = f"US English buyer intent targets {cluster}; no search-volume, trend or ranking claim is made."
        ws.cell(row_num, idx["keyword_evidence_level"]).value = "SERP_ONLY"
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R2 deep recheck without prior QA: used products_export_1.csv sha256:{admin_hash}; "
            "separated Christian butterfly blankets, reading blanket and Celtic quilt-set intents; parsed custom fields and rewrote image observations/alts from contact sheets. Still NEEDS_REVIEW."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    seen_alts = set()
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pos = next((p for p, h in handle_by_pos.items() if h == handle), None)
        if pos is None:
            continue
        image_number = int(ws.cell(row_num, idx["image_number"]).value)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation, alt = IMAGE_DETAILS[pos][image_number]
        if alt in seen_alts:
            alt = f"{alt} for {PRODUCT_UPDATES[pos][4]}"
        seen_alts.add(alt)
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = alt[:125]
        ws.cell(row_num, idx["alt_action"]).value = "SET"
        ws.cell(row_num, idx["viewed_status"]).value = "VIEWED"
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = "R2: observation and alt rechecked from contact sheet; admin image URL and current alt matched from Shopify CSV where available."
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        title, primary, secondary, _, cluster = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in secondary.split(",")]
        keyword = primary if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = cluster
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "PRODUCT"
        ws.cell(row_num, idx["decision_reason"]).value = f"{title} is separated by verified product form, artwork motif and custom fields."
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = "products_export_1.csv admin baseline + storefront customizer parse + contact-sheet inspection"
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = "YES; nearby products are differentiated by motif, product form and custom field type."
        ws.cell(row_num, idx["mapping_reason"]).value = "R2 separates Christian, reading and Celtic quilt-set pages by visible motifs and option fields."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r2"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        title, primary, _, _, cluster = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {cluster} as a specific product page."
        ws.cell(row_num, idx["jtbd_statement"]).value = f"When shopping for {cluster}, the buyer wants exact artwork, product type, size choices and available custom fields confirmed before purchase."
        ws.cell(row_num, idx["functional_motivation"]).value = "Confirm product form, visible motif, size options, custom field rules, feature panels and image accuracy before purchase."
        ws.cell(row_num, idx["emotional_social_motivation"]).value = "Choose a personalized faith, reading or Celtic-themed bedding item that matches the recipient or room theme."
        ws.cell(row_num, idx["purchase_concerns"]).value = "Confirm custom text rules; avoid unsupported material, religious promise, copyrighted identity or delivery claims."
        ws.cell(row_num, idx["source_refs"]).value = f"products_export_1.csv sha256:{admin_hash}; storefront HTML customizer parse; contact sheets 251-260"
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R2_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No independent QA, Search Console, paid keyword volume or internal site-search export supplied; proof remains SERP-only."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {primary} with title: {title}."

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in positions:
            continue
        handle = handle_by_pos[pos]
        admin_row = admin[handle]
        _, _, _, detail, _ = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = f"Original storefront/product evidence; products_export_1.csv sha256:{admin_hash}; storefront customizer parse; inspected contact sheets"
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={option_text(admin_row)}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={detail}; customizer={customizer_sentence(fields_by_pos[pos], option_fields_by_pos[pos])}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "No independent QA for batch 026 yet; r2 prepared for QA recheck."
        ws.cell(row_num, idx["confidence_and_reason"]).value = "R2 uses contact sheets, storefront customizer parse and Shopify admin CSV baseline; no approval or import created."

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_026_r2_revision", "r2", "Separate 10-product revision without prior independent QA; source workbook was not modified."),
        ("qa_batch_026_r2_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_026_r2_scope", "inventory positions 251-260", "No products outside qa_batch_026 were revised."),
        ("qa_batch_026_r2_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_026_r2_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_026_r2_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else "" for col in range(1, ws.max_column + 1)])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)
    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_evidence": str(SUMMARY_PATH.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_026 only; inventory positions 251-260",
        "revision": "r2",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Ready for independent QA of qa_batch_026_r2.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join([
            "# Revision qa_batch_026_r2",
            "",
            "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
            "",
            f"- Nguồn r1: `{SOURCE.relative_to(ROOT)}`",
            f"- Evidence summary: `{SUMMARY_PATH.relative_to(ROOT)}`",
            f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
            f"- Admin export SHA-256: `{admin_hash}`",
            f"- Workbook r2: `{OUTPUT.relative_to(ROOT)}`",
            f"- Sản phẩm sửa: {revised_products}",
            f"- Ảnh sửa/refresh: {revised_images}",
            f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
            "- Sửa trọng yếu: rà lại khi chưa có QA, tách nhóm Christian blanket, reading blanket và Celtic quilt set, parse customizer từ HTML storefront, viết lại title/meta/description/alt theo từng contact sheet.",
            "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
            "- Bước tiếp theo: gửi `qa_batch_026_r2` cho QA độc lập.",
            "",
        ]),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
