import csv
import hashlib
import html
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_025_r2"
SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "batches" / "SEO_Product_Optimization_through_batch_025.xlsx"
SUMMARY_PATH = ROOT / "seo_runs" / SHOP / RUN_ID / "batch_025_evidence_summary.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_025_r2.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_025_r2.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    241: ("Personalized Bigfoot Peace Sign Mountain Quilt Set", "personalized Bigfoot peace sign mountain quilt set", "Bigfoot quilt set, mountain landscape quilt, custom cabin bedding", "mountain landscape quilt set with a large Bigfoot figure making a peace sign", "personalized Bigfoot peace sign mountain quilt set"),
    242: ("Personalized Bigfoot Peace Sign Sunset Quilt Set", "personalized Bigfoot peace sign sunset quilt set", "Bigfoot sunset quilt set, peace sign quilt, custom cabin bedding", "sunset forest quilt set with a close Bigfoot figure making a peace sign", "personalized Bigfoot peace sign sunset quilt set"),
    243: ("Personalized Bigfoot Forest Night Quilt Set", "personalized Bigfoot forest night quilt set", "Bigfoot forest quilt set, night forest bedding, custom cabin quilt", "night forest quilt set with a large walking Bigfoot silhouette in front of a moon", "personalized Bigfoot forest night quilt set"),
    244: ("Personalized Bigfoot Snowy Mountain Quilt Set", "personalized Bigfoot snowy mountain quilt set", "Bigfoot mountain quilt set, snowy forest bedding, custom cabin quilt", "snowy mountain forest quilt set with a small Bigfoot silhouette near pine trees", "personalized Bigfoot snowy mountain quilt set"),
    245: ("Personalized Bigfoot Sunset Forest Quilt Set", "personalized Bigfoot sunset forest quilt set", "Bigfoot sunset quilt set, forest silhouette bedding, custom nature quilt", "gold sunset forest quilt set with long tree shadows and a Bigfoot silhouette", "personalized Bigfoot sunset forest quilt set"),
    246: ("Personalized Bigfoot Red Sunglasses Quilt Set", "personalized Bigfoot red sunglasses quilt set", "Bigfoot quilt set, funny cabin bedding, custom woodland quilt", "woodland quilt set with a large Bigfoot figure wearing red sunglasses", "personalized Bigfoot red sunglasses quilt set"),
    247: ("Personalized Bookshelf Reading Girl Blanket", "personalized bookshelf reading girl blanket", "personalized book lover blanket, reading girl blanket, custom reader blanket", "book-themed blanket with a black-haired reading girl, bookshelf collage, quotes and sample name", "personalized bookshelf reading girl blanket"),
    248: ("Personalized Blonde Reading Girl Blanket", "personalized blonde reading girl blanket", "personalized book lover blanket, bookshelf blanket, custom reader gift", "bookstore-style blanket with blonde reading girl, bookshelves, quotes and sample name", "personalized blonde reading girl blanket"),
    249: ("Personalized Antique Books Reading Blanket", "personalized antique books reading blanket", "personalized book lover blanket, antique books blanket, custom reading blanket", "vintage book-page blanket with girl-in-bun artwork, stacked books, quote text and sample name", "personalized antique books reading blanket"),
    250: ("Personalized Wildflower Reading Girl Blanket", "personalized wildflower reading girl blanket", "personalized book lover blanket, floral reading blanket, custom reader blanket", "wildflower book blanket with brown-haired reading girl, stack of books and quote text", "personalized wildflower reading girl blanket"),
}


IMAGE_DETAILS = {
    241: {1: ("Bed mockup of mountain quilt with Bigfoot peace sign figure and matching pillows.", "Bigfoot peace sign mountain quilt bed view"), 2: ("Second room mockup of peace sign mountain quilt with printed craft callout.", "Bigfoot peace sign mountain quilt room mockup"), 3: ("Close bed view with optional pillow shams and mountain landscape artwork.", "Bigfoot peace sign mountain quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and matching pillow image.", "Bigfoot peace sign mountain quilt fabric panel"), 5: ("Premium quilt sets sizing chart for throw through king sizes.", "Bigfoot peace sign mountain quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Bigfoot peace sign mountain quilt features panel"), 7: ("Overhead bedroom mockup of Bigfoot peace sign mountain quilt and pillows.", "Bigfoot peace sign mountain quilt overhead view")},
    242: {1: ("Bed mockup of sunset quilt with close Bigfoot peace sign figure and forest border.", "Bigfoot peace sign sunset quilt bed view"), 2: ("Second room mockup of sunset peace sign quilt with printed craft callout.", "Bigfoot peace sign sunset quilt room mockup"), 3: ("Close bed view with optional pillow shams and sunset Bigfoot artwork.", "Bigfoot peace sign sunset quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and sunset artwork sample.", "Bigfoot peace sign sunset quilt fabric panel"), 5: ("Premium quilt sets sizing chart for the sunset design.", "Bigfoot peace sign sunset quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Bigfoot peace sign sunset quilt features panel"), 7: ("Overhead bedroom mockup of sunset Bigfoot quilt and matching pillows.", "Bigfoot peace sign sunset quilt overhead view")},
    243: {1: ("Bed mockup of night forest quilt with Bigfoot silhouette and moonlit mountains.", "Bigfoot forest night quilt bed view"), 2: ("Second room mockup of moonlit forest Bigfoot quilt with craft callout.", "Bigfoot forest night quilt room mockup"), 3: ("Close bed view with optional pillow shams and moonlit silhouette art.", "Bigfoot forest night quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and matching pillow image.", "Bigfoot forest night quilt fabric panel"), 5: ("Premium quilt sets sizing chart for the night forest design.", "Bigfoot forest night quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Bigfoot forest night quilt features panel"), 7: ("Overhead bedroom mockup of moonlit Bigfoot forest quilt and pillows.", "Bigfoot forest night quilt overhead view")},
    244: {1: ("Bed mockup of snowy mountain quilt with small Bigfoot silhouette near trees.", "Bigfoot snowy mountain quilt bed view"), 2: ("Second room mockup of snowy mountain Bigfoot quilt with craft callout.", "Bigfoot snowy mountain quilt room mockup"), 3: ("Close bed view with optional pillow shams and snowy mountain artwork.", "Bigfoot snowy mountain quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and snowy artwork sample.", "Bigfoot snowy mountain quilt fabric panel"), 5: ("Premium quilt sets sizing chart for the snowy mountain design.", "Bigfoot snowy mountain quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Bigfoot snowy mountain quilt features panel"), 7: ("Overhead bedroom mockup of snowy mountain Bigfoot quilt and pillows.", "Bigfoot snowy mountain quilt overhead view")},
    245: {1: ("Bed mockup of golden sunset forest quilt with Bigfoot silhouette and long tree shadows.", "Bigfoot sunset forest quilt bed view"), 2: ("Second room mockup of sunset forest Bigfoot quilt with craft callout.", "Bigfoot sunset forest quilt room mockup"), 3: ("Close bed view with optional pillow shams and sunset forest artwork.", "Bigfoot sunset forest quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and orange forest sample.", "Bigfoot sunset forest quilt fabric panel"), 5: ("Premium quilt sets sizing chart for the sunset forest design.", "Bigfoot sunset forest quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Bigfoot sunset forest quilt features panel"), 7: ("Overhead bedroom mockup of sunset forest Bigfoot quilt and pillows.", "Bigfoot sunset forest quilt overhead view")},
    246: {1: ("Bed mockup of woodland quilt with Bigfoot wearing red sunglasses.", "Bigfoot red sunglasses quilt bed view"), 2: ("Second room mockup of red sunglasses Bigfoot quilt with craft callout.", "Bigfoot red sunglasses quilt room mockup"), 3: ("Close bed view with optional pillow shams and woodland Bigfoot artwork.", "Bigfoot red sunglasses quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and Bigfoot artwork sample.", "Bigfoot red sunglasses quilt fabric panel"), 5: ("Premium quilt sets sizing chart for the red sunglasses design.", "Bigfoot red sunglasses quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Bigfoot red sunglasses quilt features panel"), 7: ("Overhead bedroom mockup of red sunglasses Bigfoot quilt and pillows.", "Bigfoot red sunglasses quilt overhead view")},
    247: {1: ("Person holds book lover blanket with black-haired reading girl, bookshelf collage, quotes and sample name Jessica.", "Personalized bookshelf reading girl blanket held view"), 2: ("Book lover blanket draped on sofa with bookshelf collage and quote blocks.", "Bookshelf reading girl blanket on sofa"), 3: ("Collage of print close-ups showing reading girl artwork and white fleece texture.", "Bookshelf reading blanket detail collage"), 4: ("Lifestyle reading scene with blanket, open book and small pet on bed.", "Bookshelf reading blanket lifestyle scene"), 5: ("Size and use panel for the reading blanket with sofa, office, bed and travel icons.", "Bookshelf reading blanket size and use panel"), 6: ("Close-up of reading girl illustration on blanket beside an open book.", "Bookshelf reading girl blanket close-up"), 7: ("Feature panel with machine washable, soft and warm, dense stitching and fleece callouts.", "Bookshelf reading blanket feature panel")},
    248: {1: ("Held blanket with blonde reading girl, bookshelves, quote blocks and sample name Jennifer.", "Personalized blonde reading girl blanket held view"), 2: ("Lifestyle image of blonde reader wrapped in the bookshelf blanket by a window.", "Blonde reading girl blanket reading lifestyle"), 3: ("Multiple purpose panel with bed, sofa, office, travel and pet examples.", "Blonde reading blanket multi-use panel"), 4: ("Person holds full blonde reading girl blanket in a living room.", "Blonde reading girl blanket room held view"), 5: ("Lifestyle reading scene with blanket, open book and small pet.", "Blonde reading blanket bed lifestyle scene"), 6: ("Blonde reading girl blanket draped across sofa under a window.", "Blonde reading girl blanket on sofa"), 7: ("Feature panel over bookshelf artwork with fluffy and no-shedding callouts.", "Blonde reading blanket feature panel"), 8: ("Adult and child reading under blanket with bookshelf quote artwork.", "Blonde reading blanket family lifestyle image"), 9: ("Blanket size chart showing four sizes with scale figure.", "Blonde reading blanket size chart")},
    249: {1: ("Person holds vintage book-page blanket with girl-in-bun artwork, stacked books and sample name Rebecca.", "Personalized antique books reading blanket held view"), 2: ("Flat antique books blanket design with quote text and stacked books.", "Antique books reading blanket flat design"), 3: ("Person holds antique books reading blanket in living room setting.", "Antique books reading blanket room held view"), 4: ("Size chart panel showing four blanket sizes for the antique books design.", "Antique books reading blanket size chart"), 5: ("Antique books reading blanket draped across sofa.", "Antique books reading blanket on sofa"), 6: ("Second sofa mockup of antique books reading blanket beneath round mirror.", "Antique books reading blanket sofa mockup"), 7: ("Feature panel over antique book-page artwork with fabric callouts.", "Antique books reading blanket feature panel")},
    250: {1: ("Held blanket with brown-haired reading girl, wildflowers, book stack and quote text.", "Personalized wildflower reading girl blanket held view"), 2: ("Wildflower reading girl blanket draped on sofa with white fleece edge.", "Wildflower reading girl blanket on sofa"), 3: ("Detail collage showing girl, books, flowers and white fleece texture.", "Wildflower reading blanket detail collage"), 4: ("Lifestyle reading scene with wildflower blanket, open book and small pet.", "Wildflower reading blanket lifestyle scene"), 5: ("Feature panel with machine washable, soft and warm, dense stitching and fleece callouts.", "Wildflower reading blanket feature panel"), 6: ("Close-up of wildflower reading blanket near an open book.", "Wildflower reading blanket close-up"), 7: ("Size and use panel for reading blanket with sofa, office, bed and travel icons.", "Wildflower reading blanket size and use panel")},
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def parse_customizer(summary):
    raw = Path(summary["html"]["html_file"]).read_text(encoding="utf-8", errors="ignore")
    page = html.unescape(raw)
    fields = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"minLength":\s*(\d+),\s*"maxLength":\s*(\d+)',
        page,
    )
    if not fields:
        fields = re.findall(
            r'"label"\s*:\s*"([^"]+)"\s*,\s*"required"\s*:\s*(true|false)\s*,\s*"minLength"\s*:\s*(\d+)\s*,\s*"maxLength"\s*:\s*(\d+)',
            page,
        )
    options = re.findall(r'"label"\s*:\s*"([^"]+)"\s*,\s*"required"\s*:\s*(true|false)\s*,\s*"defaultOptionId"', page)
    return fields, options


def customizer_sentence(fields, options):
    parts = []
    for label, required, min_length, max_length in fields:
        requirement = "required" if required == "true" else "optional"
        parts.append(f"Personalization uses a {requirement} {label} field, {min_length}-{max_length} characters.")
    visible_options = [label for label, _ in options if label and "confirmation" not in label.lower()]
    if visible_options:
        parts.append(f"Visible required choice fields include: {', '.join(visible_options)}.")
    labels = " ".join(label for label, *_ in fields).lower()
    if "number" in labels:
        parts.append("Names and numbers visible in mockups are sample artwork unless the matching input is entered.")
    elif fields:
        parts.append("Visible custom text in mockups is sample artwork unless the matching input is entered.")
    else:
        parts.append("No shopper text-entry field is described for this product.")
    return " ".join(parts)


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    missing = sorted(set(handles) - set(admin))
    if missing:
        raise RuntimeError(f"Missing handles in admin export: {missing}")
    return admin


def option_text(admin_row):
    values = [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]]
    values = [v for v in values if v]
    return ", ".join(values) if values else "available size options"


def description(pos, admin_row, ctext):
    title, _, _, detail, _ = PRODUCT_UPDATES[pos]
    return (
        f"<p>{title} features {detail}. "
        "The copy stays specific to the visible artwork, product form and selectable options for this exact item.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork: {detail}.</li>"
        "<li>Gallery images include the main mockup plus size, feature, care or lifestyle panels where shown.</li></ul>"
        "<h3>Options and Fit</h3>"
        f"<ul><li>Available option groups: {option_text(admin_row)}.</li>"
        f"<li>{ctext}</li>"
        "<li>Select the product type and size shown on the product page before checkout.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    summaries = {int(item["inventory_row"]["inventory_position"]): item for item in json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))}
    positions = set(PRODUCT_UPDATES)
    handles = {summaries[pos]["inventory_row"]["Handle"] for pos in positions}
    admin = load_admin(handles)
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    pos_by_key = {summaries[pos]["inventory_row"]["product_key"]: pos for pos in positions}
    handle_by_pos = {pos: summaries[pos]["inventory_row"]["Handle"] for pos in positions}
    fields_by_pos = {}
    option_fields_by_pos = {}
    for pos in positions:
        fields_by_pos[pos], option_fields_by_pos[pos] = parse_customizer(summaries[pos])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        handle = handle_by_pos[pos]
        title, primary, secondary, detail, cluster = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        meta = f"Shop {title.lower()} with {detail}, selectable sizes and verified custom text options."[:155]
        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or summaries[pos]["inventory_row"]["product_type"]
        ws.cell(row_num, idx["title_proposed"]).value = title
        ws.cell(row_num, idx["meta_title_seo"]).value = title
        ws.cell(row_num, idx["meta_title_length"]).value = len(title)
        ws.cell(row_num, idx["meta_description_seo"]).value = meta
        ws.cell(row_num, idx["meta_description_length"]).value = len(meta)
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, customizer_sentence(fields_by_pos[pos], option_fields_by_pos[pos]))
        ws.cell(row_num, idx["primary_keyword"]).value = primary
        ws.cell(row_num, idx["secondary_keywords"]).value = secondary
        ws.cell(row_num, idx["meta_keyword"]).value = primary
        ws.cell(row_num, idx["keyword_strategy"]).value = f"Target the specific product motif: {cluster}; keep broader quilt, blanket or gift terms for collection pages."
        ws.cell(row_num, idx["buyer_search_summary"]).value = f"US English buyer intent targets {cluster}; no search-volume, trend or ranking claim is made."
        ws.cell(row_num, idx["keyword_evidence_level"]).value = "SERP_ONLY"
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R2 deep recheck without prior QA: used products_export_1.csv sha256:{admin_hash}; "
            "separated Bigfoot quilt-set and personalized reading blanket intents, parsed custom fields and rewrote image observations/alts from contact sheets. Still NEEDS_REVIEW."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    seen_alts = set()
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pos = next((p for p, h in handle_by_pos.items() if h == handle), None)
        if pos is None:
            continue
        image_number = int(ws.cell(row_num, idx["image_number"]).value)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation, alt = IMAGE_DETAILS[pos][image_number]
        if alt in seen_alts:
            alt = f"{alt} for {PRODUCT_UPDATES[pos][4]}"
        seen_alts.add(alt)
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = alt[:125]
        ws.cell(row_num, idx["alt_action"]).value = "SET"
        ws.cell(row_num, idx["viewed_status"]).value = "VIEWED"
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = "R2: observation and alt rechecked from contact sheet; admin image URL and current alt matched from Shopify CSV where available."
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        title, primary, secondary, _, cluster = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in secondary.split(",")]
        keyword = primary if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = cluster
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "PRODUCT"
        ws.cell(row_num, idx["decision_reason"]).value = f"{title} is separated by verified product form, artwork motif and custom fields."
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = "products_export_1.csv admin baseline + storefront customizer parse + contact-sheet inspection"
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = "YES; nearby products are differentiated by motif, product form and custom field type."
        ws.cell(row_num, idx["mapping_reason"]).value = "R2 separates Bigfoot quilt-set designs from personalized reading blanket designs."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r2"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        title, primary, _, _, cluster = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {cluster} as a specific product page."
        ws.cell(row_num, idx["jtbd_statement"]).value = f"When shopping for {cluster}, the buyer wants exact artwork, product type, size choices and available custom text or choice fields confirmed before purchase."
        ws.cell(row_num, idx["functional_motivation"]).value = "Confirm product form, visible motif, size options, custom field rules, feature panels and image accuracy before purchase."
        ws.cell(row_num, idx["emotional_social_motivation"]).value = "Choose a personalized cabin, nature, reading or book-themed bedding item that matches the recipient or room theme."
        ws.cell(row_num, idx["purchase_concerns"]).value = "Confirm custom text rules where fields are available; avoid unsupported material, character identity, copyrighted story or delivery claims."
        ws.cell(row_num, idx["source_refs"]).value = f"products_export_1.csv sha256:{admin_hash}; storefront HTML customizer parse; contact sheets 241-250"
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R2_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No independent QA, Search Console, paid keyword volume or internal site-search export supplied; proof remains SERP-only."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {primary} with title: {title}."

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in positions:
            continue
        handle = handle_by_pos[pos]
        admin_row = admin[handle]
        _, _, _, detail, _ = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = f"Original storefront/product evidence; products_export_1.csv sha256:{admin_hash}; storefront customizer parse; inspected contact sheets"
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={option_text(admin_row)}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={detail}; customizer={customizer_sentence(fields_by_pos[pos], option_fields_by_pos[pos])}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "No independent QA for batch 025 yet; r2 prepared for QA recheck."
        ws.cell(row_num, idx["confidence_and_reason"]).value = "R2 uses contact sheets, storefront customizer parse and Shopify admin CSV baseline; no approval or import created."

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_025_r2_revision", "r2", "Separate 10-product revision without prior independent QA; source workbook was not modified."),
        ("qa_batch_025_r2_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_025_r2_scope", "inventory positions 241-250", "No products outside qa_batch_025 were revised."),
        ("qa_batch_025_r2_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_025_r2_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_025_r2_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else "" for col in range(1, ws.max_column + 1)])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)
    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_evidence": str(SUMMARY_PATH.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_025 only; inventory positions 241-250",
        "revision": "r2",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Ready for independent QA of qa_batch_025_r2.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join([
            "# Revision qa_batch_025_r2",
            "",
            "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
            "",
            f"- Nguồn r1: `{SOURCE.relative_to(ROOT)}`",
            f"- Evidence summary: `{SUMMARY_PATH.relative_to(ROOT)}`",
            f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
            f"- Admin export SHA-256: `{admin_hash}`",
            f"- Workbook r2: `{OUTPUT.relative_to(ROOT)}`",
            f"- Sản phẩm sửa: {revised_products}",
            f"- Ảnh sửa/refresh: {revised_images}",
            f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
            "- Sửa trọng yếu: rà lại khi chưa có QA, tách nhóm Bigfoot quilt set và personalized reading blanket, parse customizer từ HTML storefront, viết lại title/meta/description/alt theo từng contact sheet.",
            "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
            "- Bước tiếp theo: gửi `qa_batch_025_r2` cho QA độc lập.",
            "",
        ]),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
