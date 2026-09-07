import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "batches" / "SEO_Product_Optimization_through_batch_034.xlsx"
OUT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions"
RUN_OUT_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions"
OUTPUT = OUT_DIR / "SEO_Product_Optimization_after_QA_batch_001_007_r2.xlsx"
SUMMARY = RUN_OUT_DIR / "revision_after_QA_batch_001_007_r2_summary.json"
QA_ROOT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa"
QA_RUNS = [
    "20260907_083534",
    "20260907_091545",
    "20260907_094127",
    "20260907_095643",
    "20260907_101050",
    "20260907_102634",
    "20260907_104027",
]


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(cell):
    return "" if cell is None else str(cell)


def strip_custom_claims(value):
    value = text(value)
    replacements = [
        (r"\b[Pp]ersonalized\s+", ""),
        (r"\b[Pp]ersonalised\s+", ""),
        (r"\b[Cc]ustomizable\s+", ""),
        (r"\b[Cc]ustom\s+", ""),
        (r"\bwith custom name and number\b", "with sports artwork"),
        (r"\bwith name and number\b", "with sports artwork"),
        (r"\bwith custom name\b", "with themed artwork"),
        (r"\bwith name\b", "with themed artwork"),
        (r"\bcustomization\b", "product option"),
        (r"\bpersonalization\b", "product option"),
        (r"\b[Pp]ersonalize\b", "Style"),
        (r"\b[Cc]ustomize\b", "Choose"),
        (r"\bcustom-made\b", "themed"),
        (r"\bcustom\b", "themed"),
        (r"\bpersonalized\b", "themed"),
    ]
    for pattern, replacement in replacements:
        value = re.sub(pattern, replacement, value)
    value = re.sub(r"\s{2,}", " ", value).strip()
    return value


def clean_phrase(value):
    value = text(value)
    value = re.sub(r"Review size and personalization options before checkout", "Review size and pillowcase options before checkout", value, flags=re.I)
    value = re.sub(r"Review size and options before checkout", "Review size and pillowcase options before checkout", value, flags=re.I)
    value = re.sub(r"\bDragonfly\b", "", value, flags=re.I)
    value = re.sub(r"\bChristian Knight Templar\b", "Christian inspirational", value, flags=re.I)
    value = re.sub(r"\bD14\b", "", value)
    value = re.sub(r"\bD[0-9]{1,2}\b", "", value)
    value = re.sub(r"\bsnowman\b", "gingerbread figure", value, flags=re.I)
    value = re.sub(r"\bsnowman Christmas\b", "gingerbread Christmas", value, flags=re.I)
    value = re.sub(r"\balligator\b", "crocodile", value, flags=re.I)
    value = re.sub(r"\s{2,}", " ", value).strip(" ,;.-")
    return value


def extract_detail(desc_html, title):
    desc = text(desc_html)
    match = re.search(r"Bring distinctive (.*?) artwork", desc, flags=re.I | re.S)
    if match:
        detail = re.sub(r"<.*?>", " ", match.group(1))
    else:
        match = re.search(r"<li>(.*?)</li>", desc, flags=re.I | re.S)
        detail = re.sub(r"<.*?>", " ", match.group(1)) if match else title
    detail = clean_phrase(detail)
    detail = re.sub(r"\s+", " ", detail)
    return detail[:260] or title


def product_kind(product_type):
    raw = text(product_type).strip().lower()
    if "comforter" in raw:
        return "comforter set"
    if "blanket" in raw:
        return "blanket"
    if "quilt" in raw:
        return "quilt set"
    return "bedding item"


def build_description(row, idx, flagged_no_custom):
    title = clean_phrase(text(row[idx["title_proposed"] - 1]))
    detail = extract_detail(row[idx["description_proposed_html"] - 1], title)
    if flagged_no_custom:
        title = strip_custom_claims(title)
        detail = strip_custom_claims(detail)
    kind = product_kind(row[idx["product_type"] - 1])
    return (
        f"<p>Add themed style to the bedroom with this Jeminise {kind} featuring {detail}. "
        "The artwork is shown across the main bedding mockup and coordinated pillow or sham images where included.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork focus: {detail}.</li>"
        "<li>Gallery images show the main bedding view, detail views, fabric or care graphics, and included-component visuals where available.</li></ul>"
        "<h3>Product Options</h3>"
        "<ul><li>Choose from the available size and pillowcase options shown on the product page.</li>"
        "<li>Review the live product options before checkout to confirm the selected configuration.</li></ul>"
    )


def load_qa_data():
    product_positions = {}
    no_custom = set()
    product_issues = defaultdict(list)
    image_fixes = {}
    image_observations = {}
    issue_counts = defaultdict(lambda: defaultdict(int))

    for qa_run in QA_RUNS:
        data = json.loads((QA_ROOT / qa_run / "qa_dataset.json").read_text(encoding="utf-8"))
        for product in data["QA_Products"]:
            product_positions[product["product_key"]] = int(product["inventory_position"])
        qa_image_to_media = {}
        for image in data["QA_Images"]:
            key = (image["product_key"], text(image["media_id"]))
            qa_image_to_media[image["qa_image_key"]] = key
            image_observations[key] = image["qa_observation"]
        for issue in data["QA_Issues"]:
            pk = issue["product_key"]
            issue_counts[pk][issue["severity"]] += 1
            product_issues[pk].append(issue["issue_id"])
            if issue["severity"] == "CRITICAL" and re.search(r"personal|custom", issue["field"] + " " + issue["reason"], re.I):
                no_custom.add(pk)
            if issue["field"].startswith("image_") and issue["qa_image_key"]:
                key = qa_image_to_media.get(issue["qa_image_key"])
                if key and issue["recommended_fix"]:
                    image_fixes[key] = issue["recommended_fix"][:125]
    return product_positions, no_custom, product_issues, image_fixes, image_observations, issue_counts


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    wb = load_workbook(OUTPUT)
    now = datetime.now().astimezone().isoformat()
    product_positions, no_custom, product_issues, image_fixes, image_observations, issue_counts = load_qa_data()
    allowed_positions = set(range(1, 71))
    revised_products = 0
    revised_images = 0

    ws = wb["SEO_Products"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        pos = product_positions.get(pk)
        if pos not in allowed_positions:
            continue
        row_values = [ws.cell(row_num, col).value for col in range(1, ws.max_column + 1)]
        flagged = pk in no_custom
        for col_name in ["title_proposed", "meta_title_seo", "meta_description_seo", "primary_keyword", "secondary_keywords", "meta_keyword", "keyword_strategy", "buyer_search_summary"]:
            value = clean_phrase(ws.cell(row_num, idx[col_name]).value)
            if flagged:
                value = strip_custom_claims(value)
            ws.cell(row_num, idx[col_name]).value = value

        if pos == 24:
            ws.cell(row_num, idx["primary_keyword"]).value = "gingerbread Christmas quilt"
            ws.cell(row_num, idx["meta_keyword"]).value = "gingerbread Christmas quilt"
            ws.cell(row_num, idx["secondary_keywords"]).value = "gingerbread bedding, Christmas village quilt, holiday quilt set"
            ws.cell(row_num, idx["title_proposed"]).value = "Gingerbread Christmas Village Quilt Set"
            ws.cell(row_num, idx["meta_title_seo"]).value = "Gingerbread Christmas Village Quilt Set"
            ws.cell(row_num, idx["keyword_strategy"]).value = "Target the visible gingerbread Christmas village motif; do not use snowman wording for this product."

        ws.cell(row_num, idx["description_proposed_html"]).value = build_description(row_values, idx, flagged)
        ws.cell(row_num, idx["meta_title_length"]).value = len(text(ws.cell(row_num, idx["meta_title_seo"]).value))
        ws.cell(row_num, idx["meta_description_length"]).value = len(text(ws.cell(row_num, idx["meta_description_seo"]).value))
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R2 revision after QA batches 001-007: removed internal SEO/QA language; "
            f"{'removed unverified personalization/custom claims; ' if flagged else ''}"
            "remaining admin-export, material, source-mapping and cannibalization issues require QA recheck. "
            f"Original QA issue refs: {', '.join(product_issues[pk])}."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = None
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        # Image_Audit lacks product_key; map via Handle against the product rows.
        # Build it lazily from workbook rows after SEO_Products revisions.
        # Filled below using handle_to_pk.
    handle_to_pk = {}
    pidx = headers(wb["SEO_Products"])
    for row in wb["SEO_Products"].iter_rows(min_row=2, values_only=True):
        handle_to_pk[text(row[pidx["Handle"] - 1])] = text(row[pidx["product_key"] - 1])

    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        pos = product_positions.get(pk)
        if pos not in allowed_positions:
            continue
        media_id = text(ws.cell(row_num, idx["media_id"]).value)
        key = (pk, media_id)
        if key in image_observations:
            ws.cell(row_num, idx["observed_visual_details"]).value = image_observations[key]
        if key in image_fixes:
            ws.cell(row_num, idx["alt_proposed"]).value = image_fixes[key]
        else:
            alt = clean_phrase(ws.cell(row_num, idx["alt_proposed"]).value)
            if pk in no_custom:
                alt = strip_custom_claims(alt)
            ws.cell(row_num, idx["alt_proposed"]).value = alt[:125]
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R2 revision after QA: alt/observation refreshed from QA where an image issue had a recommended fix; "
            "still requires QA recheck and admin media-alt export before approval."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        pos = product_positions.get(pk)
        if pos not in allowed_positions:
            continue
        keyword = clean_phrase(ws.cell(row_num, idx["keyword"]).value)
        if pk in no_custom:
            keyword = strip_custom_claims(keyword)
        if pos == 24 and ws.cell(row_num, idx["keyword_role"]).value == "PRIMARY":
            keyword = "gingerbread Christmas quilt"
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["decision_reason"]).value = "R2: revised from QA issue set; product-level query must match visible artwork and verified purchase features."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r2"

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        pos = product_positions.get(pk)
        if pos not in allowed_positions:
            continue
        for col_name in ["purchase_context", "jtbd_statement", "functional_motivation", "emotional_social_motivation", "purchase_concerns", "seo_application"]:
            value = clean_phrase(ws.cell(row_num, idx[col_name]).value)
            if pk in no_custom:
                value = strip_custom_claims(value)
            ws.cell(row_num, idx[col_name]).value = value
        ws.cell(row_num, idx["limitations"]).value = text(ws.cell(row_num, idx["limitations"]).value) + " R2 note: QA batches 001-007 found issues requiring recheck before approval."

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        # evidence_id includes the original inventory position as the final component.
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        m = re.search(r"_(\d{3})$", evidence_id)
        if not m or int(m.group(1)) not in allowed_positions:
            continue
        ws.cell(row_num, idx["confidence_and_reason"]).value = "R2 revised from QA batches 001-007; content still needs independent QA recheck and human approval before import."

    ws = wb["README_QA"]
    h = headers(ws)
    for metric, value, definition in [
        ("revision_after_qa_001_007", "r2", "Created a separate revised workbook for products 1-70 after QA batches 001-007; source workbook was not modified."),
        ("r2_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("r2_revised_images", str(revised_images), "Image_Audit rows refreshed and kept NEEDS_REVIEW."),
    ]:
        ws.append([metric if c == h["metric"] else value if c == h["value"] else definition if c == h["definition"] else "" for c in range(1, ws.max_column + 1)])

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    summary = {
        "created_at": now,
        "source_workbook": str(SOURCE.relative_to(ROOT)),
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_001 through qa_batch_007; inventory positions 1-70",
        "revision": "r2",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "no_custom_claims_removed_products": len(no_custom),
        "image_recommended_fixes_applied": len(image_fixes),
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "issue_counts_by_product": {pk: dict(counts) for pk, counts in issue_counts.items()},
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(SUMMARY.relative_to(ROOT))
    print(f"revised_products={revised_products} revised_images={revised_images} no_custom_products={len(no_custom)} image_fixes={len(image_fixes)}")


if __name__ == "__main__":
    main()
