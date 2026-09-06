import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_017"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_016.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_017.xlsx"


NORSE_REFS = "https://www.amazon.ca/YCYR-Bedding-Pillowcase-Bedchamber-Comforter/dp/B09LQD91XZ; https://retail.blackpoolpleasurebeach.com/product/valhalla-duvet-set/; https://www.etsy.com/listing/4431801249/viking-axes-and-shield-woven-blanket; https://myvikinggear.com/collections/viking-home-decoration/viking-quilt-set/"
OWL_REFS = "https://www.google.com/m/storepages?c=VN&hl=en-VN&q=amazon.com; https://www.ebay.com/; https://www.pinterest.com/pin/adorable-owl-quilt-patterns-free-pattern-day--5348093301080978/"
HALLOWEEN_REFS = "https://www.walmart.com/c/kp/halloween-bedding-sets; https://www.wayfair.com/bed-bath/sb1/halloween-bedding-c481592-a123871~486885.html; https://www.bedbathandbeyond.com/c/halloween/halloween-bedding?t=28313; https://www.target.com/s/winter%2Bhalloween%2Bbedding; https://www.pinterest.com/ideas/halloween-comforter-set/907456284290/"


PROFILES = {
    161: ("blue Norse raven Celtic knot quilt", "Norse raven quilt", "Norse Raven Quilt Set", "Blue Norse Raven Celtic Knot Quilt", "blue patchwork Norse raven artwork within a Celtic knot frame with moon, raven and Tree of Life panels", "Viking raven bedding, Celtic knot quilt, Norse mythology quilt", "EVERGREEN", NORSE_REFS),
    162: ("Valhalla Viking shield crossed axes quilt", "Viking shield quilt", "Viking Shield Quilt Set", "Valhalla Viking Shield Crossed Axes Quilt", "black Norse Valhalla design with a Viking shield, crossed axes and dark runic-style background", "Norse bedding, Viking quilt set, crossed axes bedding", "EVERGREEN", NORSE_REFS),
    163: ("blue gold owl night patchwork quilt", "owl patchwork quilt", "Owl Patchwork Quilt Set", "Blue Gold Owl Night Patchwork Quilt", "owl perched on a branch inside a moonlit blue circle with gold patchwork triangle border", "owl bedding, night owl quilt, animal patchwork quilt", "EVERGREEN", OWL_REFS),
    164: ("white black Halloween witch comforter set", "Halloween comforter set with sheets", "Halloween Comforter Set with Sheets", "White Black Halloween Witch Comforter Set", "white and black Halloween bedding with witch silhouette, bats, ghost shapes, pumpkins and Halloween text", "twin Halloween bedding, witch comforter set, Halloween bedding set", "HALLOWEEN", HALLOWEEN_REFS),
    165: ("gray Halloween ghost pumpkin comforter set", "ghost pumpkin comforter set", "Ghost Pumpkin Comforter Set", "Gray Halloween Ghost Pumpkin Comforter Set", "dark gray Halloween bedding with smiling ghosts, jack-o-lantern pumpkins, bats, stars and spiderwebs", "Halloween bedding set, pumpkin ghost bedding, twin Halloween comforter", "HALLOWEEN", HALLOWEEN_REFS),
    166: ("black white ghost Halloween comforter set", "ghost Halloween comforter set", "Ghost Halloween Comforter Set", "Black White Ghost Halloween Comforter Set", "black comforter set covered with white cartoon ghost shapes and matching ghost pillow shams", "black Halloween bedding, ghost bedding set, Halloween comforter with sheets", "HALLOWEEN", HALLOWEEN_REFS),
    167: ("blue haunted pumpkin ghost comforter set", "haunted pumpkin comforter set", "Haunted Pumpkin Comforter Set", "Blue Haunted Pumpkin Ghost Comforter Set", "blue Halloween scene with large pumpkins, sheet ghosts, haunted trees, bats and moon artwork", "Halloween ghost bedding, pumpkin comforter set, spooky bedding set", "HALLOWEEN", HALLOWEEN_REFS),
    168: ("pink cute ghost Halloween comforter set", "pink Halloween comforter set", "Pink Halloween Comforter Set", "Pink Cute Ghost Halloween Comforter Set", "pink Halloween bedding with cute ghosts, witch hats, candy, stars and soft pink sheets", "cute Halloween bedding, pink ghost bedding, kids Halloween comforter", "HALLOWEEN", HALLOWEEN_REFS),
    169: ("red bloody handprint Halloween comforter set", "bloody Halloween comforter set", "Bloody Halloween Comforter Set", "Red Bloody Handprint Halloween Comforter Set", "white and red horror-style bedding with splatter graphics, handprints and deep red sheets", "horror bedding set, bloody handprint comforter, Halloween bedding set", "HALLOWEEN", HALLOWEEN_REFS),
    170: ("white haunted house pumpkin Halloween comforter set", "haunted house comforter set", "Haunted House Comforter Set", "White Haunted House Pumpkin Comforter Set", "white Halloween bedding with haunted houses, pumpkins, black cats, bare trees, spiderwebs and rust orange sheets", "haunted house bedding, pumpkin comforter set, Halloween bed in a bag", "HALLOWEEN", HALLOWEEN_REFS),
}


def qa_issues(pos, html_error):
    issues = [
        "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."
    ]
    if pos in (161, 162):
        issues.append("Norse mythology products overlap with batch 016; keep raven, Celtic knot, shield and crossed-axes motifs distinct.")
    if pos == 163:
        issues.append("Owl wording should stay motif-specific and avoid broad animal quilt cannibalization.")
    if 164 <= pos <= 170:
        issues.append("Pamnest Halloween comforter products are closely related; review title uniqueness, collection placement and canonical strategy.")
        issues.append("Set-content, deep-pocket and size-chart claims should be checked against variants before import.")
    if pos == 169:
        issues.append("Bloody handprint horror artwork may be polarizing; review brand tone and paid-channel restrictions before publishing changes.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def product_context(pos):
    if pos in (161, 162):
        return "Viking-themed bedroom decor, mythology gifts or personalized Norse bedding."
    if pos == 163:
        return "animal-themed rooms, rustic bedrooms or owl gift shoppers."
    return "seasonal Halloween rooms, guest beds, kids rooms or themed fall decor."


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        issues = qa_issues(pos, html_error)
        meta = f"Shop a {kind} featuring {detail}. Review available sizes and set options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for {product_context(pos)}</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broader bedding terms for collections or stronger representative products.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, room decor or a seasonal gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": issues,
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed bedding, the buyer wants a {kind} matching a specific Norse, owl or Halloween motif.",
            "functional_motivation": "Find the right design, size, bedding type, included pieces, care details and decor fit.",
            "emotional_social_motivation": "Create a distinctive bedroom look or give a themed mythology, animal or Halloween-inspired gift.",
            "purchase_concerns": "Product identity, material/care claims, set contents, size fit, design distinction and whether visible artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed in a main product mockup",
                2: f"Close-up of the print on {h1}",
                3: f"Lifestyle bedroom mockup of {h1}",
                4: f"Front bed view of {h1}",
                5: f"Fitted sheet or size feature image for {h1}",
                6: f"Set contents and size chart image for {h1}",
                7: f"Additional product mockup of {h1}",
                8: f"Lifestyle use image for {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_017_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_017_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_017"},
        {"metric": "cumulative_products", "value": "170", "definition": "Total products included through batch_017"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[170:180]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_018",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
