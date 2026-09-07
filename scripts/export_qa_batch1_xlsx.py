from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
QA_RUN_ID = "20260907_083534"
RUN_ID = "20260906_234129"
DATA = ROOT / "seo_runs" / "jeminise.com" / RUN_ID / "qa" / QA_RUN_ID / "qa_workbook_payload.json"
OUTPUT = ROOT / "resutls" / "jeminise.com" / RUN_ID / "qa" / QA_RUN_ID / "SEO_QA_qa_batch_001.xlsx"
SHEETS = ("QA_Summary", "QA_Products", "QA_Criteria", "QA_Images", "QA_Issues")

NAVY = "17365D"
BLUE = "2F75B5"
PALE_BLUE = "D9EAF7"
WHITE = "FFFFFF"
TEXT = "1F2937"
GRID = "CBD5E1"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "FCE4D6"
PURPLE = "E4DFEC"


def safe(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, separators=(", ", ": "))
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def col_map(ws):
    return {cell.value: cell.column for cell in ws[1]}


def cell_ref(ws, header, row, absolute=False):
    col = get_column_letter(col_map(ws)[header])
    return f"${col}${row}" if absolute else f"{col}{row}"


def sheet_col(ws, header, last_row):
    col = get_column_letter(col_map(ws)[header])
    return f"'{'%s' % ws.title}'!${col}$2:${col}${last_row}"


def write_rows(ws, rows):
    headers = list(rows[0].keys()) if rows else []
    ws.append(headers)
    for row in rows:
        ws.append([safe(row.get(h)) for h in headers])


def add_formulas(wb):
    wp = wb["QA_Products"]
    wc = wb["QA_Criteria"]
    wi = wb["QA_Images"]
    ws = wb["QA_Summary"]
    wq = wb["QA_Issues"]

    # Per-image formulas, fully traceable to IM1–IM4.
    weights = {"IM1": 40, "IM2": 30, "IM3": 20, "IM4": 10}
    for r in range(2, wi.max_row + 1):
        point_terms, weight_terms = [], []
        for cid, weight in weights.items():
            ref = cell_ref(wi, cid, r)
            point_terms.append(f'IF({ref}="FULL",{weight},IF({ref}="PARTIAL",{weight}/2,0))')
            weight_terms.append(f'IF({ref}="NOT_CHECKED",0,{weight})')
        wi[cell_ref(wi, "image_verified_points", r)] = "=" + "+".join(point_terms)
        wi[cell_ref(wi, "image_assessed_weight", r)] = "=" + "+".join(weight_terms)
        vp = cell_ref(wi, "image_verified_points", r)
        aw = cell_ref(wi, "image_assessed_weight", r)
        wi[cell_ref(wi, "image_final_score", r)] = f'=IF({aw}=100,{vp},"")'
        wi[cell_ref(wi, "image_score_lower_bound", r)] = f"={vp}"
        wi[cell_ref(wi, "image_score_upper_bound", r)] = f"={vp}+(100-{aw})"

    # Product criteria formulas; I1 derives from QA_Images rather than manual scoring.
    im_pk = sheet_col(wi, "product_key", wi.max_row)
    im_pts = sheet_col(wi, "image_verified_points", wi.max_row)
    im_aw = sheet_col(wi, "image_assessed_weight", wi.max_row)
    for r in range(2, wc.max_row + 1):
        cid = cell_ref(wc, "criterion_id", r)
        ass = cell_ref(wc, "assessment", r)
        weight = cell_ref(wc, "weight", r)
        pk = cell_ref(wc, "product_key", r)
        wc[cell_ref(wc, "rating", r)] = f'=IF({cid}="I1",AVERAGEIF({im_pk},{pk},{im_pts})/100,IF({ass}="FULL",1,IF({ass}="PARTIAL",0.5,IF({ass}="FAIL",0,""))))'
        rating = cell_ref(wc, "rating", r)
        wc[cell_ref(wc, "earned_points", r)] = f'=IF({rating}="",0,{weight}*{rating})'
        wc[cell_ref(wc, "assessed_weight", r)] = f'=IF({cid}="I1",20*AVERAGEIF({im_pk},{pk},{im_aw})/100,IF({ass}="NOT_CHECKED",0,{weight}))'

    # Product rollups and blocking-status formula.
    c_pk = sheet_col(wc, "product_key", wc.max_row)
    c_ep = sheet_col(wc, "earned_points", wc.max_row)
    c_aw = sheet_col(wc, "assessed_weight", wc.max_row)
    q_pk = sheet_col(wq, "product_key", wq.max_row)
    q_sev = sheet_col(wq, "severity", wq.max_row)
    for r in range(2, wp.max_row + 1):
        pk = cell_ref(wp, "product_key", r)
        vp = cell_ref(wp, "verified_points", r)
        aw = cell_ref(wp, "assessed_weight", r)
        lower = cell_ref(wp, "score_lower_bound", r)
        upper = cell_ref(wp, "score_upper_bound", r)
        final = cell_ref(wp, "final_score", r)
        coverage = cell_ref(wp, "image_coverage", r)
        inventory_ok = cell_ref(wp, "image_inventory_complete", r)
        critical = cell_ref(wp, "critical_count", r)
        major = cell_ref(wp, "major_count", r)
        wp[vp] = f"=SUMIF({c_pk},{pk},{c_ep})"
        wp[aw] = f"=SUMIF({c_pk},{pk},{c_aw})"
        wp[lower] = f"={vp}"
        wp[upper] = f"={vp}+(100-{aw})"
        wp[final] = f'=IF(AND({aw}=100,{coverage}=1,{inventory_ok}=TRUE),{vp},"")'
        wp[critical] = f'=COUNTIFS({q_pk},{pk},{q_sev},"CRITICAL")'
        wp[major] = f'=COUNTIFS({q_pk},{pk},{q_sev},"MAJOR")'
        wp[cell_ref(wp, "minor_count", r)] = f'=COUNTIFS({q_pk},{pk},{q_sev},"MINOR")'
        wp[cell_ref(wp, "limitation_count", r)] = f'=COUNTIFS({q_pk},{pk},{q_sev},"LIMITATION")'
        wp[cell_ref(wp, "qa_status", r)] = f'=IF({critical}>0,"QA_FAIL",IF({final}="","QA_INCOMPLETE",IF({final}<70,"QA_FAIL",IF(OR({final}<85,{major}>0),"QA_REVISE","QA_PASS"))))'

    # Summary metrics reference product formulas rather than duplicating results.
    metrics = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    p_final = sheet_col(wp, "final_score", wp.max_row)
    p_status = sheet_col(wp, "qa_status", wp.max_row)
    if "batch_final_score" in metrics:
        ws.cell(metrics["batch_final_score"], 2, f"=AVERAGE({p_final})")
    if "status_counts" in metrics:
        ws.cell(metrics["status_counts"], 2, f'="PASS="&COUNTIF({p_status},"QA_PASS")&"; REVISE="&COUNTIF({p_status},"QA_REVISE")&"; FAIL="&COUNTIF({p_status},"QA_FAIL")&"; INCOMPLETE="&COUNTIF({p_status},"QA_INCOMPLETE")')
    if "batch_result" in metrics:
        ws.cell(metrics["batch_result"], 2, f'=IF(COUNTIF({p_status},"QA_PASS")=ROWS({p_status}),"PASSED","NOT_PASSED")')


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 32
    thin = Side(style="thin", color=GRID)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos Display", size=11, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=BLUE))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=10, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
    # Practical widths: identifiers/URLs stay inspectable, prose wraps.
    for c in range(1, ws.max_column + 1):
        header = str(ws.cell(1, c).value or "")
        if header in {"product_key", "submitted_value", "source_observation", "reason", "recommended_fix", "supporting_evidence", "recheck_condition", "qa_observation", "submitted_observation", "alt_effective", "evidence_refs", "issue_refs", "definition", "value"}:
            width = 42
        elif "url" in header or "workbook" in header or "evidence" in header:
            width = 38
        elif header in {"qa_image_key", "media_id", "workbook_image_id", "criterion_id", "severity", "qa_status"}:
            width = 20
        else:
            max_len = max(len(str(ws.cell(r, c).value or "")) for r in range(1, min(ws.max_row, 80) + 1))
            width = min(max(max_len + 2, 11), 24)
        ws.column_dimensions[get_column_letter(c)].width = width
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 84 if ws.title == "QA_Issues" else (66 if ws.title == "QA_Images" else 42)

    # Clickable single-URL fields.
    for header in ("url", "image_url_source", "image_url_workbook"):
        cmap = col_map(ws)
        if header in cmap:
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(r, cmap[header])
                if isinstance(cell.value, str) and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

    cmap = col_map(ws)
    if "qa_status" in cmap:
        col = get_column_letter(cmap["qa_status"])
        rng = f"{col}2:{col}{ws.max_row}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}2="QA_PASS"'], fill=PatternFill("solid", fgColor=GREEN)))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}2="QA_REVISE"'], fill=PatternFill("solid", fgColor=AMBER)))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}2="QA_FAIL"'], fill=PatternFill("solid", fgColor=RED)))
    if "severity" in cmap:
        col = get_column_letter(cmap["severity"])
        rng = f"{col}2:{col}{ws.max_row}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}2="CRITICAL"'], fill=PatternFill("solid", fgColor="F4CCCC")))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}2="MAJOR"'], fill=PatternFill("solid", fgColor=RED)))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}2="MINOR"'], fill=PatternFill("solid", fgColor=AMBER)))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}2="LIMITATION"'], fill=PatternFill("solid", fgColor=PURPLE)))

    for header in ("verified_points", "assessed_weight", "score_lower_bound", "score_upper_bound", "final_score", "rating", "earned_points", "image_verified_points", "image_assessed_weight", "image_final_score", "image_score_lower_bound", "image_score_upper_bound"):
        if header in cmap:
            for r in range(2, ws.max_row + 1):
                ws.cell(r, cmap[header]).number_format = "0.0"
    if "image_coverage" in cmap:
        for r in range(2, ws.max_row + 1):
            ws.cell(r, cmap["image_coverage"]).number_format = "0.0%"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = {"QA_Products": 2, "QA_Images": 3, "QA_Issues": 2}.get(ws.title, 1)
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.print_options.horizontalCentered = False
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5


def main():
    payload = json.loads(DATA.read_text(encoding="utf-8"))
    if list(payload) != list(SHEETS):
        raise ValueError(f"Expected exactly {SHEETS}; got {tuple(payload)}")
    wb = Workbook()
    wb.remove(wb.active)
    for name in SHEETS:
        ws = wb.create_sheet(name)
        write_rows(ws, payload[name])
    add_formulas(wb)
    for ws in wb.worksheets:
        style_sheet(ws)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)

    # Structural verification after reopening.
    check = load_workbook(OUTPUT, data_only=False, read_only=False)
    assert check.sheetnames == list(SHEETS)
    # Validate against the payload so the same audited exporter can be reused by
    # later fixed-size batches (the final batch and image counts may differ).
    for name in SHEETS:
        assert check[name].max_row == len(payload[name]) + 1
    formulas = sum(1 for ws in check.worksheets for row in ws.iter_rows() for c in row if c.data_type == "f")
    assert formulas >= 500
    print(json.dumps({"output": str(OUTPUT), "sheets": check.sheetnames,
                      "rows": {ws.title: ws.max_row - 1 for ws in check.worksheets}, "formulas": formulas}, indent=2))


if __name__ == "__main__":
    main()
