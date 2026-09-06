import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_019"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_018.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_019.xlsx"


BASEBALL_REFS = "https://www.etsy.com/listing/1648190989/personalized-baseball-bedding-set-for; https://www.amazon.com/Baseball-Personalized-Comforter-Customizable-Competitive/dp/B0GDDZ4JPF; https://www.walmart.com/c/kp/baseball-bedding-set; https://www.youcustomizeit.com/p/Baseball-Duvet-Cover-Personalized/100020; https://www.ohaprints.com/products/personalized-baseball-duvet-cover-set-baseball-ball-glove-america-flag-player-gift-duvet-cover-pillowcases-custom-name-number-bedding-set; https://www.amorcustomgifts.com/en-au/products/baseball-american-flag-personalized-name-040921dkn-bedding-set1"


PROFILES = {
    181: ("brown baseball glove custom name bedding", "custom name baseball bedding", "Custom Name Baseball Bedding", "Brown Baseball Glove Custom Name Bedding", "brown baseball bedding with glove leather textures, large baseball, script custom name and jersey number", "baseball glove bedding, personalized baseball comforter, custom sports bedding", "EVERGREEN", BASEBALL_REFS),
    182: ("American flag baseball bedding with script name", "American flag baseball bedding", "American Flag Baseball Bedding", "American Flag Baseball Bedding with Custom Name", "red white and blue American flag bedding with oversized baseball artwork and script custom name", "personalized baseball bedding, baseball flag comforter, custom name sports bedding", "EVERGREEN", BASEBALL_REFS),
    183: ("baseball flag bedding with block name and number", "baseball name number bedding", "Baseball Name Number Bedding", "Baseball Flag Bedding with Name and Number", "distressed American flag baseball bedding with glove corners, oversized baseball, block custom name and jersey number", "personalized baseball bedding, baseball number comforter, American flag baseball bedding", "EVERGREEN", BASEBALL_REFS),
    184: ("close-up baseball glove bedding with name on ball", "baseball glove bedding", "Baseball Glove Bedding Set", "Close-Up Baseball Glove Bedding with Custom Name", "close-up baseball glove bedding with a baseball in the pocket and custom name printed on the ball", "personalized baseball comforter, baseball ball bedding, custom baseball duvet", "EVERGREEN", BASEBALL_REFS),
    185: ("gray baseball American flag bedding with custom name", "gray baseball flag bedding", "Gray Baseball Flag Bedding", "Gray Baseball American Flag Bedding with Custom Name", "distressed American flag bedding with gray baseball glove artwork, large custom name and matching baseball pillow shams", "custom baseball bedding, American flag comforter, personalized sports bedding", "EVERGREEN", BASEBALL_REFS),
    186: ("red white blue baseball pattern bedding with name", "baseball pattern bedding", "Baseball Pattern Bedding Set", "Red White Blue Baseball Pattern Bedding with Custom Name", "red, white and blue baseball bedding with repeated baseball pattern, stripes, jersey number and large custom name", "personalized baseball bedding, patriotic baseball comforter, custom name bedding", "EVERGREEN", BASEBALL_REFS),
    187: ("pitcher silhouette baseball bedding with custom name", "pitcher baseball bedding", "Pitcher Baseball Bedding Set", "Pitcher Silhouette Baseball Bedding with Custom Name", "cream baseball bedding with large pitcher silhouette, red baseball stitching and bold custom name", "baseball player bedding, personalized baseball comforter, custom sports bedding", "EVERGREEN", BASEBALL_REFS),
    188: ("black baseball glove bedding with name and number", "black baseball bedding", "Black Baseball Bedding Set", "Black Baseball Glove Bedding with Name and Number", "black baseball bedding with grayscale glove and ball artwork, script custom name and jersey number", "custom baseball bedding, baseball glove comforter, personalized sports bedding", "EVERGREEN", BASEBALL_REFS),
    189: ("blue striped baseball bedding with name on ball", "blue baseball bedding", "Blue Baseball Bedding Set", "Blue Striped Baseball Bedding with Custom Name", "light blue striped bedding with oversized baseball graphic and custom name printed across the ball", "personalized baseball bedding, baseball ball comforter, custom name bedding", "EVERGREEN", BASEBALL_REFS),
    190: ("galaxy baseball bedding with custom name and number", "galaxy baseball bedding", "Galaxy Baseball Bedding Set", "Galaxy Baseball Bedding with Custom Name and Number", "dark blue galaxy baseball bedding with large baseball, script custom name and bold jersey number", "personalized baseball comforter, baseball number bedding, custom sports bedding", "EVERGREEN", BASEBALL_REFS),
}


def qa_issues(pos, html_error):
    issues = [
        "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment.",
        "Personalized baseball products are very close; review custom-name placement, jersey number accuracy and duplicate keyword cannibalization before import.",
        "Feature images mention zipper closure, machine washable and high-density weaving; verify actual product type, fabric and set configuration before deployment.",
    ]
    if pos in (181, 184, 188):
        issues.append("Glove-and-ball motifs overlap; keep color, close-up angle and name/number placement distinct.")
    if pos in (182, 183, 185, 186):
        issues.append("American flag baseball motifs overlap; review title uniqueness and collection/canonical strategy.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for baseball rooms, sports gifts, kids bedrooms or personalized fan bedding.</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broad baseball bedding terms for collections or the strongest representative product.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, a sports room accent or a personalized gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": qa_issues(pos, html_error),
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for personalized sports bedding, the buyer wants a {kind} matching a baseball motif, custom name, number or patriotic design.",
            "functional_motivation": "Find the right baseball design, size, personalization fields, care details and room fit.",
            "emotional_social_motivation": "Create a baseball-themed bedroom or give a custom sports gift to a player or fan.",
            "purchase_concerns": "Personalization spelling, jersey number, product identity, material/care claims, size fit and whether artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed in a main product mockup",
                2: f"Feature or care image for {h1}",
                3: f"Product detail or closure image for {h1}",
                4: f"Fabric or high-density weaving image for {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_019_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_019_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_019"},
        {"metric": "cumulative_products", "value": "190", "definition": "Total products included through batch_019"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[190:200]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_020",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
