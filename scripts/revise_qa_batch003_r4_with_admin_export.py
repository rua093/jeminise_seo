import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import revise_qa_batch002_r4_with_admin_export as base


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_003_r4"
QA_RUN_ID = "20260907_140123"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_003_r3" / "SEO_Product_Optimization_qa_batch_003_r3.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
CUSTOMIZER_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_audit.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_003_r4.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_003_r4.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    21: {
        "title": "Personalized Gingerbread Candy Cane Christmas Quilt Set",
        "meta_title": "Personalized Gingerbread Candy Cane Christmas Quilt Set",
        "meta_description": "Add optional custom text to a gingerbread candy cane Christmas quilt with holiday artwork, sizes and pillowcase options.",
        "primary": "personalized gingerbread candy cane Christmas quilt",
        "secondary": "gingerbread candy cane quilt set, custom Christmas quilt, holiday quilt bedding",
        "cluster": "personalized gingerbread candy cane Christmas quilt",
        "intent_role": "Design 05 page with verified optional Custom Your Name field up to 200 characters.",
        "detail": "gingerbread and candy cane Christmas quilt artwork with holiday bedding details",
    },
    22: {
        "title": "Personalized Black Christmas Tree Quilt Set",
        "meta_title": "Personalized Black Christmas Tree Quilt Set",
        "meta_description": "Add optional custom text to a black Christmas tree quilt set with holiday artwork, quilt sizes and pillowcase options.",
        "primary": "personalized black Christmas tree quilt set",
        "secondary": "black Christmas tree quilt, custom Christmas quilt set, holiday tree bedding",
        "cluster": "personalized black Christmas tree quilt set",
        "intent_role": "Design 06 page with verified optional Custom Your Name field up to 200 characters.",
        "detail": "black Christmas tree holiday quilt artwork with coordinated bedding visuals",
    },
    23: {
        "title": "Personalized Winter Cardinal Birdhouse Quilt Set",
        "meta_title": "Personalized Winter Cardinal Birdhouse Quilt Set",
        "meta_description": "Add optional custom text to a winter cardinal birdhouse quilt with Christmas artwork, sizes and pillowcase options.",
        "primary": "personalized winter cardinal birdhouse quilt set",
        "secondary": "cardinal birdhouse quilt, custom Christmas cardinal quilt, winter bird quilt bedding",
        "cluster": "personalized winter cardinal birdhouse quilt set",
        "intent_role": "Design 08 page with verified optional Custom Your Name field up to 200 characters.",
        "detail": "winter cardinal birdhouse Christmas quilt artwork with snowy holiday details",
    },
    24: {
        "title": "Personalized Gingerbread Christmas Village Quilt Set",
        "meta_title": "Personalized Gingerbread Christmas Village Quilt Set",
        "meta_description": "Add optional custom text to a gingerbread Christmas village quilt with festive artwork, sizes and pillowcase options.",
        "primary": "personalized gingerbread Christmas village quilt",
        "secondary": "gingerbread village quilt set, custom Christmas quilt, holiday village bedding",
        "cluster": "personalized gingerbread Christmas village quilt",
        "intent_role": "Design 09 page with verified optional Custom Your Name field up to 200 characters.",
        "detail": "gingerbread Christmas village holiday quilt artwork with festive bedding styling",
    },
    25: {
        "title": "Personalized Snowman and Cardinals Christmas Quilt Set",
        "meta_title": "Personalized Snowman and Cardinals Christmas Quilt Set",
        "meta_description": "Add optional custom text to a snowman and cardinals Christmas quilt with winter artwork, sizes and pillowcase options.",
        "primary": "personalized snowman cardinal Christmas quilt",
        "secondary": "snowman cardinal quilt set, custom Christmas snowman quilt, winter cardinal bedding",
        "cluster": "personalized snowman cardinal Christmas quilt",
        "intent_role": "Design 10 page corrected from the r2 motif mismatch to verified snowman and cardinals artwork.",
        "detail": "snowman and cardinals Christmas quilt artwork with winter holiday bedding details",
    },
    26: {
        "title": "Personalized White Christmas Tree Quilt Set",
        "meta_title": "Personalized White Christmas Tree Quilt Set",
        "meta_description": "Add optional custom text to a white Christmas tree quilt set with holiday artwork, sizes and pillowcase options.",
        "primary": "personalized white Christmas tree quilt set",
        "secondary": "white Christmas tree quilt, custom holiday tree quilt, Christmas bedding quilt",
        "cluster": "personalized white Christmas tree quilt set",
        "intent_role": "Design 11 page with verified optional Custom Your Name field up to 200 characters.",
        "detail": "white Christmas tree holiday quilt artwork with coordinated bedding visuals",
    },
    27: {
        "title": "Personalized Two Snowmen and Cardinals Christmas Quilt Set",
        "meta_title": "Personalized Two Snowmen Cardinals Christmas Quilt Set",
        "meta_description": "Add optional custom text to a two snowmen and cardinals Christmas quilt with winter artwork, sizes and pillowcase options.",
        "primary": "personalized two snowmen cardinal Christmas quilt",
        "secondary": "two snowmen cardinal quilt, custom Christmas snowman quilt, cardinal winter bedding",
        "cluster": "personalized two snowmen cardinal Christmas quilt",
        "intent_role": "Design 12 page corrected from the r2 motif mismatch to verified two snowmen and cardinals artwork.",
        "detail": "two snowmen and cardinals Christmas quilt artwork with winter holiday bedding details",
    },
    28: {
        "title": "Personalized Vintage Christmas Tree Quilt Set",
        "meta_title": "Personalized Vintage Christmas Tree Quilt Set",
        "meta_description": "Add optional custom text to a vintage Christmas tree quilt set with holiday artwork, sizes and pillowcase options.",
        "primary": "personalized vintage Christmas tree quilt set",
        "secondary": "vintage Christmas tree quilt, custom holiday quilt set, Christmas tree bedding",
        "cluster": "personalized vintage Christmas tree quilt set",
        "intent_role": "Christmas bedding page with verified optional Custom Your Name field up to 200 characters.",
        "detail": "vintage Christmas tree holiday quilt artwork with coordinated bedding visuals",
    },
    29: {
        "title": "Cardinal Memorial Christmas Quilt Set",
        "meta_title": "Cardinal Memorial Christmas Quilt Set",
        "meta_description": "Shop a cardinal memorial Christmas quilt with I Am Always With You text, winter patchwork artwork, sizes and pillowcase options.",
        "primary": "cardinal memorial Christmas quilt",
        "secondary": "I Am Always With You cardinal quilt, Christmas cardinal quilt set, memorial cardinal bedding",
        "cluster": "cardinal memorial Christmas quilt",
        "intent_role": "Cardinal memorial quilt page with no verified personalization input in the customizer audit.",
        "detail": "Christmas cardinal on branch patchwork quilt with I Am Always With You text",
    },
    30: {
        "title": "Cardinal Christmas Wreath Quilt Set",
        "meta_title": "Cardinal Christmas Wreath Quilt Set",
        "meta_description": "Shop a cardinal Christmas wreath quilt with pine, holly berries, pinecones, sample Sophia artwork, sizes and pillowcase options.",
        "primary": "cardinal Christmas wreath quilt",
        "secondary": "cardinal wreath quilt set, Christmas cardinal quilt, winter cardinal bedding",
        "cluster": "cardinal Christmas wreath quilt",
        "intent_role": "Cardinal wreath quilt page; Sophia is treated as sample artwork because no personalization input was verified.",
        "detail": "cardinal perched with pine, holly berries, pinecones and Sophia shown as sample artwork",
    },
}


IMAGE_DETAILS = {
    21: [
        ("Gingerbread candy cane Christmas quilt bed mockup with cream background, green border and matching shams.", "Gingerbread candy cane Christmas quilt set on bed"),
        ("Bedroom mockup showing gingerbread candy cane quilt with coordinated pillow shams.", "Gingerbread candy cane quilt bedroom mockup"),
        ("Close-up of gingerbread face, green bow, candy cane stripes and visible quilt stitching.", "Gingerbread candy cane quilt stitching close-up"),
        ("Matching pillow sham with gingerbread cookie, candy cane and decorative green border.", "Gingerbread candy cane Christmas pillow sham"),
        ("Size and component panel showing one premium quilt and two optional standard shams.", "Gingerbread candy cane quilt size and sham options"),
    ],
    22: [
        ("Black Christmas tree quilt bed mockup with decorated tree, gift boxes, dark border and matching shams.", "Black Christmas tree quilt set on bed"),
        ("Bedroom mockup showing black Christmas tree quilt with presents and coordinated pillow shams.", "Black Christmas tree quilt bedroom mockup"),
        ("Matching pillow sham with Christmas tree, presents and dark holiday border.", "Black Christmas tree quilt pillow sham"),
        ("Close-up of decorated Christmas tree branches, ornaments and visible quilt stitching on a dark background.", "Black Christmas tree quilt stitching close-up"),
        ("Size and component panel showing one premium quilt and two optional standard shams.", "Black Christmas tree quilt size and sham options"),
    ],
    23: [
        ("Winter cardinal birdhouse quilt bed mockup with red cardinals, snowy birdhouse and matching shams.", "Winter cardinal birdhouse quilt set on bed"),
        ("Bedroom mockup showing red cardinals around a snowy birdhouse on a pale winter quilt.", "Winter cardinal birdhouse quilt bedroom mockup"),
        ("Matching pillow sham with red cardinals, snowy birdhouse and winter branches.", "Winter cardinal birdhouse pillow sham"),
        ("Close-up of red cardinal, snow-covered birdhouse, berries and visible quilt stitching.", "Winter cardinal birdhouse quilt stitching close-up"),
        ("Size and component panel showing one premium quilt and two optional standard shams.", "Winter cardinal birdhouse quilt size options"),
    ],
    24: [
        ("Gingerbread Christmas village quilt bed mockup with snowman, presents, village lights and matching shams.", "Gingerbread Christmas village quilt set on bed"),
        ("Bedroom mockup showing gingerbread Christmas village quilt with colorful gifts and snowy night sky.", "Gingerbread village quilt bedroom mockup"),
        ("Matching pillow sham with gingerbread figure, moonlit village, presents and snowy holiday scene.", "Gingerbread Christmas village pillow sham"),
        ("Close-up of gingerbread snowman scarf, village houses and visible quilt stitching.", "Gingerbread village quilt stitching close-up"),
        ("Size and component panel showing one premium quilt and two optional standard shams.", "Gingerbread village quilt size and sham options"),
    ],
    25: [
        ("Snowman and cardinals Christmas quilt bed mockup with red background, pine trees and matching shams.", "Snowman cardinals Christmas quilt set on bed"),
        ("Bedroom mockup showing snowman cardinal Christmas quilt with red border and coordinated pillow shams.", "Snowman cardinals quilt bedroom mockup"),
        ("Matching pillow sham with snowman, cardinals, pine trees and red winter background.", "Snowman cardinals Christmas pillow sham"),
        ("Close-up of snowman face, red scarf, cardinal and visible quilt stitching.", "Snowman cardinals quilt stitching close-up"),
        ("Size and component panel showing one premium quilt and two optional standard shams.", "Snowman cardinals quilt size and sham options"),
    ],
    26: [
        ("White Christmas tree quilt bed mockup with red and gold ornaments, star topper and matching shams.", "White Christmas tree quilt set on bed"),
        ("Bedroom mockup showing white Christmas tree quilt with black check border and coordinated pillow shams.", "White Christmas tree quilt bedroom mockup"),
        ("Matching pillow sham with decorated Christmas tree on a light background.", "White Christmas tree quilt pillow sham"),
        ("Close-up of Christmas tree branches, red and gold ornaments, star shapes and visible quilt stitching.", "White Christmas tree quilt stitching close-up"),
        ("Size and component panel showing one premium quilt and two optional standard shams.", "White Christmas tree quilt size and sham options"),
    ],
    27: [
        ("Two snowmen and cardinals Christmas quilt bed mockup with red background, wreath frame and matching shams.", "Two snowmen cardinals Christmas quilt set"),
        ("Bedroom mockup showing two snowmen with cardinals on a red Christmas quilt.", "Two snowmen cardinals quilt bedroom mockup"),
        ("Matching pillow sham with two snowmen, cardinal pair and wreath frame on red background.", "Two snowmen cardinals Christmas pillow sham"),
        ("Close-up of cardinals perched near two snowmen with visible quilt stitching.", "Two snowmen cardinals quilt stitching close-up"),
        ("Size and component panel showing one premium quilt and two optional standard shams.", "Two snowmen cardinals quilt size options"),
    ],
    28: [
        ("Vintage Christmas tree quilt bed mockup with cream background, gifts, ornaments and matching shams.", "Vintage Christmas tree quilt set on bed"),
        ("Bedroom mockup showing vintage Christmas tree quilt with presents and coordinated pillow shams.", "Vintage Christmas tree quilt bedroom mockup"),
        ("Matching pillow sham with decorated Christmas tree, presents and vintage border artwork.", "Vintage Christmas tree quilt pillow sham"),
        ("Close-up of decorated tree branches, ornaments, garland and visible quilt stitching.", "Vintage Christmas tree quilt stitching close-up"),
        ("Size and component panel showing one premium quilt and two optional standard shams.", "Vintage Christmas tree quilt size and sham options"),
    ],
    29: [
        ("Cardinal memorial Christmas quilt bed mockup with red patchwork, poinsettias, hearts, bells and I Am Always With You text.", "Cardinal memorial Christmas quilt set on bed"),
        ("Close-up of red cardinal and I Am Always With You text with berries and visible quilt stitching.", "Cardinal memorial quilt text close-up"),
        ("Matching pillow sham with cardinal memorial patchwork, poinsettias, hearts and bells.", "Cardinal memorial Christmas pillow sham"),
        ("Bedroom mockup showing cardinal memorial patchwork quilt with coordinated pillow shams.", "Cardinal memorial quilt bedroom mockup"),
        ("Size and component panel showing one premium quilt and two optional standard shams.", "Cardinal memorial quilt size and sham options"),
    ],
    30: [
        ("Cardinal Christmas wreath quilt bed mockup with Sophia sample name, pine, holly berries, pinecones and matching shams.", "Cardinal Christmas wreath quilt set on bed"),
        ("Close-up of red cardinal with pine, holly berries, pinecones and visible quilt stitching.", "Cardinal wreath quilt stitching close-up"),
        ("Matching pillow sham with cardinal wreath artwork and Sophia sample name.", "Cardinal Christmas wreath pillow sham"),
        ("Bedroom mockup showing cardinal wreath quilt with Sophia sample artwork and coordinated shams.", "Cardinal wreath quilt bedroom mockup"),
        ("Size and component panel showing one premium quilt and two optional standard shams.", "Cardinal wreath quilt size and sham options"),
    ],
}


def patch_base_globals():
    base.ROOT = ROOT
    base.SHOP = SHOP
    base.RUN_ID = RUN_ID
    base.BATCH_ID = BATCH_ID
    base.QA_RUN_ID = QA_RUN_ID
    base.SOURCE = SOURCE
    base.QA_DATASET = QA_DATASET
    base.CUSTOMIZER_AUDIT = CUSTOMIZER_AUDIT
    base.ADMIN_EXPORT = ADMIN_EXPORT
    base.RESULT_DIR = RESULT_DIR
    base.RUN_DIR = RUN_DIR
    base.OUTPUT = OUTPUT
    base.REPORT = REPORT
    base.SUMMARY = SUMMARY
    base.PRODUCT_UPDATES = PRODUCT_UPDATES


def safe_description(pos, admin_row, customizer_text):
    item = PRODUCT_UPDATES[pos]
    option_names = [v for v in [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]] if v]
    options = ", ".join(option_names) if option_names else "available product selectors"
    options_heading = "Options and Customization" if pos <= 28 else "Options"
    product_type = (admin_row["Type"] or "quilt").lower()
    customization_line = f"<li>{customizer_text}</li>" if pos <= 28 else ""
    return (
        f"<p>{item['detail'].capitalize()} gives this Jeminise {product_type} a clear product-specific holiday bedding focus.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork focus: {item['detail']}.</li>"
        "<li>Gallery images show the main bed mockup plus close-up, sham, size or component graphics where present.</li>"
        "<li>The size panel lists quilt choices and two optional standard shams when shown.</li></ul>"
        f"<h3>{options_heading}</h3>"
        f"<ul><li>Available selectors cover {options}.</li>"
        f"{customization_line}"
        "<li>Use the live selectors to confirm quilt size and pillowcase choices before checkout.</li></ul>"
    )


def main():
    patch_base_globals()
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)

    qa = base.load_qa()
    admin = base.load_admin(set(qa["handle_by_key"].values()))
    admin_hash = base.sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()
    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = base.headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = base.text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        customizer_text = base.customizer_sentence(pos, qa["customizer_by_pos"].get(pos, {}))
        if pos == 29:
            customizer_text = "No name-entry field is used in the SEO claim for this revision."
        if pos == 30:
            customizer_text = "No name-entry field is used in the SEO claim for this revision; Sophia is described only as sample artwork."

        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or ws.cell(row_num, idx["product_type"]).value
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = safe_description(pos, admin_row, customizer_text)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"Buyer intent targets {item['cluster']} for US English product search; no search-volume claim is made."
        )
        ws.cell(row_num, idx["revision"]).value = "r4"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R4 deep audit after qa_batch_003_r3: used products_export_1.csv (sha256:{admin_hash}) as admin baseline; "
            "refreshed product copy and all image observations/alts from contact sheets, preserved optional name-field claims only for products 21-28, "
            "and kept unsupported personalization claims out of products 29-30. Still NEEDS_REVIEW; no APPROVED/import."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = base.headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    for row_num in range(2, ws.max_row + 1):
        handle = base.text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        media_id = base.text(ws.cell(row_num, idx["media_id"]).value)
        key = (pk, media_id)
        image_url = base.text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(base.normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        pos = qa["pos_by_key"][pk]
        image_number = ws.cell(row_num, idx["image_number"]).value
        try:
            image_number = int(image_number)
        except (TypeError, ValueError):
            image_number = revised_images + 1
        details = IMAGE_DETAILS.get(pos, [])
        if 1 <= image_number <= len(details):
            observation, proposed_alt = details[image_number - 1]
        else:
            observation = qa["observation_by_key"].get(key) or f"{PRODUCT_UPDATES[pos]['detail']} shown in gallery image {image_number}."
            proposed_alt = base.compact_observation_to_alt(observation, f"{PRODUCT_UPDATES[pos]['title']} image {image_number}")
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = proposed_alt[:125]
        ws.cell(row_num, idx["revision"]).value = "r4"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R4: image-level observation and alt refreshed from contact sheet review; current admin alt and image URL matched from products_export_1.csv where possible."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = base.headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = base.text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = base.text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "Product"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = "Sang re-QA r2 SERP evidence + products_export_1.csv admin baseline"
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = "R4 separates sibling product intent by visible motif, gallery evidence and verified customizer requirements."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r4"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = base.headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = base.text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as a specific Christmas quilt product."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm the design, quilt/pillowcase choices "
            "and whether the name field is available."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = "Confirm product type, size, pillowcase choices, care/fabric panels and customization requirement."
        ws.cell(row_num, idx["emotional_social_motivation"]).value = "Choose a clear holiday bedding gift with artwork that matches the page images."
        ws.cell(row_num, idx["purchase_concerns"]).value = "Name-field rules, exact bedding type, included components, fabric/care claims and image accuracy."
        ws.cell(row_num, idx["source_refs"]).value = f"products_export_1.csv sha256:{admin_hash}; Sang QA r2 {QA_RUN_ID}"
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R4_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No search-volume source supplied; Shopify admin CSV baseline is now available for stored title/meta/body/image alt."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = base.headers(ws)
    for row_num in range(2, ws.max_row + 1):
        evidence_id = base.text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = qa["product_keys"][pos - 21]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = f"Storefront/product.js from Sang QA r2; products_export_1.csv sha256:{admin_hash}"
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={admin_row['Option1 Name']}, "
            f"{admin_row['Option2 Name']}, {admin_row['Option3 Name']}; status={admin_row['Status']}; "
            f"image_count={len(admin_row['images'])}; design={item['detail']}."
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "R4 deepened qa_batch_003_r3 with image-level evidence and admin export cross-check; needs independent QA recheck."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R4 uses contact-sheet image review, live Customizer audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = base.headers(ws)
    for metric, value, definition in [
        ("qa_batch_003_r4_revision", "r4", "Deep review of latest qa_batch_003_r3; source r3 workbook was not modified."),
        ("qa_batch_003_r4_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_003_r4_scope", "inventory positions 21-30", "No products outside qa_batch_003 were revised."),
        ("qa_batch_003_r4_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_003_r4_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_003_r4_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa_r2": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_003 only; inventory positions 21-30",
        "revision": "r4",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "r4_changes": [
            "Deepened image-level observed_visual_details and alt_proposed for all 50 batch images.",
            "Kept products 25 and 27 locked to snowman/cardinal motifs, not gingerbread.",
            "Kept optional Enter Name personalization only for products 21-28.",
            "Kept products 29-30 free of unsupported personalization claims.",
            "Removed customer-facing Shopify export/admin wording from description_proposed_html.",
        ],
        "next_step": "Sang QA lại qa_batch_003_r4 before any approval or import.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_003_r4",
                "",
                "Artifact rà sâu chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r3: `{SOURCE.relative_to(ROOT)}`",
                f"- QA r2 của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r4: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Sửa trọng yếu: product 25/27 giữ đúng snowman/cardinal; product 30 chỉ ghi Sophia là sample artwork; 50 ảnh có alt riêng.",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Bước tiếp theo: Sang QA lại `qa_batch_003_r4`.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
