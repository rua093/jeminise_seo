from pathlib import Path
import json

import export_qa_batch1_xlsx as exporter
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RUN_ID, QA_RUN_ID = "20260906_234129", "20260907_102634"
exporter.DATA = ROOT / "seo_runs" / "jeminise.com" / RUN_ID / "qa" / QA_RUN_ID / "qa_workbook_payload.json"
exporter.OUTPUT = ROOT / "resutls" / "jeminise.com" / RUN_ID / "qa" / QA_RUN_ID / "SEO_QA_qa_batch_006.xlsx"


if __name__ == "__main__":
    exporter.main()
    wb = load_workbook(exporter.OUTPUT, data_only=False)
    audit = {
        "path": str(exporter.OUTPUT),
        "sheet_names": wb.sheetnames,
        "expected_sheet_names": ["QA_Summary", "QA_Products", "QA_Criteria", "QA_Images", "QA_Issues"],
        "row_counts": {ws.title: ws.max_row - 1 for ws in wb.worksheets},
        "freeze_panes": {ws.title: str(ws.freeze_panes) for ws in wb.worksheets},
        "auto_filters": {ws.title: ws.auto_filter.ref for ws in wb.worksheets},
        "formula_cells": sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")),
        "hyperlink_cells": sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.hyperlink),
    }
    audit["passed"] = audit["sheet_names"] == audit["expected_sheet_names"] and audit["row_counts"] == {"QA_Summary": 13, "QA_Products": 10, "QA_Criteria": 110, "QA_Images": 78, "QA_Issues": 86} and all(audit["freeze_panes"].values()) and all(audit["auto_filters"].values())
    path = ROOT / "seo_runs" / "jeminise.com" / RUN_ID / "qa" / QA_RUN_ID / "spreadsheet_validation.json"
    path.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))
