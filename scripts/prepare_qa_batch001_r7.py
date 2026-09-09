from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import subprocess
from datetime import datetime
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
QA_RUN_ID = "20260908_230550"
BATCH_ID = "qa_batch_001_r7"
SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID / f"SEO_Product_Optimization_{BATCH_ID}.xlsx"
INVENTORY = ROOT / "seo_runs" / SHOP / RUN_ID / "inventory.csv"
QA_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID
OUT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "qa" / QA_RUN_ID
EXPECTED_ADMIN_HASH = "729CBD9B3E0248F5D9196D999A49222DBDACA85C2F3CF8524BEEC549B57D738C"
ADMIN_COMMIT = "f3e38d54ae3d1f45297089cb313ede4696eecfe5"


def digest_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest().upper()


def digest_file(path: Path) -> str:
    return digest_bytes(path.read_bytes())


def rows(ws):
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def fetch(url: str) -> tuple[requests.Response, bool]:
    headers = {"User-Agent": "Mozilla/5.0 independent SEO QA batch001-r7"}
    try:
        response = requests.get(url, headers=headers, timeout=45)
        response.raise_for_status()
        return response, True
    except requests.exceptions.SSLError:
        response = requests.get(url, headers=headers, timeout=45, verify=False)
        response.raise_for_status()
        return response, False


def safe_name(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", value)[:100]


def main():
    QA_DIR.mkdir(parents=True, exist_ok=False)
    OUT_DIR.mkdir(parents=True, exist_ok=False)
    for name in ("source_snapshot", "live", "images"):
        (QA_DIR / name).mkdir()

    source_hash = digest_file(SOURCE)
    snapshot = QA_DIR / "source_snapshot" / SOURCE.name
    shutil.copy2(SOURCE, snapshot)
    shutil.copy2(INVENTORY, QA_DIR / "source_snapshot" / INVENTORY.name)

    historical = subprocess.check_output(["git", "show", f"{ADMIN_COMMIT}:products_export_1.csv"], cwd=ROOT)
    historical_path = QA_DIR / "source_snapshot" / "products_export_1_expected_hash.csv"
    historical_path.write_bytes(historical)
    if digest_bytes(historical) != EXPECTED_ADMIN_HASH:
        raise RuntimeError("Historical admin export hash does not match revision reference")
    current_admin = ROOT / "products_export_1.csv"
    shutil.copy2(current_admin, QA_DIR / "source_snapshot" / "products_export_1_current.csv")

    with INVENTORY.open(encoding="utf-8-sig", newline="") as fh:
        inventory = list(csv.DictReader(fh))[:10]
    keys = {x["product_key"] for x in inventory}
    handles = {x["Handle"] for x in inventory}
    urls = {x["product_url"] for x in inventory}
    positions = {x["product_key"]: int(x["inventory_position"]) for x in inventory}

    wb = load_workbook(snapshot, read_only=True, data_only=True)
    products = [x for x in rows(wb["SEO_Products"]) if x.get("product_key") in keys]
    images = [x for x in rows(wb["Image_Audit"]) if x.get("Handle") in handles]
    evidence = [x for x in rows(wb["Product_Evidence"]) if x.get("product_url") in urls]
    keywords = [x for x in rows(wb["Keyword_Map"]) if x.get("product_key") in keys]
    buyers = [x for x in rows(wb["Buyer_Search_Research"]) if x.get("product_key") in keys]
    products.sort(key=lambda x: positions[x["product_key"]])
    by_handle = {x["Handle"]: x for x in products}

    live_records = []
    for inv in inventory:
        pos = int(inv["inventory_position"])
        url = inv["product_url"]
        html_response, html_verified = fetch(url)
        js_response, js_verified = fetch(url + ".js")
        html_path = QA_DIR / "live" / f"{pos}_{safe_name(inv['Handle'])}.html"
        js_path = QA_DIR / "live" / f"{pos}_{safe_name(inv['Handle'])}.js.json"
        html_path.write_bytes(html_response.content)
        js_path.write_bytes(js_response.content)
        soup = BeautifulSoup(html_response.text, "html.parser")
        product_js = js_response.json()
        decoded = html_response.text
        customizer_hits = []
        for pattern in ("Enter Name", "Custom Your Name", "Custom Your Number", "Customize Your Quilt", "personaliz"):
            for match in re.finditer(pattern, decoded, flags=re.I):
                excerpt = decoded[max(0, match.start() - 800):match.end() + 1200]
                customizer_hits.append({"pattern": pattern, "excerpt": excerpt[:2200]})
                break
        live_records.append({
            "inventory_position": pos,
            "product_key": inv["product_key"],
            "handle": inv["Handle"],
            "product_id_inventory": inv["product_id"],
            "checked_at": datetime.now().astimezone().isoformat(),
            "status_code": html_response.status_code,
            "ssl_verified": html_verified and js_verified,
            "html_sha256": digest_bytes(html_response.content),
            "json_sha256": digest_bytes(js_response.content),
            "html_path": str(html_path),
            "json_path": str(js_path),
            "title_element": soup.title.get_text(" ", strip=True) if soup.title else "",
            "h1": soup.find("h1").get_text(" ", strip=True) if soup.find("h1") else "",
            "meta_description": (soup.find("meta", attrs={"name": "description"}) or {}).get("content", ""),
            "canonical": (soup.find("link", rel="canonical") or {}).get("href", ""),
            "product_js": {
                "id": str(product_js.get("id", "")),
                "handle": product_js.get("handle", ""),
                "title": product_js.get("title", ""),
                "description": product_js.get("description", ""),
                "vendor": product_js.get("vendor", ""),
                "type": product_js.get("type", ""),
                "options": product_js.get("options", []),
                "variants": product_js.get("variants", []),
                "images": product_js.get("images", []),
                "media": product_js.get("media", []),
            },
            "customizer_hits": customizer_hits,
        })

    image_manifest = []
    for image in sorted(images, key=lambda x: (positions[next(p["product_key"] for p in products if p["Handle"] == x["Handle"])], int(x.get("image_number") or 0))):
        product = by_handle[image["Handle"]]
        pos = positions[product["product_key"]]
        number = int(image.get("image_number") or 0)
        url = image.get("image_url_export") or image.get("image_url")
        response, verified = fetch(url)
        suffix = Path(url.split("?", 1)[0]).suffix or ".jpg"
        local = QA_DIR / "images" / f"{pos:03d}_{number:02d}_{safe_name(str(image.get('media_id') or 'media'))}{suffix}"
        local.write_bytes(response.content)
        with Image.open(local) as opened:
            size = opened.size
            mode = opened.mode
            opened.verify()
        url_hash = hashlib.sha256(url.encode("utf-8")).hexdigest()[:16]
        media_token = str(image.get("media_id") or url_hash)
        qa_key = f"{product['product_key']}|{media_token}|{image.get('image_location') or 'GALLERY'}|{number}"
        image_manifest.append({
            **image,
            "product_key": product["product_key"],
            "position": pos,
            "qa_image_key": qa_key,
            "local_path": str(local),
            "download_sha256": digest_file(local),
            "width": size[0],
            "height": size[1],
            "mode": mode,
            "ssl_verified": verified,
            "downloaded_at": datetime.now().astimezone().isoformat(),
        })

    raw = {
        "products": [{**x, "position": positions[x["product_key"]]} for x in products],
        "images": image_manifest,
        "sheets": {"Product_Evidence": evidence, "Keyword_Map": keywords, "Buyer_Search_Research": buyers},
        "inventory": inventory,
        "live": live_records,
    }
    (QA_DIR / "qa_raw_source.json").write_text(json.dumps(raw, ensure_ascii=False, indent=2), encoding="utf-8")
    (QA_DIR / "live_source_comparison.json").write_text(json.dumps(live_records, ensure_ascii=False, indent=2), encoding="utf-8")
    (QA_DIR / "image_download_manifest.json").write_text(json.dumps(image_manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    manifest = {
        "qa_run_id": QA_RUN_ID,
        "qa_batch_id": BATCH_ID,
        "revision": "r7",
        "shop": SHOP,
        "market": "United States",
        "language": "English",
        "source_workbook": str(SOURCE),
        "source_sha256_start": source_hash,
        "snapshot_sha256": digest_file(snapshot),
        "inventory_positions": "1-10",
        "product_keys": [x["product_key"] for x in inventory],
        "expected_products": 10,
        "expected_images": 65,
        "expected_keyword_rows": 40,
        "expected_buyer_rows": 10,
        "expected_evidence_rows": 10,
        "historical_admin_export_commit": ADMIN_COMMIT,
        "historical_admin_export_sha256": digest_file(historical_path),
        "current_admin_export_sha256": digest_file(current_admin),
        "rubric": "prompt_qa.md v1.0; prompt.md v2.4",
        "started_at": datetime.now().astimezone().isoformat(),
        "status": "EVIDENCE_COLLECTED",
        "awaiting_confirmation": False,
    }
    (QA_DIR / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    (QA_DIR / "qa_progress.json").write_text(json.dumps({
        "qa_run_id": QA_RUN_ID,
        "qa_batch_id": BATCH_ID,
        "current_stage": "EVIDENCE_COLLECTED",
        "completed_image_keys": [],
        "awaiting_confirmation": False,
        "last_saved_at": datetime.now().astimezone().isoformat(),
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "qa_run_id": QA_RUN_ID,
        "source_hash": source_hash,
        "products": len(products),
        "images": len(image_manifest),
        "keyword_rows": len(keywords),
        "buyer_rows": len(buyers),
        "evidence_rows": len(evidence),
        "live_pages": len(live_records),
        "ssl_unverified": sum(not x["ssl_verified"] for x in live_records),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
