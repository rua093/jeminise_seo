import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_031"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_030.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_031.xlsx"


CACTUS_REFS = "https://www.amazon.com/cactus-duvet-cover/s?k=cactus+duvet+cover; https://www.wayfair.com/keyword.php?keyword=cactus+bedding; https://www.pinterest.com/ideas/colorful-cactus-bedding-set/901450680955/; https://alphaquilt.com/collections/cactus"
CELTIC_REFS = "https://www.amazon.com/clp/B0DNSVY2QJ; https://luvingift.com/product/viking-quilt-set-norse-celtic-yggdrasil-tree-of-life/; https://myvikinggear.com/collections/viking-home-decoration/viking-quilt-set/; https://www.pinterest.com/pin/562668547209120403/"
READING_REFS = "https://www.etsy.com/market/book_lover_blanket; https://www.etsy.com/listing/1749170765/just-a-girl-who-loves-books-blanket; https://www.amazon.com/Just-Girl-Loves-Books-Blanket/dp/B0DMVXGZKG"
BASKETBALL_REFS = "https://stock.adobe.com/search?k=watercolor+basketball; https://www.pinterest.com/pin/personalized-basketball-quilt-set-basketball-ball-watercolor-quilt-blanket-with-pillowcases-custom-name-and-number-qui--999939923509377708/; https://www.etsy.com/market/personalized_basketball_comforter"
DEER_REFS = "https://www.etsy.com/market/personalized_deer_blanket_couples; https://www.youcustomizeit.com/p/Deer-Comforters-Personalized/177732; https://foter.com/wildlife-comforter-sets; https://www.pinterest.com/pin/eternity-couple-deer-in-the-forest-couple-quilt-and-bedding-set-3d-printed-personalized-red--4598245680246284800/"


PROFILES = {
    301: ("vibrant cactus flower quilt", "cactus flower quilt", "Cactus Flower Quilt Set", "Vibrant Cactus Flower Quilt Set", "cactus quilt with large red and yellow flowers, tall cactus columns and textured southwestern garden artwork", "cactus bedding, floral cactus quilt, southwest quilt set", "EVERGREEN", CACTUS_REFS),
    302: ("embroidered desert cactus quilt", "embroidered cactus quilt", "Embroidered Cactus Quilt Set", "Vibrant Embroidered Desert Cactus Quilt Set", "bright embroidered-style cactus quilt with tall teal cactus plants, pink and yellow flowers and warm desert colors", "desert cactus bedding, colorful cactus quilt, southwest bedding", "EVERGREEN", CACTUS_REFS),
    303: ("flowering cactus desert scene quilt", "flowering cactus quilt", "Flowering Cactus Desert Quilt Set", "Vibrant Flowering Cactus Desert Quilt Set", "desert scene quilt with flowering cactus, layered hills, colorful blossoms and dark lower border", "cactus bedding, desert flower quilt, southwest quilt", "EVERGREEN", CACTUS_REFS),
    304: ("potted cactus garden quilt", "potted cactus quilt", "Potted Cactus Quilt Set", "Vibrant Potted Cactus Garden Quilt Set", "yellow quilt with repeated potted cactus and succulent plants in colorful pots and a warm garden grid layout", "succulent bedding, cactus quilt set, potted plant bedding", "EVERGREEN", CACTUS_REFS),
    305: ("neutral potted cactus quilt", "potted cactus bedding", "Potted Cactus Bedding Set", "Vibrant Potted Cactus Quilt Set", "neutral cream quilt with potted cactus and succulent plants, patterned pots and soft southwestern accent colors", "cactus bedding set, succulent quilt, potted cactus quilt", "EVERGREEN", CACTUS_REFS),
    306: ("Viking Celtic Tree of Life quilt", "Viking Tree of Life quilt", "Viking Tree of Life Quilt Set", "Viking Celtic Tree of Life Quilt Set", "blue Viking quilt with custom name, Celtic knot border, rune-style symbols and Tree of Life compass medallion", "Celtic bedding, Yggdrasil quilt, Viking quilt set", "EVERGREEN", CELTIC_REFS),
    307: ("vintage reading girl flowers blanket", "vintage reading girl blanket", "Vintage Reading Girl Blanket", "Vintage Reading Girl with Flowers Blanket", "vintage newspaper-style reading blanket with flowers, cartoon reading girl, custom name and soft beige background", "book lover blanket, personalized reading blanket, vintage book blanket", "EVERGREEN", READING_REFS),
    308: ("vintage Yggdrasil Tree of Life quilt", "vintage Yggdrasil quilt", "Vintage Yggdrasil Quilt Set", "Vintage Yggdrasil Tree of Life Quilt Set", "black and ivory quilt with Yggdrasil Tree of Life, sun and moon medallion and vintage botanical border", "Tree of Life bedding, Celtic quilt, Yggdrasil bedding", "EVERGREEN", CELTIC_REFS),
    309: ("watercolor basketball dunking blanket", "watercolor basketball blanket", "Watercolor Basketball Blanket", "Watercolor Basketball Players Dunking Blanket", "personalized basketball blanket with watercolor dunking players, hoop, red and yellow stars, class text and jersey number", "custom basketball blanket, basketball player gift, personalized sports blanket", "EVERGREEN", BASKETBALL_REFS),
    310: ("two deer forest romance comforter", "deer romance comforter", "Deer Romance Comforter Set", "Two Deer Forest Romance Comforter", "dark forest comforter with two deer standing together, floral accents, You and Me text and custom couple names", "deer couple bedding, wildlife comforter, personalized deer comforter", "EVERGREEN", DEER_REFS),
}


def qa_issues(pos, html_error):
    issues = ["Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."]
    if 301 <= pos <= 305:
        issues.append("Cactus quilt products are close variants and overlap batches 028-030; review motif separation, title uniqueness and canonical strategy.")
        issues.append("Verify quilt/bedspread construction, optional pillow shams and machine-washable claims.")
    if pos in (306, 308):
        issues.append("Viking/Celtic/Yggdrasil products overlap earlier Tree of Life variants; keep Viking compass, rune and vintage sun-moon motifs distinct.")
    if pos == 307:
        issues.append("Reading blanket overlaps earlier book lover products; verify character option, sample name, size variants and fleece/care claims.")
    if pos == 309:
        issues.append("Basketball blanket overlaps sports bedding batches; verify class/number personalization, product type and blanket size claims.")
    if pos == 310:
        issues.append("Deer romance comforter overlaps deer couple products; verify custom names, comforter versus quilt naming and optional shams.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def product_context(pos):
    if 301 <= pos <= 305:
        return "southwestern bedrooms, desert decor, cactus-themed rooms or garden-inspired bedding."
    if pos in (306, 308):
        return "Viking-inspired rooms, Celtic bedrooms, spiritual decor or Tree of Life gift bedding."
    if pos == 307:
        return "reading nooks, bedrooms, book clubs, librarian gifts or personalized reader gifts."
    if pos == 309:
        return "basketball rooms, sports bedrooms, player gifts or personalized fan bedding."
    return "rustic bedrooms, hunting cabins, anniversary gifts or personalized couple bedding."


def alt_for_image(h1, pos, ipos):
    if pos in (301, 302, 303, 304, 305, 306, 308, 310):
        return {
            1: f"{h1} displayed on a bed in the main quilt mockup",
            2: f"Angled bedroom or flat product mockup for {h1}",
            3: f"Close-up pillow sham and artwork detail for {h1}",
            4: f"Fabric and layered bedding feature image for {h1}",
            5: f"Size chart for {h1}",
            6: f"Bedspread construction and care feature image for {h1}",
            7: f"Overhead bedroom mockup of {h1}",
            8: f"Additional bedroom mockup for {h1}",
        }.get(ipos, f"Product detail image for {h1}")[:125]
    return {
        1: f"{h1} shown as a personalized throw blanket",
        2: f"Main lifestyle or flat product image for {h1}",
        3: f"Size chart or lifestyle image for {h1}",
        4: f"Fabric and fleece feature image for {h1}",
        5: f"Couch or close-up mockup for {h1}",
        6: f"Use or fabric guide image for {h1}",
        7: f"Size guide or product option image for {h1}",
        8: f"Family or lifestyle image for {h1}",
    }.get(ipos, f"Product detail image for {h1}")[:125]


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for {product_context(pos)}</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broad cactus, Viking, reading, basketball and deer bedding terms for collection pages or strongest representatives.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, room decor or a personalized gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": qa_issues(pos, html_error),
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed bedding, the buyer wants a {kind} matching a specific cactus, Viking, reading, basketball or deer motif.",
            "functional_motivation": "Find the right design, size, product type, personalization fields, care details and room fit.",
            "emotional_social_motivation": "Give a personalized sports, reader or couple gift, or decorate with desert, Celtic or wildlife artwork.",
            "purchase_concerns": "Personalization spelling, product identity, material/care claims, size fit, included components and whether artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt_for_image(h1, pos, ipos), "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_031_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_031_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_031"},
        {"metric": "cumulative_products", "value": "310", "definition": "Total products included through batch_031"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[310:320]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_032",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
