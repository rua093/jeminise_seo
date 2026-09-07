import csv
import hashlib
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_007_r3"
QA_RUN_ID = "20260907_164400"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_007_r2" / "SEO_Product_Optimization_qa_batch_007_r2.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
CUSTOMIZER_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_audit.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_007_r3.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_007_r3.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    61: {
        "title": "Personalized Football Yard Line Comforter Set",
        "meta_title": "Personalized Football Yard Line Comforter Set",
        "meta_description": "Customize a football yard line comforter with required name, optional number, product type, size and bedding options.",
        "primary": "personalized football yard line comforter",
        "secondary": "custom football yard line bedding, football comforter with name, football field comforter",
        "cluster": "personalized football yard line comforter",
        "intent_role": "Yard-line perspective football page; verified required Enter Name up to 25 characters and optional Enter Number up to 5 characters.",
        "detail": "football above yard lines perspective comforter artwork with sample name and number only as examples",
    },
    62: {
        "title": "Personalized American Flag Football Comforter Set",
        "meta_title": "Personalized American Flag Football Comforter Set",
        "meta_description": "Customize an American flag football comforter with required name, optional number, product type, size and bedding options.",
        "primary": "personalized American flag football comforter",
        "secondary": "custom American flag football bedding, football flag comforter with name, patriotic football comforter",
        "cluster": "personalized American flag football comforter",
        "intent_role": "Close-up football on American flag page; separated from general patriotic and helmet flag variants.",
        "detail": "football close-up on American flag comforter artwork with sample name and number only as examples",
    },
    63: {
        "title": "Personalized Flaming Lightning Football Comforter Set",
        "meta_title": "Personalized Flaming Lightning Football Comforter Set",
        "meta_description": "Customize a flaming lightning football comforter with required name, optional number, product type, size and bedding options.",
        "primary": "personalized flaming lightning football comforter",
        "secondary": "custom fire football comforter, lightning football bedding with name, flaming football bedding",
        "cluster": "personalized flaming lightning football comforter",
        "intent_role": "Flames and lightning football page; separated from helmet-flame and flag variants.",
        "detail": "football engulfed in flames and lightning comforter artwork with sample name and number only as examples",
    },
    64: {
        "title": "Personalized Football Helmet Flag Comforter Set",
        "meta_title": "Personalized Football Helmet Flag Comforter Set",
        "meta_description": "Customize a football helmet flag comforter with required name, optional number, product type, size and bedding options.",
        "primary": "personalized football helmet flag comforter",
        "secondary": "custom football helmet bedding, football helmet flag comforter, football comforter with name",
        "cluster": "personalized football helmet flag comforter",
        "intent_role": "Football helmet on American flag page; separated from close-up flag and flaming helmet variants.",
        "detail": "football helmet on American flag vintage comforter artwork with sample name and number only as examples",
    },
    65: {
        "title": "Personalized Flaming Football Helmet Comforter Set",
        "meta_title": "Personalized Flaming Football Helmet Comforter Set",
        "meta_description": "Customize a flaming football helmet comforter with required name, optional number, product type, size and bedding options.",
        "primary": "personalized flaming football helmet comforter",
        "secondary": "custom flaming helmet football bedding, football helmet comforter with name, flaming football comforter",
        "cluster": "personalized flaming football helmet comforter",
        "intent_role": "Helmet plus flaming football page; image 4/5 panel order corrected from Sang QA r2.",
        "detail": "football helmet with flaming football comforter artwork with sample MATTHEW name and 06 number only as examples",
    },
    66: {
        "title": "Personalized Football Flag Comforter Set",
        "meta_title": "Personalized Football Flag Comforter Set",
        "meta_description": "Customize a football flag comforter with required name, optional number, product type, size and bedding options.",
        "primary": "personalized football flag comforter set",
        "secondary": "custom football flag bedding, American flag football comforter, football comforter with name",
        "cluster": "personalized football flag comforter set",
        "intent_role": "Football over American flag background page; broader flag-set angle separated from product 62.",
        "detail": "football over American flag background comforter artwork with sample name and number only as examples",
    },
    67: {
        "title": "Personalized Football Player Close Up Comforter Set",
        "meta_title": "Personalized Football Player Close Up Comforter Set",
        "meta_description": "Customize a football player close-up comforter with required name, optional number, product type, size and bedding options.",
        "primary": "personalized football player close up comforter",
        "secondary": "custom football player comforter, football player bedding with name, player holding ball comforter",
        "cluster": "personalized football player close up comforter",
        "intent_role": "Close-up player holding ball page; separated from helmet and collage player pages.",
        "detail": "football player holding ball close-up comforter artwork with sample name and number only as examples",
    },
    68: {
        "title": "Personalized Football Player Helmet Comforter Set",
        "meta_title": "Personalized Football Player Helmet Comforter Set",
        "meta_description": "Customize a football player helmet comforter with required name, optional number, product type, size and bedding options.",
        "primary": "personalized football player helmet comforter",
        "secondary": "custom football player helmet bedding, football player comforter with name, helmet football bedding",
        "cluster": "personalized football player helmet comforter",
        "intent_role": "Player close-up with helmet page; separated from product 67 by helmet modifier.",
        "detail": "football player holding ball close-up comforter with football helmet artwork and sample name/number only as examples",
    },
    69: {
        "title": "Personalized Football Player Collage Comforter Set",
        "meta_title": "Personalized Football Player Collage Comforter Set",
        "meta_description": "Customize a football player collage comforter with required name, optional number, product type, size and bedding options.",
        "primary": "personalized football player collage comforter",
        "secondary": "custom football collage bedding, football player collage comforter, football comforter with name",
        "cluster": "personalized football player collage comforter",
        "intent_role": "Player collage plus helmet page; separated from single-player close-up pages.",
        "detail": "football player holding ball collage comforter with football helmet artwork and sample name/number only as examples",
    },
    70: {
        "title": "Personalized Camo Football Player Comforter Set",
        "meta_title": "Personalized Camo Football Player Comforter Set",
        "meta_description": "Customize a camo football player comforter with required name, optional number, product type, size and bedding options.",
        "primary": "personalized camo football player comforter",
        "secondary": "custom camouflage football bedding, camo football comforter with name, football player camo comforter",
        "cluster": "personalized camo football player comforter",
        "intent_role": "Camouflage flag football player page; separated by camo/player modifier.",
        "detail": "football player running on camouflage flag comforter artwork with sample name and number only as examples",
    },
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def load_qa():
    data = json.loads(QA_DATASET.read_text(encoding="utf-8"))
    products = data["QA_Products"]
    handle_by_key = {item["product_key"]: item["handle"] for item in products}
    pos_by_key = {item["product_key"]: int(item["inventory_position"]) for item in products}
    image_by_qa_key = {img["qa_image_key"]: img for img in data["QA_Images"]}
    image_fix_by_key = {}
    observation_by_key = {}
    for image in data["QA_Images"]:
        observation_by_key[(image["product_key"], text(image["media_id"]))] = image["qa_observation"]
    for issue in data["QA_Issues"]:
        if issue["field"].startswith("image_") and issue["recommended_fix"]:
            image = image_by_qa_key.get(issue["qa_image_key"])
            if image:
                image_fix_by_key[(image["product_key"], text(image["media_id"]))] = issue["recommended_fix"]
    customizer = json.loads(CUSTOMIZER_AUDIT.read_text(encoding="utf-8"))
    customizer_by_pos = {int(item["inventory_position"]): item for item in customizer}
    return {
        "data": data,
        "product_keys": [item["product_key"] for item in products],
        "handle_by_key": handle_by_key,
        "pos_by_key": pos_by_key,
        "image_fix_by_key": image_fix_by_key,
        "observation_by_key": observation_by_key,
        "customizer_by_pos": customizer_by_pos,
    }


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Body (HTML)": first.get("Body (HTML)", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    return admin


def customizer_sentence(audit):
    nodes = audit.get("personalization_nodes", [])
    parts = []
    for node in nodes:
        label = node.get("label", "text field")
        required = "required" if node.get("required") else "optional"
        limit = node.get("maxLength")
        if limit:
            parts.append(f"Live Customizer shows a {required} {label} field up to {limit} characters.")
        else:
            parts.append(f"Live Customizer shows a {required} {label} field.")
    parts.append("Names and numbers visible in mockups are treated as samples.")
    return " ".join(parts)


def compact_observation_to_alt(observation, fallback):
    value = re.sub(r"[:;].*$", "", text(observation)).strip()
    value = re.sub(r"\s+", " ", value)
    return (value or fallback)[:125]


def description(pos, admin_row, ctext):
    item = PRODUCT_UPDATES[pos]
    options = ", ".join([v for v in [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]] if v])
    return (
        f"<p>{item['detail'].capitalize()} gives this Jeminise {admin_row['Type'].lower()} a distinct football bedding focus.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork focus: {item['detail']}.</li>"
        f"<li>Current Shopify export title: {admin_row['Title']}.</li>"
        "<li>Gallery images include the main mockup plus feature, care, close-up, size or material panels where shown.</li></ul>"
        "<h3>Options and Customization</h3>"
        f"<ul><li>Available option groups from Shopify export: {options}.</li>"
        f"<li>{ctext}</li>"
        "<li>Use Customize to enter the name and optional number, then confirm product type, size, pillowcases and sheet-cover choices before checkout.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    qa = load_qa()
    admin = load_admin(set(qa["handle_by_key"].values()))
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ctext = customizer_sentence(qa["customizer_by_pos"].get(pos, {}))

        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or ws.cell(row_num, idx["product_type"]).value
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, ctext)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"Buyer intent targets {item['cluster']} for US English product search; no search-volume claim is made."
        )
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R3 after Sang re-QA r2: used products_export_1.csv (sha256:{admin_hash}) as admin baseline; "
            "rewrote generic descriptions, mapped verified name/number rules, shortened meta copy, refreshed image observations/alt, "
            "and separated football keyword clusters. Still NEEDS_REVIEW; no APPROVED/import."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        media_id = text(ws.cell(row_num, idx["media_id"]).value)
        key = (pk, media_id)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation = qa["observation_by_key"].get(key)
        if observation:
            ws.cell(row_num, idx["observed_visual_details"]).value = observation
        fallback = f"{PRODUCT_UPDATES[pos]['title']} product image"
        ws.cell(row_num, idx["alt_proposed"]).value = (
            qa["image_fix_by_key"].get(key) or compact_observation_to_alt(observation, fallback)
        )[:125]
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R3: Sang QA r2 image observation/alt applied; current admin alt and image URL matched from products_export_1.csv where possible."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "Product"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = "Sang re-QA r2 SERP evidence + products_export_1.csv admin baseline + customizer_audit.json"
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = "R3 separates football pages by visual motif and verified name/number Customizer rules."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r3"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as a specific personalized football bedding product."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm the motif, size and bedding choices, "
            "and how to enter the required name plus optional number."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = "Confirm product type, size, pillowcase or sheet-cover choices, verified name/number fields and image accuracy."
        ws.cell(row_num, idx["emotional_social_motivation"]).value = "Create a personalized football bedding gift with artwork that matches the player, flag, helmet or field design."
        ws.cell(row_num, idx["purchase_concerns"]).value = "Name and number rules, exact bedding type, included components, fabric/care claims and image accuracy."
        ws.cell(row_num, idx["source_refs"]).value = f"products_export_1.csv sha256:{admin_hash}; Sang QA r2 {QA_RUN_ID}; customizer_audit.json"
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No search-volume source supplied; Customizer-to-cart persistence still needs controlled test before approval."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = qa["product_keys"][pos - 61]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = f"Storefront/product.js from Sang QA r2; products_export_1.csv sha256:{admin_hash}; customizer_audit.json"
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={admin_row['Option1 Name']}, "
            f"{admin_row['Option2 Name']}, {admin_row['Option3 Name']}; status={admin_row['Status']}; "
            f"image_count={len(admin_row['images'])}; design={item['detail']}; customizer={customizer_sentence(qa['customizer_by_pos'].get(pos, {}))}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "R3 applied Sang QA r2 fixes; needs independent QA recheck."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R3 uses QA-observed images, live Customizer audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_007_r3_revision", "r3", "Separate 10-product revision after Sang re-QA r2; source r2 workbook was not modified."),
        ("qa_batch_007_r3_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_007_r3_scope", "inventory positions 61-70", "No products outside qa_batch_007 were revised."),
        ("qa_batch_007_r3_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_007_r3_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_007_r3_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa_r2": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_007 only; inventory positions 61-70",
        "revision": "r3",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Sang QA lại qa_batch_007_r3 before any approval or import.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_007_r3",
                "",
                "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r2: `{SOURCE.relative_to(ROOT)}`",
                f"- QA r2 của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r3: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Sửa trọng yếu: mapped required name/optional number fields, shortened meta copy, separated football keyword clusters.",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Bước tiếp theo: Sang QA lại `qa_batch_007_r3`.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
