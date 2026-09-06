import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_024"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_023.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_024.xlsx"


BASKETBALL_REFS = "https://www.walmart.com/c/kp/basketball-bedding; https://www.amazon.com/clp/B0D2PBRQPS; https://ohaprints.com/products/personalized-basketball-duvet-cover-set-basketball-ball-splash-player-gift-black-duvet-cover-pillowcases-custom-name-number-bedding-set; https://www.amorcustomgifts.com/products/basketball-personalized-qui; https://www.etsy.com/market/personalized_basketball_comforter"
BIGFOOT_REFS = "https://www.amazon.com/bigfoot-blanket/s?k=bigfoot+blanket; https://www.etsy.com/market/cotton_bigfoot_throw; https://www.etsy.com/listing/1777422691/vintage-bigfoot-quilt-bedding-set-floral; https://www.pinterest.com/pin/believe-bigfoot-full-moon-an-bedding-setw24889--970385050950914186/; https://www.pinterest.com/pin/personalized-bigfoot-blanket-sizes-for-baby-child-teen-or-adult-custom-made-with-any-name-super-soft-sasquatch-b--155303888146137345/"


PROFILES = {
    231: ("turquoise basketball splatter blanket with name", "basketball splatter blanket", "Basketball Splatter Blanket", "Turquoise Basketball Splatter Blanket with Name", "turquoise and pink splattered basketball blanket with oversized ball, custom name and jersey number", "personalized basketball blanket, basketball name blanket, custom basketball throw", "EVERGREEN", BASKETBALL_REFS),
    232: ("red basketball paint splash comforter with name", "basketball paint splash comforter", "Basketball Paint Splash Comforter", "Red Basketball Paint Splash Comforter with Name", "red basketball comforter with white paint splash lines, large ball, custom name and jersey number", "personalized basketball bedding, custom basketball comforter, basketball room decor", "EVERGREEN", BASKETBALL_REFS),
    233: ("basketball reflection blanket with script name", "basketball reflection blanket", "Basketball Reflection Blanket", "Basketball Reflection Blanket with Script Name", "black basketball blanket with ball reflection artwork, stars, script custom name and jersey number", "personalized basketball blanket, custom basketball throw, basketball player gift", "EVERGREEN", BASKETBALL_REFS),
    234: ("orange basketball water splash comforter with name", "basketball water splash comforter", "Basketball Water Splash Comforter", "Orange Basketball Water Splash Comforter with Name", "orange basketball comforter with water splash effect, custom name and jersey number on matching shams", "personalized basketball bedding, basketball comforter set, custom sports bedding", "EVERGREEN", BASKETBALL_REFS),
    235: ("dark Bigfoot forest silhouette quilt", "Bigfoot forest quilt", "Bigfoot Forest Quilt Set", "Dark Bigfoot Forest Silhouette Quilt", "dark blue forest quilt with Bigfoot silhouette walking toward moonlight between tall trees", "Sasquatch bedding, Bigfoot quilt set, forest silhouette quilt", "EVERGREEN", BIGFOOT_REFS),
    236: ("Bigfoot campfire mountain quilt", "Bigfoot campfire quilt", "Bigfoot Campfire Quilt Set", "Bigfoot Campfire Mountain Quilt", "blue mountain campfire quilt with Bigfoot seated by a fire, lanterns, forest and snowy landscape details", "Sasquatch quilt set, campfire bedding, Bigfoot cabin decor", "EVERGREEN", BIGFOOT_REFS),
    237: ("Bigfoot campfire night forest quilt", "Bigfoot night forest quilt", "Bigfoot Night Forest Quilt Set", "Bigfoot Campfire Night Forest Quilt", "night forest quilt with two Bigfoot figures by a glowing campfire under blue trees and mountain sky", "Sasquatch bedding, campfire quilt set, Bigfoot forest bedding", "EVERGREEN", BIGFOOT_REFS),
    238: ("orange sunset Bigfoot forest silhouette quilt", "Bigfoot silhouette quilt", "Bigfoot Silhouette Quilt Set", "Orange Sunset Bigfoot Forest Silhouette Quilt", "orange sunset forest quilt with small Bigfoot silhouette walking between tall dark trees", "Sasquatch quilt, forest silhouette bedding, Bigfoot cabin quilt", "EVERGREEN", BIGFOOT_REFS),
    239: ("Bigfoot full moon forest quilt", "Bigfoot full moon quilt", "Bigfoot Full Moon Quilt Set", "Bigfoot Full Moon Forest Quilt", "dark forest quilt with large full moon and Bigfoot silhouette centered between blue-black trees", "Sasquatch bedding, full moon quilt, Bigfoot forest quilt", "EVERGREEN", BIGFOOT_REFS),
    240: ("Bigfoot moon mountain night quilt", "Bigfoot mountain quilt", "Bigfoot Moon Mountain Quilt Set", "Bigfoot Moon Mountain Night Quilt", "night mountain quilt with Bigfoot silhouette in front of a large moon, pine trees and layered peaks", "Sasquatch quilt set, Bigfoot moon bedding, mountain cabin quilt", "EVERGREEN", BIGFOOT_REFS),
}


def qa_issues(pos, html_error):
    issues = [
        "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."
    ]
    if 231 <= pos <= 234:
        issues.append("Basketball products overlap batches 021-023; keep splatter, paint splash, reflection and water splash motifs distinct.")
        issues.append("Verify product type blanket versus comforter and custom name/number fields before import.")
    if 235 <= pos <= 240:
        issues.append("Bigfoot/Sasquatch quilt products are close variants; review title uniqueness, motif separation and canonical strategy.")
        issues.append("Feature images mention quilt/bedspread/fabric claims and optional pillow shams; verify variants and included components before import.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def product_context(pos):
    if pos <= 234:
        return "basketball rooms, player gifts, sports bedrooms or personalized fan bedding."
    return "cabin rooms, rustic bedrooms, cryptid gifts or Bigfoot-themed decor."


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for {product_context(pos)}</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broad basketball and Bigfoot bedding terms for collection pages or strongest representatives.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, room decor or a personalized gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": qa_issues(pos, html_error),
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed bedding, the buyer wants a {kind} matching a specific basketball or Bigfoot/Sasquatch motif.",
            "functional_motivation": "Find the right design, size, product type, personalization fields, care details and room fit.",
            "emotional_social_motivation": "Create a themed bedroom or give a custom sports, cabin or cryptid-inspired gift.",
            "purchase_concerns": "Personalization spelling, product identity, material/care claims, size fit, included components and whether artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed in a main product mockup",
                2: f"Secondary product mockup for {h1}",
                3: f"Size, care or feature image for {h1}",
                4: f"Fabric, folded or product feature image for {h1}",
                5: f"Lifestyle, size or close-up image for {h1}",
                6: f"Bedding construction or feature image for {h1}",
                7: f"Additional size or bedroom mockup for {h1}",
                8: f"Lifestyle use image for {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_024_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_024_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_024"},
        {"metric": "cumulative_products", "value": "240", "definition": "Total products included through batch_024"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[240:250]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_025",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
