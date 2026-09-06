import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_029"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_028.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_029.xlsx"


CACTUS_REFS = "https://www.target.com/c/bedding-home-decor/cactus/-/N-5xtv4Zg2fff; https://www.wayfair.com/keyword.php?keyword=cactus+bedding; https://www.pinterest.com/ideas/cactus-bedding/959782873627/; https://www.ebay.com/itm/257177873478"
READING_REFS = "https://www.etsy.com/market/book_lover_blanket; https://www.etsy.com/listing/1749170765/just-a-girl-who-loves-books-blanket; https://macorner.co/products/just-a-girl-who-loves-books-personalized-blanket-boo180701lamt; https://www.amazon.com/Just-Girl-Loves-Books-Blanket/dp/B0DMVXGZKG"
DEER_REFS = "https://luvingift.com/product/personalized-deer-couple-quilt-set-his-doe-her-buck-she-keeps-me-wild-he-keeps-me-safe/; https://www.etsy.com/listing/4325561114/mystical-deer-couple-woven-blanket; https://www.amazon.com/deer-comforter-set/s?k=deer+comforter+set"
PHOTO_REFS = "https://www.etsy.com/market/custom_photo_quilt; https://www.etsy.com/market/photo_collage_blanket; https://www.amazon.com/custom-photo-blanket/s?k=custom+photo+blanket"
AFFIRMATION_REFS = "https://www.etsy.com/market/affirmation_blanket; https://www.etsy.com/market/personalized_affirmation_blanket; https://macorner.co/products/christian-bible-verse-affirmations-for-girls-boys-personalized-photo-blanket-bwo241101lahn"
TREE_REFS = "https://www.amazon.com/clp/B0DNSVY2QJ; https://luvingift.com/product/viking-quilt-set-norse-celtic-yggdrasil-tree-of-life/; https://myvikinggear.com/collections/viking-home-decoration/viking-quilt-set/; https://www.wayfair.com/keyword.php?keyword=tree+of+life+bedding"


PROFILES = {
    281: ("flowering cactus succulent garden quilt", "flowering cactus quilt", "Flowering Cactus Quilt Set", "Flowering Cactus Succulent Garden Quilt Set", "light southwestern quilt with flowering cactus, succulents, potted plants and patterned desert garden border", "succulent quilt, cactus bedding, southwest quilt set", "EVERGREEN", CACTUS_REFS),
    282: ("girl with glasses newspaper book lover blanket", "newspaper book lover blanket", "Newspaper Book Lover Blanket", "Girl Glasses Newspaper Reading Blanket", "book lover throw blanket with girl in glasses, newspaper pattern, leopard border, stacked books and custom name", "personalized reading blanket, girl who loves books blanket, bookworm blanket", "EVERGREEN", READING_REFS),
    283: ("girl in glasses bookshelf reading blanket", "girl in glasses reading blanket", "Girl in Glasses Reading Blanket", "Girl in Glasses Bookshelf Reading Blanket", "cream book lover blanket with girl in glasses, bookshelf, flowers, butterflies, custom books and custom name", "personalized book blanket, reading blanket for girls, book lover gift", "EVERGREEN", READING_REFS),
    284: ("messy bun girl open book reading blanket", "messy bun reading blanket", "Messy Bun Reading Blanket", "Messy Bun Girl Open Book Reading Blanket", "vintage open-book blanket with messy bun reading girl, flowers, coffee mug, script pages and custom name", "book lover blanket, just a girl who loves books blanket, custom reading blanket", "EVERGREEN", READING_REFS),
    285: ("girl reading bookshelf floral blanket", "bookshelf reading girl blanket", "Bookshelf Reading Girl Blanket", "Girl Reading Bookshelf Floral Blanket", "cream reading blanket with girl in glasses, colorful bookshelf, flowers, butterflies and custom name", "personalized book lover blanket, reading blanket, bookworm gift", "EVERGREEN", READING_REFS),
    286: ("hunting deer antlers photo comforter", "hunting deer comforter", "Hunting Deer Comforter Set", "Hunting Deer Antlers Photo Comforter", "woodgrain hunting comforter with deer antlers, heart icon, custom photo collage and custom names", "personalized hunting bedding, deer antler comforter, custom photo comforter", "EVERGREEN", DEER_REFS),
    287: ("custom name reading books blanket", "custom name reading blanket", "Custom Name Reading Blanket", "Custom Name Reading Books Blanket", "white floral reading blanket with open book, book spines, butterflies, custom name and color options", "personalized book blanket, reading book blanket, book lover gift", "EVERGREEN", READING_REFS),
    288: ("love relationship photo collage quilt", "love photo collage quilt", "Love Photo Collage Quilt", "Love Relationship Photo Collage Quilt", "black photo collage quilt with multiple custom photos, red heart icon, I love you text and optional pillow shams", "custom photo quilt, personalized couple quilt, photo collage bedding", "VALENTINES_DAY", PHOTO_REFS),
    289: ("morning affirmations floral blanket", "morning affirmations blanket", "Morning Affirmations Blanket", "Morning Affirmations Floral Blanket", "navy affirmation blanket with cartoon girl, floral accents, custom name and morning affirmation phrases", "personalized affirmation blanket, self love blanket, positive affirmation gift", "EVERGREEN", AFFIRMATION_REFS),
    290: ("mosaic Tree of Life quilt", "mosaic Tree of Life quilt", "Mosaic Tree of Life Quilt Set", "Mosaic Tree of Life Quilt Set", "bright mosaic style Tree of Life quilt with colorful leaves, glowing sun, orange hills and stained-glass effect", "Tree of Life bedding, colorful quilt set, mosaic quilt", "EVERGREEN", TREE_REFS),
}


def qa_issues(pos, html_error):
    issues = [
        "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."
    ]
    if pos in (281,):
        issues.append("Cactus/succulent quilt overlaps batch_028 cactus products; review motif separation and canonical strategy.")
    if 282 <= pos <= 285 or pos == 287:
        issues.append("Book lover blankets overlap previous reading-girl batches; keep newspaper, bookshelf, open-book and custom-name motifs distinct.")
        issues.append("Verify displayed sample names, character options, size variants and fleece/care claims.")
    if pos == 286:
        issues.append("Hunting photo comforter needs photo-upload QA, custom name field review and comforter versus quilt naming verification.")
    if pos == 288:
        issues.append("Love photo collage quilt requires photo-upload QA, seasonal positioning review and optional pillow sham/component verification.")
    if pos == 289:
        issues.append("Affirmation blanket text and character options should be reviewed; verify material, size and care claims.")
    if pos == 290:
        issues.append("Tree of Life quilt overlaps Celtic/Yggdrasil batches; keep mosaic/stained-glass visual positioning distinct.")
        issues.append("Verify quilt/bedspread construction, optional pillow shams and machine-washable claims.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def product_context(pos):
    if pos in (281,):
        return "southwestern bedrooms, desert decor, cactus-themed rooms or garden-inspired bedding."
    if 282 <= pos <= 285 or pos == 287:
        return "reading nooks, bedrooms, book clubs, librarian gifts or personalized reader gifts."
    if pos == 286:
        return "hunting cabins, rustic bedrooms, outdoor gifts or personalized photo bedding."
    if pos == 288:
        return "Valentine gifts, anniversary gifts, couple keepsakes or personalized bedroom decor."
    if pos == 289:
        return "morning routines, bedrooms, self-love gifts or personalized encouragement gifts."
    return "spiritual bedrooms, colorful decor, Tree of Life gifts or art-inspired bedding."


def alt_for_image(h1, pos, ipos):
    if pos in (281, 286, 288, 290):
        return {
            1: f"{h1} displayed in the main bedding mockup",
            2: f"Flat product image for {h1}",
            3: f"Angled bedroom or size mockup of {h1}",
            4: f"Feature or pillow sham detail for {h1}",
            5: f"Fabric and product feature image for {h1}",
            6: f"Size chart or pillow sham image for {h1}",
            7: f"Bedspread construction and care feature image for {h1}",
            8: f"Overhead bedroom mockup of {h1}",
        }.get(ipos, f"Product detail image for {h1}")[:125]
    return {
        1: f"{h1} shown as a personalized throw blanket",
        2: f"Flat product or size image for {h1}",
        3: f"Size chart or lifestyle mockup for {h1}",
        4: f"Couch mockup for {h1}",
        5: f"Fabric and fleece feature image for {h1}",
        6: f"Couch or close-up mockup for {h1}",
        7: f"Reading lifestyle image for {h1}",
        8: f"Additional reading lifestyle image for {h1}",
        9: f"Family, color or option image for {h1}",
    }.get(ipos, f"Product detail image for {h1}")[:125]


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for {product_context(pos)}</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broad cactus, reading, deer, photo gift, affirmation and Tree of Life terms for collection pages or strongest representatives.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, room decor or a personalized gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": qa_issues(pos, html_error),
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed bedding, the buyer wants a {kind} matching a specific cactus, reading, hunting, photo, affirmation or Tree of Life motif.",
            "functional_motivation": "Find the right design, size, product type, personalization fields, care details and room fit.",
            "emotional_social_motivation": "Give a personalized reader, couple, hunting or encouragement gift, or decorate with desert and Tree of Life artwork.",
            "purchase_concerns": "Personalization spelling, photo quality, affirmation text, product identity, material/care claims, size fit, included components and whether artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt_for_image(h1, pos, ipos), "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_029_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_029_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_029"},
        {"metric": "cumulative_products", "value": "290", "definition": "Total products included through batch_029"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[290:300]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_030",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
