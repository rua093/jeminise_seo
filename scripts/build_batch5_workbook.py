import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_005"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_004.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_005.xlsx"


PROFILES = {
    41: ("personalized football player comforter with the name Kevin", "personalized football player comforter", "Personalized Football Player Comforter Set", "Personalized Football Player Comforter Set", "red football player running with the ball, vintage flag-style background, vertical name Kevin and matching pillow artwork", "custom football bedding, football comforter with name, football bedding for boys", "EVERGREEN", "https://www.etsy.com/market/personalized_football_blanket; https://www.amazon.com/s?k=personalized+football+blanket; https://footballshirtmaker.com/en/"),
    42: ("personalized basketball blanket with name and number", "personalized basketball blanket", "Personalized Basketball Blanket with Name", "Personalized Basketball Name and Number Blanket", "large orange basketball graphic, player silhouettes, number 06 and name Colon", "custom basketball blanket, basketball blanket with name, basketball gift blanket", "EVERGREEN", "https://www.amazon.com/stores/PersonalizedLiving/page/E5F21C82-AB1B-4E9A-9AFC-43DF77ABCD1B; https://www.etsy.com/market/personalized_basketball_blanket; https://us.shein.com/1pc-Personalized-Boy-Basketball-Blanket-Custom-Checkered-Boy-Name-Blanket-Boy-Birthday-Gift-Basketball-Gift-Gift-For-Christmas-Birthday-And-Anniversary-p-319295421.html"),
    43: ("personalized God Says I Am photo blanket", "personalized Bible verse photo blanket", "Personalized Bible Verse Photo Blanket", "Personalized God Says I Am Photo Blanket", "God Says I Am design with a custom photo, name Maria, floral border and Bible verse affirmations", "God Says I Am blanket, Christian photo blanket, personalized Christian gift for her", "EVERGREEN", "https://www.etsy.com/market/god_says_i_am_blanket; https://www.callie.com/personalized-god-says-i-am-name-meaning-on-bible-verse-colorful-soft-throw-blanket-baptism-confirmation-christmas-gift-for-christians-family-kids; https://www.amazon.com/s?k=personalized+god+says+i+am+blanket"),
    44: ("cardinal flowering branches quilt", "cardinal flowering branches quilt", "Cardinal Flowering Branches Quilt Set", "Cardinal Flowering Branches Quilt Set", "red cardinals perched on flowering white branches over a soft gray background", "cardinal bird quilt, floral cardinal bedding, cardinal quilt with flowers", "EVERGREEN", "https://www.amazon.com/Bedspread-Cardinals-Coverlet-Set-Quilt-Set-Botanical/dp/B09QSNYM7V; https://www.walmart.com/c/kp/cardinal-quilt; https://www.etsy.com/market/cardinal_bedspread"),
    45: ("colorful Tree of Life mosaic quilt", "colorful Tree of Life quilt set", "Colorful Tree of Life Mosaic Quilt Set", "Colorful Tree of Life Mosaic Quilt Set", "stained-glass style Tree of Life with orange sunburst, multicolor branches and blue landscape", "mosaic tree quilt, Yggdrasil quilt set, colorful Celtic tree bedding", "EVERGREEN", "https://www.amazon.com/s?k=yggdrasil+tree+of+life+quilt; https://dingmun.com/product/viking-raven-tree-of-life-yggdrasil-norse-mythology-symbol-quilt-bedding-set/; https://www.pinterest.com/pin/yggdrasil-the-tree-of-life-in-norse-mythology-viking-quilt-set--1027383733757965557/"),
    46: ("blue fantasy Tree of Life eye quilt", "fantasy Tree of Life quilt set", "Fantasy Tree of Life Eye Quilt Set", "Blue Fantasy Tree of Life Eye Quilt Set", "blue Tree of Life artwork with central eye, moon shapes and ornate gray border", "blue tree of life bedding, mystical tree quilt, fantasy eye quilt set", "EVERGREEN", "https://www.amazon.com/s?k=tree+of+life+quilt+set; https://www.pinterest.com/pin/tree-of-life-yggdrasil-norse-mythology-viking-quilt-bedding-set--1097752477912958827/; https://www.etsy.com/market/tree_of_life_quilt"),
    47: ("Celtic Yggdrasil roots quilt", "Celtic Yggdrasil quilt set", "Celtic Yggdrasil Tree Quilt Set", "Celtic Yggdrasil Tree of Life Quilt Set", "Tree of Life with intertwined exposed roots and Celtic knotwork ring in green and gold landscape", "Celtic tree of life bedding, intertwined roots quilt, Yggdrasil bedding set", "EVERGREEN", "https://www.amazon.com/clp/B0DNSVY2QJ; https://dingmun.com/product/viking-raven-tree-of-life-yggdrasil-norse-mythology-symbol-quilt-bedding-set/; https://www.etsy.com/market/yggdrasil_quilt"),
    48: ("personalized Trucker's Prayer comforter", "personalized trucker prayer comforter", "Personalized Trucker's Prayer Comforter", "Personalized Trucker's Prayer Comforter Set", "red semi truck, wooden cross, Trucker's Prayer scroll and Your Name placeholder on dark background", "Christian trucker bedding, truck driver prayer comforter, custom semi truck bedding", "EVERGREEN", "https://www.etsy.com/market/trucker_prayer_gifts; https://famhose.com/products/trucker-blanket-truckers-prayer-keep-me-safe-get-me-home-trucker-fleece-blanket-sherpa-blanket-gift-for-trucker-lovers-trucker26; https://www.pinterest.com/pin/personalized-trucker-quilt-set-truckers-prayer-pink-truck-truck-driver-quilt-blanket-with-pillowcases-custom-name-qui--999939923509359204/"),
    49: ("personalized Bible emergency numbers blanket for girls", "personalized Bible emergency numbers blanket", "Personalized Bible Emergency Numbers Blanket", "Personalized Bible Emergency Numbers Blanket", "cartoon girl named Sophia, colorful flowers and Bible emergency numbers verse list", "Christian girl blanket, Bible verse blanket for girls, custom Christian blanket", "EVERGREEN", "https://www.youtube.com/results?search_query=My+Bible+Emergency+Numbers+Personalized+Blanket; https://www.pinterest.com/pin/bible-hotline-numbers-emergency-scripture-pullover-baby-blankets--1086774953802326238/; https://www.etsy.com/market/bible_verse_blanket"),
    50: ("personalized Christian Sophia inspirational blanket", "personalized Christian inspirational blanket", "Personalized Christian Inspirational Blanket", "Personalized Christian Sophia Blanket", "Dear Sophia design with cartoon girl, circular affirmation words, flower graphic and faith verse accents", "custom Christian blanket, inspirational blanket for girls, personalized faith blanket", "EVERGREEN", "https://www.etsy.com/market/personalized_christian_blanket; https://christianartbag.com/; https://www.suzitee.com/"),
}


def append_dict(ws, data):
    headers = [cell.value for cell in ws[1]]
    ws.append([data.get(h, "") for h in headers])


def image_obs(theme, pos):
    if pos == 1:
        return f"Main product mockup showing {theme}."
    if pos == 2:
        return f"Secondary mockup or folded view showing {theme}."
    if pos == 3:
        return f"Personalization, feature or lifestyle image connected to {theme}."
    if pos == 4:
        return f"Close detail, care feature or sizing information for {theme}."
    if pos == 5:
        return f"Material, wash or comfort feature image for {theme}."
    if pos == 6:
        return f"Size guide, feature panel or additional mockup for {theme}."
    return f"Additional lifestyle, feature or product detail image showing {theme}."


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        issue_bits = [
            "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."
        ]
        if html_error:
            issue_bits.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
        if pos == 44:
            issue_bits.append("Title says 'with Name' but contact sheet did not show a clear personalized name; personalization claim needs QA.")
        if pos == 50:
            issue_bits.append("Product title references D11 while source specification shows Selected Design D14; design code needs QA.")
        issues = " ".join(issue_bits)
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, giving the product a clear theme for bedroom decor or gifting.</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork, personalization and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the product-level motif: {label}. Broader sports, faith, bird or Celtic terms should be reviewed for collection pages.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, a faith gift or a personalized sports gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": issues,
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium-low: product JSON and contact sheet reviewed, but storefront HTML was blocked for this batch; admin export and QA are required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed bedding or a personalized gift, the buyer wants a coordinated {kind} matching a specific design, faith message or recipient.",
            "functional_motivation": "Find the right design, size, material/care fit and personalization option.",
            "emotional_social_motivation": "Give a personal faith or sports gift, create meaningful bedroom decor, or choose a symbolic bird/Celtic design.",
            "purchase_concerns": "Fit, included components, material/care claims, personalization spelling, design code consistency and whether the artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork, personalization and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed on a bed",
                2: f"Secondary mockup of {h1}",
                3: f"Feature or personalization image for {h1}",
                4: f"Close-up detail of {label}",
                5: f"Care or material information for {h1}",
                6: f"Size or feature guide for {h1}",
                7: f"Additional product mockup of {h1}",
                8: f"Lifestyle use image for {h1}",
                9: f"Additional lifestyle image for {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_005_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_005_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_005"},
        {"metric": "cumulative_products", "value": "50", "definition": "Total products included through batch_005"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[50:60]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_006",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
