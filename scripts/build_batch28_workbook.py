import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_028"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_027.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_028.xlsx"


DEER_REFS = "https://luvingift.com/product/personalized-deer-couple-quilt-set-his-doe-her-buck-she-keeps-me-wild-he-keeps-me-safe/; https://www.etsy.com/listing/4325561114/mystical-deer-couple-woven-blanket; https://kr.pinterest.com/pin/874753927613759157/; https://jp.pinterest.com/pin/1027383733731175391/"
CACTUS_REFS = "https://www.target.com/c/bedding-home-decor/cactus/-/N-5xtv4Zg2fff; https://www.wayfair.com/keyword.php?keyword=cactus+bedding; https://www.pinterest.com/ideas/cactus-bedding/959782873627/; https://www.ebay.com/itm/257177873478"
CHRISTIAN_REFS = "https://www.etsy.com/listing/1880095899/personalized-god-says-you-are-christian; https://macorner.co/products/christian-bible-verse-affirmations-for-girls-boys-personalized-photo-blanket-bwo241101lahn; https://www.etsy.com/market/christian_blanket"
READING_REFS = "https://www.etsy.com/listing/1749170765/just-a-girl-who-loves-books-blanket; https://macorner.co/products/just-a-girl-who-loves-books-personalized-blanket-boo180701lamt; https://www.etsy.com/market/book_lover_blanket; https://stock.adobe.com/images/girl-holding-a-book-reading-laying-down-on-a-blanket-line-art-sketch-vector/138370951"


PROFILES = {
    271: ("deer couple heart frame comforter", "deer couple comforter", "Deer Couple Comforter Set", "Deer Couple Heart Frame Comforter", "rustic buck comforter with camo background, heart frame, love quote and custom couple names", "buck and doe bedding, personalized deer bedding, hunting couple comforter", "EVERGREEN", DEER_REFS),
    272: ("deer silhouettes facing each other comforter", "buck and doe comforter", "Buck and Doe Comforter Set", "Deer Silhouettes Facing Each Other Comforter", "white and blue forest comforter with buck and doe silhouettes facing each other, small heart and custom names", "deer couple bedding, hunting comforter set, personalized buck doe bedding", "EVERGREEN", DEER_REFS),
    273: ("deer silhouettes heart outline comforter", "deer heart comforter", "Deer Heart Comforter Set", "Deer Silhouettes Heart Outline Comforter", "brown camo comforter with buck and doe silhouettes forming a heart outline and custom couple names", "buck doe comforter, hunting couple bedding, personalized deer comforter", "EVERGREEN", DEER_REFS),
    274: ("desert cactus sunrise quilt", "desert cactus quilt", "Desert Cactus Quilt Set", "Desert Cactus Sunrise Quilt Set", "southwestern quilt with orange sunrise, layered desert hills, cactus plants and teal sky", "cactus bedding, desert quilt set, southwest quilt", "EVERGREEN", CACTUS_REFS),
    275: ("pink desert cactus floral quilt", "cactus floral quilt", "Cactus Floral Quilt Set", "Desert Sunset Cactus Floral Quilt Set", "bright pink desert quilt with flowering saguaro cactus, sunset stripes, birds and southwest landscape", "desert cactus bedding, floral cactus quilt, southwest bedding", "EVERGREEN", CACTUS_REFS),
    276: ("floral cross Christian affirmation blanket", "floral cross Christian blanket", "Floral Cross Christian Blanket", "Floral Cross Christian Affirmation Blanket", "white Christian blanket with floral cross, affirmation words, Bible references and custom name", "Christian affirmation blanket, Bible verse blanket, personalized faith blanket", "EVERGREEN", CHRISTIAN_REFS),
    277: ("floral cross God Says blanket", "God Says I Am blanket", "God Says I Am Blanket", "Floral Cross God Says Throw Blanket", "white Christian blanket with floral cross, custom photo, God Says I Am text and affirmation verses", "personalized Christian blanket, Bible verse photo blanket, faith gift blanket", "EVERGREEN", CHRISTIAN_REFS),
    278: ("floral cross personalized text blanket", "personalized floral cross blanket", "Personalized Floral Cross Blanket", "Floral Cross Personalized Text Blanket", "white personalized blanket with floral cross, I am chosen text, affirmation words and Bible references", "Christian cross blanket, personalized faith blanket, Bible verse blanket", "EVERGREEN", CHRISTIAN_REFS),
    279: ("floral line art girl reading blanket", "line art book lover blanket", "Line Art Book Lover Blanket", "Floral Line Art Girl Reading Blanket", "minimal beige book lover blanket with line-art girl reading, flowers, butterflies, custom name and quote text", "just a girl who loves books blanket, reading blanket, personalized book lover blanket", "EVERGREEN", READING_REFS),
    280: ("flowering cactus succulent garden quilt", "succulent cactus quilt", "Succulent Cactus Quilt Set", "Flowering Cactus Succulent Garden Quilt Set", "garden quilt with blooming cactus, succulents, colorful flowers, pots and dark background", "cactus bedding, succulent quilt, floral cactus bedding", "EVERGREEN", CACTUS_REFS),
}


def qa_issues(pos, html_error):
    issues = [
        "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."
    ]
    if 271 <= pos <= 273:
        issues.append("Deer couple comforters are close variants; keep heart frame, facing silhouettes and heart outline motifs distinct.")
        issues.append("Verify personalization fields, comforter versus quilt naming, optional pillow shams and hunting/couple wording before import.")
    if pos in (274, 275, 280):
        issues.append("Cactus/desert quilt products overlap RV/desert products from batch_027; review canonical and motif separation.")
        issues.append("Verify quilt/bedspread construction, machine-washable claims and optional pillow shams.")
    if 276 <= pos <= 278:
        issues.append("Christian floral cross products need verse text, custom name/photo fields and faith wording reviewed before import.")
        issues.append("Feature images mention fleece, washing, stitching and softness claims; confirm exact material, GSM and care details.")
    if pos == 279:
        issues.append("Book lover blanket overlaps previous reading-girl products; keep line-art style and quote text distinct.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def product_context(pos):
    if 271 <= pos <= 273:
        return "hunting cabins, rustic bedrooms, anniversary gifts or personalized couple bedding."
    if pos in (274, 275, 280):
        return "southwestern bedrooms, desert decor, cactus-themed rooms or warm-weather gift bedding."
    if 276 <= pos <= 278:
        return "faith gifts, prayer spaces, bedrooms, church gifts or personalized encouragement gifts."
    return "reading nooks, bedrooms, book clubs, librarian gifts or personalized reader gifts."


def alt_for_image(h1, pos, ipos):
    if pos <= 275 or pos == 280:
        return {
            1: f"{h1} displayed on a bed in the main bedding mockup",
            2: f"Flat product image for {h1}",
            3: f"Angled bedroom mockup of {h1}",
            4: f"Pillow sham and artwork detail for {h1}",
            5: f"Fabric and product feature image for {h1}",
            6: f"Size chart for {h1}",
            7: f"Bedspread construction and care feature image for {h1}",
            8: f"Overhead bedroom mockup of {h1}",
        }.get(ipos, f"Product detail image for {h1}")[:125]
    return {
        1: f"{h1} shown as a personalized throw blanket",
        2: f"Flat product or fleece feature image for {h1}",
        3: f"Lifestyle or reading mockup for {h1}",
        4: f"Size guide or couch mockup for {h1}",
        5: f"Fabric and fleece feature image for {h1}",
        6: f"Couch or close-up mockup for {h1}",
        7: f"Reading lifestyle or feature image for {h1}",
        8: f"Additional couch lifestyle image for {h1}",
    }.get(ipos, f"Product detail image for {h1}")[:125]


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for {product_context(pos)}</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broad deer, cactus, Christian and book lover bedding terms for collection pages or strongest representatives.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, room decor or a personalized gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": qa_issues(pos, html_error),
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed bedding, the buyer wants a {kind} matching a specific deer couple, cactus, Christian or reading motif.",
            "functional_motivation": "Find the right design, size, product type, personalization fields, care details and room fit.",
            "emotional_social_motivation": "Give a personalized couple, faith or reader gift, or decorate with rustic, desert or bookish artwork.",
            "purchase_concerns": "Personalization spelling, verse accuracy, product identity, material/care claims, size fit, included components and whether artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt_for_image(h1, pos, ipos), "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_028_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_028_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_028"},
        {"metric": "cumulative_products", "value": "280", "definition": "Total products included through batch_028"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[280:290]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_029",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
