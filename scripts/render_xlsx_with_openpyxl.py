from __future__ import annotations

import argparse
import math
import textwrap
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


def font(size: int, bold: bool = False):
    name = "arialbd.ttf" if bold else "arial.ttf"
    path = Path("C:/Windows/Fonts") / name
    return ImageFont.truetype(str(path), size) if path.exists() else ImageFont.load_default()


def display_value(cell):
    value = cell.value
    if value is None:
        return ""
    return str(value)


def render_sheet(ws, output: Path):
    scale = 5.3
    widths = []
    for c in range(1, ws.max_column + 1):
        raw = ws.column_dimensions[ws.cell(1, c).column_letter].width or 12
        widths.append(max(70, min(235, int(raw * scale))))
    row_heights = [38]
    for r in range(2, ws.max_row + 1):
        configured = ws.row_dimensions[r].height or 42
        row_heights.append(max(34, min(108, int(configured))))
    total_w = sum(widths) + 2
    total_h = sum(row_heights) + 2
    canvas = Image.new("RGB", (total_w, total_h), "white")
    draw = ImageDraw.Draw(canvas)
    body, head = font(12), font(12, True)
    y = 1
    for r in range(1, ws.max_row + 1):
        h = row_heights[r - 1]
        x = 1
        for c in range(1, ws.max_column + 1):
            w = widths[c - 1]
            fill = "#17365D" if r == 1 else ("#F8FAFC" if r % 2 == 0 else "#FFFFFF")
            draw.rectangle((x, y, x + w, y + h), fill=fill, outline="#CBD5E1")
            value = display_value(ws.cell(r, c))
            chars = max(8, int(w / 7.0))
            lines = []
            for paragraph in value.splitlines() or [""]:
                lines.extend(textwrap.wrap(paragraph, width=chars, break_long_words=True, break_on_hyphens=False) or [""])
            max_lines = max(1, math.floor((h - 8) / 15))
            if len(lines) > max_lines:
                lines = lines[:max_lines]
                lines[-1] = (lines[-1][:-1] + "…") if lines[-1] else "…"
            color = "#FFFFFF" if r == 1 else ("#0563C1" if ws.cell(r, c).hyperlink else "#1F2937")
            draw.multiline_text((x + 4, y + 4), "\n".join(lines), font=head if r == 1 else body, fill=color, spacing=2)
            x += w
        y += h
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("output_dir")
    args = parser.parse_args()
    wb = load_workbook(args.workbook, data_only=False, read_only=False)
    out = Path(args.output_dir)
    for ws in wb.worksheets:
        render_sheet(ws, out / f"{ws.title}.png")
        print(f"{ws.title}|{ws.max_row - 1}|{ws.max_column}|{out / (ws.title + '.png')}")


if __name__ == "__main__":
    main()
