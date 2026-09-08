import csv
import hashlib
import html
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_023_r3"
QA_RUN_ID = "20260908_143500"
SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_023_r2" / "SEO_Product_Optimization_qa_batch_023_r2.xlsx"
QA_REPORT = ROOT / "resutls" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "SEO_QA_qa_batch_023_r2.md"
SUMMARY_PATH = ROOT / "seo_runs" / SHOP / RUN_ID / "batch_023_evidence_summary.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_023_r3.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_023_r3.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    221: ("Custom Basketball Hoop Blanket", "custom basketball hoop blanket", "basketball hoop blanket, personalized basketball blanket, basketball player gift blanket", "dark blanket artwork with a close-up glowing basketball hoop, white net, sample name and jersey number", "custom basketball hoop blanket with optional name number"),
    222: ("Custom Basketball Shoes Blanket", "custom basketball shoes blanket", "basketball shoes blanket, personalized basketball blanket, basketball gift blanket", "dark blanket artwork with a basketball beside worn athletic shoes, script sample name and jersey number", "custom basketball shoes blanket with optional name number"),
    223: ("Custom Court Lines Basketball Comforter", "custom court lines basketball comforter", "basketball court comforter, personalized basketball bedding, custom basketball comforter", "black comforter with gold court lines, basketball graphics, cracked texture, sample name and jersey number", "custom name number basketball court lines comforter"),
    224: ("Custom Basketball Court Perspective Comforter", "custom basketball court perspective comforter", "basketball court perspective comforter, personalized basketball bedding, custom sports comforter", "black and red court-perspective comforter with a large basketball, sample name and jersey number", "custom name number basketball court perspective comforter"),
    225: ("Custom Hardwood Court Basketball Comforter", "custom hardwood court basketball comforter", "hardwood basketball comforter, personalized basketball bedding, custom court bedding", "blue and gold hardwood-court comforter with a large basketball, sample name and jersey number", "custom name number hardwood basketball court comforter"),
    226: ("Custom Neon Basketball Player Blanket", "custom neon basketball player blanket", "basketball player blanket, personalized basketball blanket, custom sports blanket", "dark blanket with neon-style dunking player silhouette, hoop, sample name and jersey number", "custom neon basketball player blanket with optional name number"),
    227: ("Custom Flaming Basketball Player Blanket", "custom flaming basketball player blanket", "basketball player blanket, custom basketball throw blanket, personalized sports blanket", "red and navy blanket with a front-facing dribbling player, flame effect, sample name and jersey number", "custom flaming basketball player blanket with optional name number"),
    228: ("Custom Dribbling Silhouette Basketball Blanket", "custom dribbling silhouette basketball blanket", "dribbling basketball blanket, personalized basketball player blanket, basketball gift blanket", "gray and orange blanket with a black dribbling player silhouette, oversized ball, sample name and jersey number", "custom dribbling basketball silhouette blanket with optional name number"),
    229: ("Custom Basketball Collage Comforter", "custom basketball collage comforter", "basketball collage comforter, personalized basketball bedding, custom sports bedding", "black, gray and orange comforter collage with basketballs, hoops, player silhouettes, sample name and jersey number", "custom name number basketball collage comforter"),
    230: ("Custom Shattered Glass Basketball Comforter", "custom shattered glass basketball comforter", "shattered glass basketball comforter, personalized basketball bedding, custom basketball bedding", "dark shattered-glass comforter with a large basketball, script sample name and jersey number on pillow shams", "custom name number shattered glass basketball comforter"),
}

META_DESCRIPTIONS = {
    221: "Custom basketball hoop blanket with glowing rim and net artwork. Add optional name and number in fleece or sherpa.",
    222: "Personalized basketball shoes blanket with court sneakers and ball graphics. Add optional name and number in fleece or sherpa.",
    223: "Custom court lines basketball comforter with gold court graphics. Personalize with name and number on soft bedding.",
    224: "Custom basketball court perspective comforter in black and red sports styling. Personalize with name and number.",
    225: "Custom hardwood court basketball comforter with blue gold arena artwork. Personalize with player name and number.",
    226: "Custom neon basketball player blanket with dunking silhouette and hoop graphic. Add optional name and number.",
    227: "Custom flaming basketball player blanket with dribbler artwork and fire effects. Add optional name and number.",
    228: "Custom dribbling silhouette basketball blanket with gray orange graphics. Add optional name and number.",
    229: "Custom basketball collage comforter with hoops, balls and player silhouettes. Personalize with name and number.",
    230: "Custom shattered glass basketball comforter with 3D cracked artwork. Add name and number; optional shams available.",
}


IMAGE_DETAILS = {
    221: {
        1: ("Blanket flat mockup with close-up hoop, white net, sample name Justin and number 23.", "Custom basketball hoop blanket flat view"),
        2: ("Person holds the basketball hoop blanket in a living room with the same sample name and number.", "Basketball hoop blanket held in living room"),
        3: ("Blanket size chart showing 40x30 through 80x60 options with scale figure.", "Basketball hoop blanket size chart"),
        4: ("Basketball hoop blanket draped across a sofa under a window.", "Basketball hoop blanket on sofa"),
        5: ("Feature panel overlays close-up blanket artwork with fluffy, quality and no-shedding callouts.", "Basketball hoop blanket feature panel"),
        6: ("Folded blanket and fabric/care collage with breathable, skin-friendly and machine washable callouts.", "Basketball hoop blanket fabric care panel"),
        7: ("Custom blanket size guide showing bed placement examples and width/length table.", "Basketball hoop blanket bed size guide"),
        8: ("Lifestyle image with adult and child reading under the basketball hoop blanket.", "Basketball hoop blanket family lifestyle image"),
    },
    222: {
        1: ("Blanket flat mockup with dark basketball and athletic shoes artwork, sample name Emery and number 33.", "Custom basketball shoes blanket flat view"),
        2: ("Person holds basketball shoes blanket in a living room setting.", "Basketball shoes blanket held in living room"),
        3: ("Size chart panel showing four blanket sizes with basketball shoes artwork.", "Basketball shoes blanket size chart"),
        4: ("Feature panel over ball-and-shoes artwork with fluffy, high-quality and no-pilling callouts.", "Basketball shoes blanket feature panel"),
        5: ("Basketball shoes blanket draped on sofa beneath round wall mirror.", "Basketball shoes blanket on sofa"),
        6: ("Folded blanket and fabric/care collage for the basketball shoes design.", "Basketball shoes blanket fabric care panel"),
        7: ("Custom blanket bed size guide with width and length table.", "Basketball shoes blanket bed size guide"),
        8: ("Adult and child reading under a blanket showing the script sample name.", "Basketball shoes blanket family lifestyle image"),
    },
    223: {
        1: ("Bedroom mockup of black court-lines basketball comforter with sample name T Mack and number 5.", "Custom court lines basketball comforter"),
        2: ("Second bedroom mockup of the court-lines comforter with folded white duvet edge.", "Court lines basketball bedding mockup"),
        3: ("Panel comparing duvet cover set and comforter set for all-season use.", "Basketball bedding type comparison panel"),
        4: ("Feature panel with basketball court print close-ups, microfiber and 3D printed pattern callouts.", "Court lines basketball comforter feature panel"),
        5: ("Low-angle bed mockup with court-lines basketball bedding and comfort icon strip.", "Court lines basketball bedding close view"),
        6: ("Easy care panel with pillows and wrinkle-free, stain-proof, anti-pilling callouts.", "Basketball bedding easy care panel"),
        7: ("Size dimension panel showing twin, full, queen and king bed sizes.", "Basketball comforter size dimension panel"),
    },
    224: {
        1: ("Bedroom mockup of black and red basketball court perspective comforter with sample name Anthony and number 8.", "Custom basketball court perspective comforter"),
        2: ("White zipper close-up panel labeled bottom zippered closure.", "Bottom zippered closure bedding panel"),
        3: ("High-density weaving panel with fabric icons and weave comparison.", "High-density weaving bedding panel"),
        4: ("Machine washable panel with washing machine and laundry baskets.", "Machine washable bedding care panel"),
        5: ("Blue table size panel for duvet cover and pillowcase dimensions with bed icons.", "Basketball bedding size chart panel"),
    },
    225: {
        1: ("Bedroom mockup of blue and gold hardwood court basketball comforter with sample name Matthew and number 2.", "Custom hardwood court basketball comforter"),
        2: ("White zipper close-up panel labeled bottom zippered closure.", "Bottom zippered closure bedding panel for hardwood court comforter"),
        3: ("High-density weaving feature panel with fabric and breathability icons.", "High-density weaving panel for hardwood basketball bedding"),
        4: ("Machine washable care panel with washing machine and laundry baskets.", "Hardwood basketball bedding easy care panel"),
        5: ("Size chart panel listing US twin, full, queen and king duvet cover and pillowcase dimensions.", "Hardwood basketball bedding size chart"),
    },
    226: {
        1: ("Blanket flat mockup with neon dunking player silhouette, hoop, sample name Fernando and number 51.", "Custom neon basketball player blanket flat view"),
        2: ("Person holds neon basketball player blanket in a living room.", "Neon basketball player blanket held in living room"),
        3: ("Size chart panel showing four blanket sizes for the dunking player design.", "Neon basketball player blanket size chart"),
        4: ("Feature panel over player-and-hoop artwork with fluffy and no-pilling callouts.", "Neon basketball player blanket feature panel"),
        5: ("Neon basketball player blanket draped across a sofa.", "Neon basketball player blanket on sofa"),
        6: ("Folded blanket and fabric/care collage for the neon player design.", "Neon basketball player blanket fabric care panel"),
        7: ("Custom blanket bed size guide with example placements and size table.", "Neon basketball player blanket bed size guide"),
        8: ("Adult and child reading under the neon basketball player blanket.", "Neon basketball player blanket family lifestyle image"),
    },
    227: {
        1: ("Blanket flat mockup with front dribbling basketball player, flame effect, sample name King David and number 23.", "Custom flaming basketball player blanket flat view"),
        2: ("Person holds flaming basketball player blanket in a living room.", "Flaming basketball player blanket held in living room"),
        3: ("Blanket size chart with four sizes and player artwork thumbnails.", "Flaming basketball player blanket size chart"),
        4: ("Feature panel over flame player artwork with fluffy, quality and no-shedding callouts.", "Flaming basketball player blanket feature panel"),
        5: ("Flaming basketball player blanket draped over sofa.", "Flaming basketball player blanket on sofa"),
        6: ("Folded blanket and fabric/care collage for the flaming player design.", "Flaming basketball player blanket fabric care panel"),
        7: ("Custom blanket bed size guide with width and length table.", "Flaming basketball player blanket bed size guide"),
        8: ("Adult and child reading under blanket showing flame player artwork.", "Flaming basketball player blanket family lifestyle image"),
    },
    228: {
        1: ("Blanket flat mockup with black dribbling player silhouette, orange ball graphic, sample name Connor and number 3.", "Custom dribbling silhouette basketball blanket flat view"),
        2: ("Person holds dribbling silhouette basketball blanket in a living room.", "Dribbling silhouette basketball blanket held in living room"),
        3: ("Size chart panel showing four blanket sizes with silhouette artwork thumbnails.", "Dribbling silhouette basketball blanket size chart"),
        4: ("Feature panel over gray and orange basketball artwork with fabric callouts.", "Dribbling silhouette basketball blanket feature panel"),
        5: ("Dribbling silhouette basketball blanket draped on sofa.", "Dribbling silhouette basketball blanket on sofa"),
        6: ("Folded blanket and fabric/care collage for the silhouette basketball design.", "Dribbling silhouette basketball blanket fabric care panel"),
        7: ("Custom blanket bed size guide with width and length table.", "Dribbling silhouette basketball blanket bed size guide"),
        8: ("Adult and child reading under blanket showing dribbling silhouette artwork.", "Dribbling silhouette basketball blanket family lifestyle image"),
    },
    229: {
        1: ("Bedroom mockup of basketball collage comforter with balls, hoops, player silhouettes, sample name Bray and number 12.", "Custom basketball collage comforter"),
        2: ("Second bed mockup of basketball collage bedding with folded white duvet edge.", "Basketball collage bedding mockup"),
        3: ("Duvet cover set versus comforter set comparison panel.", "Basketball collage bedding type panel"),
        4: ("Feature panel with collage print close-ups, microfiber and 3D printed pattern callouts.", "Basketball collage comforter feature panel"),
        5: ("Low-angle bed mockup with basketball collage print and comfort icon strip.", "Basketball collage bedding close view"),
        6: ("Easy care panel with pillow stack and care callouts.", "Basketball collage bedding easy care panel"),
        7: ("Size dimension panel showing twin, full, queen and king bed diagrams.", "Basketball collage comforter size panel"),
    },
    230: {
        1: ("Bedroom mockup of shattered-glass basketball comforter with script sample name Aiden and number 17 on pillows.", "Custom shattered glass basketball comforter"),
        2: ("Second bed mockup of shattered-glass basketball bedding with folded white duvet edge.", "Shattered glass basketball bedding mockup"),
        3: ("Duvet cover set versus comforter set comparison panel.", "Shattered glass basketball bedding type panel"),
        4: ("Feature panel with basketball print close-ups, microfiber and 3D printed pattern callouts.", "Shattered glass basketball comforter feature panel"),
        5: ("Low-angle bed mockup of shattered-glass basketball bedding with comfort icon strip.", "Shattered glass basketball bedding close view"),
        6: ("Size dimension panel showing twin, full, queen and king bed diagrams.", "Shattered glass basketball comforter size panel"),
    },
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def load_summaries():
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    return {int(item["inventory_row"]["inventory_position"]): item for item in summaries}


def parse_customizer(summary):
    html_file = Path(summary["html"]["html_file"])
    raw = html_file.read_text(encoding="utf-8", errors="ignore")
    page = html.unescape(raw)
    fields = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"minLength":\s*(\d+),\s*"maxLength":\s*(\d+)',
        page,
    )
    if not fields:
        fields = re.findall(
            r'"label"\s*:\s*"([^"]+)"\s*,\s*"required"\s*:\s*(true|false)\s*,\s*"minLength"\s*:\s*(\d+)\s*,\s*"maxLength"\s*:\s*(\d+)',
            page,
        )
    options = re.findall(r'"label"\s*:\s*"([^"]+)"\s*,\s*"required"\s*:\s*(true|false)\s*,\s*"defaultOptionId"', page)
    return fields, options


def customizer_sentence(pos):
    if pos in {221, 222, 226, 227, 228}:
        return (
            "Custom Name is optional and supports 1 to 200 characters. "
            "Custom Number is optional and supports 1 to 20 characters."
        )
    return (
        'Customize Your Name is required and supports 1 to 30 characters. '
        'Customize Your Number is required and supports 1 to 5 characters. '
        'Enter "NO" in either field if that printed detail is not wanted.'
    )


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    missing = sorted(set(handles) - set(admin))
    if missing:
        raise RuntimeError(f"Missing handles in admin export: {missing}")
    return admin


def option_text(admin_row):
    values = [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]]
    values = [v for v in values if v]
    return ", ".join(values) if values else "available size options"


def description(pos, admin_row, ctext):
    title, _, _, detail, _ = PRODUCT_UPDATES[pos]
    if pos in {221, 222, 226, 227, 228}:
        return (
            f"<p>The {title} features {detail} on a soft basketball blanket design with optional name and number personalization.</p>"
            "<h3>Artwork and Personalization</h3>"
            f"<ul><li>{ctext}</li>"
            "<li>Sample names and numbers shown in gallery mockups are examples and are replaced by submitted custom text when provided.</li>"
            '<li>Choose a blanket size from the size selector shown on the product page.</li></ul>'
            "<h3>Fabric Options and Care</h3>"
            "<ul><li>Fleece and sherpa-style blanket options are shown in the product and size panels where available.</li>"
            "<li>Gallery panels show the main blanket mockup, size guide, fabric or feature details and lifestyle views.</li>"
            "<li>Machine wash cold on gentle cycle and tumble dry low.</li></ul>"
        )
    return (
        f"<p>The {title} features {detail} for a custom basketball bedding design with name and number personalization.</p>"
        "<h3>Personalization and Product Type</h3>"
        f"<ul><li>{ctext}</li>"
        "<li>Choose Comforter or Duvet Cover where the product-type selector is shown.</li>"
        "<li>The Duvet Cover option includes a bottom zippered closure shown in the gallery feature panel.</li></ul>"
        "<h3>Sizes, Add-ons and Care</h3>"
        '<ul><li>Available size panels list Twin 68" x 86", Full 79" x 90", Queen 90" x 90", and King 90" x 104" where shown.</li>'
        "<li>Choose None, 1 Pillowcase or 2 Pillowcases where the pillowcase selector is shown.</li>"
        "<li>An optional matching flat sheet is available where the sheet-cover selector is shown.</li>"
        "<li>For the shattered glass comforter, pillow shams shown in mockups are optional add-ons when selected.</li>"
        "<li>Feature panels show microfiber fabric, bedding type, care, size and zipper details.</li>"
        "<li>Machine wash cold on gentle cycle and tumble dry low.</li></ul>"
    )


def meta_description(pos, fields):
    return META_DESCRIPTIONS[pos]


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)

    summaries = load_summaries()
    batch_positions = set(PRODUCT_UPDATES)
    handles = {summaries[pos]["inventory_row"]["Handle"] for pos in batch_positions}
    admin = load_admin(handles)
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    pos_by_key = {summaries[pos]["inventory_row"]["product_key"]: pos for pos in batch_positions}
    handle_by_pos = {pos: summaries[pos]["inventory_row"]["Handle"] for pos in batch_positions}
    fields_by_pos = {}
    option_fields_by_pos = {}
    for pos in batch_positions:
        fields_by_pos[pos], option_fields_by_pos[pos] = parse_customizer(summaries[pos])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        handle = handle_by_pos[pos]
        title, primary, secondary, _, cluster = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ctext = customizer_sentence(pos)
        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or summaries[pos]["inventory_row"]["product_type"]
        ws.cell(row_num, idx["title_proposed"]).value = title
        ws.cell(row_num, idx["meta_title_seo"]).value = title
        ws.cell(row_num, idx["meta_title_length"]).value = len(title)
        meta = meta_description(pos, fields_by_pos[pos])
        ws.cell(row_num, idx["meta_description_seo"]).value = meta
        ws.cell(row_num, idx["meta_description_length"]).value = len(meta)
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, ctext)
        ws.cell(row_num, idx["primary_keyword"]).value = primary
        ws.cell(row_num, idx["secondary_keywords"]).value = secondary
        ws.cell(row_num, idx["meta_keyword"]).value = primary
        ws.cell(row_num, idx["keyword_strategy"]).value = (
            f"Target the specific basketball motif: {cluster}; keep broader personalized sports bedding terms for collections."
        )
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"US English buyer intent targets {cluster}; no search-volume, trend or ranking claim is made."
        )
        ws.cell(row_num, idx["keyword_evidence_level"]).value = "SERP_ONLY"
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R3 after qa_batch_023_r2 QA: used products_export_1.csv sha256:{admin_hash}; "
            "rewrote truncated meta descriptions, removed process-style description copy, separated blanket versus comforter option logic, "
            "and restored product-type, size, add-on and care details. Still NEEDS_REVIEW."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    seen_alts = set()
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pos = next((p for p, h in handle_by_pos.items() if h == handle), None)
        if pos is None:
            continue
        image_number = int(ws.cell(row_num, idx["image_number"]).value)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation, alt = IMAGE_DETAILS[pos][image_number]
        if alt in seen_alts:
            alt = f"{alt} for {PRODUCT_UPDATES[pos][4]}"
        seen_alts.add(alt)
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = alt[:125]
        ws.cell(row_num, idx["alt_action"]).value = "SET"
        ws.cell(row_num, idx["viewed_status"]).value = "VIEWED"
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R3: observation and alt carried forward from r2 checked image audit; admin image URL and current alt matched from Shopify CSV where available."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        title, primary, secondary, _, cluster = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in secondary.split(",")]
        keyword = primary if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = cluster
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "PRODUCT"
        ws.cell(row_num, idx["decision_reason"]).value = f"{title} is separated by verified sport, product form, layout and custom fields."
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = (
            "qa_batch_023_r2 QA report + products_export_1.csv admin baseline + r3 meta/description cleanup"
        )
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = "YES; nearby basketball designs are differentiated by motif and product form."
        ws.cell(row_num, idx["mapping_reason"]).value = "R3 separates nearby pages by visible motif, blanket/comforter form and verified custom fields."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r3"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        title, primary, _, _, cluster = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {cluster} as a specific product page."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {cluster}, the buyer wants the page to confirm exact artwork, product type, size choices, "
            "and available name/number fields before purchase."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = (
            "Confirm product form, basketball motif, visible layout, custom text fields, care or size panels and image accuracy before purchase."
        )
        ws.cell(row_num, idx["emotional_social_motivation"]).value = (
            "Choose a basketball bedding or blanket item that feels specific to a player, fan, sport room or gift recipient."
        )
        ws.cell(row_num, idx["purchase_concerns"]).value = (
            "Confirm spelling and number where fields are available; avoid unsupported team, player, league, material or delivery claims."
        )
        ws.cell(row_num, idx["source_refs"]).value = (
            f"products_export_1.csv sha256:{admin_hash}; QA report {QA_RUN_ID}; checked images 221-230"
        )
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R3_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No Search Console, paid keyword volume or internal site-search export supplied; proof remains SERP-only."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {primary} with title: {title}."

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in batch_positions:
            continue
        handle = handle_by_pos[pos]
        admin_row = admin[handle]
        _, _, _, detail, _ = PRODUCT_UPDATES[pos]
        ctext = customizer_sentence(pos)
        ws.cell(row_num, idx["sources_accessed"]).value = (
            f"Original storefront/product evidence; QA report {QA_RUN_ID}; products_export_1.csv sha256:{admin_hash}; checked image audit"
        )
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={option_text(admin_row)}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={detail}; customizer={ctext}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "qa_batch_023_r2 QA D1/D2 issues addressed in r3; independent QA recheck still required."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R3 uses checked image audit, QA report recommendations and Shopify admin CSV baseline; no approval or import file created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_023_r3_revision", "r3", "Separate 10-product revision after qa_batch_023_r2 QA; source workbook was not modified."),
        ("qa_batch_023_r3_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_023_r3_scope", "inventory positions 221-230", "No products outside qa_batch_023 were revised."),
        ("qa_batch_023_r3_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_023_r3_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_023_r3_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_evidence": str(SUMMARY_PATH.relative_to(ROOT)),
        "source_qa_report": str(QA_REPORT.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_023 only; inventory positions 221-230",
        "revision": "r3",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "r3_changes": [
            "Rewrote all 10 meta descriptions to avoid hard truncation and broken endings.",
            "Rewrote all 10 product descriptions to remove process/boilerplate copy flagged by qa_batch_023_r2 QA.",
            "Separated Blanket products with optional name and number fields from Comforter products with required name and number fields.",
            "Removed product-type selection language from Blanket products and kept it only for Comforter/Duvet Cover products.",
            "Clarified that pillow shams for position 230 are optional add-ons."
        ],
        "next_step": "Sang QA lại qa_batch_023_r3 before any approval or import.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join([
            "# Revision qa_batch_023_r3",
            "",
            "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
            "",
            f"- Nguồn r2: `{SOURCE.relative_to(ROOT)}`",
            f"- QA của Sang: `{QA_REPORT.relative_to(ROOT)}`",
            f"- Evidence summary: `{SUMMARY_PATH.relative_to(ROOT)}`",
            f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
            f"- Admin export SHA-256: `{admin_hash}`",
            f"- Workbook r3: `{OUTPUT.relative_to(ROOT)}`",
            f"- Sản phẩm sửa: {revised_products}",
            f"- Ảnh sửa/refresh: {revised_images}",
            f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
            "- Sửa trọng yếu: xử lý D1/D2 theo QA r2, viết lại meta không bị cắt cụt, bỏ boilerplate trong mô tả, phân biệt Blanket tùy chọn Name/Number với Comforter bắt buộc Name/Number, bỏ hướng dẫn chọn product type khỏi Blanket và làm rõ shams tùy chọn ở pos 230.",
            "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
            "- Bước tiếp theo: gửi Sang QA lại `qa_batch_023_r3` trước khi duyệt hay import.",
            "",
        ]),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
