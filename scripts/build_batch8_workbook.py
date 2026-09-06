import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_008"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_007.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_008.xlsx"


PROFILES = {
    71: ("personalized running football player comforter with Brian number 3", "personalized football player comforter", "Personalized Football Player Comforter", "Personalized Running Football Player Comforter", "teal and black running football player artwork with orange flame trails, name Brian and #3", "custom football bedding, football comforter with name, football bedding for boys", "EVERGREEN", "https://www.amazon.com/Football-Personalized-Blanket-Pillowcases-Bedspread/dp/B0BFGNX4HW; https://www.etsy.com/listing/1634020192/personalized-football-blanket-football; https://www.walmart.com/ip/Personalized-Sports-Themed-Blanket-Customized-Balls-Blankets-Custom-Name-Throws-Boys-Girls-Teen-Athletes-Fun-Birthday-Gift-Print-finish-USA-2nd-day-s/13557065996"),
    72: ("personalized electric football player comforter with Michael number 10", "custom football player comforter", "Custom Football Player Comforter Set", "Custom Electric Football Player Comforter", "blue and red electric football player running with the ball, name Michael and #10", "football bedding with name, personalized football comforter, sports comforter for boys", "EVERGREEN", "https://www.amazon.com/clp/B0CG1XRNMS; https://www.pinterest.com/ideas/football-player-pillowcase-with-name/946038768408/; https://www.ebay.com/itm/165689760708"),
    73: ("personalized American flag football player comforter with David number 33", "American flag football player comforter", "American Flag Football Player Comforter", "Personalized American Flag Football Player Comforter", "football player shown from behind over a torn American flag and metal texture background with David and 33", "patriotic football bedding, custom football player bedding, football comforter with number", "EVERGREEN", "https://www.amazon.com/Football-Personalized-Blanket-Pillowcases-Bedspread/dp/B0BFGNX4HW; https://www.amazon.com/Erosebridal-American-Football-Comforter-Federations/dp/B09L4T6CQG; https://www.walmart.com/c/kp/american-flag-bedding"),
    74: ("personalized burning football player comforter with Michael number 24", "burning football player comforter", "Burning Football Player Comforter", "Personalized Burning Football Player Comforter", "black football player holding the ball with number 24, name Michael and burning football background", "custom football bedding, football player comforter, personalized sports bedding", "EVERGREEN", "https://www.amazon.com/clp/B0CG1XRNMS; https://www.pinterest.com/ideas/football-bedding/932913586866/; https://www.ebay.com/itm/165689760708"),
    75: ("personalized football player collage comforter with David", "football player collage comforter", "Football Player Collage Comforter", "Personalized Football Player Collage Comforter", "black comforter with large football close-up, multiple football player silhouettes and the name David", "football bedding with name, football player bedding, custom sports comforter", "EVERGREEN", "https://www.amazon.com/clp/B0CG1XRNMS; https://www.pinterest.com/pin/personalized-football-bedding-set-custom-name-football-bedding-personalized-sports-bed-set-duvet-covers--413557178297079798/; https://www.ebay.com/itm/165689760708"),
    76: ("personalized football laces texture comforter with Michael number 33", "football laces comforter", "Football Laces Comforter Set", "Personalized Football Laces Comforter", "vintage brown football leather texture with prominent white laces, name Michael and number 33", "football texture bedding, custom football comforter, football bedding with name", "EVERGREEN", "https://www.amazon.com/s?k=football+comforter+set; https://www.pinterest.com/ideas/football-bedding/932913586866/; https://www.walmart.com/ip/Personalized-Sports-Themed-Blanket-Customized-Balls-Blankets-Custom-Name-Throws-Boys-Girls-Teen-Athletes-Fun-Birthday-Gift-Print-finish-USA-2nd-day-s/13557065996"),
    77: ("personalized glowing soccer ball comforter with Jackson number 10", "personalized soccer comforter", "Personalized Soccer Comforter Set", "Personalized Glowing Soccer Ball Comforter", "black and blue glowing soccer ball with sweeping light trails, name Jackson and number 10", "custom soccer bedding, soccer comforter with name, soccer bedding set for boys", "EVERGREEN", "https://www.amazon.com/Personalized-Bedding-Daughter-Friends-Birthday/dp/B0C3VYL11C; https://www.etsy.com/listing/956705304/personalized-soccer-ball-comforter-flame; https://www.pinterest.com/pin/personalized-soccer-duvet-cover-set-soccer-fire-ball-player-gift-idea-black-duvet-cover-pillowcases-custom--999939923509327748/"),
    78: ("personalized God says you are floral blanket for Jessica", "personalized God says you are blanket", "Personalized God Says You Are Blanket", "Personalized God Says You Are Floral Blanket", "soft floral blanket with butterflies, the name Jessica and God says you are Bible verse affirmations", "Christian affirmation blanket, personalized Bible verse blanket, floral faith blanket", "EVERGREEN", "https://www.walmart.com/ip/Customizaholic-Personalized-God-Says-You-Are-Blanket-Custom-Name-Bible-Verses-Floral-Christian-Faith-Gift-Encouragement-Birthday-Special-Moments/18727801098; https://www.etsy.com/listing/4514112948/god-says-you-are-personalized-floral; https://www.amazon.com/s?k=personalized+christian+blanket"),
    79: ("personalized floral butterfly Bible verse blanket for Haley", "personalized floral Bible verse blanket", "Personalized Floral Bible Verse Blanket", "Personalized Floral Butterfly Bible Verse Blanket", "cream blanket with Haley name, colorful flowers, butterflies and vertical you are affirmation words with Bible verse references", "Christian flower blanket, custom scripture blanket, personalized inspirational blanket", "EVERGREEN", "https://www.etsy.com/market/proverbs_31_throw_blankets; https://www.etsy.com/listing/4514112948/god-says-you-are-personalized-floral; https://www.amazon.com/s?k=personalized+christian+blanket"),
    80: ("personalized God Says I Am butterfly blanket for Elizabeth", "personalized God Says I Am butterfly blanket", "Personalized God Says I Am Butterfly Blanket", "Personalized God Says I Am Butterfly Blanket", "brown and tan God Says I Am design with monarch butterflies, the name Elizabeth and Bible verse affirmation blocks", "Christian butterfly blanket, custom Bible verse blanket, personalized faith blanket", "EVERGREEN", "https://www.etsy.com/market/god_says_i_am_personalized_blanket; https://www.facebook.com/CallieForGifts/videos/personalized-god-says-i-am-name-meaning-on-bible-verse-colorful-soft-throw-blank/2116242565948940/; https://christianartbag.com/products/christianart-blanket-god-says-i-am-christian-blanket-bible-verse-blanket-personalized-blanket-christmas-gift-cabbk01111123"),
}


def append_dict(ws, data):
    headers = [cell.value for cell in ws[1]]
    ws.append([data.get(h, "") for h in headers])


def image_obs(theme, pos):
    if pos == 1:
        return f"Main product mockup showing {theme}."
    if pos == 2:
        return f"Secondary flat or angled mockup showing {theme}."
    if pos == 3:
        return f"Bedding type, size or lifestyle information related to {theme}."
    if pos == 4:
        return f"Close room/detail, feature panel or size information for {theme}."
    if pos == 5:
        return f"Easy-care, material detail or additional mockup for {theme}."
    if pos == 6:
        return f"Size guide, feature image or bedroom detail for {theme}."
    return f"Additional lifestyle, feature or product detail image showing {theme}."


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        issue_bits = [
            "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."
        ]
        if 71 <= pos <= 76:
            issue_bits.append("Closely related football comforters should be reviewed for keyword cannibalization and title uniqueness.")
        if html_error:
            issue_bits.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
        issues = " ".join(issue_bits)
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for themed bedroom decor, faith gifting or personalized sports bedding.</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork, personalization and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broader sports bedding or Christian gift terms for collections or stronger representatives.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as a personalized sports item, themed bedding or faith gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": issues,
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for personalized bedding, the buyer wants a coordinated {kind} matching a specific sport, name, number or faith message.",
            "functional_motivation": "Find the right design, size, bedding type, care details and personalization fit.",
            "emotional_social_motivation": "Create a sports-themed bedroom, give a personal player gift, or choose an encouraging faith keepsake.",
            "purchase_concerns": "Fit, included components, material/care claims, personalization spelling, number accuracy and whether similar designs are distinct enough.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork, personalization and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed in a main product mockup",
                2: f"Secondary mockup of {h1}",
                3: f"Product option or lifestyle image for {h1}",
                4: f"Feature or size information for {h1}",
                5: f"Care, material or additional mockup for {h1}",
                6: f"Size guide or feature image for {h1}",
                7: f"Additional product mockup of {h1}",
                8: f"Lifestyle use image for {h1}",
                9: f"Additional lifestyle image for {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_008_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_008_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_008"},
        {"metric": "cumulative_products", "value": "80", "definition": "Total products included through batch_008"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[80:90]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_009",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
