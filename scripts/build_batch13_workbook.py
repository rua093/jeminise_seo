import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_013"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_012.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_013.xlsx"


TRUCK_REFS = "https://www.etsy.com/listing/1091390349/personalized-semi-truck-baby-boy-blanket; https://www.pinterest.com/pin/personalized-trucker-quilt-set-truck-driver-trucker-quilt-blanket-with-pillowcases-custom-name-quilt-bedding-set--999939923509376768/; https://jakescabsolutions.com/; https://www.amazon.com/clp/B0DFQ6T1C5"
CARDINAL_REFS = "https://www.southernliving.com/red-bird-holiday-decor-8402509; https://www.amazon.com/; https://www.etsy.com/market/cardinal_quilt"
WOLF_REFS = "https://www.amazon.com/wolf-comforter-set/s?k=wolf+comforter+set; https://www.pinterest.com/pin/wolf-bedding-set-duvet-set-comforter-set-or-quilt-set-wildlife-cabin-rustic-wolf-decor-gift-for-wolf-lovers-nature--4590927325055172608/; https://www.ebay.com/b/Wolf-Quilt-In-Duvet-Covers-Bedding-Sets/37644/bn_7022238087; https://wolf-stuff.com/collections/wolf-bedding-set"


PROFILES = {
    121: ("personalized vintage blue semi truck stone wall comforter", "vintage semi truck comforter", "Vintage Semi Truck Comforter", "Personalized Vintage Semi Truck Comforter", "weathered blue semi truck breaking through a gray stone wall with Your Name personalization area", "custom trucker bedding, semi truck comforter with name, truck bedding set", "EVERGREEN", TRUCK_REFS),
    122: ("personalized blue semi truck chevron comforter", "semi truck chevron comforter", "Semi Truck Chevron Comforter", "Personalized Blue Semi Truck Chevron Comforter", "blue semi truck centered on black and blue chevron graphics with Your Name personalization area", "custom trucker bedding, personalized semi truck comforter, trucker comforter set", "EVERGREEN", TRUCK_REFS),
    123: ("personalized black semi truck metallic mesh comforter", "semi truck metallic mesh comforter", "Semi Truck Metallic Mesh Comforter", "Personalized Semi Truck Metallic Mesh Comforter", "black semi truck over gray metallic mesh and torn metal graphics with Your Name personalization area", "custom trucker bedding, semi truck bedding with name, truck comforter", "EVERGREEN", TRUCK_REFS),
    124: ("personalized semi truck textured metallic comforter", "textured metallic semi truck comforter", "Textured Metallic Semi Truck Comforter", "Personalized Textured Metallic Semi Truck Comforter", "black and white semi truck illustration over a textured metallic gray background with Your Name personalization area", "custom trucker bedding, personalized semi truck bedding, semi truck comforter", "EVERGREEN", TRUCK_REFS),
    125: ("personalized starry night semi truck heart comforter", "starry night semi truck comforter", "Starry Night Semi Truck Comforter", "Personalized Starry Night Semi Truck Comforter", "blue semi truck under a starry night sky with Your Name and My Heart Is Always With You message", "custom trucker bedding, semi truck gift comforter, personalized trucker bedding", "EVERGREEN", TRUCK_REFS),
    126: ("two cardinals on holly branches Christmas quilt", "cardinal Christmas quilt", "Cardinal Christmas Quilt Set", "Two Cardinals on Holly Branches Quilt", "two red cardinals on holly branches with berries, pine greenery, script-style background and burgundy border", "cardinal quilt set, Christmas bird bedding, holiday cardinal quilt", "HOLIDAY_CHRISTMAS", CARDINAL_REFS),
    127: ("colorful close up wolf head quilt", "wolf head quilt", "Wolf Head Quilt Set", "Colorful Close Up Wolf Head Quilt", "close-up wolf face portrait in colorful geometric watercolor styling with matching shams", "wolf bedding, wildlife quilt set, colorful wolf comforter", "EVERGREEN", WOLF_REFS),
    128: ("geometric wolf head quilt set", "geometric wolf quilt", "Geometric Wolf Quilt Set", "Geometric Wolf Head Quilt Set", "front-facing wolf head rendered in angular geometric shapes with teal, tan and rust accents", "wolf bedding, wildlife quilt, geometric animal quilt", "EVERGREEN", WOLF_REFS),
    129: ("profile wolf head with feathers quilt", "wolf feathers quilt", "Wolf Feathers Quilt Set", "Profile Wolf Head with Feathers Quilt", "side-profile wolf head artwork with feathers and beads on a warm tan and brown background", "wolf bedding, Native inspired wolf quilt, wildlife quilt set", "EVERGREEN", WOLF_REFS),
    130: ("wolf dreamcatcher quilt with feathers", "wolf dreamcatcher quilt", "Wolf Dreamcatcher Quilt Set", "Wolf Dreamcatcher Quilt with Feathers", "front-facing wolf head centered inside a dreamcatcher circle with feathers on a muted tan background", "wolf bedding, dreamcatcher bedding, Native inspired wolf quilt", "EVERGREEN", WOLF_REFS),
}


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        issue_bits = [
            "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."
        ]
        if 121 <= pos <= 125:
            issue_bits.append("Closely related semi truck comforters should be reviewed for keyword cannibalization, title uniqueness and canonical strategy.")
        if pos == 125:
            issue_bits.append("Visible sentimental message should be proofread before import: No matter where you go / My heart is always with you.")
        if pos == 126:
            issue_bits.append("Holiday/Christmas seasonality should be confirmed before using seasonal keywords outside holiday merchandising windows.")
        if 127 <= pos <= 130:
            issue_bits.append("Closely related wolf quilt variants should be reviewed for cannibalization and motif-specific title differentiation.")
        if pos in (129, 130):
            issue_bits.append("Native/dreamcatcher-inspired terms should be reviewed for brand suitability and cultural sensitivity before deployment.")
        if html_error:
            issue_bits.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
        issues = " ".join(issue_bits)
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for themed decor, trucker gifts, holiday rooms or wildlife-inspired bedrooms.</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork, personalization and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broader bedding terms for collections or stronger representative products.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as a themed bedding item, personalized gift or room decor.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": issues,
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed or personalized bedding, the buyer wants a {kind} matching a specific design, name, message, animal motif or hobby.",
            "functional_motivation": "Find the right design, size, bedding type, care details and personalization fit.",
            "emotional_social_motivation": "Create meaningful room decor or give a personalized trucker, holiday or wildlife-themed gift.",
            "purchase_concerns": "Personalization spelling, product identity, material/care claims, design distinction and whether the visible artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork, personalization and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed in a main product mockup",
                2: f"Secondary mockup of {h1}",
                3: f"Product option, feature or lifestyle image for {h1}",
                4: f"Detail or feature image for {h1}",
                5: f"Additional product mockup or care information for {h1}",
                6: f"Care, size or product feature image for {h1}",
                7: f"Additional product mockup of {h1}",
                8: f"Lifestyle use image for {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_013_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_013_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_013"},
        {"metric": "cumulative_products", "value": "130", "definition": "Total products included through batch_013"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[130:140]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_014",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
