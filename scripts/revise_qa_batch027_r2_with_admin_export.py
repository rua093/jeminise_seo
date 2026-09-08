import csv
import hashlib
import html
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_027_r2"
SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "batches" / "SEO_Product_Optimization_through_batch_027.xlsx"
SUMMARY_PATH = ROOT / "seo_runs" / SHOP / RUN_ID / "batch_027_evidence_summary.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_027_r2.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_027_r2.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    261: ("Personalized Christian Woman Affirmation Blanket", "personalized Christian woman affirmation blanket", "Christian affirmation blanket, Bible verse throw blanket, faith gift blanket", "wood-panel affirmation throw blanket with seated woman illustration, Bible verse labels and sample name Rebecca", "personalized Christian woman affirmation blanket"),
    262: ("Personalized Christian Photo Scripture Blanket", "personalized Christian photo scripture blanket", "Christian photo blanket, personalized scripture blanket, faith keepsake blanket", "vintage beige scripture blanket with circular photo area, faith quotes and sample portrait", "personalized Christian photo scripture blanket"),
    263: ("Personalized Colorful Celtic Tree Quilt Set", "personalized colorful Celtic Tree quilt set", "Celtic Tree of Life quilt, colorful Yggdrasil bedding, custom Celtic quilt", "colorful teal and orange quilt set with central Tree of Life medallion and Celtic knot border", "personalized colorful Celtic Tree quilt set"),
    264: ("Personalized Vintage Yggdrasil Quilt Set", "personalized vintage Yggdrasil quilt set", "Yggdrasil quilt set, Tree of Life bedding, custom Celtic quilt", "sunset-toned quilt set with large Tree of Life artwork and Celtic border", "personalized vintage Yggdrasil quilt set"),
    265: ("Personalized Desert RV Sunset Quilt Set", "personalized desert RV sunset quilt set", "desert RV quilt set, cactus sunset bedding, custom camper quilt", "colorful desert quilt set with RV camper, cactus, flowers, mountains and sunset", "personalized desert RV sunset quilt set"),
    266: ("Personalized Rainbow Tree of Life Quilt Set", "personalized rainbow Tree of Life quilt set", "colorful Tree of Life quilt, rainbow Yggdrasil bedding, custom nature quilt", "dark quilt set with rainbow leaf Tree of Life artwork and glowing sun", "personalized rainbow Tree of Life quilt set"),
    267: ("Personalized Celtic Border Yggdrasil Quilt Set", "personalized Celtic border Yggdrasil quilt set", "Celtic border quilt set, Yggdrasil quilt, custom Tree of Life bedding", "black quilt set with green Celtic border, orange corner knots and central Yggdrasil tree", "personalized Celtic border Yggdrasil quilt set"),
    268: ("Personalized Couple Photo Quilt Set", "personalized couple photo quilt set", "custom couple quilt, photo quilt set, personalized Valentine's quilt", "couple portrait quilt set with custom photo badge and Happy Valentine's Day sample text", "personalized couple photo quilt set"),
    269: ("Personalized Cozy Room Bookworm Blanket", "personalized cozy room bookworm blanket", "personalized bookworm blanket, reading girl blanket, custom reader blanket", "cozy room reading blanket with bookshelf, green chair, cartoon reader, sample name and background options", "personalized cozy room bookworm blanket"),
    270: ("Personalized Dark Hair Book Lover Blanket", "personalized dark hair book lover blanket", "personalized book lover blanket, reading girl throw blanket, custom reader gift", "book lover blanket with dark-haired girl, stacked books, floral beige background and So Many Books So Little Time text", "personalized dark hair book lover blanket"),
}


IMAGE_DETAILS = {
    261: {1: ("Person holds wood-panel Christian affirmation blanket with seated woman, Bible verse labels and sample name Rebecca.", "Christian woman affirmation blanket held view"), 2: ("Flat view of wood-panel affirmation blanket with God Says I Am heading and verse blocks.", "Christian woman affirmation blanket flat view"), 3: ("Christian affirmation blanket draped on sofa under a window.", "Christian woman affirmation blanket on sofa"), 4: ("Second sofa mockup of Christian affirmation blanket beneath round mirror.", "Christian woman affirmation blanket sofa mockup"), 5: ("Feature panel over affirmation artwork with fluffy and no-shedding callouts.", "Christian woman affirmation blanket feature panel")},
    262: {1: ("Person holds vintage Christian scripture blanket with circular photo area and faith quotes.", "Christian photo scripture blanket held view"), 2: ("Feature panel showing white fleece edge and machine washable, soft and warm callouts.", "Christian photo scripture blanket feature panel"), 3: ("Close-up of photo area and scripture text beside an open book.", "Christian photo scripture blanket close-up"), 4: ("Size and use panel for vintage Christian photo scripture blanket.", "Christian photo scripture blanket size and use panel"), 5: ("Detail collage showing printed photo area, scripture typography and white fleece texture.", "Christian photo scripture blanket detail collage"), 6: ("Christian photo scripture blanket draped over an armchair.", "Christian photo scripture blanket chair mockup"), 7: ("Lifestyle reading scene with Christian photo scripture blanket, open book and small pet.", "Christian photo scripture blanket reading lifestyle")},
    263: {1: ("Bed mockup of colorful Celtic Tree quilt with teal orange medallion and matching pillows.", "Colorful Celtic Tree quilt bed view"), 2: ("Second room mockup of colorful Celtic Tree quilt with printed craft callout.", "Colorful Celtic Tree quilt room mockup"), 3: ("Close bed view with optional pillow shams and Tree of Life artwork.", "Colorful Celtic Tree quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and Celtic tree artwork sample.", "Colorful Celtic Tree quilt fabric panel"), 5: ("Premium quilt sets sizing chart for colorful Celtic Tree design.", "Colorful Celtic Tree quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Colorful Celtic Tree quilt features panel"), 7: ("Overhead bedroom mockup of colorful Celtic Tree quilt and matching pillows.", "Colorful Celtic Tree quilt overhead view")},
    264: {1: ("Bed mockup of sunset Yggdrasil quilt with large tree and Celtic border.", "Vintage Yggdrasil quilt bed view"), 2: ("Second room mockup of sunset Yggdrasil quilt with craft callout.", "Vintage Yggdrasil quilt room mockup"), 3: ("Close bed view with optional pillow shams and sunset tree artwork.", "Vintage Yggdrasil quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and yellow tree artwork sample.", "Vintage Yggdrasil quilt fabric panel"), 5: ("Premium quilt sets sizing chart for vintage Yggdrasil design.", "Vintage Yggdrasil quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Vintage Yggdrasil quilt features panel"), 7: ("Overhead bedroom mockup of vintage Yggdrasil quilt and matching pillows.", "Vintage Yggdrasil quilt overhead view")},
    265: {1: ("Bed mockup of colorful desert RV quilt with cactus, flowers, mountains and sunset.", "Desert RV sunset quilt bed view"), 2: ("Second room mockup of desert RV sunset quilt with craft callout.", "Desert RV sunset quilt room mockup"), 3: ("Close bed view with optional pillow shams and desert camper artwork.", "Desert RV sunset quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and desert artwork sample.", "Desert RV sunset quilt fabric panel"), 5: ("Premium quilt sets sizing chart for desert RV sunset design.", "Desert RV sunset quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Desert RV sunset quilt features panel"), 7: ("Overhead bedroom mockup of desert RV sunset quilt and pillows.", "Desert RV sunset quilt overhead view")},
    266: {1: ("Bed mockup of rainbow Tree of Life quilt with glowing sun and colorful leaves.", "Rainbow Tree of Life quilt bed view"), 2: ("Second room mockup of rainbow Tree of Life quilt with craft callout.", "Rainbow Tree of Life quilt room mockup"), 3: ("Close bed view with optional pillow shams and rainbow leaf tree artwork.", "Rainbow Tree of Life quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and rainbow artwork sample.", "Rainbow Tree of Life quilt fabric panel"), 5: ("Premium quilt sets sizing chart for rainbow Tree of Life design.", "Rainbow Tree of Life quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Rainbow Tree of Life quilt features panel"), 7: ("Overhead bedroom mockup of rainbow Tree of Life quilt and pillows.", "Rainbow Tree of Life quilt overhead view")},
    267: {1: ("Bed mockup of black Yggdrasil quilt with green Celtic border and orange corner knots.", "Celtic border Yggdrasil quilt bed view"), 2: ("Second room mockup of Celtic border Yggdrasil quilt with craft callout.", "Celtic border Yggdrasil quilt room mockup"), 3: ("Close bed view with optional pillow shams and Celtic border artwork.", "Celtic border Yggdrasil quilt pillow sham panel"), 4: ("High-quality fabric panel with quilt layers and green tree artwork sample.", "Celtic border Yggdrasil quilt fabric panel"), 5: ("Premium quilt sets sizing chart for Celtic border Yggdrasil design.", "Celtic border Yggdrasil quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Celtic border Yggdrasil quilt features panel"), 7: ("Overhead bedroom mockup of Celtic border Yggdrasil quilt and pillows.", "Celtic border Yggdrasil quilt overhead view")},
    268: {1: ("Overhead bed mockup of couple photo quilt with Happy Valentine's Day sample text.", "Personalized couple photo quilt bed view"), 2: ("Flat couple portrait quilt with custom photo badge and Valentine's Day text.", "Couple photo quilt flat view"), 3: ("Premium quilt sets sizing chart with couple photo quilt example.", "Couple photo quilt size chart"), 4: ("Bedspread features panel showing microfiber layers and care icons.", "Couple photo quilt features panel"), 5: ("Room mockup of couple photo quilt with printed craft callout.", "Couple photo quilt room mockup"), 6: ("Close bed view with optional pillow shams and couple portrait artwork.", "Couple photo quilt pillow sham panel"), 7: ("High-quality fabric panel with quilt layers and couple photo artwork sample.", "Couple photo quilt fabric panel")},
    269: {1: ("Held bookworm blanket with Sophia Bookworm text, bookshelf, green chair and cartoon reader.", "Cozy room bookworm blanket held view"), 2: ("Family reading lifestyle image with cozy room bookworm blanket.", "Cozy room bookworm blanket family lifestyle"), 3: ("Blanket size chart showing four sizes for Sophia Bookworm design.", "Cozy room bookworm blanket size chart"), 4: ("Feature panel over bookshelf artwork with fluffy and no-shedding callouts.", "Cozy room bookworm blanket feature panel"), 5: ("Multiple purpose panel with bed, sofa, office, travel and pet examples.", "Cozy room bookworm blanket multi-use panel"), 6: ("Cozy room bookworm blanket draped across sofa under a window.", "Cozy room bookworm blanket on sofa"), 7: ("Lifestyle reading scene with cozy room bookworm blanket and small pet.", "Cozy room bookworm blanket reading lifestyle"), 8: ("Reader wrapped in cozy room bookworm blanket by a window.", "Cozy room bookworm blanket wrapped reading view"), 9: ("Background option panel showing four room designs for the bookworm blanket.", "Cozy room bookworm blanket background options")},
    270: {1: ("Held book lover blanket with dark-haired girl, books and So Many Books So Little Time text.", "Dark hair book lover blanket held view"), 2: ("Book lover blanket draped on sofa with white fleece edge.", "Dark hair book lover blanket on sofa"), 3: ("Detail collage showing girl illustration, books, print texture and white fleece.", "Dark hair book lover blanket detail collage"), 4: ("Lifestyle reading scene with book lover blanket, open book and small pet.", "Dark hair book lover blanket reading lifestyle"), 5: ("Feature panel with machine washable, soft and warm, dense stitching and fleece callouts.", "Dark hair book lover blanket feature panel"), 6: ("Close-up of book lover blanket near an open book.", "Dark hair book lover blanket close-up"), 7: ("Size and use panel for book lover blanket with sofa, office, bed and travel icons.", "Dark hair book lover blanket size and use panel")},
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def parse_customizer(summary):
    raw = Path(summary["html"]["html_file"]).read_text(encoding="utf-8", errors="ignore")
    page = html.unescape(raw)
    fields = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"minLength":\s*(\d+),\s*"maxLength":\s*(\d+)',
        page,
    )
    if not fields:
        fields = re.findall(
            r'"label"\s*:\s*"([^"]+)"\s*,\s*"required"\s*:\s*(true|false)\s*,\s*"minLength"\s*:\s*(\d+)\s*,\s*"maxLength"\s*:\s*(\d+)',
            page,
        )
    options = re.findall(r'"label"\s*:\s*"([^"]+)"\s*,\s*"required"\s*:\s*(true|false)\s*,\s*"defaultOptionId"', page)
    return fields, options


def customizer_sentence(fields, options):
    parts = []
    for label, required, min_length, max_length in fields:
        requirement = "required" if required == "true" else "optional"
        parts.append(f"Personalization uses a {requirement} {label} field, {min_length}-{max_length} characters.")
    visible_options = [label for label, _ in options if label and "confirmation" not in label.lower() and "comfirmation" not in label.lower()]
    if visible_options:
        parts.append(f"Visible required choice fields include: {', '.join(visible_options)}.")
    if fields:
        parts.append("Visible custom text, photos or example selections in mockups are sample artwork unless the matching input is entered.")
    else:
        parts.append("No shopper text-entry field is described for this product.")
    return " ".join(parts)


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {"src": row["Image Src"], "alt": row.get("Image Alt Text", "")}
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    missing = sorted(set(handles) - set(admin))
    if missing:
        raise RuntimeError(f"Missing handles in admin export: {missing}")
    return admin


def option_text(admin_row):
    values = [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]]
    values = [v for v in values if v]
    return ", ".join(values) if values else "available size options"


def description(pos, admin_row, ctext):
    title, _, _, detail, _ = PRODUCT_UPDATES[pos]
    return (
        f"<p>{title} features {detail}. The copy stays specific to the visible artwork, product form and selectable options for this exact item.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork: {detail}.</li><li>Gallery images include the main mockup plus size, feature, care, fabric or lifestyle panels where shown.</li></ul>"
        "<h3>Options and Fit</h3>"
        f"<ul><li>Available option groups: {option_text(admin_row)}.</li><li>{ctext}</li><li>Select the product type and size shown on the product page before checkout.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    summaries = {int(item["inventory_row"]["inventory_position"]): item for item in json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))}
    positions = set(PRODUCT_UPDATES)
    handles = {summaries[pos]["inventory_row"]["Handle"] for pos in positions}
    admin = load_admin(handles)
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()
    wb = load_workbook(OUTPUT)

    pos_by_key = {summaries[pos]["inventory_row"]["product_key"]: pos for pos in positions}
    handle_by_pos = {pos: summaries[pos]["inventory_row"]["Handle"] for pos in positions}
    fields_by_pos, option_fields_by_pos = {}, {}
    for pos in positions:
        fields_by_pos[pos], option_fields_by_pos[pos] = parse_customizer(summaries[pos])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        handle = handle_by_pos[pos]
        title, primary, secondary, detail, cluster = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        meta = f"Shop {title.lower()} with {detail}, selectable sizes and verified custom options."[:155]
        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or summaries[pos]["inventory_row"]["product_type"]
        ws.cell(row_num, idx["title_proposed"]).value = title
        ws.cell(row_num, idx["meta_title_seo"]).value = title
        ws.cell(row_num, idx["meta_title_length"]).value = len(title)
        ws.cell(row_num, idx["meta_description_seo"]).value = meta
        ws.cell(row_num, idx["meta_description_length"]).value = len(meta)
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, customizer_sentence(fields_by_pos[pos], option_fields_by_pos[pos]))
        ws.cell(row_num, idx["primary_keyword"]).value = primary
        ws.cell(row_num, idx["secondary_keywords"]).value = secondary
        ws.cell(row_num, idx["meta_keyword"]).value = primary
        ws.cell(row_num, idx["keyword_strategy"]).value = f"Target the specific product motif: {cluster}; keep broad faith, Celtic, camping, couple gift or reader gift terms for collection pages."
        ws.cell(row_num, idx["buyer_search_summary"]).value = f"US English buyer intent targets {cluster}; no search-volume, trend or ranking claim is made."
        ws.cell(row_num, idx["keyword_evidence_level"]).value = "SERP_ONLY"
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R2 deep recheck without prior QA: used products_export_1.csv sha256:{admin_hash}; "
            "separated Christian, Celtic, desert RV, couple photo quilt and reader blanket intents; parsed custom fields and rewrote image observations/alts from contact sheets. Still NEEDS_REVIEW."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    seen_alts = set()
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pos = next((p for p, h in handle_by_pos.items() if h == handle), None)
        if pos is None:
            continue
        image_number = int(ws.cell(row_num, idx["image_number"]).value)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation, alt = IMAGE_DETAILS[pos][image_number]
        if alt in seen_alts:
            alt = f"{alt} for {PRODUCT_UPDATES[pos][4]}"
        seen_alts.add(alt)
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = alt[:125]
        ws.cell(row_num, idx["alt_action"]).value = "SET"
        ws.cell(row_num, idx["viewed_status"]).value = "VIEWED"
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = "R2: observation and alt rechecked from contact sheet; admin image URL and current alt matched from Shopify CSV where available."
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        title, primary, secondary, _, cluster = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in secondary.split(",")]
        keyword = primary if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = cluster
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "PRODUCT"
        ws.cell(row_num, idx["decision_reason"]).value = f"{title} is separated by verified product form, artwork motif and custom fields."
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = "products_export_1.csv admin baseline + storefront customizer parse + contact-sheet inspection"
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = "YES; nearby products are differentiated by motif, product form and custom field type."
        ws.cell(row_num, idx["mapping_reason"]).value = "R2 separates mixed faith, Celtic, camping, photo quilt and reading blanket pages by visible motif and option fields."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r2"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        title, primary, _, _, cluster = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {cluster} as a specific product page."
        ws.cell(row_num, idx["jtbd_statement"]).value = f"When shopping for {cluster}, the buyer wants exact artwork, product type, size choices and available custom fields confirmed before purchase."
        ws.cell(row_num, idx["functional_motivation"]).value = "Confirm product form, visible motif, size options, custom field rules, feature panels and image accuracy before purchase."
        ws.cell(row_num, idx["emotional_social_motivation"]).value = "Choose a personalized faith, Celtic, camping, couple or reader-themed bedding item that matches the recipient or room theme."
        ws.cell(row_num, idx["purchase_concerns"]).value = "Confirm custom text or photo rules; avoid unsupported material, religious promise, copyrighted identity or delivery claims."
        ws.cell(row_num, idx["source_refs"]).value = f"products_export_1.csv sha256:{admin_hash}; storefront HTML customizer parse; contact sheets 261-270"
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R2_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No independent QA, Search Console, paid keyword volume or internal site-search export supplied; proof remains SERP-only."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {primary} with title: {title}."

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in positions:
            continue
        handle = handle_by_pos[pos]
        admin_row = admin[handle]
        _, _, _, detail, _ = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = f"Original storefront/product evidence; products_export_1.csv sha256:{admin_hash}; storefront customizer parse; inspected contact sheets"
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={option_text(admin_row)}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={detail}; customizer={customizer_sentence(fields_by_pos[pos], option_fields_by_pos[pos])}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "No independent QA for batch 027 yet; r2 prepared for QA recheck."
        ws.cell(row_num, idx["confidence_and_reason"]).value = "R2 uses contact sheets, storefront customizer parse and Shopify admin CSV baseline; no approval or import created."

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_027_r2_revision", "r2", "Separate 10-product revision without prior independent QA; source workbook was not modified."),
        ("qa_batch_027_r2_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_027_r2_scope", "inventory positions 261-270", "No products outside qa_batch_027 were revised."),
        ("qa_batch_027_r2_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_027_r2_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_027_r2_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else "" for col in range(1, ws.max_column + 1)])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)
    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_evidence": str(SUMMARY_PATH.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_027 only; inventory positions 261-270",
        "revision": "r2",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Ready for independent QA of qa_batch_027_r2.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join([
            "# Revision qa_batch_027_r2",
            "",
            "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
            "",
            f"- Nguồn r1: `{SOURCE.relative_to(ROOT)}`",
            f"- Evidence summary: `{SUMMARY_PATH.relative_to(ROOT)}`",
            f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
            f"- Admin export SHA-256: `{admin_hash}`",
            f"- Workbook r2: `{OUTPUT.relative_to(ROOT)}`",
            f"- Sản phẩm sửa: {revised_products}",
            f"- Ảnh sửa/refresh: {revised_images}",
            f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
            "- Sửa trọng yếu: rà lại khi chưa có QA, tách nhóm Christian, Celtic/Yggdrasil, desert RV, couple photo quilt và book lover blanket; parse customizer từ HTML storefront; viết lại title/meta/description/alt theo từng contact sheet.",
            "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
            "- Bước tiếp theo: gửi `qa_batch_027_r2` cho QA độc lập.",
            "",
        ]),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
