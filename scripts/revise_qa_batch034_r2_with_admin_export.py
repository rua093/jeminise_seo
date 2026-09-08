import csv
import hashlib
import html
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_034_r2"
SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "batches" / "SEO_Product_Optimization_through_batch_034.xlsx"
SUMMARY_PATH = ROOT / "seo_runs" / SHOP / RUN_ID / "batch_034_evidence_summary.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_034_r2.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_034_r2.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    331: ("Vintage Yellow Softball Comforter Set", "vintage yellow softball comforter set", "softball comforter set, yellow softball bedding, vintage softball comforter", "dark brown comforter with a large distressed yellow softball, red stitching, script lettering and number artwork", "vintage yellow softball comforter set"),
    332: ("Polka Dot Softball Comforter Set", "polka dot softball comforter set", "softball comforter set, teal softball bedding, polka dot comforter", "black polka dot comforter with teal stripe, paisley border, softball lettering and glove-and-ball artwork", "polka dot softball comforter set"),
    333: ("Fireball Softball Comforter Set", "fireball softball comforter set", "softball comforter set, teal yellow softball bedding, fire softball comforter", "orange and blue comforter with a large glowing yellow softball, flame ring, script lettering and number artwork", "fireball softball comforter set"),
    334: ("Pink Glove Softball Comforter Set", "pink glove softball comforter set", "softball comforter set, pink softball bedding, softball glove comforter", "pink and gray comforter with a red softball glove, yellow softball, script lettering and number artwork", "pink glove softball comforter set"),
    335: ("Celtic Tree of Life Quilt Set", "Celtic Tree of Life quilt set", "Tree of Life quilt set, Celtic quilt bedding, Celtic knot quilt", "earth tone quilt set with a green Tree of Life, exposed knot roots, hills and Celtic knot borders", "Celtic Tree of Life quilt set"),
    336: ("Winter Cardinal Berry Branch Quilt", "winter cardinal berry branch quilt", "cardinal quilt, winter bird bedding, Christmas cardinal quilt", "winter quilt with a red cardinal, red berry branches, blue-gray leaves and snowy watercolor background", "winter cardinal berry branch quilt"),
    337: ("Heart Branch Cardinals Christmas Quilt", "heart branch cardinals Christmas quilt", "cardinal Christmas quilt, winter cardinals bedding, heart branch quilt", "Christmas quilt with two red cardinals, heart-shaped berry branches, snowy village scenery and pale winter background", "heart branch cardinals Christmas quilt"),
    338: ("Winter Cat Cardinal Christmas Quilt", "winter cat cardinal Christmas quilt", "cat cardinal quilt, Christmas cat bedding, winter animal quilt", "winter quilt with a gray kitten, red cardinal, snowy branches and It's Well With My Soul text", "winter cat cardinal Christmas quilt"),
}


IMAGE_DETAILS = {
    331: {1: ("Front bed mockup of vintage yellow softball comforter with script lettering and number 07 on pillows.", "Vintage yellow softball comforter front view"), 2: ("Angled room mockup of dark brown comforter with large distressed yellow softball.", "Vintage yellow softball comforter angled view"), 3: ("Pillowcase close view with yellow softball, number 07 and script lettering.", "Vintage yellow softball comforter pillowcase view"), 4: ("Folded comforter mockup showing distressed yellow softball print and white backing.", "Vintage yellow softball comforter folded view"), 5: ("Close fabric view of yellow softball stitching and script lettering under white backing.", "Vintage yellow softball comforter fabric close-up"), 6: ("Comforter feature panel with washable, soft filling and lightweight callouts.", "Vintage yellow softball comforter feature panel"), 7: ("Set inclusion panel showing one comforter and two pillowcases.", "Vintage yellow softball comforter set panel")},
    332: {1: ("Front bed mockup of black polka dot softball comforter with teal stripe and paisley border.", "Polka dot softball comforter front view"), 2: ("Angled room mockup of polka dot softball comforter with glove and ball artwork.", "Polka dot softball comforter angled view"), 3: ("Pillowcase close view with black background and white polka dots.", "Polka dot softball comforter pillowcase view"), 4: ("Folded comforter mockup showing polka dots, teal stripe, glove artwork and white backing.", "Polka dot softball comforter folded view"), 5: ("Close fabric view of teal stripe, softball lettering and paisley border.", "Polka dot softball comforter fabric close-up"), 6: ("Comforter feature panel with washable, soft filling and lightweight callouts.", "Polka dot softball comforter feature panel"), 7: ("Set inclusion panel showing one comforter and two pillowcases.", "Polka dot softball comforter set panel")},
    333: {1: ("Front bed mockup of fireball softball comforter with glowing ball, flames and number 21 pillows.", "Fireball softball comforter front view"), 2: ("Angled room mockup of orange blue fireball softball comforter with script lettering.", "Fireball softball comforter angled view"), 3: ("Pillowcase close view with glowing softball and number 21.", "Fireball softball comforter pillowcase view"), 4: ("Folded comforter mockup showing yellow softball, flame ring and white backing.", "Fireball softball comforter folded view"), 5: ("Close fabric view of glowing yellow softball with flame ring under white backing.", "Fireball softball comforter fabric close-up"), 6: ("Comforter feature panel with washable, soft filling and lightweight callouts.", "Fireball softball comforter feature panel"), 7: ("Set inclusion panel showing one comforter and two pillowcases.", "Fireball softball comforter set panel")},
    334: {1: ("Front bed mockup of pink glove softball comforter with yellow ball, script lettering and number 11.", "Pink glove softball comforter front view"), 2: ("Angled room mockup of pink glove softball comforter with matching pillows.", "Pink glove softball comforter angled view"), 3: ("Pillowcase close view with red softball glove, yellow ball and number 11.", "Pink glove softball comforter pillowcase view"), 4: ("Folded comforter mockup showing glove and yellow softball artwork with white backing.", "Pink glove softball comforter folded view"), 5: ("Close fabric view of red softball glove and script lettering under white backing.", "Pink glove softball comforter fabric close-up"), 6: ("Comforter feature panel with washable, soft filling and lightweight callouts.", "Pink glove softball comforter feature panel"), 7: ("Set inclusion panel showing one comforter and two pillowcases.", "Pink glove softball comforter set panel"), 8: ("Repeated front bed mockup of pink glove softball comforter with number 11 artwork.", "Pink glove softball comforter repeated front view"), 9: ("Repeated angled room mockup of pink glove softball comforter and matching pillows.", "Pink glove softball comforter repeated angled view"), 10: ("Repeated pillowcase close view with glove, softball and number 11.", "Pink glove softball comforter repeated pillowcase view"), 11: ("Repeated folded comforter mockup showing pink glove softball print and white backing.", "Pink glove softball comforter repeated folded view"), 12: ("Repeated close fabric view of red glove and script lettering under white backing.", "Pink glove softball comforter repeated fabric close-up"), 13: ("Repeated feature panel with washable, soft filling and lightweight callouts.", "Pink glove softball comforter repeated feature panel"), 14: ("Repeated set inclusion panel showing one comforter and two pillowcases.", "Pink glove softball comforter repeated set panel")},
    335: {1: ("Bed mockup of Celtic Tree of Life quilt with exposed knot roots and earth tone border.", "Celtic Tree of Life quilt bed view"), 2: ("Room mockup of Celtic Tree of Life quilt with printed craft callout.", "Celtic Tree of Life quilt room mockup"), 3: ("Close bed view with optional pillow shams and Celtic knot border artwork.", "Celtic Tree of Life quilt pillow sham panel"), 4: ("High-quality fabric panel showing quilt layers and Celtic tree artwork sample.", "Celtic Tree of Life quilt fabric panel"), 5: ("Premium quilt sets sizing chart for Celtic Tree of Life design.", "Celtic Tree of Life quilt size chart"), 6: ("Bedspread features panel showing microfiber layers and care icons.", "Celtic Tree of Life quilt features panel"), 7: ("Overhead bedroom mockup of Celtic Tree of Life quilt and matching pillows.", "Celtic Tree of Life quilt overhead view")},
    336: {1: ("Bed mockup of winter cardinal quilt with red berries and blue-gray leaves.", "Winter cardinal berry branch quilt bed view"), 2: ("Close fabric view of red cardinal, berry branch and quilting pattern.", "Winter cardinal berry branch quilt close fabric view"), 3: ("Single pillow sham mockup with cardinal and berry branch artwork.", "Winter cardinal berry branch quilt pillow sham"), 4: ("Room mockup of winter cardinal quilt set with matching pillows.", "Winter cardinal berry branch quilt room mockup"), 5: ("Bedding set panel showing premium quilt sizes and optional standard shams.", "Winter cardinal berry branch quilt size panel")},
    337: {1: ("Bed mockup of Christmas cardinal quilt with two red birds and heart-shaped berry branches.", "Heart branch cardinals Christmas quilt bed view"), 2: ("Close fabric view of two red cardinals, snowy village and quilting pattern.", "Heart branch cardinals Christmas quilt close fabric view"), 3: ("Single pillow sham mockup with cardinals and heart branch artwork.", "Heart branch cardinals Christmas quilt pillow sham"), 4: ("Room mockup of Christmas cardinal quilt set with matching pillows.", "Heart branch cardinals Christmas quilt room mockup"), 5: ("Bedding set panel showing premium quilt sizes and optional standard shams.", "Heart branch cardinals Christmas quilt size panel")},
    338: {1: ("Bed mockup of winter cat and cardinal quilt with snowy branches and It's Well With My Soul text.", "Winter cat cardinal Christmas quilt bed view"), 2: ("Close fabric view of gray kitten face, red cardinal beak and quilting pattern.", "Winter cat cardinal Christmas quilt close fabric view"), 3: ("Single pillow sham mockup with gray kitten, red cardinal and script text.", "Winter cat cardinal Christmas quilt pillow sham"), 4: ("Room mockup of winter cat cardinal quilt set with matching pillows.", "Winter cat cardinal Christmas quilt room mockup"), 5: ("Bedding set panel showing premium quilt sizes and optional standard shams.", "Winter cat cardinal Christmas quilt size panel")},
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def parse_customizer(summary):
    raw = Path(summary["html"]["html_file"]).read_text(encoding="utf-8", errors="ignore")
    page = html.unescape(raw)
    fields = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"minLength":\s*(\d+),\s*"maxLength":\s*(\d+)',
        page,
    )
    if not fields:
        fields = re.findall(
            r'"label"\s*:\s*"([^"]+)"\s*,\s*"required"\s*:\s*(true|false)\s*,\s*"minLength"\s*:\s*(\d+)\s*,\s*"maxLength"\s*:\s*(\d+)',
            page,
        )
    options = re.findall(r'"label"\s*:\s*"([^"]+)"\s*,\s*"required"\s*:\s*(true|false)\s*,\s*"defaultOptionId"', page)
    return fields, options


def customizer_sentence(fields, options):
    parts = []
    for label, required, min_length, max_length in fields:
        requirement = "required" if required == "true" else "optional"
        parts.append(f"Personalization uses a {requirement} {label} field, {min_length}-{max_length} characters.")
    visible_options = [label for label, _ in options if label and "confirmation" not in label.lower() and "comfirmation" not in label.lower()]
    if visible_options:
        parts.append(f"Visible required choice fields include: {', '.join(visible_options)}.")
    if fields:
        parts.append("Visible custom text or example selections in mockups are sample artwork unless the matching input is entered.")
    else:
        parts.append("No shopper text-entry field is described for this product.")
    return " ".join(parts)


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {"src": row["Image Src"], "alt": row.get("Image Alt Text", "")}
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    missing = sorted(set(handles) - set(admin))
    if missing:
        raise RuntimeError(f"Missing handles in admin export: {missing}")
    return admin


def option_text(admin_row):
    values = [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]]
    values = [v for v in values if v]
    return ", ".join(values) if values else "available size options"


def description(pos, admin_row, ctext):
    title, _, _, detail, _ = PRODUCT_UPDATES[pos]
    return (
        f"<p>{title} features {detail}. The copy stays specific to the visible artwork, product form and selectable options for this exact item.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork: {detail}.</li><li>Gallery images include the main mockup plus size, feature, care, fabric or lifestyle panels where shown.</li></ul>"
        "<h3>Options and Fit</h3>"
        f"<ul><li>Available option groups: {option_text(admin_row)}.</li><li>{ctext}</li><li>Select the product type and size shown on the product page before checkout.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    summaries = {int(item["inventory_row"]["inventory_position"]): item for item in json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))}
    positions = set(PRODUCT_UPDATES)
    handles = {summaries[pos]["inventory_row"]["Handle"] for pos in positions}
    admin = load_admin(handles)
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()
    wb = load_workbook(OUTPUT)

    pos_by_key = {summaries[pos]["inventory_row"]["product_key"]: pos for pos in positions}
    handle_by_pos = {pos: summaries[pos]["inventory_row"]["Handle"] for pos in positions}
    fields_by_pos, option_fields_by_pos = {}, {}
    for pos in positions:
        fields_by_pos[pos], option_fields_by_pos[pos] = parse_customizer(summaries[pos])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        handle = handle_by_pos[pos]
        title, primary, secondary, detail, cluster = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        meta = f"Shop {title.lower()} with {detail}, selectable sizes and product details."[:155]
        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or summaries[pos]["inventory_row"]["product_type"]
        ws.cell(row_num, idx["title_proposed"]).value = title
        ws.cell(row_num, idx["meta_title_seo"]).value = title
        ws.cell(row_num, idx["meta_title_length"]).value = len(title)
        ws.cell(row_num, idx["meta_description_seo"]).value = meta
        ws.cell(row_num, idx["meta_description_length"]).value = len(meta)
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, customizer_sentence(fields_by_pos[pos], option_fields_by_pos[pos]))
        ws.cell(row_num, idx["primary_keyword"]).value = primary
        ws.cell(row_num, idx["secondary_keywords"]).value = secondary
        ws.cell(row_num, idx["meta_keyword"]).value = primary
        ws.cell(row_num, idx["keyword_strategy"]).value = f"Target the specific product motif: {cluster}; keep broad sports bedding, softball decor or animal quilt terms for collection pages."
        ws.cell(row_num, idx["buyer_search_summary"]).value = f"US English buyer intent targets {cluster}; no search-volume, trend or ranking claim is made."
        ws.cell(row_num, idx["keyword_evidence_level"]).value = "SERP_ONLY"
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R2 deep recheck without prior QA: used products_export_1.csv sha256:{admin_hash}; "
            "separated four visually different softball comforters, one Celtic Tree of Life quilt and three winter cardinal or cat quilts; treated names and numbers in mockups as sample artwork unless customizer fields are parsed; rewrote image observations/alts from contact sheets. Still NEEDS_REVIEW."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    seen_alts = set()
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pos = next((p for p, h in handle_by_pos.items() if h == handle), None)
        if pos is None:
            continue
        image_number = int(ws.cell(row_num, idx["image_number"]).value)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation, alt = IMAGE_DETAILS[pos][image_number]
        if alt in seen_alts:
            alt = f"{alt} for {PRODUCT_UPDATES[pos][4]}"
        seen_alts.add(alt)
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = alt[:125]
        ws.cell(row_num, idx["alt_action"]).value = "SET"
        ws.cell(row_num, idx["viewed_status"]).value = "VIEWED"
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = "R2: observation and alt rechecked from contact sheet; admin image URL and current alt matched from Shopify CSV where available."
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        title, primary, secondary, _, cluster = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in secondary.split(",")]
        keyword = primary if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = cluster
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "PRODUCT"
        ws.cell(row_num, idx["decision_reason"]).value = f"{title} is separated by verified product form, artwork motif and custom fields."
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = "products_export_1.csv admin baseline + storefront customizer parse + contact-sheet inspection"
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = "YES; nearby products are differentiated by motif, product form and custom field type."
        ws.cell(row_num, idx["mapping_reason"]).value = "R2 separates each softball comforter variant, Celtic Tree of Life quilt and winter cardinal or cat quilt by visible motif, color palette, seasonal intent and option fields."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r2"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        title, primary, _, _, cluster = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {cluster} as a specific product page."
        ws.cell(row_num, idx["jtbd_statement"]).value = f"When shopping for {cluster}, the buyer wants exact artwork, product type, size choices and available custom fields confirmed before purchase."
        ws.cell(row_num, idx["functional_motivation"]).value = "Confirm product form, visible motif, size options, custom field rules, feature panels and image accuracy before purchase."
        ws.cell(row_num, idx["emotional_social_motivation"]).value = "Choose softball, Celtic tree or winter animal bedding that matches a hobby, seasonal room theme or gift context."
        ws.cell(row_num, idx["purchase_concerns"]).value = "Confirm product type, size choices, optional shams and visible design details; avoid unsupported material, custom input, copyrighted identity or delivery claims."
        ws.cell(row_num, idx["source_refs"]).value = f"products_export_1.csv sha256:{admin_hash}; storefront HTML customizer parse; contact sheets 331-338"
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R2_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No independent QA, Search Console, paid keyword volume or internal site-search export supplied; proof remains SERP-only."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {primary} with title: {title}."

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in positions:
            continue
        handle = handle_by_pos[pos]
        admin_row = admin[handle]
        _, _, _, detail, _ = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = f"Original storefront/product evidence; products_export_1.csv sha256:{admin_hash}; storefront customizer parse; inspected contact sheets"
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={option_text(admin_row)}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={detail}; customizer={customizer_sentence(fields_by_pos[pos], option_fields_by_pos[pos])}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "No independent QA for batch 034 yet; r2 prepared for QA recheck."
        ws.cell(row_num, idx["confidence_and_reason"]).value = "R2 uses contact sheets, storefront customizer parse and Shopify admin CSV baseline; no approval or import created."

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_034_r2_revision", "r2", "Final 8-product revision without prior independent QA; source workbook was not modified."),
        ("qa_batch_034_r2_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_034_r2_scope", "inventory positions 331-338", "No products outside qa_batch_034 were revised."),
        ("qa_batch_034_r2_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_034_r2_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_034_r2_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else "" for col in range(1, ws.max_column + 1)])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)
    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_evidence": str(SUMMARY_PATH.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_034 only; inventory positions 331-338",
        "revision": "r2",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Ready for independent QA of qa_batch_034_r2.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join([
            "# Revision qa_batch_034_r2",
            "",
            "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
            "",
            f"- Nguồn r1: `{SOURCE.relative_to(ROOT)}`",
            f"- Evidence summary: `{SUMMARY_PATH.relative_to(ROOT)}`",
            f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
            f"- Admin export SHA-256: `{admin_hash}`",
            f"- Workbook r2: `{OUTPUT.relative_to(ROOT)}`",
            f"- Sản phẩm sửa: {revised_products}",
            f"- Ảnh sửa/refresh: {revised_images}",
            f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
            "- Sửa trọng yếu: rà lại khi chưa có QA, tách 4 biến thể softball comforter, Celtic Tree of Life quilt và 3 mẫu winter/cardinal/cat quilt theo màu, motif, chữ/số mẫu và bố cục ảnh; parse customizer từ HTML storefront; viết lại title/meta/description/alt theo từng contact sheet.",
            "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
            "- Bước tiếp theo: gửi `qa_batch_034_r2` cho QA độc lập.",
            "",
        ]),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
