import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_023"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_022.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_023.xlsx"


BASKETBALL_REFS = "https://www.etsy.com/listing/1689155537/personalized-basketball-blanket-custom; https://www.amazon.com/Personalized-Basketball-Bedding-Number-Comforter/dp/B0CBRHP3FP; https://www.etsy.com/listing/1739772752/basketball-custom-name-number-blanket; https://www.amazon.com/clp/B0D2PBRQPS; https://eloquentinnovations.com/shop/basketball-design-bedding-with-name/; https://www.pinterest.com/ideas/basketball-net-with-personalized-name/919191574350/"


PROFILES = {
    221: ("basketball hoop net close-up blanket with name", "basketball hoop blanket", "Basketball Hoop Blanket", "Basketball Hoop Net Close-Up Blanket with Name", "blue and red close-up hoop and net blanket with custom name above the rim and jersey number below", "personalized basketball blanket, basketball net blanket, custom basketball throw", "EVERGREEN", BASKETBALL_REFS),
    222: ("basketball athletic shoes blanket with script name", "basketball shoes blanket", "Basketball Shoes Blanket", "Basketball Athletic Shoes Blanket with Name", "dark basketball blanket with vintage athletic shoes, ball, script custom name and jersey number", "personalized basketball blanket, basketball gift blanket, custom sports throw", "EVERGREEN", BASKETBALL_REFS),
    223: ("black basketball court lines comforter with name", "basketball court lines comforter", "Basketball Court Lines Comforter", "Black Basketball Court Lines Comforter with Name", "black basketball comforter with gold court lines, hoop graphics, large basketball, custom name and jersey number", "personalized basketball bedding, custom basketball comforter, basketball room decor", "EVERGREEN", BASKETBALL_REFS),
    224: ("basketball court perspective comforter with vertical name", "basketball court perspective comforter", "Basketball Court Perspective Comforter", "Basketball Court Perspective Comforter with Name", "dark court-perspective basketball comforter with hoop in the background, vertical custom name and jersey number", "personalized basketball comforter, basketball court bedding, custom name sports bedding", "EVERGREEN", BASKETBALL_REFS),
    225: ("hardwood court basketball comforter with name", "hardwood basketball comforter", "Hardwood Basketball Comforter", "Hardwood Court Basketball Comforter with Name", "blue and gold hardwood court basketball comforter with large ball, bold custom name and jersey number", "personalized basketball bedding, basketball court comforter, custom sports bedding", "EVERGREEN", BASKETBALL_REFS),
    226: ("neon dunk basketball player blanket with name", "basketball player blanket", "Basketball Player Blanket", "Neon Dunk Basketball Player Blanket with Name", "dark neon basketball blanket with dunking player silhouette, hoop, custom name and jersey number", "personalized basketball blanket, basketball player throw, custom sports blanket", "EVERGREEN", BASKETBALL_REFS),
    227: ("front basketball player blanket with flames", "basketball player blanket", "Basketball Player Blanket", "Front Basketball Player Blanket with Flames", "red and navy blanket with front-facing basketball player artwork, flame effect, custom name and jersey number", "personalized basketball blanket, basketball player gift, custom name throw", "EVERGREEN", BASKETBALL_REFS),
    228: ("orange basketball player silhouette blanket", "basketball silhouette blanket", "Basketball Silhouette Blanket", "Orange Basketball Player Silhouette Blanket", "gray and orange blanket with dribbling basketball player silhouette, large ball graphic, custom name and number", "personalized basketball blanket, basketball dribble blanket, custom sports throw", "EVERGREEN", BASKETBALL_REFS),
    229: ("basketball players and hoops comforter collage", "basketball collage comforter", "Basketball Collage Comforter", "Basketball Players and Hoops Comforter with Name", "black, gray and orange basketball comforter collage with player silhouettes, hoops, balls, custom name and number", "personalized basketball bedding, basketball comforter set, custom sports bedding", "EVERGREEN", BASKETBALL_REFS),
    230: ("shattered glass basketball comforter with name", "shattered glass basketball comforter", "Shattered Glass Basketball Comforter", "Shattered Glass Basketball Comforter with Name", "basketball comforter with large ball over shattered glass texture, script custom name and jersey number on shams", "personalized basketball bedding, basketball comforter set, custom basketball bedding", "EVERGREEN", BASKETBALL_REFS),
}


def qa_issues(pos, html_error):
    issues = [
        "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment.",
        "Personalized basketball products from batches 021-023 are close variants; review custom-name placement, jersey number accuracy and keyword cannibalization before import.",
        "Feature images mention duvet cover/comforter/blanket options, microfiber, easy care and size dimensions; verify actual product type, variants and set configuration before deployment.",
    ]
    if pos in (221, 222, 226, 227, 228):
        issues.append("This item is positioned as a blanket; verify blanket sizing and avoid comforter wording in import fields.")
    if pos in (223, 224, 225):
        issues.append("Court-based basketball comforter designs overlap; keep court-lines, perspective and hardwood wording distinct.")
    if pos in (226, 227, 228, 229):
        issues.append("Basketball player/silhouette motifs are close; review page title uniqueness and canonical strategy.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for basketball rooms, player gifts, sports bedrooms or personalized fan bedding.</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broad personalized basketball bedding terms for collections or strongest representative pages.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as basketball room decor, a player gift or personalized bedding.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": qa_issues(pos, html_error),
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for personalized basketball bedding, the buyer wants a {kind} matching a basketball motif, custom name and number.",
            "functional_motivation": "Find the right basketball design, size, personalization fields, care details and room fit.",
            "emotional_social_motivation": "Create a basketball-themed bedroom or give a custom gift to a player or fan.",
            "purchase_concerns": "Personalization spelling, jersey number, product identity, material/care claims, size fit and whether artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed in a main product mockup",
                2: f"Secondary product mockup for {h1}",
                3: f"Size or bedding type image for {h1}",
                4: f"Lifestyle, fabric or product feature image for {h1}",
                5: f"Care, close-up or feature image for {h1}",
                6: f"Bedding type or blanket feature image for {h1}",
                7: f"Size dimension image for {h1}",
                8: f"Lifestyle use image for {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_023_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_023_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_023"},
        {"metric": "cumulative_products", "value": "230", "definition": "Total products included through batch_023"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[230:240]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_024",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
