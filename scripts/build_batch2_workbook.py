import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_002"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
RESULTS_DIR = ROOT / "resutls" / SHOP / RUN_ID
BATCHES_DIR = RESULTS_DIR / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_001.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_002.xlsx"


PROFILES = {
    11: {
        "theme": "orange God Says I Am Christian design with the name Jessica, roses and butterflies",
        "visual": "orange and white personalized Christian bedding with God Says I Am text, the name Jessica, rose flowers, butterflies, Bible verse affirmation blocks and product information images",
        "primary": "personalized God Says I Am comforter set",
        "secondary": "Christian name bedding, God Says I Am bedding, personalized Christian bedding for women",
        "season": "EVERGREEN",
        "evidence_level": "SERP_ONLY",
        "seo_title": "Personalized God Says I Am Comforter Set",
        "h1": "Personalized God Says I Am Christian Comforter Set",
        "meta": "Create a faith-filled bedroom with personalized God Says I Am bedding featuring name text, roses, butterflies and Bible verse affirmations.",
        "buyer": "Customer may be shopping for a personalized Christian bedding gift with name text and Scripture-based affirmations.",
    },
    12: {
        "theme": "white personalized Christian affirmation design with the name Evelyn and Bible verse words",
        "visual": "white Christian bedding with the name Evelyn, cross icons, dove and fish symbols, many Bible affirmation words and product information images",
        "primary": "personalized Bible verse comforter set",
        "secondary": "Christian affirmation bedding, custom name Christian comforter, God says I am bedding",
        "season": "EVERGREEN",
        "evidence_level": "SERP_ONLY",
        "seo_title": "Personalized Bible Verse Comforter Set",
        "h1": "Personalized Christian Bible Verse Comforter Set",
        "meta": "Personalize a Christian bedroom with a white Bible verse comforter set featuring name text, cross symbols and affirmation words.",
        "buyer": "Customer may want a custom Christian comforter that makes a name and faith affirmations central to the design.",
    },
    13: {
        "theme": "black and gold Christian design with the name Olivia, butterflies and God is within her text",
        "visual": "black, brown and gold Christian bedding with the name Olivia, butterflies, floral vine artwork, God is within her she will not fall text and affirmation blocks",
        "primary": "personalized Christian comforter for women",
        "secondary": "God is within her bedding, Christian bedding for women, custom name faith comforter",
        "season": "EVERGREEN",
        "evidence_level": "SERP_ONLY",
        "seo_title": "Personalized Christian Comforter for Women",
        "h1": "Personalized God Is Within Her Christian Comforter",
        "meta": "A black and gold Christian comforter design with personalized name text, butterflies and God Is Within Her Scripture-inspired artwork.",
        "buyer": "Customer may be looking for a faith gift for a woman or girl with personalized name text and Psalm-inspired encouragement.",
    },
    14: {
        "theme": "blue God Says I Am Christian design with the name Emily and floral border",
        "visual": "white and blue personalized Christian bedding with the name Emily, God Says I Am text, blue floral border and affirmation blocks",
        "primary": "blue personalized Christian comforter set",
        "secondary": "God Says I Am comforter, Christian name bedding set, Bible verse comforter for girls",
        "season": "EVERGREEN",
        "evidence_level": "SERP_ONLY",
        "seo_title": "Blue Personalized Christian Comforter Set",
        "h1": "Blue God Says I Am Christian Comforter Set",
        "meta": "Personalized Christian bedding in blue and white with God Says I Am text, name artwork and Bible verse affirmation blocks.",
        "buyer": "Customer may prefer a cleaner blue Christian bedding design for a daughter, teen, woman or faith-based gift.",
    },
    15: {
        "theme": "brown Christian butterfly design with the name Charlotte and God is within her text",
        "visual": "brown personalized Christian bedding with the name Charlotte, butterfly artwork, God is within her she will not fall text and vertical affirmation blocks",
        "primary": "personalized Christian butterfly comforter",
        "secondary": "God is within her comforter, Christian butterfly bedding, custom name Christian bedding",
        "season": "EVERGREEN",
        "evidence_level": "SERP_ONLY",
        "seo_title": "Personalized Christian Butterfly Comforter",
        "h1": "Personalized Christian Butterfly Comforter Set",
        "meta": "A brown Christian comforter set with personalized name text, butterfly artwork and God Is Within Her affirmation design.",
        "buyer": "Customer may want a custom faith bedding gift with warm colors, butterflies and Scripture-inspired encouragement.",
    },
    16: {
        "theme": "black Christian warrior lion design with the name David and large white cross",
        "visual": "black personalized Christian comforter with the name David, large white cross, lion and armored warrior artwork, and text about being chosen and made for a purpose",
        "primary": "personalized Christian comforter for men",
        "secondary": "Christian warrior bedding, lion cross comforter, custom name Bible verse comforter",
        "season": "EVERGREEN",
        "evidence_level": "SERP_ONLY",
        "seo_title": "Personalized Christian Comforter for Men",
        "h1": "Personalized Christian Warrior Comforter Set",
        "meta": "Bold black Christian bedding with personalized name text, a white cross, lion artwork and faith affirmations for a purpose-driven design.",
        "buyer": "Customer may be shopping for masculine Christian bedding or a custom faith gift with warrior, lion and cross imagery.",
    },
    17: {
        "theme": "Christmas tree patchwork quilt in red, navy, green and cream squares",
        "visual": "Christmas quilt set with patchwork squares, tree silhouettes, red, navy, olive and cream fabric-look blocks, matching sham, close-up texture and size image",
        "primary": "Christmas tree patchwork quilt set",
        "secondary": "Christmas tree quilt set, patchwork Christmas bedding, holiday quilt with shams",
        "season": "CHRISTMAS",
        "evidence_level": "SERP_ONLY",
        "seo_title": "Christmas Tree Patchwork Quilt Set",
        "h1": "Christmas Tree Patchwork Quilt Set",
        "meta": "Decorate for the holidays with a Christmas tree patchwork quilt set in red, navy, green and cream with matching sham artwork.",
        "buyer": "Customer may be looking for classic holiday bedding with Christmas tree patchwork and coordinated shams.",
    },
    18: {
        "theme": "Santa, snowman, reindeer and Christmas tree patchwork quilt",
        "visual": "Christmas quilt set with Santa faces, snowmen, reindeer, green trees, snowflakes, red and green patchwork blocks, matching sham and size image",
        "primary": "Santa snowman Christmas quilt set",
        "secondary": "Christmas character quilt, Santa bedding set, snowman holiday quilt",
        "season": "CHRISTMAS",
        "evidence_level": "SERP_ONLY",
        "seo_title": "Santa Snowman Christmas Quilt Set",
        "h1": "Santa and Snowman Christmas Quilt Set",
        "meta": "Make a festive bedroom with a Christmas quilt set featuring Santa, snowmen, reindeer, trees and red-green patchwork blocks.",
        "buyer": "Customer may want playful holiday bedding with recognizable Christmas characters for a family or guest room.",
    },
    19: {
        "theme": "gingerbread man Christmas quilt with candy cane and holly on a red background",
        "visual": "red Christmas quilt set with a large gingerbread man, candy cane, holly leaves, peppermint details, matching pillow sham and size image",
        "primary": "gingerbread Christmas quilt set",
        "secondary": "gingerbread bedding, Christmas gingerbread quilt, red holiday quilt set",
        "season": "CHRISTMAS",
        "evidence_level": "SERP_ONLY",
        "seo_title": "Gingerbread Christmas Quilt Set",
        "h1": "Gingerbread Christmas Quilt Set",
        "meta": "Add sweet holiday style with a red gingerbread Christmas quilt set featuring candy cane, holly and matching pillow sham artwork.",
        "buyer": "Customer may be shopping for whimsical Christmas bedding centered on gingerbread and candy-cane decor.",
    },
    20: {
        "theme": "two gingerbread figures with Christmas gifts, candy cane and holly",
        "visual": "light Christmas quilt set with two gingerbread figures, wrapped gifts, candy cane, holly greenery, red bows, matching pillow sham and size image",
        "primary": "gingerbread Christmas bedding set",
        "secondary": "gingerbread quilt set, Christmas gift quilt, holiday gingerbread bedding",
        "season": "CHRISTMAS",
        "evidence_level": "SERP_ONLY",
        "seo_title": "Gingerbread Christmas Bedding Set",
        "h1": "Gingerbread Gift Christmas Quilt Set",
        "meta": "Brighten holiday decor with a gingerbread Christmas quilt set featuring two cookie figures, wrapped gifts, candy cane and holly artwork.",
        "buyer": "Customer may want cute Christmas bedding with gingerbread characters and gift imagery for seasonal bedroom decor.",
    },
}

SEARCH_REFS = {
    "christian": "https://www.etsy.com/market/god_says_i_am_blanket; https://www.amazon.com/clp/B0F1Y2J2FD; https://www.etsy.com/listing/4410818176/christian-name-blanket-god-says-i-am",
    "christmas": "https://www.amazon.com/Christmas-Quilt-Sets/s?k=Christmas+Quilt+Sets; https://www.bedbathandbeyond.com/c/christmas-bedding/christmas-bedding-sets?t=28287; https://www.homedepot.com/p/MarCielo-BY218-3-Pieces-Red-Patchwork-Christmas-Tree-Queen-Size-Polyester-Quilt-Set-Reversible-Holiday-Bedspread-BY218-Q/337970050",
    "gingerbread": "https://www.amazon.com/gingerbread-christmas-quilt-set/s?k=gingerbread+christmas+quilt+set; https://www.target.com/s/gingerbread%2Bbedding; https://www.greenlandhomefashions.com/product/gingerbread-lane-quilt-set/",
}


def append_dict(ws, data):
    headers = [cell.value for cell in ws[1]]
    ws.append([data.get(h, "") for h in headers])


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    image_count = 0
    keyword_count = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        profile = PROFILES[pos]
        product_key = row["product_key"]
        local_idx = pos
        evidence_id = f"evidence_{BATCH_ID}_{local_idx:03d}"
        research_id = f"research_{BATCH_ID}_{local_idx:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{local_idx:03d}_{int(img['position']):02d}" for img in images)
        ref_key = "gingerbread" if pos in (19, 20) else ("christmas" if pos in (17, 18) else "christian")
        refs = SEARCH_REFS[ref_key]
        description_html = (
            f"<p>Bring {profile['theme']} into the bedroom with this Jeminise {kind}. "
            f"The coordinated artwork appears across the main bedding piece and matching sham images.</p>"
            f"<h3>Design Details</h3><ul><li>{profile['visual'].capitalize()}.</li>"
            f"<li>Available product details from the source include: {facts[:650]}.</li></ul>"
            f"<h3>Shopping Context</h3><p>{profile['buyer']} Review this draft against the live product page and Shopify export before approval.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": profile["primary"], "secondary_keywords": profile["secondary"],
            "keyword_evidence_level": profile["evidence_level"],
            "keyword_strategy": f"Use a product-level query tied to {profile['theme']}; avoid collapsing the series into one generic Christian or Christmas term.",
            "season": profile["season"], "buyer_search_summary": profile["buyer"],
            "title_action": "SET", "title_proposed": profile["h1"],
            "meta_title_action": "SET", "meta_title_seo": profile["seo_title"], "meta_title_length": len(profile["seo_title"]),
            "meta_description_action": "SET", "meta_description_seo": profile["meta"], "meta_description_length": len(profile["meta"]),
            "description_action": "SET", "description_proposed_html": description_html,
            "meta_keyword": profile["primary"], "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": "DRAFTED", "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": "Admin export not provided; Shopify stored SEO fields, current image alt text and exact admin media mapping should be verified before deployment.",
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images),
            "image_audit_references": image_refs, "SERP_evidence_references": refs,
            "buyer_research_references": research_id, "processing_status": "DRAFTED",
            "confidence_and_reason": "Medium: page JSON, HTML and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": profile["buyer"],
            "jtbd_statement": f"When shopping for {profile['theme']} bedding, the buyer wants a coordinated {kind} that matches a faith, name-personalized or holiday room theme.",
            "functional_motivation": f"Find a {kind} with the right design, size options and included bedding components.",
            "emotional_social_motivation": "Give a personalized faith gift or create seasonal holiday bedroom decor.",
            "purchase_concerns": "Fit, included components, material/care claims, personalization accuracy and design match.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": profile["evidence_level"],
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; source refs show public marketplace/SERP comparables, not volume.",
            "seo_application": f"Use {profile['primary']} as candidate primary keyword for this product page.",
        })

        keywords = [(profile["primary"], "PRIMARY")] + [(kw.strip(), "SECONDARY") for kw in profile["secondary"].split(",")]
        for keyword, role in keywords:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": profile["theme"],
                "intent": "purchase", "target_page_type": "PRODUCT", "target_url": row["product_url"],
                "keyword_role": role, "decision_reason": "Specific to visible artwork, personalization or holiday product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": profile["evidence_level"],
                "season": profile["season"], "research_period": "2026-09-06", "validation_source": refs,
                "checked_at": now, "representative_SERP_URLs": refs,
                "possible_overlap_with_other_products": "YES", "mapping_reason": "Product page query includes attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })
            keyword_count += 1

        for image in images:
            ipos = int(image["position"])
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now,
                "observed_visual_details": base.image_observation(profile, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": base.alt_for(profile, ipos, kind), "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            image_count += 1

    for row in [
        {"metric": "batch_002_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_002_images", "value": str(image_count), "definition": "Gallery images downloaded and viewed for batch_002"},
        {"metric": "cumulative_products", "value": "20", "definition": "Total products included through batch_002"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    import csv
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[20:30]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_003",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={image_count} appended_keywords={keyword_count}")


if __name__ == "__main__":
    main()
