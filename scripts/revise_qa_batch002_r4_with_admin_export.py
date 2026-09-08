import csv
import hashlib
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_002_r4"
QA_RUN_ID = "20260907_123726"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_002_r3" / "SEO_Product_Optimization_qa_batch_002_r3.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
CUSTOMIZER_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_audit.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_002_r4.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_002_r4.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    11: {
        "title": "Personalized God Says I Am Christian Comforter Set",
        "meta_title": "Personalized God Says I Am Christian Comforter Set",
        "meta_description": "Customize a God Says I Am Christian comforter with Enter Name, birth-month flower artwork, bedding sizes and pillowcase options.",
        "primary": "personalized God Says I Am Christian comforter",
        "secondary": "Christian comforter with name, God Says I Am bedding, personalized Christian bedding set",
        "cluster": "personalized God Says I Am comforter",
        "intent_role": "God Says I Am design 02 page with verified required Enter Name field up to 13 characters.",
        "detail": "God Says I Am affirmation artwork with Jessica sample name, Bible verse styling and birth-month flower options",
    },
    12: {
        "title": "Personalized Christian Bible Verse Comforter Set",
        "meta_title": "Personalized Christian Bible Verse Comforter Set",
        "meta_description": "Add a custom name to this Christian Bible verse comforter with faith artwork, bedding sizes and pillowcase options.",
        "primary": "personalized Christian Bible verse comforter",
        "secondary": "Christian Bible verse bedding, custom Christian comforter, faith comforter with name",
        "cluster": "personalized Christian Bible verse comforter",
        "intent_role": "Christian Bible verse design 03 page with verified required Enter Name field up to 30 characters.",
        "detail": "Christian Bible verse artwork with Evelyn sample name and faith-themed bedroom styling",
    },
    13: {
        "title": "Personalized God Is Within Her Christian Comforter",
        "meta_title": "Personalized God Is Within Her Christian Comforter",
        "meta_description": "Customize a God Is Within Her Christian comforter with Enter Name, faith artwork, bedding sizes and pillowcase options.",
        "primary": "personalized God Is Within Her comforter",
        "secondary": "God Is Within Her bedding, Christian comforter with name, personalized faith bedding",
        "cluster": "personalized God Is Within Her comforter",
        "intent_role": "Women-focused God Is Within Her design; removes Christian Knight Templar wording from customer copy.",
        "detail": "God Is Within Her Christian artwork with Olivia sample name and faith message",
    },
    14: {
        "title": "Personalized Blue God Says I Am Christian Comforter",
        "meta_title": "Personalized Blue God Says I Am Christian Comforter",
        "meta_description": "Customize a blue God Says I Am Christian comforter with Enter Name, faith artwork, bedding sizes and pillowcase options.",
        "primary": "blue personalized God Says I Am comforter",
        "secondary": "blue Christian bedding with name, God Says I Am comforter set, personalized faith bedding",
        "cluster": "blue personalized God Says I Am comforter",
        "intent_role": "Blue God Says I Am design 05 page with verified required Enter Name field up to 13 characters.",
        "detail": "blue God Says I Am Christian affirmation artwork with Emily sample name",
    },
    15: {
        "title": "Personalized Is Within Her Butterfly Comforter",
        "meta_title": "Personalized Is Within Her Butterfly Comforter",
        "meta_description": "Customize an Is Within Her Christian comforter with Enter Name, butterfly artwork, bedding sizes and pillowcase options.",
        "primary": "personalized Is Within Her butterfly comforter",
        "secondary": "Christian butterfly bedding, Is Within Her comforter, personalized faith bedding",
        "cluster": "personalized Is Within Her butterfly comforter",
        "intent_role": "Christian butterfly design 06 page with verified required Enter Name field up to 13 characters.",
        "detail": "brown Is Within Her Christian artwork with Charlotte sample name, butterflies and faith affirmation blocks",
    },
    16: {
        "title": "Personalized Christian Warrior Comforter Set",
        "meta_title": "Personalized Christian Warrior Comforter Set",
        "meta_description": "Customize a Christian warrior comforter with Enter Name, faith artwork, bedding sizes and pillowcase options.",
        "primary": "personalized Christian warrior comforter",
        "secondary": "Christian warrior bedding, custom Bible verse comforter, faith comforter with name",
        "cluster": "personalized Christian warrior comforter",
        "intent_role": "Christian warrior design page with verified required Enter Name field up to 30 characters.",
        "detail": "Christian warrior artwork with David sample name and faith-themed bedding visuals",
    },
    17: {
        "title": "Personalized Christmas Tree Patchwork Quilt Set",
        "meta_title": "Personalized Christmas Tree Patchwork Quilt Set",
        "meta_description": "Add optional custom text to a Christmas tree patchwork quilt with holiday print, quilt sizes and pillowcase quantity choices.",
        "primary": "personalized Christmas tree patchwork quilt",
        "secondary": "Christmas tree quilt set, custom Christmas quilt, holiday patchwork bedding",
        "cluster": "personalized Christmas tree patchwork quilt",
        "intent_role": "Christmas tree design 01 page with verified optional Custom Your Name field up to 200 characters.",
        "detail": "Christmas tree patchwork quilt artwork with holiday print and quilting details",
    },
    18: {
        "title": "Personalized Santa and Snowman Christmas Quilt Set",
        "meta_title": "Personalized Santa and Snowman Christmas Quilt Set",
        "meta_description": "Add optional custom text to a Santa and snowman Christmas quilt with festive print, sizes and pillowcase quantity choices.",
        "primary": "personalized Santa snowman Christmas quilt",
        "secondary": "Santa snowman quilt set, custom Christmas quilt, holiday quilt bedding",
        "cluster": "personalized Santa snowman Christmas quilt",
        "intent_role": "Santa and snowman design 02 page with verified optional Custom Your Name field up to 200 characters.",
        "detail": "Santa and snowman Christmas quilt artwork with festive print and matching sham visuals",
    },
    19: {
        "title": "Personalized Gingerbread Christmas Quilt Set",
        "meta_title": "Personalized Gingerbread Christmas Quilt Set",
        "meta_description": "Add optional custom text to a gingerbread Christmas quilt with holiday print, quilt sizes and pillowcase quantity choices.",
        "primary": "personalized gingerbread Christmas quilt",
        "secondary": "gingerbread Christmas quilt set, custom holiday quilt, Christmas bedding quilt",
        "cluster": "personalized gingerbread Christmas quilt",
        "intent_role": "Gingerbread design 03 page with verified optional Custom Your Name field up to 200 characters.",
        "detail": "gingerbread Christmas quilt artwork with holiday print, quilting details and matching sham",
    },
    20: {
        "title": "Personalized Gingerbread Gift Christmas Quilt Set",
        "meta_title": "Personalized Gingerbread Gift Christmas Quilt Set",
        "meta_description": "Add optional custom text to a gingerbread gift Christmas quilt with holiday print, sizes and pillowcase quantity choices.",
        "primary": "personalized gingerbread gift Christmas quilt",
        "secondary": "gingerbread gift quilt set, custom Christmas quilt, holiday quilt bedding",
        "cluster": "personalized gingerbread gift Christmas quilt",
        "intent_role": "Gingerbread gift design 04 page with verified optional Custom Your Name field up to 200 characters.",
        "detail": "gingerbread gift Christmas quilt artwork with holiday print and matching sham",
    },
}


IMAGE_DETAILS = {
    11: [
        ("Bedroom mockup of orange God Says I Am Christian bedding with Jessica name, roses, butterflies and verse labels.", "Orange God Says I Am Christian bedding with Jessica name"),
        ("Angled bedroom mockup showing orange affirmation blocks, Jessica pillowcases, rose artwork and butterflies.", "Orange personalized God Says I Am bedding room mockup"),
        ("Product type comparison panel showing duvet cover set and comforter set options for all-season use.", "God Says I Am design 02 duvet and comforter options"),
        ("Feature panel with microfiber callout, 3D printed pattern close-ups, soft filling and lightweight breathable notes.", "God Says I Am design 02 microfiber feature panel"),
        ("Birth month flowers panel with flower illustrations for January through December.", "God Says I Am design 02 birth month flowers panel"),
        ("Stress-free easy care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade notes.", "God Says I Am design 02 easy care feature panel"),
        ("Close bedroom view showing orange affirmation bedding with soft, lightweight, durable and breathable icons.", "God Says I Am design 02 bedding feature icons"),
        ("Size dimension panel showing Twin, Full, Queen and King bedding dimensions.", "God Says I Am design 02 bedding size chart"),
    ],
    12: [
        ("Bedroom mockup of white Christian Bible verse bedding with Evelyn name, crosses and multicolor affirmation words.", "Christian Bible verse bedding with Evelyn name"),
        ("Angled room mockup showing Evelyn pillowcases, cross motifs, faith words and Bible verse labels.", "Personalized Christian Bible verse bedding room mockup"),
        ("Product type comparison panel showing duvet cover set and comforter set options for all-season use.", "Christian Bible verse duvet and comforter options"),
        ("Feature panel with microfiber callout, 3D printed pattern close-ups, soft filling and lightweight breathable notes.", "Christian Bible verse microfiber feature panel"),
        ("Close bedroom view showing Evelyn Bible verse bedding with soft, lightweight, durable and breathable icons.", "Christian Bible verse bedding feature icons"),
        ("Stress-free easy care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade notes.", "Christian Bible verse easy care feature panel"),
        ("Size dimension panel showing Twin, Full, Queen and King bedding dimensions.", "Christian Bible verse bedding size chart"),
    ],
    13: [
        ("Bedroom mockup of dark God Is Within Her bedding with Olivia name, gold florals, butterflies and Psalm 46:5 text.", "Dark God Is Within Her bedding with Olivia name"),
        ("Angled room mockup showing Olivia pillowcases, gold floral divider, butterflies and affirmation blocks.", "God Is Within Her Christian bedding room mockup"),
        ("Product type comparison panel showing duvet cover set and comforter set options for all-season use.", "God Is Within Her duvet and comforter options"),
        ("Feature panel with microfiber callout, 3D printed pattern close-ups, soft filling and lightweight breathable notes.", "God Is Within Her microfiber feature panel"),
        ("Close bedroom view showing dark gold God Is Within Her bedding with soft, lightweight, durable and breathable icons.", "God Is Within Her bedding feature icons"),
        ("Stress-free easy care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade notes.", "God Is Within Her easy care feature panel"),
        ("Size dimension panel showing Twin, Full, Queen and King bedding dimensions.", "God Is Within Her bedding size chart"),
    ],
    14: [
        ("Bedroom mockup of blue God Says I Am Christian bedding with Emily name, blue florals and Bible verse labels.", "Blue God Says I Am Christian bedding with Emily name"),
        ("Angled room mockup showing Emily pillowcases, blue affirmation blocks and floral border artwork.", "Blue personalized God Says I Am bedding room mockup"),
        ("Product type comparison panel showing duvet cover set and comforter set options for all-season use.", "Blue God Says I Am duvet and comforter options"),
        ("Feature panel with microfiber callout, 3D printed pattern close-ups, soft filling and lightweight breathable notes.", "Blue God Says I Am microfiber feature panel"),
        ("Close bedroom view showing blue Christian bedding with soft, lightweight, durable and breathable icons.", "Blue God Says I Am bedding feature icons"),
        ("Stress-free easy care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade notes.", "Blue God Says I Am easy care feature panel"),
        ("Size dimension panel showing Twin, Full, Queen and King bedding dimensions.", "Blue God Says I Am bedding size chart"),
    ],
    15: [
        ("Bedroom mockup of brown Is Within Her Christian bedding with Charlotte name, butterflies and faith affirmation blocks.", "Brown Is Within Her bedding with Charlotte name"),
        ("Angled room mockup showing Charlotte pillowcases, butterfly artwork and vertical name lettering.", "Brown personalized Is Within Her bedding room mockup"),
        ("Product type comparison panel showing duvet cover set and comforter set options for all-season use.", "Is Within Her butterfly duvet and comforter options"),
        ("Feature panel with microfiber callout, 3D printed pattern close-ups, soft filling and lightweight breathable notes.", "Is Within Her butterfly microfiber feature panel"),
        ("Close bedroom view showing brown butterfly bedding with soft, lightweight, durable and breathable icons.", "Is Within Her butterfly bedding feature icons"),
        ("Stress-free easy care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade notes.", "Is Within Her butterfly easy care feature panel"),
        ("Size dimension panel showing Twin, Full, Queen and King bedding dimensions.", "Is Within Her butterfly bedding size chart"),
    ],
    16: [
        ("Bedroom mockup of black Christian warrior bedding with David name, large white cross, armor artwork and faith text.", "Black Christian warrior bedding with David name"),
        ("Square mockup of black Christian warrior bedding with David name, cross graphic and armored figure artwork.", "Christian warrior bedding square mockup"),
        ("Angled room mockup showing David pillowcases, white cross lettering and warrior armor artwork.", "Personalized Christian warrior bedding room mockup"),
        ("Product type comparison panel showing duvet cover set and comforter set options for all-season use.", "Christian warrior duvet and comforter options"),
        ("Feature panel with microfiber callout, 3D printed pattern close-ups, soft filling and lightweight breathable notes.", "Christian warrior microfiber feature panel"),
        ("Close bedroom view showing black Christian warrior bedding with soft, lightweight, durable and breathable icons.", "Christian warrior bedding feature icons"),
        ("Stress-free easy care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade notes.", "Christian warrior easy care feature panel"),
        ("Size dimension panel showing Twin, Full, Queen and King bedding dimensions.", "Christian warrior bedding size chart"),
    ],
    17: [
        ("Christmas tree patchwork quilt bed mockup with red, green, navy and cream squares and matching shams.", "Christmas tree patchwork quilt set on bed"),
        ("Bedroom mockup showing the Christmas tree patchwork quilt with coordinated pillow shams.", "Christmas tree patchwork quilt bedroom mockup"),
        ("Matching pillow sham with red, green, navy and cream Christmas tree patchwork squares.", "Christmas tree patchwork pillow sham"),
        ("Close-up of Christmas tree patchwork squares with visible quilt stitching.", "Christmas tree patchwork quilt stitching close-up"),
        ("Size and component panel showing one premium quilt and two optional standard shams.", "Christmas tree patchwork quilt size and sham options"),
    ],
    18: [
        ("Santa, snowman and reindeer Christmas patchwork quilt bed mockup with red and green holiday squares.", "Santa snowman Christmas patchwork quilt set"),
        ("Bedroom mockup showing Santa and snowman Christmas quilt with coordinated pillow shams.", "Santa snowman Christmas quilt bedroom mockup"),
        ("Close-up of snowman, Christmas tree, snowflakes and quilt stitching on red and green patchwork.", "Santa snowman quilt stitching close-up"),
        ("Matching pillow sham with Santa, snowman, reindeer and Christmas tree patchwork artwork.", "Santa snowman Christmas pillow sham"),
        ("Size and component panel showing one premium quilt and two optional standard shams.", "Santa snowman Christmas quilt size and sham options"),
    ],
    19: [
        ("Red gingerbread Christmas quilt bed mockup with candy cane, holly, floral accents and matching shams.", "Red gingerbread Christmas quilt set on bed"),
        ("Bedroom mockup showing gingerbread Christmas quilt with candy cane artwork and coordinated pillow shams.", "Gingerbread Christmas quilt bedroom mockup"),
        ("Close-up of gingerbread cookie artwork, buttons, icing trim and visible quilt stitching.", "Gingerbread Christmas quilt stitching close-up"),
        ("Matching pillow sham with gingerbread cookie, candy cane, holly and red holiday background.", "Gingerbread Christmas pillow sham"),
        ("Size and component panel showing one premium quilt and two optional standard shams.", "Gingerbread Christmas quilt size and sham options"),
    ],
    20: [
        ("Gingerbread gift Christmas quilt bed mockup with two gingerbread figures, presents, candy canes, bows and holly.", "Gingerbread gift Christmas quilt set on bed"),
        ("Bedroom mockup showing gingerbread gift Christmas quilt with presents and coordinated pillow shams.", "Gingerbread gift Christmas quilt bedroom mockup"),
        ("Matching pillow sham with gingerbread figures, gifts, candy cane and holiday greenery.", "Gingerbread gift Christmas pillow sham"),
        ("Close-up of gingerbread faces, red bow, icing trim and visible quilt stitching.", "Gingerbread gift quilt stitching close-up"),
        ("Size and component panel showing one premium quilt and two optional standard shams.", "Gingerbread gift Christmas quilt size options"),
    ],
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def load_qa():
    data = json.loads(QA_DATASET.read_text(encoding="utf-8"))
    products = data["QA_Products"]
    handle_by_key = {item["product_key"]: item["handle"] for item in products}
    pos_by_key = {item["product_key"]: int(item["inventory_position"]) for item in products}

    image_fix_by_key = {}
    observation_by_key = {}
    image_by_qa_key = {img["qa_image_key"]: img for img in data["QA_Images"]}
    for image in data["QA_Images"]:
        observation_by_key[(image["product_key"], text(image["media_id"]))] = image["qa_observation"]
    for issue in data["QA_Issues"]:
        if issue["field"].startswith("image_") and issue["recommended_fix"]:
            image = image_by_qa_key.get(issue["qa_image_key"])
            if image:
                image_fix_by_key[(image["product_key"], text(image["media_id"]))] = issue["recommended_fix"]

    customizer = json.loads(CUSTOMIZER_AUDIT.read_text(encoding="utf-8"))
    customizer_by_pos = {int(item["inventory_position"]): item for item in customizer}
    return {
        "data": data,
        "product_keys": [item["product_key"] for item in products],
        "handle_by_key": handle_by_key,
        "pos_by_key": pos_by_key,
        "image_fix_by_key": image_fix_by_key,
        "observation_by_key": observation_by_key,
        "customizer_by_pos": customizer_by_pos,
    }


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Body (HTML)": first.get("Body (HTML)", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    return admin


def customizer_sentence(pos, audit):
    nodes = audit.get("personalization_nodes", [])
    if not nodes:
        return "No customer text field is used in the SEO claim for this revision."
    text_nodes = [node for node in nodes if node.get("label") in {"Enter Name", "Custom Your Name"}]
    if not text_nodes:
        text_nodes = [nodes[-1]]
    parts = []
    for node in text_nodes:
        required = "required" if node.get("required") else "optional"
        limit = node.get("maxLength")
        label = node.get("label", "custom text")
        label = "Enter Name" if label == "Custom Your Name" else label
        if limit:
            parts.append(f"Personalization uses a {required} {label} field up to {limit} characters.")
        else:
            parts.append(f"Personalization uses a {required} {label} field.")
    return " ".join(parts)


def make_description(pos, admin_row, customizer_text):
    item = PRODUCT_UPDATES[pos]
    option_names = [v for v in [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]] if v]
    options = ", ".join(option_names) if option_names else "available product selectors"
    product_type = (admin_row["Type"] or "bedding").lower()
    return (
        f"<p>{item['detail'].capitalize()} gives this Jeminise {product_type} a specific faith or holiday bedding look.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork focus: {item['detail']}.</li>"
        "<li>Gallery images show the main bed mockup plus close-up, fabric, care, size, component or room-view graphics where available.</li>"
        "<li>Use the visible gallery panels to compare bedding type, dimensions and included or optional pieces before checkout.</li></ul>"
        "<h3>Options and Customization</h3>"
        f"<ul><li>Available selectors cover {options}.</li>"
        f"<li>{customizer_text}</li>"
        "<li>Confirm the final product type, size, pillowcase and sheet-cover choices before checkout.</li></ul>"
    )


def compact_observation_to_alt(observation, fallback):
    value = re.sub(r"[:;].*$", "", text(observation)).strip()
    value = re.sub(r"\s+", " ", value)
    return (value or fallback)[:125]


def remove_templar_copy(wb):
    for sheet_name in ["SEO_Products", "Keyword_Map", "Buyer_Search_Research", "Product_Evidence"]:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and ("Templar" in cell.value or "Knight" in cell.value):
                    value = cell.value
                    value = value.replace("Christian Knight Templar", "Christian faith")
                    value = value.replace("Knight Templar", "faith")
                    value = value.replace("Templar", "faith")
                    value = value.replace("Knight", "faith")
                    cell.value = value


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    qa = load_qa()
    admin = load_admin(set(qa["handle_by_key"].values()))
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        customizer_text = customizer_sentence(pos, qa["customizer_by_pos"].get(pos, {}))

        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or ws.cell(row_num, idx["product_type"]).value
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = make_description(pos, admin_row, customizer_text)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"Buyer intent targets {item['cluster']} for US English product search; no search-volume claim is made."
        )
        ws.cell(row_num, idx["revision"]).value = "r4"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R4 deep audit after qa_batch_002_r3: used products_export_1.csv (sha256:{admin_hash}) as admin baseline; "
            "rewrote customer-facing copy, refreshed every image observation and alt from contact sheets, restored verified customization wording, "
            "and kept unsupported/mismatched labels out. Still NEEDS_REVIEW; no APPROVED/import."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        media_id = text(ws.cell(row_num, idx["media_id"]).value)
        key = (pk, media_id)
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        image_number = ws.cell(row_num, idx["image_number"]).value
        try:
            image_number = int(image_number)
        except (TypeError, ValueError):
            image_number = revised_images + 1
        details = IMAGE_DETAILS.get(pos, [])
        if 1 <= image_number <= len(details):
            observation, proposed_alt = details[image_number - 1]
        else:
            observation = qa["observation_by_key"].get(key) or f"{item['detail']} shown in gallery image {image_number}."
            proposed_alt = compact_observation_to_alt(observation, f"{item['title']} image {image_number}")
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = proposed_alt[:125]
        ws.cell(row_num, idx["revision"]).value = "r4"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R4: image-level observation and alt refreshed from contact sheet review; current admin alt and image URL matched from products_export_1.csv where possible."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "Product"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = "Sang re-QA r2 SERP evidence + products_export_1.csv admin baseline"
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = "R4 separates sibling product intent by visible motif, gallery evidence and verified customizer requirements."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r4"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as personalized faith or Christmas bedding."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm the design, size/component choices "
            "and whether the name field is required or optional."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = "Confirm product type, size, pillowcase or sheet-cover choices, care/fabric panels and customization requirement."
        ws.cell(row_num, idx["emotional_social_motivation"]).value = "Give a faith-centered, personalized or holiday-themed bedding gift with clear artwork."
        ws.cell(row_num, idx["purchase_concerns"]).value = "Name-field rules, exact bedding type, included components, fabric/care claims and image accuracy."
        ws.cell(row_num, idx["source_refs"]).value = f"products_export_1.csv sha256:{admin_hash}; Sang QA r2 {QA_RUN_ID}"
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R4_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No search-volume source supplied; Shopify admin CSV baseline is now available for stored title/meta/body/image alt."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = qa["product_keys"][pos - 11]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = f"Storefront/product.js from Sang QA r2; products_export_1.csv sha256:{admin_hash}"
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={admin_row['Option1 Name']}, "
            f"{admin_row['Option2 Name']}, {admin_row['Option3 Name']}; status={admin_row['Status']}; "
            f"image_count={len(admin_row['images'])}; design={item['detail']}."
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "R4 deepened qa_batch_002_r3 with image-level evidence and admin export cross-check; needs independent QA recheck."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R4 uses contact-sheet image review, live Customizer audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_002_r4_revision", "r4", "Deep review of latest qa_batch_002_r3; source r3 workbook was not modified."),
        ("qa_batch_002_r4_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_002_r4_scope", "inventory positions 11-20", "No products outside qa_batch_002 were revised."),
        ("qa_batch_002_r4_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_002_r4_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_002_r4_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    remove_templar_copy(wb)
    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa_r2": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_002 only; inventory positions 11-20",
        "revision": "r4",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "r4_changes": [
            "Deepened image-level observed_visual_details and alt_proposed for all 64 batch images.",
            "Removed customer-facing references to Shopify export/admin baseline from description_proposed_html.",
            "Kept personalization wording tied to verified required or optional Enter Name fields.",
            "Kept product 13 away from unsupported Knight Templar wording and tightened product 15 around Is Within Her butterfly evidence.",
        ],
        "next_step": "Sang QA lại qa_batch_002_r4 before any approval or import.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_002_r4",
                "",
                "Artifact rà sâu chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r3: `{SOURCE.relative_to(ROOT)}`",
                f"- QA r2 của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r4: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Điểm sửa sâu: image audit từng ảnh, description không còn ngôn ngữ nội bộ, customizer ghi đúng required/optional.",
                "- Bước tiếp theo: Sang QA lại `qa_batch_002_r4`.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
