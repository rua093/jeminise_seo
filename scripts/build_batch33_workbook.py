import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_033"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_032.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_033.xlsx"


TURTLE_REFS = "https://www.etsy.com/market/sea_turtle_quilt; https://www.wayfair.com/keyword.php?keyword=sea+turtle+bedding; https://www.amazon.com/sea-turtle-quilt/s?k=sea+turtle+quilt; https://www.pinterest.com/ideas/sea-turtle-quilt/904740390159/"
SOFTBALL_REFS = "https://www.amazon.com/softball-comforter/s?k=softball+comforter; https://www.etsy.com/market/softball_bedding; https://www.pinterest.com/ideas/softball-bedding/918246657688/; https://www.wayfair.com/keyword.php?keyword=softball+bedding"
PERSONALIZED_SOFTBALL_REFS = "https://www.etsy.com/market/personalized_softball_blanket; https://www.amazon.com/personalized-softball-blanket/s?k=personalized+softball+blanket; https://www.pinterest.com/search/pins/?q=personalized%20softball%20bedding; https://www.zazzle.com/softball+blankets"


PROFILES = {
    321: ("sea turtle patchwork quilt", "sea turtle quilt set", "Sea Turtle Quilt Set", "Sea Turtle Patchwork Quilt Set", "watercolor sea turtle quilt with blue and gold shell, floral garden artwork and matching turtle pillow shams", "sea turtle bedding, turtle patchwork quilt, coastal animal bedding", "EVERGREEN", TURTLE_REFS),
    322: ("striped softball ball comforter", "softball comforter set", "Softball Comforter Set", "Striped Softball Ball Comforter Set", "gray striped comforter with large yellow softball graphic, softball wordmark and matching pillowcases", "softball bedding, yellow softball comforter, girls softball bedding", "EVERGREEN", SOFTBALL_REFS),
    323: ("personalized softball glove comforter", "personalized softball comforter", "Personalized Softball Comforter Set", "Personalized Softball Glove Comforter Set", "black and gray softball comforter with glove, yellow ball, Payton name and number 09", "custom softball bedding, softball bedding with name, softball player comforter", "EVERGREEN", PERSONALIZED_SOFTBALL_REFS),
    324: ("softball quote patchwork comforter", "softball quote comforter", "Softball Quote Comforter Set", "Softball Quote Patchwork Comforter Set", "yellow, black and red patchwork comforter with softball quotes, ball graphics and sports typography", "softball lover bedding, softball patchwork comforter, softball bedroom decor", "EVERGREEN", SOFTBALL_REFS),
    325: ("sunflower softball comforter", "sunflower softball comforter", "Sunflower Softball Comforter Set", "Sunflower Softball Player Comforter Set", "black comforter with large yellow sunflowers, softball player silhouette, butterflies, number 19 and Taylor signature", "softball bedding for girls, sunflower softball bedding, personalized softball comforter", "EVERGREEN", PERSONALIZED_SOFTBALL_REFS),
    326: ("pink softball quote comforter", "pink softball comforter", "Pink Softball Comforter Set", "Pink Softball Quote Comforter Set", "pink and bright yellow softball comforter with player silhouettes, Gabby name and Look Pretty Play Dirty quote", "girls softball bedding, softball bedroom decor, softball quote bedding", "EVERGREEN", PERSONALIZED_SOFTBALL_REFS),
    327: ("vintage softball glove comforter", "softball fan comforter", "Softball Fan Comforter Set", "Vintage Softball Glove Comforter Set", "brown vintage softball comforter with large ball, leather glove artwork and Ally name on matching shams", "softball fan bedding, softball glove comforter, personalized softball bedding", "EVERGREEN", PERSONALIZED_SOFTBALL_REFS),
    328: ("softball flag comforter", "softball lover comforter", "Softball Lover Comforter Set", "Softball Flag Player Comforter Set", "black, white and yellow softball comforter with flag stripes, player silhouettes, number 13 and Sarah name", "softball player bedding, softball lover comforter, custom softball bedding", "EVERGREEN", PERSONALIZED_SOFTBALL_REFS),
    329: ("pastel softball patchwork comforter", "softball patchwork comforter", "Pastel Softball Patchwork Comforter Set", "Pastel Softball Patchwork Comforter Set", "mint, pink, white and black patchwork comforter with softball player silhouettes, ball pattern squares and matching shams", "softball bedding for girls, pastel softball bedding, softball comforter set", "EVERGREEN", SOFTBALL_REFS),
    330: ("yellow softball number comforter", "yellow softball comforter", "Yellow Softball Number Comforter Set", "Yellow Softball Number Comforter Set", "full yellow softball comforter with red stitching, number 47, Cayleigh name and matching numbered pillowcases", "personalized softball bedding, softball number comforter, yellow softball comforter", "EVERGREEN", PERSONALIZED_SOFTBALL_REFS),
}


def qa_issues(pos, html_error):
    issues = ["Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."]
    if pos == 321:
        issues.append("Sea turtle quilt should be checked for coastal versus animal patchwork targeting, quilt construction and optional sham wording.")
    if 322 <= pos <= 330:
        issues.append("Softball comforters are close variants; review keyword cannibalization, title uniqueness and whether each personalization example should stay in SEO copy.")
        issues.append("Several current storefront titles appear truncated; proposed titles should be checked against admin field limits before import.")
    if pos in (323, 325, 326, 327, 328, 330):
        issues.append("Visible sample names or numbers appear in imagery; verify personalization options and avoid implying the sample name is fixed.")
    if pos == 324:
        issues.append("Quote patchwork design contains multiple slogans; verify readable quote text before using any exact phrase in final copy.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def product_context(pos):
    if pos == 321:
        return "coastal bedrooms, sea turtle decor, animal-themed rooms or beach house bedding."
    return "softball bedrooms, teen sports rooms, player gifts, team-family gifts or personalized fan bedding."


def alt_for_image(h1, pos, ipos):
    if pos == 321:
        return {
            1: f"{h1} displayed on a bed in the main quilt mockup",
            2: f"Angled bedroom mockup for {h1}",
            3: f"Close-up pillow sham and turtle artwork detail for {h1}",
            4: f"Fabric and layered bedding feature image for {h1}",
            5: f"Size chart for {h1}",
            6: f"Bedspread construction and care feature image for {h1}",
            7: f"Overhead bedroom mockup of {h1}",
        }.get(ipos, f"Product detail image for {h1}")[:125]
    return {
        1: f"{h1} displayed on a bed in the main comforter mockup",
        2: f"Angled bedroom mockup for {h1}",
        3: f"Matching pillowcase detail for {h1}",
        4: f"Folded comforter mockup for {h1}",
        5: f"Fabric close-up for {h1}",
        6: f"Soft microfiber and care feature image for {h1}",
        7: f"Included comforter and pillowcase information for {h1}",
    }.get(ipos, f"Product detail image for {h1}")[:125]


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for {product_context(pos)}</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broad softball bedding and sea turtle bedding terms for collection pages or strongest representative products.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, room decor or a personalized gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": qa_issues(pos, html_error),
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed bedding, the buyer wants a {kind} matching a specific sea turtle or softball motif.",
            "functional_motivation": "Find the right design, size, product type, personalization fields, care details and room fit.",
            "emotional_social_motivation": "Give a personalized sports gift or decorate with coastal animal, softball team or player artwork.",
            "purchase_concerns": "Personalization spelling, sample names or numbers, product identity, material/care claims, size fit, included components and whether artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt_for_image(h1, pos, ipos), "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_033_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_033_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_033"},
        {"metric": "cumulative_products", "value": "330", "definition": "Total products included through batch_033"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[330:340]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_034",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
