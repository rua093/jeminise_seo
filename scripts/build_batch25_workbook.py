import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_025"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_024.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_025.xlsx"


BIGFOOT_REFS = "https://www.etsy.com/market/bigfoot_bedding; https://www.etsy.com/market/bigfoot_comforter_set; https://www.amazon.com/bigfoot-blanket/s?k=bigfoot+blanket; https://zumbamboo.com/collections/bigfoot-quilt-bedding-set; https://us.shein.com/CAPOZEN-Bigfoot-Bedding-Set-Cozy-Sasquatch-Quilt-Comforter-Perfect-For-Rustic-Cabin-Bedroom-Gift-Available-In-Twin-Full-Queen-King-Bigfoot-Style-05-p-350426124.html"
READING_REFS = "https://www.etsy.com/market/book_lover_blanket; https://www.etsy.com/ca/market/reading_blanket_for_girl; https://www.amazon.com/Just-Girl-Loves-Books-Blanket/dp/B0DMVXGZKG; https://www.amazon.com/clp/B0DJ32WQ8X; https://linovagoods.com/products/personalized-book-blanket-custom-name-reading-throw-floral-bookshelf-gift-for-book-lovers-cozy-library-decor; https://us.shein.com/Personalized-Reading-Blanket-For-Book-Lovers-Woven-Throw-Blanket-With-Custom-Quote-Gift-For-Her-Reader-Gift-Cozy-Reading-Nook-Decor-p-145337647.html"


PROFILES = {
    241: ("Bigfoot peace sign mountain landscape quilt", "Bigfoot peace sign quilt", "Bigfoot Peace Sign Quilt Set", "Bigfoot Peace Sign Mountain Quilt Set", "rustic mountain quilt with a Bigfoot figure giving a peace sign beside blue peaks and autumn trees", "Sasquatch quilt set, Bigfoot mountain bedding, cabin quilt set", "EVERGREEN", BIGFOOT_REFS),
    242: ("Bigfoot peace sign sunset quilt", "Bigfoot peace sign sunset quilt", "Bigfoot Peace Sign Sunset Quilt Set", "Bigfoot Peace Sign Sunset Quilt Set", "warm sunset quilt with close-up Bigfoot giving a peace sign over orange forest and teal mountain scenery", "Sasquatch bedding, Bigfoot quilt set, peace sign bedding", "EVERGREEN", BIGFOOT_REFS),
    243: ("Bigfoot silhouette moon forest quilt", "Bigfoot moon silhouette quilt", "Bigfoot Moon Silhouette Quilt Set", "Bigfoot Moon Forest Silhouette Quilt Set", "dark blue forest quilt with a large moon, mountain backdrop and Bigfoot silhouette walking forward", "Sasquatch forest bedding, Bigfoot quilt set, moon forest quilt", "EVERGREEN", BIGFOOT_REFS),
    244: ("Bigfoot snowy mountain forest quilt", "Bigfoot snowy mountain quilt", "Bigfoot Snowy Mountain Quilt Set", "Bigfoot Snowy Mountain Forest Quilt Set", "purple-blue snowy mountain quilt with pine forest, moonlit sky and small Bigfoot silhouette", "Sasquatch mountain bedding, Bigfoot cabin quilt, snowy forest quilt", "EVERGREEN", BIGFOOT_REFS),
    245: ("Bigfoot sunset forest silhouette quilt", "Bigfoot sunset forest quilt", "Bigfoot Sunset Forest Quilt Set", "Bigfoot Sunset Forest Silhouette Quilt Set", "orange sunset forest quilt with tall tree shadows and Bigfoot silhouette walking through the woods", "Sasquatch quilt, Bigfoot forest bedding, rustic cabin bedspread", "EVERGREEN", BIGFOOT_REFS),
    246: ("Bigfoot red sunglasses quilt", "Bigfoot red sunglasses quilt", "Bigfoot Red Sunglasses Quilt Set", "Bigfoot with Red Sunglasses Quilt Set", "vintage-style Bigfoot quilt with a red sunglasses character standing on a rock above pine trees", "Sasquatch bedding, Bigfoot quilt set, funny Bigfoot gift", "EVERGREEN", BIGFOOT_REFS),
    247: ("black hair reading girl bookshelf blanket", "personalized reading girl blanket", "Personalized Reading Girl Blanket", "Black Hair Reading Girl Bookshelf Blanket", "book lover throw blanket with a black-haired reading girl, bookshelf panels, quote blocks and custom name", "book lover blanket, reading blanket for girls, personalized book blanket", "EVERGREEN", READING_REFS),
    248: ("blonde reading girl bookshelf blanket", "personalized book lover blanket", "Personalized Book Lover Blanket", "Blonde Reading Girl Bookshelf Blanket", "bookshelf throw blanket with a blonde reading girl, stacked books, reading quotes and custom name", "reading blanket, bookworm blanket, custom name book blanket", "EVERGREEN", READING_REFS),
    249: ("bun hair girl antique books blanket", "girl who loves books blanket", "Girl Who Loves Books Blanket", "Bun Hair Girl Antique Books Blanket", "vintage book-page throw blanket with bun-haired girl, antique book stacks, flowers and custom name", "book lover gift blanket, personalized reading blanket, bookish throw blanket", "EVERGREEN", READING_REFS),
    250: ("brown hair girl wildflowers reading blanket", "just a girl who loves books blanket", "Just a Girl Who Loves Books Blanket", "Brown Hair Girl Wildflowers Reading Blanket", "floral book-page blanket with brown-haired girl, stacked books, wildflowers and reading quote artwork", "book lover blanket, reading gift for women, personalized throw blanket", "EVERGREEN", READING_REFS),
}


def qa_issues(pos, html_error):
    issues = [
        "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."
    ]
    if 241 <= pos <= 246:
        issues.append("Bigfoot/Sasquatch quilt products are close variants; review title uniqueness, motif separation and canonical strategy.")
        issues.append("Verify quilt/bedspread/fabric claims, machine-washable claims and optional pillow shams before import.")
    if 247 <= pos <= 250:
        issues.append("Reading girl blanket products are close personalized gift variants; verify character hair/style, displayed sample name and personalization fields.")
        issues.append("Feature images mention fleece, stitching, size and care claims; confirm exact blanket material and size variants before import.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def product_context(pos):
    if pos <= 246:
        return "cabin rooms, rustic bedrooms, cryptid gifts or Bigfoot-themed decor."
    return "reading nooks, bedrooms, book clubs, librarian gifts or personalized reader gifts."


def alt_for_image(h1, pos, ipos):
    if pos <= 246:
        return {
            1: f"{h1} displayed on a bed in the main quilt mockup",
            2: f"Angled bedroom mockup of {h1}",
            3: f"Close-up pillow sham and quilt artwork for {h1}",
            4: f"Fabric and layered bedding feature image for {h1}",
            5: f"Size chart for {h1}",
            6: f"Bedspread construction and care feature image for {h1}",
            7: f"Overhead bedroom mockup of {h1}",
        }.get(ipos, f"Product detail image for {h1}")[:125]
    return {
        1: f"{h1} shown as a personalized book lover throw blanket",
        2: f"Lifestyle couch mockup of {h1}",
        3: f"Close-up fabric and artwork detail for {h1}",
        4: f"Reading lifestyle mockup for {h1}",
        5: f"Size or multi-use guide for {h1}",
        6: f"Close-up reading scene artwork for {h1}",
        7: f"Fleece blanket feature image for {h1}",
        8: f"Family or lifestyle use image for {h1}",
        9: f"Blanket size chart for {h1}",
    }.get(ipos, f"Product detail image for {h1}")[:125]


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for {product_context(pos)}</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broad Bigfoot bedding and book lover blanket terms for collection pages or strongest representatives.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, room decor or a personalized gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": qa_issues(pos, html_error),
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed bedding, the buyer wants a {kind} matching a specific Bigfoot/Sasquatch or book lover motif.",
            "functional_motivation": "Find the right design, size, product type, personalization fields, care details and room fit.",
            "emotional_social_motivation": "Create a themed bedroom, cozy reading space or give a personalized cabin, cryptid or reader gift.",
            "purchase_concerns": "Personalization spelling, product identity, material/care claims, size fit, included components and whether artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt_for_image(h1, pos, ipos), "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_025_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_025_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_025"},
        {"metric": "cumulative_products", "value": "250", "definition": "Total products included through batch_025"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[250:260]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_026",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
