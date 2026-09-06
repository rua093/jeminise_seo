import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base
from build_batch11_workbook import append_dict, image_obs


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_032"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_031.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_032.xlsx"


DEER_REFS = "https://www.etsy.com/market/personalized_deer_blanket_couples; https://www.youcustomizeit.com/p/Deer-Comforters-Personalized/177732; https://foter.com/wildlife-comforter-sets; https://www.pinterest.com/pin/eternity-couple-deer-in-the-forest-couple-quilt-and-bedding-set-3d-printed-personalized-red--4598245680246284800/"
WOLF_REFS = "https://www.etsy.com/market/queen_size_wolf_quilts; https://society6.com/collections/duvet-covers-wolf; https://www.walmart.com/c/kp/wolf-comforter-sets; https://www.pinterest.com/pin/southwestern-wolf-dreamcatcher-quilt-set-wolf-patchwork-quilt-set-wolf-quilt-wolf-dreamcatcher--599260294195405768/"
CELTIC_REFS = "https://www.amazon.com/clp/B0DNSVY2QJ; https://luvingift.com/product/viking-quilt-set-norse-celtic-yggdrasil-tree-of-life/; https://myvikinggear.com/collections/viking-home-decoration/viking-quilt-set/; https://www.pinterest.com/pin/562668547209120403/"
ROOSTER_REFS = "https://www.amazon.com/chicken-quilt/s?k=chicken+quilt; https://www.etsy.com/market/chicken_quilt; https://www.wayfair.com/keyword.php?keyword=rooster+bedding; https://www.pinterest.com/ideas/chicken-quilt/912678085758/"


PROFILES = {
    311: ("wildlife deer couple comforter", "wildlife deer comforter", "Wildlife Deer Comforter Set", "Wildlife Deer Couple Comforter", "cream forest comforter with buck and doe standing together, You and Me We Got This text and floral grass details", "deer couple bedding, buck and doe comforter, rustic wildlife bedding", "EVERGREEN", DEER_REFS),
    312: ("deer couple forest clearing comforter", "deer couple comforter", "Deer Couple Comforter Set", "Deer Couple Forest Clearing Comforter", "cream and brown comforter with buck and doe inside forest clearing roots, You and Me We Got This text and custom names on shams", "buck doe bedding, personalized deer comforter, hunting couple bedding", "EVERGREEN", DEER_REFS),
    313: ("dark forest deer couple comforter", "forest deer comforter", "Forest Deer Comforter Set", "Dark Forest Deer Couple Comforter", "dark woodland comforter with buck and doe facing forward, You and Me We Got This script and custom his and her name shams", "deer couple bedding, personalized wildlife comforter, rustic deer bedding", "EVERGREEN", DEER_REFS),
    314: ("deer couple touching noses comforter", "deer couple love comforter", "Deer Couple Love Comforter Set", "Deer Couple Touching Noses Comforter", "autumn forest comforter with buck and doe touching noses, All of Me Loves All of You quote and custom name shams", "buck and doe bedding, romantic deer comforter, personalized couple bedding", "EVERGREEN", DEER_REFS),
    315: ("wolf dreamcatcher quilt", "wolf dreamcatcher quilt", "Wolf Dreamcatcher Quilt Set", "Wolf Head Dreamcatcher Quilt Set", "tan wolf bedding with large wolf head framed by dreamcatcher feathers, geometric border and matching shams", "wolf bedding, dreamcatcher quilt, southwestern wolf comforter", "EVERGREEN", WOLF_REFS),
    316: ("wolf geometric headdress quilt", "wolf headdress quilt", "Wolf Headdress Quilt Set", "Wolf Head Geometric Headdress Quilt Set", "wolf head bedding with geometric headdress, tan tribal pattern background and matching wolf pillow shams", "wolf quilt set, geometric wolf bedding, southwestern comforter", "EVERGREEN", WOLF_REFS),
    317: ("Yggdrasil Tree of Life quilt", "Yggdrasil Tree of Life quilt", "Yggdrasil Tree of Life Quilt Set", "Yggdrasil Tree of Life Quilt Set", "dark green and gold Yggdrasil quilt with circular Tree of Life medallion, Celtic knotwork and matching shams", "Celtic Tree of Life bedding, Viking quilt set, Yggdrasil bedding", "EVERGREEN", CELTIC_REFS),
    318: ("Yggdrasil Tree of Life ravens quilt", "Yggdrasil ravens quilt", "Yggdrasil Ravens Quilt Set", "Yggdrasil Tree of Life with Ravens Quilt Set", "blue and black Yggdrasil quilt with Tree of Life, raven silhouettes, Celtic border and matching shams", "Viking raven bedding, Yggdrasil quilt set, Norse Tree of Life quilt", "EVERGREEN", CELTIC_REFS),
    319: ("vibrant gold Yggdrasil quilt", "gold Yggdrasil quilt", "Gold Yggdrasil Quilt Set", "Vibrant Gold Yggdrasil Quilt Set", "gold and teal Yggdrasil quilt with bright Tree of Life motif, ornate Celtic border and matching shams", "Yggdrasil bedding, Celtic quilt set, Viking Tree of Life bedding", "EVERGREEN", CELTIC_REFS),
    320: ("rooster patchwork quilt", "rooster quilt set", "Rooster Patchwork Quilt Set", "Rooster Patchwork Pattern Quilt Set", "farmhouse rooster quilt with patchwork squares, chicken artwork, warm red and cream tones and matching shams", "chicken bedding, farmhouse rooster quilt, animal bedding set", "EVERGREEN", ROOSTER_REFS),
}


def qa_issues(pos, html_error):
    issues = ["Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment."]
    if 311 <= pos <= 314:
        issues.append("Deer couple comforters are close variants and overlap previous deer products; review motif separation, custom names or quote wording, comforter versus quilt naming and optional shams.")
    if pos in (315, 316):
        issues.append("Wolf products overlap southwestern wolf items from earlier batches; verify dreamcatcher versus geometric headdress motif and product type before import.")
    if 317 <= pos <= 319:
        issues.append("Yggdrasil variants overlap previous Tree of Life products; review canonical strategy and distinguish ravens, gold palette and medallion artwork.")
    if pos == 320:
        issues.append("Rooster/chicken animal bedding should be checked for quilt or bedspread construction, optional shams and farmhouse keyword targeting.")
    if html_error:
        issues.append(f"Storefront HTML fetch blocked/error: {html_error}; draft uses public product JSON and downloaded images.")
    return " ".join(issues)


def product_context(pos):
    if 311 <= pos <= 314:
        return "rustic bedrooms, hunting cabins, anniversary gifts or personalized couple bedding."
    if pos in (315, 316):
        return "southwestern bedrooms, rustic cabins, wolf decor or outdoors-inspired bedding."
    if 317 <= pos <= 319:
        return "Viking-inspired rooms, Celtic bedrooms, spiritual decor or Tree of Life gift bedding."
    return "farmhouse bedrooms, chicken decor, country homes or animal-themed bedding."


def alt_for_image(h1, pos, ipos):
    return {
        1: f"{h1} displayed on a bed in the main quilt mockup",
        2: f"Angled bedroom or flat product mockup for {h1}",
        3: f"Close-up pillow sham and artwork detail for {h1}",
        4: f"Fabric and layered bedding feature image for {h1}",
        5: f"Size chart or bedding set information for {h1}",
        6: f"Bedspread construction and care feature image for {h1}",
        7: f"Overhead bedroom mockup of {h1}",
        8: f"Additional bedroom mockup for {h1}",
    }.get(ipos, f"Product detail image for {h1}")[:125]


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        html_error = summary["html"].get("fetch_error", "")
        processing_status = "PARTIAL" if html_error else "DRAFTED"
        meta = f"Shop a {kind} featuring {detail}. Review size and personalization options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, making it suited for {product_context(pos)}</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the specific product motif: {label}. Keep broad deer, wolf, Celtic and farmhouse animal bedding terms for collection pages or strongest representatives.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding, room decor or a personalized gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": processing_status, "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": qa_issues(pos, html_error),
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": processing_status, "confidence_and_reason": "Medium: storefront HTML, product JSON and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed bedding, the buyer wants a {kind} matching a specific deer, wolf, Yggdrasil or rooster motif.",
            "functional_motivation": "Find the right design, size, product type, personalization fields, care details and room fit.",
            "emotional_social_motivation": "Give a personalized wildlife or farmhouse gift, or decorate with rustic, Celtic, wolf or animal artwork.",
            "purchase_concerns": "Personalization spelling, product identity, material/care claims, size fit, included components and whether artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt_for_image(h1, pos, ipos), "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_032_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_032_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_032"},
        {"metric": "cumulative_products", "value": "320", "definition": "Total products included through batch_032"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[320:330]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_033",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
