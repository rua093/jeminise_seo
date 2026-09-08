import csv
import hashlib
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_016_r3"
QA_RUN_ID = "20260907_210204"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_016_r2" / "SEO_Product_Optimization_qa_batch_016_r2.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
CUSTOMIZER_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_audit.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_016_r3.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_016_r3.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    151: {
        "title": "Blue Mosaic Giraffe Quilt Set",
        "meta_title": "Blue Mosaic Giraffe Quilt Set",
        "meta_description": "Shop a blue mosaic giraffe quilt set with orange giraffe portrait, matching shams, fabric panel and size chart.",
        "primary": "blue mosaic giraffe quilt set",
        "secondary": "giraffe patchwork quilt, animal patchwork bedding, safari giraffe quilt",
        "cluster": "blue orange mosaic giraffe quilt set",
        "detail": "large orange giraffe portrait built from blue and orange mosaic-style patchwork shapes",
        "intent_role": "Animal quilt page separated by blue mosaic background and close-up giraffe portrait.",
        "customizer": "The visible field is a general Customize Your Quilt text box; no name or number field is claimed.",
    },
    152: {
        "title": "Vintage Circle Horse Quilt Set",
        "meta_title": "Vintage Circle Horse Quilt Set",
        "meta_description": "Shop a vintage horse quilt set with circular red-orange frame, flowing mane artwork, matching shams and size chart.",
        "primary": "vintage circle horse quilt set",
        "secondary": "vintage horse quilt, horse patchwork bedding, equestrian quilt set",
        "cluster": "red orange circular horse patchwork quilt set",
        "detail": "horse head artwork with flowing mane inside a red-orange circular patchwork frame",
        "intent_role": "Horse page separated by circular frame, red-orange palette and flowing mane artwork.",
        "customizer": "The visible field is a general Customize Your Quilt text box; no name or number field is claimed.",
    },
    153: {
        "title": "Yellow Horse Patchwork Quilt Set",
        "meta_title": "Yellow Horse Patchwork Quilt Set",
        "meta_description": "Shop a yellow horse patchwork quilt with colorful horse portrait, black mane, tree silhouettes, shams and size chart.",
        "primary": "yellow horse patchwork quilt set",
        "secondary": "colorful horse quilt, farmhouse horse bedding, animal patchwork quilt",
        "cluster": "yellow colorful horse portrait patchwork quilt set",
        "detail": "colorful horse portrait on a bright yellow background with black mane and tree silhouettes",
        "intent_role": "Horse page separated by bright yellow base, geometric color blocks and black mane.",
        "customizer": "The visible field is a general Customize Your Quilt text box; no name or number field is claimed.",
    },
    154: {
        "title": "I Am Who He Says Blanket",
        "meta_title": "I Am Who He Says Blanket",
        "meta_description": "Shop a Christian scripture blanket with I Am Who He Says I Am text, Bible references, green gold typography and fleece panel.",
        "primary": "I Am Who He Says blanket",
        "secondary": "Christian scripture blanket, Bible verse blanket, inspirational Christian blanket",
        "cluster": "I Am Who He Says I Am Christian scripture blanket",
        "detail": "white blanket with green, gold and gray I Am Who He Says I Am scripture typography and Bible references",
        "intent_role": "Christian blanket page separated by visible scripture wording and blanket/fleece feature panels.",
        "customizer": "The customizer shows a required Choose Style control; no image personalization field is claimed.",
    },
    155: {
        "title": "Cream Twisted Tree Quilt Set",
        "meta_title": "Cream Twisted Tree Quilt Set",
        "meta_description": "Shop a cream Tree of Life quilt set with twisted trunk, curled roots, green leaves, matching shams and size chart.",
        "primary": "cream twisted tree quilt set",
        "secondary": "twisted tree of life quilt, green leaf tree bedding, nature tree quilt set",
        "cluster": "cream green twisted Tree of Life quilt set",
        "detail": "twisted brown Tree of Life trunk with curled roots and green leaves on a cream background",
        "intent_role": "Tree page separated by cream base, curled roots and green leaf artwork.",
        "customizer": "The visible field is a general Customize Your Quilt text box; no name or number field is claimed.",
    },
    156: {
        "title": "Blooming Tree of Life Quilt Set",
        "meta_title": "Blooming Tree of Life Quilt Set",
        "meta_description": "Shop a blooming Tree of Life quilt set with white flowers, exposed roots, black floral border, shams and size chart.",
        "primary": "blooming tree of life quilt set",
        "secondary": "floral tree quilt, black flower tree bedding, nature tree quilt set",
        "cluster": "black blooming floral Tree of Life quilt set",
        "detail": "blooming tree with white flowers, exposed roots, meadow scene and black floral border",
        "intent_role": "Tree page separated by blooming white flowers and black garden-style border.",
        "customizer": "The visible field is a general Customize Your Quilt text box; no name or number field is claimed.",
    },
    157: {
        "title": "Stained Glass Tree Landscape Quilt Set",
        "meta_title": "Stained Glass Tree Landscape Quilt Set",
        "meta_description": "Shop a stained-glass Tree of Life quilt with sunset landscape, round frame, colorful hills, matching shams and size chart.",
        "primary": "stained glass tree landscape quilt set",
        "secondary": "Tree of Life stained glass quilt, sunset tree bedding, landscape tree quilt set",
        "cluster": "stained-glass Tree of Life sunset landscape quilt",
        "detail": "Tree of Life landscape inside a stained-glass style round frame with sunset sky and colorful hills",
        "intent_role": "Tree page separated by stained-glass frame, sunset background and landscape hills.",
        "customizer": "The visible field is a general Customize Your Quilt text box; no name or number field is claimed.",
    },
    158: {
        "title": "Mjolnir Skull Celtic Quilt",
        "meta_title": "Mjolnir Skull Celtic Quilt",
        "meta_description": "Shop a Norse Mjolnir skull quilt with circular Celtic knot artwork, raven border, black silver palette and size chart.",
        "primary": "Mjolnir skull Celtic quilt",
        "secondary": "Norse mythology quilt, Viking hammer bedding, black silver Celtic quilt",
        "cluster": "black silver Mjolnir skull Celtic quilt",
        "detail": "Mjolnir hammer and skull motif inside a circular Celtic knot design with black raven border",
        "intent_role": "Norse page separated by Mjolnir skull center, Celtic ring and black-silver palette.",
        "customizer": "The visible field is a general Customize Your Item text box; no name or number field is claimed.",
    },
    159: {
        "title": "Norse Ravens Runic Quilt",
        "meta_title": "Norse Ravens Runic Quilt",
        "meta_description": "Shop a Norse ravens quilt set with two black ravens, runic compass artwork, beige background, shams and size panel.",
        "primary": "Norse ravens runic quilt",
        "secondary": "Norse raven quilt, runic compass bedding, Viking raven quilt set",
        "cluster": "beige Norse ravens runic compass quilt set",
        "detail": "two black ravens flanking a runic compass symbol on a beige quilt background",
        "intent_role": "Norse page separated by two-raven composition, beige base and runic compass center.",
        "customizer": "The visible field is a general Customize Your Item text box; no name or number field is claimed.",
    },
    160: {
        "title": "Pink Runic Raven Comforter",
        "meta_title": "Pink Runic Raven Comforter",
        "meta_description": "Shop a dark Norse raven comforter with pink runic circle, centered raven emblem, matching shams and bedding panels.",
        "primary": "pink runic raven comforter",
        "secondary": "Norse raven comforter, runic circle bedding, dark raven bedding set",
        "cluster": "dark blue pink runic raven comforter",
        "detail": "centered raven emblem inside a bright pink runic circle on a dark blue comforter",
        "intent_role": "Raven page separated by single centered raven, pink runic circle and comforter product type.",
        "customizer": "The visible field is a general Customize Your Item text box; no name or number field is claimed.",
    },
}


COMMON_QUILT_INCLUDED = ("Three-piece bedding set panel listing one premium quilt and two optional standard shams.", "Three piece quilt set information panel")
COMMON_SIZE = ("Premium quilt set size chart showing available bed sizes and measurements.", "Quilt set size chart")
COMMON_BEDSPREAD = ("Bedspread features panel showing quilt layers and anti-wrinkle, anti-static and machine-washable icons.", "Quilt construction feature panel")


def quilt7(label, alt_root):
    return {
        1: (f"Bedroom mockup of {label} with matching shams.", alt_root),
        2: (f"Room mockup of {label} with premium printed craft callout and fabric inset.", f"{alt_root} room mockup"),
        3: (f"Optional pillow shams panel showing {label}.", f"{alt_root} pillow shams"),
        4: (f"High-quality fabric panel showing {label} with quilt detail and pillow mockup.", f"{alt_root} fabric detail panel"),
        5: COMMON_SIZE,
        6: COMMON_BEDSPREAD,
        7: (f"Second bedroom mockup of {label} with matching shams.", f"{alt_root} bedroom mockup"),
    }


IMAGE_DETAILS = {
    151: quilt7("blue mosaic giraffe quilt with orange giraffe portrait", "Blue mosaic giraffe quilt set"),
    152: quilt7("vintage horse quilt with red-orange circular frame and flowing mane", "Vintage circle horse quilt set"),
    153: quilt7("yellow horse patchwork quilt with colorful horse portrait and black mane", "Yellow horse patchwork quilt"),
    154: {
        1: ("Person holding white Christian scripture blanket with I Am Who He Says I Am text.", "I Am Who He Says I Am scripture blanket"),
        2: ("Christian scripture blanket draped on sofa with green and gold typography.", "Christian scripture blanket on sofa"),
        3: ("Blanket style and size panel showing I Am Who He Says I Am design and use-case icons.", "Scripture blanket style and size panel"),
        4: ("Close-up collage showing printed scripture typography and white fleece texture.", "Christian blanket fleece close-up"),
        5: ("I Am Who He Says I Am blanket draped over a living room sofa.", "Inspirational Christian blanket on sofa"),
        6: ("Lifestyle image of scripture blanket used while reading with a small dog nearby.", "Scripture blanket reading lifestyle image"),
        7: ("Close view of scripture blanket beside an open book and cup.", "Bible verse blanket close-up"),
        8: ("Fleece blanket feature panel listing machine washable, soft and warm, dense stitching and 260GSM fleece.", "Fleece blanket feature panel"),
    },
    155: {
        1: ("Bedroom mockup of cream twisted Tree of Life quilt with green leaves and matching shams.", "Cream twisted Tree of Life quilt set"),
        2: ("Optional pillow shams panel showing cream Tree of Life artwork and feature icons.", "Cream Tree of Life pillow shams"),
        3: ("High-quality fabric panel showing cream Tree of Life quilt detail and pillow mockup.", "Cream Tree of Life quilt fabric panel"),
        4: COMMON_SIZE,
        5: COMMON_BEDSPREAD,
        6: ("Second bedroom mockup of cream twisted Tree of Life quilt with matching shams.", "Cream twisted tree quilt bedroom mockup"),
        7: ("Bedroom view of cream twisted Tree of Life quilt with green leaves and curled roots.", "Cream Tree of Life quilt bedroom view"),
    },
    156: quilt7("blooming Tree of Life quilt with white flowers, exposed roots and black floral border", "Blooming Tree of Life quilt set"),
    157: quilt7("stained-glass Tree of Life quilt with sunset landscape and colorful hills", "Stained glass Tree landscape quilt"),
    158: quilt7("black silver Norse Mjolnir skull quilt with Celtic ring and raven border", "Mjolnir skull Celtic quilt"),
    159: {
        1: ("Bedroom mockup of beige Norse ravens quilt with two black ravens and runic compass artwork.", "Norse ravens runic quilt set"),
        2: ("Feature mockup of Norse ravens quilt with cozy, breathable, lightweight and washable icons.", "Norse ravens quilt feature panel"),
        3: ("All seasons panel showing Norse ravens quilt details and magnified texture circles.", "Norse ravens quilt detail panel"),
        4: ("High-quality fabric panel showing raven and rune print close-ups.", "Norse ravens quilt fabric panel"),
        5: COMMON_SIZE,
        6: COMMON_BEDSPREAD,
        7: ("Second bedroom mockup of beige Norse ravens quilt with matching shams.", "Norse ravens quilt bedroom mockup"),
    },
    160: {
        1: ("Bedroom mockup of dark blue comforter with pink runic circle and centered raven emblem.", "Pink runic raven comforter"),
        2: ("Feature mockup of pink runic raven comforter with cozy, breathable, lightweight and washable icons.", "Pink runic raven comforter feature panel"),
        3: ("All seasons panel showing dark blue raven comforter with pink runic circle and magnified details.", "Runic raven comforter detail panel"),
        4: ("Precision stitching panel showing raven emblem and dark blue fabric close-ups.", "Runic raven comforter stitching panel"),
        5: ("Layer panel showing microfiber woven top, filling microfiber and back microfiber.", "Runic raven comforter layer panel"),
        6: ("Bedspread and pillow shams size panel for pink runic raven comforter set.", "Runic raven comforter size panel"),
        7: ("Optional pillow shams panel showing pink runic raven emblem on dark blue bedding.", "Runic raven pillow shams"),
    },
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def load_qa():
    data = json.loads(QA_DATASET.read_text(encoding="utf-8"))
    products = data["QA_Products"]
    customizer = json.loads(CUSTOMIZER_AUDIT.read_text(encoding="utf-8"))
    return {
        "product_keys": [item["product_key"] for item in products],
        "handle_by_key": {item["product_key"]: item["handle"] for item in products},
        "pos_by_key": {item["product_key"]: int(item["inventory_position"]) for item in products},
        "customizer_by_pos": {int(item["inventory_position"]): item for item in customizer},
    }


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    return admin


def option_text(admin_row):
    values = [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]]
    values = [v for v in values if v]
    return ", ".join(values) if values else "available size options"


def customizer_sentence(audit):
    excerpt = audit.get("config_excerpt", "")
    text_matches = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"minLength":\s*(\d+),\s*"maxLength":\s*(\d+)',
        excerpt,
    )
    option_matches = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"defaultOptionId"',
        excerpt,
    )
    parts = []
    for label, required, min_length, max_length in text_matches:
        requirement = "required" if required == "true" else "optional"
        parts.append(f"Personalization uses a {requirement} {label} field, {min_length}-{max_length} characters.")
    useful_options = [label for label, _ in option_matches if label and "confirmation" not in label.lower()]
    if useful_options:
        parts.append(f"Additional visible customization option: {', '.join(useful_options)}.")
    if not text_matches:
        parts.append("No shopper text-entry field is described for this product.")
    labels = " ".join(label for label, *_ in text_matches).lower()
    if "number" in labels:
        parts.append("Names and numbers visible in mockups are sample artwork unless the matching input is entered.")
    elif "name" in labels:
        parts.append("Visible names in mockups are sample artwork unless the matching input is entered.")
    elif text_matches:
        parts.append("Visible custom text in mockups is sample artwork unless the matching input is entered.")
    else:
        parts.append("Displayed wording or artwork is part of the selected style option, not a shopper-entered text claim.")
    return " ".join(parts)


def description(pos, admin_row, ctext):
    item = PRODUCT_UPDATES[pos]
    return (
        f"<p>{item['title']} features {item['detail']}. "
        "The wording focuses on the visible artwork, bedding format and selectable options for this exact design.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork: {item['detail']}.</li>"
        "<li>Gallery images show the main product mockup plus detail, feature, care, size or included-item panels where present.</li></ul>"
        "<h3>Options and Fit</h3>"
        f"<ul><li>Available option groups: {option_text(admin_row)}.</li>"
        f"<li>{ctext}</li>"
        "<li>Choose the product type and size shown on the product page before checkout.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    qa = load_qa()
    admin = load_admin(set(qa["handle_by_key"].values()))
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ctext = customizer_sentence(qa["customizer_by_pos"].get(pos, {}))
        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or "Bedding"
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, ctext)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"US English buyer intent targets {item['cluster']}; no volume or ranking claim is made."
        )
        ws.cell(row_num, idx["keyword_evidence_level"]).value = "SERP_ONLY"
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R3 deep recheck after Sang QA batch 016: used products_export_1.csv sha256:{admin_hash}; "
            "removed internal workflow copy, filled publishable H1/title values, avoided unsupported customization/photo claims, "
            "split animal/tree/Norse intents by visible motif, revalidated optional/required custom fields and rewrote image observations/alts from inspected contact sheets. "
            "Still NEEDS_REVIEW."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    seen_alts = set()
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        image_number = int(ws.cell(row_num, idx["image_number"]).value)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation, alt = IMAGE_DETAILS[pos][image_number]
        if alt in seen_alts:
            alt = f"{alt} for {PRODUCT_UPDATES[pos]['cluster']}"
        seen_alts.add(alt)
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = alt[:125]
        ws.cell(row_num, idx["alt_action"]).value = "SET"
        ws.cell(row_num, idx["viewed_status"]).value = "VIEWED"
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R3: observation and alt rechecked from inspected contact sheet; admin image URL and current alt matched from Shopify CSV where available."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "PRODUCT"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = (
            "Sang QA batch 016 + products_export_1.csv admin baseline + customizer_audit.json + contact-sheet inspection"
        )
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = (
            "R3 separates nearby pages by visible motif, wording, color palette and proof level for customization."
        )
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r3"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as a specific product page."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm exact artwork, product type, size choices, "
            "and only customization options that are actually supported."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = (
            "Confirm product type, visible motif, size/care panels, supported custom fields and image accuracy before purchase."
        )
        ws.cell(row_num, idx["emotional_social_motivation"]).value = (
            "Choose an animal, Christian, Tree of Life or Norse bedding item without confusing it with another similar design."
        )
        ws.cell(row_num, idx["purchase_concerns"]).value = (
            "Avoid unsupported name, number, photo, upload, cultural-origin or material claims; confirm exact size and product type."
        )
        ws.cell(row_num, idx["source_refs"]).value = (
            f"products_export_1.csv sha256:{admin_hash}; Sang QA {QA_RUN_ID}; customizer_audit.json; contact sheets 151-160"
        )
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R3_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = (
            "No Search Console, paid keyword volume or internal site-search export supplied; proof remains SERP-only."
        )
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    pk_by_pos = {qa["pos_by_key"][pk]: pk for pk in qa["product_keys"]}
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = pk_by_pos[pos]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ctext = customizer_sentence(qa["customizer_by_pos"].get(pos, {}))
        ws.cell(row_num, idx["sources_accessed"]).value = (
            f"Original storefront/product evidence; Sang QA {QA_RUN_ID}; products_export_1.csv sha256:{admin_hash}; "
            "customizer_audit.json; inspected contact sheets"
        )
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={option_text(admin_row)}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={item['detail']}; "
            f"customizer={ctext}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "Sang QA issues addressed in r3; independent QA recheck still required."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R3 uses inspected contact sheets, customizer audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_016_r3_revision", "r3", "Separate 10-product revision after Sang QA; source workbook was not modified."),
        ("qa_batch_016_r3_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_016_r3_scope", "inventory positions 151-160", "No products outside qa_batch_016 were revised."),
        ("qa_batch_016_r3_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_016_r3_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_016_r3_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_016 only; inventory positions 151-160",
        "revision": "r3",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Hold for combined QA handoff after revisions through qa_batch_020, per user request.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_016_r3",
                "",
                "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r2: `{SOURCE.relative_to(ROOT)}`",
                f"- QA của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r3: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Sửa trọng yếu: rà lại theo rút kinh nghiệm các batch trước, bỏ nội dung nội bộ, giữ wording publishable, tránh claim custom/photo không có bằng chứng, tách intent cho animal/tree/Norse/Christian và viết lại mô tả/alt theo contact sheets.",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Bước tiếp theo: giữ lại để gửi Sang QA chung sau khi revision đến `qa_batch_020`, theo yêu cầu người dùng.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
