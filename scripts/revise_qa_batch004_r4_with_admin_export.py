import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import revise_qa_batch002_r4_with_admin_export as base


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_004_r4"
QA_RUN_ID = "20260907_144539"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_004_r3" / "SEO_Product_Optimization_qa_batch_004_r3.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
CUSTOMIZER_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_audit.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_004_r4.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_004_r4.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    31: {
        "title": "Christmas Cardinal Birdhouse Quilt Set",
        "meta_title": "Christmas Cardinal Birdhouse Quilt Set",
        "meta_description": "Shop a Christmas cardinal birdhouse quilt with snowy patchwork artwork, quilt sizes and pillowcase options.",
        "primary": "Christmas cardinal birdhouse quilt",
        "secondary": "cardinal birdhouse quilt set, snowy cardinal quilt, Christmas bird quilt bedding",
        "cluster": "Christmas cardinal birdhouse quilt",
        "intent_role": "Cardinal birdhouse quilt page; no shopper text-entry field is used in the SEO claim.",
        "detail": "cardinals near a snowy birdhouse with floral patchwork Christmas quilt artwork",
    },
    32: {
        "title": "Christmas Cardinals Snowy Branches Quilt Set",
        "meta_title": "Christmas Cardinals Snowy Branches Quilt Set",
        "meta_description": "Shop a Christmas cardinals snowy branches quilt with winter patchwork artwork, sizes and pillowcase options.",
        "primary": "Christmas cardinals snowy branches quilt",
        "secondary": "cardinal winter quilt set, snowy branch cardinal quilt, Christmas cardinal bedding",
        "cluster": "Christmas cardinals snowy branches quilt",
        "intent_role": "Snowy branches cardinal quilt page; image order corrected for close-up, sham and angled bedroom views.",
        "detail": "red cardinals on snowy branches with winter patchwork quilt artwork",
    },
    33: {
        "title": "Cow Landscape Farmhouse Quilt Set",
        "meta_title": "Cow Landscape Farmhouse Quilt Set",
        "meta_description": "Shop a cow landscape farmhouse quilt set with animal patchwork artwork, quilt sizes and separate pillowcase options.",
        "primary": "cow landscape farmhouse quilt set",
        "secondary": "cow patchwork quilt, farmhouse animal quilt set, cow quilt bedding",
        "cluster": "cow landscape farmhouse quilt set",
        "intent_role": "Cow farmhouse quilt-set page; dragonfly wording remains removed and no personalization claim is used.",
        "detail": "cow and landscape farmhouse patchwork quilt artwork",
    },
    34: {
        "title": "Crocodile Patchwork Animal Quilt Set",
        "meta_title": "Crocodile Patchwork Animal Quilt Set",
        "meta_description": "Shop a crocodile patchwork animal quilt set with wildlife artwork, quilt sizes and separate pillowcase options.",
        "primary": "crocodile animal quilt set",
        "secondary": "crocodile patchwork quilt, crocodile bedding quilt, animal patchwork quilt set",
        "cluster": "crocodile animal quilt set",
        "intent_role": "Broader crocodile bedding phrase retained as SERP_ONLY because exact finished-product evidence is limited.",
        "detail": "crocodile wildlife patchwork quilt artwork",
    },
    35: {
        "title": "Personalized Retro Football Flag Comforter Set",
        "meta_title": "Personalized Retro Football Flag Comforter Set",
        "meta_description": "Customize a retro football flag comforter with required name, optional number, product type, size, pillowcase and sheet-cover choices.",
        "primary": "personalized retro football flag comforter",
        "secondary": "custom football flag comforter, football comforter with name, football bedding with number",
        "cluster": "personalized retro football flag comforter",
        "intent_role": "Football flag page with verified required Enter Name up to 25 characters and optional Enter Number up to 5 characters.",
        "detail": "retro American football close-up on flag comforter artwork with sample name and number only as examples",
    },
    36: {
        "title": "Personalized Grunge Football Comforter Set",
        "meta_title": "Personalized Grunge Football Comforter Set",
        "meta_description": "Customize a grunge football comforter with required name, optional number, product type, size, pillowcase and sheet-cover choices.",
        "primary": "personalized grunge football comforter",
        "secondary": "custom grunge football bedding, football comforter with name, vintage football comforter",
        "cluster": "personalized grunge football comforter",
        "intent_role": "Natural personalized grunge football comforter phrase; verified required name and optional number fields restored.",
        "detail": "American football close-up on grunge vintage comforter artwork with sample name and number only as examples",
    },
    37: {
        "title": "Personalized Cosmic Football Comforter Set",
        "meta_title": "Personalized Cosmic Football Comforter Set",
        "meta_description": "Customize a cosmic football comforter with required name, optional number, product type, size, pillowcase and sheet-cover choices.",
        "primary": "personalized cosmic football comforter",
        "secondary": "custom galaxy football bedding, football comforter with name, cosmic football bedding set",
        "cluster": "personalized cosmic football comforter",
        "intent_role": "Cosmic football page separated from generic football bedding; verified name and optional number fields restored.",
        "detail": "American football on cosmic background vintage comforter artwork with sample name and number only as examples",
    },
    38: {
        "title": "Personalized USA Flag Football Comforter Set",
        "meta_title": "Personalized USA Flag Football Comforter Set",
        "meta_description": "Customize a USA flag football comforter with required name, optional number, product type, size, pillowcase and sheet-cover choices.",
        "primary": "personalized USA flag football comforter",
        "secondary": "custom USA football comforter, football bedding with name and number, American flag football bedding",
        "cluster": "personalized USA flag football comforter",
        "intent_role": "USA flag football angle separated from product 40's broader patriotic flag angle.",
        "detail": "American football on USA flag background comforter artwork with sample name and number only as examples",
    },
    39: {
        "title": "Personalized Paint Splash Football Comforter Set",
        "meta_title": "Personalized Paint Splash Football Comforter Set",
        "meta_description": "Customize a paint splash football comforter with required name, optional number, product type, size, pillowcase and sheet-cover choices.",
        "primary": "personalized paint splash football comforter",
        "secondary": "custom paint splash football bedding, vintage football comforter with name, football bedding with number",
        "cluster": "personalized paint splash football comforter",
        "intent_role": "Paint splash football page with verified required name and optional number fields restored.",
        "detail": "American football with paint splash vintage comforter artwork with sample name and number only as examples",
    },
    40: {
        "title": "Personalized Patriotic Football Comforter Set",
        "meta_title": "Personalized Patriotic Football Comforter Set",
        "meta_description": "Customize a patriotic football comforter with required name, optional number, product type, size, pillowcase and sheet-cover choices.",
        "primary": "personalized patriotic football comforter",
        "secondary": "custom patriotic football bedding, football flag comforter with name, football bedding with number",
        "cluster": "personalized patriotic football comforter",
        "intent_role": "Patriotic flag football angle separated from product 38's USA flag background phrase.",
        "detail": "American football with patriotic flag comforter artwork with sample name and number only as examples",
    },
}


IMAGE_DETAILS = {
    31: [
        ("Christmas cardinal birdhouse quilt bed mockup with red cardinals, snowy birdhouse, white flowers and red floral border.", "Christmas cardinal birdhouse quilt set on bed"),
        ("Bedroom mockup with cardinal birdhouse quilt and premium printed craft callout.", "Cardinal birdhouse quilt premium craft mockup"),
        ("Optional pillow shams panel showing matching cardinal birdhouse pillow artwork and feature icons.", "Cardinal birdhouse quilt optional pillow shams"),
        ("Fabric panel showing cardinal birdhouse quilt layers and soft breathable skin-friendly fabric callout.", "Cardinal birdhouse quilt fabric feature panel"),
        ("Premium quilt set size chart with cardinal birdhouse artwork and quilt dimensions.", "Cardinal birdhouse quilt size chart"),
        ("Bedspread features panel showing microfiber layers and anti-wrinkle, anti-static, fade resistant, shrink resistant and machine washable icons.", "Cardinal birdhouse quilt bedspread features"),
        ("Overhead bedroom mockup showing cardinal birdhouse quilt with matching pillows.", "Cardinal birdhouse quilt overhead bedroom view"),
    ],
    32: [
        ("Christmas cardinals snowy branches quilt bed mockup with red cardinals, berries, poinsettias and matching shams.", "Christmas cardinals snowy branches quilt set"),
        ("Close-up of a red cardinal head with snowy branch background and visible quilt stitching.", "Cardinal snowy branches quilt close-up"),
        ("Matching pillow sham with two red cardinals, snowy branches, berries and poinsettias.", "Cardinals snowy branches pillow sham"),
        ("Bedroom mockup showing cardinals snowy branches quilt with coordinated pillow shams.", "Cardinals snowy branches quilt bedroom mockup"),
        ("Size and component panel showing one premium quilt and two optional standard shams.", "Cardinals snowy branches quilt size options"),
    ],
    33: [
        ("Cow farmhouse quilt bed mockup with black-and-white cow, sunflowers, blue sky and matching pillow shams.", "Cow sunflower farmhouse quilt set on bed"),
        ("Bedroom mockup with cow sunflower quilt and premium printed craft callout.", "Cow sunflower farmhouse quilt room mockup"),
        ("Optional pillow shams panel showing matching cow and sunflower pillow artwork.", "Cow sunflower quilt optional pillow shams"),
        ("Fabric panel showing cow quilt layers and soft breathable skin-friendly fabric callout.", "Cow sunflower quilt fabric feature panel"),
        ("Premium quilt set size chart with cow sunflower artwork and quilt dimensions.", "Cow sunflower farmhouse quilt size chart"),
        ("Bedspread features panel showing microfiber layers and care-resistance icons.", "Cow sunflower quilt bedspread features"),
        ("Overhead bedroom mockup showing cow sunflower quilt with matching pillows.", "Cow sunflower quilt overhead bedroom view"),
    ],
    34: [
        ("Crocodile patchwork quilt bed mockup with colorful crocodile, white background, flowers and matching pillow shams.", "Crocodile patchwork animal quilt set on bed"),
        ("Bedroom mockup with crocodile patchwork quilt and premium printed craft callout.", "Crocodile patchwork quilt room mockup"),
        ("Optional pillow shams panel showing matching crocodile patchwork pillow artwork.", "Crocodile patchwork quilt optional pillow shams"),
        ("Fabric panel showing crocodile quilt layers and soft breathable skin-friendly fabric callout.", "Crocodile patchwork quilt fabric feature panel"),
        ("Premium quilt set size chart with crocodile artwork and quilt dimensions.", "Crocodile patchwork quilt size chart"),
        ("Bedspread features panel showing microfiber layers and care-resistance icons.", "Crocodile patchwork quilt bedspread features"),
        ("Overhead bedroom mockup showing crocodile patchwork quilt with matching pillows.", "Crocodile patchwork quilt overhead bedroom view"),
    ],
    35: [
        ("Retro football flag comforter bed mockup with Kevin sample name, number 15 and football close-up artwork.", "Personalized retro football flag comforter"),
        ("Angled bedroom mockup showing Kevin sample name, number 15 pillowcases and football flag artwork.", "Retro football flag comforter room mockup"),
        ("Product type comparison panel showing duvet cover set and comforter set options for all-season use.", "Retro football comforter bedding type options"),
        ("Close bedroom view of retro football flag comforter with soft, lightweight, durable and breathable icons.", "Retro football flag comforter feature icons"),
        ("Stress-free easy care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade notes.", "Retro football comforter easy care panel"),
        ("Size dimension panel showing Twin, Full, Queen and King bedding dimensions.", "Retro football comforter size chart"),
    ],
    36: [
        ("Grunge football comforter bed mockup with Brian sample name, number 15 and vintage red-white stripe background.", "Personalized grunge football comforter"),
        ("Angled bedroom mockup showing Brian sample name, number 15 pillowcases and grunge football artwork.", "Grunge football comforter room mockup"),
        ("Product type comparison panel showing duvet cover set and comforter set options for all-season use.", "Grunge football comforter bedding type options"),
        ("Close bedroom view of grunge football comforter with soft, lightweight, durable and breathable icons.", "Grunge football comforter feature icons"),
        ("Stress-free easy care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade notes.", "Grunge football comforter easy care panel"),
        ("Size dimension panel showing Twin, Full, Queen and King bedding dimensions.", "Grunge football comforter size chart"),
    ],
    37: [
        ("Cosmic football comforter bed mockup with Kevin sample name, number 20 and space-color background.", "Personalized cosmic football comforter"),
        ("Angled bedroom mockup showing Kevin sample name, number 20 pillowcases and cosmic football artwork.", "Cosmic football comforter room mockup"),
        ("Product type comparison panel showing duvet cover set and comforter set options for all-season use.", "Cosmic football comforter bedding type options"),
        ("Close bedroom view of cosmic football comforter with soft, lightweight, durable and breathable icons.", "Cosmic football comforter feature icons"),
        ("Stress-free easy care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade notes.", "Cosmic football comforter easy care panel"),
        ("Size dimension panel showing Twin, Full, Queen and King bedding dimensions.", "Cosmic football comforter size chart"),
    ],
    38: [
        ("USA flag football comforter bed mockup with Michael sample name, number 30 and American flag background.", "Personalized USA flag football comforter"),
        ("Angled bedroom mockup showing Michael sample name, number 30 pillowcases and USA flag football artwork.", "USA flag football comforter room mockup"),
        ("Product type comparison panel showing duvet cover set and comforter set options for all-season use.", "USA flag football comforter bedding type options"),
        ("Close bedroom view of USA flag football comforter with soft, lightweight, durable and breathable icons.", "USA flag football comforter feature icons"),
        ("Stress-free easy care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade notes.", "USA flag football comforter easy care panel"),
        ("Size dimension panel showing Twin, Full, Queen and King bedding dimensions.", "USA flag football comforter size chart"),
    ],
    39: [
        ("Paint splash football comforter bed mockup with Mark sample name, number 15 and brown-gold football artwork.", "Personalized paint splash football comforter"),
        ("Angled bedroom mockup showing Mark sample name, number 15 pillowcases and paint splash football artwork.", "Paint splash football comforter room mockup"),
        ("Product type comparison panel showing duvet cover set and comforter set options for all-season use.", "Paint splash football bedding type options"),
        ("Close bedroom view of paint splash football comforter with soft, lightweight, durable and breathable icons.", "Paint splash football comforter feature icons"),
        ("Stress-free easy care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade notes.", "Paint splash football easy care panel"),
        ("Size dimension panel showing Twin, Full, Queen and King bedding dimensions.", "Paint splash football comforter size chart"),
    ],
    40: [
        ("Patriotic football comforter bed mockup with Matthew sample name, number 08 and football printed with American flag artwork.", "Personalized patriotic football comforter"),
        ("Angled bedroom mockup showing Matthew sample name, number 08 pillowcases and patriotic football artwork.", "Patriotic football comforter room mockup"),
        ("Product type comparison panel showing duvet cover set and comforter set options for all-season use.", "Patriotic football bedding type options"),
        ("Close bedroom view of patriotic football comforter with soft, lightweight, durable and breathable icons.", "Patriotic football comforter feature icons"),
        ("Stress-free easy care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade notes.", "Patriotic football easy care panel"),
        ("Size dimension panel showing Twin, Full, Queen and King bedding dimensions.", "Patriotic football comforter size chart"),
    ],
}


def patch_base_globals():
    base.ROOT = ROOT
    base.SHOP = SHOP
    base.RUN_ID = RUN_ID
    base.BATCH_ID = BATCH_ID
    base.QA_RUN_ID = QA_RUN_ID
    base.SOURCE = SOURCE
    base.QA_DATASET = QA_DATASET
    base.CUSTOMIZER_AUDIT = CUSTOMIZER_AUDIT
    base.ADMIN_EXPORT = ADMIN_EXPORT
    base.RESULT_DIR = RESULT_DIR
    base.RUN_DIR = RUN_DIR
    base.OUTPUT = OUTPUT
    base.REPORT = REPORT
    base.SUMMARY = SUMMARY
    base.PRODUCT_UPDATES = PRODUCT_UPDATES


def customizer_text(pos, qa):
    if pos < 35:
        return "No shopper text-entry field is described for this product."
    return (
        "Personalization uses a required Enter Name field up to 25 characters and "
        "optional Enter Number field up to 5 characters. Names and numbers visible in mockups are treated as samples."
    )


def description(pos, admin_row, customizer):
    item = PRODUCT_UPDATES[pos]
    option_names = [v for v in [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]] if v]
    options = ", ".join(option_names) if option_names else "available product selectors"
    heading = "Options and Personalization" if pos >= 35 else "Options"
    product_type = (admin_row["Type"] or "bedding").lower()
    return (
        f"<p>{item['detail'].capitalize()} gives this Jeminise {product_type} a product-specific bedding focus.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork focus: {item['detail']}.</li>"
        "<li>Gallery images include the main mockup plus close-up, sham, care, comparison, size or material graphics where shown.</li>"
        "<li>Feature panels should be used to confirm fabric, care, size and component details before checkout.</li></ul>"
        f"<h3>{heading}</h3>"
        f"<ul><li>Available selectors cover {options}.</li>"
        f"<li>{customizer}</li>"
        "<li>Use the live selectors to confirm product type, size, pillowcases and sheet-cover choices before checkout.</li></ul>"
    )


def main():
    patch_base_globals()
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)

    qa = base.load_qa()
    admin = base.load_admin(set(qa["handle_by_key"].values()))
    admin_hash = base.sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()
    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = base.headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = base.text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ctext = customizer_text(pos, qa)

        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or ws.cell(row_num, idx["product_type"]).value
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, ctext)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"Buyer intent targets {item['cluster']} for US English product search; no search-volume claim is made."
        )
        ws.cell(row_num, idx["revision"]).value = "r4"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R4 deep audit after qa_batch_004_r3: used products_export_1.csv (sha256:{admin_hash}) as admin baseline; "
            "rewrote customer-facing descriptions, refreshed every image observation/alt from contact sheets, restored verified football personalization, "
            "and separated overlapping keyword intent. Still NEEDS_REVIEW; no APPROVED/import."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = base.headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    for row_num in range(2, ws.max_row + 1):
        handle = base.text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        media_id = base.text(ws.cell(row_num, idx["media_id"]).value)
        key = (pk, media_id)
        image_url = base.text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(base.normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        pos = qa["pos_by_key"][pk]
        image_number = ws.cell(row_num, idx["image_number"]).value
        try:
            image_number = int(image_number)
        except (TypeError, ValueError):
            image_number = revised_images + 1
        details = IMAGE_DETAILS.get(pos, [])
        if 1 <= image_number <= len(details):
            observation, proposed_alt = details[image_number - 1]
        else:
            observation = qa["observation_by_key"].get(key) or f"{PRODUCT_UPDATES[pos]['detail']} shown in gallery image {image_number}."
            proposed_alt = base.compact_observation_to_alt(observation, f"{PRODUCT_UPDATES[pos]['title']} image {image_number}")
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = proposed_alt[:125]
        ws.cell(row_num, idx["revision"]).value = "r4"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R4: image-level observation and alt refreshed from contact sheet review; current admin alt and image URL matched from products_export_1.csv where possible."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = base.headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = base.text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = base.text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "Product"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = "Sang re-QA r2 SERP evidence + products_export_1.csv admin baseline"
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = "R4 separates sibling product intent by visible motif, gallery evidence and verified customizer requirements."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r4"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = base.headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = base.text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as a specific bedding product."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm the artwork, product type, size/component choices "
            "and any verified name or number requirements."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = "Confirm product type, size, pillowcase or sheet-cover choices, care/fabric panels and customization requirement."
        ws.cell(row_num, idx["emotional_social_motivation"]).value = "Choose a holiday, animal, farmhouse or football-themed bedding gift with clear artwork."
        ws.cell(row_num, idx["purchase_concerns"]).value = "Name and number rules, exact bedding type, included components, fabric/care claims and image accuracy."
        ws.cell(row_num, idx["source_refs"]).value = f"products_export_1.csv sha256:{admin_hash}; Sang QA r2 {QA_RUN_ID}"
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R4_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No search-volume source supplied; Shopify admin CSV baseline is now available for stored title/meta/body/image alt."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = base.headers(ws)
    for row_num in range(2, ws.max_row + 1):
        evidence_id = base.text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = qa["product_keys"][pos - 31]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = f"Storefront/product.js from Sang QA r2; products_export_1.csv sha256:{admin_hash}"
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={admin_row['Option1 Name']}, "
            f"{admin_row['Option2 Name']}, {admin_row['Option3 Name']}; status={admin_row['Status']}; "
            f"image_count={len(admin_row['images'])}; design={item['detail']}."
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "R4 deepened qa_batch_004_r3 with image-level evidence and admin export cross-check; needs independent QA recheck."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R4 uses contact-sheet image review, live Customizer audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = base.headers(ws)
    for metric, value, definition in [
        ("qa_batch_004_r4_revision", "r4", "Deep review of latest qa_batch_004_r3; source r3 workbook was not modified."),
        ("qa_batch_004_r4_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_004_r4_scope", "inventory positions 31-40", "No products outside qa_batch_004 were revised."),
        ("qa_batch_004_r4_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_004_r4_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_004_r4_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa_r2": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_004 only; inventory positions 31-40",
        "revision": "r4",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "r4_changes": [
            "Deepened image-level observed_visual_details and alt_proposed for all 62 batch images.",
            "Restored football personalization claims only where verified: required Enter Name 1-25 characters and optional Enter Number 1-5 characters.",
            "Kept products 31-34 free of unsupported personalization claims.",
            "Separated product 38 USA flag football intent from product 40 patriotic football intent.",
            "Removed customer-facing Shopify export/admin wording from description_proposed_html.",
        ],
        "next_step": "Sang QA lại qa_batch_004_r4 before any approval or import.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_004_r4",
                "",
                "Artifact rà sâu chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r3: `{SOURCE.relative_to(ROOT)}`",
                f"- QA r2 của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r4: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Sửa trọng yếu: football products 35-40 khôi phục personalization đã xác minh; products 38/40 tách intent; 62 ảnh có alt riêng.",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Bước tiếp theo: Sang QA lại `qa_batch_004_r4`.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
