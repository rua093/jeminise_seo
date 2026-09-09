from copy import copy
from datetime import datetime
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
SRC = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_032_r4" / "SEO_Product_Optimization_qa_batch_032_r4.xlsx"
OUT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_032_r5"
OUT_XLSX = OUT_DIR / "SEO_Product_Optimization_qa_batch_032_r5.xlsx"
OUT_MD = OUT_DIR / "SEO_Product_Optimization_qa_batch_032_r5.md"


FIXES = {
    "personalized-wildlife-deer-comforter-with-buck-6532972aba-6532972aba": {
        "primary_keyword": "You and Me deer quilt set",
        "secondary_keywords": "deer quilt set; woodland deer bedding; buck doe quilt",
        "title_proposed": "You and Me Deer Quilt Set",
        "meta_title_seo": "You and Me Deer Quilt Set",
        "meta_description_seo": "Shop You and Me deer quilt set with beige woodland artwork, two spotted deer, orange flowers and statement text.",
        "description_proposed_html": """
<p>Two spotted deer, orange flower accents and You and Me We Got This text give this quilt set a warm woodland couple theme. The beige background keeps the deer artwork soft and easy to read.</p>
<h2>Design Details</h2>
<ul>
<li>Beige woodland quilt artwork with two spotted deer.</li>
<li>You and Me We Got This text and orange flower accents are visible.</li>
<li>Gallery includes the quilt mockup plus size, care, fabric or lifestyle panels where shown.</li>
</ul>
<h2>Personalization and Options</h2>
<p>Choose Quilt Size and Pillowcase Quantity options are listed. Customize Your Quilt is optional and accepts 1-1000 characters. Visible custom text or example selections are samples unless matching input is entered.</p>
""",
    },
    "personalized-wildlife-deer-couple-in-forest-clearing-comforter-with-buck-4a23dec50b-4a23dec50b": {
        "primary_keyword": "forest clearing deer couple quilt set",
        "secondary_keywords": "deer couple quilt; forest deer bedding; rustic woodland quilt",
        "title_proposed": "Forest Clearing Deer Couple Quilt Set",
        "meta_title_seo": "Forest Clearing Deer Couple Quilt Set",
        "meta_description_seo": "Shop forest clearing deer couple quilt set with tan woodland artwork, buck and doe, tree roots and sample names.",
        "description_proposed_html": """
<p>A buck and doe stand between tree roots in this tan forest clearing quilt design. You and Me text and sample names make the woodland scene feel like a personalized couple keepsake.</p>
<h2>Design Details</h2>
<ul>
<li>Tan forest clearing quilt artwork with buck and doe.</li>
<li>Tree roots, You and Me text and sample names are visible.</li>
<li>Gallery includes the main quilt view plus size, care, fabric or lifestyle panels where shown.</li>
</ul>
<h2>Personalization and Options</h2>
<p>Choose Quilt Size and Pillowcase Quantity options are available. Customize Your Quilt is optional and accepts 1-1000 characters. Visible names or example selections are samples unless matching input is entered.</p>
""",
    },
    "personalized-wildlife-deer-couple-in-forest-comforter-0b240654a2-0b240654a2": {
        "primary_keyword": "dark forest deer couple quilt set",
        "secondary_keywords": "deer couple quilt; woodland deer bedding; rustic forest quilt",
        "title_proposed": "Dark Forest Deer Couple Quilt Set",
        "meta_title_seo": "Dark Forest Deer Couple Quilt Set",
        "meta_description_seo": "Shop dark forest deer couple quilt set with two deer, butterfly accents, sample name pillows and statement text.",
        "description_proposed_html": """
<p>This dark forest quilt set pairs two deer with butterfly accents and statement text. Sample name pillows appear in the product imagery, giving the design a personalized bedroom look.</p>
<h2>Design Details</h2>
<ul>
<li>Dark forest quilt artwork with two deer.</li>
<li>You and Me We Got This text, butterfly accents and sample name pillows are visible.</li>
<li>Gallery includes the quilt mockup plus size, care, fabric or lifestyle panels where shown.</li>
</ul>
<h2>Personalization and Options</h2>
<p>Choose Quilt Size and Pillowcase Quantity options are listed. Customize Your Quilt is optional and accepts 1-1000 characters. Any visible custom text is sample artwork unless matching input is entered.</p>
""",
    },
    "personalized-wildlife-deer-couple-touching-noses-comforter-with-buck-af760d7fee-af760d7fee": {
        "primary_keyword": "deer couple touching noses quilt",
        "secondary_keywords": "deer couple quilt; romantic deer bedding; woodland quilt set",
        "title_proposed": "Deer Couple Touching Noses Quilt",
        "meta_title_seo": "Deer Couple Touching Noses Quilt",
        "meta_description_seo": "Shop deer couple touching noses quilt with autumn woodland artwork, romantic text and sample name pillows.",
        "description_proposed_html": """
<p>Two deer touching noses create the central romantic moment on this autumn woodland quilt. All of Me Loves All of You text and sample name pillows reinforce the couple-focused design.</p>
<h2>Design Details</h2>
<ul>
<li>Autumn woodland quilt with two deer touching noses.</li>
<li>Romantic text and sample name pillows are visible in the artwork.</li>
<li>Gallery includes the main quilt view plus size, care, fabric or lifestyle panels where shown.</li>
</ul>
<h2>Personalization and Options</h2>
<p>Choose Quilt Size and Pillowcase Quantity options are available. Customize Your Quilt is optional and accepts 1-1000 characters. Visible text and names are samples unless matching input is entered.</p>
""",
    },
}

ALT_REPLACEMENTS = {
    "comforter set": "quilt set",
    "comforter": "quilt",
}

KEYWORD_REPLACEMENTS = {
    "You and Me deer comforter set": "You and Me deer quilt set",
    "deer comforter set": "deer quilt set",
    "buck doe comforter": "buck doe quilt",
    "forest clearing deer couple comforter set": "forest clearing deer couple quilt set",
    "deer couple comforter": "deer couple quilt",
    "rustic woodland comforter": "rustic woodland quilt",
    "dark forest deer couple comforter set": "dark forest deer couple quilt set",
    "rustic forest comforter": "rustic forest quilt",
    "deer couple touching noses comforter": "deer couple touching noses quilt",
    "woodland comforter set": "woodland quilt set",
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def replace_terms(text):
    if text is None:
        return text
    out = str(text)
    pairs = [
        ("Comforter Set", "Quilt Set"),
        ("comforter set", "quilt set"),
        ("Comforter", "Quilt"),
        ("comforter", "quilt"),
    ]
    for old, new in pairs:
        out = out.replace(old, new)
    return out


def clean_html(html):
    soup = BeautifulSoup(html.strip(), "html.parser")
    return str(soup)


def main():
    wb = load_workbook(SRC)

    ws = wb["SEO_Products"]
    h = headers(ws)
    changed_products = []
    for row in range(2, ws.max_row + 1):
        handle = ws.cell(row, h["Handle"]).value
        if handle not in FIXES:
            continue
        fix = FIXES[handle]
        for col, value in fix.items():
            ws.cell(row, h[col]).value = clean_html(value) if col == "description_proposed_html" else value
        ws.cell(row, h["meta_title_length"]).value = len(fix["meta_title_seo"])
        ws.cell(row, h["meta_description_length"]).value = len(fix["meta_description_seo"])
        ws.cell(row, h["meta_keyword"]).value = fix["primary_keyword"]
        ws.cell(row, h["revision"]).value = "r5"
        ws.cell(row, h["issues"]).value = "Resolved QA r4 MAJOR product-type wording: proposal now uses Quilt/Quilt Set consistently; pending QA recheck."
        changed_products.append(handle)

    ws = wb["Image_Audit"]
    h = headers(ws)
    changed_images = 0
    for row in range(2, ws.max_row + 1):
        handle = ws.cell(row, h["Handle"]).value
        if handle not in FIXES:
            continue
        for col in ["observed_visual_details", "alt_proposed"]:
            ws.cell(row, h[col]).value = replace_terms(ws.cell(row, h[col]).value)
        ws.cell(row, h["revision"]).value = "r5"
        ws.cell(row, h["issues"]).value = "Resolved QA r4 product-type wording: alt/observation uses Quilt instead of Comforter."
        changed_images += 1

    ws = wb["Product_Evidence"]
    h = headers(ws)
    fixed_ids = {f"evidence_batch_032_{n}" for n in range(311, 315)}
    for row in range(2, ws.max_row + 1):
        eid = ws.cell(row, h["evidence_id"]).value
        if eid not in fixed_ids:
            continue
        for col in ["verified_product_facts", "confidence_and_reason", "proposed_field_to_fact_map"]:
            ws.cell(row, h[col]).value = replace_terms(ws.cell(row, h[col]).value)
        ws.cell(row, h["factual_conflicts"]).value = "QA r4 found source conflict: live title/handle used Comforter while admin Type, size options, admin SEO baseline and contact-sheet info panels support Quilt/Quilt Set. R5 standardizes proposed SEO fields to Quilt/Quilt Set."

    ws = wb["Keyword_Map"]
    h = headers(ws)
    changed_keywords = 0
    for row in range(2, ws.max_row + 1):
        pk = str(ws.cell(row, h["product_key"]).value or "")
        if not any(handle in pk for handle in FIXES):
            continue
        kw = ws.cell(row, h["keyword"]).value
        if kw in KEYWORD_REPLACEMENTS:
            ws.cell(row, h["keyword"]).value = KEYWORD_REPLACEMENTS[kw]
            changed_keywords += 1
        for col in ["decision_reason", "supporting_fact_ids", "demand_evidence", "mapping_reason"]:
            ws.cell(row, h[col]).value = replace_terms(ws.cell(row, h[col]).value)
        ws.cell(row, h["mapping_version"]).value = "r5"

    ws = wb["Buyer_Search_Research"]
    h = headers(ws)
    for row in range(2, ws.max_row + 1):
        pk = str(ws.cell(row, h["product_key"]).value or "")
        if not any(handle in pk for handle in FIXES):
            continue
        for col in ["purchase_context", "jtbd_statement", "functional_motivation", "customer_language", "seo_application", "limitations"]:
            ws.cell(row, h[col]).value = replace_terms(ws.cell(row, h[col]).value)

    ws = wb["README_QA"]
    h = headers(ws)
    ws.append([
        "revision_note_qa_batch_032_r5",
        datetime.now().astimezone().isoformat(timespec="seconds"),
        "R5 revises only batch 032 products 311-314 after QA r4: source evidence supports Quilt/Quilt Set over Comforter for proposed SEO fields and image alt text.",
    ])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    OUT_MD.write_text(
        """# SEO Product Optimization Revision - qa_batch_032_r5

Scope: batch 032, products 311-320.

Source revision: `qa_batch_032_r4`.

QA source: `resutls/jeminise.com/20260906_234129/qa/20260909_132406/SEO_QA_qa_batch_032_r4.md`.

Changes made:
- Revised only the 4 non-passing products `311-314`.
- Confirmed the best public SEO product type should be `Quilt`/`Quilt Set`, not `Comforter`, because live `.js` product type, Shopify admin export `Type`, size option labels, current admin SEO title, and contact-sheet info panels all support quilt terminology.
- Updated `title_proposed`, `meta_title_seo`, `meta_description_seo`, `description_proposed_html`, `primary_keyword`, `secondary_keywords`, and `meta_keyword` for those 4 products.
- Updated all 32 related `Image_Audit` rows so observations and alt text use `quilt` instead of `comforter`.
- Updated related `Keyword_Map`, `Product_Evidence`, and `Buyer_Search_Research` wording for consistency.
- Preserved the 6 products that already passed QA.
- No `APPROVED` status was created.

QA focus for next check:
- Recheck that products `311-314` no longer trigger product-type consistency MAJOR.
- Recheck all 70 images in batch 032 and confirm no alt still uses `comforter` for the deer quilt products.
""",
        encoding="utf-8-sig",
    )
    print(
        {
            "output": str(OUT_XLSX),
            "summary": str(OUT_MD),
            "products_changed": len(changed_products),
            "images_changed": changed_images,
            "keywords_changed": changed_keywords,
        }
    )


if __name__ == "__main__":
    main()
