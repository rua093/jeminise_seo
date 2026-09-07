import csv
import hashlib
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_006_r3"
QA_RUN_ID = "20260907_161639"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_006_r2" / "SEO_Product_Optimization_qa_batch_006_r2.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
CUSTOMIZER_FIELDS_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_fields_audit.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_006_r3.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_006_r3.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    51: {
        "title": "Personalized Christian Scripture Floral Blanket",
        "meta_title": "Personalized Christian Scripture Floral Blanket",
        "meta_description": "Customize a Christian scripture floral blanket with required name, pink EMILY artwork, blanket sizes and verified text fields.",
        "primary": "personalized Christian scripture floral blanket",
        "secondary": "custom Bible verse floral blanket, Christian blanket with name, personalized scripture blanket",
        "cluster": "personalized Christian scripture floral blanket",
        "intent_role": "Design D10 with verified required Custom Name field up to 1000 characters.",
        "detail": "pink EMILY Christian scripture floral blanket artwork",
    },
    52: {
        "title": "Personalized Emily God Says I Am Blanket",
        "meta_title": "Personalized Emily God Says I Am Blanket",
        "meta_description": "Customize an Emily God Says I Am blanket with required name, required color choice, blanket sizes and verified text fields.",
        "primary": "personalized Emily God Says I Am blanket",
        "secondary": "God Says I Am blanket with name, custom Christian affirmation blanket, Emily scripture blanket",
        "cluster": "personalized Emily God Says I Am blanket",
        "intent_role": "Design D9 with verified required Custom Name up to 1000 characters and required Choose Color group.",
        "detail": "blue EMILY God Says I Am scripture blanket artwork with selectable color choices",
    },
    53: {
        "title": "Custom Exploding Soccer Ball Comforter Set",
        "meta_title": "Custom Exploding Soccer Ball Comforter Set",
        "meta_description": "Customize an exploding soccer ball comforter set with optional custom text, product type, size, pillowcase and sheet-cover choices.",
        "primary": "custom exploding soccer ball comforter set",
        "secondary": "exploding soccer comforter, custom soccer bedding, soccer ball comforter set",
        "cluster": "custom exploding soccer ball comforter set",
        "intent_role": "A10 exploding soccer page with verified optional Customize Your Item field up to 1000 characters.",
        "detail": "JACKSON exploding soccer ball comforter artwork with sample text only as an example",
    },
    54: {
        "title": "Custom Fiery Soccer Ball Comforter Set",
        "meta_title": "Custom Fiery Soccer Ball Comforter Set",
        "meta_description": "Customize a fiery soccer ball comforter set with optional custom text, product type, size, pillowcase and sheet-cover choices.",
        "primary": "custom fiery soccer ball comforter set",
        "secondary": "fiery soccer comforter, custom soccer ball bedding, soccer comforter with text",
        "cluster": "custom fiery soccer ball comforter set",
        "intent_role": "Design 5 fiery soccer page with verified optional Customize Your Item field up to 1000 characters; source title typo avoided.",
        "detail": "KENZO fiery soccer ball comforter artwork with sample text only as an example",
    },
    55: {
        "title": "Custom Flaming Soccer Ball Comforter Set",
        "meta_title": "Custom Flaming Soccer Ball Comforter Set",
        "meta_description": "Customize a flaming soccer ball comforter set with optional custom text, product type, size, pillowcase and sheet-cover choices.",
        "primary": "custom flaming soccer ball comforter set",
        "secondary": "flaming soccer ball bedding, custom soccer comforter set, soccer comforter with name",
        "cluster": "custom flaming soccer ball comforter set",
        "intent_role": "A12 flaming soccer ball page separated from product 56 by explicit ball-focused wording.",
        "detail": "Michael 07 flaming soccer ball comforter artwork with sample text only as an example",
    },
    56: {
        "title": "Custom Flaming Soccer Bedding Set",
        "meta_title": "Custom Flaming Soccer Bedding Set",
        "meta_description": "Customize a flaming soccer bedding set with optional custom text, product type, size, pillowcase and sheet-cover choices.",
        "primary": "custom flaming soccer bedding set",
        "secondary": "flaming soccer bedding, custom soccer comforter with name, MICHAEL 07 soccer bedding",
        "cluster": "custom flaming soccer bedding set",
        "intent_role": "Flaming soccer bedding page separated from product 55 by broader bedding-set wording.",
        "detail": "MICHAEL 07 flaming soccer bedding artwork with sample text only as an example",
    },
    57: {
        "title": "Personalized Proverbs 31 Floral Butterfly Blanket",
        "meta_title": "Personalized Proverbs 31 Floral Butterfly Blanket",
        "meta_description": "Customize a Proverbs 31 floral butterfly blanket with required name, inspirational artwork and blanket size choices.",
        "primary": "personalized Proverbs 31 floral butterfly blanket",
        "secondary": "Proverbs 31 blanket with name, custom floral butterfly blanket, personalized inspirational blanket",
        "cluster": "personalized Proverbs 31 floral butterfly blanket",
        "intent_role": "Design 8 with verified required Custom Name field up to 1000 characters.",
        "detail": "Proverbs 31 floral butterfly inspirational blanket artwork with sample Sophia name",
    },
    58: {
        "title": "Personalized Floral Bible Verse Blanket",
        "meta_title": "Personalized Floral Bible Verse Blanket",
        "meta_description": "Customize a floral Bible verse blanket with required name, SARA artwork, blanket sizes and verified text fields.",
        "primary": "personalized floral Bible verse blanket",
        "secondary": "custom Bible verse blanket, floral Christian blanket with name, personalized inspirational blanket",
        "cluster": "personalized floral Bible verse blanket",
        "intent_role": "Design D3 with verified required Custom Name field up to 1000 characters.",
        "detail": "SARA floral butterfly Bible verse blanket artwork",
    },
    59: {
        "title": "Custom Purple Floral Cross Bible Verse Blanket",
        "meta_title": "Custom Purple Floral Cross Bible Verse Blanket",
        "meta_description": "Customize a purple floral cross Bible verse blanket with optional name, Christian artwork and blanket size choices.",
        "primary": "custom purple floral cross Bible verse blanket",
        "secondary": "purple floral cross blanket, custom Christian cross blanket, Bible verse blanket with name",
        "cluster": "custom purple floral cross Bible verse blanket",
        "intent_role": "Design D13 with verified optional Custom Name field up to 1000 characters.",
        "detail": "purple floral cross Bible verse blanket artwork",
    },
    60: {
        "title": "Custom Floral Cross Butterfly Bible Verse Blanket",
        "meta_title": "Custom Floral Cross Butterfly Bible Verse Blanket",
        "meta_description": "Customize a floral cross butterfly Bible verse blanket with optional name, Christian artwork and blanket size choices.",
        "primary": "custom floral cross butterfly Bible verse blanket",
        "secondary": "floral cross butterfly blanket, custom Bible verse blanket, Christian butterfly blanket",
        "cluster": "custom floral cross butterfly Bible verse blanket",
        "intent_role": "Design D6 with verified optional Custom Name field up to 1000 characters.",
        "detail": "floral cross with butterflies Bible verse blanket artwork",
    },
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def load_qa():
    data = json.loads(QA_DATASET.read_text(encoding="utf-8"))
    products = data["QA_Products"]
    handle_by_key = {item["product_key"]: item["handle"] for item in products}
    pos_by_key = {item["product_key"]: int(item["inventory_position"]) for item in products}
    image_by_qa_key = {img["qa_image_key"]: img for img in data["QA_Images"]}
    image_fix_by_key = {}
    observation_by_key = {}
    for image in data["QA_Images"]:
        observation_by_key[(image["product_key"], text(image["media_id"]))] = image["qa_observation"]
    for issue in data["QA_Issues"]:
        if issue["field"].startswith("image_") and issue["recommended_fix"]:
            image = image_by_qa_key.get(issue["qa_image_key"])
            if image:
                image_fix_by_key[(image["product_key"], text(image["media_id"]))] = issue["recommended_fix"]
    return {
        "data": data,
        "product_keys": [item["product_key"] for item in products],
        "handle_by_key": handle_by_key,
        "pos_by_key": pos_by_key,
        "image_fix_by_key": image_fix_by_key,
        "observation_by_key": observation_by_key,
    }


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Body (HTML)": first.get("Body (HTML)", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    return admin


def load_field_audit():
    return {int(item["inventory_position"]): item for item in json.loads(CUSTOMIZER_FIELDS_AUDIT.read_text(encoding="utf-8"))}


def field_sentence(pos, field_audit):
    item = field_audit.get(pos, {})
    fields = item.get("fields", [])
    text_inputs = [entry["fields"] for entry in fields if ".textInputs[" in entry.get("path", "")]
    option_groups = [entry["fields"] for entry in fields if re.search(r"\.optionGroups\[\d+\]$", entry.get("path", ""))]
    parts = []
    for text_input in text_inputs:
        required = "required" if text_input.get("required") else "optional"
        label = text_input.get("label", "text field")
        limit = text_input.get("maxLength")
        if limit:
            parts.append(f"Live Customizer shows a {required} {label} field up to {limit} characters.")
        else:
            parts.append(f"Live Customizer shows a {required} {label} field.")
    if option_groups:
        labels = ", ".join(group.get("label", "") for group in option_groups[:4] if group.get("label"))
        if labels:
            parts.append(f"Additional verified option groups include: {labels}.")
    return " ".join(parts) or "No text-entry field is used in the SEO claim for this revision."


def compact_observation_to_alt(observation, fallback):
    value = re.sub(r"[:;].*$", "", text(observation)).strip()
    value = re.sub(r"\s+", " ", value)
    return (value or fallback)[:125]


def description(pos, admin_row, ctext):
    item = PRODUCT_UPDATES[pos]
    options = ", ".join([v for v in [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]] if v])
    product_type = admin_row["Type"].lower() if admin_row["Type"] else "product"
    component_note = (
        "Use the live selectors to confirm product type, size, pillowcase and sheet-cover choices before checkout."
        if "comforter" in product_type
        else "Use the live selector to confirm blanket size and any required or optional text field before checkout."
    )
    return (
        f"<p>{item['detail'].capitalize()} gives this Jeminise {product_type} a product-specific design focus.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork focus: {item['detail']}.</li>"
        f"<li>Current Shopify export title: {admin_row['Title']}.</li>"
        "<li>Gallery images include the main mockup plus detail, lifestyle, care, feature, size or material panels where shown.</li></ul>"
        "<h3>Options and Customization</h3>"
        f"<ul><li>Available option groups from Shopify export: {options}.</li>"
        f"<li>{ctext}</li>"
        f"<li>{component_note}</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    qa = load_qa()
    field_audit = load_field_audit()
    admin = load_admin(set(qa["handle_by_key"].values()))
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ctext = field_sentence(pos, field_audit)

        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or ws.cell(row_num, idx["product_type"]).value
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, ctext)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"Buyer intent targets {item['cluster']} for US English product search; no search-volume claim is made."
        )
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R3 after Sang re-QA r2: used products_export_1.csv (sha256:{admin_hash}) as admin baseline; "
            "rewrote generic descriptions, removed blanket pillowcase/sham wording, mapped verified Customizer rules, "
            "refreshed image observations/alt, and separated overlapping keyword clusters. Still NEEDS_REVIEW; no APPROVED/import."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        media_id = text(ws.cell(row_num, idx["media_id"]).value)
        key = (pk, media_id)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation = qa["observation_by_key"].get(key)
        if observation:
            ws.cell(row_num, idx["observed_visual_details"]).value = observation
        fallback = f"{PRODUCT_UPDATES[pos]['title']} product image"
        ws.cell(row_num, idx["alt_proposed"]).value = (
            qa["image_fix_by_key"].get(key) or compact_observation_to_alt(observation, fallback)
        )[:125]
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R3: Sang QA r2 image observation/alt applied; current admin alt and image URL matched from products_export_1.csv where possible."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "Product"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = "Sang re-QA r2 SERP evidence + products_export_1.csv admin baseline + customizer_fields_audit.json"
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = "R3 separates sibling product intent by motif and verified Customizer rules."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r3"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as a specific bedding product."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm artwork, product type, size choices "
            "and verified customizer rules."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = "Confirm size, product type, verified text fields, sample values, care panels and image accuracy."
        ws.cell(row_num, idx["emotional_social_motivation"]).value = "Choose a faith, soccer or inspirational bedding gift with clear artwork."
        ws.cell(row_num, idx["purchase_concerns"]).value = "Exact Customizer rules, included components, fabric/care claims, blanket versus comforter options and image accuracy."
        ws.cell(row_num, idx["source_refs"]).value = f"products_export_1.csv sha256:{admin_hash}; Sang QA r2 {QA_RUN_ID}; customizer_fields_audit.json"
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No search-volume source supplied; Shopify admin CSV baseline is now available for stored title/meta/body/image alt."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = qa["product_keys"][pos - 51]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = f"Storefront/product.js from Sang QA r2; products_export_1.csv sha256:{admin_hash}; customizer_fields_audit.json"
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={admin_row['Option1 Name']}, "
            f"{admin_row['Option2 Name']}, {admin_row['Option3 Name']}; status={admin_row['Status']}; "
            f"image_count={len(admin_row['images'])}; design={item['detail']}; customizer={field_sentence(pos, field_audit)}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "R3 applied Sang QA r2 fixes; needs independent QA recheck."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R3 uses QA-observed images, detailed Customizer field audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_006_r3_revision", "r3", "Separate 10-product revision after Sang re-QA r2; source r2 workbook was not modified."),
        ("qa_batch_006_r3_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_006_r3_scope", "inventory positions 51-60", "No products outside qa_batch_006 were revised."),
        ("qa_batch_006_r3_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_006_r3_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_006_r3_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa_r2": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_006 only; inventory positions 51-60",
        "revision": "r3",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Sang QA lại qa_batch_006_r3 before any approval or import.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_006_r3",
                "",
                "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r2: `{SOURCE.relative_to(ROOT)}`",
                f"- QA r2 của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r3: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Sửa trọng yếu: mapped customizer fields, removed blanket pillowcase wording, separated 51-52 and 53-56 keyword clusters.",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Bước tiếp theo: Sang QA lại `qa_batch_006_r3`.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
