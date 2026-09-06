import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_003"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_002.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_003.xlsx"


PROFILES = {
    21: ("cream gingerbread man and candy cane Christmas quilt", "gingerbread candy cane quilt set", "Gingerbread Candy Cane Christmas Quilt Set", "Gingerbread Candy Cane Christmas Quilt Set", "Cream Christmas quilt with a gingerbread man, striped candy cane, green bow and decorative red-green border.", "gingerbread Christmas quilt, candy cane bedding, Christmas cookie quilt set"),
    22: ("black Christmas tree quilt with ornaments and gifts", "black Christmas tree quilt set", "Black Christmas Tree Quilt Set", "Black Christmas Tree Quilt Set", "Dark holiday quilt with a decorated Christmas tree, ornaments, wrapped gifts and matching sham artwork.", "Christmas tree bedding, holiday tree quilt, Christmas gift quilt set"),
    23: ("winter cardinal birdhouse quilt in soft blue and cream", "winter cardinal quilt set", "Winter Cardinal Birdhouse Quilt Set", "Winter Cardinal Birdhouse Quilt Set", "Soft winter quilt design with red cardinals, a snowy birdhouse, wreath detail and matching pillow sham artwork.", "cardinal birdhouse quilt, winter bird quilt, Christmas cardinal bedding"),
    24: ("snowman Christmas village quilt with gifts and cookies", "snowman Christmas quilt set", "Snowman Christmas Quilt Set", "Snowman Christmas Village Quilt Set", "Colorful Christmas quilt with a large snowman, wrapped gifts, snowy village background and cookie detail.", "snowman bedding set, Christmas village quilt, holiday snowman quilt"),
    25: ("red snowman quilt with cardinal birds and pine trees", "red snowman Christmas quilt", "Red Snowman Christmas Quilt Set", "Red Snowman and Cardinal Christmas Quilt", "Red-toned Christmas quilt featuring a snowman in a top hat and scarf, cardinal birds, pine trees and holly accents.", "snowman cardinal bedding, red Christmas quilt, snowman holiday quilt"),
    26: ("white Christmas tree quilt with red and gold stars", "white Christmas tree quilt set", "White Christmas Tree Quilt Set", "White Christmas Tree Quilt Set", "Light Christmas quilt with a black and gray tree, red and gold star ornaments and black-white check border.", "black white Christmas quilt, Christmas star quilt, holiday tree bedding"),
    27: ("red snowman pair and cardinal Christmas quilt", "snowman cardinal Christmas quilt", "Snowman Cardinal Christmas Quilt Set", "Snowman and Cardinal Christmas Quilt Set", "Deep red Christmas quilt with two snowmen in a round frame, red and brown cardinals, pine greenery and matching sham.", "snowman cardinal bedding, red snowman quilt, Christmas cardinal quilt"),
    28: ("vintage cream Christmas tree quilt with wrapped gifts", "vintage Christmas tree quilt set", "Vintage Christmas Tree Quilt Set", "Vintage Christmas Tree Quilt Set", "Cream vintage-style Christmas quilt with a decorated tree, wrapped gifts, ornate border and matching pillow sham.", "cream Christmas tree bedding, Christmas gift quilt, holiday quilt set"),
    29: ("Christmas cardinal memorial quilt with I am Always With You text", "Christmas cardinal memorial quilt", "Christmas Cardinal Memorial Quilt Set", "Christmas Cardinal Memorial Patchwork Quilt", "Christmas patchwork quilt with a red cardinal, I am Always With You text, red flowers, bells, hearts and berries.", "I am always with you quilt, cardinal remembrance quilt, memorial Christmas bedding"),
    30: ("personalized Christmas cardinal quilt with the name Sophia", "personalized Christmas cardinal quilt", "Personalized Christmas Cardinal Quilt Set", "Personalized Cardinal Christmas Quilt Set", "Personalized Christmas quilt with the name Sophia, red cardinal, pine branches, berries, pinecones and matching sham.", "custom name cardinal bedding, personalized cardinal quilt, Christmas bird quilt set"),
}

REFS = {
    "gingerbread": "https://www.amazon.com/christmas-candy-bedding/s?k=christmas+candy+bedding; https://www.target.com/s/gingerbread%2Bbedding; https://levtexhome.com/products/gingerbread-village-quilt",
    "tree": "https://www.amazon.com/; https://retrobarn.com/; https://www.bedbathandbeyond.com/c/christmas-bedding/christmas-bedding-sets?t=28287",
    "cardinal": "https://www.etsy.com/listing/4505275205/personalized-cardinal-memorial-quilt; https://www.walmart.com/ip/Lotusprinthandmade-Personalized-Cardinal-Memorial-Quilt-Bedding-Set-Fold-Hem-5-Sizes-Made-in-Vietnam-70x80-cardinal-memorial-quilt-bed-set/16670953823; https://www.ebay.com/",
    "snowman": "https://www.amazon.com/Christmas-Quilt-Sets/s?k=Christmas+Quilt+Sets; https://www.bedbathandbeyond.com/c/christmas-bedding/christmas-bedding-sets?t=28287; https://www.ebay.com/",
}


def append_dict(ws, data):
    headers = [cell.value for cell in ws[1]]
    ws.append([data.get(h, "") for h in headers])


def ref_for(pos):
    if pos in (21,):
        return REFS["gingerbread"]
    if pos in (22, 26, 28):
        return REFS["tree"]
    if pos in (23, 29, 30):
        return REFS["cardinal"]
    return REFS["snowman"]


def image_obs(theme, pos):
    if pos == 1:
        return f"Main bed mockup showing {theme}."
    if pos == 2:
        return f"Angled bedroom mockup showing the {theme} quilt and matching shams."
    if pos == 3:
        return f"Matching pillow sham with {theme} artwork."
    if pos == 4:
        return f"Close-up view of the printed/quilted texture and {theme} detail."
    return "Included-components and size information image for the quilt set."


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        theme, primary, seo_title, h1, meta_detail, secondary = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        refs = ref_for(pos)
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        meta = f"{meta_detail} Review size and sham options before checkout."
        desc = (
            f"<p>Give the bedroom a festive seasonal update with this Jeminise {kind} featuring {theme}. "
            f"The same artwork appears across the quilt and matching sham images for a coordinated holiday look.</p>"
            f"<h3>Design Details</h3><ul><li>{meta_detail}</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>Search Intent</h3><p>This draft targets shoppers looking for a specific Christmas quilt design rather than a broad bedding category. It requires QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the visible design: {theme}. Keep broad Christmas bedding terms for collection pages.",
            "season": "CHRISTMAS", "buyer_search_summary": f"Buyer is likely looking for Christmas bedding with {theme} for seasonal decor or gifting.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": "DRAFTED", "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment.",
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": "DRAFTED", "confidence_and_reason": "Medium: product JSON, HTML and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for Christmas bedding with {theme}.",
            "jtbd_statement": f"When decorating for Christmas, the buyer wants a coordinated {kind} with a specific holiday motif that fits the room and gift occasion.",
            "functional_motivation": "Find the right holiday design, size and included sham options.",
            "emotional_social_motivation": "Create a festive, sentimental or personalized bedroom setting.",
            "purchase_concerns": "Fit, included components, material/care claims, personalization accuracy and whether the motif matches the intended room.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": theme, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible holiday motif and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": "CHRISTMAS",
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            h = h1
            alt = {
                1: f"{h} displayed on a bed",
                2: f"Angled bedroom view of {h}",
                3: f"Matching pillow sham for {h}",
                4: f"Close-up of {theme} on the quilt",
                5: f"Size guide for {h}",
            }.get(ipos, f"Product detail image for {h}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(theme, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_003_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_003_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_003"},
        {"metric": "cumulative_products", "value": "30", "definition": "Total products included through batch_003"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[30:40]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_004",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
