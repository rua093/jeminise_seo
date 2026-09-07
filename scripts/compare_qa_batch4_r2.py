from pathlib import Path
import json
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / "seo_runs" / "jeminise.com" / "20260906_234129"
ORIGINAL = ROOT / "resutls" / "jeminise.com" / "20260906_234129" / "batches" / "SEO_Product_Optimization_through_batch_034.xlsx"
REVISION = ROOT / "resutls" / "jeminise.com" / "20260906_234129" / "revisions" / "qa_batch_004_r2" / "SEO_Product_Optimization_qa_batch_004_r2.xlsx"
OUTPUT = RUN / "qa" / "20260907_144539" / "revision_diff.json"

def rows(path, sheet):
    wb = load_workbook(path, read_only=True, data_only=False)
    ws = wb[sheet]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    return [dict(zip(headers, (c.value for c in row))) for row in ws.iter_rows(min_row=2)]

def main():
    inventory = json.loads((RUN / "inventory.json").read_text(encoding="utf-8-sig"))
    selected = inventory[30:40]
    keys = {x["product_key"] for x in selected}; handles = {x["Handle"] for x in selected}; urls = {x["product_url"] for x in selected}
    specs = {
        "SEO_Products": (lambda r: r.get("product_key") in keys, lambda r: r.get("product_key")),
        "Image_Audit": (lambda r: r.get("Handle") in handles, lambda r: (r.get("Handle"), r.get("media_id"), r.get("image_number"))),
        "Product_Evidence": (lambda r: r.get("product_url") in urls, lambda r: r.get("evidence_id")),
        "Keyword_Map": (lambda r: r.get("product_key") in keys, lambda r: (r.get("product_key"), r.get("keyword"), r.get("keyword_role"))),
        "Buyer_Search_Research": (lambda r: r.get("product_key") in keys, lambda r: r.get("research_id")),
    }
    result = {"source": str(ORIGINAL), "revision": str(REVISION), "sheets": {}}
    for sheet, (include, identify) in specs.items():
        before = {str(identify(r)): r for r in rows(ORIGINAL, sheet) if include(r)}
        after = {str(identify(r)): r for r in rows(REVISION, sheet) if include(r)}
        changes = []
        for rid in sorted(set(before) | set(after)):
            old, new = before.get(rid, {}), after.get(rid, {})
            fields = [{"field": f, "before": old.get(f), "after": new.get(f)} for f in sorted(set(old) | set(new)) if old.get(f) != new.get(f)]
            if fields: changes.append({"row_id": rid, "changes": fields})
        result["sheets"][sheet] = {"before_rows": len(before), "after_rows": len(after), "changed_rows": len(changes), "changes": changes}
    OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({s:{k:v[k] for k in ("before_rows","after_rows","changed_rows")} for s,v in result["sheets"].items()}, indent=2))

if __name__ == "__main__":
    main()
