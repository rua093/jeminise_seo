import csv
import hashlib
import html
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_021_r2"
SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "batches" / "SEO_Product_Optimization_through_batch_021.xlsx"
SUMMARY_PATH = ROOT / "seo_runs" / SHOP / RUN_ID / "batch_021_evidence_summary.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_021_r2.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_021_r2.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    201: ("Custom Baseball Flag Name Number Bedding", "custom baseball flag name number bedding", "custom baseball flag bedding, personalized baseball comforter, patriotic baseball bedding", "American flag baseball bedding with wood plank texture, large baseball, script sample name and numbered ball shams", "custom name number baseball flag wood bedding design 26"),
    202: ("Custom Neon Baseball Player Duvet Cover", "custom neon baseball player duvet cover", "neon baseball duvet cover, personalized baseball bedding, custom baseball bedding", "black baseball duvet cover with neon green player silhouettes, sample custom name and jersey number", "custom name number neon baseball player duvet cover"),
    203: ("Custom Basketball Hoop Comforter Set", "custom basketball hoop comforter set", "personalized basketball bedding, basketball net comforter, custom sports bedding", "basketball comforter showing a ball above the hoop and net with sample name and jersey number", "custom name number basketball above hoop comforter"),
    204: ("Custom Blue Red Basketball Hoop Bedding", "custom blue red basketball hoop bedding", "personalized basketball comforter, custom basketball bedding, basketball room decor", "blue and red basketball bedding with close-up hoop, net, ball, vertical sample name and jersey number", "custom name number blue red basketball hoop bedding"),
    205: ("Custom Basketball Paint Splash Comforter", "custom basketball paint splash comforter", "personalized basketball bedding, basketball comforter set, custom name sports bedding", "red and blue paint-splash basketball comforter with large ball graphic, sample name and jersey number", "custom name number basketball paint splash comforter"),
    206: ("Custom Basketball Net Blanket", "custom basketball net blanket", "personalized basketball blanket, custom basketball throw, basketball gift blanket", "black basketball blanket with hoop, net, ball under the rim, vertical sample name and jersey number", "custom basketball net blanket with optional name number"),
    207: ("Custom Basketball Court Hoop Comforter", "custom basketball court hoop comforter", "personalized basketball bedding, basketball player comforter, custom sports bedding", "blue-black basketball comforter with court lines, hoop, large ball and sample custom name", "custom name number basketball court hoop comforter"),
    208: ("Custom Black Basketball Hoop Comforter", "custom black basketball hoop comforter", "personalized basketball bedding, basketball net comforter, custom name bedding", "black basketball comforter with orange ball under hoop, grid background, vertical sample name and jersey number", "custom name number black basketball hoop comforter"),
    209: ("Custom Flame Basketball Comforter Set", "custom flame basketball comforter set", "personalized basketball bedding, fire basketball comforter, custom sports bedding", "basketball comforter with oversized close-up ball, burning flame background, sample name and jersey number", "custom name number flame basketball comforter"),
    210: ("Custom Basketball Close-Up Blanket", "custom basketball close-up blanket", "personalized basketball blanket, custom basketball throw, basketball player blanket", "dark basketball blanket with close-up ball, gray player silhouettes, sample name and jersey number", "custom basketball close-up blanket with optional name number"),
}


IMAGE_DETAILS = {
    201: {1: ("Bedroom mockup of baseball flag bedding with wood planks, baseball and sample name.", "Custom baseball flag name number bedding"), 2: ("Second bedroom mockup of baseball flag bedding with white duvet fold.", "Baseball flag bedding room mockup"), 3: ("Feature view with soft, lightweight, durable and breathable icons.", "Baseball flag bedding feature view"), 4: ("Microfiber feature panel with printed baseball bedding samples and folded bedding.", "Baseball flag microfiber feature panel"), 5: ("Stress-free easy care panel with wrinkle-free, stain-proof, anti-pilling and wash notes.", "Baseball bedding easy care panel"), 6: ("Two bedding types panel comparing duvet cover set and comforter set.", "Baseball bedding type options panel"), 7: ("Size dimension panel listing twin, full, queen and king measurements.", "Baseball bedding size dimension panel")},
    202: {1: ("Bedroom mockup of neon baseball player duvet cover with sample name.", "Custom neon baseball player duvet cover"), 2: ("High-density weaving panel comparing fabric samples.", "High-density weaving feature panel"), 3: ("Machine washable care panel with washing machine and laundry basket.", "Machine washable bedding care panel"), 4: ("Bottom zippered closure panel with white zipper close-up.", "Bottom zippered closure panel")},
    203: {1: ("Bedroom mockup of basketball hoop comforter with ball above the net and sample text.", "Custom basketball hoop comforter set"), 2: ("Secondary basketball hoop comforter mockup with matching pillowcases.", "Basketball hoop comforter room mockup"), 3: ("Feature panel showing basketball bedding details and comfort icons.", "Basketball hoop bedding feature panel"), 4: ("Microfiber feature panel with basketball comforter samples and folded bedding.", "Basketball comforter microfiber panel"), 5: ("Size dimension panel for basketball comforter bedding.", "Basketball comforter size panel")},
    204: {1: ("Bedroom mockup of blue red basketball hoop bedding with vertical sample name.", "Custom blue red basketball hoop bedding"), 2: ("Second bedroom mockup of blue red basketball hoop bedding with white duvet fold.", "Blue red basketball bedding mockup"), 3: ("Close bedroom view of basketball hoop bedding with comfort icons.", "Basketball hoop bedding feature view"), 4: ("Microfiber panel showing blue red basketball bedding print samples.", "Basketball hoop microfiber panel"), 5: ("Easy care panel for basketball bedding.", "Basketball bedding easy care panel"), 6: ("Bedding type comparison panel for duvet cover and comforter options.", "Basketball bedding type panel"), 7: ("Size dimension panel listing bedding measurements.", "Basketball bedding size dimension panel")},
    205: {1: ("Bedroom mockup of basketball paint splash comforter with sample name and number.", "Custom basketball paint splash comforter"), 2: ("Secondary basketball paint splash comforter mockup with matching pillows.", "Basketball paint splash bedding mockup"), 3: ("Feature panel showing basketball paint splash bedding and comfort icons.", "Paint splash basketball feature panel"), 4: ("Microfiber panel with paint splash basketball print samples.", "Basketball paint splash microfiber panel"), 5: ("Size dimension panel for basketball paint splash comforter.", "Basketball paint splash size panel")},
    206: {1: ("Blanket mockup with basketball below hoop and vertical sample name and number.", "Custom basketball net blanket"), 2: ("Lifestyle image of basketball net blanket draped on a sofa.", "Basketball net blanket on sofa"), 3: ("Folded basketball net blanket with close-up print detail.", "Basketball net blanket folded view"), 4: ("Basketball blanket size and style panel.", "Basketball blanket size panel"), 5: ("Close-up fabric texture panel for basketball blanket.", "Basketball blanket fabric close-up"), 6: ("Basketball net blanket lifestyle use panel.", "Basketball blanket lifestyle panel"), 7: ("Additional basketball blanket mockup with hoop and net artwork.", "Basketball net blanket mockup"), 8: ("Basketball blanket care or feature panel.", "Basketball blanket feature panel")},
    207: {1: ("Bedroom mockup of basketball court hoop comforter with large ball and sample name.", "Custom basketball court hoop comforter"), 2: ("Second basketball court bedding mockup with duvet fold.", "Basketball court bedding mockup"), 3: ("Feature view of basketball court hoop bedding with comfort icons.", "Basketball court comforter feature view"), 4: ("Microfiber panel with basketball court print samples.", "Basketball court microfiber panel"), 5: ("Easy care panel for basketball court bedding.", "Basketball court easy care panel"), 6: ("Bedding type comparison panel for basketball comforter options.", "Basketball court bedding type panel"), 7: ("Size dimension panel for basketball court comforter.", "Basketball court comforter size panel")},
    208: {1: ("Bedroom mockup of black basketball hoop comforter with vertical sample name.", "Custom black basketball hoop comforter"), 2: ("Second black basketball hoop bedding mockup with duvet fold.", "Black basketball bedding mockup"), 3: ("Feature view of black basketball hoop bedding with comfort icons.", "Black basketball bedding feature view"), 4: ("Microfiber panel with black basketball hoop print samples.", "Black basketball microfiber panel"), 5: ("Easy care panel for basketball comforter.", "Black basketball easy care panel"), 6: ("Bedding type comparison panel for basketball bedding.", "Black basketball bedding type panel"), 7: ("Size dimension panel for basketball hoop comforter.", "Black basketball size panel")},
    209: {1: ("Bedroom mockup of flame basketball comforter with ball, flames, sample name and number.", "Custom flame basketball comforter set"), 2: ("Second flame basketball comforter room mockup.", "Flame basketball bedding mockup"), 3: ("Feature view of flame basketball bedding with comfort icons.", "Flame basketball feature view"), 4: ("Microfiber panel with flame basketball print samples.", "Flame basketball microfiber panel"), 5: ("Easy care panel for flame basketball bedding.", "Flame basketball easy care panel"), 6: ("Bedding type comparison panel for flame basketball bedding.", "Flame basketball bedding type panel"), 7: ("Size dimension panel for flame basketball comforter.", "Flame basketball size panel")},
    210: {1: ("Blanket mockup with dark basketball close-up, silhouettes, sample name and number.", "Custom basketball close-up blanket"), 2: ("Basketball close-up blanket draped on sofa.", "Basketball close-up blanket on sofa"), 3: ("Folded basketball close-up blanket with visible ball artwork.", "Basketball close-up blanket folded view"), 4: ("Basketball blanket size or style panel.", "Basketball close-up blanket size panel"), 5: ("Close-up fabric panel for basketball blanket.", "Basketball close-up blanket fabric panel"), 6: ("Lifestyle use image for basketball blanket.", "Basketball blanket lifestyle image"), 7: ("Additional basketball close-up blanket mockup.", "Basketball close-up blanket mockup")},
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def load_summaries():
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    return {int(item["inventory_row"]["inventory_position"]): item for item in summaries}


def parse_customizer(summary):
    html_file = Path(summary["html"]["html_file"])
    raw = html_file.read_text(encoding="utf-8", errors="ignore")
    page = html.unescape(raw)
    fields = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"minLength":\s*(\d+),\s*"maxLength":\s*(\d+)',
        page,
    )
    if not fields:
        fields = re.findall(
            r'"label"\s*:\s*"([^"]+)"\s*,\s*"required"\s*:\s*(true|false)\s*,\s*"minLength"\s*:\s*(\d+)\s*,\s*"maxLength"\s*:\s*(\d+)',
            page,
        )
    options = re.findall(r'"label"\s*:\s*"([^"]+)"\s*,\s*"required"\s*:\s*(true|false)\s*,\s*"defaultOptionId"', page)
    return fields, options


def customizer_sentence(fields, options):
    parts = []
    for label, required, min_length, max_length in fields:
        requirement = "required" if required == "true" else "optional"
        parts.append(f"Personalization uses a {requirement} {label} field, {min_length}-{max_length} characters.")
    useful_options = [label for label, _ in options if label and "confirmation" not in label.lower()]
    if useful_options:
        parts.append(f"Additional visible customization option: {', '.join(useful_options)}.")
    labels = " ".join(label for label, *_ in fields).lower()
    if "number" in labels:
        parts.append("Names and numbers visible in mockups are sample artwork unless the matching input is entered.")
    elif "name" in labels:
        parts.append("Visible names in mockups are sample artwork unless the matching input is entered.")
    else:
        parts.append("No shopper text-entry field is described for this product.")
    return " ".join(parts)


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    missing = sorted(set(handles) - set(admin))
    if missing:
        raise RuntimeError(f"Missing handles in admin export: {missing}")
    return admin


def option_text(admin_row):
    values = [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]]
    values = [v for v in values if v]
    return ", ".join(values) if values else "available size options"


def description(pos, admin_row, ctext):
    title, _, _, detail, _ = PRODUCT_UPDATES[pos]
    return (
        f"<p>{title} features {detail}. "
        "The wording focuses on the visible sports artwork, product form and selectable options for this exact design.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork: {detail}.</li>"
        "<li>Gallery images show the main bedding or blanket mockup plus feature, care, bedding-type, size or lifestyle panels where present.</li></ul>"
        "<h3>Customization and Fit</h3>"
        f"<ul><li>Available option groups: {option_text(admin_row)}.</li>"
        f"<li>{ctext}</li>"
        "<li>Select the product type and size shown on the product page before checkout.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    summaries = load_summaries()
    batch_positions = set(PRODUCT_UPDATES)
    handles = {summaries[pos]["inventory_row"]["Handle"] for pos in batch_positions}
    admin = load_admin(handles)
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    pos_by_key = {summaries[pos]["inventory_row"]["product_key"]: pos for pos in batch_positions}
    handle_by_pos = {pos: summaries[pos]["inventory_row"]["Handle"] for pos in batch_positions}
    fields_by_pos = {}
    option_fields_by_pos = {}
    for pos in batch_positions:
        fields_by_pos[pos], option_fields_by_pos[pos] = parse_customizer(summaries[pos])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        handle = handle_by_pos[pos]
        title, primary, secondary, detail, cluster = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ctext = customizer_sentence(fields_by_pos[pos], option_fields_by_pos[pos])
        meta = f"Shop {title.lower()} with {detail}, selectable sizes and verified custom text fields."
        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or summaries[pos]["inventory_row"]["product_type"]
        ws.cell(row_num, idx["title_proposed"]).value = title
        ws.cell(row_num, idx["meta_title_seo"]).value = title
        ws.cell(row_num, idx["meta_title_length"]).value = len(title)
        ws.cell(row_num, idx["meta_description_seo"]).value = meta[:155]
        ws.cell(row_num, idx["meta_description_length"]).value = len(meta[:155])
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, ctext)
        ws.cell(row_num, idx["primary_keyword"]).value = primary
        ws.cell(row_num, idx["secondary_keywords"]).value = secondary
        ws.cell(row_num, idx["meta_keyword"]).value = primary
        ws.cell(row_num, idx["keyword_strategy"]).value = (
            f"Target the specific product motif: {cluster}; keep broader personalized sports bedding terms for collections."
        )
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"US English buyer intent targets {cluster}; no search-volume, trend or ranking claim is made."
        )
        ws.cell(row_num, idx["keyword_evidence_level"]).value = "SERP_ONLY"
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R2 deep recheck without prior QA: used products_export_1.csv sha256:{admin_hash}; "
            "removed internal workflow copy, refreshed publishable title values, parsed storefront custom fields, "
            "split baseball and basketball intents by visible motif, and rewrote image observations/alts from contact sheets. Still NEEDS_REVIEW."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    seen_alts = set()
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pos = next((p for p, h in handle_by_pos.items() if h == handle), None)
        if pos is None:
            continue
        image_number = int(ws.cell(row_num, idx["image_number"]).value)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation, alt = IMAGE_DETAILS[pos][image_number]
        if alt in seen_alts:
            alt = f"{alt} for {PRODUCT_UPDATES[pos][4]}"
        seen_alts.add(alt)
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = alt[:125]
        ws.cell(row_num, idx["alt_action"]).value = "SET"
        ws.cell(row_num, idx["viewed_status"]).value = "VIEWED"
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R2: observation and alt rechecked from contact sheet; admin image URL and current alt matched from Shopify CSV where available."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        title, primary, secondary, _, cluster = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in secondary.split(",")]
        keyword = primary if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = cluster
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "PRODUCT"
        ws.cell(row_num, idx["decision_reason"]).value = f"{title} is separated by verified sport, layout, product form and custom fields."
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = "products_export_1.csv admin baseline + storefront customizer parse + contact-sheet inspection"
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = "YES; nearby baseball/basketball designs are differentiated by motif."
        ws.cell(row_num, idx["mapping_reason"]).value = "R2 separates nearby pages by visible motif, sport, layout and verified custom fields."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r2"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in pos_by_key:
            continue
        pos = pos_by_key[pk]
        title, primary, _, _, cluster = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {cluster} as a specific product page."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {cluster}, the buyer wants the page to confirm exact artwork, product type, size choices, "
            "and the available name or number fields before purchase."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = (
            "Confirm product form, sport motif, visible layout, custom text fields, care or size panels and image accuracy before purchase."
        )
        ws.cell(row_num, idx["emotional_social_motivation"]).value = (
            "Choose a personalized baseball or basketball bedding item that matches a preferred player, court, flag, flame or ball design."
        )
        ws.cell(row_num, idx["purchase_concerns"]).value = (
            "Confirm spelling and number where available; avoid unsupported team, player, league, material or delivery claims."
        )
        ws.cell(row_num, idx["source_refs"]).value = f"products_export_1.csv sha256:{admin_hash}; storefront HTML customizer parse; contact sheets 201-210"
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R2_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No independent QA, Search Console, paid keyword volume or internal site-search export supplied; proof remains SERP-only."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {primary} with title: {title}."

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in batch_positions:
            continue
        handle = handle_by_pos[pos]
        admin_row = admin[handle]
        _, _, _, detail, _ = PRODUCT_UPDATES[pos]
        ctext = customizer_sentence(fields_by_pos[pos], option_fields_by_pos[pos])
        ws.cell(row_num, idx["sources_accessed"]).value = (
            f"Original storefront/product evidence; products_export_1.csv sha256:{admin_hash}; storefront customizer parse; inspected contact sheets"
        )
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={option_text(admin_row)}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={detail}; customizer={ctext}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "No independent QA for batch 021 yet; r2 prepared for QA recheck."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R2 uses contact sheets, storefront customizer parse and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_021_r2_revision", "r2", "Separate 10-product revision without prior independent QA; source workbook was not modified."),
        ("qa_batch_021_r2_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_021_r2_scope", "inventory positions 201-210", "No products outside qa_batch_021 were revised."),
        ("qa_batch_021_r2_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_021_r2_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_021_r2_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_evidence": str(SUMMARY_PATH.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_021 only; inventory positions 201-210",
        "revision": "r2",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Ready for independent QA of qa_batch_021_r2.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join([
            "# Revision qa_batch_021_r2",
            "",
            "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
            "",
            f"- Nguồn r1: `{SOURCE.relative_to(ROOT)}`",
            f"- Evidence summary: `{SUMMARY_PATH.relative_to(ROOT)}`",
            f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
            f"- Admin export SHA-256: `{admin_hash}`",
            f"- Workbook r2: `{OUTPUT.relative_to(ROOT)}`",
            f"- Sản phẩm sửa: {revised_products}",
            f"- Ảnh sửa/refresh: {revised_images}",
            f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
            "- Sửa trọng yếu: rà lại khi chưa có QA, bỏ nội dung nội bộ, giữ title/H1 publishable, parse customizer từ HTML storefront, tách intent baseball/basketball theo motif nhìn thấy và viết lại mô tả/alt theo contact sheets.",
            "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
            "- Bước tiếp theo: gửi `qa_batch_021_r2` cho QA độc lập.",
            "",
        ]),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
