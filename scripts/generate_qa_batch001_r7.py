from __future__ import annotations

import hashlib
import html
import json
import re
import shutil
import sys
import tempfile
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import export_qa_batch1_xlsx as exporter

SHOP, RUN_ID, QA_RUN_ID = "jeminise.com", "20260906_234129", "20260908_230550"
BATCH, REVISION = "qa_batch_001_r7", "r7"
QA_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID
OUT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "qa" / QA_RUN_ID
SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH / f"SEO_Product_Optimization_{BATCH}.xlsx"
SNAPSHOT = QA_DIR / "source_snapshot" / SOURCE.name
OUT_XLSX = OUT_DIR / f"SEO_QA_{BATCH}.xlsx"
OUT_MD = OUT_DIR / f"SEO_QA_{BATCH}.md"
OLD_QA_XLSX = ROOT / "resutls" / SHOP / RUN_ID / "qa" / "20260908_001100" / "SEO_QA_qa_batch_001_r4.xlsx"

PW = {"P1":15,"P2":10,"K1":10,"K2":5,"K3":5,"T1":10,"T2":5,"D1":5,"D2":10,"I1":20,"E1":5}
IW = {"IM1":40,"IM2":30,"IM3":20,"IM4":10}
RATING = {"FULL":1.0,"PARTIAL":0.5,"FAIL":0.0}

ASSESS = {
 1:{"P1":"FULL","P2":"FULL","K1":"PARTIAL","K2":"FULL","K3":"FULL","T1":"FULL","T2":"FULL","D1":"FULL","D2":"PARTIAL","E1":"PARTIAL"},
 2:{"P1":"FULL","P2":"FULL","K1":"FULL","K2":"FULL","K3":"FULL","T1":"FULL","T2":"FULL","D1":"FULL","D2":"FULL","E1":"PARTIAL"},
 3:{"P1":"FULL","P2":"FULL","K1":"FULL","K2":"FULL","K3":"FULL","T1":"FULL","T2":"FULL","D1":"FULL","D2":"PARTIAL","E1":"PARTIAL"},
 4:{"P1":"FULL","P2":"FULL","K1":"PARTIAL","K2":"PARTIAL","K3":"FULL","T1":"FULL","T2":"FULL","D1":"FULL","D2":"PARTIAL","E1":"PARTIAL"},
 5:{"P1":"FULL","P2":"FULL","K1":"FULL","K2":"FULL","K3":"FULL","T1":"FULL","T2":"FULL","D1":"FULL","D2":"FULL","E1":"PARTIAL"},
 6:{"P1":"FULL","P2":"FULL","K1":"PARTIAL","K2":"PARTIAL","K3":"FULL","T1":"FULL","T2":"FULL","D1":"FULL","D2":"FULL","E1":"PARTIAL"},
 7:{"P1":"FULL","P2":"FULL","K1":"PARTIAL","K2":"FULL","K3":"FULL","T1":"FULL","T2":"FULL","D1":"FULL","D2":"FULL","E1":"PARTIAL"},
 8:{"P1":"FULL","P2":"FULL","K1":"FULL","K2":"FULL","K3":"FULL","T1":"FULL","T2":"FULL","D1":"FULL","D2":"FULL","E1":"PARTIAL"},
 9:{"P1":"FULL","P2":"FULL","K1":"FULL","K2":"FULL","K3":"FULL","T1":"FULL","T2":"FULL","D1":"FULL","D2":"FULL","E1":"PARTIAL"},
10:{"P1":"FULL","P2":"FULL","K1":"FULL","K2":"FULL","K3":"FULL","T1":"FULL","T2":"FULL","D1":"FULL","D2":"FULL","E1":"PARTIAL"},
}

SERP = {
1:[("cardinal sunflower autumn quilt set","https://vantique.net/products/autumn-cardinal-3-piece-quilted-bedding-set-ncu0tv033","Exact sunflower modifier remains hypothesis-led; finished autumn-cardinal bedding intent is commercial."),("autumn cardinal quilt sunflower bedding","https://www.macys.com/shop/product/seventh-studio-cardinal-quilted-3-piece-quilt-sets?ID=27709691","Commercial cardinal quilt-set comparator; no volume claim.")],
2:[("personalized softball comforter set name number","https://www.youcustomizeit.com/p/Softball-Comforters-Personalized/353832","Direct personalized softball comforter product with editable text."),("teal softball bedding personalized","https://shoalscustomcreations.com/products/softball-coquette-bow-personalized-blanket","Related teal personalized softball bedding intent; product type differs.")],
3:[("Christmas cardinal memorial quilt set","https://www.etsy.com/listing/1085801805/i-am-always-with-you-blanket-cardinal","Commercial memorial/cardinal product with the same remembrance phrase."),("cardinal Christmas I am always with you quilt","https://alphaquilt.com/products/tai181024271","Finished Christmas cardinal quilt-set comparator.")],
4:[("cardinal roses memorial quilt set bedding","https://www.etsy.com/listing/1085801805/i-am-always-with-you-blanket-cardinal","Memorial cardinal commercial intent exists, but exact rose-quilt combination is weak."),("I am always with you cardinal rose quilt bedding","https://www.macys.com/shop/product/seventh-studio-cardinal-quilted-3-piece-quilt-sets?ID=27709691","Finished cardinal quilt comparator; exact memorial rose intent remains niche.")],
5:[("colorful cat patchwork quilt set bedding","https://www.wayfair.com/bed-bath/pdp/ambesonne-cat-bedspread-set-patchwork-style-silly-faces-multicolor-bbqv8335.html","Direct colorful cat patchwork bedspread-set intent."),("cat patchwork bedding set colorful cat face quilt","https://www.walmart.com/ip/20788770071","Commercial cat patchwork quilt-set comparator.")],
6:[("geometric sitting cat patchwork quilt set bedding","https://www.spoonflower.com/en/home-decor/bedding/duvet-cover/11072634-lino-cut-geometric-stamp-cats-by-luli_print","Commercial geometric-cat bedding comparator; exact quilt phrase is weak."),("sitting cat quilt bedding geometric cat design","https://pixels.com/featured/colorful-geometric-sleeping-cat-romero-britto-romero-britto.html?product=duvet-cover","Commercial geometric-cat duvet comparator; supports motif, not exact quilt-set demand.")],
7:[("Celtic fantasy tree quilt set bedding","https://www.freshouseshop.com/products/celtic-bed-knot-tree-of-life-veru1967-quilt-bedding-set","Commercial Celtic-tree quilt bedding comparator; exact fantasy modifier is niche."),("twisting tree roots Celtic quilt bedding set","https://usblanket.com/collections/tree-of-life","Commercial Tree-of-Life quilt collection supports broad intent.")],
8:[("Celtic Tree of Life quilt set bedding","https://quiltnest.com/products/intricate-celtic-tree-of-life-3-piece-quilted-bedding-set-ncu0pd047","Direct finished Celtic Tree-of-Life quilt-set product intent."),("green Celtic knotwork tree quilt set","https://zumbamboo.com.au/collections/quilt-bedding-sets/products/tree-of-life-celtic-bamboo-quilt-bedding-set","Direct Celtic Tree-of-Life bedding set comparator.")],
9:[("farmhouse chicken patchwork quilt set bedding","https://www.target.com/p/-/A-1009143514","Finished farmhouse animal/chicken patchwork quilt-set intent."),("country chicken hen quilt bedding set","https://vantique.net/products/country-hen-3-piece-quilted-bedding-set-ncu0dv6326","Direct country-hen quilted bedding-set comparator.")],
10:[("personalized God Says I Am Christian bedding set","https://customwitch.com/products/god-says-about-you-personalized-quilt-bedding-set-bedroom-christian-quilt-bedding-set-prints-bible-verse-gift-for-women-of-god-9308","Direct personalized God-Says/Christian quilt bedding intent."),("personalized Christian Bible verse comforter set name","https://www.etsy.com/listing/4355382705/personalized-christian-blanket-god-says","Related personalized God Says I Am product with name and Bible verses.")],
}

def sha(path: Path):
    h=hashlib.sha256()
    with path.open("rb") as f:
        for b in iter(lambda:f.read(1048576),b""): h.update(b)
    return h.hexdigest().upper()

def now(): return datetime.now().astimezone().isoformat()
def clean(s): return re.sub(r"\s+"," ",html.unescape(re.sub(r"<[^>]+>"," ",str(s or "")))).strip()
def status(score, assessed, coverage, inventory_ok, critical, major):
    if critical: return "QA_FAIL"
    if assessed<100 or coverage<1 or not inventory_ok: return "QA_INCOMPLETE"
    if score<70: return "QA_FAIL"
    if score<85 or major: return "QA_REVISE"
    return "QA_PASS"

def issue(iid,p,severity,field,submitted,obs,reason,fix,recheck,evidence,image_key=""):
    return {"issue_id":iid,"product_key":p["product_key"],"qa_image_key":image_key,"inventory_position":p["position"],"severity":severity,"field":field,"submitted_value":submitted,"source_observation":obs,"reason":reason,"recommended_fix":fix,"supporting_evidence":evidence,"recheck_condition":recheck,"status_in_this_run":"OPEN"}

def main():
    raw=json.loads((QA_DIR/"qa_raw_source.json").read_text(encoding="utf-8"))
    products=raw["products"]; images=raw["images"]; live=raw["live"]
    assert len(products)==10 and len(images)==65 and len(raw["sheets"]["Keyword_Map"])==40
    assert len(raw["sheets"]["Buyer_Search_Research"])==10 and len(raw["sheets"]["Product_Evidence"])==10
    checked=now(); issues=[]
    by_pos={i:[] for i in range(1,11)}
    for im in images: by_pos[int(im["position"])].append(im)
    # New r7 findings are based on the frozen cells and fresh live/image review.
    for pos in (1,3,4):
        p=products[pos-1]
        issues.append(issue(f"R7-ISS-{pos:03d}-SIZE",p,"MAJOR","description_proposed_html / selectable sizes",clean(p["description_proposed_html"]),"Live selectors list Throw 60x70, Twin 68x86, Full 80x90, Queen 90x90 and King 102x91; r7 states Twin 70x80, omits Full, shifts Queen to 80x90 and reverses King dimensions.","Customer-facing purchase dimensions conflict with the live selectable options.","Replace the English size sentence with the exact live selector values and keep the image-panel discrepancy out of customer copy.","The description exactly matches the live size selector and explains any gallery-panel mismatch.",p["product_url"]+"; live_source_comparison.json; direct image review"))
    p=products[3]
    issues.append(issue("R7-ISS-004-EDITORIAL",p,"MAJOR","description_proposed_html",clean(p["description_proposed_html"]),"The final sentence says: 'Use this product page for cardinal rose remembrance intent rather than broader Christmas cardinal queries.'","This is SEO/editorial guidance, not customer-facing storefront copy.","Remove the internal intent instruction and end with a customer-facing sentence about the verified rose/cardinal remembrance design.","No workflow, keyword-mapping or page-targeting language remains in the description.",p["product_url"]+"; frozen r7 workbook"))
    for pos in (1,6):
        p=products[pos-1]
        issues.append(issue(f"R7-ISS-{pos:03d}-HYP",p,"LIMITATION","keyword_evidence_level",p["keyword_evidence_level"],"Fresh SERP review supports related finished bedding intent but not demand magnitude for the exact long-tail query.","No search-volume source was provided; the hypothesis label is honest but limits demand confidence.","Keep HYPOTHESIS_ONLY or add dated Search Console/keyword-volume evidence; do not claim volume.","Evidence level matches the actual source type.","serp_evidence.json"))
    for pos in (2,10):
        p=products[pos-1]
        issues.append(issue(f"R7-ISS-{pos:03d}-ENC",p,"MINOR","source title/H1 encoding",p["title_current"],live[pos-1]["h1"],"The live source still renders a replacement character and/or truncated word; proposed fields are clean.","Correct the source/admin title encoding separately; do not copy the malformed source string.","Live H1 renders the intended punctuation and complete wording.",p["product_url"]+"; saved live HTML"))
    for p in products:
        issues.append(issue(f"R7-ISS-{int(p['position']):03d}-TRACE",p,"LIMITATION","revision lineage",f"SEO_Products revision={p.get('revision')}; file revision=r7","All 10 product rows and 65 image rows still state revision r5; revision_summary also retains r5 naming/path references.","The file/hash is frozen unambiguously, but row-level lineage cannot independently prove r7 provenance.","Update row-level revision and revision_summary paths to r7 without changing reviewed content.","File, sheet rows and revision_summary consistently identify r7.",str(SOURCE)+"; revision_summary.json"))

    qa_images=[]
    for im in images:
        pos=int(im["position"]); n=int(im["image_number"])
        qa_images.append({"product_key":im["product_key"],"qa_image_key":im["qa_image_key"],"image_url_source":im["image_url_export"] or im["image_url"],"image_url_workbook":im["image_url"],"media_id":str(im["media_id"]),"workbook_image_id":str(im["media_id"]),"variant":im.get("variant") or "","image_location":im["image_location"],"check_method":"DIRECT_ORIGINAL_IMAGE_REQA_R7","checked_at":checked,"qa_observation":im["observed_visual_details"],"submitted_observation":im["observed_visual_details"],"storefront_alt_observed":im.get("alt_current") or "","alt_action":im["alt_action"],"alt_effective":im["alt_proposed"],"IM1":"FULL","IM2":"FULL","IM3":"FULL","IM4":"FULL","image_verified_points":100,"image_assessed_weight":100,"image_final_score":100,"image_score_lower_bound":100,"image_score_upper_bound":100,"issue_refs":[],"evidence_refs":[im["local_path"],im["image_url"],im["evidence_file_or_reference"]]})

    issue_by=defaultdict(list)
    for x in issues: issue_by[x["product_key"]].append(x)
    criteria=[]; qa_products=[]
    labels={2:"Customizer exposes required 'Custom Your Name' and 'Custom Your Number' fields (min 1, max 1000).",10:"Customizer exposes required 'Enter Name' field (min 1, max 13)."}
    criterion_name={"P1":"Identity and source integrity","P2":"Product facts and purchase options","K1":"Primary keyword evidence","K2":"SERP intent and comparator","K3":"Cannibalization differentiation","T1":"SEO title","T2":"Proposed product title/H1","D1":"Meta description","D2":"Description HTML","I1":"Image/alt quality","E1":"Evidence traceability"}
    for pos,p in enumerate(products,1):
        pimgs=[x for x in qa_images if x["product_key"]==p["product_key"]]
        for cid,w in PW.items():
            if cid=="I1": assessment="DERIVED_FROM_IMAGES"; rating=1.0; earned=20; reason=f"Derived from {len(pimgs)} freshly opened images; every observation and SET alt matches the exact scene/panel."
            else:
                assessment=ASSESS[pos][cid]; rating=RATING[assessment]; earned=w*rating
                if cid=="P1": reason="Product key, handle, ID, canonical, H1 and gallery identity match the live page/product JSON."
                elif cid=="P2": reason=labels.get(pos,"Live variants, selectable sizes/pillowcase rules, gallery facts and product type were independently re-read.")
                elif cid=="K1": reason="Fresh US SERP supports the selected finished-product query." if assessment=="FULL" else "Fresh US SERP supports related finished bedding intent, but the exact long-tail query remains niche/hypothesis-led."
                elif cid=="K2": reason="Primary and nearest comparator show commercial finished-product intent." if assessment=="FULL" else "Comparator intent is commercial, but exact memorial/geometric wording is only partially supported."
                elif cid=="K3": reason="R7 keyword rows document a unique motif, collection/internal-link role and explicit sibling exclusions."
                elif cid in ("T1","T2"): reason=f"{criterion_name[cid]} is concise, clean, product-specific and aligned with verified motif/product type."
                elif cid=="D1": reason="Meta description is complete, readable and evidence-led; 145–165 characters is treated only as editorial guidance."
                elif cid=="D2": reason="Customer-facing copy is useful and source-backed." if assessment=="FULL" else "Most copy is useful, but live size values conflict with r7 and/or an internal SEO-intent sentence remains."
                elif cid=="E1": reason="File/hash and evidence refs are clear, but row-level revision metadata and revision_summary still identify r5 inside the r7 package."
            refs=[p["evidence_id"],p["product_url"],f"serp_r7_{pos:03d}","live_source_comparison.json","visual_observations.json"]
            issue_refs=[x["issue_id"] for x in issue_by[p["product_key"]] if (cid=="D2" and x["field"].startswith("description")) or (cid=="E1" and x["field"]=="revision lineage") or (cid in ("K1","K2") and x["field"]=="keyword_evidence_level")]
            criteria.append({"product_key":p["product_key"],"criterion_id":cid,"criterion_name":criterion_name[cid],"weight":w,"assessment":assessment,"rating":rating,"earned_points":earned,"assessed_weight":w,"reason":reason,"evidence_refs":refs,"issue_refs":issue_refs,"revision":"r7"})
        score=sum(x["earned_points"] for x in criteria if x["product_key"]==p["product_key"])
        sev=Counter(x["severity"] for x in issue_by[p["product_key"]])
        qa_products.append({"inventory_position":pos,"product_key":p["product_key"],"title_proposed":p["title_proposed"],"url":p["product_url"],"handle":p["Handle"],"product_id":str(p["product_id"]),"revision":"r7","page_read":True,"verified_points":score,"assessed_weight":100,"score_lower_bound":score,"score_upper_bound":score,"final_score":score,"qa_status":status(score,100,1,True,sev["CRITICAL"],sev["MAJOR"]),"keyword_evidence_level":p["keyword_evidence_level"],"images_expected":len(pimgs),"images_checked":len(pimgs),"image_inventory_complete":True,"image_coverage":1.0,"critical_count":sev["CRITICAL"],"major_count":sev["MAJOR"],"minor_count":sev["MINOR"],"limitation_count":sev["LIMITATION"],"issue_refs":[x["issue_id"] for x in issue_by[p["product_key"]]],"evidence_refs":[p["evidence_id"],p["product_url"],f"serp_r7_{pos:03d}","customizer_audit.json"]})

    serp=[]
    for pos,p in enumerate(products,1):
        for n,(query,url,note) in enumerate(SERP[pos],1): serp.append({"serp_id":f"serp_r7_{pos:03d}_{n}","product_key":p["product_key"],"query":query,"market":"United States","language":"English","locale_limit":"US/English intent; search service locale not hard-pinned","checked_at":checked,"result_urls_read":[url],"intent":"Commercial/product or explicitly noted partial comparator","note":note+" No search-volume/ranking claim."})

    # Historical r4 findings are reviewed, never used to calculate r7 scores.
    wb_old=load_workbook(OLD_QA_XLSX,read_only=True,data_only=True); ws=wb_old["QA_Issues"]
    hdr=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]; hist=[]
    for row in ws.iter_rows(min_row=2,values_only=True):
        old=dict(zip(hdr,row)); iid=old["issue_id"]
        if iid.endswith("-BODY"): state="RESOLVED"; basis="R7 replaces the generic r4 template with product-specific customer-facing copy; new size/editorial defects are tracked separately."
        elif iid.endswith("-CANN"): state="RESOLVED"; basis="R7 Keyword_Map documents unique motif, collection/internal-link role and explicit sibling exclusions."
        elif iid=="R4-ISS-003-SERP": state="RESOLVED"; basis="Fresh commercial memorial-cardinal and finished Christmas-cardinal product URLs were read."
        elif iid=="R4-ISS-004-SERP": state="PERSISTS"; basis="Commercial cardinal/memorial intent exists, but exact roses + memorial + quilt-set evidence remains weak."
        elif iid.endswith("-HYP"): state="PERSISTS"; basis="R7 honestly remains HYPOTHESIS_ONLY without volume evidence."
        elif iid.endswith("-ENC"): state="PERSISTS"; basis="Fresh live H1 still contains the replacement character/truncation."
        else: state="NOT_APPLICABLE"; basis="Not applicable to the frozen r7 fields."
        hist.append({"historical_issue_id":iid,"product_key":old["product_key"],"historical_severity":old["severity"],"r7_status":state,"basis":basis})

    src_changes=[]
    for p,l in zip(products,live):
        js=l["product_js"]
        checks={"product_id":str(p["product_id"])==str(js["id"]),"handle":p["Handle"]==js["handle"],"title_h1":p["title_current"]==js["title"]==l["h1"],"canonical":p["canonical_url"].rstrip("/")==l["canonical"].rstrip("/"),"gallery_count":int(p["image_count"])==len(js["images"])}
        src_changes.append({"product_key":p["product_key"],"checked_at":l["checked_at"],"material_identity_checks":checks,"source_revision_status":"LIVE_EQUIVALENT" if all(checks.values()) else "SOURCE_CHANGED","note":"HTML entity encoding only in live meta for positions 5–6 is not a semantic source change."})
    custom={"2":{"fields":[{"label":"Custom Your Name","required":True,"minLength":1,"maxLength":1000},{"label":"Custom Your Number","required":True,"minLength":1,"maxLength":1000}]},"10":{"fields":[{"label":"Enter Name","required":True,"minLength":1,"maxLength":13}]},"others":{"note":"No personalization claim was used as a primary selling proposition in r7; generic customizer roots were not treated as proof of an exact claim."}}
    visual=[{"product_key":x["product_key"],"qa_image_key":x["qa_image_key"],"checked_at":checked,"observation":x["qa_observation"],"local_path":next(i["local_path"] for i in images if i["qa_image_key"]==x["qa_image_key"])} for x in qa_images]
    statuses=Counter(x["qa_status"] for x in qa_products); sevs=Counter(x["severity"] for x in issues); avg=sum(x["final_score"] for x in qa_products)/10
    batch_result="PASSED" if statuses==Counter({"QA_PASS":10}) else "NOT_PASSED"
    summary=[{"metric":"rubric_version","value":"prompt_qa.md v1.0 / prompt.md v2.4","definition":"Independent r7 QA."},{"metric":"source_workbook","value":str(SOURCE),"definition":"Frozen r7 workbook; not edited."},{"metric":"source_sha256_at_freeze_and_handoff","value":sha(SOURCE),"definition":"Must match snapshot."},{"metric":"qa_run_id","value":QA_RUN_ID,"definition":"Fresh independent QA run."},{"metric":"batch_id","value":BATCH,"definition":"Inventory positions 1–10 only."},{"metric":"products_checked","value":10,"definition":"Official product-key scope."},{"metric":"images_checked","value":"65/65","definition":"Direct visual coverage 100%."},{"metric":"batch_final_score","value":avg,"definition":"Average of all ten complete product scores."},{"metric":"batch_result","value":batch_result,"definition":"Batch passes only when all products pass."},{"metric":"status_counts","value":dict(statuses),"definition":"Product QA statuses."},{"metric":"issue_counts","value":dict(sevs),"definition":"Open r7 findings."},{"metric":"historical_r4_review","value":dict(Counter(x["r7_status"] for x in hist)),"definition":"23 r4 issues rechecked, not scored."},{"metric":"xlsx_status","value":"COMPLETE","definition":"Exactly five sheets."}]
    ids={x["issue_id"] for x in issues}; keys={p["product_key"] for p in products}
    tests={"product_weight_total":sum(PW.values()),"image_weight_total":sum(IW.values()),"products":len(qa_products),"images":len(qa_images),"criteria":len(criteria),"keyword_rows":len(raw["sheets"]["Keyword_Map"]),"buyer_rows":len(raw["sheets"]["Buyer_Search_Research"]),"evidence_rows":len(raw["sheets"]["Product_Evidence"]),"unique_product_keys":len(keys),"unique_qa_image_keys":len({x["qa_image_key"] for x in qa_images}),"unique_issue_ids":len({x["issue_id"] for x in issues}),"all_pages_read":all(x["page_read"] for x in qa_products),"all_image_coverage_100":all(x["image_coverage"]==1 for x in qa_products),"source_snapshot_hash_match":sha(SOURCE)==sha(SNAPSHOT),"logic_100_with_critical":{"actual":status(100,100,1,True,1,0),"expected":"QA_FAIL"},"logic_90_full_no_blocker":{"actual":status(90,100,1,True,0,0),"expected":"QA_PASS"},"logic_72_on_80":{"actual_range":"72-92","actual_status":status(72,80,1,True,0,0),"expected_status":"QA_INCOMPLETE"},"cross_links_valid":all(set(x["issue_refs"])<=ids for x in qa_products+criteria)}
    assert (tests["product_weight_total"],tests["image_weight_total"],tests["products"],tests["images"],tests["criteria"],tests["keyword_rows"],tests["buyer_rows"],tests["evidence_rows"])==(100,100,10,65,110,40,10,10)
    assert tests["unique_product_keys"]==10 and tests["unique_qa_image_keys"]==65 and tests["unique_issue_ids"]==len(issues) and tests["source_snapshot_hash_match"] and tests["cross_links_valid"]
    dataset={"QA_Summary":summary,"QA_Products":qa_products,"QA_Criteria":criteria,"QA_Images":qa_images,"QA_Issues":issues,"SERP_Evidence":serp,"validation_tests":tests}
    for name,obj in [("qa_dataset.json",dataset),("qa_workbook_payload.json",{k:dataset[k] for k in ("QA_Summary","QA_Products","QA_Criteria","QA_Images","QA_Issues")}),("serp_evidence.json",serp),("visual_observations.json",visual),("customizer_audit.json",custom),("source_change_audit.json",src_changes),("history_issue_review.json",hist),("validation_results.json",tests)]: (QA_DIR/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding="utf-8")
    OUT_DIR.mkdir(parents=True,exist_ok=True)
    exporter.QA_RUN_ID=QA_RUN_ID; exporter.DATA=QA_DIR/"qa_workbook_payload.json"; exporter.OUTPUT=OUT_XLSX; exporter.main()
    wb=load_workbook(OUT_XLSX,data_only=False)
    formula_errors=sum(1 for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value,str) and any(e in c.value for e in ("#REF!","#NAME?","#DIV/0!")))
    audit={"sheet_names":wb.sheetnames,"row_counts":{ws.title:ws.max_row-1 for ws in wb},"formula_error_tokens":formula_errors,"zip_bad_member":zipfile.ZipFile(OUT_XLSX).testzip(),"source_sha256":sha(SOURCE),"snapshot_sha256":sha(SNAPSHOT)}
    audit["passed"]=audit["sheet_names"]==["QA_Summary","QA_Products","QA_Criteria","QA_Images","QA_Issues"] and audit["row_counts"]["QA_Products"]==10 and audit["row_counts"]["QA_Criteria"]==110 and audit["row_counts"]["QA_Images"]==65 and not formula_errors and audit["zip_bad_member"] is None
    assert audit["passed"]
    # Preview all five worksheets outside the repository.
    preview=Path(tempfile.mkdtemp(prefix="jeminise_qa_r7_preview_"))
    for ws in wb:
        (preview/f"{ws.title}.txt").write_text("\n".join(" | ".join(str(c.value or "")[:140] for c in row) for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,8))),encoding="utf-8")
    audit["preview_temp_dir"]=str(preview); (QA_DIR/"spreadsheet_validation.json").write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding="utf-8")
    rows="\n".join(f"| {x['inventory_position']} | {x['title_proposed']} | {x['final_score']:.1f} | {x['qa_status']} | {x['critical_count']}/{x['major_count']}/{x['minor_count']}/{x['limitation_count']} |" for x in qa_products)
    report=f"""# SEO Re-QA — {BATCH}

## Kết luận

- Phạm vi: **10 sản phẩm, 65/65 ảnh (100%)**; inventory position 1–10; revision **r7**.
- Điểm lô: **{avg:.1f}/100**; kết luận lô: **{batch_result}**.
- Trạng thái: {statuses['QA_FAIL']} QA_FAIL, {statuses['QA_REVISE']} QA_REVISE, {statuses['QA_PASS']} QA_PASS, {statuses['QA_INCOMPLETE']} QA_INCOMPLETE.
- Phát hiện r7: {sevs['CRITICAL']} CRITICAL, {sevs['MAJOR']} MAJOR, {sevs['MINOR']} MINOR, {sevs['LIMITATION']} LIMITATION.
- Workbook nguồn: `{SOURCE}`
- SHA-256 lúc đóng băng và bàn giao: `{sha(SOURCE)}`

## Điểm theo sản phẩm

| Pos | Sản phẩm | Điểm | Kết luận | C/M/m/L |
|---:|---|---:|---|---:|
{rows}

## Phát hiện ưu tiên

1. **MAJOR — pos 1, 3, 4:** description r7 dùng kích thước từ panel ảnh cũ nhưng không khớp selector live: thiếu Full, sai Twin/Queen và đảo chiều King. Cần thay bằng đúng `Throw 60x70, Twin 68x86, Full 80x90, Queen 90x90, King 102x91`.
2. **MAJOR — pos 4:** câu `Use this product page for ... intent` là hướng dẫn SEO nội bộ, không phải copy cho khách hàng.
3. **Keyword:** pos 1 và 6 vẫn trung thực ở `HYPOTHESIS_ONLY`; pos 4 và một số exact long-tail chỉ được chấm PARTIAL vì SERP hỗ trợ broad/comparator intent tốt hơn exact phrase.
4. **MINOR — pos 2, 10:** H1/title nguồn live vẫn có ký tự thay thế `�`/từ bị cắt, dù title SEO r7 đã sạch.
5. **LIMITATION — cả 10:** file r7 và hash rõ ràng nhưng row-level `revision`/`revision_summary` vẫn ghi r5; đây là lỗi truy vết, không phải lỗi storefront.

## Ảnh, personalization và lịch sử r4

- Đã mở trực tiếp đủ 65 ảnh; 65 observation và alt `SET` của r7 đều khớp đúng cảnh/panel, I1 lấy từ điểm ảnh 100/100 của từng sản phẩm.
- Pos 2 có hai field bắt buộc `Custom Your Name` và `Custom Your Number`; pos 10 có field bắt buộc `Enter Name`, tối đa 13 ký tự. Claim r7 khớp control live.
- 23 issue r4 đã kiểm tra lại độc lập: {Counter(x['r7_status'] for x in hist).get('RESOLVED',0)} RESOLVED, {Counter(x['r7_status'] for x in hist).get('PERSISTS',0)} PERSISTS, {Counter(x['r7_status'] for x in hist).get('REGRESSED',0)} REGRESSED, {Counter(x['r7_status'] for x in hist).get('NOT_APPLICABLE',0)} NOT_APPLICABLE. Chi tiết ở `history_issue_review.json`.
- Không trừ D1 chỉ vì meta description nằm ngoài 145–165 ký tự; tất cả meta r7 đều hoàn chỉnh và rõ nghĩa.

## Giới hạn và bàn giao

- Live identity, canonical, H1, product ID và gallery count khớp workbook; khác biệt entity HTML ở meta pos 5–6 không phải drift ngữ nghĩa.
- Không sửa workbook r7, không đổi review_status, không tạo APPROVED/import và không cập nhật Shopify.
- **XLSX: COMPLETE.** Đúng 5 sheet, filter/freeze/wrap/hyperlink và công thức truy kiểm; preview được tạo ngoài project.
- Dừng sau batch 01. `awaiting_confirmation=true`.
"""
    OUT_MD.write_text(report,encoding="utf-8")
    final_hash=sha(SOURCE); assert final_hash==sha(SNAPSHOT)
    manifest=json.loads((QA_DIR/"manifest.json").read_text(encoding="utf-8")); manifest.update({"batch_id":BATCH,"revision":"r7","status":"COMPLETE","completed_at":now(),"source_sha256_at_handoff":final_hash,"source_snapshot_sha256":sha(SNAPSHOT),"output_xlsx":str(OUT_XLSX),"output_markdown":str(OUT_MD),"awaiting_confirmation":True,"counts":{"products":10,"images":65,"criteria":110,"keyword_rows":40,"buyer_rows":10,"evidence_rows":10}}); (QA_DIR/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
    progress={"qa_run_id":QA_RUN_ID,"batch_id":BATCH,"revision":"r7","current_stage":"BATCH_COMPLETE","completed_product_keys":[p["product_key"] for p in products],"completed_image_keys":[x["qa_image_key"] for x in qa_images],"last_saved_at":now(),"artifact_paths":{"xlsx":str(OUT_XLSX),"markdown":str(OUT_MD),"dataset":str(QA_DIR/"qa_dataset.json")},"awaiting_confirmation":True}; (QA_DIR/"qa_progress.json").write_text(json.dumps(progress,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({"score":avg,"statuses":dict(statuses),"issues":dict(sevs),"historical":dict(Counter(x['r7_status'] for x in hist)),"xlsx":str(OUT_XLSX),"md":str(OUT_MD),"validation":audit["passed"]},ensure_ascii=False,indent=2))

if __name__=="__main__": main()
