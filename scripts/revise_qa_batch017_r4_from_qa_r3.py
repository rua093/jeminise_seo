import csv
import hashlib
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_017_r4"
QA_RUN_ID = "20260908_103500"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_017_r3" / "SEO_Product_Optimization_qa_batch_017_r3.xlsx"
QA_REPORT = ROOT / "resutls" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "SEO_QA_qa_batch_017_r3.md"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_017_r4.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_017_r4.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    161: {
        "title": "Blue Raven Celtic Knot Quilt",
        "meta_title": "Blue Raven Celtic Knot Quilt",
        "meta_description": "Shop a blue Norse raven quilt with Celtic knot frame, moon patchwork panels, matching shams and size chart.",
        "primary": "blue raven Celtic knot quilt",
        "secondary": "Norse raven Celtic quilt, Viking raven quilt set, Celtic knot raven bedding",
        "cluster": "blue Norse raven Celtic knot quilt",
        "detail": "blue and black raven artwork inside a Celtic knot frame with moon and patchwork panels",
        "intent_role": "Product page for blue raven Celtic knot quilt; link from Norse raven collection and avoid Valhalla shield, owl and Halloween comforter targets.",
        "customizer": "The visible field is a general Customize Your Item text box; no name, number or upload field is claimed.",
    },
    162: {
        "title": "Valhalla Viking Shield Quilt",
        "meta_title": "Valhalla Viking Shield Quilt",
        "meta_description": "Shop a black Valhalla Viking shield quilt with crossed axes, runic artwork, matching shams and size chart.",
        "primary": "Valhalla Viking shield quilt",
        "secondary": "Viking shield quilt, crossed axes bedding, Norse Valhalla quilt",
        "cluster": "black Valhalla Viking shield crossed axes quilt",
        "detail": "Viking shield and crossed axes artwork with Valhalla text on a black runic background",
        "intent_role": "Product page for Valhalla Viking shield quilt; link from Viking quilt collection and avoid raven, owl and Halloween comforter targets.",
        "customizer": "The visible field is a general Customize Your Item text box; no name, number or upload field is claimed.",
    },
    163: {
        "title": "Blue Gold Owl Night Quilt Set",
        "meta_title": "Blue Gold Owl Night Quilt Set",
        "meta_description": "Shop a blue gold owl night quilt set with moon frame artwork, starry patchwork border, shams and size chart.",
        "primary": "blue gold owl night quilt set",
        "secondary": "owl patchwork quilt, owl moon bedding, animal quilt set",
        "cluster": "blue gold owl moon patchwork quilt set",
        "detail": "owl perched on a branch inside a moon-style round frame with navy and gold patchwork details",
        "intent_role": "Product page for blue gold owl night quilt set; link from animal quilt collection and avoid raven, Viking shield and Halloween comforter targets.",
        "customizer": "The visible field is a general Customize Your Quilt text box; no name, number or upload field is claimed.",
    },
    164: {
        "title": "Witch Moon Halloween Comforter Set",
        "meta_title": "Witch Moon Halloween Comforter Set",
        "meta_description": "Shop a witch moon Halloween comforter set with black white artwork, bats, pumpkins, sheets and size chart.",
        "primary": "witch moon Halloween comforter set",
        "secondary": "black white Halloween bedding, Halloween comforter set with sheets, twin Halloween comforter set",
        "cluster": "black white witch moon Halloween comforter set",
        "detail": "black and white Halloween design with flying witch, moon, bats, pumpkins and HALLOWEEN text",
        "intent_role": "Product page for witch moon Halloween comforter set; link from Halloween comforter collection and avoid ghost, handprint and haunted house targets.",
        "customizer": "No verified personalization, name, number or upload field is claimed for this product.",
    },
    165: {
        "title": "Ghost Pumpkin Halloween Comforter Set",
        "meta_title": "Ghost Pumpkin Halloween Comforter Set",
        "meta_description": "Shop a ghost pumpkin Halloween comforter set with dark blue pattern, spiderwebs, bats, sheets and size chart.",
        "primary": "ghost pumpkin Halloween comforter set",
        "secondary": "dark blue Halloween bedding, pumpkin ghost comforter, twin Halloween comforter set",
        "cluster": "dark blue ghost pumpkin Halloween comforter set",
        "detail": "dark blue and gray Halloween pattern with white ghosts, orange pumpkins, spiderwebs, bats and stars",
        "intent_role": "Product page for ghost pumpkin Halloween comforter set; link from Halloween comforter collection and avoid witch moon, black ghost and haunted house targets.",
        "customizer": "No verified personalization, name, number or upload field is claimed for this product.",
    },
    166: {
        "title": "Black Ghost Halloween Comforter Set",
        "meta_title": "Black Ghost Halloween Comforter Set",
        "meta_description": "Shop a black ghost Halloween comforter set with white ghost pattern, matching pieces, fitted sheet panel and size chart.",
        "primary": "black ghost Halloween comforter set",
        "secondary": "white ghost bedding set, spooky ghost comforter, twin Halloween comforter set",
        "cluster": "black white ghost Halloween comforter set",
        "detail": "black Halloween comforter design covered with simple white ghost faces and matching pillow pattern",
        "intent_role": "Product page for black ghost Halloween comforter set; link from Halloween comforter collection and avoid pink ghost or ghost pumpkin targets.",
        "customizer": "No verified personalization, name, number or upload field is claimed for this product.",
    },
    167: {
        "title": "Haunted Pumpkin Ghost Comforter Set",
        "meta_title": "Haunted Pumpkin Ghost Comforter Set",
        "meta_description": "Shop a haunted pumpkin ghost comforter set with moon, graveyard silhouettes, sheets and Halloween size chart.",
        "primary": "haunted pumpkin ghost comforter set",
        "secondary": "blue Halloween bedding set, pumpkin ghost bedding, twin Halloween comforter set",
        "cluster": "blue haunted pumpkin ghost Halloween comforter set",
        "detail": "blue Halloween scene with ghosts, jack-o-lantern pumpkins, yellow moon, bare trees and graveyard silhouettes",
        "intent_role": "Product page for haunted pumpkin ghost comforter set; link from Halloween comforter collection and avoid simple ghost pumpkin or haunted house targets.",
        "customizer": "No verified personalization, name, number or upload field is claimed for this product.",
    },
    168: {
        "title": "Pink Cute Ghost Halloween Comforter Set",
        "meta_title": "Pink Cute Ghost Halloween Comforter Set",
        "meta_description": "Shop a pink cute ghost Halloween comforter set with witch hats, candy motifs, sheets and size chart.",
        "primary": "pink cute ghost Halloween comforter set",
        "secondary": "pink Halloween bedding, cute ghost bedding set, twin Halloween comforter set",
        "cluster": "pink cute ghost witch hat Halloween comforter set",
        "detail": "pink Halloween pattern with cute white ghosts, witch hats, candy, stars and small seasonal icons",
        "intent_role": "Product page for pink cute ghost Halloween comforter set; link from Halloween comforter collection and avoid black ghost or dark ghost pumpkin targets.",
        "customizer": "No verified personalization, name, number or upload field is claimed for this product.",
    },
    169: {
        "title": "Red Handprint Halloween Comforter Set",
        "meta_title": "Red Handprint Halloween Comforter Set",
        "meta_description": "Shop a red handprint Halloween comforter set with splatter-style artwork, sheets, fitted sheet panel and size chart.",
        "primary": "red handprint Halloween comforter set",
        "secondary": "red splatter bedding set, scary Halloween comforter, twin Halloween comforter set",
        "cluster": "red handprint splatter Halloween comforter set",
        "detail": "white comforter design with red splatter-style marks and red handprints plus matching red sheet pieces",
        "intent_role": "Product page for red handprint Halloween comforter set; link from Halloween comforter collection and avoid ghost, witch and pumpkin targets.",
        "customizer": "No verified personalization, name, number or upload field is claimed for this product.",
    },
    170: {
        "title": "Haunted House Pumpkin Comforter Set",
        "meta_title": "Haunted House Pumpkin Comforter Set",
        "meta_description": "Shop a haunted house pumpkin comforter set with black cat, spiderwebs, autumn leaves, sheets and size chart.",
        "primary": "haunted house pumpkin comforter set",
        "secondary": "white Halloween bedding set, pumpkin cat comforter, twin Halloween comforter set",
        "cluster": "white haunted house pumpkin cat Halloween comforter set",
        "detail": "white Halloween pattern with haunted houses, pumpkins, black cats, spiderwebs, bare trees and autumn leaves",
        "intent_role": "Product page for haunted house pumpkin comforter set; link from Halloween comforter collection and avoid ghost pumpkin, witch and handprint targets.",
        "customizer": "No verified personalization, name, number or upload field is claimed for this product.",
    },
}


COMMON_COMFORTER_MATTRESS = ("Fitted sheet panel showing standard and deep mattress fit up to 14 inches.", "Fitted sheet panel for 14 inch mattress depth")


IMAGE_DETAILS = {
    161: {
        1: ("Bedroom mockup of blue raven quilt with Celtic knot frame, moon panel and matching shams.", "Blue raven Celtic knot quilt on bed with shams"),
        2: ("Feature mockup of Norse raven quilt with cozy, breathable, lightweight and washable icons.", "Blue raven quilt feature panel"),
        3: ("All seasons panel showing blue raven Celtic knot quilt details and magnified texture circles.", "Norse raven Celtic knot detail panel"),
        4: ("Precision stitching panel showing raven, moon and Celtic knot close-ups.", "Blue raven quilt stitching panel"),
        5: ("Layer panel showing microfiber woven top, filling microfiber and back microfiber.", "Raven quilt microfiber layer panel"),
        6: ("Premium quilt set size chart for blue raven Celtic knot quilt with bed-size measurements.", "Size chart for blue raven Celtic knot quilt"),
        7: ("Fabric detail panel showing blue raven Celtic knot print and quilt surface close-ups.", "Blue raven Celtic quilt fabric panel"),
    },
    162: {
        1: ("Bedroom mockup of black Valhalla Viking shield quilt with crossed axes and matching shams.", "Valhalla Viking shield quilt on bed with shams"),
        2: ("Feature mockup of Viking shield quilt with cozy, breathable, lightweight and washable icons.", "Viking shield quilt feature panel"),
        3: ("All seasons panel showing Valhalla shield, crossed axes and runic background details.", "Valhalla shield quilt detail panel"),
        4: ("Precision stitching panel showing Viking shield, crossed axes and black fabric close-ups.", "Viking shield quilt stitching panel"),
        5: ("Layer panel showing microfiber woven top, filling microfiber and back microfiber.", "Viking shield quilt microfiber layer panel"),
        6: ("Premium quilt set size chart for Valhalla Viking shield quilt with bed-size measurements.", "Size chart for Valhalla Viking shield quilt"),
        7: ("Fabric detail panel showing black Valhalla shield quilt print and quilt surface close-ups.", "Valhalla Viking quilt fabric panel"),
    },
    163: {
        1: ("Bedroom mockup of blue gold owl night quilt set with matching shams.", "Blue gold owl night quilt set"),
        2: ("Feature mockup of owl moon quilt with cozy, breathable, lightweight and washable icons.", "Owl moon quilt feature panel"),
        3: ("All seasons panel showing owl, moon frame and starry blue patchwork details.", "Owl night quilt detail panel"),
        4: ("Precision stitching panel showing owl, branch and blue gold patchwork close-ups.", "Owl patchwork quilt stitching panel"),
        5: ("Layer panel showing microfiber woven top, filling microfiber and back microfiber.", "Owl quilt microfiber layer panel"),
        6: ("Premium quilt set size chart for blue gold owl night quilt with bed-size measurements.", "Size chart for blue gold owl night quilt"),
        7: ("Fabric detail panel showing owl night print and quilt surface close-ups.", "Blue gold owl quilt fabric panel"),
    },
    164: {
        1: ("Bedroom mockup of black white witch moon Halloween comforter set with bats and pumpkins.", "Witch moon Halloween comforter set"),
        2: ("Close-up of black comforter folded over white witch moon Halloween print.", "Witch Halloween comforter close-up"),
        3: ("Warm and cozy room mockup of witch moon Halloween bedding with pumpkins.", "Witch moon Halloween bedding room mockup"),
        4: ("Front bedroom mockup showing witch, moon, bats, pumpkins and HALLOWEEN text.", "Black white witch Halloween bedding"),
        5: COMMON_COMFORTER_MATTRESS,
        6: ("Package include panel with witch moon Halloween bedding and size table.", "Witch Halloween comforter package panel"),
    },
    165: {
        1: ("Bedroom mockup of dark blue ghost pumpkin Halloween comforter set with matching pillows.", "Ghost pumpkin Halloween comforter set"),
        2: ("Close-up of dark blue ghost pumpkin bedding pattern with spiderwebs and bats.", "Ghost pumpkin comforter close-up"),
        3: ("Warm and cozy room mockup of dark blue Halloween bedding with ghosts and pumpkins.", "Dark blue Halloween bedding room mockup"),
        4: COMMON_COMFORTER_MATTRESS,
        5: ("Package include panel with ghost pumpkin Halloween bedding and size table.", "Ghost pumpkin comforter package panel"),
    },
    166: {
        1: ("Bedroom mockup of black Halloween comforter set covered with white ghost faces.", "Black ghost Halloween comforter set"),
        2: ("Close-up of black comforter folded over repeated white ghost pattern.", "Black ghost comforter close-up"),
        3: ("Warm and cozy room mockup of black white ghost Halloween bedding.", "Black white ghost Halloween bedding"),
        4: ("Front bedroom mockup showing black ghost comforter and matching pillow pattern.", "White ghost bedding set front view"),
        5: COMMON_COMFORTER_MATTRESS,
        6: ("Package include panel with black ghost Halloween bedding and size table.", "Black ghost comforter package panel"),
    },
    167: {
        1: ("Bedroom mockup of blue haunted pumpkin ghost comforter set with moon and graveyard silhouettes.", "Haunted pumpkin ghost comforter set"),
        2: ("Close-up of blue Halloween bedding with ghost, pumpkin, moon and bare tree details.", "Haunted ghost comforter close-up"),
        3: ("Warm and cozy room mockup of blue Halloween bedding with pumpkins and ghosts.", "Blue Halloween bedding room mockup"),
        4: ("Front bedroom mockup showing moon, ghosts, jack-o-lanterns, trees and graveyard silhouettes.", "Pumpkin ghost Halloween bedding"),
        5: COMMON_COMFORTER_MATTRESS,
        6: ("Package include panel with haunted pumpkin ghost bedding and size table.", "Haunted pumpkin comforter package panel"),
    },
    168: {
        1: ("Bedroom mockup of pink cute ghost Halloween comforter set with matching pillows.", "Pink cute ghost Halloween comforter set"),
        2: ("Close-up of pink ghost bedding pattern with witch hats, candy and small stars.", "Pink ghost comforter close-up"),
        3: ("Warm and cozy room mockup of pink Halloween bedding with cute ghosts.", "Pink Halloween bedding room mockup"),
        4: ("Front bedroom mockup showing pink ghost pattern with witch hats and candy motifs.", "Cute ghost Halloween bedding"),
        5: COMMON_COMFORTER_MATTRESS,
        6: ("Package include panel with pink cute ghost Halloween bedding and size table.", "Pink ghost comforter package panel"),
    },
    169: {
        1: ("Bedroom mockup of white red handprint Halloween comforter set with red sheet pieces.", "Red handprint Halloween comforter set"),
        2: ("Close-up of red handprint and splatter-style Halloween bedding pattern.", "Red handprint comforter close-up"),
        3: ("Warm and cozy room mockup of red handprint Halloween bedding.", "Red handprint Halloween bedding room mockup"),
        4: ("Front bedroom mockup showing red handprints and splatter-style marks on white bedding.", "Red splatter Halloween bedding"),
        5: COMMON_COMFORTER_MATTRESS,
        6: ("Package include panel with red handprint Halloween bedding and size table.", "Red handprint comforter package panel"),
    },
    170: {
        1: ("Bedroom mockup of white haunted house pumpkin comforter set with black cats and rust sheets.", "Haunted house pumpkin comforter set"),
        2: ("Close-up of white Halloween bedding with haunted houses, pumpkins, cats and spiderwebs.", "Haunted house comforter close-up"),
        3: ("Warm and cozy room mockup of white Halloween bedding with haunted houses and pumpkins.", "White Halloween bedding room mockup"),
        4: ("Front bedroom mockup showing haunted houses, pumpkins, black cats, spiderwebs and leaves.", "Haunted house pumpkin bedding"),
        5: COMMON_COMFORTER_MATTRESS,
        6: ("Package include panel with haunted house pumpkin bedding and size table.", "Haunted house comforter package panel"),
    },
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def load_qa():
    wb = load_workbook(SOURCE, data_only=False)
    ws = wb["SEO_Products"]
    idx = headers(ws)
    products = []
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos in PRODUCT_UPDATES:
            products.append({
                "product_key": text(ws.cell(row_num, idx["product_key"]).value),
                "handle": text(ws.cell(row_num, idx["Handle"]).value),
                "inventory_position": pos,
            })
    products.sort(key=lambda item: item["inventory_position"])
    return {
        "product_keys": [item["product_key"] for item in products],
        "handle_by_key": {item["product_key"]: item["handle"] for item in products},
        "pos_by_key": {item["product_key"]: int(item["inventory_position"]) for item in products},
        "customizer_by_pos": {},
    }


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    missing = sorted(set(handles) - set(admin))
    if missing:
        raise RuntimeError(f"Missing handles in admin export: {missing}")
    return admin


def option_text(admin_row):
    values = [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]]
    values = [v for v in values if v]
    return ", ".join(values) if values else "available size options"


def customizer_sentence(pos):
    if pos in {161, 162, 158, 159, 160}:
        return "A general optional Customize Your Item text box supports custom text from 1 to 1000 characters."
    if pos == 163:
        return "A general optional Customize Your Quilt text box supports custom text from 1 to 1000 characters."
    return "No shopper text-entry field is claimed for this comforter set."


def description(pos, admin_row, ctext):
    item = PRODUCT_UPDATES[pos]
    if pos <= 163:
        return (
            f"<p>The {item['title']} features {item['detail']} for a distinct quilt design.</p>"
            "<h3>Product Features</h3>"
            "<ul><li>Soft microfiber fabric details are shown in the product feature panels.</li>"
            "<li>Lightweight all-season construction is shown with breathable comfort and care icons.</li>"
            "<li>Gallery panels show the main bed mockup, detail views, fabric/layer panels and size chart.</li></ul>"
            "<h3>Sizing and Options</h3>"
            "<ul><li>Choose from the available quilt size options shown on the product page.</li>"
            "<li>Matching pillowcases are available separately where the selector is shown.</li>"
            f"<li>{ctext}</li></ul>"
        )
    return (
        f"<p>The {item['title']} features {item['detail']} with coordinated sheets for a complete Halloween bedding look.</p>"
        "<h3>Package Contents by Size</h3>"
        '<ul><li>Twin 5-piece set: 1 comforter 68" x 90", 1 pillow sham 20" x 30", 1 pillowcase 20" x 30", 1 fitted sheet 75" x 39" + 14" pocket and 1 flat sheet 92" x 66".</li>'
        '<li>Full 7-piece set: 1 comforter 80" x 90", 2 pillow shams 20" x 30", 2 pillowcases 20" x 30", 1 fitted sheet 75" x 54" + 14" pocket and 1 flat sheet 96" x 81".</li>'
        '<li>Queen 7-piece set: 1 comforter 90" x 90", 2 pillow shams 20" x 30", 2 pillowcases 20" x 30", 1 fitted sheet 80" x 60" + 14" pocket and 1 flat sheet 102" x 90".</li>'
        '<li>King 7-piece set: 1 comforter 104" x 90", 2 pillow shams 20" x 36", 2 pillowcases 20" x 36", 1 fitted sheet 80" x 78" + 14" pocket and 1 flat sheet 108" x 102".</li></ul>'
        "<h3>Material and Care</h3>"
        "<ul><li>The fitted sheet is shown with 360-degree elastic for mattresses up to 14 inches deep.</li>"
        "<li>Brushed microfiber fabric and polyester filling provide a soft, lightweight bedding feel.</li>"
        f"<li>{ctext}</li>"
        "<li>Care details include machine washing cold on gentle cycle and tumble drying low.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    qa = load_qa()
    admin = load_admin(set(qa["handle_by_key"].values()))
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ctext = customizer_sentence(pos)
        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or "Bedding"
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, ctext)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"US English buyer intent targets {item['cluster']}; no search-volume, trend or ranking claim is made."
        )
        ws.cell(row_num, idx["keyword_evidence_level"]).value = "SERP_ONLY"
        ws.cell(row_num, idx["revision"]).value = "r4"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R4 after qa_batch_017_r3 QA: used products_export_1.csv sha256:{admin_hash}; "
            "replaced process-style descriptions with product-ready package/specification copy, avoided unsupported name/photo/upload claims, "
            "and kept Norse, owl and Halloween intents separated by visible motif. "
            "Still NEEDS_REVIEW."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    seen_alts = set()
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        image_number = int(ws.cell(row_num, idx["image_number"]).value)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation, alt = IMAGE_DETAILS[pos][image_number]
        if alt in seen_alts:
            alt = f"{alt} for {PRODUCT_UPDATES[pos]['cluster']}"
        seen_alts.add(alt)
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = alt[:125]
        ws.cell(row_num, idx["alt_action"]).value = "SET"
        ws.cell(row_num, idx["viewed_status"]).value = "VIEWED"
        ws.cell(row_num, idx["revision"]).value = "r4"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R4: observation and alt carried forward from r3 checked image audit; admin image URL and current alt matched from Shopify CSV where available."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "PRODUCT"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = (
            "qa_batch_017_r3 QA report + products_export_1.csv admin baseline + r4 description/specification cleanup"
        )
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = (
            "R4 separates nearby pages by visible motif, wording, color palette, product type and verified option level."
        )
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r4"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as a specific product page."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm exact artwork, product type, size choices, "
            "and only options that are actually supported."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = (
            "Confirm product type, visible motif, size or fit panels, supported options and image accuracy before purchase."
        )
        ws.cell(row_num, idx["emotional_social_motivation"]).value = (
            "Choose a Norse, owl or Halloween bedding item without confusing it with another similar design."
        )
        ws.cell(row_num, idx["purchase_concerns"]).value = (
            "Avoid unsupported name, number, upload, cultural-origin or material claims; confirm exact size and product type."
        )
        ws.cell(row_num, idx["source_refs"]).value = (
            f"products_export_1.csv sha256:{admin_hash}; QA report {QA_RUN_ID}; checked images 161-170"
        )
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R4_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = (
            "No Search Console, paid keyword volume or internal site-search export supplied; proof remains SERP-only."
        )
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    pk_by_pos = {qa["pos_by_key"][pk]: pk for pk in qa["product_keys"]}
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = pk_by_pos[pos]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ctext = customizer_sentence(pos)
        ws.cell(row_num, idx["sources_accessed"]).value = (
            f"Original storefront/product evidence; QA report {QA_RUN_ID}; products_export_1.csv sha256:{admin_hash}; "
            "checked image audit"
        )
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={option_text(admin_row)}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={item['detail']}; "
            f"customizer={ctext}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "qa_batch_017_r3 QA description issues addressed in r4; independent QA recheck still required."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R4 uses checked image audit, QA report recommendations and Shopify admin CSV baseline; no approval or import file created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_017_r4_revision", "r4", "Separate 10-product revision after qa_batch_017_r3 QA; source r3 workbook was not modified."),
        ("qa_batch_017_r4_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_017_r4_scope", "inventory positions 161-170", "No products outside qa_batch_017 were revised."),
        ("qa_batch_017_r4_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_017_r4_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_017_r4_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa_report": str(QA_REPORT.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_017 only; inventory positions 161-170",
        "revision": "r4",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "r4_changes": [
            "Rewrote all 10 product descriptions to remove process/meta commentary flagged by qa_batch_017_r3 QA.",
            "Added verified package contents by size for Halloween comforter sets, including Twin 5-piece and Full/Queen/King 7-piece details.",
            "Added material, fitted sheet depth and care details where QA identified them as verified from images 5 and 6.",
            "Kept no unsupported name, number, upload or personalization placement claims.",
        ],
        "next_step": "Sang QA lại qa_batch_017_r4 before any approval or import.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_017_r4",
                "",
                "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r3: `{SOURCE.relative_to(ROOT)}`",
                f"- QA r3 report: `{QA_REPORT.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r4: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Sửa trọng yếu: thay body copy kiểu quy trình bằng mô tả publish-ready, bổ sung package contents theo size cho comforter, giữ đúng material/care/fitted-sheet facts và không claim name/photo/upload.",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo approval; không tạo file import Shopify.",
                "- Bước tiếp theo: Sang QA lại `qa_batch_017_r4`.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
