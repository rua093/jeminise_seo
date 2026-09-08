import csv
import hashlib
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_013_r3"
QA_RUN_ID = "20260907_182000"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_013_r2" / "SEO_Product_Optimization_qa_batch_013_r2.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
CUSTOMIZER_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_audit.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_013_r3.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_013_r3.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    121: {
        "title": "Vintage Blue Semi Truck Comforter",
        "meta_title": "Vintage Blue Semi Truck Comforter",
        "meta_description": "Shop a vintage blue semi truck comforter with stone wall artwork, YOUR NAME sample text, care panels and size guide.",
        "primary": "vintage blue semi truck comforter",
        "secondary": "custom trucker bedding, blue truck comforter with name, stone wall semi truck bedding",
        "cluster": "vintage blue semi truck stone wall comforter",
        "detail": "weathered blue semi truck artwork breaking through gray stone blocks with YOUR NAME sample text",
        "intent_role": "Truck page separated by vintage/weathered blue cab, stone wall blocks and confirmed Enter Name field.",
        "customizer": "A required Enter Name field supports up to 35 characters, with a required product-confirmation choice.",
    },
    122: {
        "title": "Blue Semi Truck Chevron Comforter",
        "meta_title": "Blue Semi Truck Chevron Comforter",
        "meta_description": "Shop a blue semi truck comforter with black chevron panels, flag side borders, YOUR NAME sample text and size guide.",
        "primary": "blue semi truck chevron comforter",
        "secondary": "chevron trucker bedding, custom blue semi truck comforter, truck comforter with name",
        "cluster": "blue semi truck chevron flag comforter",
        "detail": "blue semi truck centered on black and blue chevron panels with flag side borders and Your Name sample text",
        "intent_role": "Truck page separated by chevron center panel, flag borders and confirmed Enter Name field.",
        "customizer": "A required Enter Name field supports up to 35 characters, with a required product-confirmation choice.",
    },
    123: {
        "title": "Metallic Mesh Semi Truck Comforter",
        "meta_title": "Metallic Mesh Semi Truck Comforter",
        "meta_description": "Shop a black and silver semi truck comforter with metallic mesh background, map texture, YOUR NAME panel and size guide.",
        "primary": "metallic mesh semi truck comforter",
        "secondary": "silver trucker bedding, semi truck mesh comforter, custom truck comforter",
        "cluster": "black silver semi truck metallic mesh comforter",
        "detail": "black and silver semi truck artwork on metallic mesh with map-like texture, stars and vertical YOUR NAME panel",
        "intent_role": "Truck page separated by metallic mesh, grayscale map texture and confirmed Enter Name field.",
        "customizer": "A required Enter Name field supports up to 35 characters, with a required product-confirmation choice.",
    },
    124: {
        "title": "Textured Metallic Semi Truck Comforter",
        "meta_title": "Textured Metallic Semi Truck Comforter",
        "meta_description": "Shop a textured metallic semi truck comforter with line-art truck design, YOUR NAME sample text, care panels and size guide.",
        "primary": "textured metallic semi truck comforter",
        "secondary": "line art trucker bedding, metallic truck comforter, semi truck comforter with name",
        "cluster": "line-art semi truck textured metallic comforter",
        "detail": "white line-art semi truck on a dark textured metallic panel with YOUR NAME sample text",
        "intent_role": "Truck page separated by line-art truck style, dark textured panel and confirmed Enter Name field.",
        "customizer": "A required Enter Name field supports up to 35 characters, with a required product-confirmation choice.",
    },
    125: {
        "title": "Starry Night Semi Truck Comforter",
        "meta_title": "Starry Night Semi Truck Comforter",
        "meta_description": "Shop a starry night semi truck comforter with blue-purple truck art, YOUR NAME field, romantic text bands and size guide.",
        "primary": "starry night semi truck comforter",
        "secondary": "blue purple trucker bedding, semi truck heart comforter, custom truck comforter with name",
        "cluster": "starry night semi truck heart message comforter",
        "detail": "blue-purple semi truck under a starry sky with YOUR NAME sample text and My Heart Is Always With You wording",
        "intent_role": "Truck page separated by starry sky, heart-message wording and confirmed Enter Name field.",
        "customizer": "A required Enter Name field supports up to 35 characters, with a required product-confirmation choice.",
    },
    126: {
        "title": "Cardinal Christmas Quilt Set",
        "meta_title": "Cardinal Christmas Quilt Set",
        "meta_description": "Shop a Christmas cardinal quilt set with two red birds, holly branches, berry artwork, matching shams and size chart.",
        "primary": "cardinal Christmas quilt set",
        "secondary": "red cardinal holiday quilt, holly cardinal bedding, Christmas bird quilt set",
        "cluster": "Christmas cardinal holly quilt set",
        "detail": "two red cardinals on holly branches with berries, script-style background and dark red border",
        "intent_role": "Quilt page separated by Christmas cardinal and holly branch artwork.",
        "customizer": "The visible field is a general Customize Your Quilt text box; no name or number field is claimed.",
    },
    127: {
        "title": "Colorful Wolf Head Quilt Set",
        "meta_title": "Colorful Wolf Head Quilt Set",
        "meta_description": "Shop a colorful wolf head quilt set with close-up portrait artwork, yellow eyes, matching shams and bedding size panel.",
        "primary": "colorful wolf head quilt set",
        "secondary": "wolf portrait bedding, colorful wildlife quilt, wolf face quilt set",
        "cluster": "colorful close-up wolf portrait quilt set",
        "detail": "close-up wolf head portrait with yellow eyes and multicolor abstract fur artwork",
        "intent_role": "Wolf page separated by colorful close-up portrait artwork and optional Custom Your Name field.",
        "customizer": "A Custom Your Name field is present but not required; no number field is claimed.",
    },
    128: {
        "title": "Geometric Wolf Head Quilt Set",
        "meta_title": "Geometric Wolf Head Quilt Set",
        "meta_description": "Shop a geometric wolf head quilt set with teal, cream and orange artwork, close-up eye detail and matching shams.",
        "primary": "geometric wolf head quilt set",
        "secondary": "geometric wolf bedding, teal wolf quilt set, wolf face quilt",
        "cluster": "geometric teal cream wolf head quilt set",
        "detail": "geometric wolf face artwork in teal, cream and orange tones with close-up eye detail",
        "intent_role": "Wolf page separated by geometric face styling; no personalization claim is used because labels were not verified.",
        "customizer": "No clear customizer label was verified for this product, so name or number customization is not claimed.",
    },
    129: {
        "title": "Profile Wolf Feathers Quilt Set",
        "meta_title": "Profile Wolf Feathers Quilt Set",
        "meta_description": "Shop a profile wolf quilt set with feather accents, neutral tan artwork, matching shams and 3-piece bedding panel.",
        "primary": "profile wolf feathers quilt set",
        "secondary": "wolf feathers bedding, tan wolf quilt set, wildlife wolf quilt",
        "cluster": "profile wolf head feather accent quilt set",
        "detail": "side-profile wolf head artwork with feather accents on a tan and brown background",
        "intent_role": "Wolf page separated by side-profile pose, neutral palette and feather accents.",
        "customizer": "A Custom Your Name field is present but not required; no number field is claimed.",
    },
    130: {
        "title": "Wolf Dreamcatcher Quilt Set",
        "meta_title": "Wolf Dreamcatcher Quilt Set",
        "meta_description": "Shop a wolf dreamcatcher quilt set with centered wolf head, feather accents, earth-tone artwork and matching shams.",
        "primary": "wolf dreamcatcher quilt set",
        "secondary": "wolf dreamcatcher bedding, wolf feather quilt, earth tone wolf quilt set",
        "cluster": "centered wolf dreamcatcher feather quilt set",
        "detail": "centered wolf head inside a dreamcatcher-style ring with feather accents and earth-tone background",
        "intent_role": "Wolf page separated by centered dreamcatcher-style ring and feather accents.",
        "customizer": "A Custom Your Name field is present but not required; no number field is claimed.",
    },
}


COMMON_DUVET_PANEL = ("Comparison panel explaining duvet cover set and comforter set bedding options.", "Duvet cover and comforter set comparison panel")
COMMON_CARE = ("Stress-free easy-care panel listing wrinkle-free, stain-proof, anti-pilling and wash durability.", "Easy care panel for comforter bedding")
COMMON_SIZE = ("Size dimension chart showing twin, full, queen and king bedding measurements.", "Comforter set size dimension chart")
COMMON_QUILT_INCLUDED = ("Three-piece bedding set panel listing one premium quilt and two optional standard shams.", "Three piece quilt set information panel")


def truck_images(label, feature_label):
    return {
        1: (f"Bedroom mockup of {label}.", feature_label),
        2: (f"Second bedroom mockup of {label} with matching shams and folded top sheet.", f"{feature_label} bedding set"),
        3: COMMON_DUVET_PANEL,
        4: (f"Feature panel showing {feature_label.lower()}, microfiber label and 3D printed pattern callouts.", f"{feature_label} feature panel"),
        5: (f"Low-angle bedding mockup of {label} with soft, lightweight, durable and breathable icons.", f"{feature_label} with feature icons"),
        6: COMMON_CARE,
        7: COMMON_SIZE,
    }


IMAGE_DETAILS = {
    121: truck_images("weathered blue semi truck breaking through gray stone blocks with YOUR NAME text", "Vintage blue semi truck comforter"),
    122: truck_images("blue semi truck on black and blue chevron panels with flag borders and Your Name text", "Blue semi truck chevron comforter"),
    123: truck_images("black and silver semi truck on metallic mesh and map-texture background with vertical YOUR NAME panel", "Metallic mesh semi truck comforter"),
    124: truck_images("line-art semi truck on dark textured metallic panel with YOUR NAME text", "Textured metallic semi truck comforter"),
    125: truck_images("blue-purple semi truck under a starry sky with YOUR NAME and My Heart Is Always With You text", "Starry night semi truck comforter"),
    126: {
        1: ("Bedroom mockup of Christmas cardinal quilt with two red birds, holly branches, berries and matching shams.", "Cardinal Christmas quilt set"),
        2: ("Room mockup of cardinal quilt with premium printed craft callout and fabric inset.", "Cardinal holiday quilt room mockup"),
        3: ("Optional pillow shams panel showing red cardinal and holly branch artwork.", "Cardinal Christmas pillow shams"),
        4: ("High-quality fabric panel showing cardinal quilt detail and pillow mockup.", "Cardinal quilt fabric detail panel"),
        5: ("Premium quilt set size chart with cardinal quilt thumbnail and bed measurements.", "Cardinal quilt set size chart"),
        6: ("Bedspread features panel showing quilt layers and anti-wrinkle, anti-static and machine-washable icons.", "Cardinal quilt construction feature panel"),
        7: ("Bedroom mockup of cardinal Christmas quilt with dark red border and matching shams.", "Cardinal Christmas quilt bedroom mockup"),
    },
    127: {
        1: ("Bedroom mockup of colorful close-up wolf head quilt with yellow eyes and matching shams.", "Colorful wolf head quilt set"),
        2: ("Close-up view of colorful wolf face artwork with yellow eyes and abstract fur lines.", "Colorful wolf face quilt close-up"),
        3: ("Pillow sham mockup showing colorful wolf head portrait artwork.", "Colorful wolf head pillow sham"),
        4: ("Room mockup of colorful wolf portrait quilt set on a modern bed.", "Colorful wolf portrait bedding set"),
        5: COMMON_QUILT_INCLUDED,
    },
    128: {
        1: ("Bedroom mockup of geometric wolf head quilt in teal, cream and orange tones with matching shams.", "Geometric wolf head quilt set"),
        2: ("Close-up view of geometric wolf face artwork with brown eye and teal-orange shapes.", "Geometric wolf face quilt close-up"),
        3: ("Pillow sham mockup showing geometric wolf head artwork.", "Geometric wolf head pillow sham"),
        4: ("Room mockup of geometric wolf head quilt set on a modern bed.", "Geometric wolf bedding set"),
        5: COMMON_QUILT_INCLUDED,
    },
    129: {
        1: ("Bedroom mockup of side-profile wolf head quilt with feather accents and matching shams.", "Profile wolf feathers quilt set"),
        2: ("Close-up view of wolf fur, eye and feather accent artwork.", "Wolf feathers quilt close-up"),
        3: ("Pillow sham mockup showing side-profile wolf head with feather accents.", "Profile wolf pillow sham with feathers"),
        4: ("Room mockup of tan side-profile wolf quilt set on a modern bed.", "Profile wolf quilt bedding set"),
        5: COMMON_QUILT_INCLUDED,
    },
    130: {
        1: ("Bedroom mockup of centered wolf head quilt inside a dreamcatcher-style ring with feather accents.", "Wolf dreamcatcher quilt set"),
        2: ("Close-up view of wolf face, blue eyes and quilt stitching texture.", "Wolf dreamcatcher quilt close-up"),
        3: ("Pillow sham mockup showing centered wolf head and dreamcatcher-style ring.", "Wolf dreamcatcher pillow sham"),
        4: ("Room mockup of wolf dreamcatcher quilt set with matching shams.", "Wolf dreamcatcher bedding set"),
        5: COMMON_QUILT_INCLUDED,
    },
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def load_qa():
    data = json.loads(QA_DATASET.read_text(encoding="utf-8"))
    products = data["QA_Products"]
    customizer = json.loads(CUSTOMIZER_AUDIT.read_text(encoding="utf-8"))
    return {
        "product_keys": [item["product_key"] for item in products],
        "handle_by_key": {item["product_key"]: item["handle"] for item in products},
        "pos_by_key": {item["product_key"]: int(item["inventory_position"]) for item in products},
        "customizer_by_pos": {int(item["inventory_position"]): item for item in customizer},
    }


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Body (HTML)": first.get("Body (HTML)", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    return admin


def option_text(admin_row):
    values = [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]]
    values = [v for v in values if v]
    return ", ".join(values) if values else "available size options"


def customizer_sentence(audit):
    excerpt = audit.get("config_excerpt", "")
    text_matches = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"minLength":\s*(\d+),\s*"maxLength":\s*(\d+)',
        excerpt,
    )
    option_matches = re.findall(
        r'\{"id":\s*"[^"]+",\s*"label":\s*"([^"]+)",\s*"required":\s*(true|false),\s*"defaultOptionId"',
        excerpt,
    )
    parts = []
    for label, required, min_length, max_length in text_matches:
        requirement = "required" if required == "true" else "optional"
        parts.append(f"Personalization uses a {requirement} {label} field, {min_length}-{max_length} characters.")
    useful_options = [
        label for label, _ in option_matches
        if label and "confirmation" not in label.lower()
    ]
    if useful_options:
        parts.append(f"Additional visible customization option: {', '.join(useful_options)}.")
    if not text_matches:
        parts.append("No shopper text-entry field is described for this product.")
    labels = " ".join(label for label, *_ in text_matches).lower()
    if "number" in labels:
        parts.append("Names and numbers visible in mockups are sample artwork unless the matching input is entered.")
    elif "name" in labels:
        parts.append("Visible names in mockups are sample artwork unless the matching input is entered.")
    else:
        parts.append("Visible custom text in mockups is sample artwork unless the matching input is entered.")
    return " ".join(parts)


def description(pos, admin_row, ctext):
    item = PRODUCT_UPDATES[pos]
    return (
        f"<p>{item['title']} features {item['detail']}. "
        "This product page is written around the specific visible artwork rather than a broad template.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork: {item['detail']}.</li>"
        "<li>Gallery images show the main bed mockup plus detail, feature, care, size or bedding-type panels where present.</li></ul>"
        "<h3>Options and Fit</h3>"
        f"<ul><li>Available option groups: {option_text(admin_row)}.</li>"
        f"<li>{ctext}</li>"
        "<li>Choose the product type and size shown on the product page before checkout.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    qa = load_qa()
    admin = load_admin(set(qa["handle_by_key"].values()))
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    handles = set()
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        handles.add(handle)
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ctext = customizer_sentence(qa["customizer_by_pos"].get(pos, {}))
        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or "Bedding"
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, ctext)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"US English buyer intent targets {item['cluster']}; no volume or ranking claim is made."
        )
        ws.cell(row_num, idx["keyword_evidence_level"]).value = "SERP_ONLY"
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R3 deep recheck after Sang QA batch 013: used products_export_1.csv sha256:{admin_hash}; "
            "removed internal workflow copy, tied customization wording only to customizer evidence, "
            "split truck and wolf intents by visible motif, and rewrote image observations/alts from inspected contact sheets. "
            "Still NEEDS_REVIEW."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    seen_alts = set()
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        image_number = int(ws.cell(row_num, idx["image_number"]).value)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation, alt = IMAGE_DETAILS[pos][image_number]
        if alt in seen_alts:
            alt = f"{alt} for {PRODUCT_UPDATES[pos]['cluster']}"
        seen_alts.add(alt)
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = alt[:125]
        ws.cell(row_num, idx["alt_action"]).value = "SET"
        ws.cell(row_num, idx["viewed_status"]).value = "VIEWED"
        ws.cell(row_num, idx["revision"]).value = "r3"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R3: observation and alt rewritten from inspected contact sheet; admin image URL and current alt matched from Shopify CSV where available."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "PRODUCT"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = (
            "Sang QA batch 013 + products_export_1.csv admin baseline + customizer_audit.json + contact-sheet inspection"
        )
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = (
            "R3 separates nearby pages by visible motif, wording, color palette and proof level for customization."
        )
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r3"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as a specific product page."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm exact artwork, product type, size choices, "
            "and only customization options that are actually supported."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = (
            "Confirm product type, visible motif, size/care panels, supported custom fields and image accuracy before purchase."
        )
        ws.cell(row_num, idx["emotional_social_motivation"]).value = (
            "Choose a truck, cardinal or wolf bedding item without confusing it with another similar design in the batch."
        )
        ws.cell(row_num, idx["purchase_concerns"]).value = (
            "Avoid unsupported name, number, photo, upload, cultural-origin or material claims; confirm exact size and product type."
        )
        ws.cell(row_num, idx["source_refs"]).value = (
            f"products_export_1.csv sha256:{admin_hash}; Sang QA {QA_RUN_ID}; customizer_audit.json; contact sheets 121-130"
        )
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_R3_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = (
            "No Search Console, paid keyword volume or internal site-search export supplied; proof remains SERP-only."
        )
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    pk_by_pos = {qa["pos_by_key"][pk]: pk for pk in qa["product_keys"]}
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = pk_by_pos[pos]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = (
            f"Original storefront/product evidence; Sang QA {QA_RUN_ID}; products_export_1.csv sha256:{admin_hash}; "
            "customizer_audit.json; inspected contact sheets"
        )
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={option_text(admin_row)}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={item['detail']}; "
            f"customizer={customizer_sentence(qa['customizer_by_pos'].get(pos, {}))}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "Sang QA issues addressed in r3; independent QA recheck still required."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R3 uses inspected contact sheets, customizer audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_013_r3_revision", "r3", "Separate 10-product revision after Sang QA; source workbook was not modified."),
        ("qa_batch_013_r3_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_013_r3_scope", "inventory positions 121-130", "No products outside qa_batch_013 were revised."),
        ("qa_batch_013_r3_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_013_r3_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_013_r3_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_013 only; inventory positions 121-130",
        "revision": "r3",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Hold for combined QA handoff after revisions through qa_batch_020, per user request.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_013_r3",
                "",
                "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r2: `{SOURCE.relative_to(ROOT)}`",
                f"- QA của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r3: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Sửa trọng yếu: bỏ nội dung nội bộ, chỉ dùng claim name khi customizer audit xác nhận, gỡ claim cá nhân hóa ở sản phẩm 128, tách intent cho truck/cardinal/wolf, viết lại mô tả/alt theo contact sheets.",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Bước tiếp theo: giữ lại để gửi Sang QA chung sau khi revision đến `qa_batch_020`, theo yêu cầu người dùng.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
