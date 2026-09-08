import csv
import hashlib
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_020_r2"
QA_RUN_ID = "20260907_214100"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "batches" / "SEO_Product_Optimization_through_batch_020.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_020_r2.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_020_r2.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    191: {
        "title": "Custom Red Black Batter Bedding",
        "meta_title": "Custom Red Black Batter Bedding",
        "meta_description": "Shop custom red black baseball batter bedding with player silhouette, required name and number fields and size panel.",
        "primary": "custom red black batter bedding",
        "secondary": "baseball batter bedding, personalized baseball bedding, custom number baseball comforter",
        "cluster": "custom name number red black baseball batter bedding design 16",
        "detail": "red and black baseball bedding with batter silhouette, sample name and sample number on the comforter",
        "intent_role": "Baseball bedding page supported by required name and number fields and separated by red-black batter silhouette artwork.",
    },
    192: {
        "title": "Custom Brown Baseball Glove Bedding",
        "meta_title": "Custom Brown Baseball Glove Bedding",
        "meta_description": "Shop custom brown baseball glove bedding with close-up ball artwork, required name and number fields and size panel.",
        "primary": "custom brown baseball glove bedding",
        "secondary": "baseball glove bedding, personalized baseball bedding, custom name baseball comforter",
        "cluster": "custom name brown baseball glove close-up bedding design 17",
        "detail": "brown baseball glove bedding with close-up baseball, leather glove artwork and sample cursive name",
        "intent_role": "Baseball bedding page supported by required name and number fields and separated by brown glove close-up artwork.",
    },
    193: {
        "title": "Custom Flaming Baseball Batter Bedding",
        "meta_title": "Custom Flaming Baseball Batter Bedding",
        "meta_description": "Shop custom flaming baseball batter bedding with fire ring artwork, required name and number fields and feature panels.",
        "primary": "custom flaming baseball batter bedding",
        "secondary": "flaming baseball bedding, baseball player bedding, personalized baseball comforter",
        "cluster": "custom name number flaming baseball batter bedding design 18",
        "detail": "black bedding with flaming baseball ring, batter artwork, vertical sample name and sample number",
        "intent_role": "Baseball bedding page supported by required name and number fields and separated by flaming ball-and-batter artwork.",
    },
    194: {
        "title": "Custom Orange Fire Baseball Bedding",
        "meta_title": "Custom Orange Fire Baseball Bedding",
        "meta_description": "Shop custom orange fire baseball bedding with batter artwork, required name and number fields and feature panels.",
        "primary": "custom orange fire baseball bedding",
        "secondary": "fire baseball bedding, baseball player comforter, personalized baseball bedding",
        "cluster": "custom name number orange fire baseball player bedding design 19",
        "detail": "orange fire baseball bedding with batter player artwork, vertical sample name and sample number",
        "intent_role": "Baseball bedding page supported by required name and number fields and separated by orange fire player scene.",
    },
    195: {
        "title": "Custom Baseball Glove Flag Bedding",
        "meta_title": "Custom Baseball Glove Flag Bedding",
        "meta_description": "Shop custom baseball glove flag bedding with American flag artwork, required name and number fields and feature panels.",
        "primary": "custom baseball glove flag bedding",
        "secondary": "American flag baseball bedding, baseball glove comforter, personalized baseball bedding",
        "cluster": "custom name number baseball glove American flag bedding design 20",
        "detail": "American flag baseball bedding with glove and ball artwork, vertical sample name and sample number on shams",
        "intent_role": "Baseball bedding page supported by required name and number fields and separated by glove over flag layout.",
    },
    196: {
        "title": "Custom Black Baseball Glove Bedding",
        "meta_title": "Custom Black Baseball Glove Bedding",
        "meta_description": "Shop custom black baseball glove bedding with monochrome glove artwork, required name and number fields and feature panels.",
        "primary": "custom black baseball glove bedding",
        "secondary": "black baseball bedding, baseball glove comforter, personalized baseball bedding",
        "cluster": "custom name number black baseball glove bedding design 21",
        "detail": "black baseball bedding with monochrome glove, ball artwork, sample cursive name and sample number",
        "intent_role": "Baseball bedding page supported by required name and number fields and separated by black monochrome glove artwork.",
    },
    197: {
        "title": "Custom Lightning Baseball Bedding",
        "meta_title": "Custom Lightning Baseball Bedding",
        "meta_description": "Shop custom lightning baseball bedding with storm background, large baseball artwork and verified custom fields.",
        "primary": "custom lightning baseball bedding",
        "secondary": "lightning baseball bedding, personalized baseball comforter, baseball number bedding",
        "cluster": "custom name number lightning baseball bedding design 22",
        "detail": "stormy blue lightning baseball bedding with large baseball, sample cursive name and sample number",
        "intent_role": "Baseball bedding page supported by customizer audit and separated by lightning storm baseball artwork.",
    },
    198: {
        "title": "Custom Night Field Baseball Bedding",
        "meta_title": "Custom Night Field Baseball Bedding",
        "meta_description": "Shop custom night field baseball bedding with helmet, glove, ball artwork, required name and number fields.",
        "primary": "custom night field baseball bedding",
        "secondary": "night baseball bedding, baseball glove bedding, personalized baseball comforter",
        "cluster": "custom name number night field baseball bedding design 23",
        "detail": "night field baseball bedding with helmet, glove, bat, ball, grass edge and sample cursive name",
        "intent_role": "Baseball bedding page supported by required name and number fields and separated by night field equipment scene.",
    },
    199: {
        "title": "Custom Smoke Baseball Batter Bedding",
        "meta_title": "Custom Smoke Baseball Batter Bedding",
        "meta_description": "Shop custom smoke baseball batter bedding with dark player silhouette, required name and number fields and panels.",
        "primary": "custom smoke baseball batter bedding",
        "secondary": "smoke baseball bedding, baseball player bedding, personalized baseball comforter",
        "cluster": "custom name number smoke baseball batter bedding design 24",
        "detail": "dark smoke baseball bedding with batter silhouette, vertical sample name and sample number on pillow shams",
        "intent_role": "Baseball bedding page supported by required name and number fields and separated by smoky black batter scene.",
    },
    200: {
        "title": "Custom Baseball Stitch Name Bedding",
        "meta_title": "Custom Baseball Stitch Name Bedding",
        "meta_description": "Shop custom baseball stitch bedding with close-up laces, required name and number fields and feature panels.",
        "primary": "custom baseball stitch name bedding",
        "secondary": "baseball stitching bedding, personalized baseball bedding, baseball number comforter",
        "cluster": "custom name number baseball stitching bedding design 25",
        "detail": "close-up baseball stitching bedding with red laces, sample cursive name and sample number",
        "intent_role": "Baseball bedding page supported by required name and number fields and separated by close-up red stitching artwork.",
    },
}


COMMON_WEAVING = ("High-density weaving panel comparing BeddingOutlet fabric with other weave samples.", "High-density weaving feature panel")
COMMON_WASH = ("Machine washable care panel with washing machine, laundry basket and care text.", "Machine washable bedding care panel")
COMMON_ZIPPER = ("Bottom zippered closure panel showing white zipper close-up.", "Bottom zippered closure panel")
COMMON_EASY_CARE = ("Stress-free easy care panel listing wrinkle-free, stain-proof, anti-pilling and wash durability notes.", "Stress-free easy care bedding panel")
COMMON_BEDDING_TYPES = ("Two bedding types panel comparing duvet cover set and comforter set for all-season use.", "Duvet cover and comforter type panel")
COMMON_SIZE = ("Size dimension panel listing twin, full, queen and king measurements.", "Bedding size dimension panel")


IMAGE_DETAILS = {
    191: {
        1: ("Bedroom mockup of custom red black batter bedding with sample name and number.", "Custom red black batter bedding"),
        2: ("Second bedroom mockup of red black baseball batter bedding with white duvet fold.", "Red black baseball batter bedding mockup"),
        3: ("Close bedroom view of batter bedding with soft, lightweight, durable and breathable icons.", "Baseball batter bedding feature view"),
        4: ("Microfiber feature panel showing batter bedding sample, printed pattern insets and folded bedding.", "Baseball batter microfiber feature panel"),
        5: COMMON_EASY_CARE,
        6: COMMON_BEDDING_TYPES,
        7: COMMON_SIZE,
    },
    192: {
        1: ("Bedroom mockup of custom brown baseball glove bedding with close-up ball and sample name.", "Custom brown baseball glove bedding"),
        2: ("Second bedroom mockup of brown baseball glove bedding with white duvet fold.", "Brown baseball glove bedding mockup"),
        3: ("Close bedroom view of brown glove bedding with soft, lightweight, durable and breathable icons.", "Brown baseball glove bedding feature view"),
        4: ("Microfiber feature panel showing glove bedding sample, printed pattern insets and folded bedding.", "Baseball glove microfiber feature panel"),
        5: COMMON_EASY_CARE,
        6: COMMON_BEDDING_TYPES,
        7: COMMON_SIZE,
    },
    193: {
        1: ("Bedroom mockup of custom flaming baseball batter bedding with fire ring, sample name and number.", "Custom flaming baseball batter bedding"),
        2: COMMON_ZIPPER,
        3: COMMON_WEAVING,
        4: COMMON_WASH,
    },
    194: {
        1: ("Bedroom mockup of custom orange fire baseball bedding with batter artwork, sample name and number.", "Custom orange fire baseball bedding"),
        2: COMMON_ZIPPER,
        3: COMMON_WEAVING,
        4: COMMON_WASH,
    },
    195: {
        1: ("Bedroom mockup of custom baseball glove flag bedding with glove, ball, sample name and number.", "Custom baseball glove flag bedding"),
        2: COMMON_WASH,
        3: COMMON_ZIPPER,
        4: COMMON_WEAVING,
    },
    196: {
        1: ("Bedroom mockup of custom black baseball glove bedding with monochrome glove, sample name and number.", "Custom black baseball glove bedding"),
        2: COMMON_WEAVING,
        3: COMMON_WASH,
        4: COMMON_ZIPPER,
    },
    197: {
        1: ("Bedroom mockup of custom lightning baseball bedding with storm background, sample name and number.", "Custom lightning baseball bedding"),
        2: COMMON_ZIPPER,
        3: COMMON_WEAVING,
        4: COMMON_WASH,
    },
    198: {
        1: ("Bedroom mockup of custom night field baseball bedding with helmet, glove, bat, ball and sample name.", "Custom night field baseball bedding"),
        2: COMMON_WEAVING,
        3: COMMON_WASH,
        4: COMMON_ZIPPER,
    },
    199: {
        1: ("Bedroom mockup of custom smoke baseball batter bedding with dark player silhouette and sample name.", "Custom smoke baseball batter bedding"),
        2: COMMON_ZIPPER,
        3: COMMON_WEAVING,
        4: COMMON_WASH,
    },
    200: {
        1: ("Bedroom mockup of custom baseball stitch bedding with red laces, sample name and number.", "Custom baseball stitch name bedding"),
        2: COMMON_WEAVING,
        3: COMMON_WASH,
        4: COMMON_ZIPPER,
    },
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def load_qa():
    data = json.loads(QA_DATASET.read_text(encoding="utf-8"))
    products = data["QA_Products"]
    return {
        "product_keys": [item["product_key"] for item in products],
        "handle_by_key": {item["product_key"]: item["handle"] for item in products},
        "pos_by_key": {item["product_key"]: int(item["inventory_position"]) for item in products},
    }


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    missing = sorted(set(handles) - set(admin))
    if missing:
        raise RuntimeError(f"Missing handles in admin export: {missing}")
    return admin


def option_text(admin_row):
    values = [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]]
    values = [v for v in values if v]
    return ", ".join(values) if values else "Shopify export shows no named option group."


def description(pos, admin_row):
    item = PRODUCT_UPDATES[pos]
    return (
        f"<p>{item['title']} features {item['detail']}. "
        "The page copy focuses on the specific visible design, product type and verified custom fields.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork: {item['detail']}.</li>"
        "<li>Gallery images show the main baseball bedding mockup plus feature, care, bedding-type, size or zipper panels where present.</li></ul>"
        "<h3>Customization and Fit</h3>"
        f"<ul><li>Available option groups: {option_text(admin_row)}.</li>"
        "<li>Customizer audit captured Customize Your Name and Customize Your Number fields for this product group.</li>"
        "<li>Select the product type and size shown on the product page before checkout.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    qa = load_qa()
    admin = load_admin(set(qa["handle_by_key"].values()))
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or "Bedding"
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"US English buyer intent targets {item['cluster']}; no search-volume, trend or ranking claim is made."
        )
        ws.cell(row_num, idx["keyword_evidence_level"]).value = "SERP_ONLY"
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R2 after Sang QA batch 020: used products_export_1.csv sha256:{admin_hash}; "
            "removed internal workflow copy, filled publishable title values, kept custom name/number claims because customizer audit verifies them, "
            "split baseball intents by visible motif and rewrote image observations/alts from inspected contact sheets. Still NEEDS_REVIEW."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        image_number = int(ws.cell(row_num, idx["image_number"]).value)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation, alt = IMAGE_DETAILS[pos][image_number]
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = alt
        ws.cell(row_num, idx["alt_action"]).value = "SET"
        ws.cell(row_num, idx["viewed_status"]).value = "VIEWED"
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R2: observation and alt rewritten from inspected contact sheet; admin image URL and current alt matched from Shopify CSV where available."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "PRODUCT"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = (
            "Sang QA batch 020 + products_export_1.csv admin baseline + customizer audit + contact-sheet inspection"
        )
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = (
            "R2 separates nearby baseball pages by motif, layout, color palette and required custom fields."
        )
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r2"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as a specific product page."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm exact baseball artwork, product type, size choices, "
            "and custom name and number fields before purchase."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = (
            "Confirm product type, visible baseball motif, custom name and number fields, care panels and image accuracy before purchase."
        )
        ws.cell(row_num, idx["emotional_social_motivation"]).value = (
            "Choose a personalized baseball bedding item that matches a preferred visual style, such as batter, glove, lightning or stitching artwork."
        )
        ws.cell(row_num, idx["purchase_concerns"]).value = (
            "Confirm spelling and number before checkout; avoid unsupported team, player, league, material or delivery claims."
        )
        ws.cell(row_num, idx["source_refs"]).value = (
            f"products_export_1.csv sha256:{admin_hash}; Sang QA {QA_RUN_ID}; customizer audit; contact sheets 191-200"
        )
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = (
            "No Search Console, paid keyword volume or internal site-search export supplied; proof remains SERP-only."
        )
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    pk_by_pos = {qa["pos_by_key"][pk]: pk for pk in qa["product_keys"]}
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = pk_by_pos[pos]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = (
            f"Original storefront/product evidence; Sang QA {QA_RUN_ID}; products_export_1.csv sha256:{admin_hash}; "
            "customizer audit; inspected contact sheets"
        )
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={option_text(admin_row)}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={item['detail']}; "
            "custom_supported=true; customizer_note=Customize Your Name and Customize Your Number fields captured"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "Sang QA issues addressed in r2; independent QA recheck still required."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R2 uses inspected contact sheets, customizer audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_020_r2_revision", "r2", "Separate 10-product revision after Sang QA; source workbook was not modified."),
        ("qa_batch_020_r2_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_020_r2_scope", "inventory positions 191-200", "No products outside qa_batch_020 were revised."),
        ("qa_batch_020_r2_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_020_r2_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_020_r2_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_020 only; inventory positions 191-200",
        "revision": "r2",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Ready to hand off revisions qa_batch_008_r2 through qa_batch_020_r2 for QA, per user request.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_020_r2",
                "",
                "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r1: `{SOURCE.relative_to(ROOT)}`",
                f"- QA của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r2: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Sửa trọng yếu: bỏ nội dung nội bộ, bổ sung title/H1 publishable, giữ custom name/number theo customizer audit, tách intent cho 10 thiết kế baseball theo motif nhìn thấy và viết lại mô tả/alt theo contact sheets.",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Bước tiếp theo: có thể gửi Sang QA lại nhóm revision `qa_batch_008_r2` đến `qa_batch_020_r2`, theo yêu cầu người dùng.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
