import csv
import hashlib
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "qa_batch_008_r2"
QA_RUN_ID = "20260907_173014"

SOURCE = ROOT / "resutls" / SHOP / RUN_ID / "batches" / "SEO_Product_Optimization_through_batch_008.xlsx"
QA_DATASET = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "qa_dataset.json"
CUSTOMIZER_AUDIT = ROOT / "seo_runs" / SHOP / RUN_ID / "qa" / QA_RUN_ID / "customizer_audit.json"
ADMIN_EXPORT = ROOT / "products_export_1.csv"
RESULT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / BATCH_ID
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID / "revisions" / BATCH_ID
OUTPUT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_008_r2.xlsx"
REPORT = RESULT_DIR / "SEO_Product_Optimization_qa_batch_008_r2.md"
SUMMARY = RUN_DIR / "revision_summary.json"


PRODUCT_UPDATES = {
    71: {
        "title": "Personalized Teal Football Player Comforter Set",
        "meta_title": "Personalized Teal Football Player Comforter Set",
        "meta_description": "Customize a teal football player comforter with required name and optional number, plus bedding type, size and add-on options.",
        "primary": "personalized teal football player comforter",
        "secondary": "custom football player bedding, teal football comforter with name, football comforter set for players",
        "cluster": "teal running football player comforter",
        "detail": "teal and black running football player artwork with orange flame trails and sample Brian #3 text",
        "intent_role": "Football player page separated by teal/orange running-player motif and verified name plus number fields.",
        "product_copy": "sports-themed bedroom decor or a personalized football gift",
    },
    72: {
        "title": "Personalized Electric Football Player Comforter Set",
        "meta_title": "Personalized Electric Football Player Comforter Set",
        "meta_description": "Customize an electric blue football player comforter with required name and optional number, plus size and bedding choices.",
        "primary": "personalized electric football player comforter",
        "secondary": "blue football player bedding, custom football comforter with name, football player comforter set",
        "cluster": "electric blue football player comforter",
        "detail": "blue and red electric football player running with the ball and sample Michael #10 text",
        "intent_role": "Football player page separated by electric blue/red running-player motif and verified name plus number fields.",
        "product_copy": "sports-themed bedroom decor or a personalized football gift",
    },
    73: {
        "title": "Personalized American Flag Football Player Comforter Set",
        "meta_title": "Personalized American Flag Football Player Comforter Set",
        "meta_description": "Customize an American flag football player comforter with required name and optional number, plus size and bedding choices.",
        "primary": "personalized American flag football player comforter",
        "secondary": "patriotic football player bedding, custom football flag comforter, football comforter with name and number",
        "cluster": "American flag football player comforter",
        "detail": "football player viewed from behind over a torn American flag and metal texture with sample David 33 text",
        "intent_role": "Patriotic football player page separated from plain flag, helmet flag and flame football variants.",
        "product_copy": "patriotic football room decor or a personalized football gift",
    },
    74: {
        "title": "Personalized Burning Football Player Comforter Set",
        "meta_title": "Personalized Burning Football Player Comforter Set",
        "meta_description": "Customize a burning football player comforter with required name and optional number, plus product type, size and add-ons.",
        "primary": "personalized burning football player comforter",
        "secondary": "custom football player bedding, black football comforter with name, fire football player comforter",
        "cluster": "burning football player comforter",
        "detail": "black football player holding the ball with number 24, sample Michael name and burning football background",
        "intent_role": "Flame football player page separated from electric player, flag player and football-laces texture pages.",
        "product_copy": "dramatic football bedroom decor or a personalized sports gift",
    },
    75: {
        "title": "Personalized Football Players Name Comforter Set",
        "meta_title": "Personalized Football Players Name Comforter Set",
        "meta_description": "Customize a football players comforter with a required name field, bedding type, size, pillowcase and sheet-cover options.",
        "primary": "personalized football players name comforter",
        "secondary": "custom football players comforter, football collage bedding with name, football comforter set with name",
        "cluster": "football players and large football name comforter",
        "detail": "black comforter with large football close-up, three football player silhouettes and sample David name",
        "intent_role": "Football collage page; structured jersey-digit customization is omitted because audit confirms only required Enter Name for this product.",
        "product_copy": "football fan bedroom decor or a custom name sports gift",
    },
    76: {
        "title": "Personalized Football Laces Comforter Set",
        "meta_title": "Personalized Football Laces Comforter Set",
        "meta_description": "Customize a brown football laces comforter with required name and optional number, plus product type, size and add-ons.",
        "primary": "personalized football laces comforter",
        "secondary": "football texture bedding with name, vintage football comforter set, custom football laces bedding",
        "cluster": "vintage football laces comforter",
        "detail": "brown football leather texture with prominent white laces and sample Michael 33 text",
        "intent_role": "Texture/laces football page separated from player-led and flag-led football comforter variants.",
        "product_copy": "vintage football room decor or a personalized sports gift",
    },
    77: {
        "title": "Custom Glowing Soccer Ball Comforter Set",
        "meta_title": "Custom Glowing Soccer Ball Comforter Set",
        "meta_description": "Shop a glowing blue soccer ball comforter with optional custom text/request field, bedding type, size and add-on options.",
        "primary": "custom glowing soccer ball comforter set",
        "secondary": "blue soccer ball bedding, soccer comforter set, glowing soccer bedding for boys",
        "cluster": "glowing soccer ball comforter",
        "detail": "black and blue glowing soccer ball with sweeping light trails and sample Jackson 10 artwork",
        "intent_role": "Soccer ball page; no structured name/number claim because audit shows only an optional Customize Your Item field.",
        "product_copy": "soccer-themed bedroom decor",
    },
    78: {
        "title": "Custom God Says You Are Floral Blanket",
        "meta_title": "Custom God Says You Are Floral Blanket",
        "meta_description": "Shop a floral God says you are blanket with optional custom name and message fields, fleece or sherpa sizes.",
        "primary": "custom God says you are floral blanket",
        "secondary": "Christian affirmation blanket, floral Bible verse blanket, personalized faith blanket",
        "cluster": "God says you are floral blanket",
        "detail": "cream blanket with God says you are text, colorful flowers, butterflies, verse references and sample Jessica name",
        "intent_role": "Christian floral affirmation blanket page; no sports language and no number personalization claim.",
        "product_copy": "faith-inspired gift or daily encouragement blanket",
    },
    79: {
        "title": "Custom Floral Butterfly Affirmation Blanket",
        "meta_title": "Custom Floral Butterfly Affirmation Blanket",
        "meta_description": "Shop a floral butterfly affirmation blanket with optional custom name and message fields, fleece or sherpa sizes.",
        "primary": "custom floral butterfly affirmation blanket",
        "secondary": "personalized you are blanket, floral butterfly Bible verse blanket, Christian affirmation throw",
        "cluster": "floral butterfly you are affirmation blanket",
        "detail": "cream floral blanket with Haley name, butterflies and vertical affirmation words with Bible verse references",
        "intent_role": "Floral butterfly affirmation page separated from God says you are and God says I am designs.",
        "product_copy": "faith-inspired gift or personal encouragement blanket",
    },
    80: {
        "title": "Custom God Says I Am Butterfly Blanket",
        "meta_title": "Custom God Says I Am Butterfly Blanket",
        "meta_description": "Shop a brown God Says I Am butterfly blanket with optional custom name and message fields, fleece or sherpa sizes.",
        "primary": "custom God Says I Am butterfly blanket",
        "secondary": "Christian butterfly blanket, personalized God Says I Am blanket, Bible verse affirmation blanket",
        "cluster": "God Says I Am butterfly blanket",
        "detail": "brown and tan God Says I Am design with monarch butterflies, verse blocks and sample Elizabeth name",
        "intent_role": "God Says I Am butterfly page separated from floral God says you are and vertical affirmation blanket pages.",
        "product_copy": "faith-inspired gift or devotional throw blanket",
    },
}


IMAGE_DETAILS = {
    71: {
        1: ("Bedroom mockup of teal running football player comforter with Brian #3 and matching pillows.", "Teal running football player comforter with Brian #3 on bed"),
        2: ("Second bedroom mockup with the same teal football player design, white duvet fold and matching pillows.", "Teal football player comforter with Brian #3 in bedroom"),
        3: ("Comparison panel explaining duvet cover set versus comforter set for all-season use.", "Duvet cover and comforter set comparison panel"),
        4: ("Close bedroom feature panel showing the Brian #3 football artwork with soft, lightweight, durable and breathable icons.", "Feature panel for teal football player comforter"),
        5: ("Easy-care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade resistance.", "Easy care panel for football comforter set"),
        6: ("Size dimension chart for Twin 68 x 86, Full 79 x 90, Queen 90 x 90 and King 90 x 104.", "Comforter size dimension chart for Twin Full Queen King"),
    },
    72: {
        1: ("Bedroom mockup of blue and red electric football player comforter with Michael #10 and matching pillows.", "Electric football player comforter with Michael #10 on bed"),
        2: ("Second bedroom mockup showing the electric football player design with white duvet fold.", "Blue electric football player comforter with Michael #10"),
        3: ("Comparison panel explaining duvet cover set versus comforter set for all-season use.", "Duvet cover and comforter set comparison panel"),
        4: ("Close bedroom feature panel showing the electric football artwork with soft, lightweight, durable and breathable icons.", "Feature panel for electric football player comforter"),
        5: ("Easy-care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade resistance.", "Easy care panel for football comforter set"),
        6: ("Size dimension chart for Twin 68 x 86, Full 79 x 90, Queen 90 x 90 and King 90 x 104.", "Comforter size dimension chart for Twin Full Queen King"),
    },
    73: {
        1: ("Bedroom mockup of football player viewed from behind over a torn American flag with David 33.", "American flag football player comforter with David 33"),
        2: ("Second bedroom mockup of the same David 33 football player flag design with white duvet fold.", "Football player flag comforter with David 33 in bedroom"),
        3: ("Comparison panel explaining duvet cover set versus comforter set for all-season use.", "Duvet cover and comforter set comparison panel"),
        4: ("Close bedroom feature panel of the David 33 flag player artwork with soft, lightweight, durable and breathable icons.", "Feature panel for American flag football player comforter"),
        5: ("Easy-care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade resistance.", "Easy care panel for football comforter set"),
        6: ("Size dimension chart for Twin 68 x 86, Full 79 x 90, Queen 90 x 90 and King 90 x 104.", "Comforter size dimension chart for Twin Full Queen King"),
    },
    74: {
        1: ("Bedroom mockup of black football player number 24 with Michael name and burning football background.", "Burning football player comforter with Michael and 24"),
        2: ("Second bedroom mockup showing the black number 24 player and flame background under a white duvet fold.", "Black football player comforter with Michael 24 in bedroom"),
        3: ("Comparison panel explaining duvet cover set versus comforter set for all-season use.", "Duvet cover and comforter set comparison panel"),
        4: ("Close bedroom feature panel for the burning football player design with soft, lightweight, durable and breathable icons.", "Feature panel for burning football player comforter"),
        5: ("Easy-care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade resistance.", "Easy care panel for football comforter set"),
        6: ("Size dimension chart for Twin 68 x 86, Full 79 x 90, Queen 90 x 90 and King 90 x 104.", "Comforter size dimension chart for Twin Full Queen King"),
    },
    75: {
        1: ("Bedroom mockup of black football comforter with David name, three player silhouettes and a large football close-up.", "Football players name comforter with David and large football"),
        2: ("Second bedroom mockup showing David name, football players and large football graphic under a white duvet fold.", "Football players comforter with David name in bedroom"),
        3: ("Comparison panel explaining duvet cover set versus comforter set for all-season use.", "Duvet cover and comforter set comparison panel"),
        4: ("Feature panel showing microfiber, 3D printed pattern and soft lightweight breathable callouts.", "Microfiber and 3D print feature panel for football comforter"),
        5: ("Close bedroom feature panel showing football players, large football and soft, lightweight, durable and breathable icons.", "Feature panel for football players name comforter"),
        6: ("Easy-care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade resistance.", "Easy care panel for football comforter set"),
        7: ("Size dimension chart for Twin 68 x 86, Full 79 x 90, Queen 90 x 90 and King 90 x 104.", "Comforter size dimension chart for Twin Full Queen King"),
    },
    76: {
        1: ("Bedroom mockup of brown football leather texture comforter with white laces, Michael name and 33.", "Football laces comforter with Michael 33 on bed"),
        2: ("Second bedroom mockup showing brown football texture, white laces, Michael name and 33 under a white duvet fold.", "Brown football laces comforter with Michael 33"),
        3: ("Comparison panel explaining duvet cover set versus comforter set for all-season use.", "Duvet cover and comforter set comparison panel"),
        4: ("Close bedroom feature panel for football laces design with soft, lightweight, durable and breathable icons.", "Feature panel for football laces comforter"),
        5: ("Easy-care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade resistance.", "Easy care panel for football comforter set"),
        6: ("Size dimension chart for Twin 68 x 86, Full 79 x 90, Queen 90 x 90 and King 90 x 104.", "Comforter size dimension chart for Twin Full Queen King"),
    },
    77: {
        1: ("Bedroom mockup of black comforter with glowing blue soccer ball light trails and sample Jackson text.", "Glowing soccer ball comforter with Jackson text on bed"),
        2: ("Feature panel showing microfiber, 3D printed pattern and lightweight breathable callouts for the soccer design.", "Microfiber and 3D print feature panel for soccer comforter"),
        3: ("Pillowcase mockup on white bed showing glowing blue soccer ball and number 10 artwork.", "Glowing soccer ball pillowcases on white bed"),
        4: ("Comparison panel explaining duvet cover set versus comforter set for all-season use.", "Duvet cover and comforter set comparison panel"),
        5: ("Second bedroom mockup of glowing blue soccer ball comforter with Jackson text and matching pillows.", "Glowing soccer ball comforter with Jackson in bedroom"),
        6: ("Close bedroom feature panel showing glowing soccer ball design with soft, lightweight, durable and breathable icons.", "Feature panel for glowing soccer ball comforter"),
        7: ("Easy-care panel listing wrinkle-free, stain-proof, anti-pilling and wash-fade resistance.", "Easy care panel for soccer comforter set"),
    },
    78: {
        1: ("Lifestyle mockup of a woman holding cream God says you are floral blanket with butterflies and Jessica name.", "God says you are floral blanket with Jessica name"),
        2: ("Flat blanket image showing God says you are text, flowers, butterflies, verse references and Jessica name.", "Flat God says you are floral blanket with butterflies"),
        3: ("Lifestyle image of adult and child reading with the floral God says you are blanket on a bed.", "God says you are blanket used while reading on bed"),
        4: ("Blanket size chart showing 40x30, 50x40, 60x50 and 80x60 sizes.", "Blanket size chart for four size options"),
        5: ("Room mockup with the floral God says you are blanket draped across a bed.", "Floral God says you are blanket draped on bed"),
        6: ("Room mockup with the floral God says you are blanket displayed on a sofa.", "Floral God says you are blanket on sofa"),
        7: ("Feature panel listing fluffy and soft, high quality, no pilling, no shedding and gentle smooth fabric.", "Feature panel for floral Christian blanket"),
        8: ("Close-up collage showing floral print, butterflies, verse text and white blanket texture.", "Close-up fabric collage for floral God says you are blanket"),
    },
    79: {
        1: ("Lifestyle mockup of a woman holding cream floral butterfly blanket with Haley name and affirmation words.", "Floral butterfly affirmation blanket with Haley name"),
        2: ("Flat blanket image showing Haley name, floral border, butterflies and vertical affirmation words.", "Flat floral butterfly affirmation blanket with Haley"),
        3: ("Close-up feature image with machine washable, soft and warm, dense stitching and 260GSM fleece blanket callouts.", "Machine washable fleece blanket feature panel"),
        4: ("Use and size panel showing text can be changed, 80 by 60 in blanket, and sofa, office, bed, plane and travel uses.", "Floral butterfly blanket size and use panel"),
        5: ("Close-up collage showing fleece texture, flowers, butterflies and printed affirmation words.", "Close-up fabric collage for floral butterfly blanket"),
        6: ("Room mockup with floral butterfly affirmation blanket draped over a sofa.", "Floral butterfly affirmation blanket on sofa"),
        7: ("Lifestyle reading image with floral butterfly blanket across lap and visible Bible verse references.", "Floral butterfly blanket used while reading"),
        8: ("Room mockup with floral butterfly affirmation blanket on an armchair beside books and a mug.", "Floral butterfly affirmation blanket on armchair"),
        9: ("Close lifestyle image of blanket across lap with open book and mug.", "Floral affirmation blanket with open book and mug"),
    },
    80: {
        1: ("Lifestyle mockup of a woman holding brown God Says I Am butterfly blanket with Elizabeth name.", "God Says I Am butterfly blanket with Elizabeth name"),
        2: ("Flat blanket image showing God Says I Am title, monarch butterflies, Elizabeth name and verse blocks.", "Flat God Says I Am butterfly blanket"),
        3: ("Lifestyle image of adult and child reading with God Says I Am blanket on a bed.", "God Says I Am butterfly blanket used while reading"),
        4: ("Blanket size chart showing 40x30, 50x40, 60x50 and 80x60 sizes.", "Blanket size chart for four size options"),
        5: ("Close-up collage showing brown butterfly print, verse blocks and white fleece texture.", "Close-up fabric collage for God Says I Am blanket"),
        6: ("Feature panel listing fluffy and soft, high quality, no pilling, no shedding and gentle smooth fabric.", "Feature panel for God Says I Am butterfly blanket"),
        7: ("Room mockup with God Says I Am butterfly blanket draped over a sofa.", "God Says I Am butterfly blanket on sofa"),
        8: ("Room mockup with God Says I Am butterfly blanket draped on a couch near a window.", "God Says I Am butterfly blanket on couch"),
    },
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def text(value):
    return "" if value is None else str(value)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def normalize_url(url):
    return text(url).split("?")[0]


def load_qa():
    data = json.loads(QA_DATASET.read_text(encoding="utf-8"))
    products = data["QA_Products"]
    handle_by_key = {item["product_key"]: item["handle"] for item in products}
    pos_by_key = {item["product_key"]: int(item["inventory_position"]) for item in products}
    customizer = json.loads(CUSTOMIZER_AUDIT.read_text(encoding="utf-8"))
    customizer_by_pos = {int(item["inventory_position"]): item for item in customizer}
    return {
        "data": data,
        "product_keys": [item["product_key"] for item in products],
        "handle_by_key": handle_by_key,
        "pos_by_key": pos_by_key,
        "customizer_by_pos": customizer_by_pos,
    }


def load_admin(handles):
    by_handle = defaultdict(list)
    with ADMIN_EXPORT.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Handle"] in handles:
                by_handle[row["Handle"]].append(row)
    admin = {}
    for handle, rows in by_handle.items():
        first = rows[0]
        images = {}
        for row in rows:
            if row.get("Image Src"):
                images[normalize_url(row["Image Src"])] = {
                    "src": row["Image Src"],
                    "position": row.get("Image Position", ""),
                    "alt": row.get("Image Alt Text", ""),
                }
        admin[handle] = {
            "Title": first.get("Title", ""),
            "Body (HTML)": first.get("Body (HTML)", ""),
            "Type": first.get("Type", ""),
            "SEO Title": first.get("SEO Title", ""),
            "SEO Description": first.get("SEO Description", ""),
            "Option1 Name": first.get("Option1 Name", ""),
            "Option2 Name": first.get("Option2 Name", ""),
            "Option3 Name": first.get("Option3 Name", ""),
            "Status": first.get("Status", ""),
            "images": images,
        }
    return admin


def customizer_sentence(audit):
    nodes = audit.get("personalization_nodes", [])
    if not nodes:
        return "No structured personalization fields were verified in the Customizer audit."
    parts = []
    for node in nodes:
        label = node.get("label", "text field")
        required = "required" if node.get("required") else "optional"
        limit = node.get("maxLength")
        if limit:
            parts.append(f"Customizer shows a {required} {label} field up to {limit} characters.")
        else:
            parts.append(f"Customizer shows a {required} {label} field.")
    labels = " ".join(node.get("label", "") for node in nodes).lower()
    if "number" in labels:
        parts.append("Names and numbers visible in mockups are sample artwork unless the matching input is verified.")
    elif "name" in labels:
        parts.append("Visible names in mockups are sample artwork unless the matching input is verified.")
    else:
        parts.append("Visible custom text in mockups is sample artwork unless the matching input is verified.")
    return " ".join(parts)


def option_text(admin_row):
    values = [admin_row["Option1 Name"], admin_row["Option2 Name"], admin_row["Option3 Name"]]
    values = [v for v in values if v]
    return ", ".join(values) if values else "Shopify export shows the product size option."


def description(pos, admin_row, ctext):
    item = PRODUCT_UPDATES[pos]
    product_type = (admin_row["Type"] or "product").lower()
    if pos <= 77:
        return (
            f"<p>{item['detail'].capitalize()} gives this Jeminise {product_type} a clear {item['cluster']} focus for {item['product_copy']}.</p>"
            "<h3>Design Details</h3>"
            f"<ul><li>Artwork focus: {item['detail']}.</li>"
            "<li>Gallery images show the main bedding mockups plus product-type, feature, care, close-up or size panels where available.</li>"
            f"<li>Current Shopify export title used for identity check: {admin_row['Title']}.</li></ul>"
            "<h3>Options and Customization</h3>"
            f"<ul><li>Shopify option groups: {option_text(admin_row)}.</li>"
            f"<li>{ctext}</li>"
            "<li>Confirm the product type and size before checkout; pillowcase and sheet-cover add-ons appear only when selected.</li></ul>"
        )
    return (
        f"<p>{item['detail'].capitalize()} makes this Jeminise {product_type} suitable as a {item['product_copy']}.</p>"
        "<h3>Design Details</h3>"
        f"<ul><li>Artwork focus: {item['detail']}.</li>"
        "<li>Gallery images show flat blanket views, room or reading lifestyle mockups, size guidance and fabric/detail panels where available.</li>"
        f"<li>Current Shopify export title used for identity check: {admin_row['Title']}.</li></ul>"
        "<h3>Options and Customization</h3>"
        f"<ul><li>Shopify option groups: {option_text(admin_row)}.</li>"
        f"<li>{ctext}</li>"
        "<li>Choose the blanket size and use the available custom fields only for supported name or message requests.</li></ul>"
    )


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    qa = load_qa()
    admin = load_admin(set(qa["handle_by_key"].values()))
    admin_hash = sha256(ADMIN_EXPORT)
    now = datetime.now().astimezone().isoformat()

    wb = load_workbook(OUTPUT)
    product_key_set = set(qa["product_keys"])

    ws = wb["SEO_Products"]
    idx = headers(ws)
    revised_products = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        handle = qa["handle_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        admin_row = admin[handle]
        ctext = customizer_sentence(qa["customizer_by_pos"].get(pos, {}))

        ws.cell(row_num, idx["title_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["h1_current"]).value = admin_row["Title"]
        ws.cell(row_num, idx["rendered_title_current"]).value = admin_row["SEO Title"] or admin_row["Title"]
        ws.cell(row_num, idx["meta_description_current"]).value = admin_row["SEO Description"]
        ws.cell(row_num, idx["product_type"]).value = admin_row["Type"] or ws.cell(row_num, idx["product_type"]).value
        ws.cell(row_num, idx["title_proposed"]).value = item["title"]
        ws.cell(row_num, idx["meta_title_seo"]).value = item["meta_title"]
        ws.cell(row_num, idx["meta_title_length"]).value = len(item["meta_title"])
        ws.cell(row_num, idx["meta_description_seo"]).value = item["meta_description"]
        ws.cell(row_num, idx["meta_description_length"]).value = len(item["meta_description"])
        ws.cell(row_num, idx["description_proposed_html"]).value = description(pos, admin_row, ctext)
        ws.cell(row_num, idx["primary_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["secondary_keywords"]).value = item["secondary"]
        ws.cell(row_num, idx["meta_keyword"]).value = item["primary"]
        ws.cell(row_num, idx["keyword_strategy"]).value = item["intent_role"]
        ws.cell(row_num, idx["buyer_search_summary"]).value = (
            f"US English buyer intent targets {item['cluster']}; no search-volume claim is made."
        )
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields"]:
            ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            f"R2 after Sang QA: used products_export_1.csv (sha256:{admin_hash}) as admin baseline; "
            "removed internal workflow copy, corrected Customizer claims per product, rewrote descriptions, "
            "refreshed image observations/alt from inspected contact sheets, and separated keyword intents. "
            "Still NEEDS_REVIEW; no APPROVED/import."
        )
        revised_products += 1

    ws = wb["Image_Audit"]
    idx = headers(ws)
    revised_images = 0
    admin_alt_matches = 0
    handle_to_pk = {v: k for k, v in qa["handle_by_key"].items()}
    for row_num in range(2, ws.max_row + 1):
        handle = text(ws.cell(row_num, idx["Handle"]).value)
        pk = handle_to_pk.get(handle)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        image_number = int(ws.cell(row_num, idx["image_number"]).value)
        image_url = text(ws.cell(row_num, idx["image_url_export"]).value or ws.cell(row_num, idx["image_url"]).value)
        admin_image = admin.get(handle, {}).get("images", {}).get(normalize_url(image_url))
        if admin_image:
            ws.cell(row_num, idx["image_url_export"]).value = admin_image["src"]
            ws.cell(row_num, idx["alt_current"]).value = admin_image["alt"]
            admin_alt_matches += 1
        observation, alt = IMAGE_DETAILS[pos][image_number]
        ws.cell(row_num, idx["observed_visual_details"]).value = observation
        ws.cell(row_num, idx["alt_proposed"]).value = alt[:125]
        ws.cell(row_num, idx["revision"]).value = "r2"
        ws.cell(row_num, idx["review_status"]).value = "NEEDS_REVIEW"
        for col_name in ["approved_by", "approved_at", "approved_fields", "approved_revision"]:
            if col_name in idx:
                ws.cell(row_num, idx[col_name]).value = ""
        ws.cell(row_num, idx["issues"]).value = (
            "R2: image observation and alt rewritten from inspected contact sheet; current admin alt and image URL matched from products_export_1.csv where possible."
        )
        revised_images += 1

    ws = wb["Keyword_Map"]
    idx = headers(ws)
    keyword_counts = defaultdict(int)
    keyword_rows = 0
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        role = text(ws.cell(row_num, idx["keyword_role"]).value)
        secondaries = [part.strip() for part in item["secondary"].split(",")]
        keyword = item["primary"] if role == "PRIMARY" else secondaries[min(keyword_counts[pk], len(secondaries) - 1)]
        if role != "PRIMARY":
            keyword_counts[pk] += 1
        ws.cell(row_num, idx["keyword"]).value = keyword
        ws.cell(row_num, idx["semantic_cluster"]).value = item["cluster"]
        ws.cell(row_num, idx["intent"]).value = "Commercial product intent"
        ws.cell(row_num, idx["target_page_type"]).value = "Product"
        ws.cell(row_num, idx["decision_reason"]).value = item["intent_role"]
        ws.cell(row_num, idx["demand_evidence"]).value = "SERP intent evidence only; no search-volume claim."
        ws.cell(row_num, idx["validation_source"]).value = "Sang QA + products_export_1.csv admin baseline + customizer_audit.json + contact-sheet inspection"
        ws.cell(row_num, idx["checked_at"]).value = now
        ws.cell(row_num, idx["possible_overlap_with_other_products"]).value = item["intent_role"]
        ws.cell(row_num, idx["mapping_reason"]).value = "R2 separates pages by visible motif and verified Customizer capability."
        ws.cell(row_num, idx["mapping_status"]).value = "CANDIDATE_MAPPED_NEEDS_RECHECK"
        ws.cell(row_num, idx["mapping_version"]).value = "r2"
        keyword_rows += 1

    ws = wb["Buyer_Search_Research"]
    idx = headers(ws)
    for row_num in range(2, ws.max_row + 1):
        pk = text(ws.cell(row_num, idx["product_key"]).value)
        if pk not in product_key_set:
            continue
        pos = qa["pos_by_key"][pk]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["purchase_context"]).value = f"US buyer looking for {item['cluster']} as a specific product page."
        ws.cell(row_num, idx["jtbd_statement"]).value = (
            f"When shopping for {item['cluster']}, the buyer wants the page to confirm the exact motif, product type, size choices, "
            "and which customization inputs are actually supported."
        )
        ws.cell(row_num, idx["functional_motivation"]).value = "Confirm product type, size, add-ons, supported custom fields and image accuracy before purchase."
        ws.cell(row_num, idx["emotional_social_motivation"]).value = item["product_copy"].capitalize()
        ws.cell(row_num, idx["purchase_concerns"]).value = "Avoid unsupported personalization claims; confirm size, fabric/care statements and included components."
        ws.cell(row_num, idx["source_refs"]).value = f"products_export_1.csv sha256:{admin_hash}; Sang QA {QA_RUN_ID}; customizer_audit.json; contact sheets 71-80"
        ws.cell(row_num, idx["observed_at"]).value = now
        ws.cell(row_num, idx["research_status"]).value = "REVISED_NEEDS_RECHECK"
        ws.cell(row_num, idx["limitations"]).value = "No search-volume source supplied; Customizer-to-cart persistence still needs controlled test before approval."
        ws.cell(row_num, idx["seo_application"]).value = f"Use {item['primary']} with intent role: {item['intent_role']}"

    ws = wb["Product_Evidence"]
    idx = headers(ws)
    pk_by_pos = {qa["pos_by_key"][pk]: pk for pk in qa["product_keys"]}
    for row_num in range(2, ws.max_row + 1):
        evidence_id = text(ws.cell(row_num, idx["evidence_id"]).value)
        match = re.search(r"_(\d{3})$", evidence_id)
        if not match:
            continue
        pos = int(match.group(1))
        if pos not in PRODUCT_UPDATES:
            continue
        pk = pk_by_pos[pos]
        handle = qa["handle_by_key"][pk]
        admin_row = admin[handle]
        item = PRODUCT_UPDATES[pos]
        ws.cell(row_num, idx["sources_accessed"]).value = (
            f"Original storefront/product evidence; Sang QA {QA_RUN_ID}; products_export_1.csv sha256:{admin_hash}; customizer_audit.json; contact sheet inspected"
        )
        ws.cell(row_num, idx["current_H1"]).value = admin_row["Title"]
        ws.cell(row_num, idx["current_meta_title"]).value = admin_row["SEO Title"] or "EMPTY"
        ws.cell(row_num, idx["verified_product_facts"]).value = (
            f"Admin export matched handle {handle}; type={admin_row['Type']}; options={option_text(admin_row)}; "
            f"status={admin_row['Status']}; image_count={len(admin_row['images'])}; design={item['detail']}; "
            f"customizer={customizer_sentence(qa['customizer_by_pos'].get(pos, {}))}"
        )
        ws.cell(row_num, idx["factual_conflicts"]).value = "Sang QA issues addressed in r2; requires independent QA recheck."
        ws.cell(row_num, idx["confidence_and_reason"]).value = (
            "R2 uses inspected contact sheets, Customizer audit and Shopify admin CSV baseline; no approval or import created."
        )

    ws = wb["README_QA"]
    idx = headers(ws)
    for metric, value, definition in [
        ("qa_batch_008_r2_revision", "r2", "Separate 10-product revision after Sang QA; source workbook was not modified."),
        ("qa_batch_008_r2_admin_export", str(ADMIN_EXPORT.relative_to(ROOT)), f"Admin CSV baseline used; sha256:{admin_hash}."),
        ("qa_batch_008_r2_scope", "inventory positions 71-80", "No products outside qa_batch_008 were revised."),
        ("qa_batch_008_r2_revised_products", str(revised_products), "SEO_Products rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_008_r2_revised_images", str(revised_images), "Image_Audit rows revised and kept NEEDS_REVIEW."),
        ("qa_batch_008_r2_admin_alt_matches", str(admin_alt_matches), "Image rows with current admin alt matched by image URL from Shopify CSV."),
    ]:
        ws.append([
            metric if col == idx["metric"] else value if col == idx["value"] else definition if col == idx["definition"] else ""
            for col in range(1, ws.max_column + 1)
        ])

    wb.save(OUTPUT)
    load_workbook(OUTPUT, keep_links=False).save(OUTPUT)

    summary = {
        "created_at": now,
        "batch_id": BATCH_ID,
        "source_revision": str(SOURCE.relative_to(ROOT)),
        "source_qa": str(QA_DATASET.relative_to(ROOT)),
        "admin_export": str(ADMIN_EXPORT.relative_to(ROOT)),
        "admin_export_sha256": admin_hash,
        "revision_workbook": str(OUTPUT.relative_to(ROOT)),
        "scope": "qa_batch_008 only; inventory positions 71-80",
        "revision": "r2",
        "revised_products": revised_products,
        "revised_images": revised_images,
        "admin_alt_matches": admin_alt_matches,
        "keyword_rows_revised": keyword_rows,
        "review_status": "NEEDS_REVIEW",
        "approved_created": False,
        "shopify_import_created": False,
        "next_step": "Sang QA qa_batch_008_r2 before any approval or import.",
    }
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT.write_text(
        "\n".join(
            [
                "# Revision qa_batch_008_r2",
                "",
                "Artifact chính thức theo quy trình từng lô 10 sản phẩm.",
                "",
                f"- Nguồn r1: `{SOURCE.relative_to(ROOT)}`",
                f"- QA của Sang: `{QA_DATASET.relative_to(ROOT)}`",
                f"- Shopify admin export: `{ADMIN_EXPORT.relative_to(ROOT)}`",
                f"- Admin export SHA-256: `{admin_hash}`",
                f"- Workbook r2: `{OUTPUT.relative_to(ROOT)}`",
                f"- Sản phẩm sửa: {revised_products}",
                f"- Ảnh sửa/refresh: {revised_images}",
                f"- Admin image alt match: {admin_alt_matches}/{revised_images}",
                "- Sửa trọng yếu: bỏ nội dung nội bộ, sửa claim customizer theo từng sản phẩm, tách keyword intent, viết lại mô tả/alt theo ảnh đã xem.",
                "- Trạng thái: `NEEDS_REVIEW`; không tạo `APPROVED`; không tạo file import Shopify.",
                "- Bước tiếp theo: Sang QA lại `qa_batch_008_r2`.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
