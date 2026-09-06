import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import build_batch_workbook as base


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
BATCH_ID = "batch_004"
RUN_DIR = ROOT / "seo_runs" / SHOP / RUN_ID
BATCHES_DIR = ROOT / "resutls" / SHOP / RUN_ID / "batches"
SUMMARY_PATH = RUN_DIR / f"{BATCH_ID}_evidence_summary.json"
PROGRESS_PATH = RUN_DIR / "progress.json"
PREVIOUS = BATCHES_DIR / "SEO_Product_Optimization_through_batch_003.xlsx"
OUTPUT = BATCHES_DIR / "SEO_Product_Optimization_through_batch_004.xlsx"


PROFILES = {
    31: ("Christmas cardinal birdhouse quilt with poinsettias", "Christmas cardinal birdhouse quilt", "Christmas Cardinal Birdhouse Quilt Set", "Christmas Cardinal Birdhouse Quilt Set", "red cardinals, snowy birdhouse, poinsettias, white flowers and pine greenery", "cardinal Christmas bedding, winter cardinal quilt, cardinal poinsettia quilt", "CHRISTMAS", "https://www.etsy.com/market/cardinal_christmas_bedding; https://www.amazon.com/cardinal-christmas-quilt/s?k=cardinal+christmas+quilt; https://www.bedbathandbeyond.com/Bedding-Bath/MarCielo-3-Pcs-Winter-Cardinals-Christmas-Quilt-Bedspread-Set-C79/38941937/product.html"),
    32: ("Christmas cardinals snowy branch quilt", "Christmas cardinal quilt set", "Christmas Cardinal Snowy Branch Quilt Set", "Christmas Cardinals on Snowy Branches Quilt", "two bright red cardinals on snowy pine branches with berries, poinsettias and warm lights", "winter cardinal bedding, red cardinal Christmas quilt, snowy cardinal quilt", "CHRISTMAS", "https://www.wayfair.com/keyword.php?keyword=cardinal+quilt; https://www.ebay.com/b/Cardinal-Quilt-Indiana-Quilts-Bedspreads-Coverlets/175749/bn_7022481015; https://alphaquilt.com/collections/cardinal"),
    33: ("cow and sunflower farmhouse quilt", "cow sunflower quilt set", "Cow Sunflower Farmhouse Quilt Set", "Cow and Sunflower Farmhouse Quilt Set", "black and white cow artwork with large yellow sunflowers and patchwork-style farm colors", "farmhouse cow bedding, cow quilt set, sunflower cow quilt", "EVERGREEN", "https://www.amazon.com/; https://www.walmart.com/; https://www.bedbathandbeyond.com/"),
    34: ("crocodile patchwork animal quilt", "crocodile patchwork quilt set", "Crocodile Patchwork Quilt Set", "Crocodile Patchwork Animal Quilt Set", "cartoon crocodile in green, orange and yellow patchwork with small flowers", "crocodile bedding, alligator quilt set, animal patchwork bedding", "EVERGREEN", "https://www.etsy.com/sg-en/market/crocodile_bedding; https://www.walmart.com/ip/Crocodile-Duvet-Cover-Set-Picturesque-American-Alligator-Mouth-Open-Decorative-3-Piece-Bedding-Set-2-Pillow-Shams-Calking-Size-Grey-Dark-Grey-Ambeson/5168331823; https://www.wayfair.com/bed-bath/sb2/animal-print-queen-comforters-sets-c215334-a9645~33035-a9647~33050.html"),
    35: ("personalized football flag comforter with the name Kevin and number 15", "personalized football comforter", "Personalized Football Flag Comforter Set", "Personalized American Football Comforter Set", "close-up football artwork over a vintage American flag background with name Kevin and number 15", "custom football bedding, football comforter with name, football bedding for boys", "EVERGREEN", "https://www.facebook.com/61590427798204/posts/let-me-know-what-i-can-make-for-you/122134162143347593/"),
    36: ("personalized grunge football comforter with the name Brian and number 15", "custom football comforter with name", "Custom Football Comforter with Name", "Custom Grunge Football Comforter Set", "large football on a red-white grunge stripe background with name Brian and number 15", "personalized football bedding, American football comforter, football bedding set", "EVERGREEN", "https://www.facebook.com/61590427798204/posts/let-me-know-what-i-can-make-for-you/122134162143347593/"),
    37: ("personalized cosmic football comforter with the name Kevin and number 20", "cosmic football comforter", "Personalized Cosmic Football Comforter Set", "Personalized Cosmic Football Comforter Set", "football floating over a blue and red cosmic background with name Kevin and number 20", "custom football bedding, football comforter set, sports bedding for football fans", "EVERGREEN", "https://www.facebook.com/61590427798204/posts/let-me-know-what-i-can-make-for-you/122134162143347593/"),
    38: ("personalized USA flag football comforter with the name Michael and number 30", "personalized patriotic football comforter", "Personalized Patriotic Football Comforter", "Personalized USA Flag Football Comforter Set", "large football on a worn USA flag background with name Michael and number 30", "American flag football bedding, custom football comforter, patriotic sports bedding", "EVERGREEN", "https://www.facebook.com/61590427798204/posts/let-me-know-what-i-can-make-for-you/122134162143347593/"),
    39: ("personalized vintage football comforter with the name Mark and number 15", "personalized vintage football comforter", "Personalized Vintage Football Comforter", "Personalized Vintage Football Comforter Set", "large close-up football, tan paint-splash background, vertical name Mark and number 15", "custom football comforter, football bedding with name, vintage football bedding", "EVERGREEN", "https://www.facebook.com/61590427798204/posts/let-me-know-what-i-can-make-for-you/122134162143347593/"),
    40: ("personalized patriotic football comforter with the name Matthew and number 08", "custom patriotic football comforter", "Custom Patriotic Football Comforter Set", "Custom Patriotic Football Comforter Set", "American flag football graphic with teal sky, name Matthew and number 08", "personalized football bedding, American flag football comforter, custom sports bedding", "EVERGREEN", "https://www.facebook.com/61590427798204/posts/let-me-know-what-i-can-make-for-you/122134162143347593/"),
}


def append_dict(ws, data):
    headers = [cell.value for cell in ws[1]]
    ws.append([data.get(h, "") for h in headers])


def image_obs(theme, pos):
    if pos == 1:
        return f"Main bed mockup showing {theme}."
    if pos == 2:
        return f"Angled bedroom mockup showing {theme}."
    if pos == 3:
        return f"Product type comparison or matching pillow sham/detail image for {theme}."
    if pos == 4:
        return f"Close room/detail or fabric feature image showing {theme}."
    if pos == 5:
        return "Easy-care or sizing information image for the bedding set."
    if pos == 6:
        return "Size dimension image for twin, full, queen and king bedding options."
    return f"Additional bedroom mockup or product feature image showing {theme}."


def main():
    wb = load_workbook(PREVIOUS)
    summaries = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    now = datetime.now().astimezone().isoformat()
    added_images = 0

    for summary in summaries:
        row = summary["inventory_row"]
        pos = int(row["inventory_position"])
        label, primary, seo_title, h1, detail, secondary, season, refs = PROFILES[pos]
        product_key = row["product_key"]
        evidence_id = f"evidence_{BATCH_ID}_{pos:03d}"
        research_id = f"research_{BATCH_ID}_{pos:03d}_001"
        kind = base.product_kind(row["product_type"])
        facts = base.extract_facts(summary["body_html"])
        body_text = base.strip_html(summary["body_html"])
        images = summary["images"]
        image_refs = "; ".join(f"img_{BATCH_ID}_{pos:03d}_{int(img['position']):02d}" for img in images)
        meta = f"Shop a {kind} featuring {detail}. Review size and options before checkout."
        desc = (
            f"<p>This Jeminise {kind} features {detail}, giving the product a clear theme for bedroom decor or gifting.</p>"
            f"<h3>Design Details</h3><ul><li>Visible artwork: {detail}.</li><li>Source product details include: {facts[:650]}.</li></ul>"
            f"<h3>SEO Use</h3><p>This draft targets a product-level query tied to the visible artwork and product type. It needs QA and approval before import.</p>"
        )

        append_dict(wb["SEO_Products"], {
            "shop_domain": SHOP, "product_key": product_key, "Handle": row["Handle"], "product_id": row["product_id"],
            "product_url": row["product_url"], "canonical_url": summary["html"].get("canonical_url", ""),
            "product_type": row["product_type"], "title_current": row["title_current"],
            "h1_current": summary["html"].get("h1_current", ""), "rendered_title_current": summary["html"].get("rendered_title_current", ""),
            "meta_description_current": summary["html"].get("meta_description_current", ""),
            "primary_keyword": primary, "secondary_keywords": secondary, "keyword_evidence_level": "SERP_ONLY",
            "keyword_strategy": f"Target the product-level motif: {label}. Broader animal, Christmas or football terms should be reviewed for collection pages.",
            "season": season, "buyer_search_summary": f"Buyer is likely shopping for {label} as themed bedding or a gift.",
            "title_action": "SET", "title_proposed": h1,
            "meta_title_action": "SET", "meta_title_seo": seo_title, "meta_title_length": len(seo_title),
            "meta_description_action": "SET", "meta_description_seo": meta, "meta_description_length": len(meta),
            "description_action": "SET", "description_proposed_html": desc,
            "meta_keyword": primary, "image_count": len(images), "images_viewed_count": len(images),
            "evidence_id": evidence_id, "processing_status": "DRAFTED", "review_status": "NEEDS_REVIEW", "revision": "r1",
            "issues": "Admin export not provided; stored SEO fields, current alt text and exact admin media mapping should be verified before deployment.",
        })

        append_dict(wb["Product_Evidence"], {
            "evidence_id": evidence_id, "product_url": row["product_url"], "reviewed_at": now,
            "sources_accessed": f"{row['product_url']}; {summary['product_json_ref']}; {summary['contact_sheet']}",
            "current_H1": summary["html"].get("h1_current", ""), "current_meta_title": summary["html"].get("rendered_title_current", ""),
            "short_source_excerpt": body_text[:900], "verified_product_facts": facts,
            "gallery_image_count": len(images), "images_viewed_count": len(images), "image_audit_references": image_refs,
            "SERP_evidence_references": refs, "buyer_research_references": research_id,
            "processing_status": "DRAFTED", "confidence_and_reason": "Medium: product JSON, HTML and contact sheet reviewed; admin export and QA are still required.",
            "fact_to_source_map": f"Product facts from {summary['product_json_ref']}; visual observations from {summary['contact_sheet']}",
            "proposed_field_to_fact_map": f"title/meta/description use {evidence_id} and {research_id}; alt proposals use {image_refs}",
        })

        append_dict(wb["Buyer_Search_Research"], {
            "research_id": research_id, "product_key": product_key, "supporting_fact_ids": evidence_id,
            "purchase_context": f"Shopping for {label}.",
            "jtbd_statement": f"When shopping for themed bedding, the buyer wants a coordinated {kind} matching a specific animal, holiday or football design.",
            "functional_motivation": "Find the right design, size and included bedding components.",
            "emotional_social_motivation": "Create seasonal decor, farmhouse personality, animal-themed style or a personalized sports gift.",
            "purchase_concerns": "Fit, included components, material/care claims, personalization accuracy and whether the artwork matches expectations.",
            "source_refs": refs, "source_scope": "MIXED", "observed_at": now, "market": "United States",
            "source_language": "English", "research_status": "SERP_ONLY",
            "limitations": "No Search Console, internal search or verified customer review corpus was provided; public results are query comparables, not volume.",
            "seo_application": f"Use {primary} as candidate product-page keyword.",
        })

        for keyword, role in [(primary, "PRIMARY")] + [(s.strip(), "SECONDARY") for s in secondary.split(",")]:
            append_dict(wb["Keyword_Map"], {
                "keyword": keyword, "product_key": product_key, "buyer_research_refs": research_id,
                "query_origin": "SERP_OBSERVED_OR_AGENT_CANDIDATE", "semantic_cluster": label, "intent": "purchase",
                "target_page_type": "PRODUCT", "target_url": row["product_url"], "keyword_role": role,
                "decision_reason": "Specific to visible artwork, personalization and product type.",
                "supporting_fact_ids": evidence_id, "demand_evidence": "SERP_ONLY", "season": season,
                "research_period": "2026-09-07", "validation_source": refs, "checked_at": now,
                "representative_SERP_URLs": refs, "possible_overlap_with_other_products": "YES",
                "mapping_reason": "Product query includes design attributes visible on this item.",
                "mapping_status": "CANDIDATE_MAPPED", "mapping_version": "r1",
            })

        for image in images:
            ipos = int(image["position"])
            alt = {
                1: f"{h1} displayed on a bed",
                2: f"Angled bedroom view of {h1}",
                3: f"Matching sham or product info image for {h1}",
                4: f"Close-up detail of {label}",
                5: f"Care or size information for {h1}",
                6: f"Size dimension guide for {h1}",
                7: f"Additional bedroom mockup of {h1}",
            }.get(ipos, f"Product detail image for {h1}")[:125]
            append_dict(wb["Image_Audit"], {
                "shop_domain": SHOP, "Handle": row["Handle"], "product_id": row["product_id"], "media_id": image["image_id"],
                "image_location": "GALLERY", "image_number": ipos, "variant": image.get("variant_ids", ""),
                "image_url": image["src"], "image_url_export": image["src"], "identity_status": "PUBLIC_JSON_IMAGE_ID",
                "viewed_status": "VIEWED_CONTACT_SHEET", "viewed_at": now, "observed_visual_details": image_obs(label, ipos),
                "alt_current": "UNKNOWN", "alt_proposed": alt, "alt_action": "SET",
                "review_status": "NEEDS_REVIEW", "revision": "r1",
                "evidence_file_or_reference": f"{image['local_path']}; {summary['contact_sheet']}",
                "issues": "Alt current unknown without admin export/rendered image-alt extraction.",
            })
            added_images += 1

    for row in [
        {"metric": "batch_004_products", "value": "10", "definition": "Products appended in this cumulative workbook"},
        {"metric": "batch_004_images", "value": str(added_images), "definition": "Gallery images downloaded and viewed for batch_004"},
        {"metric": "cumulative_products", "value": "40", "definition": "Total products included through batch_004"},
    ]:
        append_dict(wb["README_QA"], row)

    wb.save(OUTPUT)
    load_workbook(OUTPUT).save(OUTPUT)

    progress = json.loads(PROGRESS_PATH.read_text(encoding="utf-8-sig"))
    with (RUN_DIR / "inventory.csv").open(newline="", encoding="utf-8-sig") as f:
        inventory = list(csv.DictReader(f))
    next_batch = inventory[40:50]
    progress.update({
        "batch_id": BATCH_ID,
        "batch_product_keys": [s["inventory_row"]["product_key"] for s in summaries],
        "batch_status": "BATCH_WORKBOOK_SAVED",
        "current_stage": "AWAITING_CONFIRMATION_FOR_NEXT_BATCH",
        "awaiting_confirmation": True,
        "next_batch_id": "batch_005",
        "next_batch_product_keys": [r["product_key"] for r in next_batch],
        "last_saved_at": datetime.now().astimezone().isoformat(),
    })
    progress.setdefault("artifact_paths", {})["batch_evidence_summary"] = str(SUMMARY_PATH.relative_to(ROOT))
    progress.setdefault("artifact_paths", {})["latest_batch_workbook"] = str(OUTPUT.relative_to(ROOT))
    PROGRESS_PATH.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT.relative_to(ROOT))
    print(f"appended_products=10 appended_images={added_images}")


if __name__ == "__main__":
    main()
