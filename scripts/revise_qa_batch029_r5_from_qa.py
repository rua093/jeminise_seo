import hashlib
from datetime import datetime
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SHOP = "jeminise.com"
RUN_ID = "20260906_234129"
SRC = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_029_r4" / "SEO_Product_Optimization_qa_batch_029_r4.xlsx"
OUT_DIR = ROOT / "resutls" / SHOP / RUN_ID / "revisions" / "qa_batch_029_r5"
OUT_XLSX = OUT_DIR / "SEO_Product_Optimization_qa_batch_029_r5.xlsx"
OUT_MD = OUT_DIR / "SEO_Product_Optimization_qa_batch_029_r5.md"

HANDLE = "personalized-hunting-and-outdoor-comforter-with-deer-antlers-4e13e688e0-4e13e688e0"
PRODUCT_KEY = f"{SHOP}+{HANDLE}"
EVIDENCE_ID = "evidence_batch_029_286"

SEO_FIX = {
    "primary_keyword": "hunting deer antlers quilt set",
    "secondary_keywords": "deer antlers quilt; hunting bedding set; rustic deer photo quilt",
    "keyword_strategy": "Target the specific product motif: hunting deer antlers quilt set; keep broader hunting bedding terms for collection pages.",
    "buyer_search_summary": "US English buyer intent targets hunting deer antlers quilt set; no search-volume, trend or ranking claim is made.",
    "title_proposed": "Hunting Deer Antlers Quilt Set",
    "meta_title_seo": "Hunting Deer Antlers Quilt Set",
    "meta_description_seo": "Shop hunting deer antlers quilt set with brown stripes, deer head silhouette, antlers, heart icon and sample names.",
    "description_proposed_html": """
<p>Brown stripes, antlers and a deer head silhouette give this quilt set a hunting-inspired look. The heart icon, photo collage artwork and sample names add a personalized couple-style layer to the design.</p>
<h3>Design Details</h3>
<ul>
<li>Brown striped quilt artwork with deer head silhouette and antlers.</li>
<li>Heart icon, photo collage artwork and sample names are visible.</li>
<li>Gallery includes the quilt mockup plus size, care, fabric or lifestyle panels where shown.</li>
</ul>
<h3>Personalization and Options</h3>
<ul>
<li>Choose Quilt Size and Pillowcase Quantity options are listed.</li>
<li>Customize Your Quilt is optional and accepts 1-1000 characters.</li>
<li>Visible custom text or example selections are samples unless matching input is entered.</li>
</ul>
""",
    "meta_keyword": "hunting deer antlers quilt set",
}

IMAGE_ALT = {
    1: ("Bed mockup of hunting deer antlers quilt with photo collage artwork and sample names.", "Hunting deer antlers quilt bed view"),
    2: ("Held quilt showing deer head silhouette, antlers, heart icon and photo collage artwork.", "Hunting deer antlers quilt held view"),
    3: ("Room mockup of hunting deer antlers quilt with printed craft callout.", "Hunting deer antlers quilt room mockup"),
    4: ("Close bed view with optional pillow shams and deer antlers photo artwork.", "Hunting deer antlers quilt pillow sham panel"),
    5: ("Fabric panel with white backing and deer antlers quilt pillow mockup.", "Hunting deer antlers quilt fabric panel"),
    6: ("Sizing and details chart for hunting deer antlers quilt set.", "Hunting deer antlers quilt size chart"),
    7: ("Bedspread features panel showing quilt layers and care icons.", "Hunting deer antlers quilt features panel"),
    8: ("Overhead bedroom mockup of hunting deer antlers quilt and matching pillows.", "Hunting deer antlers quilt overhead view"),
}

KEYWORD_ROWS = {
    "hunting deer antlers comforter set": "hunting deer antlers quilt set",
    "deer antlers comforter": "deer antlers quilt",
    "hunting bedding set": "hunting bedding set",
    "rustic photo comforter": "rustic deer photo quilt",
}


def headers(ws):
    return {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def clean_html(value):
    return str(BeautifulSoup(value.strip(), "html.parser"))


def replace_public_terms(value):
    if value is None:
        return value
    out = str(value)
    for old, new in [
        ("Hunting Deer Antlers Comforter Set", "Hunting Deer Antlers Quilt Set"),
        ("hunting deer antlers comforter set", "hunting deer antlers quilt set"),
        ("deer antlers comforter", "deer antlers quilt"),
        ("rustic photo comforter", "rustic deer photo quilt"),
        ("comforter set", "quilt set"),
        ("Comforter Set", "Quilt Set"),
        ("comforter", "quilt"),
        ("Comforter", "Quilt"),
    ]:
        out = out.replace(old, new)
    return out


def main():
    wb = load_workbook(SRC)

    ws = wb["SEO_Products"]
    h = headers(ws)
    changed_products = 0
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, h["Handle"]).value != HANDLE:
            continue
        for col, value in SEO_FIX.items():
            ws.cell(row, h[col]).value = clean_html(value) if col == "description_proposed_html" else value
        ws.cell(row, h["meta_title_length"]).value = len(SEO_FIX["meta_title_seo"])
        ws.cell(row, h["meta_description_length"]).value = len(SEO_FIX["meta_description_seo"])
        ws.cell(row, h["revision"]).value = "r5"
        ws.cell(row, h["issues"]).value = "Resolved QA r4 MAJOR: proposed public SEO fields now use Quilt/Quilt Set consistently for product 286; current handle/title kept as source data."
        changed_products += 1

    ws = wb["Image_Audit"]
    h = headers(ws)
    changed_images = 0
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, h["Handle"]).value != HANDLE:
            continue
        image_number = int(ws.cell(row, h["image_number"]).value)
        obs, alt = IMAGE_ALT[image_number]
        ws.cell(row, h["observed_visual_details"]).value = obs
        ws.cell(row, h["alt_proposed"]).value = alt
        ws.cell(row, h["revision"]).value = "r5"
        ws.cell(row, h["issues"]).value = "Resolved QA r4 MAJOR: image observation and alt use Quilt wording supported by admin/live/contact-sheet evidence."
        changed_images += 1

    ws = wb["Product_Evidence"]
    h = headers(ws)
    changed_evidence = 0
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, h["evidence_id"]).value != EVIDENCE_ID:
            continue
        ws.cell(row, h["verified_product_facts"]).value = (
            f"Admin export matched handle {HANDLE}; type=Quilt; options=Choose Quilt Size, "
            "Pillowcase Quantity; status=active; image_count=8; design=brown striped quilt "
            "with deer head silhouette, antlers, heart icon, photo collage artwork and sample "
            "names; customizer=Personalization uses an optional Customize Your Quilt field, "
            "1-1000 characters. Visible custom text or example selections in mockups are sample "
            "artwork unless the matching input is entered."
        )
        ws.cell(row, h["factual_conflicts"]).value = "QA r4 found source conflict: live title/handle used Comforter while admin Type, live JSON type, size options, live description and contact-sheet panels support Quilt/Quilt Set. R5 keeps current source fields intact but standardizes proposed SEO fields to Quilt/Quilt Set."
        for col in ["confidence_and_reason", "proposed_field_to_fact_map"]:
            ws.cell(row, h[col]).value = replace_public_terms(ws.cell(row, h[col]).value)
        changed_evidence += 1

    ws = wb["Keyword_Map"]
    h = headers(ws)
    changed_keywords = 0
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, h["product_key"]).value != PRODUCT_KEY:
            continue
        kw = ws.cell(row, h["keyword"]).value
        ws.cell(row, h["keyword"]).value = KEYWORD_ROWS.get(kw, replace_public_terms(kw))
        for col in ["semantic_cluster", "decision_reason", "supporting_fact_ids", "demand_evidence", "mapping_reason"]:
            ws.cell(row, h[col]).value = replace_public_terms(ws.cell(row, h[col]).value)
        ws.cell(row, h["mapping_version"]).value = "r5"
        changed_keywords += 1

    ws = wb["Buyer_Search_Research"]
    h = headers(ws)
    changed_research = 0
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, h["product_key"]).value != PRODUCT_KEY:
            continue
        for col in ["purchase_context", "jtbd_statement", "functional_motivation", "customer_language", "limitations", "seo_application"]:
            ws.cell(row, h[col]).value = replace_public_terms(ws.cell(row, h[col]).value)
        changed_research += 1

    ws = wb["README_QA"]
    ws.append([
        "revision_note_qa_batch_029_r5",
        datetime.now().astimezone().isoformat(timespec="seconds"),
        "R5 revises only product 286 after QA r4: admin/live/contact-sheet evidence supports Quilt/Quilt Set, so proposed SEO fields and 8 image alts were changed from comforter wording to quilt wording. No APPROVED status was created.",
    ])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    OUT_MD.write_text(
        f"""# SEO Product Optimization Revision - qa_batch_029_r5

Scope: batch 029, products 281-290.

Source revision: `qa_batch_029_r4`.

QA source: `resutls/jeminise.com/20260906_234129/qa/20260909_145706/SEO_QA_qa_batch_029_r4.md`.

Changes made:
- Revised only product `286`: `{HANDLE}`.
- Confirmed proposed public SEO product type should be `Quilt`/`Quilt Set`, because admin `Type`, live JSON type, size options, live description and contact-sheet info panels support quilt terminology.
- Kept current source fields such as handle/current title unchanged because they represent current store data.
- Updated `title_proposed`, `meta_title_seo`, `meta_description_seo`, `description_proposed_html`, `primary_keyword`, `secondary_keywords`, `keyword_strategy`, `buyer_search_summary` and `meta_keyword`.
- Updated all 8 related `Image_Audit` rows so observations and alt text use `quilt` instead of `comforter`.
- Updated related `Keyword_Map`, `Product_Evidence` and `Buyer_Search_Research` wording for consistency.
- No `APPROVED` status was created.

Verification:
- Output workbook SHA-256: `{sha256(OUT_XLSX)}`.
- Products changed: `{changed_products}`.
- Image rows changed: `{changed_images}`.
- Keyword rows changed: `{changed_keywords}`.
- Evidence rows changed: `{changed_evidence}`.
- Buyer research rows changed: `{changed_research}`.
""",
        encoding="utf-8-sig",
    )
    print({
        "output": str(OUT_XLSX),
        "summary": str(OUT_MD),
        "sha256": sha256(OUT_XLSX),
        "products_changed": changed_products,
        "images_changed": changed_images,
        "keywords_changed": changed_keywords,
        "evidence_rows_changed": changed_evidence,
        "buyer_research_rows_changed": changed_research,
    })


if __name__ == "__main__":
    main()
