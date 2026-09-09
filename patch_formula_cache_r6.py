import os, zipfile, tempfile, xml.etree.ElementTree as ET, re
from openpyxl import load_workbook

p=r'resutls\jeminise.com\20260906_234129\qa\20260909_120000\SEO_QA_qa_batch_003_r6.xlsx'
wb=load_workbook(p,data_only=False)
crit=wb['QA_Criteria']; imgs=wb['QA_Images']; prods=wb['QA_Products']
scores={}; crit_cache={}
for r in range(2,112):
    key=crit.cell(r,1).value; rating=crit.cell(r,4).value; weight=crit.cell(r,3).value
    e=1 if rating in ('FULL','DERIVED') else .5 if rating=='PARTIAL' else 0
    f=weight*e; g=20 if rating=='DERIVED' else weight
    crit_cache[r]=(e,f,g); scores[key]=scores.get(key,0)+f
score_by_pos={int(prods.cell(r,1).value):round(scores[prods.cell(r,2).value],1) for r in range(2,12)}
status_by_pos={21:'QA_PASS',22:'QA_FAIL',23:'QA_FAIL',24:'QA_FAIL',25:'QA_FAIL',26:'QA_FAIL',27:'QA_FAIL',28:'QA_FAIL',29:'QA_REVISE',30:'QA_REVISE'}

ns={'main':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
ET.register_namespace('',ns['main'])
def set_v(root,ref,val):
    cell=None
    for c in root.findall('.//main:c',ns):
        if c.attrib.get('r')==ref: cell=c; break
    if cell is None: return
    v=cell.find('main:v',ns)
    if v is None: v=ET.SubElement(cell,'{%s}v'%ns['main'])
    v.text=str(val)
    cell.attrib['t']='n' if isinstance(val,(int,float)) and not isinstance(val,bool) else 'b' if isinstance(val,bool) else 'str'
    if isinstance(val,bool): v.text='1' if val else '0'

with tempfile.TemporaryDirectory(dir=os.path.dirname(p)) as td:
    out=os.path.join(td,'patched.xlsx')
    with zipfile.ZipFile(p,'r') as zin, zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data=zin.read(info.filename)
            if info.filename.startswith('xl/worksheets/sheet') and info.filename.endswith('.xml'):
                root=ET.fromstring(data)
                # Sheet order is QA_Summary, QA_Products, QA_Criteria, QA_Images, QA_Issues.
                sheet_no=int(re.search(r'sheet(\d+)\.xml$',info.filename).group(1))
                if sheet_no==1:
                    set_v(root,'B13',73.25); set_v(root,'B14','NOT_PASSED'); set_v(root,'B15','PASS=1; REVISE=2; FAIL=7; INCOMPLETE=0')
                elif sheet_no==2:
                    for r in range(2,12):
                        pos=int(prods.cell(r,1).value); score=score_by_pos[pos]
                        for ref,val in [(f'G{r}',score),(f'H{r}',100),(f'I{r}',score),(f'J{r}',score),(f'K{r}',score),(f'L{r}',status_by_pos[pos]),(f'P{r}',True),(f'Q{r}',1),(f'R{r}',0 if pos==21 else 0),(f'S{r}',0 if pos==21 else 1),(f'T{r}',1 if 22<=pos<=28 else 0),(f'U{r}',0)]: set_v(root,ref,val)
                elif sheet_no==3:
                    for r,(e,f,g) in crit_cache.items(): set_v(root,f'E{r}',e); set_v(root,f'F{r}',f); set_v(root,f'G{r}',g)
                elif sheet_no==4:
                    for r in range(2,52):
                        for ref,val in [(f'T{r}',100),(f'U{r}',100),(f'V{r}',100),(f'W{r}',100),(f'X{r}',100)]: set_v(root,ref,val)
                data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            zout.writestr(info,data)
    with open(p,'wb') as dst, open(out,'rb') as srcf:
        dst.write(srcf.read())
print('patched formula caches',p)
