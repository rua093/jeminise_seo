import re
import shutil
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl


XLSX = Path("resutls/jeminise.com/20260906_234129/qa/20260909_154500/SEO_QA_qa_batch_004_r5.xlsx")
NS = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
ET.register_namespace("", NS["main"])


def cell_ref_to_tuple(ref):
    m = re.match(r"([A-Z]+)(\d+)", ref)
    col = 0
    for ch in m.group(1):
        col = col * 26 + ord(ch) - 64
    return int(m.group(2)), col


def set_cache(cell, value):
    for child in list(cell):
        if child.tag.split("}")[-1] == "v":
            cell.remove(child)
    formula = None
    for child in list(cell):
        if child.tag.split("}")[-1] == "f":
            formula = child
            break
    v_tag = "v"
    if formula is not None and "}" in formula.tag:
        v_tag = formula.tag.rsplit("}", 1)[0] + "}v"
    v = ET.SubElement(cell, v_tag)
    if isinstance(value, str):
        cell.set("t", "str")
        v.text = value
    elif isinstance(value, bool):
        cell.set("t", "b")
        v.text = "1" if value else "0"
    elif value is None:
        v.text = ""
    else:
        if "t" in cell.attrib:
            del cell.attrib["t"]
        v.text = str(round(value, 10)).rstrip("0").rstrip(".")


def value_text(value):
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, str):
        return value
    if value is None:
        return ""
    return str(round(float(value), 10)).rstrip("0").rstrip(".")


def patch_cache_text(xml, ref, value):
    val = value_text(value)
    pattern = re.compile(r'(<c r="' + re.escape(ref) + r'"[^>]*>.*?<v>)(.*?)(</v>)', re.S)
    return pattern.sub(lambda m: m.group(1) + val + m.group(3), xml, count=1)


def sheet_formula_cells(sheet_xml):
    tree = ET.parse(sheet_xml)
    root = tree.getroot()
    cells = {}
    for c in root.iter():
        if c.tag.split("}")[-1] != "c":
            continue
        if any(child.tag.split("}")[-1] == "f" for child in list(c)):
            cells[c.attrib["r"]] = c
    return tree, cells


def sheet_all_cells(sheet_xml):
    tree = ET.parse(sheet_xml)
    root = tree.getroot()
    cells = {}
    for c in root.iter():
        if c.tag.split("}")[-1] == "c" and "r" in c.attrib:
            cells[c.attrib["r"]] = c
    return tree, cells


def sheet_paths(tmp):
    wb_tree = ET.parse(tmp / "xl/workbook.xml")
    rel_tree = ET.parse(tmp / "xl/_rels/workbook.xml.rels")
    rel_ns = {"rel": "http://schemas.openxmlformats.org/package/2006/relationships"}
    rid_to_target = {}
    for rel in rel_tree.getroot().findall("rel:Relationship", rel_ns):
        rid_to_target[rel.attrib["Id"]] = rel.attrib["Target"]
    paths = {}
    for sh in wb_tree.getroot().findall(".//main:sheet", NS):
        name = sh.attrib["name"]
        rid = sh.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
        target = rid_to_target[rid].lstrip("/")
        if not target.startswith("xl/"):
            target = "xl/" + target
        paths[name] = target
    return paths


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=False)
    summary = {
        "B10": 10,
        "B11": 110,
        "B12": 62,
        "B16": 87.25,
        "B17": "NEEDS_REVISION",
        "B18": "PASS=3; REVISE=7; FAIL=0; INCOMPLETE=0",
    }
    product_scores = {
        31: (100, 95, 95, 95, 95, 0, 0, "QA_PASS"),
        32: (100, 95, 95, 95, 95, 0, 0, "QA_PASS"),
        33: (100, 95, 95, 95, 95, 0, 0, "QA_PASS"),
        34: (100, 77.5, 77.5, 77.5, 77.5, 0, 1, "QA_REVISE"),
        35: (100, 85, 85, 85, 85, 0, 1, "QA_REVISE"),
        36: (100, 85, 85, 85, 85, 0, 1, "QA_REVISE"),
        37: (100, 85, 85, 85, 85, 0, 1, "QA_REVISE"),
        38: (100, 85, 85, 85, 85, 0, 1, "QA_REVISE"),
        39: (100, 85, 85, 85, 85, 0, 1, "QA_REVISE"),
        40: (100, 85, 85, 85, 85, 0, 1, "QA_REVISE"),
    }

    caches_by_name = {
        "QA_Summary": summary,
        "QA_Products": {},
        "QA_Criteria": {},
        "QA_Images": {},
    }

    # QA_Products formula caches.
    for row in range(2, 12):
        pos = wb["QA_Products"][f"B{row}"].value
        assessed, verified, floor, ceiling, final, crit, major, status = product_scores[pos]
        caches_by_name["QA_Products"].update({
            f"O{row}": 1,
            f"P{row}": assessed,
            f"Q{row}": verified,
            f"R{row}": floor,
            f"S{row}": ceiling,
            f"T{row}": final,
            f"U{row}": crit,
            f"V{row}": major,
            f"W{row}": status,
        })

    # QA_Criteria formula caches.
    ws = wb["QA_Criteria"]
    for row in range(2, ws.max_row + 1):
        weight = float(ws[f"F{row}"].value)
        rating = ws[f"G{row}"].value
        if rating == "DERIVED":
            rating_value = 1
        elif rating == "FULL":
            rating_value = 1
        elif rating == "PARTIAL":
            rating_value = 0.5
        elif rating == "FAIL":
            rating_value = 0
        else:
            rating_value = 0
        caches_by_name["QA_Criteria"].update({
            f"H{row}": rating_value,
            f"I{row}": weight * rating_value,
            f"J{row}": weight if rating != "NOT_CHECKED" else 0,
        })

    # QA_Images formula caches.
    for row in range(2, wb["QA_Images"].max_row + 1):
        caches_by_name["QA_Images"].update({
            f"W{row}": 100,
            f"X{row}": 100,
            f"Y{row}": 1,
        })

    with tempfile.TemporaryDirectory(dir=str(XLSX.parent)) as tmp:
        tmp = Path(tmp)
        with zipfile.ZipFile(XLSX, "r") as zin:
            zin.extractall(tmp)
        path_by_name = sheet_paths(tmp)
        for sheet_name, vals in caches_by_name.items():
            rel = path_by_name[sheet_name]
            tree, cells = sheet_all_cells(tmp / rel)
            for ref, val in vals.items():
                if ref in cells:
                    set_cache(cells[ref], val)
            tree.write(tmp / rel, encoding="utf-8", xml_declaration=True)
            xml = (tmp / rel).read_text(encoding="utf-8")
            for ref, val in vals.items():
                xml = patch_cache_text(xml, ref, val)
            (tmp / rel).write_text(xml, encoding="utf-8")
        patched = tmp / "patched.xlsx"
        with zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as zout:
            for path in tmp.rglob("*"):
                if path.is_file() and path.name != "patched.xlsx":
                    zout.write(path, path.relative_to(tmp).as_posix())
        shutil.copyfile(patched, XLSX)
    print("patched", XLSX)


if __name__ == "__main__":
    main()
